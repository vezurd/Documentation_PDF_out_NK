"""Qt-free styled table export to xlsx.

Paint comes from :class:`ExportedCell` (typically ``MonitorCell`` via
:func:`cell_from_monitor`). QTableWidget snapshots live in
``rd_catalog.table_xlsx_qt`` so this module stays Qt-free.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from rd_catalog.kits import kit_identity_key
from rd_catalog.kits_table_layout import canonical_kits_header_name
from rd_catalog.monitor_views import MonitorCell, contrast_foreground

KITS_XLSX_BUTTON = "Сохранить Excel"
KITS_XLSX_SHEET = "Комплекты"
KITS_XLSX_DEFAULT_WIDTH_PX = 80
KITS_XLSX_MAX_KEYS = 5000

_SHEET_TITLE_INVALID = frozenset("[]:*?/\\")
_DEFAULT_SHEET_TITLE = "Лист"
_HEADER_FILL_RGB = "E8EAED"
_HEADER_FG_RGB = "202124"
_MAX_SHEET_TITLE = 31


@dataclass(frozen=True, slots=True)
class ExportedCell:
    """One exported table cell (text plus optional paint)."""

    text: str
    fill_hex: str | None = None  # #RRGGBB or None → no PatternFill
    fg_hex: str | None = None
    bold: bool = False
    underline: bool = False


@dataclass(frozen=True, slots=True)
class ExportedColumn:
    """One visual column in export order (left to right)."""

    header: str
    width_px: int
    logical_index: int = 0


@dataclass(frozen=True, slots=True)
class ExportedTable:
    """A complete table snapshot ready for :func:`write_exported_table_xlsx`."""

    columns: tuple[ExportedColumn, ...]
    rows: tuple[tuple[ExportedCell, ...], ...]
    sheet_name: str = "Лист"


def normalize_hex(color: str | None) -> str | None:
    """Return ``#RRGGBB`` (uppercase hex) or ``None`` if ``color`` is invalid.

    Args:
        color: ``#RRGGBB`` / ``RRGGBB``, or ``None``.

    Returns:
        Normalized ``#RRGGBB``, or ``None`` when missing or not six hex digits.
    """

    if color is None:
        return None
    text = str(color).strip().lstrip("#")
    if len(text) != 6:
        return None
    try:
        int(text, 16)
    except ValueError:
        return None
    return f"#{text.upper()}"


def excel_rgb(color: str | None) -> str | None:
    """Return ``RRGGBB`` (no ``#``) for openpyxl, or ``None``.

    Args:
        color: ``#RRGGBB`` / ``RRGGBB``, or ``None``.

    Returns:
        Six hex digits, or ``None`` when :func:`normalize_hex` rejects ``color``.
    """

    normalized = normalize_hex(color)
    if normalized is None:
        return None
    return normalized[1:]


def px_to_excel_width(px: int | float) -> float:
    """Convert Qt pixel width to Excel Calibri-11 character width.

    Args:
        px: Header ``sectionSize`` in pixels.

    Returns:
        Width clamped to Excel's ``1..255`` range: ``(px - 5) / 7``.
    """

    return max(1.0, min(255.0, (float(px) - 5.0) / 7.0))


def dated_xlsx_filename(stem: str, when: datetime | None = None) -> str:
    """Return ``{stem}_YYYY.MM.DD.xlsx`` using local time.

    Args:
        stem: Filename prefix (kits: ``Комплекты``).
        when: Timestamp. ``None`` uses ``datetime.now()``.

    Returns:
        Filename without a directory.
    """

    stamp = (when or datetime.now()).strftime("%Y.%m.%d")
    return f"{stem}_{stamp}.xlsx"


def cell_from_monitor(cell: MonitorCell) -> ExportedCell:
    """Copy paint from a ``MonitorCell`` (same contrast rule as Qt bind).

    Args:
        cell: Join-layer cell. Tooltip is not exported.

    Returns:
        Frozen export cell. Missing fill/fg become ``None``.
    """

    fill = cell.fill
    fg = cell.foreground
    if fill and not fg:
        fg = contrast_foreground(fill)
    return ExportedCell(
        text=cell.text or "",
        fill_hex=normalize_hex(fill),
        fg_hex=normalize_hex(fg),
        bold=cell.bold,
        underline=cell.underline,
    )


def columns_from_layout(
    order: Sequence[str],
    widths: Mapping[str, int] | None = None,
    *,
    default_px: int = KITS_XLSX_DEFAULT_WIDTH_PX,
) -> tuple[ExportedColumn, ...]:
    """Build visual columns from a saved Комплекты layout.

    Args:
        order: Header titles left to right.
        widths: Header → pixel width. Missing or non-positive use *default_px*.
        default_px: Fallback Excel-source pixel width.

    Returns:
        Columns in *order*, ``logical_index`` matching that sequence.
    """

    mapping = widths or {}
    columns: list[ExportedColumn] = []
    for index, name in enumerate(order):
        header = str(name)
        raw = mapping.get(header, 0)
        try:
            px = int(raw)
        except (TypeError, ValueError):
            px = 0
        columns.append(
            ExportedColumn(
                header=header,
                width_px=px if px > 0 else int(default_px),
                logical_index=index,
            )
        )
    return tuple(columns)


def columns_from_client(
    raw: Sequence[Mapping[str, Any]] | None,
    *,
    known_names: Sequence[str],
    fallback_order: Sequence[str],
    fallback_widths: Mapping[str, int] | None = None,
) -> tuple[ExportedColumn, ...]:
    """Prefer the client's visual columns; otherwise the saved layout.

    Args:
        raw: ``{header, width_px}`` items from Tabulator, or ``None``.
        known_names: Current Комплекты header titles.
        fallback_order: Layout order when *raw* is empty.
        fallback_widths: Layout widths when *raw* is empty.

    Returns:
        Visible columns only when *raw* is usable; otherwise the fallback.
    """

    known = {str(name) for name in known_names}
    parsed: list[ExportedColumn] = []
    seen: set[str] = set()
    for item in raw or ():
        if not isinstance(item, Mapping):
            continue
        header = canonical_kits_header_name(str(item.get("header") or "").strip())
        if header not in known or header in seen:
            continue
        seen.add(header)
        try:
            px = int(item.get("width_px") or 0)
        except (TypeError, ValueError):
            px = 0
        parsed.append(
            ExportedColumn(
                header=header,
                width_px=px if px > 0 else KITS_XLSX_DEFAULT_WIDTH_PX,
                logical_index=len(parsed),
            )
        )
    if parsed:
        return tuple(parsed)
    return columns_from_layout(fallback_order, fallback_widths)


def exported_table_from_kit_cells(
    cells_by_key: Mapping[tuple[str, str], Mapping[str, MonitorCell]],
    keys: Sequence[tuple[str, str]],
    columns: Sequence[ExportedColumn],
    *,
    sheet_name: str = KITS_XLSX_SHEET,
) -> ExportedTable:
    """Assemble an export table from painted kit cells in *keys* order.

    Unknown or duplicate identities are skipped. Paint still comes from
    ``MonitorCell``, never from the client.

    Args:
        cells_by_key: ``kit_identity_key`` → header → cell.
        keys: Requested ``(title, mark)`` in screen order.
        columns: Visual columns left to right.
        sheet_name: Workbook sheet title.

    Returns:
        Frozen table ready for :func:`write_exported_table_xlsx`.
    """

    exported_columns = tuple(columns)
    rows: list[tuple[ExportedCell, ...]] = []
    seen: set[tuple[str, str]] = set()
    for title, mark in keys:
        key = kit_identity_key(str(title), str(mark))
        if not key[0] or not key[1] or key in seen:
            continue
        seen.add(key)
        cells = cells_by_key.get(key)
        if cells is None:
            continue
        rows.append(
            tuple(
                cell_from_monitor(
                    cells.get(column.header, MonitorCell(text="—"))
                )
                for column in exported_columns
            )
        )
    return ExportedTable(
        columns=exported_columns,
        rows=tuple(rows),
        sheet_name=sheet_name,
    )


def _sheet_title(name: str) -> str:
    cleaned = "".join(" " if ch in _SHEET_TITLE_INVALID else ch for ch in name)
    cleaned = cleaned.strip()[:_MAX_SHEET_TITLE].strip()
    return cleaned or _DEFAULT_SHEET_TITLE


def _freeze_anchor(columns: tuple[ExportedColumn, ...]) -> str:
    if (
        len(columns) >= 2
        and columns[0].header == "Титул"
        and columns[1].header == "Марка"
    ):
        return "C2"
    return "A2"


def _paint_exported_table(workbook: Any, table: ExportedTable) -> None:
    """Fill an empty openpyxl workbook from *table*."""

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from utils.excel import sanitize_cell_value_for_openpyxl

    header_fill = PatternFill("solid", fgColor=_HEADER_FILL_RGB)
    header_font = Font(
        name="Calibri", size=11, bold=True, color=_HEADER_FG_RGB
    )
    header_align = Alignment(wrap_text=True, vertical="center")
    data_align = Alignment(wrap_text=False, vertical="center")
    style_cache: dict[
        tuple[str | None, str | None, bool, bool],
        tuple[PatternFill | None, Font],
    ] = {}

    sheet = workbook.active
    sheet.title = _sheet_title(table.sheet_name)
    n_cols = len(table.columns)

    for col_idx, column in enumerate(table.columns, start=1):
        header_cell = sheet.cell(
            1,
            col_idx,
            sanitize_cell_value_for_openpyxl(column.header, enabled=True),
        )
        header_cell.font = header_font
        header_cell.fill = header_fill
        header_cell.alignment = header_align
    sheet.row_dimensions[1].height = 30

    for row_idx, row in enumerate(table.rows, start=2):
        for col_idx, exported in enumerate(row, start=1):
            if n_cols and col_idx > n_cols:
                break
            value = sanitize_cell_value_for_openpyxl(
                exported.text, enabled=True
            )
            data_cell = sheet.cell(row_idx, col_idx, value)
            fill_rgb = excel_rgb(exported.fill_hex)
            fg_rgb = excel_rgb(exported.fg_hex)
            key = (fill_rgb, fg_rgb, exported.bold, exported.underline)
            cached = style_cache.get(key)
            if cached is None:
                fill_obj = (
                    PatternFill("solid", fgColor=fill_rgb) if fill_rgb else None
                )
                font_kwargs: dict[str, object] = {
                    "name": "Calibri",
                    "size": 11,
                    "bold": exported.bold,
                    "underline": "single" if exported.underline else None,
                }
                if fg_rgb:
                    font_kwargs["color"] = fg_rgb
                cached = (fill_obj, Font(**font_kwargs))
                style_cache[key] = cached
            fill_obj, font_obj = cached
            data_cell.font = font_obj
            if fill_obj is not None:
                data_cell.fill = fill_obj
            data_cell.alignment = data_align

    for col_idx, column in enumerate(table.columns, start=1):
        letter = get_column_letter(col_idx)
        sheet.column_dimensions[letter].width = px_to_excel_width(
            column.width_px
        )

    sheet.freeze_panes = _freeze_anchor(table.columns)
    if n_cols:
        last_col = get_column_letter(n_cols)
        last_row = max(1, 1 + len(table.rows))
        sheet.auto_filter.ref = f"A1:{last_col}{last_row}"


def write_exported_table_xlsx(table: ExportedTable, path: str | Path) -> None:
    """Write ``table`` to an xlsx workbook and close the file.

    Args:
        table: Visual-order columns and painted data rows.
        path: Destination workbook. Parent directories are created.

    Raises:
        OSError: If the workbook cannot be written.
    """

    import openpyxl

    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    try:
        _paint_exported_table(workbook, table)
        workbook.save(dest)
    finally:
        workbook.close()


def exported_table_xlsx_bytes(table: ExportedTable) -> bytes:
    """Return the workbook as xlsx bytes (WEB download, no disk file).

    Args:
        table: Visual-order columns and painted data rows.

    Returns:
        OOXML workbook bytes.
    """

    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    try:
        _paint_exported_table(workbook, table)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()
