"""Copy the newest RD/SQ MTO xlsx into the robot folder.

The operation is user-initiated: RD and SQ are never written. The previous
robot file is moved into a sibling ``_old_<stem>_YYYY.MM.DD_HH.MM`` folder.
Those folders are skipped by the robot scanner (substring ``old``).
"""

from __future__ import annotations

import os
import shutil
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from rd_catalog.kits import format_revision, kit_identity_key
from rd_catalog.models import (
    FileKind,
    FileRecord,
    MtoContentStatus,
    MtoIssueKind,
    MtoReadinessStatus,
    ParsedFile,
    ParseStatus,
    SourceKind,
    make_path_key,
)
from rd_catalog.mto_diff import (
    CanonicalMtoRow,
    MtoComparisonResult,
    MtoPair,
    MtoRowChange,
    RowLoader,
    compare_mto_pair,
    load_canonical_mto,
)
from rd_catalog.overlay import revision_rank
from rd_catalog.parse import parse_catalog_file
from rd_catalog.path_actions import path_is_under
from rd_catalog.perf_log import perf_span
from rd_catalog.scan import _ROBOT_VALID_SUBFOLDERS

_SOURCE_LABELS = {
    SourceKind.RD: "РД",
    SourceKind.SQ: "SQ",
    SourceKind.ROBOT: "робот",
}


@dataclass(frozen=True, slots=True)
class MtoFileRef:
    """One MTO xlsx used as a copy source or robot target."""

    path: str
    name: str
    source: SourceKind
    title: str
    mark: str
    title_system: str
    discipline_block: str | None
    revision: str | None
    appendix: str | None
    revision_text: str
    mtime_ns: int
    size: int

    @property
    def source_label(self) -> str:
        """Return the Russian source name for dialogs."""

        return _SOURCE_LABELS.get(self.source, self.source.value)


@dataclass(frozen=True, slots=True)
class RobotMtoSyncPlan:
    """Preview of one kit MTO copy/replace into ``robot_root``."""

    title: str
    mark: str
    source: MtoFileRef
    robot_files: tuple[MtoFileRef, ...]
    destination_dir: str
    destination_path: str
    robot_root: str

    @property
    def is_replace(self) -> bool:
        """Return True when at least one robot file will be archived."""

        return bool(self.robot_files)


@dataclass(frozen=True, slots=True)
class ArchivedRobotFile:
    """One robot MTO moved aside before the copy."""

    original_path: str
    archive_dir: str
    archived_path: str


@dataclass(frozen=True, slots=True)
class RobotMtoSyncResult:
    """Filesystem outcome of :func:`execute_robot_mto_sync`."""

    destination_path: str
    archived: tuple[ArchivedRobotFile, ...]
    added: bool


@dataclass(frozen=True, slots=True)
class RobotMtoComparePreview:
    """Live source↔robot comparison for the confirmation dialog."""

    summary: str
    report_path: str | None = None
    readiness: str | None = None
    appear: int = 0
    disappear: int = 0
    changed: int = 0
    source_rows: int = 0
    robot_rows: int = 0
    error: str | None = None


class RobotMtoSyncError(Exception):
    """User-visible failure while planning or executing a robot MTO copy."""


def archive_folder_name(file_stem: str, when: datetime) -> str:
    """Build ``_old_<stem>_YYYY.MM.DD_HH.MM``.

    Args:
        file_stem: Robot filename without the extension.
        when: Local timestamp of the replacement.

    Returns:
        Folder name for the archived robot file.
    """

    stamp = when.strftime("%Y.%m.%d_%H.%M")
    return f"_old_{file_stem}_{stamp}"


def format_mtime_ns(mtime_ns: int | None) -> str:
    """Format a nanosecond mtime as ``YYYY.MM.DD HH:MM``.

    Args:
        mtime_ns: Filesystem mtime in nanoseconds.

    Returns:
        Display timestamp, or ``—`` when missing/invalid.
    """

    try:
        return datetime.fromtimestamp(int(mtime_ns) / 1_000_000_000).strftime(
            "%Y.%m.%d %H:%M"
        )
    except (TypeError, ValueError, OSError):
        return "—"


def format_size(size: int | None) -> str:
    """Format a byte size for the confirmation dialog.

    Args:
        size: File size in bytes.

    Returns:
        Human-readable size in Russian units.
    """

    try:
        value = int(size)
    except (TypeError, ValueError):
        return "—"
    if value < 1024:
        return f"{value} байт"
    if value < 1024 * 1024:
        return f"{value / 1024:.1f} КБ"
    return f"{value / (1024 * 1024):.1f} МБ"


def unique_archive_folder(dest_dir: Path, file_stem: str, when: datetime) -> Path:
    """Return a non-existing archive folder path next to the robot file.

    Args:
        dest_dir: Robot file directory.
        file_stem: Filename without extension.
        when: Timestamp used in the folder name.

    Returns:
        Path that does not exist yet. A numeric suffix is added on collision.
    """

    base = archive_folder_name(file_stem, when)
    candidate = dest_dir / base
    if not candidate.exists():
        return candidate
    for index in range(2, 100):
        candidate = dest_dir / f"{base}_{index}"
        if not candidate.exists():
            return candidate
    raise RobotMtoSyncError(
        f"Не удалось подобрать имя папки архива в {dest_dir}"
    )


def _freshness_key(ref: MtoFileRef) -> tuple[tuple[int, int, str], int]:
    return revision_rank(ref.revision, ref.appendix), int(ref.mtime_ns or 0)


def _matches_kit(title: str, mark: str, other_title: str, other_mark: str) -> bool:
    return kit_identity_key(title, mark) == kit_identity_key(other_title, other_mark)


def _ref_from_record(record: FileRecord) -> MtoFileRef | None:
    if not record.present:
        return None
    if str(record.data.get("file_kind") or "") != FileKind.MTO_XLSX.value:
        return None
    if str(record.data.get("parse_status") or "") != ParseStatus.PARSED.value:
        return None
    title = str(record.data.get("title") or "").strip()
    mark = str(record.data.get("mark") or "").strip()
    if not title or not mark:
        return None
    revision = record.data.get("revision")
    appendix = record.data.get("appendix")
    revision_text = format_revision(
        str(revision) if revision not in (None, "") else None,
        str(appendix) if appendix not in (None, "") else None,
    )
    discipline = record.data.get("discipline_block")
    title_system = str(record.data.get("title_system") or f"{title}-{mark}")
    try:
        mtime_ns = int(record.data.get("mtime_ns") or 0)
        size = int(record.data.get("size") or 0)
    except (TypeError, ValueError):
        mtime_ns, size = 0, 0
    return MtoFileRef(
        path=record.path,
        name=str(record.data.get("name") or Path(record.path).name),
        source=record.source,
        title=title,
        mark=mark,
        title_system=title_system,
        discipline_block=str(discipline) if discipline else None,
        revision=str(revision) if revision not in (None, "") else None,
        appendix=str(appendix) if appendix not in (None, "") else None,
        revision_text=revision_text,
        mtime_ns=mtime_ns,
        size=size,
    )


def _ref_from_parsed(parsed: ParsedFile) -> MtoFileRef | None:
    if parsed.file_kind is not FileKind.MTO_XLSX:
        return None
    if parsed.parse_status is not ParseStatus.PARSED:
        return None
    if not parsed.title or not parsed.mark:
        return None
    return MtoFileRef(
        path=parsed.path,
        name=parsed.name,
        source=parsed.source,
        title=parsed.title,
        mark=parsed.mark,
        title_system=parsed.title_system or f"{parsed.title}-{parsed.mark}",
        discipline_block=parsed.discipline_block,
        revision=parsed.revision,
        appendix=parsed.appendix,
        revision_text=format_revision(parsed.revision, parsed.appendix),
        mtime_ns=int(parsed.mtime_ns or 0),
        size=int(parsed.size or 0),
    )


def discover_mto_in_folders(
    folders: Iterable[str],
    *,
    source: SourceKind,
    title: str,
    mark: str,
) -> list[MtoFileRef]:
    """Parse MTO xlsx files sitting next to known kit files.

    Args:
        folders: Directories to list (not recursive).
        source: RD or SQ identity assigned to discovered files.
        title: Kit title.
        mark: Kit mark.

    Returns:
        Parsed MTO refs matching ``title``/``mark``. Listing errors are skipped.
    """

    found: list[MtoFileRef] = []
    seen: set[str] = set()
    for folder in folders:
        try:
            names = os.listdir(folder)
        except OSError:
            continue
        for name in names:
            folded = name.casefold()
            if not folded.endswith(".xlsx"):
                continue
            if "mto" not in folded and "мто" not in folded:
                continue
            path = os.path.join(folder, name)
            key = make_path_key(path)
            if key in seen:
                continue
            seen.add(key)
            try:
                stat = os.stat(path, follow_symlinks=False)
                parsed = parse_catalog_file(
                    path,
                    source,
                    size=stat.st_size,
                    mtime_ns=stat.st_mtime_ns,
                )
            except OSError:
                continue
            ref = _ref_from_parsed(parsed)
            if ref is None:
                continue
            if not _matches_kit(title, mark, ref.title, ref.mark):
                continue
            found.append(ref)
    return found


def collect_kit_mto_candidates(
    *,
    title: str,
    mark: str,
    records: Iterable[FileRecord],
    detected_current_ids: set[int],
    rd_folders: Iterable[str] = (),
    sq_folders: Iterable[str] = (),
    discover_siblings: bool = True,
) -> list[MtoFileRef]:
    """Collect RD-current and SQ MTO files for one kit.

    Args:
        title: Four-digit title.
        mark: Latin AGCC mark.
        records: Persisted catalog files.
        detected_current_ids: Overlay-current RD file ids.
        rd_folders: Optional extra RD folders to list for sibling MTO.
        sq_folders: Optional SQ folders to list for sibling MTO.
        discover_siblings: When True, list ``rd_folders`` / ``sq_folders``.

    Returns:
        Deduplicated candidate refs. Robot files are not included.
    """

    by_key: dict[str, MtoFileRef] = {}
    for record in records:
        if record.source is SourceKind.ROBOT:
            continue
        ref = _ref_from_record(record)
        if ref is None or not _matches_kit(title, mark, ref.title, ref.mark):
            continue
        if record.source is SourceKind.RD:
            if record.id not in detected_current_ids:
                continue
        elif record.source is not SourceKind.SQ:
            continue
        by_key[make_path_key(ref.path)] = ref

    if discover_siblings:
        extra = discover_mto_in_folders(
            rd_folders, source=SourceKind.RD, title=title, mark=mark
        )
        extra.extend(
            discover_mto_in_folders(
                sq_folders, source=SourceKind.SQ, title=title, mark=mark
            )
        )
        for ref in extra:
            by_key.setdefault(make_path_key(ref.path), ref)
    return list(by_key.values())


def collect_robot_mto_files(
    *,
    title: str,
    mark: str,
    records: Iterable[FileRecord],
) -> list[MtoFileRef]:
    """Return present robot MTO files for one title+mark kit.

    Args:
        title: Four-digit title.
        mark: Latin AGCC mark.
        records: Persisted catalog files.

    Returns:
        Robot MTO refs for the kit, newest-first.
    """

    refs: list[MtoFileRef] = []
    for record in records:
        if record.source is not SourceKind.ROBOT:
            continue
        ref = _ref_from_record(record)
        if ref is None or not _matches_kit(title, mark, ref.title, ref.mark):
            continue
        refs.append(ref)
    refs.sort(key=_freshness_key, reverse=True)
    return refs


def robot_kit_directory(
    robot_root: str | Path,
    title: str,
    mark: str,
    *,
    flat_structure: bool,
    existing_robot_path: str | None = None,
) -> Path:
    """Choose the robot folder for a kit MTO.

    Args:
        robot_root: Configured robot MTO root.
        title: Four-digit title.
        mark: Latin AGCC mark.
        flat_structure: When True, files live directly in ``robot_root``.
        existing_robot_path: If set, reuse that file's parent folder.

    Returns:
        Destination directory for the new robot MTO.
    """

    if existing_robot_path:
        parent = Path(existing_robot_path).parent
        if str(parent):
            return parent
    root = Path(robot_root)
    if flat_structure:
        return root
    mark_folder = (mark or "").strip().upper()
    if mark_folder in _ROBOT_VALID_SUBFOLDERS:
        return root / str(title).strip() / mark_folder
    return root / str(title).strip()


def plan_robot_mto_sync(
    *,
    title: str,
    mark: str,
    records: Iterable[FileRecord],
    detected_current_ids: set[int],
    robot_root: str | Path,
    robot_flat_structure: bool,
    rd_folders: Iterable[str] = (),
    sq_folders: Iterable[str] = (),
    discover_siblings: bool = True,
) -> RobotMtoSyncPlan:
    """Build a copy/replace plan for one kits-row title+mark.

    The source is the newest present MTO among overlay-current RD files and
    SQ files (higher filename revision, then later mtime).

    Args:
        title: Four-digit title.
        mark: Latin AGCC mark.
        records: Persisted catalog files.
        detected_current_ids: Overlay-current RD file ids.
        robot_root: Configured robot MTO root.
        robot_flat_structure: Skip title/mark subfolders when True.
        rd_folders: Folders of current RD kit files (sibling MTO search).
        sq_folders: Folders of SQ kit files (sibling MTO search).
        discover_siblings: List sibling folders for MTO not yet in the DB.

    Returns:
        A validated plan that writes only under ``robot_root``.

    Raises:
        RobotMtoSyncError: If no source MTO exists or the destination is unsafe.
    """

    with perf_span("robot.plan", title=title, mark=mark):
        materialized = list(records)
        candidates = collect_kit_mto_candidates(
            title=title,
            mark=mark,
            records=materialized,
            detected_current_ids=detected_current_ids,
            rd_folders=rd_folders,
            sq_folders=sq_folders,
            discover_siblings=discover_siblings,
        )
        if not candidates:
            raise RobotMtoSyncError(
                f"Нет файла MTO в РД или SQ для {title}-{mark}."
            )
        source = max(candidates, key=_freshness_key)
        robot_root_text = str(robot_root)
        if path_is_under(source.path, robot_root_text):
            raise RobotMtoSyncError(
                "Источник MTO уже находится в папке робота — копирование не нужно."
            )
        robot_files = tuple(
            collect_robot_mto_files(title=title, mark=mark, records=materialized)
        )
        existing_path = robot_files[0].path if robot_files else None
        destination_dir = robot_kit_directory(
            robot_root_text,
            title,
            mark,
            flat_structure=robot_flat_structure,
            existing_robot_path=existing_path,
        )
        destination_path = destination_dir / source.name
        if not path_is_under(destination_dir, robot_root_text):
            raise RobotMtoSyncError(
                "Папка назначения не находится в каталоге робота."
            )
        if not path_is_under(destination_path, robot_root_text):
            raise RobotMtoSyncError(
                "Путь назначения не находится в каталоге робота."
            )
        if make_path_key(source.path) == make_path_key(destination_path):
            raise RobotMtoSyncError(
                "Источник и файл робота совпадают — копирование не нужно."
            )
        return RobotMtoSyncPlan(
            title=title,
            mark=mark,
            source=source,
            robot_files=robot_files,
            destination_dir=str(destination_dir),
            destination_path=str(destination_path),
            robot_root=robot_root_text,
        )


def _ref_from_existing_path(path: str, source: SourceKind) -> MtoFileRef:
    stat = os.stat(path, follow_symlinks=False)
    parsed = parse_catalog_file(
        path,
        source,
        size=stat.st_size,
        mtime_ns=stat.st_mtime_ns,
    )
    ref = _ref_from_parsed(parsed)
    if ref is not None:
        return ref
    name = Path(path).name
    return MtoFileRef(
        path=path,
        name=name,
        source=source,
        title="",
        mark="",
        title_system="",
        discipline_block=None,
        revision=None,
        appendix=None,
        revision_text="",
        mtime_ns=int(stat.st_mtime_ns),
        size=int(stat.st_size),
    )


def execute_robot_mto_sync(
    plan: RobotMtoSyncPlan,
    *,
    now: datetime | None = None,
) -> RobotMtoSyncResult:
    """Archive existing robot MTO files and copy the planned source.

    Args:
        plan: Validated copy/replace plan.
        now: Timestamp for archive folder names; defaults to local now.

    Returns:
        Paths of the new robot file and archived originals.

    Raises:
        RobotMtoSyncError: If the copy would write outside ``robot_root``
            or a filesystem operation fails.
    """

    with perf_span("robot.execute", title=plan.title, mark=plan.mark):
        if not path_is_under(plan.destination_dir, plan.robot_root):
            raise RobotMtoSyncError(
                "Папка назначения не находится в каталоге робота."
            )
        if not path_is_under(plan.destination_path, plan.robot_root):
            raise RobotMtoSyncError(
                "Путь назначения не находится в каталоге робота."
            )
        if make_path_key(plan.source.path) == make_path_key(plan.destination_path):
            raise RobotMtoSyncError("Нельзя копировать файл сам в себя.")
        if not os.path.isfile(plan.source.path):
            raise RobotMtoSyncError(f"Исходный файл не найден:\n{plan.source.path}")

        dest_dir = Path(plan.destination_dir)
        dest_path = Path(plan.destination_path)
        try:
            dest_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise RobotMtoSyncError(
                f"Не удалось создать папку робота:\n{dest_dir}\n{exc}"
            ) from exc

        when = now or datetime.now()
        to_archive: list[MtoFileRef] = list(plan.robot_files)
        dest_key = make_path_key(dest_path)
        archived_keys = {make_path_key(item.path) for item in to_archive}
        if dest_path.is_file() and dest_key not in archived_keys:
            try:
                to_archive.append(_ref_from_existing_path(str(dest_path), SourceKind.ROBOT))
            except OSError as exc:
                raise RobotMtoSyncError(
                    f"Не удалось прочитать существующий файл робота:\n{dest_path}\n{exc}"
                ) from exc

        archived: list[ArchivedRobotFile] = []
        for existing in to_archive:
            if not path_is_under(existing.path, plan.robot_root):
                raise RobotMtoSyncError(
                    f"Отказ перемещать файл вне каталога робота:\n{existing.path}"
                )
            if not os.path.isfile(existing.path):
                continue
            stem = Path(existing.name).stem or Path(existing.path).stem
            existing_dir = Path(existing.path).parent
            if not path_is_under(existing_dir, plan.robot_root):
                raise RobotMtoSyncError(
                    f"Папка архива вне каталога робота:\n{existing_dir}"
                )
            archive_dir = unique_archive_folder(existing_dir, stem, when)
            archived_path = archive_dir / Path(existing.path).name
            try:
                archive_dir.mkdir(parents=False, exist_ok=False)
                shutil.move(existing.path, str(archived_path))
            except OSError as exc:
                raise RobotMtoSyncError(
                    f"Не удалось отложить старый файл робота:\n{existing.path}\n{exc}"
                ) from exc
            archived.append(
                ArchivedRobotFile(
                    original_path=existing.path,
                    archive_dir=str(archive_dir),
                    archived_path=str(archived_path),
                )
            )

        try:
            shutil.copy2(plan.source.path, dest_path)
        except OSError as exc:
            raise RobotMtoSyncError(
                f"Не удалось скопировать MTO в папку робота:\n{dest_path}\n{exc}"
            ) from exc
        return RobotMtoSyncResult(
            destination_path=str(dest_path),
            archived=tuple(archived),
            added=not bool(archived),
        )


def plan_summary_text(plan: RobotMtoSyncPlan, *, now: datetime | None = None) -> str:
    """Build the confirmation-dialog body.

    Args:
        plan: Copy/replace plan.
        now: Timestamp used only to preview the archive folder name.

    Returns:
        Multi-line Russian description of source and robot target.
    """

    when = now or datetime.now()
    source = plan.source
    lines = [
        f"Комплект: {plan.title}-{plan.mark}",
        "",
        "Откуда берём файл",
        f"  Источник: {source.source_label}",
        f"  Файл: {source.name}",
        f"  Ревизия: {source.revision_text or '—'}",
        f"  Дата сохранения: {format_mtime_ns(source.mtime_ns)}",
        f"  Размер: {format_size(source.size)}",
        f"  Путь: {source.path}",
        "",
    ]
    if plan.is_replace:
        lines.append("Что заменяем у робота")
        for index, robot in enumerate(plan.robot_files, start=1):
            prefix = f"  [{index}] " if len(plan.robot_files) > 1 else "  "
            preview = archive_folder_name(Path(robot.name).stem, when)
            robot_dir = str(Path(robot.path).parent)
            lines.extend(
                [
                    f"{prefix}Файл: {robot.name}",
                    f"{prefix}Ревизия: {robot.revision_text or '—'}",
                    f"{prefix}Дата сохранения: {format_mtime_ns(robot.mtime_ns)}",
                    f"{prefix}Размер: {format_size(robot.size)}",
                    f"{prefix}Путь: {robot.path}",
                    f"{prefix}Старый файл будет перемещён в папку:",
                    f"{prefix}  {robot_dir}{os.sep}{preview}",
                ]
            )
        lines.extend(
            [
                "",
                "Новый файл робота:",
                f"  {plan.destination_path}",
            ]
        )
    else:
        lines.extend(
            [
                "Что заменяем у робота",
                "  Файла нет — MTO будет добавлен.",
                f"  Папка: {plan.destination_dir}",
                f"  Имя файла: {source.name}",
                f"  Полный путь: {plan.destination_path}",
            ]
        )
    return "\n".join(lines)


_ISSUE_LABELS = {
    MtoIssueKind.ROBOT_MISSING: "нет файла у робота",
    MtoIssueKind.ROBOT_DUPLICATE: "несколько файлов у робота",
    MtoIssueKind.ROBOT_EXTRA: "лишний файл у робота",
    MtoIssueKind.RD_UNPARSED: "источник не разобран",
    MtoIssueKind.ROBOT_UNPARSED: "файл робота не разобран",
    MtoIssueKind.LOAD_ERROR: "ошибка чтения xlsx",
    MtoIssueKind.COMPARE_ERROR: "ошибка сверки",
    MtoIssueKind.CONTENT_DIFF: "отличия содержимого",
    MtoIssueKind.REVISION_MISMATCH: "ревизия расходится",
    MtoIssueKind.MTIME_SUSPICION: "файл робота старше источника",
}
_CONTENT_LABELS = {
    MtoContentStatus.EQUAL: "совпадает",
    MtoContentStatus.DIFF: "есть отличия",
    MtoContentStatus.NOT_COMPARED: "не сравнивалось",
    MtoContentStatus.COMPARE_ERROR: "ошибка сверки",
}
_READINESS_LABELS = {
    MtoReadinessStatus.READY: "READY",
    MtoReadinessStatus.WARNING: "WARNING",
    MtoReadinessStatus.BLOCKED: "BLOCKED",
}
_COMPARE_EXAMPLES = 8


def _live_parsed_file(path: str, source: SourceKind) -> ParsedFile:
    stat = os.stat(path, follow_symlinks=False)
    return parse_catalog_file(
        path,
        source,
        size=stat.st_size,
        mtime_ns=stat.st_mtime_ns,
    )


def _row_preview(row: CanonicalMtoRow) -> str:
    name = (row.name or "").strip()
    if len(name) > 48:
        name = name[:45] + "…"
    code = row.code or "—"
    return f"{code}  {name}".rstrip()


def _field_preview(change: MtoRowChange) -> str:
    parts: list[str] = []
    for field in change.fields:
        left = field.rd_value
        right = field.robot_value
        if isinstance(left, list):
            left = "; ".join(str(item) for item in left)
        if isinstance(right, list):
            right = "; ".join(str(item) for item in right)
        parts.append(f"{field.field}: {left} → {right}")
    joined = "; ".join(parts)
    if len(joined) > 90:
        joined = joined[:87] + "…"
    return joined


def _compare_report_path(
    output_dir: Path,
    plan: RobotMtoSyncPlan,
    when: datetime,
) -> Path:
    stamp = when.strftime("%Y.%m.%d_%H.%M")
    stem = f"{plan.title}-{plan.mark}_mto_compare_{stamp}"
    candidate = output_dir / f"{stem}.xlsx"
    if not candidate.exists():
        return candidate
    for index in range(2, 100):
        candidate = output_dir / f"{stem}_{index}.xlsx"
        if not candidate.exists():
            return candidate
    raise RobotMtoSyncError(f"Не удалось подобрать имя файла сверки в {output_dir}")


def _cell_text(value: object) -> str:
    if isinstance(value, (list, tuple)):
        return "; ".join(str(item) for item in value)
    return "" if value is None else str(value)


def write_compare_workbook(
    path: str | Path,
    *,
    plan: RobotMtoSyncPlan,
    result: MtoComparisonResult,
    summary: str,
) -> str:
    """Write a local Excel preview of the live MTO comparison.

    Args:
        path: Destination xlsx under the catalog runtime directory.
        plan: Copy/replace plan (source vs robot).
        result: Live ``compare_mto_pair`` output.
        summary: Russian text already shown in the dialog.

    Returns:
        Saved workbook path.

    Raises:
        RobotMtoSyncError: If the workbook cannot be written.
    """

    import openpyxl

    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    for line in summary.splitlines() or ("",):
        summary_sheet.append([line])
    summary_sheet.append([])
    summary_sheet.append(["Источник", plan.source.path])
    summary_sheet.append(
        ["Робот", plan.robot_files[0].path if plan.robot_files else ""]
    )

    appear_sheet = workbook.create_sheet("Появятся")
    disappear_sheet = workbook.create_sheet("Пропадут")
    changed_sheet = workbook.create_sheet("Изменятся")
    headers = ["CODE", "UNITS", "VALUES", "TAGS", "NAME", "VENDOR", "TYPE_MARK"]
    appear_sheet.append(headers)
    disappear_sheet.append(headers)
    for row in result.diff.removed:
        appear_sheet.append(
            [
                row.code,
                row.units,
                row.values,
                _cell_text(row.tags),
                row.name,
                row.vendor,
                row.type_mark,
            ]
        )
    for row in result.diff.added:
        disappear_sheet.append(
            [
                row.code,
                row.units,
                row.values,
                _cell_text(row.tags),
                row.name,
                row.vendor,
                row.type_mark,
            ]
        )
    changed_sheet.append(
        ["CODE", "Поле", "Источник", "Робот", "NAME источника", "NAME робота"]
    )
    for change in result.diff.changed:
        for field in change.fields:
            changed_sheet.append(
                [
                    change.code,
                    field.field,
                    _cell_text(field.rd_value),
                    _cell_text(field.robot_value),
                    change.rd_row.name,
                    change.robot_row.name,
                ]
            )
    try:
        workbook.save(dest)
    except OSError as exc:
        raise RobotMtoSyncError(
            f"Не удалось сохранить файл сверки:\n{dest}\n{exc}"
        ) from exc
    finally:
        workbook.close()
    return str(dest)


def format_compare_summary(
    plan: RobotMtoSyncPlan,
    result: MtoComparisonResult | None,
    *,
    extra_notes: tuple[str, ...] = (),
    source_rows: int | None = None,
) -> str:
    """Build the Russian compare pane for the confirmation dialog.

    Args:
        plan: Copy/replace plan.
        result: Live comparison, or ``None`` when the robot file is missing.
        extra_notes: Additional warning lines.
        source_rows: Fallback source row count when ``result`` is absent.

    Returns:
        Multi-line summary. ``added`` in the engine is robot-only (will
        disappear); ``removed`` is source-only (will appear after replace).
    """

    lines = [
        "Сверка сейчас (файлы прочитаны заново, не кэш скана)",
        "",
        f"Источник ({plan.source.source_label}): {plan.source.name}",
        f"  рев. {plan.source.revision_text or '—'} · "
        f"{format_mtime_ns(plan.source.mtime_ns)}",
    ]
    if result is None:
        lines.extend(
            [
                "",
                "У робота нет файла MTO — сравнивать не с чем.",
                f"В источнике строк: {source_rows if source_rows is not None else '—'}",
                "После добавления сверка станет доступна.",
            ]
        )
        lines.extend(extra_notes)
        return "\n".join(lines)

    robot = plan.robot_files[0] if plan.robot_files else None
    if robot is not None:
        lines.extend(
            [
                f"Робот: {robot.name}",
                f"  рев. {robot.revision_text or '—'} · "
                f"{format_mtime_ns(robot.mtime_ns)}",
            ]
        )
    lines.extend(
        [
            "",
            f"Строк: источник {result.rd_rows}, робот {result.robot_rows}",
            f"Статус готовности: {_READINESS_LABELS.get(result.readiness, result.readiness.value)}",
            f"Содержимое: {_CONTENT_LABELS.get(result.content_status, result.content_status.value)}",
        ]
    )
    if result.issues:
        labels = ", ".join(
            _ISSUE_LABELS.get(issue, issue.value) for issue in result.issues
        )
        lines.append(f"Замечания: {labels}")
    if result.error:
        lines.append(f"Ошибка: {result.error}")
    appear = len(result.diff.removed)
    disappear = len(result.diff.added)
    changed = len(result.diff.changed)
    lines.extend(
        [
            "",
            "После замены у робота:",
            f"  появятся строк: {appear} (есть в источнике, нет у робота)",
            f"  пропадут строк: {disappear} (есть у робота, нет в источнике)",
            f"  изменятся строк: {changed}",
        ]
    )
    examples: list[str] = []
    for row in result.diff.removed[:_COMPARE_EXAMPLES]:
        examples.append(f"  + {_row_preview(row)}")
    for row in result.diff.added[:_COMPARE_EXAMPLES]:
        examples.append(f"  − {_row_preview(row)}")
    for change in result.diff.changed[:_COMPARE_EXAMPLES]:
        examples.append(f"  Δ {change.code}  {_field_preview(change)}")
    if examples:
        lines.append("")
        lines.append("Примеры:")
        lines.extend(examples[:_COMPARE_EXAMPLES])
    lines.extend(extra_notes)
    return "\n".join(lines)


def preview_robot_mto_compare(
    plan: RobotMtoSyncPlan,
    output_dir: str | Path,
    *,
    loader: RowLoader | None = None,
    now: datetime | None = None,
) -> RobotMtoComparePreview:
    """Compare the planned source MTO with the current robot file.

    Reads both xlsx files now (same algorithm as the catalog scan). The
    optional workbook is written only under ``output_dir`` (runtime), never
    to RD/SQ/robot UNC roots.

    Args:
        plan: Copy/replace plan for one title+mark.
        output_dir: Local directory for the preview workbook.
        loader: Optional test row loader; production uses ``load_mto_rows``.
        now: Timestamp used in the workbook filename.

    Returns:
        Summary text and optional report path.
    """

    extra_notes: list[str] = []
    if len(plan.robot_files) > 1:
        extra_notes.append(
            f"У робота {len(plan.robot_files)} файлов; сверка с самым новым."
        )
    if not plan.robot_files:
        source_rows = 0
        error = None
        try:
            document = load_canonical_mto(plan.source.path, loader=loader)
            source_rows = len(document.rows)
        except Exception as exc:
            error = str(exc)
            extra_notes.append(f"Не удалось прочитать источник: {exc}")
        summary = format_compare_summary(
            plan,
            None,
            extra_notes=tuple(extra_notes),
            source_rows=source_rows,
        )
        return RobotMtoComparePreview(
            summary=summary,
            source_rows=source_rows,
            error=error,
        )

    try:
        source_parsed = _live_parsed_file(plan.source.path, plan.source.source)
        robot_parsed = _live_parsed_file(
            plan.robot_files[0].path, SourceKind.ROBOT
        )
    except OSError as exc:
        summary = (
            "Сверка сейчас (файлы прочитаны заново, не кэш скана)\n\n"
            f"Не удалось открыть файл: {exc}"
        )
        return RobotMtoComparePreview(summary=summary, error=str(exc))

    title_system = (
        source_parsed.title_system
        or plan.source.title_system
        or f"{plan.title}-{plan.mark}"
    )
    discipline = source_parsed.discipline_block or plan.source.discipline_block or ""
    pair = MtoPair(
        (title_system.casefold(), discipline.casefold()),
        source_parsed,
        (robot_parsed,),
    )
    result = compare_mto_pair(pair, loader=loader)
    summary = format_compare_summary(
        plan, result, extra_notes=tuple(extra_notes)
    )
    report_path = None
    if result.content_status is not MtoContentStatus.COMPARE_ERROR:
        when = now or datetime.now()
        dest = _compare_report_path(Path(output_dir), plan, when)
        try:
            report_path = write_compare_workbook(
                dest, plan=plan, result=result, summary=summary
            )
        except RobotMtoSyncError as exc:
            extra_notes.append(str(exc))
            summary = format_compare_summary(
                plan, result, extra_notes=tuple(extra_notes)
            )
    return RobotMtoComparePreview(
        summary=summary,
        report_path=report_path,
        readiness=result.readiness.value,
        appear=len(result.diff.removed),
        disappear=len(result.diff.added),
        changed=len(result.diff.changed),
        source_rows=result.rd_rows,
        robot_rows=result.robot_rows,
        error=result.error,
    )
