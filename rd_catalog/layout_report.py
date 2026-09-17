"""Qt-free RD layout-violation report for authors.

Groups :func:`~rd_catalog.parse.list_layout_violations` by
``(title, mark, reason)`` so one row is one thing to fix. Writes Excel with
openpyxl (already a project dependency). No disk IO besides the caller-chosen
xlsx path.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path, PureWindowsPath

from rd_catalog.kits import kit_identity_key
from rd_catalog.models import FileRecord, SourceKind
from rd_catalog.parse import (
    _relative_windows_parts,
    layout_reason_label,
    list_layout_violations,
    record_has_canonical_layout,
)

_EXAMPLE_LIMIT = 3
_SHEET_SUMMARY = "Сводка"
_SHEET_LOST = "Нет в каталоге"
_SHEET_WORK = "Что исправить"


@dataclass(frozen=True, slots=True)
class LayoutReportRow:
    """One author action: a kit plus one layout reason.

    Attributes:
        title: Four-digit title from the file records.
        mark: Latin AGCC mark from the file records.
        reason: Machine token from :func:`~rd_catalog.parse.classify_rd_layout_reason`.
        reason_label: Russian instruction for the author.
        file_count: Present RD files in this group.
        example_paths: A few paths (relative to ``rd_root`` when possible).
        package_hint: Nearest ``NN_`` folder when the group is a single
            package; ``N передачи`` when several distinct packages appear;
            empty when no ``NN_`` folder is present.
        mto_file_count: Files in this group that are MTO workbooks.
    """

    title: str
    mark: str
    reason: str
    reason_label: str
    file_count: int
    example_paths: tuple[str, ...]
    package_hint: str
    mto_file_count: int = 0


@dataclass(frozen=True, slots=True)
class LayoutReport:
    """Grouped layout-violation report.

    Sort key for ``rows`` (ascending): lost kit, then has-MTO, then more
    MTO files, then more files, then title / mark / reason. Lost kits with
    MTO workbooks come first because those files silently change export.

    Attributes:
        rows: Work-list rows in the sort order above.
        total_files: Present RD files outside the canonical chain.
        total_mto_files: Subset of ``total_files`` that are MTO workbooks.
        total_kits: Distinct ``(title, mark)`` with at least one violation.
        by_reason: ``(token, file_count, mto_file_count)`` triples, highest
            file count first.
        fully_lost_kits: Kits that have present RD files but zero canonical
            files, so they vanish from the catalog.
    """

    rows: tuple[LayoutReportRow, ...]
    total_files: int
    total_mto_files: int
    total_kits: int
    by_reason: tuple[tuple[str, int, int], ...]
    fully_lost_kits: tuple[tuple[str, str], ...]


def _record_title_mark(record: FileRecord) -> tuple[str, str]:
    title = str(record.data.get("title") or "").strip()
    mark = str(record.data.get("mark") or "").strip()
    return title, mark


def _example_path(path: str, rd_root: str) -> str:
    relative = _relative_windows_parts(path, rd_root)
    if relative is None:
        return path
    return str(PureWindowsPath(*relative))


def _transfer_count_label(count: int) -> str:
    """Return a Russian count of distinct issued packages."""

    mod100 = count % 100
    mod10 = count % 10
    if mod10 == 1 and mod100 != 11:
        word = "передача"
    elif mod10 in {2, 3, 4} and mod100 not in {12, 13, 14}:
        word = "передачи"
    else:
        word = "передач"
    return f"{count} {word}"


def _package_hint_for_group(hints: Sequence[str]) -> str:
    """Return one package name, a distinct-package count, or empty.

    A single distinct ``NN_`` name is shown as-is. Several distinct names
    become ``N передачи`` so the cell does not point at a random transfer
    the examples do not show. Empty hints are ignored (they mean "no NN
    folder"), which is a different fact from "several packages".
    """

    distinct = tuple(dict.fromkeys(hint for hint in hints if hint))
    if not distinct:
        return ""
    if len(distinct) == 1:
        return distinct[0]
    return _transfer_count_label(len(distinct))


def build_layout_report(
    *,
    records: Sequence[FileRecord],
    rd_root: str | Path,
) -> LayoutReport:
    """Group layout violations into an author work list.

    Sort key, ascending:

    1. fully-lost kit (0) before kits that still have a canonical file (1);
    2. row with MTO files (0) before a row with none (1);
    3. more MTO files first;
    4. more files first;
    5. title, mark, reason.

    Args:
        records: Catalog file rows, typically ``list_files()``.
        rd_root: RD source root from :class:`~rd_catalog.config.CatalogConfig`.

    Returns:
        Grouped report. Classification is lexical (no disk IO).
    """

    root = str(rd_root or "").strip()
    violations = list_layout_violations(records=records, rd_root=root)

    canonical_keys: set[tuple[str, str]] = set()
    rd_keys: set[tuple[str, str]] = set()
    display_names: dict[tuple[str, str], tuple[str, str]] = {}
    for record in records:
        if not record.present or record.source is not SourceKind.RD:
            continue
        title, mark = _record_title_mark(record)
        key = kit_identity_key(title, mark)
        rd_keys.add(key)
        display_names.setdefault(key, (title, mark))
        if record_has_canonical_layout(record, root):
            canonical_keys.add(key)

    lost_keys = {key for key in rd_keys if key not in canonical_keys}
    fully_lost_kits = tuple(
        sorted(
            (display_names[key][0], display_names[key][1])
            for key in lost_keys
        )
    )

    grouped: dict[tuple[str, str, str], list] = defaultdict(list)
    display_group: dict[tuple[str, str, str], tuple[str, str]] = {}
    for row in violations:
        key = (row.title.casefold(), row.mark.casefold(), row.reason)
        grouped[key].append(row)
        display_group.setdefault(key, (row.title, row.mark))

    report_rows: list[LayoutReportRow] = []
    for key, items in grouped.items():
        title, mark = display_group[key]
        examples = tuple(
            _example_path(item.path, root) for item in items[:_EXAMPLE_LIMIT]
        )
        report_rows.append(
            LayoutReportRow(
                title=title,
                mark=mark,
                reason=key[2],
                reason_label=items[0].reason_label or layout_reason_label(key[2]),
                file_count=len(items),
                example_paths=examples,
                package_hint=_package_hint_for_group(
                    tuple(item.package_hint for item in items)
                ),
                mto_file_count=sum(1 for item in items if item.is_mto),
            )
        )

    lost_identity = {
        kit_identity_key(title, mark) for title, mark in fully_lost_kits
    }
    report_rows.sort(
        key=lambda item: (
            0 if kit_identity_key(item.title, item.mark) in lost_identity else 1,
            0 if item.mto_file_count else 1,
            -item.mto_file_count,
            -item.file_count,
            item.title.casefold(),
            item.mark.casefold(),
            item.reason,
        )
    )

    reason_files: Counter[str] = Counter()
    reason_mto: Counter[str] = Counter()
    for item in report_rows:
        reason_files[item.reason] += item.file_count
        reason_mto[item.reason] += item.mto_file_count
    by_reason = tuple(
        (token, reason_files[token], reason_mto[token])
        for token, _count in reason_files.most_common()
    )
    kit_ids = {
        kit_identity_key(item.title, item.mark) for item in report_rows
    }
    total_mto_files = sum(item.mto_file_count for item in report_rows)
    return LayoutReport(
        rows=tuple(report_rows),
        total_files=sum(item.file_count for item in report_rows),
        total_mto_files=total_mto_files,
        total_kits=len(kit_ids),
        by_reason=by_reason,
        fully_lost_kits=fully_lost_kits,
    )


def format_layout_report_text(report: LayoutReport) -> str:
    """Return a UTF-8 text rendering of ``report``.

    Args:
        report: Grouped layout report.

    Returns:
        Multi-line text with a summary, fully-lost kits, and the work list.
    """

    lines = [
        "Раскладка РД: ТИТУЛ / МАРКА / «Для передачи» / NN_рев.…",
        "",
        f"Файлов вне раскладки: {report.total_files}",
        (
            f"MTO вне раскладки: {report.total_mto_files}  "
            "← приоритет: без этих файлов автоматическая выгрузка MTO "
            "молча теряет комплекты"
        ),
        f"Строк (титул + марка + причина): {len(report.rows)}",
        f"Комплектов с нарушениями: {report.total_kits}",
        f"Полностью потерянных комплектов: {len(report.fully_lost_kits)}",
        "",
        "По причинам (файлов / из них MTO):",
    ]
    for token, count, mto_count in report.by_reason:
        share = 100.0 * count / report.total_files if report.total_files else 0.0
        lines.append(
            f"  {count:5d}  MTO={mto_count:4d}  {share:5.1f}%  {token}  "
            f"{layout_reason_label(token)}"
        )
    lines.extend(
        [
            "",
            "Полностью потерянные комплекты "
            "(нет ни одного файла в канонической раскладке — "
            "каталог покажет «есть в Google, нет в РД»):",
        ]
    )
    if not report.fully_lost_kits:
        lines.append("  (нет)")
    else:
        lost_file_counts: Counter[tuple[str, str]] = Counter()
        lost_mto_counts: Counter[tuple[str, str]] = Counter()
        for row in report.rows:
            key = kit_identity_key(row.title, row.mark)
            lost_file_counts[key] += row.file_count
            lost_mto_counts[key] += row.mto_file_count
        for title, mark in report.fully_lost_kits:
            key = kit_identity_key(title, mark)
            lines.append(
                f"  {title}/{mark}  файлов={lost_file_counts[key]}"
                f"  MTO={lost_mto_counts[key]}"
            )

    lines.extend(["", "Что исправить:"])
    lost_identity = {
        kit_identity_key(title, mark) for title, mark in report.fully_lost_kits
    }
    for row in report.rows:
        lost_flag = (
            "  [потерян]"
            if kit_identity_key(row.title, row.mark) in lost_identity
            else ""
        )
        lines.append(
            f"  {row.title}/{row.mark}  {row.reason}  "
            f"файлов={row.file_count}  MTO={row.mto_file_count}"
            f"{lost_flag}"
        )
        lines.append(f"    {row.reason_label}")
        if row.package_hint:
            lines.append(f"    пакет: {row.package_hint}")
        for example in row.example_paths:
            lines.append(f"    пример: {example}")
    lines.append("")
    return "\n".join(lines)


def format_layout_report_summary(report: LayoutReport) -> str:
    """Return a short confirmation text for the GUI save dialog.

    Built from :func:`format_layout_report_text`: the scale section only
    (no per-row work list). Leads with the MTO count. When there are no
    violations, states that plainly so the dialog is not empty.

    Args:
        report: Grouped layout report.

    Returns:
        Multi-line UTF-8 text suitable for ``QMessageBox``.
    """

    full = format_layout_report_text(report)
    head, _sep, _tail = full.partition("Что исправить:")
    lines = [f"MTO вне раскладки: {report.total_mto_files}"]
    if report.total_files == 0:
        lines.extend(
            [
                "",
                "Нарушений раскладки нет — все загруженные файлы РД "
                "лежат в канонической цепочке.",
            ]
        )
    lines.extend(["", head.strip()])
    return "\n".join(lines).strip() + "\n"


def layout_report_default_filename(when: datetime | None = None) -> str:
    """Return the default workbook name ``Раскладка_РД_YYYY.MM.DD.xlsx``.

    Args:
        when: Local timestamp. ``None`` uses ``datetime.now()``.

    Returns:
        Filename without a directory.
    """

    stamp = (when or datetime.now()).strftime("%Y.%m.%d")
    return f"Раскладка_РД_{stamp}.xlsx"


def write_layout_report_xlsx(report: LayoutReport, path: str | Path) -> None:
    """Write ``report`` to an xlsx workbook and close the file.

    Args:
        report: Grouped layout report.
        path: Destination workbook. Parent directories are created.

    Raises:
        OSError: If the workbook cannot be written.
    """

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)

    header_font = Font(bold=True)
    lost_fill = PatternFill("solid", fgColor="F4C7C3")
    wrap = Alignment(wrap_text=True, vertical="top")

    workbook = openpyxl.Workbook()
    try:
        summary = workbook.active
        summary.title = _SHEET_SUMMARY
        summary.append(["Раскладка РД"])
        summary.append(
            ["Канон", "ТИТУЛ / МАРКА / «Для передачи» / NN_рев.… / [PDF|DWG|BBB]"]
        )
        summary.append(["Файлов вне раскладки", report.total_files])
        summary.append(["MTO вне раскладки", report.total_mto_files])
        summary.append(
            [
                "Приоритет",
                "Эти MTO молча меняют выгрузку: без них комплект "
                "пропадает из автоматического экспорта",
            ]
        )
        summary.append(["Строк работы", len(report.rows)])
        summary.append(["Комплектов с нарушениями", report.total_kits])
        summary.append(
            ["Полностью потерянных комплектов", len(report.fully_lost_kits)]
        )
        summary.append([])
        summary.append(["Код", "Файлов", "из них MTO", "Что сделать"])
        for cell in summary[summary.max_row]:
            cell.font = header_font
        for token, count, mto_count in report.by_reason:
            summary.append(
                [token, count, mto_count, layout_reason_label(token)]
            )

        lost_sheet = workbook.create_sheet(_SHEET_LOST)
        lost_headers = [
            "Титул",
            "Марка",
            "Файлов вне раскладки",
            "из них MTO",
            "Причины",
        ]
        lost_sheet.append(lost_headers)
        for cell in lost_sheet[1]:
            cell.font = header_font
            cell.fill = lost_fill
        lost_identity = {
            kit_identity_key(title, mark) for title, mark in report.fully_lost_kits
        }
        lost_file_counts: Counter[tuple[str, str]] = Counter()
        lost_mto_counts: Counter[tuple[str, str]] = Counter()
        lost_reasons: dict[tuple[str, str], list[str]] = defaultdict(list)
        for row in report.rows:
            key = kit_identity_key(row.title, row.mark)
            if key not in lost_identity:
                continue
            lost_file_counts[key] += row.file_count
            lost_mto_counts[key] += row.mto_file_count
            if row.reason not in lost_reasons[key]:
                lost_reasons[key].append(row.reason)
        for title, mark in report.fully_lost_kits:
            key = kit_identity_key(title, mark)
            reasons = ", ".join(lost_reasons.get(key, ()))
            lost_sheet.append(
                [
                    title,
                    mark,
                    lost_file_counts.get(key, 0),
                    lost_mto_counts.get(key, 0),
                    reasons,
                ]
            )

        work = workbook.create_sheet(_SHEET_WORK)
        work_headers = [
            "Потерян",
            "Титул",
            "Марка",
            "Код",
            "Что сделать",
            "Файлов",
            "из них MTO",
            "Пакет",
            "Пример 1",
            "Пример 2",
            "Пример 3",
        ]
        work.append(work_headers)
        for cell in work[1]:
            cell.font = header_font
        for row in report.rows:
            lost = kit_identity_key(row.title, row.mark) in lost_identity
            examples = list(row.example_paths) + [""] * _EXAMPLE_LIMIT
            excel_row = [
                "да" if lost else "",
                row.title,
                row.mark,
                row.reason,
                row.reason_label,
                row.file_count,
                row.mto_file_count,
                row.package_hint,
                examples[0],
                examples[1],
                examples[2],
            ]
            work.append(excel_row)
            written = work[work.max_row]
            written[4].alignment = wrap
            if lost:
                for cell in written:
                    cell.fill = lost_fill

        for sheet, widths in (
            (summary, (28, 16, 14, 80)),
            (lost_sheet, (12, 12, 22, 14, 60)),
            (work, (12, 12, 12, 22, 70, 10, 12, 40, 50, 50, 50)),
        ):
            for index, width in enumerate(widths, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            sheet.freeze_panes = "A2"
            if sheet is not summary:
                sheet.auto_filter.ref = sheet.dimensions

        workbook.save(dest)
    finally:
        workbook.close()
