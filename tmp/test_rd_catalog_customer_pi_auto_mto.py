"""Smoke: АвтоМто filenames, Спецификация sheet, roundtrip into load_canonical_mto."""

from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rd_catalog.customer_pi import (
    COLUMN_KEYS,
    IDX_CODE_RD,
    IDX_MARK,
    IDX_NAME,
    IDX_QTY_RD,
    IDX_RD_REV,
    IDX_SPEC,
    IDX_SUPPLIER,
    IDX_TAGS,
    IDX_TITLE,
    IDX_TYPE_MARK,
    IDX_UNITS,
    N_COLS,
    CustomerPiStore,
    default_pickle_path,
    load_customer_pi,
    save_customer_pi,
)
from rd_catalog.auto_mto_compare_cache import (
    cache_key,
    entry_from_result,
    load_auto_mto_compare_cache,
    result_from_entry,
    save_auto_mto_compare_cache,
)
from rd_catalog.customer_pi_auto_mto import (
    AUTO_MTO_COMPARE_ALGORITHM_VERSION,
    AUTO_MTO_EXACT_STATUS,
    AUTO_MTO_SOFT_STATUS,
    HEADER_ROW,
    PAINT_OUTSIDE_KITS,
    SHEET_NAME,
    STATUS_EMPTY,
    STATUS_NOT_MTO,
    STATUS_QUEUED,
    STATUS_SKIPPED,
    STATUS_WRITTEN,
    STATUS_WRITING,
    AutoMtoCompareResult,
    AutoMtoFile,
    auto_mto_compare_status,
    auto_mto_filename,
    auto_mto_path,
    auto_mto_rd_path_key,
    auto_mto_rd_paths_match,
    cached_auto_mto_rd_paths_by_kit,
    compare_auto_mto_to_rd,
    format_kits_coverage_summary,
    format_spec_progress_summary,
    index_auto_mto_by_kit,
    kits_with_moved_cached_rd_mto_path,
    list_spec_progress_rows,
    mto_sheet_row,
    needs_auto_mto_compare,
    official_rd_mto_paths_changed,
    overlay_written_specs,
    plan_auto_mto_files,
    rebuild_auto_mto_catalog,
    revision_text_from_mto_filename,
    revision_texts_match,
    spec_as_mto_stem,
    spec_paint_key,
    write_auto_mto_workbook,
)
from rd_catalog.kits import kit_identity_key
from rd_catalog.mto_diff import canonicalize_mto_rows, load_canonical_mto, semantic_fingerprint
from utils.file_name_converts import parse_agcc_mto_xlsx_revision_for_chain


def _raw(**values: str) -> tuple[str, ...]:
    cells = [""] * N_COLS
    cells[IDX_SUPPLIER] = values.get("supplier", "БИ.СИ.СИ., ООО")
    cells[IDX_TITLE] = values.get("title", "8445")
    cells[IDX_MARK] = values.get("mark", "SOT")
    cells[IDX_SPEC] = values.get("spec", "AGCC.287-8445-SOT.MTO-0001")
    cells[IDX_RD_REV] = values.get("rd_revision", "04-AN01")
    cells[IDX_TAGS] = values.get("tags", "8445-S-SX-1002")
    cells[IDX_CODE_RD] = values.get("code_rd", "BCC0001")
    cells[IDX_NAME] = values.get("name", "Кабель силовой")
    cells[IDX_TYPE_MARK] = values.get("type_mark", "ВВГнг")
    cells[IDX_UNITS] = values.get("units", "м")
    cells[IDX_QTY_RD] = values.get("qty_rd", "12")
    return tuple(cells)


def _store(*rows: tuple[tuple[str, ...], tuple[str, str, str]]) -> CustomerPiStore:
    raw_rows = [item[0] for item in rows]
    parsed = [item[1] for item in rows]
    return CustomerPiStore(
        meta={
            "source_path": r"\\example\report.xlsb",
            "pickle_path": "memory",
            "parsed_at": "2026-09-10T00:00:00Z",
            "counts": {
                "specs": 2,
                "discipline": {"MTO": 2, "BOM": 1, "DS": 0},
            },
        },
        columns=list(COLUMN_KEYS),
        rows=raw_rows,
        excel_rows=list(range(3, 3 + len(raw_rows))),
        parsed=parsed,
    )


class AutoMtoCatalogSmoke(unittest.TestCase):
    def test_filename_parses_as_agcc_mto_chain(self) -> None:
        name = auto_mto_filename("AGCC.287-8445-SOT.MTO-0001", "04-AN01")
        self.assertEqual(name, "AGCC.287-8445-SOT.MTO-0001_04-AN01_RU.xlsx")
        self.assertEqual(
            parse_agcc_mto_xlsx_revision_for_chain(name),
            ("04", "AN01"),
        )
        name_plain = auto_mto_filename("AGCC.287-8445-SOT.MTO-0001", "04")
        self.assertEqual(
            parse_agcc_mto_xlsx_revision_for_chain(name_plain),
            ("04", ""),
        )
        name_ss30 = auto_mto_filename("AGCC.287-1600-SS30.MTO-0001", "0")
        self.assertEqual(name_ss30, "AGCC.287-1600-SS30.MTO-0001_0_RU.xlsx")
        self.assertEqual(parse_agcc_mto_xlsx_revision_for_chain(name_ss30), ("0", ""))
        name_ss30_an = auto_mto_filename("AGCC.287-6400-SS30.MTO-0001", "01-AN01")
        self.assertEqual(
            parse_agcc_mto_xlsx_revision_for_chain(name_ss30_an),
            ("01", "AN01"),
        )
        name_pd21 = auto_mto_filename("AGCC.287-8150-PD21.MTO-0001", "01")
        self.assertEqual(parse_agcc_mto_xlsx_revision_for_chain(name_pd21), ("01", ""))

    def test_filename_rejects_unsafe_and_empty(self) -> None:
        with self.assertRaises(ValueError):
            auto_mto_filename("", "01")
        with self.assertRaises(ValueError):
            auto_mto_filename("AGCC.287-8445-SOT.MTO-0001", "")
        with self.assertRaises(ValueError):
            auto_mto_filename("AGCC.287-8445-SOT/MTO-0001", "01")

    def test_spec_as_mto_stem_rewrites_bom_boe_ds(self) -> None:
        self.assertEqual(
            spec_as_mto_stem("AGCC.287-8445-SOT.BOM-0001"),
            "AGCC.287-8445-SOT.MTO-0001",
        )
        self.assertEqual(
            spec_as_mto_stem("AGCC.287-8445-SOT.BOE-0001"),
            "AGCC.287-8445-SOT.MTO-0001",
        )
        self.assertEqual(
            spec_as_mto_stem("AGCC.287-7450-TXM6.DS-0003"),
            "AGCC.287-7450-TXM6.MTO-0003",
        )
        self.assertEqual(
            spec_as_mto_stem("AGCC.287-8445-SOT.MTO-0001"),
            "AGCC.287-8445-SOT.MTO-0001",
        )
        self.assertEqual(
            spec_as_mto_stem("AGCC.287-0000-94A-0009"),
            "AGCC.287-0000-94A-0009",
        )

    def test_plan_converts_bom_and_nests_title_mark(self) -> None:
        store = _store(
            (
                _raw(),
                ("8445", "SOT", "MTO"),
            ),
            (
                _raw(
                    spec="AGCC.287-8445-SOT.BOM-0001",
                    name="Не MTO",
                ),
                ("8445", "SOT", "BOM"),
            ),
            (
                _raw(
                    spec="AGCC.287-8445-SOT.MTO-0002",
                    rd_revision="04",
                    code_rd="BCC0002",
                    tags="T-202",
                ),
                ("8445", "SOT", "MTO"),
            ),
        )
        planned, skipped = plan_auto_mto_files(store)
        self.assertEqual(skipped, ())
        specs = [item.spec for item in planned]
        self.assertIn("AGCC.287-8445-SOT.MTO-0001", specs)
        self.assertIn("AGCC.287-8445-SOT.MTO-0002", specs)
        self.assertTrue(all(".BOM-" not in item.spec for item in planned))
        primary = index_auto_mto_by_kit(store)[kit_identity_key("8445", "SOT")]
        self.assertEqual(primary.spec, "AGCC.287-8445-SOT.MTO-0001")
        self.assertEqual(primary.row_count, 1)
        self.assertEqual(
            primary.source_specs,
            (
                "AGCC.287-8445-SOT.MTO-0001",
                "AGCC.287-8445-SOT.BOM-0001",
            ),
        )
        self.assertEqual(primary.extras, ("AGCC.287-8445-SOT.MTO-0002 04",))
        self.assertEqual(
            primary.relpath,
            "8445/SOT/AGCC.287-8445-SOT.MTO-0001_04-AN01_RU.xlsx",
        )

    def test_plan_converts_bom_only_and_ds_stems(self) -> None:
        store = _store(
            (
                _raw(
                    spec="AGCC.287-1513-POS.BOM-0001",
                    title="1513",
                    mark="POS",
                    rd_revision="01",
                    tags="1513-P-1",
                ),
                ("1513", "POS", "BOM"),
            ),
            (
                _raw(
                    spec="AGCC.287-7450-TXM6.DS-0003",
                    title="7450",
                    mark="TXM6",
                    rd_revision="01",
                    tags="7450-T-1",
                    code_rd="BCC0009",
                ),
                ("7450", "TXM6", "DS"),
            ),
        )
        planned, skipped = plan_auto_mto_files(store)
        self.assertEqual(skipped, ())
        by_spec = {item.spec: item for item in planned}
        self.assertEqual(
            by_spec["AGCC.287-1513-POS.MTO-0001"].relpath,
            "1513/POS/AGCC.287-1513-POS.MTO-0001_01_RU.xlsx",
        )
        self.assertEqual(
            by_spec["AGCC.287-7450-TXM6.MTO-0003"].relpath,
            "7450/TXM6/AGCC.287-7450-TXM6.MTO-0003_01_RU.xlsx",
        )
        self.assertEqual(
            index_auto_mto_by_kit(store)[kit_identity_key("1513", "POS")].spec,
            "AGCC.287-1513-POS.MTO-0001",
        )

    def test_different_revisions_stay_separate_files(self) -> None:
        store = _store(
            (
                _raw(
                    spec="AGCC.287-8445-SOT1.MTO-0001",
                    title="8445",
                    mark="SOT1",
                    rd_revision="02-AN01",
                ),
                ("8445", "SOT1", "MTO"),
            ),
            (
                _raw(
                    spec="AGCC.287-8445-SOT1.DS-0001",
                    title="8445",
                    mark="SOT1",
                    rd_revision="0",
                    tags="8445-D-1",
                    code_rd="BCC0099",
                ),
                ("8445", "SOT1", "DS"),
            ),
        )
        planned, skipped = plan_auto_mto_files(store)
        self.assertEqual(skipped, ())
        revs = {(item.spec, item.rd_revision) for item in planned}
        self.assertEqual(
            revs,
            {
                ("AGCC.287-8445-SOT1.MTO-0001", "02-AN01"),
                ("AGCC.287-8445-SOT1.MTO-0001", "0"),
            },
        )
        from rd_catalog.customer_pi_auto_mto import (
            format_auto_mto_cell_text,
            list_auto_mto_files_by_kit,
            pick_auto_mto_file,
        )

        files = list_auto_mto_files_by_kit(store)[kit_identity_key("8445", "SOT1")]
        self.assertEqual(len(files), 2)
        primary = pick_auto_mto_file(files)
        self.assertEqual(primary.rd_revision, "02-AN01")
        matched = pick_auto_mto_file(files, "0")
        self.assertEqual(matched.rd_revision, "0")
        self.assertIn("0", primary.extras[0])
        self.assertEqual(format_auto_mto_cell_text(files), "02-AN01 (2)")
        self.assertEqual(format_auto_mto_cell_text(files, "0"), "0 (2)")
        self.assertEqual(format_auto_mto_cell_text(files[:1]), "02-AN01")
        self.assertEqual(format_auto_mto_cell_text(()), "")

    def test_write_sheet_and_roundtrip_fingerprint(self) -> None:
        store = _store((_raw(), ("8445", "SOT", "MTO")))
        record = store.record(0)
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "out.xlsx"
            fingerprint = write_auto_mto_workbook(dest, [record])
            from openpyxl import load_workbook

            workbook = load_workbook(dest, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, [SHEET_NAME])
                rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(tuple(rows[0]), HEADER_ROW)
            self.assertEqual(rows[1][2], "Кабель силовой")
            self.assertEqual(rows[1][4], "BCC0001")
            self.assertIn(rows[1][5], ("", None))
            loaded = load_canonical_mto(dest)
            expected = semantic_fingerprint(
                canonicalize_mto_rows(
                    [{**record.as_mto_fields(), "TAGS": ["8445-S-SX-1002"]}]
                )
            )
            self.assertEqual(loaded.fingerprint, expected)
            self.assertEqual(fingerprint, expected)
            self.assertEqual(len(loaded.rows), 1)
            self.assertEqual(loaded.rows[0].code, "BCC0001")
            self.assertEqual(loaded.rows[0].vendor, "")

    def test_rebuild_writes_manifest_and_removes_stale(self) -> None:
        store = _store((_raw(), ("8445", "SOT", "MTO")))
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            stale_rel = "8445/SOT/stale_01_RU.xlsx"
            stale = root / stale_rel
            stale.parent.mkdir(parents=True)
            stale.write_bytes(b"stale")
            unrelated = root / "keep_me.txt"
            unrelated.write_text("keep", encoding="utf-8")
            (root / "_manifest.json").write_text(
                json.dumps({"files": [{"relpath": stale_rel}]}),
                encoding="utf-8",
            )
            result = rebuild_auto_mto_catalog(store, root)
            self.assertEqual(len(result.written), 1)
            written = auto_mto_path(result.written[0], root)
            self.assertTrue(written.is_file())
            self.assertFalse(stale.exists())
            self.assertTrue(unrelated.is_file())
            self.assertIn(stale_rel, result.removed)
            payload = json.loads(result.manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(payload["format"], "rd_catalog.auto_mto")
            self.assertEqual(payload["files"][0]["spec"], "AGCC.287-8445-SOT.MTO-0001")

    def test_revision_texts_match(self) -> None:
        self.assertTrue(revision_texts_match("04-AN01", "04-AN01"))
        self.assertFalse(revision_texts_match("04-AN01", "04"))
        self.assertIsNone(revision_texts_match("", "04"))
        self.assertEqual(
            revision_text_from_mto_filename(
                "AGCC.287-8445-SOT.MTO-0001_04-AN01_RU.xlsx"
            ),
            "04-AN01",
        )

    def test_live_pickle_plans_only_mto(self) -> None:
        path = default_pickle_path()
        if not path.is_file():
            raise unittest.SkipTest(f"missing dump {path}")
        store = load_customer_pi(path)
        planned, _skipped = plan_auto_mto_files(store)
        self.assertGreater(len(planned), 200)
        self.assertTrue(all(item.spec.upper().find(".MTO-") >= 0 for item in planned))
        self.assertTrue(all(".BOM-" not in item.spec.upper() for item in planned))
        self.assertTrue(all(".DS-" not in item.spec.upper() for item in planned))
        self.assertTrue(all(".BOE-" not in item.spec.upper() for item in planned))
        parse_agcc_mto_xlsx_revision_for_chain(Path(planned[0].relpath).name)

    def test_spec_progress_rows_classify_mto_bom_ss30(self) -> None:
        store = _store(
            (_raw(), ("8445", "SOT", "MTO")),
            (
                _raw(spec="AGCC.287-8445-SOT.BOM-0001", name="Не MTO"),
                ("8445", "SOT", "BOM"),
            ),
            (
                _raw(
                    spec="AGCC.287-0000-94A-0009",
                    title="0000",
                    mark="94A",
                ),
                ("", "", "other"),
            ),
            (
                _raw(
                    spec="AGCC.287-1600-SS30.MTO-0001",
                    title="1600",
                    mark="SS30",
                    rd_revision="0",
                    tags="1600-S-SS-1",
                ),
                ("1600", "SS30", "MTO"),
            ),
            (
                _raw(spec="", rd_revision="", tags=""),
                ("", "", "other"),
            ),
        )
        rows = list_spec_progress_rows(store)
        by_spec = {row.spec: row for row in rows}
        self.assertEqual(by_spec["AGCC.287-8445-SOT.MTO-0001"].status, STATUS_QUEUED)
        self.assertEqual(by_spec["AGCC.287-8445-SOT.MTO-0001"].row_count, 1)
        self.assertEqual(by_spec["AGCC.287-8445-SOT.BOM-0001"].status, STATUS_QUEUED)
        self.assertEqual(by_spec["AGCC.287-0000-94A-0009"].status, STATUS_NOT_MTO)
        self.assertEqual(by_spec["AGCC.287-1600-SS30.MTO-0001"].status, STATUS_QUEUED)
        self.assertEqual(by_spec[""].status, STATUS_EMPTY)
        written = overlay_written_specs(
            rows, [("AGCC.287-8445-SOT.MTO-0001", "04-AN01")]
        )
        written_map = {row.spec: row.status for row in written}
        self.assertEqual(written_map["AGCC.287-8445-SOT.MTO-0001"], STATUS_WRITTEN)
        self.assertEqual(written_map["AGCC.287-8445-SOT.BOM-0001"], STATUS_WRITTEN)
        self.assertEqual(written_map["AGCC.287-1600-SS30.MTO-0001"], STATUS_QUEUED)
        summary = format_spec_progress_summary(written)
        self.assertIn("записано 2", summary)
        self.assertIn("ожидает 1", summary)
        self.assertIn("пропуск 0", summary)
        self.assertIn("не MTO 1", summary)

    def test_live_pickle_two_digit_marks_are_queued(self) -> None:
        path = default_pickle_path()
        if not path.is_file():
            raise unittest.SkipTest(f"missing dump {path}")
        store = load_customer_pi(path)
        rows = list_spec_progress_rows(store)
        tokens = ("SS30", "PD21", "PD23", "RT30")
        two_digit = [
            row for row in rows if any(token in row.spec for token in tokens)
        ]
        self.assertGreaterEqual(
            len([row for row in two_digit if "SS30" in row.spec]), 10
        )
        self.assertTrue(all(row.status == STATUS_QUEUED for row in two_digit))
        skipped = [row for row in rows if row.status == STATUS_SKIPPED]
        self.assertFalse(
            any(any(token in row.spec for token in tokens) for row in skipped)
        )
        self.assertGreater(len(rows), 150)

    def test_kits_coverage_and_outside_paint(self) -> None:
        store = _store(
            (
                _raw(),
                ("8445", "SOT", "MTO"),
            ),
            (
                _raw(
                    spec="AGCC.287-0000-94A-0009",
                    title="0000",
                    mark="94A",
                    rd_revision="0",
                    tags="",
                ),
                ("0000", "94A", "other"),
            ),
        )
        rows = list_spec_progress_rows(store)
        kit_keys = {
            kit_identity_key("8445", "SOT"),
            kit_identity_key("1600", "POS"),
        }
        summary = format_kits_coverage_summary(rows, kit_keys)
        self.assertIn("Комплекты: 2", summary)
        self.assertIn("получено МТО: 1", summary)
        self.assertIn("записано: 0", summary)
        self.assertIn("нет в ПИ: 1", summary)
        self.assertIn("спеки вне комплектов: 1", summary)
        sot = next(row for row in rows if "SOT" in row.spec)
        other = next(row for row in rows if "94A" in row.spec)
        self.assertEqual(spec_paint_key(sot, kit_keys), STATUS_QUEUED)
        self.assertEqual(spec_paint_key(other, kit_keys), PAINT_OUTSIDE_KITS)
        self.assertEqual(
            spec_paint_key(other, kit_keys, STATUS_WRITING),
            STATUS_WRITING,
        )
        cells = mto_sheet_row(store.record(0), 1)
        self.assertEqual(len(cells), len(HEADER_ROW))
        self.assertEqual(cells[0], "8445-S-SX-1002")
        self.assertEqual(cells[1], 1)
        self.assertEqual(cells[4], "BCC0001")
        self.assertEqual(cells[5], "")


class AutoMtoContentCompareSmoke(unittest.TestCase):
    def test_exact_single_skips_unreadable_candidate(self) -> None:
        store = _store(
            (
                _raw(rd_revision="03"),
                ("8445", "SOT", "MTO"),
            ),
            (
                _raw(rd_revision="02"),
                ("8445", "SOT", "MTO"),
            ),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            result = rebuild_auto_mto_catalog(store, root)
            by_revision = {item.rd_revision: item for item in result.written}
            auto_mto_path(by_revision["03"], root).write_bytes(b"broken xlsx")
            rd_path = Path(tmp) / "AGCC.287-8445-SOT.MTO-0001_04_RU.xlsx"
            write_auto_mto_workbook(rd_path, [store.record(0)])

            compared = compare_auto_mto_to_rd(result.written, rd_path, dest=root)

        self.assertEqual(AUTO_MTO_COMPARE_ALGORITHM_VERSION, 2)
        self.assertTrue(compared.matched)
        self.assertEqual(compared.match_kind, "single")
        self.assertEqual(compared.grade, "exact")
        self.assertEqual(compared.members[0].rd_revision, "02")
        self.assertEqual(compared.rd_rows, 1)
        self.assertEqual(compared.auto_rows, 1)
        self.assertEqual(compared.combinations_checked, 0)

    def test_exact_composite_mto_and_ds_union(self) -> None:
        store = _store(
            (
                _raw(
                    rd_revision="02",
                    code_rd="BCC-MTO",
                    tags="8445-MTO-1",
                    name="MTO row",
                ),
                ("8445", "SOT", "MTO"),
            ),
            (
                _raw(
                    spec="AGCC.287-8445-SOT.DS-0001",
                    rd_revision="0",
                    code_rd="BCC-DS",
                    tags="8445-DS-1",
                    name="DS row",
                ),
                ("8445", "SOT", "DS"),
            ),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            result = rebuild_auto_mto_catalog(store, root)
            rd_path = Path(tmp) / "AGCC.287-8445-SOT.MTO-0001_03_RU.xlsx"
            write_auto_mto_workbook(rd_path, [store.record(0), store.record(1)])

            compared = compare_auto_mto_to_rd(result.written, rd_path, dest=root)

        self.assertTrue(compared.matched)
        self.assertEqual(compared.match_kind, "composite")
        self.assertEqual(compared.grade, "exact")
        self.assertEqual(
            tuple(item.rd_revision for item in compared.members),
            ("02", "0"),
        )
        self.assertEqual(compared.cell_text(), "02 + 0")
        self.assertIn("суммы 2 файлов", compared.tooltip_lines()[0])
        self.assertEqual(compared.rd_rows, 2)
        self.assertEqual(compared.auto_rows, 2)
        self.assertEqual(compared.combinations_checked, 1)

    def test_same_original_source_revisions_are_not_composite(self) -> None:
        store = _store(
            (
                _raw(
                    rd_revision="02",
                    code_rd="BCC-OLD",
                    tags="8445-OLD-1",
                    name="Old row",
                ),
                ("8445", "SOT", "MTO"),
            ),
            (
                _raw(
                    rd_revision="0",
                    code_rd="BCC-NEW",
                    tags="8445-NEW-1",
                    name="New row",
                ),
                ("8445", "SOT", "MTO"),
            ),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            result = rebuild_auto_mto_catalog(store, root)
            rd_path = Path(tmp) / "AGCC.287-8445-SOT.MTO-0001_03_RU.xlsx"
            write_auto_mto_workbook(rd_path, [store.record(0), store.record(1)])

            compared = compare_auto_mto_to_rd(result.written, rd_path, dest=root)

        self.assertFalse(compared.matched)
        self.assertEqual(compared.match_kind, "no_match")
        self.assertEqual(compared.combinations_checked, 0)

    def test_different_output_stems_are_not_mixed(self) -> None:
        store = _store(
            (
                _raw(
                    code_rd="BCC-1",
                    tags="8445-MTO1-1",
                    name="Stem one",
                ),
                ("8445", "SOT", "MTO"),
            ),
            (
                _raw(
                    spec="AGCC.287-8445-SOT.MTO-0002",
                    rd_revision="03",
                    code_rd="BCC-2",
                    tags="8445-MTO2-1",
                    name="Stem two",
                ),
                ("8445", "SOT", "MTO"),
            ),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            result = rebuild_auto_mto_catalog(store, root)
            rd_path = Path(tmp) / "AGCC.287-8445-SOT.MTO-0001_05_RU.xlsx"
            write_auto_mto_workbook(rd_path, [store.record(0), store.record(1)])

            compared = compare_auto_mto_to_rd(result.written, rd_path, dest=root)

        self.assertEqual(compared.match_kind, "no_match")
        self.assertEqual(compared.combinations_checked, 0)

    def test_negative_content_is_no_match(self) -> None:
        store = _store((_raw(), ("8445", "SOT", "MTO")))
        other = _store(
            (
                _raw(code_rd="OTHER", tags="OTHER-1", name="Other row"),
                ("8445", "SOT", "MTO"),
            )
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            result = rebuild_auto_mto_catalog(store, root)
            rd_path = Path(tmp) / "AGCC.287-8445-SOT.MTO-0001_05_RU.xlsx"
            write_auto_mto_workbook(rd_path, [other.record(0)])

            compared = compare_auto_mto_to_rd(result.written, rd_path, dest=root)

        self.assertFalse(compared.matched)
        self.assertEqual(compared.match_kind, "no_match")
        self.assertEqual(compared.rd_rows, 1)

    def test_composite_checks_respect_maximum(self) -> None:
        spec = "AGCC.287-8445-SOT.MTO-0001"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            files: list[AutoMtoFile] = []
            rows_by_name: dict[str, list[dict[str, object]]] = {}
            for index in range(1, 6):
                name = f"{spec}_0{index}_RU.xlsx"
                relpath = f"8445/SOT/{name}"
                path = root / relpath
                path.parent.mkdir(parents=True, exist_ok=True)
                path.touch()
                files.append(
                    AutoMtoFile(
                        title="8445",
                        mark="SOT",
                        spec=spec,
                        rd_revision=f"0{index}",
                        relpath=relpath,
                        row_count=1,
                        fingerprint="",
                        source_specs=(f"SOURCE-{index}",),
                    )
                )
                rows_by_name[name] = [
                    {
                        "CODE": f"AUTO-{index}",
                        "UNITS": "шт",
                        "VALUES": "1",
                    }
                ]
            rd_path = Path(tmp) / f"{spec}_09_RU.xlsx"
            rd_path.touch()
            rows_by_name[rd_path.name] = [
                {"CODE": "RD-1", "UNITS": "шт", "VALUES": "1"},
                {"CODE": "RD-2", "UNITS": "шт", "VALUES": "1"},
            ]

            def loader(path: str | Path) -> list[dict[str, object]]:
                return rows_by_name[Path(path).name]

            compared = compare_auto_mto_to_rd(
                files,
                rd_path,
                dest=root,
                loader=loader,
                max_members=2,
                max_combinations=3,
            )

        self.assertEqual(compared.match_kind, "no_match")
        self.assertEqual(compared.combinations_checked, 3)


    def _grade_fixture(
        self,
        *,
        auto_row: dict[str, object],
        rd_row: dict[str, object],
    ) -> AutoMtoCompareResult:
        spec = "AGCC.287-8445-SOT.MTO-0001"
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "АвтоМто"
            relpath = f"8445/SOT/{spec}_02_RU.xlsx"
            path = root / relpath
            path.parent.mkdir(parents=True, exist_ok=True)
            path.touch()
            files = (
                AutoMtoFile(
                    title="8445",
                    mark="SOT",
                    spec=spec,
                    rd_revision="02",
                    relpath=relpath,
                    row_count=1,
                    fingerprint="",
                    source_specs=(spec,),
                ),
            )
            rd_path = Path(tmp) / f"{spec}_03_RU.xlsx"
            rd_path.touch()
            rows_by_name = {path.name: [auto_row], rd_path.name: [rd_row]}

            def loader(src: str | Path) -> list[dict[str, object]]:
                return rows_by_name[Path(src).name]

            return compare_auto_mto_to_rd(
                files, rd_path, dest=root, loader=loader
            )

    def test_exact_ignores_name_vendor_and_type_mark(self) -> None:
        compared = self._grade_fixture(
            auto_row={
                "CODE": "BCC1",
                "UNITS": "шт",
                "VALUES": "2",
                "TAGS": "T-1",
                "NAME": "short",
                "VENDOR": "",
                "TYPE_MARK": "",
            },
            rd_row={
                "CODE": "BCC1",
                "UNITS": "шт",
                "VALUES": "2",
                "TAGS": "T-1",
                "NAME": "long with vendor suffix",
                "VENDOR": "Болид",
                "TYPE_MARK": "РИП-12",
            },
        )
        self.assertEqual(compared.match_kind, "single")
        self.assertEqual(compared.grade, "exact")

    def test_soft_when_tags_or_units_differ(self) -> None:
        compared = self._grade_fixture(
            auto_row={
                "CODE": "BCC1",
                "UNITS": "м",
                "VALUES": "2",
                "TAGS": "T-AUTO",
            },
            rd_row={
                "CODE": "BCC1",
                "UNITS": "шт",
                "VALUES": "2",
                "TAGS": "T-RD",
            },
        )
        self.assertEqual(compared.match_kind, "single")
        self.assertEqual(compared.grade, "soft")
        self.assertTrue(compared.matched)

    def test_no_match_when_quantity_differs(self) -> None:
        compared = self._grade_fixture(
            auto_row={"CODE": "BCC1", "UNITS": "шт", "VALUES": "2", "TAGS": "T-1"},
            rd_row={"CODE": "BCC1", "UNITS": "шт", "VALUES": "3", "TAGS": "T-1"},
        )
        self.assertFalse(compared.matched)
        self.assertEqual(compared.match_kind, "no_match")


class AutoMtoDialogSmoke(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def test_dialog_shows_meta_without_loading_window(self) -> None:
        from rd_catalog.customer_pi_dialog import CustomerPiDialog

        store = _store((_raw(), ("8445", "SOT", "MTO")))
        with tempfile.TemporaryDirectory() as tmp:
            pickle_path = Path(tmp) / "pi.pkl"
            dest = Path(tmp) / "АвтоМто"
            save_customer_pi(store, pickle_path)
            dialog = CustomerPiDialog(
                pickle_path=pickle_path,
                xlsb_path=Path(tmp) / "report.xlsb",
                dest_dir=dest,
                store=store,
            )
            try:
                self.assertEqual(dialog.windowTitle(), "База заказчика (Авто МТО)")
                text = dialog._meta_view.toPlainText()
                self.assertIn("Строк BCC: 1", text)
                self.assertIn("АвтоМТО", text)
                self.assertFalse(dialog.is_busy())
                self.assertTrue(dialog._reload_button.isEnabled())
                self.assertTrue(dialog._catalog_button.isEnabled())
                self.assertTrue(dialog._open_catalog_button.isEnabled())
                self.assertEqual(
                    dialog._open_catalog_button.toolTip(), str(dest)
                )
                self.assertEqual(dialog._spec_table.rowCount(), 1)
                self.assertEqual(
                    dialog._spec_table.item(0, 0).text(),
                    "AGCC.287-8445-SOT.MTO-0001",
                )
                self.assertEqual(dialog._spec_table.item(0, 1).text(), "8445")
                self.assertEqual(dialog._spec_table.item(0, 2).text(), "SOT")
                self.assertEqual(dialog._spec_table.item(0, 5).text(), "1")
            finally:
                dialog.close()
                dialog.deleteLater()
                self.app.processEvents()

    def test_dialog_spec_table_paints_live_status(self) -> None:
        from rd_catalog.customer_pi_dialog import CustomerPiDialog

        store = _store(
            (_raw(), ("8445", "SOT", "MTO")),
            (
                _raw(
                    spec="AGCC.287-1600-SS30.MTO-0001",
                    title="1600",
                    mark="SS30",
                    rd_revision="0",
                ),
                ("1600", "SS30", "MTO"),
            ),
            (
                _raw(spec="AGCC.287-8445-SOT.BOM-0001"),
                ("8445", "SOT", "BOM"),
            ),
        )
        with tempfile.TemporaryDirectory() as tmp:
            pickle_path = Path(tmp) / "pi.pkl"
            dest = Path(tmp) / "АвтоМто"
            save_customer_pi(store, pickle_path)
            dialog = CustomerPiDialog(
                pickle_path=pickle_path,
                dest_dir=dest,
                store=store,
            )
            try:
                self.assertEqual(dialog._spec_table.rowCount(), 3)
                by_spec = {
                    dialog._spec_table.item(row, 0).text(): row
                    for row in range(dialog._spec_table.rowCount())
                }
                mto_row = by_spec["AGCC.287-8445-SOT.MTO-0001"]
                ss30_row = by_spec["AGCC.287-1600-SS30.MTO-0001"]
                bom_row = by_spec["AGCC.287-8445-SOT.BOM-0001"]
                self.assertEqual(dialog._spec_table.item(ss30_row, 6).text(), "ожидает")
                self.assertEqual(dialog._spec_table.item(bom_row, 6).text(), "ожидает")
                dialog._set_spec_status(
                    "AGCC.287-8445-SOT.MTO-0001", STATUS_WRITING
                )
                self.assertEqual(
                    dialog._spec_table.item(mto_row, 6).text(), "пишется"
                )
                self.assertEqual(
                    dialog._spec_table.item(bom_row, 6).text(), "пишется"
                )
                dialog._set_spec_status(
                    "AGCC.287-8445-SOT.MTO-0001", STATUS_WRITTEN
                )
                self.assertEqual(
                    dialog._spec_table.item(mto_row, 6).text(), "записано"
                )
                self.assertIn("ожидает", dialog._spec_summary_label.text())
            finally:
                dialog.close()
                dialog.deleteLater()
                self.app.processEvents()

    def test_open_catalog_button_opens_dest_dir(self) -> None:
        from unittest.mock import patch

        from rd_catalog.customer_pi_dialog import CustomerPiDialog

        store = _store((_raw(), ("8445", "SOT", "MTO")))
        with tempfile.TemporaryDirectory() as tmp:
            pickle_path = Path(tmp) / "pi.pkl"
            dest = Path(tmp) / "АвтоМто"
            dest.mkdir()
            save_customer_pi(store, pickle_path)
            dialog = CustomerPiDialog(
                pickle_path=pickle_path,
                dest_dir=dest,
                store=store,
            )
            try:
                with patch(
                    "rd_catalog.customer_pi_dialog.open_path",
                    return_value=(True, str(dest)),
                ) as opened:
                    dialog._on_open_catalog()
                opened.assert_called_once_with(dest)
            finally:
                dialog.close()
                dialog.deleteLater()
                self.app.processEvents()


class AutoMtoCompareStatusSmoke(unittest.TestCase):
    def test_labels_cover_match_gap_and_stale_cache(self) -> None:
        file = AutoMtoFile(
            title="1111",
            mark="KSB",
            spec="AGCC.287-1111-KSB.MTO-0001",
            rd_revision="02",
            relpath="1111/KSB/AGCC.287-1111-KSB.MTO-0001_02_RU.xlsx",
            row_count=1,
            fingerprint="mto",
        )
        rd_path = r"C:\rd\AGCC.287-1111-KSB.MTO-0001_03_RU.xlsx"
        self.assertEqual(
            auto_mto_compare_status().text,
            "нет в ПИ",
        )
        self.assertEqual(
            auto_mto_compare_status(files=(file,)).text,
            "нет MTO РД",
        )
        self.assertEqual(
            auto_mto_compare_status(
                files=(file,),
                rd_path=rd_path,
                rd_revision="03",
            ).kind,
            "not_compared",
        )
        rev_match = auto_mto_compare_status(
            files=(file,),
            rd_path=rd_path,
            rd_revision="02",
        )
        self.assertEqual(rev_match.kind, "rev_match")
        self.assertEqual(rev_match.text, "рев. совпала")
        self.assertTrue(rev_match.paint_match)
        matched = auto_mto_compare_status(
            files=(file,),
            rd_path=rd_path,
            rd_revision="03",
            comparison=AutoMtoCompareResult(
                match_kind="single",
                members=(file,),
                rd_rows=1,
                auto_rows=1,
                combinations_checked=0,
            ),
            comparison_rd_path=rd_path,
        )
        self.assertEqual(matched.kind, "matched")
        self.assertEqual(matched.text, f"{AUTO_MTO_EXACT_STATUS} · 02")
        self.assertTrue(matched.paint_match)
        composite = auto_mto_compare_status(
            files=(file,),
            rd_path=rd_path,
            comparison=AutoMtoCompareResult(
                match_kind="composite",
                members=(file, file),
                rd_rows=2,
                auto_rows=2,
                combinations_checked=1,
            ),
            comparison_rd_path=rd_path,
        )
        self.assertEqual(
            composite.text, f"{AUTO_MTO_EXACT_STATUS} · сумма 02 + 02"
        )
        soft = auto_mto_compare_status(
            files=(file,),
            rd_path=rd_path,
            rd_revision="03",
            comparison=AutoMtoCompareResult(
                match_kind="single",
                members=(file,),
                rd_rows=1,
                auto_rows=1,
                combinations_checked=0,
                grade="soft",
            ),
            comparison_rd_path=rd_path,
        )
        self.assertEqual(soft.kind, "soft")
        self.assertEqual(soft.text, f"{AUTO_MTO_SOFT_STATUS} · 02")
        self.assertTrue(soft.paint_match)
        self.assertFalse(needs_auto_mto_compare(soft))
        stale = auto_mto_compare_status(
            files=(file,),
            rd_path=rd_path,
            comparison=AutoMtoCompareResult(
                match_kind="single",
                members=(file,),
                rd_rows=1,
                auto_rows=1,
                combinations_checked=0,
            ),
            comparison_rd_path=r"C:\rd\other.xlsx",
            pinned=True,
        )
        self.assertEqual(stale.kind, "stale")
        self.assertEqual(stale.text, "другой файл")
        self.assertFalse(stale.paint_match)
        self.assertIn("ручным выбором", stale.tooltip)
        self.assertTrue(needs_auto_mto_compare(rev_match))
        self.assertTrue(needs_auto_mto_compare(stale))
        self.assertFalse(needs_auto_mto_compare(matched))
        dashed = rd_path.replace("-1111-", "\u20101111\u2010")
        same_file = auto_mto_compare_status(
            files=(file,),
            rd_path=rd_path,
            comparison=AutoMtoCompareResult(
                match_kind="single",
                members=(file,),
                rd_rows=1,
                auto_rows=1,
                combinations_checked=0,
            ),
            comparison_rd_path=dashed,
        )
        self.assertEqual(same_file.kind, "matched")
        self.assertFalse(needs_auto_mto_compare(same_file))
        self.assertTrue(auto_mto_rd_paths_match(rd_path, dashed))
        self.assertEqual(
            auto_mto_rd_path_key(rd_path),
            auto_mto_rd_path_key(dashed),
        )
        self.assertFalse(auto_mto_rd_paths_match(rd_path, r"C:\rd\other.xlsx"))


class OfficialRdMtoPathChangeSmoke(unittest.TestCase):
    def test_in_session_and_first_paint_vs_cache(self) -> None:
        key = kit_identity_key("7570", "SOS")
        old = r"\\bcc\eng\PrDoc\РД\7570\SOS\08_Void\AGCC.287-7570-SOS.MTO-0001_02-AN01_RU.xlsx"
        new = r"\\bcc\eng\PrDoc\РД\7570\SOS\06_рев.02-AN01\AGCC.287-7570-SOS.MTO-0001_02-AN01_RU.xlsx"
        dashed = new.replace("-7570-", "\u20107570\u2010")
        self.assertEqual(official_rd_mto_paths_changed(None, {key: new}), set())
        self.assertEqual(
            official_rd_mto_paths_changed({key: old}, {key: new}),
            {key},
        )
        self.assertEqual(
            official_rd_mto_paths_changed({key: new}, {key: dashed}),
            set(),
        )
        other = kit_identity_key("1111", "KSB")
        self.assertEqual(
            official_rd_mto_paths_changed({key: new}, {key: new, other: old}),
            {other},
        )
        self.assertEqual(
            official_rd_mto_paths_changed(
                {key: new},
                {key: new, other: old},
                include_new_keys=False,
            ),
            set(),
        )
        grouped = cached_auto_mto_rd_paths_by_kit(
            {
                "a": {"title": "7570", "mark": "SOS", "rd_path": old},
                "b": {"title": "1111", "mark": "KSB", "rd_path": old},
            }
        )
        self.assertEqual(
            kits_with_moved_cached_rd_mto_path({key: new, other: old}, grouped),
            {key},
        )
        self.assertEqual(
            kits_with_moved_cached_rd_mto_path({key: dashed}, grouped),
            {key},
        )
        grouped_new = cached_auto_mto_rd_paths_by_kit(
            {"a": {"title": "7570", "mark": "SOS", "rd_path": new}}
        )
        self.assertEqual(
            kits_with_moved_cached_rd_mto_path({key: dashed}, grouped_new),
            set(),
        )


class AutoMtoCompareCacheSmoke(unittest.TestCase):
    def test_roundtrip_skips_not_compared_and_missing_member(self) -> None:
        file = AutoMtoFile(
            title="1111",
            mark="KSB",
            spec="AGCC.287-1111-KSB.MTO-0001",
            rd_revision="02",
            relpath="1111/KSB/AGCC.287-1111-KSB.MTO-0001_02_RU.xlsx",
            row_count=1,
            fingerprint="mto",
        )
        rd_path = r"C:\rd\AGCC.287-1111-KSB.MTO-0001_02_RU.xlsx"
        result = AutoMtoCompareResult(
            match_kind="single",
            members=(file,),
            rd_rows=1,
            auto_rows=1,
            combinations_checked=0,
        )
        key = cache_key("1111", "KSB", rd_path, 123, (file,))
        entry = entry_from_result(
            result, title="1111", mark="KSB", rd_path=rd_path
        )
        self.assertIsNotNone(entry)
        skipped = entry_from_result(
            AutoMtoCompareResult(
                match_kind="not_compared",
                members=(),
                rd_rows=0,
                auto_rows=0,
                combinations_checked=0,
                error="retry",
            ),
            title="1111",
            mark="KSB",
            rd_path=rd_path,
        )
        self.assertIsNone(skipped)
        with tempfile.TemporaryDirectory(prefix="auto_mto_cache_") as temp:
            save_auto_mto_compare_cache(
                temp,
                {
                    key: entry,
                    "other": {
                        "match_kind": "not_compared",
                        "member_relpaths": [],
                        "rd_rows": 0,
                        "auto_rows": 0,
                        "combinations_checked": 0,
                        "error": "",
                    },
                },
            )
            loaded = load_auto_mto_compare_cache(temp)
        self.assertIn(key, loaded)
        self.assertNotIn("other", loaded)
        rebuilt = result_from_entry(loaded[key], (file,))
        self.assertIsNotNone(rebuilt)
        self.assertEqual(rebuilt.match_kind, "single")
        self.assertEqual(rebuilt.grade, "exact")
        self.assertEqual(rebuilt.members, (file,))
        self.assertIsNone(result_from_entry(loaded[key], ()))
        self.assertIsNone(
            result_from_entry(
                {"match_kind": "not_compared", "member_relpaths": []},
                (file,),
            )
        )


if __name__ == "__main__":
    unittest.main()
