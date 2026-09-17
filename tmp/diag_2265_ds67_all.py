"""All 2265-KSB DС67 RFP qty by code in yesterday Step4."""
from __future__ import annotations

import sys
from collections import defaultdict

from openpyxl import load_workbook

XLSX = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.09.02.13.49"
    r"\Шаг4_Сопоставление_RFP_MTO_20260902_140222.xlsx"
)


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def _norm(v: object) -> str:
    return "" if v is None else str(v).strip()


def _qty(v: object) -> float:
    if v in (None, ""):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def main() -> None:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [_norm(v) for v in next(rows)]
    idx = {n: i for i, n in enumerate(header) if n}

    def cell(row, name):
        i = idx.get(name)
        if i is None or i >= len(row):
            return None
        return row[i]

    by_code: dict[str, float] = defaultdict(float)
    by_pos: dict[str, float] = defaultdict(float)
    n = 0
    for row in rows:
        if _norm(cell(row, "Титул/Марка")) != "2265-KSB":
            continue
        if _norm(cell(row, "Порядковый ДС")) != "ДС67":
            continue
        n += 1
        code = _norm(cell(row, "Код RFP")) or "(empty RFP code)"
        pos = _norm(cell(row, "№ позиции")) or "(no pos)"
        q = _qty(cell(row, "Кол-во RFP"))
        by_code[code] += q
        by_pos[f"{pos}|{code}"] += q
        if pos in {"25206", "25207", "146", "138", "139"} or q in {2.0, 3.0}:
            _safe_print(
                f"pos={pos!r} code={code!r} qty={q:g} tags={_norm(cell(row, 'Теги RFP'))!r} "
                f"mto={_norm(cell(row, 'Код MTO'))!r} mto_qty={cell(row, 'Количество MTO')!r} "
                f"status={_norm(cell(row, 'MTO, Статус позиции'))!r} "
                f"ul={_norm(cell(row, 'Статус УЛ'))!r}"
            )
    _safe_print(f"ds67_rows={n}")
    _safe_print("by_code:")
    for k, v in sorted(by_code.items(), key=lambda kv: -kv[1]):
        _safe_print(f"  {k} = {v:g}")
    _safe_print(f"total_rfp_ds67={sum(by_code.values()):g}")
    _safe_print("pos 25206/25207:")
    for k, v in sorted(by_pos.items()):
        if k.startswith("25206") or k.startswith("25207"):
            _safe_print(f"  {k} = {v:g}")
    wb.close()


if __name__ == "__main__":
    main()
