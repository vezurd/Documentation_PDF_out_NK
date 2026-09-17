"""Smoke tests for 1C UPD upload loader (compact/wide layouts, split sheets)."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.ds_compare.upd_load import (
    LAYOUT_COMPACT,
    LAYOUT_WIDE,
    build_upd_header,
    classify_upd_row,
    detect_upd_header,
    is_primary_upd_sheet,
    load_and_merge_upd,
)
from base.base_classes import RowType


def _compact_headers() -> tuple[list[str], list[str]]:
    group = [""] * 24
    group[0] = "Нпп ***"
    group[1] = "Позиции спецификации"
    group[5] = "Этап по Договору"
    group[6] = "Товарная позиция"
    group[11] = "Себестоимость остатков на складе (RUB)"
    group[15] = "в вал.сделки"
    detail = [""] * 24
    detail[1] = "Позиция"
    detail[2] = "Кол-во (остаток)"
    detail[3] = "Ед.изм."
    detail[4] = "Стоимость (остаток) - в валюте специф. ****"
    detail[6] = "Код / производителя"
    detail[7] = "Наименование"
    detail[8] = "Код 1С"
    detail[9] = "Ед.изм."
    detail[10] = "Остаток *"
    detail[11] = "Цена  (без НДС)"
    detail[12] = "Сумма (без НДС) *"
    detail[13] = "Цена БУ (без НДС)"
    detail[14] = "Сумма БУ (без НДС) *"
    detail[15] = "Сумма продажи (с НДС) **"
    detail[16] = "Раздел (код)"
    detail[17] = "Раздел"
    detail[18] = "Номер строки БТП"
    detail[19] = "CRM_GUID (строка спецификации проектировщика)"
    detail[20] = "Валюта спецификации"
    return group, detail


def _wide_headers() -> tuple[list[str], list[str]]:
    group = [""] * 30
    group[1] = "Структура договорной спецификации"
    group[6] = "Позиции спецификации по договору"
    group[11] = "Закупленные товарные позиции"
    group[16] = "Себестоимость остатков на складе (RUB)"
    group[20] = "в вал.сделки"
    detail = [""] * 30
    detail[0] = "Нпп ***"
    detail[1] = "Уровень 1"
    detail[2] = "Уровень 2"
    detail[3] = "Уровень 3"
    detail[4] = "Уровень 4"
    detail[5] = "Уровень 5"
    detail[6] = "Позиция"
    detail[7] = "Кол-во (остаток)"
    detail[8] = "Ед.изм."
    detail[9] = "Стоимость (остаток) - в валюте специф. ****"
    detail[10] = "Этап по Договору"
    detail[11] = "Код / производителя"
    detail[12] = "Наименование"
    detail[13] = "Код 1С"
    detail[14] = "Ед.изм."
    detail[15] = "Остаток *"
    detail[16] = "Цена  (без НДС)"
    detail[17] = "Сумма (без НДС) *"
    detail[18] = "Цена БУ (без НДС)"
    detail[19] = "Сумма БУ (без НДС) *"
    detail[20] = "Сумма продажи (с НДС) **"
    detail[21] = "Раздел (код)"
    detail[22] = "Раздел"
    detail[23] = "Номер строки БТП"
    detail[24] = "CRM_GUID (строка спецификации проектировщика)"
    detail[25] = "Валюта спецификации"
    detail[29] = "Id строки спецификации"
    return group, detail


def _write_row(ws, row_idx: int, values: list[object]) -> None:
    for col, value in enumerate(values, start=1):
        if value not in (None, ""):
            ws.cell(row_idx, col, value)


class UpdHeaderSmokeTest(unittest.TestCase):
    def test_primary_sheet_names(self) -> None:
        self.assertTrue(is_primary_upd_sheet("Лист1"))
        self.assertTrue(is_primary_upd_sheet("Лист_1"))
        self.assertTrue(is_primary_upd_sheet("Лист_1 (2)"))
        self.assertTrue(is_primary_upd_sheet("общая"))
        self.assertTrue(is_primary_upd_sheet("общий"))
        self.assertFalse(is_primary_upd_sheet("6600"))
        self.assertFalse(is_primary_upd_sheet("Лист2"))
        self.assertFalse(is_primary_upd_sheet("спец40"))

    def test_compact_header_maps_npp_and_two_units(self) -> None:
        group, detail = _compact_headers()
        header = build_upd_header(detail, group, detail_row=4)
        self.assertIsNotNone(header)
        assert header is not None
        self.assertEqual(header.layout, LAYOUT_COMPACT)
        self.assertEqual(header.column_keys[0], "npp")
        self.assertEqual(header.column_keys[1], "spec_position")
        self.assertEqual(header.column_keys[3], "spec_units")
        self.assertEqual(header.column_keys[5], "contract_stage")
        self.assertEqual(header.column_keys[7], "name")
        self.assertEqual(header.column_keys[8], "code_1c")
        self.assertEqual(header.column_keys[9], "goods_units")

    def test_wide_header_has_levels(self) -> None:
        group, detail = _wide_headers()
        header = build_upd_header(detail, group, detail_row=4)
        self.assertIsNotNone(header)
        assert header is not None
        self.assertEqual(header.layout, LAYOUT_WIDE)
        self.assertEqual(header.column_keys[1], "level_1")
        self.assertEqual(header.column_keys[6], "spec_position")
        self.assertEqual(header.column_keys[12], "name")
        self.assertEqual(header.column_keys[29], "spec_line_id")

    def test_detect_header_skips_banner_rows(self) -> None:
        group, detail = _compact_headers()
        preview = [
            ["", "Остатки (без ТолькоУУ ...)"],
            ["", '(форма отчета имеет нужный формат ...)'],
            group,
            detail,
        ]
        header = detect_upd_header(preview)
        self.assertIsNotNone(header)
        assert header is not None
        self.assertEqual(header.detail_row, 4)
        self.assertEqual(header.layout, LAYOUT_COMPACT)

    def test_classify_position_and_empty(self) -> None:
        self.assertEqual(
            classify_upd_row({"spec_position": "Скоба", "spec_qty": 2}),
            RowType.position_row,
        )
        self.assertEqual(classify_upd_row({"name": "", "code_1c": ""}), RowType.empty_row)
        self.assertEqual(
            classify_upd_row(
                {
                    "name": "Наименование",
                    "code_1c": "Код 1С",
                    "spec_position": "Позиция",
                }
            ),
            RowType.other_row,
        )


class UpdMergeSmokeTest(unittest.TestCase):
    def _make_compact_book(self, path: Path, *, extra_split: bool) -> None:
        group, detail = _compact_headers()
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Лист1"
        ws["B1"] = "Остатки (без ТолькоУУ и товаров на отв.хранении)"
        ws["B2"] = '(форма отчета имеет нужный формат для закачки ... "РеализацияТоваров2_0")'
        _write_row(ws, 3, group)
        _write_row(ws, 4, detail)
        data = [""] * 21
        data[0] = "7412-38"
        data[1] = "Металлорукав"
        data[2] = 400
        data[3] = "Метр"
        data[4] = 94958.4
        data[5] = "спецификация 8"
        data[6] = "zeta42313"
        data[7] = "Металлорукав МРПИ"
        data[8] = "00279319"
        data[9] = "Метр"
        data[10] = 400
        data[11] = 81.9
        data[18] = "241111-902-004-0004"
        _write_row(ws, 5, data)
        if extra_split:
            split = wb.create_sheet("6600")
            _write_row(split, 3, group)
            _write_row(split, 4, detail)
            dup = list(data)
            dup[0] = "6600-99"
            dup[1] = "Только на вкладке 6600"
            _write_row(split, 5, dup)
        wb.save(path)
        wb.close()

    def _make_wide_book(self, path: Path) -> None:
        group, detail = _wide_headers()
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Лист_1"
        ws["G1"] = "Остатки (без ТолькоУУ ...)"
        ws["G2"] = "(форма отчета имеет нужный формат ...)"
        _write_row(ws, 3, group)
        _write_row(ws, 4, detail)
        data = [""] * 30
        data[0] = "2"
        data[1] = "Доп. соглашение № 2"
        data[6] = "Секция двусторонняя"
        data[7] = 3
        data[8] = "Штука"
        data[11] = "PERCo-STD-01"
        data[12] = "Секция турникета"
        data[13] = "00303760"
        data[14] = "Штука"
        data[15] = 3
        data[29] = "70c6f798-9847-4e5a-88ee-77c34b43f4fb"
        _write_row(ws, 5, data)
        wb.save(path)
        wb.close()

    def _make_invoice_book(self, path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Лист_1"
        ws["B1"] = "Универсальный передаточный документ"
        ws["I1"] = "Счет-фактура №"
        ws["A5"] = "№"
        ws["B5"] = "Наименование товара"
        wb.save(path)
        wb.close()

    def test_merge_compact_wide_skips_split_and_invoice(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            root = Path(tmp) / "src"
            out = Path(tmp) / "out"
            ds13 = root / "ДС 13"
            ds15 = root / "ДС 15"
            ds13.mkdir(parents=True)
            ds15.mkdir(parents=True)
            self._make_compact_book(ds13 / "PL_compact.xlsx", extra_split=True)
            self._make_wide_book(ds15 / "PL_wide.xlsx")
            self._make_invoice_book(ds15 / "invoice.xlsx")
            result = load_and_merge_upd(root, output_dir=out, workers=1)
            self.assertTrue(result.success)
            self.assertEqual(result.stats.position_rows, 2)
            self.assertEqual(result.stats.files_ok, 2)
            self.assertEqual(result.stats.files_format_mismatch, 1)
            self.assertEqual(result.stats.layout_compact, 1)
            self.assertEqual(result.stats.layout_wide, 1)
            self.assertGreaterEqual(result.stats.sheets_skipped, 1)
            self.assertTrue(Path(result.summary_path).is_file())
            loaded_sheets = {
                hit.sheet_name
                for hit in result.stats.sheet_hits
                if hit.status == "loaded"
            }
            self.assertIn("Лист1", loaded_sheets)
            self.assertIn("Лист_1", loaded_sheets)
            self.assertNotIn("6600", loaded_sheets)
            from openpyxl import load_workbook

            wb = load_workbook(result.summary_path, read_only=True, data_only=True)
            try:
                ws = wb.active
                assert ws is not None
                rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
            finally:
                wb.close()
            headers = list(rows[0])
            self.assertIn("Имя файла", headers)
            self.assertIn("Путь к файлу", headers)
            self.assertIn("Наименование", headers)
            self.assertIn("Код 1С", headers)
            names = {row[headers.index("Имя файла")] for row in rows[1:]}
            self.assertEqual(names, {"PL_compact.xlsx", "PL_wide.xlsx"})
            codes = {row[headers.index("Код 1С")] for row in rows[1:]}
            self.assertEqual(codes, {"00279319", "00303760"})


if __name__ == "__main__":
    unittest.main(verbosity=2)
