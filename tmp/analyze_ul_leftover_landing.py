"""Classify leftover «Только в УЛ» keys vs Step4 result (read-only)."""

from __future__ import annotations

import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from RFQ.packing_list_provider import normalize_packing_key_part, packing_match_key
from RFQ.tags_rfp_compare.step4.step4_packing_compare import parse_composite_title

XLSX = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.08.25.11.56"
    r"\Шаг4_Сопоставление_RFP_MTO_20260825_120001.xlsx"
)

STATUS_UL = "Только в УЛ"
STATUS_OPEN = "Незакрыто по УПД"
STATUS_RFP = "Только в RFP (лишняя)"
STATUS_MTO = "Только в МТО"
STATUS_VO = "Только в РКД"
STATUS_SHORT = "Недопоставка"
STATUS_OK = "Поставка комплектна"

MISS = {STATUS_OPEN, STATUS_RFP, STATUS_MTO, STATUS_VO, STATUS_SHORT}


def _num(value: object) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return 0.0


def _title_parts(display: object) -> tuple[str, str]:
    text = str(display or "").strip()
    parsed = parse_composite_title(text)
    if parsed is not None:
        return parsed
    if "-" in text:
        left, right = text.rsplit("-", 1)
        return left, right
    return text, ""


def main() -> None:
    print("loading", XLSX)
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Сопоставление RFP и MTO"]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows)]
    idx = {name: i for i, name in enumerate(header) if name}

    def cell(row: tuple, name: str) -> object:
        i = idx.get(name)
        return None if i is None or i >= len(row) else row[i]

    leftover: list[dict] = []
    result_by_key: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    result_by_code: dict[str, list[dict]] = defaultdict(list)
    status_qty = Counter()
    leftover_by_title = Counter()
    leftover_qty_by_title = Counter()

    for row in rows:
        status = str(cell(row, "Статус УЛ") or "").strip()
        if not status:
            continue
        title_disp = str(cell(row, "Титул/Марка") or "").strip()
        title, system = _title_parts(title_disp)
        code_rfp = normalize_packing_key_part(cell(row, "Код RFP"))
        code_mto = normalize_packing_key_part(cell(row, "Код MTO"))
        code_ul = normalize_packing_key_part(cell(row, "Код УЛ"))
        qty_ul = _num(cell(row, "Кол-во по УЛ"))
        qty_rfp = _num(cell(row, "Кол-во RFP"))
        qty_mto = _num(cell(row, "Количество MTO"))
        rec = {
            "title": title_disp,
            "t": title,
            "s": system,
            "status": status,
            "code_rfp": code_rfp,
            "code_mto": code_mto,
            "code_ul": code_ul,
            "qty_ul": qty_ul,
            "qty_rfp": qty_rfp,
            "qty_mto": qty_mto,
            "name_ul": str(cell(row, "Наименование УЛ") or "")[:80],
            "name_rfp": str(cell(row, "Наименование RFP") or "")[:80],
            "src": str(cell(row, "Источник УЛ (файл · вкладка · строка)") or "")[:120],
        }
        status_qty[status] += qty_ul if status == STATUS_UL else 1
        if status == STATUS_UL:
            leftover.append(rec)
            leftover_by_title[title_disp or "<пусто>"] += 1
            leftover_qty_by_title[title_disp or "<пусто>"] += qty_ul
            continue
        for code in (code_rfp, code_mto, code_ul):
            if not code:
                continue
            key = packing_match_key(title, system, code)
            if all(key):
                result_by_key[key].append(rec)
            result_by_code[code].append(rec)
    wb.close()

    print("\n=== leftover rows", len(leftover), "qty", round(sum(r["qty_ul"] for r in leftover), 3))
    print("top titles by leftover rows:")
    for title, n in leftover_by_title.most_common(15):
        print(f"  {n:5d} rows  qty={leftover_qty_by_title[title]:g}  {title}")

    buckets = Counter()
    bucket_qty = Counter()
    samples: dict[str, list[str]] = defaultdict(list)
    title_mismatch_pairs = Counter()
    miss_status_when_code_hit = Counter()

    for rec in leftover:
        code = rec["code_ul"]
        key = packing_match_key(rec["t"], rec["s"], code) if code else ("", "", "")
        same_key = result_by_key.get(key, []) if all(key) else []
        same_code = result_by_code.get(code, []) if code else []
        qty = rec["qty_ul"]
        if not code:
            bucket = "no_ul_code"
        elif same_key:
            statuses = Counter(r["status"] for r in same_key)
            if any(r["status"] == STATUS_OK for r in same_key):
                bucket = "same_key_already_complete"  # true extra UL / overdelivery
            elif any(r["status"] in MISS for r in same_key):
                bucket = "same_key_still_miss"  # queue leftover + miss on same key
                miss_status_when_code_hit.update(statuses)
            else:
                bucket = "same_key_other"
        elif same_code:
            other_titles = sorted({r["title"] for r in same_code if r["title"] != rec["title"]})
            bucket = "same_code_other_title"
            for other in other_titles[:3]:
                title_mismatch_pairs[(rec["title"] or "?", other or "?")] += 1
            miss_status_when_code_hit.update(r["status"] for r in same_code)
        else:
            bucket = "code_absent_in_result"
        buckets[bucket] += 1
        bucket_qty[bucket] += qty
        if len(samples[bucket]) < 8:
            samples[bucket].append(
                f"{rec['title']!r} code={code!r} qty={qty:g} "
                f"name={rec['name_ul']!r} src={rec['src']!r}"
            )

    print("\n=== leftover classification (rows / qty) ===")
    for bucket, n in buckets.most_common():
        print(f"  {n:5d}  qty={bucket_qty[bucket]:g}  {bucket}")

    print("\n=== samples ===")
    for bucket, lines in samples.items():
        print(f"\n[{bucket}]")
        for line in lines:
            print(" ", line)

    print("\n=== title mismatch pairs (leftover title -> result title), top 20 ===")
    for pair, n in title_mismatch_pairs.most_common(20):
        print(f"  {n:5d}  {pair[0]!r}  ->  {pair[1]!r}")

    print("\n=== statuses of result rows sharing leftover code (any title) ===")
    for status, n in miss_status_when_code_hit.most_common():
        print(f"  {n:5d}  {status}")

    leftover_codes = {r["code_ul"] for r in leftover if r["code_ul"]}
    miss_codes = set()
    # Re-open is expensive; miss codes already in result_by_code via non-UL rows.
    print("\nunique leftover codes", len(leftover_codes))


if __name__ == "__main__":
    main()
