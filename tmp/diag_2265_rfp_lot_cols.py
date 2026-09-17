"""Print Lot vs RD qty for DС67 2265-KSB BCC0000642."""
from __future__ import annotations

import sys

from openpyxl import load_workbook

RFP = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_RFP\RFP_Зиновьев\ДС67. AGCC.287-0000-12.4.1-RFP-0032_0_RU.xlsx"
)


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def main() -> None:
    wb = load_workbook(RFP, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = list(row)
        blob = " ".join(str(v) for v in cells if v is not None)
        if "BCC0000642" not in blob.upper():
            continue
        title = str(cells[1] if len(cells) > 1 else "")
        _safe_print(f"\nexcel_row={i} title={title!r}")
        for idx, val in enumerate(cells):
            if val is None or val == "":
                continue
            _safe_print(f"  col[{idx}]={val!r}")
    wb.close()


if __name__ == "__main__":
    main()
