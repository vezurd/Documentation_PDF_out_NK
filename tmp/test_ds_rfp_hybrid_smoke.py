"""Offline smoke for DS↔RFP group reconciliation and atomic hybrid overlay."""

from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from collections import Counter
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import base.t_comm_initial_classes as t_com_init_cls
from RFQ.rfp_parts.ds_baseline import IdentityDsUnitsConverter, build_ds_baseline
from RFQ.rfp_parts.ds_registry import (
    MODE_NEEDS_SPLIT,
    MODE_NO_UL,
    MODE_WHOLE,
    REGISTRY_SHEET_NAME,
    DsRegistryRelation,
    DsRegistryRow,
    STATUS_ACTIVE,
    load_registry,
    parse_registry_ds_number,
    write_registry_workbook,
)
from RFQ.rfp_parts.ds_rfp_hybrid import (
    ALGORITHM_VERSION,
    HYBRID_REPORT_PREFIX,
    HYBRID_XLSX_NAME,
    ISSUE_AMBIGUOUS_MAPPING,
    ISSUE_DUPLICATE_RFP_KEY,
    ISSUE_UNPARSED_RFP,
    SOURCE_DS,
    SOURCE_RFP,
    STATUS_BLOCKED,
    STATUS_DS_ONLY,
    STATUS_MATCH,
    STATUS_MISMATCH,
    STATUS_RFP_ONLY,
    IdentityRfpUnitsConverter,
    build_ds_rfp_hybrid,
    collect_rfp_workbooks,
)
from RFQ.rfp_parts.ds_rfp_tag_placement import MIX_MIXED, MIX_SEPARATE, mixed_ds_label
from RFQ.tags_rfp_compare.rfp_supply_status import CANONICAL_EXCLUDED_FROM_SUPPLY
from RFQ.units_convert.models import GoogleUnitsIndex
from base.base_classes import RowType, TableComments
from base.base_mto import get_std_from_excel_file
from base.tables_columns import (
    CODE,
    DS_NAME,
    RFP_SUPPLY_STATUS,
    TAGS,
    UNITS,
    VALUES,
)

DS_HEADER = [
    "№ п/п",
    "Титул",
    "Раздел",
    "Спецификация",
    "RFQ",
    "Наименование Позиций Товара по РД",
    "Код 1С СОУ",
    "Код РД",
    "Наименование Позиций Товара Поставщика",
    "Технические требования (ГОСТ/ ТУ и др.)",
    "Ед. изм.",
    "Кол-во",
]

STAMP = "20260102_030405"
TAG_A = "8529-SS-01-S-UZ-0711"
TAG_B = "8529-SS-01-S-UZ-0712"
TAG_61 = "1600-XX-01-S-UZ-0001"
CLUSTER_RFP_NAME = "ДС15_61. AGCC.xlsx"
CODE_MATCH = "BCC0000001"
CODE_ONLY = "BCC0000002"
CODE_BLOCK = "BCC0000003"
CODE_NOUL = "BCC0000004"
CODE_MM = "BCC0000005"
CODE_MIX = "BCC0000009"
CODE_UNIT = "BCC0000006"
CODE_LIVE = "BCC0000007"
CODE_EXCL = "BCC0000008"


def _load_xlsx(path: Path):
    return load_workbook(BytesIO(path.read_bytes()))


def _ds_row(
    *,
    npp: object = 1,
    title: str = "8529",
    system: str = "SOS",
    spec: str = "spec-1",
    rfq: str = "RFQ-1",
    name: str = "Кабель",
    code_1c: str = "1C-1",
    code: str = CODE_MATCH,
    supplier: str = "Vendor",
    type_mark: str = "NYM",
    units: str = "шт",
    qty: object = 1,
) -> list[object]:
    return [
        npp,
        title,
        system,
        spec,
        rfq,
        name,
        code_1c,
        code,
        supplier,
        type_mark,
        units,
        qty,
    ]


def _write_xlsx(path: Path, rows: list[list[object]], *, sheet: str = "Спека") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def _rfp_header(*, status: bool = False) -> list[object]:
    row = [""] * 20
    row[0] = "№ п/п"
    row[1] = "Титул"
    row[2] = "Спецификация"
    row[3] = "Линия, Tag-номер"
    row[4] = "Код 1С"
    row[5] = "Код РД"
    row[6] = "Наименование МТР"
    row[7] = "Технические характеристики"
    row[16] = "Кол-во"
    row[17] = "Ед. изм"
    if status:
        row[19] = "MTO, Статус позиции"
    return row


def _rfp_data(
    *,
    number: int,
    code: str,
    qty: object,
    units: str = "шт",
    tag: str = "",
    status: object = "",
    title: str = "8529-SOS",
    name: str = "Кабель",
) -> list[object]:
    row = [""] * 20
    row[0] = number
    row[1] = title
    row[2] = "AGCC.287-8529-SOS.MTO-0001"
    row[3] = tag
    row[5] = code
    row[6] = name
    row[7] = "NYM"
    row[16] = qty
    row[17] = units
    row[19] = status
    return row


def _write_rfp(
    path: Path,
    rows: list[list[object]],
    *,
    status: bool = False,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Перечень материалов"
    ws.append(["титульный"])
    ws.append(_rfp_header(status=status))
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def _rel(
    *,
    group_id: str,
    rfp_key: str,
    ul_folder: str,
    mode: str,
    block_index: int = 1,
    rfp_file: str = "",
) -> DsRegistryRelation:
    return DsRegistryRelation(
        group_id=group_id,
        rfp_key=rfp_key,
        ul_folder=ul_folder,
        mode=mode,
        block_index=block_index,
        rfp_file=rfp_file,
    )


def _write_registry(path: Path) -> None:
    rows = [
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="13",
            relations=(
                _rel(
                    group_id="ДС13",
                    rfp_key="13",
                    ul_folder="согл УЛ ДС13",
                    mode=MODE_WHOLE,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="47",
            relations=(
                _rel(
                    group_id="ДС47",
                    rfp_key="47",
                    ul_folder="согл УЛ ДС47",
                    mode=MODE_WHOLE,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="88",
            relations=(
                _rel(
                    group_id="ДС88",
                    rfp_key="88",
                    ul_folder="согл УЛ ДС88",
                    mode=MODE_WHOLE,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="8",
            relations=(
                _rel(
                    group_id="ДС8",
                    rfp_key="8",
                    ul_folder="согл УЛ ДС8",
                    mode=MODE_NEEDS_SPLIT,
                    block_index=1,
                ),
                _rel(
                    group_id="ДС81",
                    rfp_key="81",
                    ul_folder="согл УЛ ДС81",
                    mode=MODE_NEEDS_SPLIT,
                    block_index=2,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="101",
            relations=(
                _rel(
                    group_id="NO_UL:101",
                    rfp_key="101",
                    ul_folder="",
                    mode=MODE_NO_UL,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="99",
            relations=(
                _rel(
                    group_id="ДС99",
                    rfp_key="99",
                    ul_folder="согл УЛ ДС99",
                    mode=MODE_WHOLE,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="55",
            relations=(
                _rel(
                    group_id="ДС55",
                    rfp_key="55",
                    ul_folder="согл УЛ ДС55",
                    mode=MODE_WHOLE,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="66",
            relations=(
                _rel(
                    group_id="ДС66",
                    rfp_key="66",
                    ul_folder="согл УЛ ДС66",
                    mode=MODE_WHOLE,
                ),
            ),
        ),
    ]
    write_registry_workbook(path, rows)


def _google_index() -> GoogleUnitsIndex:
    codes = (
        CODE_MATCH,
        CODE_ONLY,
        CODE_BLOCK,
        CODE_NOUL,
        CODE_MM,
        CODE_MIX,
        CODE_UNIT,
        CODE_LIVE,
        CODE_EXCL,
    )
    return GoogleUnitsIndex(
        units_by_code={code: "шт" for code in codes},
        display_units_by_code={code: "шт" for code in codes},
        display_code_by_code={code: code for code in codes},
    )


def _write_ds_tree(ds_root: Path) -> None:
    _write_xlsx(ds_root / "ДС13.xlsx", [DS_HEADER, _ds_row(npp=1, qty=2)])
    _write_xlsx(
        ds_root / "ДС_47" / "spec.xlsx",
        [DS_HEADER, _ds_row(npp=2, qty=1)],
    )
    _write_xlsx(
        ds_root / "ДС88.xlsx",
        [DS_HEADER, _ds_row(npp=3, code=CODE_ONLY, qty=4, title="1600", system="POS")],
    )
    _write_xlsx(
        ds_root / "ДС8.xlsx",
        [DS_HEADER, _ds_row(npp=4, code=CODE_BLOCK, qty=5, title="8000", system="AAA")],
    )
    _write_xlsx(
        ds_root / "ДС101.xlsx",
        [DS_HEADER, _ds_row(npp=5, code=CODE_NOUL, qty=1, title="1010", system="NO")],
    )
    _write_xlsx(
        ds_root / "ДС99.xlsx",
        [
            DS_HEADER,
            _ds_row(npp=6, code=CODE_MM, qty=2, title="9900", system="MM"),
            _ds_row(npp=7, code=CODE_MIX, qty=1, title="9900", system="MM"),
        ],
    )
    _write_xlsx(
        ds_root / "ДС55.xlsx",
        [DS_HEADER, _ds_row(npp=8, code=CODE_UNIT, qty=2, title="5500", system="UN")],
    )
    _write_xlsx(
        ds_root / "ДС66.xlsx",
        [DS_HEADER, _ds_row(npp=9, code=CODE_LIVE, qty=2, title="6600", system="EX")],
    )


def _write_rfp_tree(rfp_root: Path, *, duplicate: bool = False) -> None:
    _write_rfp(
        rfp_root / "ДС13. AGCC.xlsx",
        [
            _rfp_data(
                number=1,
                code=CODE_MATCH,
                qty=2,
                tag=f"{TAG_A}, {TAG_B}",
            )
        ],
    )
    _write_rfp(
        rfp_root / "ДС99. AGCC.xlsx",
        [_rfp_data(number=1, code=CODE_MM, qty=9, tag="9900-XX-01-S-UZ-0001")],
    )
    _write_rfp(
        rfp_root / "ДС55. AGCC.xlsx",
        [_rfp_data(number=1, code=CODE_UNIT, qty=2, units="кг")],
    )
    _write_rfp(
        rfp_root / "ДС66. AGCC.xlsx",
        [
            _rfp_data(number=1, code=CODE_LIVE, qty=2),
            _rfp_data(
                number=2,
                code=CODE_EXCL,
                qty=10,
                status=CANONICAL_EXCLUDED_FROM_SUPPLY,
            ),
        ],
        status=True,
    )
    _write_rfp(
        rfp_root / "ДС8. AGCC.xlsx",
        [_rfp_data(number=1, code=CODE_BLOCK, qty=5)],
    )
    _write_rfp(
        rfp_root / "ДС77. AGCC.xlsx",
        [_rfp_data(number=1, code="BCC9999999", qty=3)],
    )
    nested = rfp_root / "nested"
    _write_rfp(
        nested / "ДС88. AGCC.xlsx",
        [_rfp_data(number=1, code=CODE_ONLY, qty=4)],
    )
    (rfp_root / "~$lock.xlsx").write_bytes(b"PK\x03\x04lock")
    _write_xlsx(rfp_root / "rfp_parts_net.xlsx", [["net"]])
    _write_xlsx(rfp_root / HYBRID_XLSX_NAME, [["hybrid"]])
    if duplicate:
        shutil.copyfile(rfp_root / "ДС13. AGCC.xlsx", rfp_root / "ДС13. copy.xlsx")


def _run(
    root: Path,
    *,
    duplicate: bool = False,
    extra_ds: list[tuple[str, list[list[object]]]] | None = None,
):
    ds_root = root / "ds"
    rfp_root = root / "rfp"
    out_dir = root / "out"
    registry_path = root / "registry.xlsx"
    _write_registry(registry_path)
    _write_ds_tree(ds_root)
    if extra_ds:
        for name, rows in extra_ds:
            _write_xlsx(ds_root / name, rows)
    _write_rfp_tree(rfp_root, duplicate=duplicate)
    registry = load_registry(registry_path)
    baseline = build_ds_baseline(
        ds_root,
        registry,
        out_dir,
        write_baseline=True,
        converter=IdentityDsUnitsConverter(),
        google_index=_google_index(),
        stamp=STAMP,
    )
    hybrid = build_ds_rfp_hybrid(
        baseline,
        registry,
        rfp_root,
        out_dir,
        write_hybrid=True,
        converter=IdentityRfpUnitsConverter(),
        stamp=STAMP,
    )
    return baseline, hybrid, rfp_root, out_dir


def _qty_sum(rows, *, ds_name: str | None = None, code: str | None = None) -> Decimal:
    total = Decimal("0")
    for item in rows:
        if ds_name is not None and item.record.ds_name != ds_name:
            continue
        if code is not None and item.record.code != code:
            continue
        total += item.record.values
    return total


def _write_cluster_registry(
    path: Path,
    *,
    ul15: str = "согл УЛ ДС15",
    ul61: str = "согл УЛ ДС61",
    ds15_file: str = "",
    ds61_file: str = "",
) -> None:
    rows = [
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="15",
            ds_file=ds15_file,
            relations=(
                _rel(
                    group_id="ДС15",
                    rfp_key="15",
                    ul_folder=ul15,
                    mode=MODE_WHOLE,
                    rfp_file=CLUSTER_RFP_NAME,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="61",
            ds_file=ds61_file,
            relations=(
                _rel(
                    group_id="ДС61",
                    rfp_key="61",
                    ul_folder=ul61,
                    mode=MODE_WHOLE,
                    rfp_file=CLUSTER_RFP_NAME,
                ),
            ),
        ),
    ]
    write_registry_workbook(path, rows)


def _run_cluster(
    root: Path,
    *,
    mix_mode: str,
    ds15: list[list[object]],
    ds61: list[list[object]],
    rfp_rows: list[list[object]],
    rfp_status: bool = False,
    registry_kw: dict[str, str] | None = None,
    source_bytes: list[bytes] | None = None,
):
    ds_root = root / "ds"
    rfp_root = root / "rfp"
    out_dir = root / "out"
    registry_path = root / "registry.xlsx"
    _write_cluster_registry(registry_path, **(registry_kw or {}))
    if source_bytes is not None:
        source_bytes.append(registry_path.read_bytes())
    _write_xlsx(ds_root / "ДС15.xlsx", [DS_HEADER, *ds15])
    _write_xlsx(ds_root / "ДС61.xlsx", [DS_HEADER, *ds61])
    _write_rfp(rfp_root / CLUSTER_RFP_NAME, rfp_rows, status=rfp_status)
    registry = load_registry(registry_path)
    baseline = build_ds_baseline(
        ds_root,
        registry,
        out_dir,
        write_baseline=True,
        converter=IdentityDsUnitsConverter(),
        google_index=_google_index(),
        stamp=STAMP,
    )
    hybrid = build_ds_rfp_hybrid(
        baseline,
        registry,
        rfp_root,
        out_dir,
        write_hybrid=True,
        converter=IdentityRfpUnitsConverter(),
        stamp=STAMP,
        mix_mode=mix_mode,
    )
    return baseline, hybrid


class DsRfpHybridCollectSmokeTest(unittest.TestCase):
    def test_root_only_skips_locks_outputs_and_nested(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            rfp_root = root / "rfp"
            _write_rfp_tree(rfp_root)
            files, skipped = collect_rfp_workbooks(rfp_root)
            names = {item.relpath for item in files}
            self.assertIn("ДС13. AGCC.xlsx", names)
            self.assertNotIn("rfp_parts_net.xlsx", names)
            self.assertNotIn(HYBRID_XLSX_NAME, names)
            self.assertTrue(all("/" not in item.relpath for item in files))
            nested = rfp_root / "nested" / "ДС88. AGCC.xlsx"
            self.assertTrue(nested.is_file())
            self.assertNotIn("ДС88. AGCC.xlsx", names)
            reasons = {item.reason for item in skipped}
            self.assertIn("excel_lock", reasons)
            self.assertIn("own_output", reasons)


class DsRfpHybridSmokeTest(unittest.TestCase):
    def test_group_overlay_matrix_tags_and_no_mixing(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            baseline, hybrid, rfp_root, out_dir = _run(Path(raw))
            self.assertFalse(baseline.blocking, baseline.summary_line())
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            self.assertIsNotNone(hybrid.hybrid_path)
            self.assertEqual(hybrid.hybrid_path.name, HYBRID_XLSX_NAME)
            self.assertTrue(hybrid.hybrid_path.is_file())
            self.assertTrue(
                hybrid.report_path.name.startswith(HYBRID_REPORT_PREFIX)
            )
            self.assertTrue(hybrid.report_path.name.endswith(".xlsx"))
            self.assertTrue(hybrid.report_path.is_file())
            self.assertEqual(hybrid.algorithm_version, ALGORITHM_VERSION)
            self.assertTrue(hybrid.fingerprint)
            leftovers = list(out_dir.glob(".*.tmp.xlsx"))
            self.assertEqual(leftovers, [])

            by_id = {item.group_id: item for item in hybrid.groups}
            self.assertEqual(by_id["ДС13"].status, STATUS_MATCH)
            self.assertEqual(by_id["ДС13"].selected_source, "RFP")
            self.assertEqual(by_id["ДС99"].status, STATUS_MISMATCH)
            self.assertEqual(by_id["ДС99"].selected_source, "ДС")
            self.assertEqual(by_id["ДС88"].status, STATUS_DS_ONLY)
            self.assertEqual(by_id["ДС88"].selected_source, "ДС")
            self.assertEqual(by_id["ДС101"].status, STATUS_DS_ONLY)
            self.assertEqual(by_id["ДС55"].status, STATUS_MISMATCH)
            self.assertEqual(by_id["ДС66"].status, STATUS_MATCH)
            self.assertEqual(by_id["ДС8"].status, STATUS_MATCH)
            self.assertEqual(by_id["ДС8"].selected_source, "RFP")
            rfp_only = [
                item for item in hybrid.groups if item.status == STATUS_RFP_ONLY
            ]
            self.assertTrue(rfp_only)
            self.assertTrue(all(item.selected_source == "" for item in rfp_only))

            t_com = TableComments(
                file_full_path=str(hybrid.hybrid_path),
                dir_path="-1",
                tabel_class=t_com_init_cls.RFP_AGGREGATED,
            )
            loaded = get_std_from_excel_file(t_com)
            pos = [row for row in loaded if row.row_type == RowType.position_row]
            self.assertTrue(pos)
            by_ds: dict[str, list] = {}
            for row in pos:
                by_ds.setdefault(str(row.el[DS_NAME].value), []).append(row)

            match_rows = by_ds["ДС13"]
            self.assertEqual(len(match_rows), 2)
            tags_blob = " ".join(str(row.el[TAGS].value or "") for row in match_rows)
            self.assertIn(TAG_A, tags_blob)
            self.assertIn(TAG_B, tags_blob)
            for row in match_rows:
                self.assertEqual(int(row.el[VALUES].value), 1)
                self.assertEqual(str(row.el[CODE].value), CODE_MATCH)
            self.assertNotIn("13", by_ds)
            self.assertNotIn("47", by_ds)

            mismatch_rows = by_ds["ДС99"]
            codes = {str(row.el[CODE].value) for row in mismatch_rows}
            self.assertEqual(codes, {CODE_MM, CODE_MIX})
            for row in mismatch_rows:
                self.assertTrue(str(row.el[TAGS].value or "") == "")
                self.assertNotEqual(int(row.el[VALUES].value), 9)
            mix_row = next(
                row for row in mismatch_rows if str(row.el[CODE].value) == CODE_MIX
            )
            self.assertEqual(int(mix_row.el[VALUES].value), 1)

            only_rows = by_ds["ДС88"]
            self.assertEqual(len(only_rows), 1)
            self.assertEqual(str(only_rows[0].el[CODE].value), CODE_ONLY)
            self.assertEqual(int(only_rows[0].el[VALUES].value), 4)

            self.assertNotIn("ДС77", by_ds)
            rfp_only_codes = {
                str(row.el[CODE].value)
                for row in pos
                if str(row.el[CODE].value) == "BCC9999999"
            }
            self.assertFalse(rfp_only_codes)

            blocked_rows = by_ds["ДС8"]
            self.assertEqual(len(blocked_rows), 1)
            self.assertEqual(str(blocked_rows[0].el[CODE].value), CODE_BLOCK)
            self.assertEqual(str(blocked_rows[0].el[TAGS].value or ""), "")

            excl_group = by_ds["ДС66"]
            live = [
                row
                for row in excl_group
                if str(row.el[CODE].value) == CODE_LIVE
            ]
            excluded = [
                row
                for row in excl_group
                if str(row.el[CODE].value) == CODE_EXCL
            ]
            self.assertEqual(len(live), 1)
            self.assertEqual(int(live[0].el[VALUES].value), 2)
            self.assertEqual(len(excluded), 1)
            self.assertEqual(int(excluded[0].el[VALUES].value), 10)
            self.assertEqual(
                str(excluded[0].el[RFP_SUPPLY_STATUS].value),
                CANONICAL_EXCLUDED_FROM_SUPPLY,
            )

            origin_sources = {item.source for item in hybrid.hybrid_rows}
            self.assertIn(SOURCE_RFP, origin_sources)
            self.assertIn(SOURCE_DS, origin_sources)
            match_origin = [
                item
                for item in hybrid.hybrid_rows
                if item.record.ds_name == "ДС13"
            ]
            self.assertTrue(all(item.source == SOURCE_RFP for item in match_origin))
            mismatch_origin = [
                item
                for item in hybrid.hybrid_rows
                if item.record.ds_name == "ДС99"
            ]
            self.assertTrue(all(item.source == SOURCE_DS for item in mismatch_origin))

            report = _load_xlsx(hybrid.report_path)
            try:
                self.assertEqual(
                    set(report.sheetnames),
                    {
                        "Сводка",
                        "Группы поставки",
                        "Разница по кодам",
                        "Покрытие файлов",
                        "Проблемы",
                        "Происхождение",
                    },
                )
                summary = report["Сводка"]
                labels = [
                    row[0]
                    for row in summary.iter_rows(min_row=2, values_only=True)
                    if row and row[0]
                ]
                self.assertIn("Алгоритм", labels)
                self.assertIn("Fingerprint", labels)
                groups_ws = report["Группы поставки"]
                self.assertTrue(groups_ws.auto_filter.ref)
                self.assertEqual(groups_ws.freeze_panes, "A2")
                cover = report["Покрытие файлов"]
                linked = False
                for row in cover.iter_rows(min_row=2, max_col=3):
                    for cell in row:
                        if cell.hyperlink is not None:
                            linked = True
                self.assertTrue(linked)
                origin = report["Происхождение"]
                self.assertGreaterEqual(origin.max_row, 2)
            finally:
                report.close()

            nested = rfp_root / "nested" / "ДС88. AGCC.xlsx"
            self.assertTrue(nested.is_file())
            collected = {item.relpath for item in hybrid.files}
            self.assertNotIn("ДС88. AGCC.xlsx", collected)

    def test_duplicate_rfp_key_blocks_group_and_keeps_ds(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            baseline, hybrid, _rfp_root, _out = _run(Path(raw), duplicate=True)
            self.assertFalse(baseline.blocking, baseline.summary_line())
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            by_id = {item.group_id: item for item in hybrid.groups}
            self.assertEqual(by_id["ДС13"].status, STATUS_BLOCKED)
            self.assertEqual(by_id["ДС13"].selected_source, "ДС")
            self.assertTrue(
                any(item.code == ISSUE_DUPLICATE_RFP_KEY for item in hybrid.issues)
            )
            dup_rows = [
                item
                for item in hybrid.hybrid_rows
                if item.record.ds_name == "ДС13"
            ]
            self.assertEqual(len(dup_rows), 1)
            self.assertEqual(dup_rows[0].record.values, Decimal("2"))
            self.assertTrue(all(item.source == SOURCE_DS for item in dup_rows))
            self.assertTrue(all(not item.record.tags for item in dup_rows))

    def test_unparsed_rfp_is_global_blocker_report_still_written(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            rfp_root = root / "rfp"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            _write_ds_tree(ds_root)
            _write_rfp(
                rfp_root / "без_номера.xlsx",
                [_rfp_data(number=1, code=CODE_MATCH, qty=1)],
            )
            registry = load_registry(registry_path)
            baseline = build_ds_baseline(
                ds_root,
                registry,
                out_dir,
                converter=IdentityDsUnitsConverter(),
                google_index=_google_index(),
                stamp=STAMP,
            )
            hybrid = build_ds_rfp_hybrid(
                baseline,
                registry,
                rfp_root,
                out_dir,
                converter=IdentityRfpUnitsConverter(),
                stamp=STAMP,
            )
            self.assertTrue(hybrid.blocking)
            self.assertIsNone(hybrid.hybrid_path)
            self.assertFalse((out_dir / HYBRID_XLSX_NAME).exists())
            self.assertTrue(hybrid.report_path.is_file())
            self.assertTrue(
                any(item.code == ISSUE_UNPARSED_RFP for item in hybrid.issues)
            )

    def test_baseline_blockers_write_report_without_hybrid(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            rfp_root = root / "rfp"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            _write_xlsx(
                ds_root / "unmapped.xlsx",
                [DS_HEADER, _ds_row(code="")],
            )
            _write_rfp(
                rfp_root / "ДС13. AGCC.xlsx",
                [_rfp_data(number=1, code=CODE_MATCH, qty=1)],
            )
            registry = load_registry(registry_path)
            baseline = build_ds_baseline(
                ds_root,
                registry,
                out_dir,
                converter=IdentityDsUnitsConverter(),
                google_index=_google_index(),
                stamp=STAMP,
            )
            self.assertTrue(baseline.blocking)
            self.assertIsNone(baseline.baseline_path)
            hybrid = build_ds_rfp_hybrid(
                baseline,
                registry,
                rfp_root,
                out_dir,
                converter=IdentityRfpUnitsConverter(),
                stamp=STAMP,
            )
            self.assertTrue(hybrid.blocking)
            self.assertIsNone(hybrid.hybrid_path)
            self.assertTrue(hybrid.report_path.is_file())
            report = _load_xlsx(hybrid.report_path)
            try:
                self.assertIn("Сводка", report.sheetnames)
            finally:
                report.close()

    def test_real_extract_reads_lot_and_status(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            _baseline, hybrid, _rfp, _out = _run(Path(raw))
            extracted = {
                item.source.relpath: item for item in hybrid.extracts
            }
            match_file = extracted["ДС13. AGCC.xlsx"]
            self.assertFalse(match_file.extract_error)
            self.assertEqual(len(match_file.records), 1)
            record = match_file.records[0]
            self.assertEqual(record.code, CODE_MATCH)
            self.assertEqual(record.values, Decimal("2"))
            self.assertIn(TAG_A, record.tags)
            excl_file = extracted["ДС66. AGCC.xlsx"]
            statuses = {item.rfp_supply_status for item in excl_file.records}
            self.assertIn(CANONICAL_EXCLUDED_FROM_SUPPLY, statuses)
            self.assertIn("", statuses)

            types = Counter(
                str(item.record.ds_name) for item in hybrid.hybrid_rows
            )
            self.assertGreaterEqual(types["ДС13"], 2)
            net_wb = load_workbook(hybrid.hybrid_path, data_only=True)
            try:
                ws = net_wb.active
                headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                self.assertEqual(headers[4], "Теги / Tag")
                self.assertEqual(headers[9], "Кол-во")
            finally:
                net_wb.close()


class DsRfpClusterPlacementSmokeTest(unittest.TestCase):
    def test_mixed_ds_label_sorts_digit_ids_numerically(self) -> None:
        self.assertEqual(mixed_ds_label(("61", "15")), "ДС15/61")
        self.assertEqual(mixed_ds_label(("102", "13", "47")), "ДС13/47/102")
        self.assertNotIn("_", mixed_ds_label(("61", "15")))

    def test_separate_unique_keys_plant_on_member_names(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            baseline, hybrid = _run_cluster(
                Path(raw),
                mix_mode=MIX_SEPARATE,
                ds15=[_ds_row(npp=1, qty=2)],
                ds61=[
                    _ds_row(
                        npp=1,
                        code=CODE_ONLY,
                        qty=1,
                        title="1600",
                        system="POS",
                    )
                ],
                rfp_rows=[
                    _rfp_data(
                        number=1,
                        code=CODE_MATCH,
                        qty=2,
                        tag=f"{TAG_A}, {TAG_B}",
                        title="8529-SOS",
                    ),
                    _rfp_data(
                        number=2,
                        code=CODE_ONLY,
                        qty=1,
                        tag=TAG_61,
                        title="1600-POS",
                    ),
                ],
            )
            self.assertFalse(baseline.blocking, baseline.summary_line())
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            self.assertEqual(hybrid.global_blocker_count, 0)
            self.assertFalse(
                any(item.code == ISSUE_AMBIGUOUS_MAPPING for item in hybrid.issues)
            )
            names = {item.record.ds_name for item in hybrid.hybrid_rows}
            self.assertEqual(names, {"ДС15", "ДС61"})
            self.assertTrue(
                all("/" not in item.record.ds_name for item in hybrid.hybrid_rows)
            )
            tags_15 = " ".join(
                item.record.tags or ""
                for item in hybrid.hybrid_rows
                if item.record.ds_name == "ДС15"
            )
            tags_61 = " ".join(
                item.record.tags or ""
                for item in hybrid.hybrid_rows
                if item.record.ds_name == "ДС61"
            )
            self.assertIn(TAG_A, tags_15)
            self.assertIn(TAG_B, tags_15)
            self.assertIn(TAG_61, tags_61)
            self.assertEqual(_qty_sum(hybrid.hybrid_rows), Decimal("3"))
            self.assertEqual(
                _qty_sum(hybrid.hybrid_rows, ds_name="ДС15", code=CODE_MATCH),
                Decimal("2"),
            )
            self.assertEqual(
                _qty_sum(hybrid.hybrid_rows, ds_name="ДС61", code=CODE_ONLY),
                Decimal("1"),
            )
            cluster = hybrid.groups[0]
            self.assertTrue(cluster.reason.startswith("без смешения"))
            extracts = [
                item
                for item in hybrid.extracts
                if item.source.path.name == CLUSTER_RFP_NAME
            ]
            self.assertEqual(len(extracts), 1)
            self.assertEqual(extracts[0].cluster_source_ids, frozenset({"15", "61"}))

    def test_separate_shared_key_keeps_ds_qty_sum_without_slash(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            _baseline, hybrid = _run_cluster(
                Path(raw),
                mix_mode=MIX_SEPARATE,
                ds15=[_ds_row(npp=1, qty=2)],
                ds61=[_ds_row(npp=1, qty=3)],
                rfp_rows=[
                    _rfp_data(
                        number=1,
                        code=CODE_MATCH,
                        qty=5,
                        tag=f"{TAG_A}, {TAG_B}",
                        title="8529-SOS",
                    )
                ],
            )
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            self.assertTrue(
                all("/" not in item.record.ds_name for item in hybrid.hybrid_rows)
            )
            key_rows = [
                item
                for item in hybrid.hybrid_rows
                if item.record.code == CODE_MATCH
            ]
            self.assertEqual(_qty_sum(key_rows), Decimal("5"))
            self.assertTrue(all(not (item.record.tags or "") for item in key_rows))
            self.assertTrue(hybrid.groups[0].reason.startswith("без смешения"))
            self.assertIn("8529-SOS", hybrid.groups[0].reason)

    def test_mixed_unique_keys_use_slash_name_once(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds15 = [_ds_row(npp=1, qty=2)]
            ds61 = [
                _ds_row(
                    npp=1,
                    code=CODE_ONLY,
                    qty=1,
                    title="1600",
                    system="POS",
                )
            ]
            rfp_rows = [
                _rfp_data(
                    number=1,
                    code=CODE_MATCH,
                    qty=2,
                    tag=f"{TAG_A}, {TAG_B}",
                    title="8529-SOS",
                ),
                _rfp_data(
                    number=2,
                    code=CODE_ONLY,
                    qty=1,
                    tag=TAG_61,
                    title="1600-POS",
                ),
            ]
            _baseline, hybrid = _run_cluster(
                root,
                mix_mode=MIX_MIXED,
                ds15=ds15,
                ds61=ds61,
                rfp_rows=rfp_rows,
            )
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            names = {item.record.ds_name for item in hybrid.hybrid_rows}
            self.assertEqual(names, {"ДС15/61"})
            self.assertNotIn("ДС15", names)
            self.assertNotIn("ДС61", names)
            self.assertEqual(_qty_sum(hybrid.hybrid_rows), Decimal("3"))
            tags_blob = " ".join(item.record.tags or "" for item in hybrid.hybrid_rows)
            self.assertIn(TAG_A, tags_blob)
            self.assertIn(TAG_61, tags_blob)
            cluster = hybrid.groups[0]
            self.assertEqual(cluster.group_id, "ДС15/61")
            self.assertTrue(cluster.reason.startswith("смешение ДС15/61"))
            separate = build_ds_rfp_hybrid(
                _baseline,
                load_registry(root / "registry.xlsx"),
                root / "rfp",
                root / "out_separate",
                converter=IdentityRfpUnitsConverter(),
                stamp=STAMP,
                mix_mode=MIX_SEPARATE,
            )
            self.assertNotEqual(hybrid.fingerprint, separate.fingerprint)

    def test_mixed_qty_mismatch_uses_ds_sum_not_rfp(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            _baseline, hybrid = _run_cluster(
                Path(raw),
                mix_mode=MIX_MIXED,
                ds15=[_ds_row(npp=1, qty=2)],
                ds61=[_ds_row(npp=1, qty=1)],
                rfp_rows=[
                    _rfp_data(
                        number=1,
                        code=CODE_MATCH,
                        qty=10,
                        tag=f"{TAG_A}, {TAG_B}",
                        title="8529-SOS",
                    )
                ],
            )
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            key_rows = [
                item
                for item in hybrid.hybrid_rows
                if item.record.code == CODE_MATCH
            ]
            self.assertTrue(key_rows)
            self.assertTrue(all(item.record.ds_name == "ДС15/61" for item in key_rows))
            self.assertEqual(_qty_sum(key_rows), Decimal("3"))
            self.assertTrue(all(not (item.record.tags or "") for item in key_rows))
            reason = hybrid.groups[0].reason
            self.assertTrue(reason.startswith("смешение ДС15/61"))
            self.assertIn("3", reason)
            self.assertIn("10", reason)

    def test_cluster_named_file_is_not_global_blocker(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            _baseline, hybrid = _run_cluster(
                Path(raw),
                mix_mode=MIX_SEPARATE,
                ds15=[_ds_row(npp=1, qty=1)],
                ds61=[
                    _ds_row(
                        npp=1,
                        code=CODE_ONLY,
                        qty=1,
                        title="1600",
                        system="POS",
                    )
                ],
                rfp_rows=[
                    _rfp_data(
                        number=1,
                        code=CODE_MATCH,
                        qty=1,
                        title="8529-SOS",
                    ),
                    _rfp_data(
                        number=2,
                        code=CODE_ONLY,
                        qty=1,
                        title="1600-POS",
                    ),
                ],
            )
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            self.assertEqual(hybrid.global_blocker_count, 0)
            self.assertFalse(
                any(item.code == ISSUE_AMBIGUOUS_MAPPING for item in hybrid.issues)
            )

    def test_mixed_derived_registry_merges_shared_folder_keeps_source(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            registry_path = root / "registry.xlsx"
            shared = "согл УЛ ДС15"
            held: list[bytes] = []
            _baseline, hybrid = _run_cluster(
                root,
                mix_mode=MIX_MIXED,
                ds15=[_ds_row(npp=1, qty=2)],
                ds61=[
                    _ds_row(
                        npp=1,
                        code=CODE_ONLY,
                        qty=1,
                        title="1600",
                        system="POS",
                    )
                ],
                rfp_rows=[
                    _rfp_data(
                        number=1,
                        code=CODE_MATCH,
                        qty=2,
                        tag=f"{TAG_A}, {TAG_B}",
                        title="8529-SOS",
                    ),
                    _rfp_data(
                        number=2,
                        code=CODE_ONLY,
                        qty=1,
                        tag=TAG_61,
                        title="1600-POS",
                    ),
                ],
                registry_kw={
                    "ul15": shared,
                    "ul61": shared,
                    "ds15_file": "ДС15.xlsx",
                    "ds61_file": "ДС61.xlsx",
                },
                source_bytes=held,
            )
            self.assertTrue(held)
            self.assertEqual(registry_path.read_bytes(), held[0])
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            self.assertIsNotNone(hybrid.derived_registry_path)
            derived_path = hybrid.derived_registry_path
            assert derived_path is not None
            self.assertEqual(derived_path.name, registry_path.name)
            self.assertEqual(derived_path.parent, hybrid.hybrid_path.parent)
            self.assertNotEqual(derived_path.resolve(), registry_path.resolve())

            loaded = load_registry(derived_path)
            ids = {row.source_id for row in loaded.rows}
            self.assertEqual(ids, {"15/61"})
            mixed = loaded.rows[0]
            self.assertEqual(mixed.status, STATUS_ACTIVE)
            folders = [rel.ul_folder for rel in mixed.relations if rel.ul_folder]
            self.assertEqual(folders, [shared])
            rfp_files = [rel.rfp_file for rel in mixed.relations if rel.rfp_file]
            self.assertEqual(len(rfp_files), 1)
            self.assertIn(CLUSTER_RFP_NAME, rfp_files[0])
            self.assertIn("ДС15.xlsx", mixed.ds_file)
            self.assertIn("ДС61.xlsx", mixed.ds_file)

            wb = load_workbook(derived_path)
            try:
                ws = wb[REGISTRY_SHEET_NAME]
                folder_hits = 0
                parsed = []
                cells = []
                for excel_row in range(2, ws.max_row + 1):
                    source_id = parse_registry_ds_number(
                        ws.cell(row=excel_row, column=2).value
                    )
                    if source_id != "15/61":
                        continue
                    parsed.append(source_id)
                    cells.append(ws.cell(row=excel_row, column=2).value)
                    folder = str(ws.cell(row=excel_row, column=5).value or "")
                    if folder == shared:
                        folder_hits += 1
                self.assertTrue(parsed)
                self.assertTrue(all(item == "15/61" for item in parsed))
                self.assertTrue(all("ДС15/61" in str(item) for item in cells))
                self.assertEqual(folder_hits, 1)
            finally:
                wb.close()

    def test_separate_derived_registry_keeps_ids_and_comment(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            _baseline, hybrid = _run_cluster(
                Path(raw),
                mix_mode=MIX_SEPARATE,
                ds15=[_ds_row(npp=1, qty=2)],
                ds61=[
                    _ds_row(
                        npp=1,
                        code=CODE_ONLY,
                        qty=1,
                        title="1600",
                        system="POS",
                    )
                ],
                rfp_rows=[
                    _rfp_data(
                        number=1,
                        code=CODE_MATCH,
                        qty=2,
                        tag=f"{TAG_A}, {TAG_B}",
                        title="8529-SOS",
                    ),
                    _rfp_data(
                        number=2,
                        code=CODE_ONLY,
                        qty=1,
                        tag=TAG_61,
                        title="1600-POS",
                    ),
                ],
            )
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            derived_path = hybrid.derived_registry_path
            self.assertIsNotNone(derived_path)
            assert derived_path is not None
            loaded = load_registry(derived_path)
            ids = {row.source_id for row in loaded.rows}
            self.assertEqual(ids, {"15", "61"})
            wb = load_workbook(derived_path)
            try:
                ws = wb[REGISTRY_SHEET_NAME]
                comments = []
                for excel_row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=excel_row, column=2)
                    if cell.comment is not None:
                        comments.append(cell.comment.text)
                self.assertTrue(any("без смешения" in text for text in comments))
            finally:
                wb.close()

    def test_derived_registry_write_failure_does_not_raise(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            with mock.patch(
                "RFQ.rfp_parts.ds_rfp_hybrid.write_registry_workbook",
                side_effect=OSError("locked"),
            ):
                _baseline, hybrid = _run_cluster(
                    Path(raw),
                    mix_mode=MIX_SEPARATE,
                    ds15=[_ds_row(npp=1, qty=1)],
                    ds61=[
                        _ds_row(
                            npp=1,
                            code=CODE_ONLY,
                            qty=1,
                            title="1600",
                            system="POS",
                        )
                    ],
                    rfp_rows=[
                        _rfp_data(
                            number=1,
                            code=CODE_MATCH,
                            qty=1,
                            title="8529-SOS",
                        ),
                        _rfp_data(
                            number=2,
                            code=CODE_ONLY,
                            qty=1,
                            title="1600-POS",
                        ),
                    ],
                )
            self.assertFalse(hybrid.blocking, hybrid.summary_line())
            self.assertIsNotNone(hybrid.hybrid_path)
            self.assertTrue(hybrid.hybrid_path.is_file())
            self.assertIsNone(hybrid.derived_registry_path)


if __name__ == "__main__":
    unittest.main()
