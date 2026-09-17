"""Extra readonly counts for 9110-KSB1, BCC0000805, title-side presence."""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

STEP4 = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.08.25.11.56"
    r"\Шаг4_Сопоставление_RFP_MTO_20260825_120001.xlsx"
)
OUT = Path(__file__).resolve().parent / "diag_step4_customer_rows_extra.md"


def _t(v) -> str:
    return "" if v is None else str(v).replace("\n", " ").strip()


def _p(s: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((s + "\n").encode(enc, errors="backslashreplace"))


def main() -> int:
    wb = load_workbook(STEP4, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    titles_9110: Counter[str] = Counter()
    titles_watch = {
        "1757-POS": Counter(),
        "1600-SOT": Counter(),
        "8630-KSB1": Counter(),
        "8445-SOT1": Counter(),
        "9110-KSB1": Counter(),
        "9110-KSB": Counter(),
        "9110-KSB2": Counter(),
    }
    ksb1_samples: list[str] = []
    code_0805: list[str] = []
    code_0515_unmatched: list[str] = []
    rfp_vs_mto = Counter()  # for 9110-KSB1: has_rfp/has_mto
    for i, cells in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            headers = [_t(c) for c in cells]
            idx = {h: n for n, h in enumerate(headers) if h}
            continue
        title = _t(cells[idx["Титул/Марка"]]) if "Титул/Марка" in idx else ""
        rfp = _t(cells[idx["Код RFP"]]) if "Код RFP" in idx else ""
        mto = _t(cells[idx["Код MTO"]]) if "Код MTO" in idx else ""
        name_r = _t(cells[idx.get("Наименование RFP", 0)]) if "Наименование RFP" in idx else ""
        name_m = _t(cells[idx.get("Наименование MTO", 0)]) if "Наименование MTO" in idx else ""
        ul = _t(cells[idx.get("Статус УЛ", 0)]) if "Статус УЛ" in idx else ""
        if title.upper().startswith("9110-KSB"):
            titles_9110[title] += 1
        if title in titles_watch:
            key = ("RFP" if rfp else "-") + "/" + ("MTO" if mto else "-")
            titles_watch[title][key] += 1
        if title == "9110-KSB1" and len(ksb1_samples) < 15:
            ksb1_samples.append(
                f"| {i} | {title} | {rfp} | {mto} | {ul} | {name_r[:60]} | {name_m[:60]} |"
            )
        ru, mu = rfp.upper(), mto.upper()
        if "BCC0000805" in (ru, mu) and len(code_0805) < 20:
            code_0805.append(f"| {i} | {title} | {rfp} | {mto} | {ul} | {name_r[:50]} | {name_m[:50]} |")
        if ru == "BCC0000515" and not mto and len(code_0515_unmatched) < 8:
            code_0515_unmatched.append(f"| {i} | {title} | {rfp} | {ul} |")
        if mu == "BCC0000805" and not rfp and len(code_0515_unmatched) < 16:
            code_0515_unmatched.append(f"| {i} | {title} | MTO={mto} | {ul} |")
    wb.close()

    lines = [
        "# Extra Step4 dump",
        "",
        "## 9110-KSB* counts",
        "",
    ]
    for t, n in titles_9110.most_common():
        lines.append(f"- `{t}` = {n}")
    if not titles_9110:
        lines.append("- нет титулов 9110-KSB*")
    lines.append("")
    lines.append("## Presence RFP/MTO by title")
    lines.append("")
    for title, counter in titles_watch.items():
        lines.append(f"- `{title}`: " + ", ".join(f"{k}={v}" for k, v in counter.most_common()))
    lines.append("")
    lines.append("## 9110-KSB1 samples")
    lines.append("")
    if ksb1_samples:
        lines.append("| excel_row | title | RFP | MTO | UL | name RFP | name MTO |")
        lines.append("| --- | --- | --- | --- | --- | --- | --- |")
        lines.extend(ksb1_samples)
    else:
        lines.append("Титул `9110-KSB1` в Step4 **не найден**.")
    lines.append("")
    lines.append("## BCC0000805")
    lines.append("")
    if code_0805:
        lines.append("| excel_row | title | RFP | MTO | UL | name RFP | name MTO |")
        lines.append("| --- | --- | --- | --- | --- | --- | --- |")
        lines.extend(code_0805)
    else:
        lines.append("BCC0000805 не найден.")
    lines.append("")
    lines.append("## BCC0000515 unmatched / BCC0000805 MTO-only (partial)")
    lines.append("")
    lines.extend(code_0515_unmatched or ["(нет в выборке)"])
    lines.append("")
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    _p(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
