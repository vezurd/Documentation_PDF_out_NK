"""Smoke tests for RFP parts per-file status JSON and Russian messages."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.rfp_parts.file_status import (
    FILE_STATUS_JSON_NAME,
    STATUS_ERROR,
    STATUS_OK,
    build_file_status_payload,
    humanize_parts_message,
    is_tag_remark_message,
    load_file_status_payload,
    reconstruct_file_status_payload,
    write_file_status_json,
)
from RFQ.rfp_parts.parts_net_preflight import SOURCES_JSON_NAME


@dataclass
class _Stats:
    kind: str
    file_name: str
    ds_name: str = ""
    sheet: str = ""
    rows: int = 0


class RfpPartsFileStatusSmokeTest(unittest.TestCase):
    def test_humanize_english_diagnostics(self) -> None:
        self.assertEqual(
            humanize_parts_message("required columns were not found: NAME, CODE"),
            "не найдены обязательные столбцы: Наименование МТР, Код РД",
        )
        self.assertIn(
            "шапка RFP не найдена",
            humanize_parts_message(
                "RFP header was not found in first 100 rows of sheet 'Table 1'"
            ),
        )
        self.assertTrue(
            humanize_parts_message(
                "row 34: cannot parse Lot VALUES (Excel №17)='193,000'"
            ).startswith("строка 34: не разобрать количество лота")
        )
        self.assertEqual(
            humanize_parts_message("cannot open workbook: PermissionError: locked"),
            "не удалось открыть книгу: PermissionError: locked",
        )

    def test_rfq_in_tag_is_gui_error_not_tag_remark(self) -> None:
        msg = (
            "строка 483: в поле Tag недопустима подстрока RFQ "
            "('AGCC.287-0000-12.4.1-RFQ-0028')"
        )
        self.assertFalse(is_tag_remark_message(msg))
        payload = build_file_status_payload(
            [_Stats("parts", "a.xlsx", rows=1)],
            [("ERROR", "a.xlsx", msg)],
        )
        self.assertEqual(payload["error"], 1)
        self.assertIn("RFQ", payload["files"][0]["detail"])

    def test_humanize_russian_is_stable(self) -> None:
        text = "не найдены обязательные столбцы: Наименование МТР"
        self.assertEqual(humanize_parts_message(text), text)

    def test_build_payload_orders_errors_first(self) -> None:
        stats = [
            _Stats("parts", "ok.xlsx", ds_name="ДС1", rows=10),
            _Stats("parts", "bad.xlsx", ds_name="ДС2", rows=0),
            _Stats("parts", "warn.xlsx", ds_name="ДС3", rows=4),
            _Stats("summary", "Сводная RFP.xlsx", rows=99),
        ]
        warnings = [
            ("ERROR", "bad.xlsx", "required columns were not found: NAME"),
            ("WARN", "warn.xlsx", "в книге 2 рабочих листов, позиции прочитаны только с 'A'"),
        ]
        payload = build_file_status_payload(stats, warnings)
        self.assertEqual(payload["error"], 1)
        self.assertEqual(payload["ok"], 2)
        names = [item["file_name"] for item in payload["files"]]
        self.assertEqual(names, ["bad.xlsx", "ok.xlsx", "warn.xlsx"])
        self.assertEqual(payload["files"][0]["status"], STATUS_ERROR)
        self.assertIn("Наименование МТР", payload["files"][0]["detail"])
        self.assertEqual(payload["files"][1]["status"], STATUS_OK)
        self.assertEqual(payload["files"][1]["detail"], "10 позиций")
        self.assertEqual(payload["files"][2]["status"], STATUS_OK)
        self.assertEqual(payload["files"][2]["detail"], "4 позиций")
        self.assertNotIn("Сводная RFP.xlsx", names)
        self.assertNotIn("рабочих листов", payload["files"][2]["detail"])

    def test_tag_remarks_are_not_gui_errors(self) -> None:
        self.assertTrue(
            is_tag_remark_message(
                "строка 12: теги=['A']; количество тегов (1) не совпадает "
                "с количеством лота (2)"
            )
        )
        self.assertTrue(
            is_tag_remark_message(
                "строка 4: тег '8950-POS-001' дублируется (встречается в 2 строках)"
            )
        )
        self.assertFalse(is_tag_remark_message("не найдены обязательные столбцы: NAME"))
        stats = [_Stats("parts", "ok.xlsx", rows=5)]
        warnings = [
            (
                "WARN",
                "ok.xlsx",
                "строка 12: теги=['A']; количество тегов (1) не совпадает "
                "с количеством лота (2)",
            )
        ]
        payload = build_file_status_payload(stats, warnings)
        self.assertEqual(payload["ok"], 1)
        self.assertEqual(payload["error"], 0)
        self.assertEqual(payload["files"][0]["status"], STATUS_OK)
        self.assertEqual(payload["files"][0]["detail"], "5 позиций")
        self.assertGreaterEqual(payload["tag_remarks"], 1)

    def test_collapse_repeated_row_errors(self) -> None:
        stats = [_Stats("parts", "a.xlsx", rows=3)]
        warnings = [
            ("ERROR", "a.xlsx", "row 10: cannot parse Lot VALUES (Excel №17)='x'"),
            ("ERROR", "a.xlsx", "row 11: cannot parse Lot VALUES (Excel №17)='x'"),
            ("ERROR", "a.xlsx", "row 12: cannot parse Lot VALUES (Excel №17)='x'"),
        ]
        payload = build_file_status_payload(stats, warnings)
        detail = payload["files"][0]["detail"]
        self.assertIn("строки 10, 11, 12:", detail)
        self.assertEqual(detail.count("не разобрать количество лота"), 1)

    def test_json_roundtrip(self) -> None:
        stats = [_Stats("parts", "ok.xlsx", rows=2)]
        payload = build_file_status_payload(stats, [])
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            path = write_file_status_json(folder, payload)
            self.assertEqual(path.name, FILE_STATUS_JSON_NAME)
            loaded = load_file_status_payload(folder)
            self.assertEqual(loaded["ok"], 1)
            self.assertEqual(loaded["files"][0]["file_name"], "ok.xlsx")

    def test_reconstruct_from_sources_and_diagnostics(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            (folder / SOURCES_JSON_NAME).write_text(
                json.dumps(
                    {
                        "parts_dir": "x",
                        "files": [
                            {"name": "ok.xlsx", "mtime_ns": 1, "size": 1},
                            {"name": "bad.xlsx", "mtime_ns": 1, "size": 1},
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            wb = Workbook()
            ws = wb.active
            ws.title = "legend"
            ws.append(["Описание"])
            diag = wb.create_sheet("diagnostics")
            diag.append(["Уровень", "Папка RFP", "Файл", "Строка", "Сообщение"])
            diag.append(
                [
                    "Ошибка",
                    "части",
                    "bad.xlsx",
                    "",
                    "required columns were not found: NAME",
                ]
            )
            wb.save(folder / "rfp_parts_diagnostics.xlsx")
            wb.close()
            payload = reconstruct_file_status_payload(folder)
            by_name = {item["file_name"]: item for item in payload["files"]}
            self.assertEqual(by_name["ok.xlsx"]["status"], STATUS_OK)
            self.assertEqual(by_name["bad.xlsx"]["status"], STATUS_ERROR)
            self.assertIn("Наименование МТР", by_name["bad.xlsx"]["detail"])
            loaded = load_file_status_payload(folder)
            self.assertEqual(loaded["error"], 1)
            self.assertEqual(loaded["ok"], 1)


class RfpPartsDuplicateTagsXlsxSmokeTest(unittest.TestCase):
    def test_collect_and_write_duplicate_tags_xlsx(self) -> None:
        from decimal import Decimal

        from RFQ.rfp_parts.analyze_rfp_parts import (
            DUPLICATE_TAGS_XLSX_NAME,
            RfpRecord,
            _collect_tag_remarks,
            _write_duplicate_tags_xlsx,
        )

        def rec(**kwargs: object) -> RfpRecord:
            data: dict[str, object] = {
                "kind": "parts",
                "file_name": "a.xlsx",
                "sheet": "Перечень материалов",
                "excel_row": 10,
                "ds_name": "ДС1",
                "ds_number": "1",
                "ds_title": "8950",
                "ds_specification": "",
                "tags": "TAG-1",
                "ds_code_1c": "",
                "code": "BCC1",
                "name": "насос",
                "type_mark": "",
                "values": Decimal("1"),
                "units": "шт",
            }
            data.update(kwargs)
            return RfpRecord(**data)  # type: ignore[arg-type]

        records = [
            rec(file_name="a.xlsx", excel_row=10, tags="DUP-1", code="BCC1"),
            rec(file_name="b.xlsx", excel_row=20, tags="DUP-1", code="BCC2"),
            rec(
                file_name="c.xlsx",
                excel_row=30,
                tags="ONE, TWO",
                values=Decimal("1"),
                code="BCC3",
            ),
            rec(
                file_name="d.xlsx",
                excel_row=40,
                tags="SOLO",
                values=Decimal("3"),
                code="BCC4",
            ),
        ]
        remarks = _collect_tag_remarks(records)
        kinds = {item["kind"] for item in remarks}
        self.assertIn("duplicate", kinds)
        self.assertIn("tags_gt_lot", kinds)
        self.assertIn("lot_gt_tags", kinds)
        self.assertNotIn("count_mismatch", kinds)
        dup_rows = [item for item in remarks if item["kind"] == "duplicate"]
        self.assertEqual(len(dup_rows), 2)
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            path = folder / DUPLICATE_TAGS_XLSX_NAME
            written = _write_duplicate_tags_xlsx(path, remarks)
            self.assertEqual(written, len(remarks))
            self.assertTrue(path.is_file())
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                self.assertIn("Дубли тегов", wb.sheetnames)
                ws = wb["Дубли тегов"]
                self.assertGreater(ws.max_row, 1)
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
