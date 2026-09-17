"""Writer tests for RD catalog styled table xlsx export (no Qt)."""

from __future__ import annotations

import io
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rd_catalog.kits import kit_identity_key
from rd_catalog.monitor_views import KITS_HEADERS, MonitorCell, contrast_foreground
from rd_catalog.table_xlsx import (
    ExportedCell,
    ExportedColumn,
    ExportedTable,
    cell_from_monitor,
    columns_from_client,
    columns_from_layout,
    dated_xlsx_filename,
    exported_table_from_kit_cells,
    px_to_excel_width,
    write_exported_table_xlsx,
)


def _fill_rgb(cell) -> str | None:
    fill = cell.fill
    fill_type = getattr(fill, "fill_type", None) or getattr(
        fill, "patternType", None
    )
    if not fill_type:
        return None
    color = fill.fgColor
    rgb = getattr(color, "rgb", None)
    if rgb is None:
        return None
    text = str(rgb)
    if len(text) == 8:
        text = text[2:]
    return text


def _load_written_workbook(path: Path):
    import openpyxl

    return openpyxl.load_workbook(io.BytesIO(path.read_bytes()), data_only=False)


def test_px_to_excel_width_formula_and_clamp() -> None:
    assert abs(px_to_excel_width(57) - (57.0 - 5.0) / 7.0) < 1e-9
    assert px_to_excel_width(0) == 1.0
    assert px_to_excel_width(10_000) == 255.0


def test_dated_xlsx_filename_kits_stem() -> None:
    name = dated_xlsx_filename("Комплекты", datetime(2026, 9, 16))
    assert name == "Комплекты_2026.09.16.xlsx"


def test_write_round_trip_visual_order_freeze_a2() -> None:
    table = ExportedTable(
        columns=(
            ExportedColumn(header="Марка", width_px=70, logical_index=1),
            ExportedColumn(header="Титул", width_px=57, logical_index=0),
            ExportedColumn(header="Ок", width_px=40, logical_index=2),
        ),
        rows=(
            (
                ExportedCell(
                    text="KSB",
                    fill_hex="#E2F2E1",
                    fg_hex="#202124",
                    bold=True,
                ),
                ExportedCell(text="9110"),
                ExportedCell(
                    text="да",
                    fill_hex="#F7E8BE",
                    fg_hex="#202124",
                ),
            ),
            (
                ExportedCell(text="POS", fill_hex="#E2F2E1", fg_hex="#202124"),
                ExportedCell(text="1757"),
                ExportedCell(text="нет"),
            ),
        ),
        sheet_name="Комплекты",
    )
    from openpyxl.utils import get_column_letter

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "kits.xlsx"
        write_exported_table_xlsx(table, path)
        workbook = _load_written_workbook(path)
        try:
            sheet = workbook.active
            assert sheet.title == "Комплекты"
            assert [sheet.cell(1, col).value for col in range(1, 4)] == [
                "Марка",
                "Титул",
                "Ок",
            ]
            assert [sheet.cell(2, col).value for col in range(1, 4)] == [
                "KSB",
                "9110",
                "да",
            ]
            assert [sheet.cell(3, col).value for col in range(1, 4)] == [
                "POS",
                "1757",
                "нет",
            ]
            assert _fill_rgb(sheet.cell(2, 1)).lower() == "e2f2e1"
            assert _fill_rgb(sheet.cell(2, 2)) is None
            assert _fill_rgb(sheet.cell(2, 3)).lower() == "f7e8be"
            assert sheet.cell(2, 1).font.bold is True
            assert sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref == "A1:C3"
            for col_idx, column in enumerate(table.columns, start=1):
                letter = get_column_letter(col_idx)
                width = sheet.column_dimensions[letter].width
                assert abs(width - px_to_excel_width(column.width_px)) < 0.05
        finally:
            workbook.close()


def test_write_freeze_c2_when_title_mark_first() -> None:
    table = ExportedTable(
        columns=(
            ExportedColumn(header="Титул", width_px=57, logical_index=0),
            ExportedColumn(header="Марка", width_px=70, logical_index=1),
            ExportedColumn(header="Ок", width_px=40, logical_index=2),
        ),
        rows=((ExportedCell(text="9110"), ExportedCell(text="KSB"), ExportedCell(text="да")),),
        sheet_name="Комплекты",
    )
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "kits_title_first.xlsx"
        write_exported_table_xlsx(table, path)
        workbook = _load_written_workbook(path)
        try:
            assert workbook.active.freeze_panes == "C2"
        finally:
            workbook.close()


def test_cell_from_monitor_contrast_foreground() -> None:
    cell = cell_from_monitor(MonitorCell(text="да", fill="#E2F2E1"))
    assert cell.text == "да"
    assert cell.fill_hex == "#E2F2E1"
    assert cell.fg_hex == contrast_foreground("#E2F2E1")
    assert cell.fg_hex == "#202124"


def test_empty_rows_write_header_and_autofilter() -> None:
    table = ExportedTable(
        columns=(
            ExportedColumn(header="Марка", width_px=70),
            ExportedColumn(header="Титул", width_px=57),
            ExportedColumn(header="Ок", width_px=40),
        ),
        rows=(),
        sheet_name="Комплекты",
    )
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "kits_empty.xlsx"
        write_exported_table_xlsx(table, path)
        workbook = _load_written_workbook(path)
        try:
            sheet = workbook.active
            assert [sheet.cell(1, col).value for col in range(1, 4)] == [
                "Марка",
                "Титул",
                "Ок",
            ]
            assert sheet.cell(2, 1).value is None
            assert sheet.auto_filter.ref == "A1:C1"
            assert sheet.freeze_panes == "A2"
        finally:
            workbook.close()


def test_columns_from_layout_and_client() -> None:
    laid = columns_from_layout(
        ("Титул", "Марка"),
        {"Титул": 57},
    )
    assert laid[0].width_px == 57
    assert laid[1].width_px == 80
    client = columns_from_client(
        [
            {"header": "Марка", "width_px": 99},
            {"header": "АН", "width_px": 12},
            {"header": "unknown", "width_px": 8},
        ],
        known_names=KITS_HEADERS,
        fallback_order=("Титул", "Марка"),
        fallback_widths={"Титул": 10},
    )
    assert [col.header for col in client] == ["Марка", "АН МТО"]
    assert client[0].width_px == 99
    fallback = columns_from_client(
        [],
        known_names=KITS_HEADERS,
        fallback_order=("Титул", "Марка"),
        fallback_widths={"Титул": 57},
    )
    assert [col.header for col in fallback] == ["Титул", "Марка"]
    assert fallback[0].width_px == 57


def test_exported_table_from_kit_cells_skips_unknown_and_duplicates() -> None:
    key = kit_identity_key("9110", "KSB")
    cells = {
        key: {
            "Титул": MonitorCell(text="9110", fill="#E2F2E1"),
            "Марка": MonitorCell(text="KSB"),
            "Ок": MonitorCell(text="да", fill="#F7E8BE", bold=True),
        }
    }
    columns = columns_from_layout(
        ("Марка", "Титул", "Ок"),
        {"Марка": 70, "Титул": 57, "Ок": 40},
    )
    table = exported_table_from_kit_cells(
        cells,
        (
            ("9110", "KSB"),
            ("9110", "ksb"),
            ("0000", "NONE"),
            ("9110", "KSB"),
        ),
        columns,
    )
    assert len(table.rows) == 1
    assert [cell.text for cell in table.rows[0]] == ["KSB", "9110", "да"]
    assert table.rows[0][2].bold is True
    assert table.rows[0][0].fill_hex is None
    assert table.rows[0][1].fill_hex == "#E2F2E1"


def main() -> None:
    """Run styled-xlsx assertions without Qt or a live database."""

    test_px_to_excel_width_formula_and_clamp()
    test_dated_xlsx_filename_kits_stem()
    test_write_round_trip_visual_order_freeze_a2()
    test_write_freeze_c2_when_title_mark_first()
    test_cell_from_monitor_contrast_foreground()
    test_empty_rows_write_header_and_autofilter()
    test_columns_from_layout_and_client()
    test_exported_table_from_kit_cells_skips_unknown_and_duplicates()
    print("RD catalog table xlsx: OK")


if __name__ == "__main__":
    main()
