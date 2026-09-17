"""Quick check of the 16.35 Step4 Excel: leftover titles and status mix."""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parents[1] / "tmp" / "ul_title_diag_20260818_1635.txt"
STEP4 = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.08.18.16.35"
    r"\Шаг4_Сопоставление_RFP_MTO_20260818_163719.xlsx"
)
COMPOSITE_RE = re.compile(r"^[0-9]+-[A-Za-zА-Яа-я0-9]+$")
TITLE_ONLY_RE = re.compile(r"^[0-9]+$")


def main() -> int:
    lines: list[str] = []
    lines.append(f"exists={STEP4.exists()} size={STEP4.stat().st_size}")
    wb = openpyxl.load_workbook(STEP4, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value
    }
    title_col = headers["Титул/Марка"]
    status_col = headers["Статус УЛ"]
    status_counts: Counter[str] = Counter()
    leftover_kinds: Counter[str] = Counter()
    leftover_samples: list[str] = []
    leftover_title_only: list[str] = []
    leftover_titles: Counter[str] = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        status = str(row[status_col - 1] or "").strip()
        status_counts[status or "(empty)"] += 1
        if status != "Только в УЛ":
            continue
        title = str(row[title_col - 1] or "").strip()
        leftover_titles[title] += 1
        if COMPOSITE_RE.fullmatch(title):
            leftover_kinds["composite"] += 1
            if len(leftover_samples) < 8:
                leftover_samples.append(title)
        elif TITLE_ONLY_RE.fullmatch(title):
            leftover_kinds["title_only"] += 1
            if len(leftover_title_only) < 8:
                leftover_title_only.append(title)
        elif not title:
            leftover_kinds["empty"] += 1
        else:
            leftover_kinds["other"] += 1
            if len(leftover_samples) < 8:
                leftover_samples.append(title)
    wb.close()
    lines.append("STATUS:")
    for key, value in status_counts.most_common():
        lines.append(f"  {key}: {value}")
    lines.append("Только в УЛ Титул/Марка kinds:")
    for key, value in leftover_kinds.most_common():
        lines.append(f"  {key}: {value}")
    lines.append("composite samples: " + ", ".join(leftover_samples))
    lines.append("title_only samples: " + ", ".join(leftover_title_only))
    lines.append("top leftover titles:")
    for key, value in leftover_titles.most_common(12):
        lines.append(f"  {key!r}: {value}")
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {OUT}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
