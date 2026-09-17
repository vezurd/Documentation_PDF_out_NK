"""Read-only: RFP / MTO / UL source qty for 2265-KSB BCC0000642."""
from __future__ import annotations

import os
import sys

from openpyxl import load_workbook

CODE = "BCC0000642"
TITLE = "2265-KSB"
TSD_ROOT = (
    r"\\bcc\root\CurProjects\di_manegers\Пятницкий_П"
    r"\Амурский ГХК\Поставки\ТСД по всем ДС"
)

RFP = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_RFP\RFP_Зиновьев\ДС67. AGCC.287-0000-12.4.1-RFP-0032_0_RU.xlsx"
)
MTO = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_ГОТОВЫЕ_ДЛЯ_РОБОТА\2265\AGCC.287-2265-KSB.MTO-0001_01-AN02_RU.xlsx"
)
UL = os.path.join(
    TSD_ROOT,
    r"согл УЛ ДС67",
    "Packing list PL_2076961.32_3705.01 - ред.220526.xlsx",
)


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def _norm(v: object) -> str:
    return "" if v is None else str(v).strip()


def dump_hits(path: str, extra_rows: set[int] | None = None) -> None:
    _safe_print(f"\nFILE exists={os.path.exists(path)} {ascii(path)}")
    if not os.path.exists(path):
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            _safe_print(f"-- sheet {ascii(sheet_name)}")
            header = None
            qty_sum = 0.0
            n_hits = 0
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if i == 1:
                    header = [_norm(v) for v in row]
                    continue
                joined = " | ".join(_norm(v) for v in row)
                force = extra_rows is not None and i in extra_rows
                if CODE not in joined.upper() and TITLE not in joined and not force:
                    continue
                n_hits += 1
                _safe_print(f"r{i}: {joined[:450]}")
                for v in row:
                    if isinstance(v, (int, float)) and v not in (0, 1) and 0 < v < 10000:
                        pass
            _safe_print(f"hits={n_hits}")
    finally:
        wb.close()


def main() -> int:
    _safe_print(f"TSD_ROOT={ascii(TSD_ROOT)}")
    dump_hits(RFP)
    dump_hits(MTO)
    dump_hits(UL, extra_rows={110, 111, 112, 113, 114, 115})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
