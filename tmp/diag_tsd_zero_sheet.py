"""Dump why a zero-position Single sheet still has no rows after VENDOR fix."""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl

from base.base_classes import RowType
from base.tables_columns import CODE, NAME, SPECIFICATION_NAME, UNITS, VALUES, VENDOR, TITLE
from RFQ.ds_compare.tsd_packing_load import (
    DEFAULT_TSD_PACKING_ROOT,
    collect_tsd_files,
    load_tsd_file_rows,
)

OUT = ROOT / "tmp" / "tsd_zero_sheet_diag.txt"
TARGET = "PL_2076961.12_2893.16.xlsx"


def main() -> int:
    docs = collect_tsd_files(DEFAULT_TSD_PACKING_ROOT)
    cand = next(d for d in docs if TARGET in d.file_name)
    fp = cand.file_full_path
    lines = [f"file: {fp}"]
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    lines.append(f"sheets: {wb.sheetnames}")
    for sheet in wb.sheetnames:
        if sheet.lower().startswith("master"):
            continue
        ws = wb[sheet]
        lines.append(f"--- {sheet!r} raw rows 1..25 ---")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=25, max_col=14, values_only=True)):
            cells = [(j, v) for j, v in enumerate(row) if v not in (None, "")]
            if cells:
                lines.append(f"{i + 1}: {cells}")
    wb.close()

    pos, plan, all_rows, hits = load_tsd_file_rows(fp, root=DEFAULT_TSD_PACKING_ROOT)
    lines.append(f"pos={len(pos)} hits={[(h.sheet_name, h.position_rows) for h in hits]}")
    lines.append(f"types={dict(Counter(r.row_type for r in all_rows))}")
    # count near-misses
    miss = Counter()
    for r in all_rows:
        if r.row_type != RowType.other_row:
            continue
        code = str(r.el[CODE].value or "").strip()
        name = str(r.el[NAME].value or "").strip()
        values = str(r.el[VALUES].value or "").strip()
        units = str(r.el[UNITS].value or "").strip()
        flags = (
            f"c{int(bool(code))}n{int(bool(name))}v{int(bool(values))}u{int(bool(units))}"
        )
        miss[flags] += 1
    lines.append(f"other_row presence patterns code/name/values/units: {dict(miss)}")
    shown = 0
    for r in all_rows:
        if r.row_type != RowType.other_row:
            continue
        code = str(r.el[CODE].value or "").strip()
        if not code.startswith("BCC") and "AGCC" not in str(r.el[SPECIFICATION_NAME].value or ""):
            continue
        lines.append(
            {
                CODE: repr(r.el[CODE].value)[:60],
                SPECIFICATION_NAME: repr(r.el[SPECIFICATION_NAME].value)[:60],
                NAME: repr(r.el[NAME].value)[:60],
                VALUES: repr(r.el[VALUES].value)[:40],
                UNITS: repr(r.el[UNITS].value)[:40],
                VENDOR: repr(r.el[VENDOR].value)[:40],
            }
        )
        shown += 1
        if shown >= 8:
            break
    OUT.write_text("\n".join(str(x) for x in lines) + "\n", encoding="utf-8")
    print("wrote", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
