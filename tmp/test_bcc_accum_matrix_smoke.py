"""Smoke tests for the Step4 BCC accumulation matrix (no replacements/tags)."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from base.base_classes import RowStd, RowType
from base.tables_columns import (
    ANNOTATION,
    CODE,
    DS_NAME,
    DS_NUMBER,
    NAME,
    TYPE_MARK,
    UNITS,
    VALUES,
    VENDOR,
)
from RFQ.packing_list_provider import PackingDataset, PackingQualityLevel
from RFQ.tags_rfp_compare.step4.step4_bcc_accum_matrix import (
    build_bcc_accum_matrix,
    save_bcc_accum_matrix_to_excel,
)


def _pos(
    *,
    code: str = "",
    qty: object = 1,
    units: str = "шт",
    name: str = "",
    type_mark: str = "",
    vendor: str = "",
    ds_name: str = "",
    ds_number: str = "",
    annotation: str = "",
) -> RowStd:
    row = RowStd()
    row.row_type = RowType.position_row
    row.el[CODE].value = code
    row.el[VALUES].value = qty
    row.el[UNITS].value = units
    row.el[NAME].value = name
    row.el[TYPE_MARK].value = type_mark
    row.el[VENDOR].value = vendor
    row.el[DS_NAME].value = ds_name
    row.el[DS_NUMBER].value = ds_number
    row.el[ANNOTATION].value = annotation
    return row


def _google(
    *,
    code: str,
    name: str,
    type_mark: str,
    units: str,
    vendor: str,
) -> RowStd:
    row = RowStd()
    row.el[CODE].value = code
    row.el[NAME].value = name
    row.el[TYPE_MARK].value = type_mark
    row.el[UNITS].value = units
    row.el[VENDOR].value = vendor
    return row


def _packing(rows: list[RowStd]) -> PackingDataset:
    return PackingDataset(
        rows=rows,
        meta=None,
        quality=PackingQualityLevel.OK,
        issues=[],
        cache_path="",
    )


class BccAccumMatrixSmokeTest(unittest.TestCase):
    def test_one_code_three_mto_titles(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[],
            mto_data={
                "8445-SOT1": [_pos(code="BCC0001", qty=2, name="from-mto")],
                "9110-KSB": [_pos(code="BCC0001", qty=3, name="from-mto")],
                "1757-POS": [_pos(code="BCC0001", qty=5, name="from-mto")],
            },
            packing_dataset=None,
        )
        self.assertEqual(len(matrix.rows), 1)
        row = matrix.rows[0]
        self.assertEqual(row.code, "BCC0001")
        self.assertEqual(row.qty_mto, 10)
        self.assertEqual(row.mto_by_title["8445-SOT1"], 2)
        self.assertEqual(row.mto_by_title["9110-KSB"], 3)
        self.assertEqual(row.mto_by_title["1757-POS"], 5)
        self.assertEqual(sorted(matrix.mto_title_columns), ["1757-POS", "8445-SOT1", "9110-KSB"])
        self.assertEqual(
            row.mto_titles_info,
            "1757-POS - 5\n8445-SOT1 - 2\n9110-KSB - 3",
        )

    def test_google_identity_overrides_mto_text(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[],
            mto_data={
                "T1": [
                    _pos(
                        code="BCC0002",
                        qty=1,
                        name="mto-name",
                        type_mark="mto-mark",
                        units="шт",
                        vendor="mto-vendor",
                    )
                ]
            },
            packing_dataset=None,
            google_rows=[
                _google(
                    code="BCC0002",
                    name="google-name",
                    type_mark="google-mark",
                    units="компл",
                    vendor="google-vendor",
                )
            ],
        )
        row = matrix.rows[0]
        self.assertEqual(row.name, "google-name")
        self.assertEqual(row.type_mark, "google-mark")
        self.assertEqual(row.units, "компл")
        self.assertEqual(row.vendor, "google-vendor")

    def test_empty_code_position_rows_not_collapsed(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[],
            mto_data={
                "T1": [
                    _pos(code="", qty=1, name="Bolt A", units="шт"),
                    _pos(code="", qty=4, name="Bolt B", units="шт"),
                ]
            },
            packing_dataset=None,
        )
        names = sorted(row.name for row in matrix.rows)
        self.assertEqual(names, ["Bolt A", "Bolt B"])
        by_name = {row.name: row for row in matrix.rows}
        self.assertEqual(by_name["Bolt A"].qty_mto, 1)
        self.assertEqual(by_name["Bolt B"].qty_mto, 4)
        self.assertEqual(by_name["Bolt A"].code, "")

    def test_empty_code_same_identity_sums(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[],
            mto_data={
                "T1": [_pos(code="", qty=2, name="Gasket", units="шт")],
                "T2": [_pos(code="", qty=3, name="Gasket", units="шт")],
            },
            packing_dataset=None,
        )
        self.assertEqual(len(matrix.rows), 1)
        self.assertEqual(matrix.rows[0].qty_mto, 5)
        self.assertEqual(matrix.rows[0].mto_by_title["T1"], 2)
        self.assertEqual(matrix.rows[0].mto_by_title["T2"], 3)

    def test_outer_join_ds_and_ul_only(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[
                _pos(code="BCC0003", qty=7, ds_name="ДС15", ds_number="12", name="ds-name")
            ],
            mto_data={},
            packing_dataset=_packing(
                [_pos(code="BCC0004", qty=9, annotation=r"C:\ul\PL-001.xlsx")]
            ),
        )
        codes = sorted(row.code for row in matrix.rows)
        self.assertEqual(codes, ["BCC0003", "BCC0004"])
        by_code = {row.code: row for row in matrix.rows}
        self.assertEqual(by_code["BCC0003"].qty_ds, 7)
        self.assertEqual(by_code["BCC0003"].qty_mto, 0)
        self.assertEqual(by_code["BCC0003"].ds_sources_info, "ДС15 - 7")
        self.assertEqual(by_code["BCC0004"].qty_ul, 9)
        self.assertEqual(by_code["BCC0004"].ul_sources_info, "PL-001.xlsx - 9")

    def test_units_mismatch_blank_diff(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[_pos(code="BCC0005", qty=1, units="кг", ds_name="ДС1")],
            mto_data={"T1": [_pos(code="BCC0005", qty=2, units="шт")]},
            packing_dataset=None,
        )
        row = matrix.rows[0]
        self.assertTrue(row.units_mismatch)
        self.assertIsNone(row.diff_mto_ds)
        self.assertEqual(row.qty_mto, 2)
        self.assertEqual(row.qty_ds, 1)

    def test_compatible_diff(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[_pos(code="BCC0006", qty=4, units="шт.", ds_name="ДС1")],
            mto_data={"T1": [_pos(code="BCC0006", qty=10, units="шт")]},
            packing_dataset=_packing(
                [_pos(code="BCC0006", qty=3, units="шт", annotation="ul1.xlsx")]
            ),
        )
        row = matrix.rows[0]
        self.assertFalse(row.units_mismatch)
        self.assertEqual(row.diff_mto_ds, 6)
        self.assertEqual(row.diff_ds_ul, 1)
        self.assertEqual(row.diff_mto_ul, 7)

    def test_section_rows_skipped(self) -> None:
        section = _pos(code="BCC0007", qty=99, name="section")
        section.row_type = RowType.section_row
        matrix = build_bcc_accum_matrix(
            rfp_rows=[],
            mto_data={"T1": [section, _pos(code="BCC0007", qty=1)]},
            packing_dataset=None,
        )
        self.assertEqual(len(matrix.rows), 1)
        self.assertEqual(matrix.rows[0].qty_mto, 1)

    def test_excel_headers_and_google_row(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[_pos(code="BCC0008", qty=1, ds_name="ДС31", units="шт")],
            mto_data={"8445-SOT1": [_pos(code="BCC0008", qty=2, units="шт")]},
            packing_dataset=_packing(
                [_pos(code="BCC0008", qty=1, units="шт", annotation="UL-A.xlsx")]
            ),
            google_rows=[
                _google(
                    code="BCC0008",
                    name="G-name",
                    type_mark="G-mark",
                    units="шт",
                    vendor="G-vendor",
                )
            ],
        )
        tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        try:
            path = save_bcc_accum_matrix_to_excel(matrix, tmp.name)
            self.assertTrue(os.path.isfile(path))
            self.assertTrue(os.path.basename(path).startswith("Шаг4_Матрица_BCC_"))
            wb = load_workbook(path, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                header_row = next(ws.iter_rows(min_row=1, max_row=1))
                data_row = next(ws.iter_rows(min_row=2, max_row=2))
                headers = [cell.value for cell in header_row]
                data = [cell.value for cell in data_row]
                sheetnames = list(wb.sheetnames)
            finally:
                wb.close()
        finally:
            tmp.cleanup()
        self.assertEqual(
            headers,
            [
                "Код",
                "Наименование",
                "Тип / марка",
                "Вендор",
                "Ед. изм.",
                "Σ МТО",
                "Σ ДС",
                "Σ УЛ",
                "МТО − ДС",
                "ДС − УЛ",
                "МТО − УЛ",
                "Титулы МТО",
                "Источники ДС",
                "Источники УЛ",
            ],
        )
        self.assertEqual(data[0], "BCC0008")
        self.assertEqual(data[3], "G-vendor")
        self.assertEqual(data[4], "шт")
        self.assertEqual(data[5], 2)
        self.assertEqual(data[6], 1)
        self.assertEqual(data[7], 1)
        self.assertEqual(data[11], "8445-SOT1 - 2")
        self.assertEqual(data[12], "ДС31 - 1")
        self.assertEqual(data[13], "UL-A.xlsx - 1")
        self.assertNotIn("8445-SOT1", headers)
        self.assertNotIn("ДС31", headers)
        self.assertIn("Сводка", sheetnames)

    def test_summary_kpi_only_ul_and_all_three(self) -> None:
        matrix = build_bcc_accum_matrix(
            rfp_rows=[_pos(code="BCC_ALL", qty=2, ds_name="ДС1", units="шт")],
            mto_data={"T1": [_pos(code="BCC_ALL", qty=3, units="шт")]},
            packing_dataset=_packing(
                [
                    _pos(code="BCC_ALL", qty=1, units="шт", annotation="all.xlsx"),
                    _pos(code="BCC_UL", qty=5, units="шт", annotation="only-ul.xlsx"),
                ]
            ),
        )
        tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        try:
            path = save_bcc_accum_matrix_to_excel(matrix, tmp.name)
            wb = load_workbook(path, data_only=True)
            try:
                ws = wb["Сводка"]
                kpi: dict[str, object] = {}
                for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=2):
                    label = row[0].value
                    if label:
                        kpi[str(label)] = row[1].value
            finally:
                wb.close()
        finally:
            tmp.cleanup()
        self.assertEqual(kpi["Только УЛ"], 1)
        self.assertEqual(kpi["Во всех трёх контурах"], 1)


if __name__ == "__main__":
    unittest.main()
