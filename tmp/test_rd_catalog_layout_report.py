"""Lexical checks for RD layout classification and the author report."""

from __future__ import annotations

import io
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rd_catalog.layout_report import (
    build_layout_report,
    format_layout_report_summary,
    format_layout_report_text,
    layout_report_default_filename,
    write_layout_report_xlsx,
)
from rd_catalog.models import FileKind, FileRecord, ReviewState, SourceKind
from rd_catalog.parse import (
    classify_rd_layout_reason,
    layout_reason_label,
    list_layout_violations,
)

RD_ROOT = r"C:\rd"


def _record(
    path: str,
    *,
    title: str = "",
    mark: str = "",
    source: SourceKind = SourceKind.RD,
    present: bool = True,
    file_id: int = 1,
) -> FileRecord:
    return FileRecord(
        id=file_id,
        path=path,
        path_key=path.casefold(),
        source=source,
        present=present,
        review_state=ReviewState.ACKNOWLEDGED,
        first_seen_run_id=1,
        last_seen_run_id=1,
        data={
            "title": title,
            "mark": mark,
            "parse_status": "parsed",
            "file_kind": FileKind.PDF.value,
        },
    )


def _reason(path: str) -> str:
    return classify_rd_layout_reason(path, RD_ROOT)


def test_new_tokens_on_synthetic_paths() -> None:
    cases = {
        "bad_title_folder": (
            r"C:\rd\_SQ список по титулам\2815\SKUD"
            r"\AGCC.287-2815-SKUD.OD-0001_0_RU.pdf"
        ),
        "gate_under_title": (
            r"C:\rd\9000\03_Для передачи\04_рев.0-AN02_AGCC.287-9000-KSB"
            r"\PDF\AGCC.287-9000-KSB.OD-0001_0_RU.pdf"
        ),
        "bad_mark_folder": (
            r"C:\rd\1757\МАРКАП\СКУД"
            r"\AGCC.287-1757-SKUD.LAY-0005_A_RU.pdf"
        ),
        "working_folder": (
            r"C:\rd\1715\14_POS\Резерв\02_рев.0_AGCC.287-1715-POS"
            r"\DWG\AGCC.287-1715-POS.WIR-0001_0_RU.dwg"
        ),
        "extra_before_gate": (
            r"C:\rd\8630\14_KSB4\вложенная\для передачи"
            r"\03_рев.0_AGCC.287-8630-KSB4"
            r"\PDF\AGCC.287-8630-KSB4.OD-0001_0_RU.pdf"
        ),
        "bad_package_name": (
            r"C:\rd\8442\13_KBI\Rev.1_AGCC.287-8442-KBI"
            r"\BBB\AGCC.287-8442-KBI.BOM-0001_01_RU.xlsx"
        ),
        "duplicate_level": (
            r"C:\rd\8630\14_KSB4\03_KSB4\для передачи"
            r"\03_рев.AN-01_AGCC.287-8630-KSB4"
            r"\PDF\AGCC.287-8630-KSB4.OD-0001_0_RU.pdf"
        ),
        "package_nested": (
            r"C:\rd\9110\KSB\вложенная\01_рев.0_AGCC.287-9110-KSB"
            r"\PDF\AGCC.287-9110-KSB.OD-0001_0_RU.pdf"
        ),
        "outside_gate": (
            r"C:\rd\8950\24_SOO2\ЭМ"
            r"\AGCC.287-8950-SOO2.OD-0001_0_RU.pdf"
        ),
        "other": r"D:\elsewhere\AGCC.287-9000-KSB.OD-0001_0_RU.pdf",
    }
    for token, path in cases.items():
        assert _reason(path) == token, (token, path, _reason(path))
        assert layout_reason_label(token)
        assert layout_reason_label(token) != token


def test_working_and_gate_case_and_unicode_hyphen() -> None:
    reserve_cases = (
        r"C:\rd\1715\14_POS\Резерв\file.dwg",
        r"C:\rd\1715\14_POS\резерв\file.dwg",
        r"C:\rd\1715\14_POS\РЕЗЕРВ\file.dwg",
    )
    reasons = {_reason(path) for path in reserve_cases}
    assert reasons == {"working_folder"}

    extra_ascii = (
        r"C:\rd\8630\14_KSB4\вложенная\для передачи"
        r"\03_рев.0_x\PDF\file.pdf"
    )
    extra_title = (
        r"C:\rd\8630\14_KSB4\вложенная\Для передачи"
        r"\03_рев.0_x\PDF\file.pdf"
    )
    extra_hyphen = (
        r"C:\rd\8630\14_KSB4\вложенная\для"
        "\u2010"
        r"передачи\03_рев.0_x\PDF\file.pdf"
    )
    assert _reason(extra_ascii) == "extra_before_gate"
    assert _reason(extra_title) == "extra_before_gate"
    assert _reason(extra_hyphen) == "extra_before_gate"

    gate_under_lower = (
        r"C:\rd\9000\03_для передачи\04_рев.0_x\PDF\file.pdf"
    )
    gate_under_hyphen = (
        r"C:\rd\9000\03_Для"
        "\u2010"
        r"передачи\04_рев.0_x\PDF\file.pdf"
    )
    assert _reason(gate_under_lower) == "gate_under_title"
    assert _reason(gate_under_hyphen) == "gate_under_title"


def test_canonical_path_produces_no_row() -> None:
    path = (
        r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
        r"\PDF\AGCC.287-9110-KSB.OD-0001_0_RU.pdf"
    )
    lower_gate = (
        r"C:\rd\9110\KSB\для передачи\05_рев.0_AGCC.287-9110-KSB"
        r"\PDF\AGCC.287-9110-KSB.OD-0001_0_RU.pdf"
    )
    hyphen_gate = (
        r"C:\rd\9110\KSB\Для"
        "\u2010"
        r"передачи\05_рев.0_AGCC.287-9110-KSB"
        r"\PDF\AGCC.287-9110-KSB.OD-0001_0_RU.pdf"
    )
    for candidate in (path, lower_gate, hyphen_gate):
        assert _reason(candidate) == ""
        rows = list_layout_violations(
            records=[_record(candidate, title="9110", mark="KSB")],
            rd_root=RD_ROOT,
        )
        assert rows == ()


def test_grouping_collapses_same_kit_and_reason() -> None:
    paths = (
        r"C:\rd\1715\14_POS\Резерв\a.dwg",
        r"C:\rd\1715\14_POS\Резерв\b.dwg",
        r"C:\rd\1715\14_POS\Резерв\c.dwg",
        r"C:\rd\1715\14_POS\от 2023.05.28\d.dwg",
    )
    records = [
        _record(path, title="1715", mark="POS", file_id=index)
        for index, path in enumerate(paths, start=1)
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    assert report.total_files == 4
    by_reason = {row.reason: row for row in report.rows}
    assert set(by_reason) == {"working_folder", "bad_package_name"}
    assert by_reason["working_folder"].file_count == 3
    assert len(by_reason["working_folder"].example_paths) == 3
    assert by_reason["bad_package_name"].file_count == 1
    assert by_reason["working_folder"].mto_file_count == 0
    assert report.total_mto_files == 0
    assert report.total_kits == 1
    assert len(report.rows) == 2
    working_tally = next(
        item for item in report.by_reason if item[0] == "working_folder"
    )
    assert working_tally == ("working_folder", 3, 0)


def test_fully_lost_kits_and_partial_kit() -> None:
    lost_a = r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.OD-0001_0_RU.pdf"
    lost_b = r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.WIR-0001_0_RU.pdf"
    canonical = (
        r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
        r"\PDF\AGCC.287-9110-KSB.OD-0001_0_RU.pdf"
    )
    extra = (
        r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
        r"\DWG\test\file.dwg"
    )
    records = [
        _record(lost_a, title="1715", mark="PD", file_id=1),
        _record(lost_b, title="1715", mark="PD", file_id=2),
        _record(canonical, title="9110", mark="KSB", file_id=3),
        _record(extra, title="9110", mark="KSB", file_id=4),
        _record(
            r"C:\sq\1715\PD\file.pdf",
            title="1715",
            mark="PD",
            source=SourceKind.SQ,
            file_id=5,
        ),
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    assert report.fully_lost_kits == (("1715", "PD"),)
    lost_row = next(row for row in report.rows if row.title == "1715")
    assert lost_row.file_count == 2
    assert lost_row.reason == "working_folder"
    ksb_rows = [row for row in report.rows if row.mark == "KSB"]
    assert len(ksb_rows) == 1
    assert ksb_rows[0].reason == "extra_subfolder"
    assert ksb_rows[0].file_count == 1
    text = format_layout_report_text(report)
    assert "1715/PD" in text
    assert "Полностью потерянные" in text


def test_write_layout_report_xlsx_has_lost_sheet() -> None:
    records = [
        _record(
            r"C:\rd\1715\PD\Рабочая\file.pdf",
            title="1715",
            mark="PD",
            file_id=1,
        )
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    with tempfile.TemporaryDirectory(prefix="rd_layout_xlsx_") as raw:
        path = Path(raw) / "layout.xlsx"
        write_layout_report_xlsx(report, path)
        payload = path.read_bytes()
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(payload))
        try:
            assert workbook.sheetnames == [
                "Сводка",
                "Нет в каталоге",
                "Что исправить",
            ]
            lost = workbook["Нет в каталоге"]
            values = [tuple(row) for row in lost.iter_rows(values_only=True)]
            assert values[1][:2] == ("1715", "PD")
            work = workbook["Что исправить"]
            headers = next(work.iter_rows(min_row=1, max_row=1, values_only=True))
            assert headers[5:8] == ("Файлов", "из них MTO", "Пакет")
            summary = workbook["Сводка"]
            summary_rows = [tuple(row) for row in summary.iter_rows(values_only=True)]
            assert any(row and row[0] == "MTO вне раскладки" for row in summary_rows)
        finally:
            workbook.close()


def test_mto_counts_and_summary() -> None:
    records = [
        _record(
            r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.MTO-0001_B_RU.xlsx",
            title="1715",
            mark="PD",
            file_id=1,
        ),
        _record(
            r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.OD-0001_B_RU.pdf",
            title="1715",
            mark="PD",
            file_id=2,
        ),
        _record(
            r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.WIR-0001_B_RU.dwg",
            title="1715",
            mark="PD",
            file_id=3,
        ),
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    assert report.total_files == 3
    assert report.total_mto_files == 1
    assert len(report.rows) == 1
    row = report.rows[0]
    assert row.mto_file_count == 1
    assert row.file_count == 3
    assert report.by_reason == (("working_folder", 3, 1),)
    text = format_layout_report_text(report)
    assert "MTO вне раскладки: 1" in text
    assert "MTO=1" in text


def test_mto_rows_sort_before_non_mto_within_lost_split() -> None:
    records = [
        _record(
            r"C:\rd\8441\SKUD\Сила\AGCC.287-8441-EN.OD-0001_0_RU.pdf",
            title="8441",
            mark="EN",
            file_id=1,
        ),
        _record(
            r"C:\rd\8441\SKUD\Сила\AGCC.287-8441-EN.LAY-0001_0_RU.pdf",
            title="8441",
            mark="EN",
            file_id=2,
        ),
        _record(
            r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.MTO-0001_B_RU.xlsx",
            title="1715",
            mark="PD",
            file_id=3,
        ),
        _record(
            r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
            r"\PDF\AGCC.287-9110-KSB.OD-0001_0_RU.pdf",
            title="9110",
            mark="KSB",
            file_id=4,
        ),
        _record(
            r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
            r"\DWG\test\AGCC.287-9110-KSB.MTO-0001_0_RU.xlsx",
            title="9110",
            mark="KSB",
            file_id=5,
        ),
        _record(
            r"C:\rd\9200\SOT\Для передачи\01_рев.0_AGCC.287-9200-SOT"
            r"\PDF\AGCC.287-9200-SOT.OD-0001_0_RU.pdf",
            title="9200",
            mark="SOT",
            file_id=6,
        ),
        _record(
            r"C:\rd\9200\SOT\Для передачи\01_рев.0_AGCC.287-9200-SOT"
            r"\DWG\test\file.dwg",
            title="9200",
            mark="SOT",
            file_id=7,
        ),
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    keys = [(row.title, row.mark, row.mto_file_count > 0) for row in report.rows]
    assert keys == [
        ("1715", "PD", True),
        ("8441", "EN", False),
        ("9110", "KSB", True),
        ("9200", "SOT", False),
    ]


def test_multi_package_hint_shows_count_not_a_name() -> None:
    records = [
        _record(
            r"C:\rd\8630\14_KSB4\03_KSB4\для передачи"
            r"\01_рев.0_AGCC.287-8630-KSB4"
            r"\PDF\AGCC.287-8630-KSB4.OD-0001_0_RU.pdf",
            title="8630",
            mark="KSB4",
            file_id=1,
        ),
        _record(
            r"C:\rd\8630\14_KSB4\03_KSB4\для передачи"
            r"\03_рев.AN-01_AGCC.287-8630-KSB4"
            r"\PDF\AGCC.287-8630-KSB4.OD-0001_AN01_RU.pdf",
            title="8630",
            mark="KSB4",
            file_id=2,
        ),
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    assert len(report.rows) == 1
    assert report.rows[0].reason == "duplicate_level"
    assert report.rows[0].package_hint == "2 передачи"
    assert "01_рев.0_" not in report.rows[0].package_hint
    assert "03_рев.AN-01_" not in report.rows[0].package_hint


def test_single_package_hint_keeps_folder_name() -> None:
    records = [
        _record(
            r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
            r"\DWG\test\a.dwg",
            title="9110",
            mark="KSB",
            file_id=1,
        ),
        _record(
            r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
            r"\DWG\test\b.dwg",
            title="9110",
            mark="KSB",
            file_id=2,
        ),
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    assert len(report.rows) == 1
    assert report.rows[0].package_hint == "05_рев.0_AGCC.287-9110-KSB"


def test_zero_violations_summary_is_plain() -> None:
    path = (
        r"C:\rd\9110\KSB\Для передачи\05_рев.0_AGCC.287-9110-KSB"
        r"\PDF\AGCC.287-9110-KSB.OD-0001_0_RU.pdf"
    )
    report = build_layout_report(
        records=[_record(path, title="9110", mark="KSB")],
        rd_root=RD_ROOT,
    )
    assert report.total_files == 0
    assert report.total_mto_files == 0
    assert report.fully_lost_kits == ()
    text = format_layout_report_summary(report)
    assert text.startswith("MTO вне раскладки: 0")
    assert "Нарушений раскладки нет" in text
    assert "Файлов вне раскладки: 0" in text
    assert "Полностью потерянных комплектов: 0" in text
    assert text.strip()


def test_summary_leads_with_mto_count() -> None:
    records = [
        _record(
            r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.MTO-0001_B_RU.xlsx",
            title="1715",
            mark="PD",
            file_id=1,
        ),
        _record(
            r"C:\rd\1715\PD\Рабочая\AGCC.287-1715-PD.OD-0001_B_RU.pdf",
            title="1715",
            mark="PD",
            file_id=2,
        ),
    ]
    report = build_layout_report(records=records, rd_root=RD_ROOT)
    text = format_layout_report_summary(report)
    assert text.startswith("MTO вне раскладки: 1")
    assert "Файлов вне раскладки: 2" in text
    assert "Полностью потерянных комплектов: 1" in text
    assert "Что исправить:" not in text


def test_default_filename_uses_dotted_date() -> None:
    name = layout_report_default_filename(datetime(2026, 9, 10))
    assert name == "Раскладка_РД_2026.09.10.xlsx"


def main() -> None:
    """Run layout-report assertions without UNC or a live database."""

    test_new_tokens_on_synthetic_paths()
    test_working_and_gate_case_and_unicode_hyphen()
    test_canonical_path_produces_no_row()
    test_grouping_collapses_same_kit_and_reason()
    test_fully_lost_kits_and_partial_kit()
    test_write_layout_report_xlsx_has_lost_sheet()
    test_mto_counts_and_summary()
    test_mto_rows_sort_before_non_mto_within_lost_split()
    test_multi_package_hint_shows_count_not_a_name()
    test_single_package_hint_keeps_folder_name()
    test_zero_violations_summary_is_plain()
    test_summary_leads_with_mto_count()
    test_default_filename_uses_dotted_date()
    print("RD catalog layout report: OK")


if __name__ == "__main__":
    main()
