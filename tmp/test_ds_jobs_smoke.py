"""Offline smoke for DS cockpit jobs and the RFP · Сбор частей panel."""

from __future__ import annotations

import ast
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.ds_compare.ds_compare_config import (
    get_default_gui_paths,
    normalize_gui_paths,
)
from RFQ.rfp_parts.ds_baseline import IdentityDsUnitsConverter
from RFQ.rfp_parts.ds_jobs import (
    CockpitRow,
    DsCockpitSnapshot,
    get_last_ds_cockpit,
    run_ds_baseline_job,
    run_ds_coverage_job,
    run_ds_hybrid_job,
    run_ds_registry_check_job,
)
from RFQ.rfp_parts.ds_registry import (
    CANONICAL_REGISTRY_NAME,
    DATED_REGISTRY_RE,
    DEFAULT_REGISTRY_PATH,
    MODE_WHOLE,
    MIGRATION_REPORT_PREFIX,
    OLD_REGISTRY_NAME,
    STATUS_ACTIVE,
    DsRegistryRelation,
    DsRegistryRow,
    detect_registry_format,
    install_working_registry,
    resolve_latest_registry,
    write_registry_workbook,
)
from RFQ.rfp_parts.ds_rfp_hybrid import (
    HYBRID_REPORT_PREFIX,
    HYBRID_XLSX_NAME,
    IdentityRfpUnitsConverter,
)
from RFQ.units_convert.models import GoogleUnitsIndex

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

DS_HEADER = [
    "№ п/п",
    "Титул",
    "Раздел",
    "Спецификация",
    "RFQ",
    "Наименование Позиций Товара по РД",
    "Код 1С СОУ",
    "Код РД",
    "Наименование Позиций Товара Поставщика",
    "Технические требования (ГОСТ/ ТУ и др.)",
    "Ед. изм.",
    "Кол-во",
]


def _write_legacy(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "Актуальный ДС"
    ws["B1"] = "Старый ДС"
    ws["C1"] = "УЛ"
    ws["D1"] = "Примечание"
    ws["A2"] = 13
    ws["B2"] = ""
    ws["C2"] = "согл УЛ ДС13"
    ws["D2"] = ""
    wb.save(path)
    wb.close()


def _write_registry(path: Path) -> None:
    write_registry_workbook(
        path,
        [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                    ),
                ),
            )
        ],
        spare_rows=1,
    )


def _write_ds_xlsx(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Спека"
    ws.append(DS_HEADER)
    ws.append(
        [1, "8529", "SOS", "spec", "RFQ", "Кабель", "1C", "BCC0000001", "Vendor", "NYM", "шт", 2]
    )
    wb.save(path)
    wb.close()


def _local_ul(root: Path) -> Path:
    ul = root / "ul"
    ul.mkdir(parents=True, exist_ok=True)
    return ul


def _google() -> GoogleUnitsIndex:
    return GoogleUnitsIndex(
        units_by_code={"BCC0000001": "шт"},
        display_units_by_code={"BCC0000001": "шт"},
        display_code_by_code={"BCC0000001": "BCC0000001"},
    )


class DsJobsSmokeTest(unittest.TestCase):
    def test_gui_paths_have_ds_cockpit_keys(self) -> None:
        defaults = get_default_gui_paths()
        self.assertIn("last_ds_trusted_folder", defaults)
        self.assertIn("last_ds_registry_file", defaults)
        self.assertEqual(defaults["last_ds_registry_file"], str(DEFAULT_REGISTRY_PATH))
        normalized = normalize_gui_paths({"last_ds_file": "x.xlsx"})
        self.assertEqual(normalized["last_ds_registry_file"], str(DEFAULT_REGISTRY_PATH))
        self.assertEqual(normalized["last_ds_file"], "x.xlsx")
        emptied = normalize_gui_paths({"last_ds_trusted_folder": ""})
        self.assertEqual(emptied["last_ds_trusted_folder"], "")

    def test_jobs_module_never_calls_unc_replace(self) -> None:
        import RFQ.rfp_parts.ds_jobs as jobs

        tree = ast.parse(Path(jobs.__file__).read_text(encoding="utf-8"))
        imported: list[str] = []
        for node in ast.walk(tree):
            if isinstance(node, ast.ImportFrom):
                imported.extend(alias.name for alias in node.names)
            elif isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
                imported.append(node.func.id)
        self.assertNotIn("backup_and_replace_registry", imported)
        self.assertNotIn("backup_and_replace_registry", jobs.__all__)

    def test_registry_check_new_format(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            registry = root / "Реестр_ДС_УЛ.xlsx"
            reports = root / "reports"
            _write_registry(registry)
            result = run_ds_registry_check_job(registry, reports, root / "missing-ul")
            self.assertTrue(result.success)
            self.assertIn("реестр", result.message.lower())
            self.assertTrue(registry.is_file())

    def test_registry_legacy_migrates_copy_not_source(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            legacy = root / "Реестр_ДС_УЛ.xlsx"
            reports = root / "reports"
            _write_legacy(legacy)
            before = legacy.read_bytes()
            self.assertEqual(detect_registry_format(legacy), "legacy")
            result = run_ds_registry_check_job(legacy, reports, _local_ul(root))
            self.assertTrue(result.success, result.message)
            self.assertEqual(legacy.read_bytes(), before)
            working = reports / "Реестр_ДС_УЛ.xlsx"
            self.assertTrue(working.is_file())
            self.assertEqual(detect_registry_format(working), "new")
            self.assertNotEqual(working.resolve(), DEFAULT_REGISTRY_PATH)
            self.assertFalse((reports / "Реестр_ДС_УЛ_migrated.xlsx").exists())
            self.assertTrue(
                any(path.name.startswith(MIGRATION_REPORT_PREFIX) for path in reports.iterdir())
            )
            self.assertIn("формат=new", result.message)
            cockpit = get_last_ds_cockpit()
            self.assertIsNotNone(cockpit)
            self.assertIn("Рабочий реестр один", cockpit.next_step)
            self.assertIn(working.name, cockpit.next_step)

    def test_legacy_in_home_is_renamed_old(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            home = Path(raw)
            legacy = home / CANONICAL_REGISTRY_NAME
            _write_legacy(legacy)
            before = legacy.read_bytes()
            result = run_ds_registry_check_job(legacy, home, _local_ul(home))
            self.assertTrue(result.success, result.message)
            old = home / OLD_REGISTRY_NAME
            self.assertTrue(old.is_file())
            self.assertEqual(old.read_bytes(), before)
            self.assertEqual(detect_registry_format(legacy), "new")
            self.assertEqual(resolve_latest_registry(home), legacy)
            cockpit = get_last_ds_cockpit()
            self.assertIsNotNone(cockpit)
            self.assertIn(OLD_REGISTRY_NAME, cockpit.next_step)

    def test_locked_canonical_keeps_five_dated_copies(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            home = Path(raw)
            real_replace = os.replace

            def fake_replace(src, dst):
                if Path(dst).name == CANONICAL_REGISTRY_NAME:
                    raise PermissionError(5, "locked")
                return real_replace(src, dst)

            row = DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                    ),
                ),
            )
            with patch("RFQ.rfp_parts.ds_registry.os.replace", side_effect=fake_replace):
                last = None
                for _ in range(6):
                    last = install_working_registry([row], home)
            self.assertIsNotNone(last)
            assert last is not None
            self.assertEqual(last.mode, "dated")
            dated = [
                path
                for path in home.iterdir()
                if DATED_REGISTRY_RE.match(path.name)
            ]
            self.assertEqual(len(dated), 5)
            self.assertEqual(resolve_latest_registry(home), last.path)
            self.assertFalse((home / CANONICAL_REGISTRY_NAME).exists())

    def test_coverage_warns_if_rfp_missing(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            registry = root / "registry.xlsx"
            ds_root = root / "ds"
            ds_root.mkdir()
            _write_registry(registry)
            _write_ds_xlsx(ds_root / "ДС_13" / "spec.xlsx")
            missing_rfp = root / "no_rfp"
            result = run_ds_coverage_job(
                ds_root, registry, None, _local_ul(root), missing_rfp
            )
            self.assertTrue(result.success, result.message)
            self.assertIn("WARN", result.message)

    def test_baseline_emits_ds_progress(self) -> None:
        import io
        from contextlib import redirect_stdout

        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            registry = root / "registry.xlsx"
            ds_root = root / "ds"
            reports = root / "out"
            _write_registry(registry)
            _write_ds_xlsx(ds_root / "ДС_13" / "spec.xlsx")
            buf = io.StringIO()
            with redirect_stdout(buf):
                result = run_ds_baseline_job(
                    ds_root,
                    registry,
                    reports,
                    _local_ul(root),
                    converter=IdentityDsUnitsConverter(),
                    google_index=_google(),
                )
            self.assertTrue(result.success, result.message)
            log = buf.getvalue()
            self.assertIn("[ds progress] START: аудит ДС", log)
            self.assertIn("[ds progress] PROGRESS: ДС", log)
            self.assertIn("[ds progress] FRACTION:", log)

    def test_baseline_and_hybrid_with_temp_dirs(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            registry = root / "registry.xlsx"
            ds_root = root / "ds"
            reports = root / "out"
            rfp_root = root / "rfp"
            rfp_root.mkdir()
            _write_registry(registry)
            _write_ds_xlsx(ds_root / "ДС_13" / "spec.xlsx")
            conv = IdentityDsUnitsConverter()
            google = _google()
            baseline = run_ds_baseline_job(
                ds_root,
                registry,
                reports,
                _local_ul(root),
                converter=conv,
                google_index=google,
            )
            self.assertTrue(baseline.success, baseline.message)
            self.assertIsNotNone(baseline.result_path)
            hybrid = run_ds_hybrid_job(
                ds_root,
                registry,
                reports,
                _local_ul(root),
                rfp_root,
                converter=conv,
                rfp_converter=IdentityRfpUnitsConverter(),
                google_index=google,
            )
            self.assertTrue(hybrid.success, hybrid.message)
            self.assertTrue((reports / HYBRID_XLSX_NAME).is_file())
            self.assertTrue(
                any(
                    path.name.startswith(HYBRID_REPORT_PREFIX)
                    for path in reports.rglob("*.xlsx")
                )
            )


class RfpPartsDsCockpitPanelSmokeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def test_instantiate_and_fill_snapshot(self) -> None:
        from RFQ.rfp_parts import ds_jobs as ds_jobs_mod
        from ds_compare_center.rfp_parts_panel import RfpPartsPanel

        ds_jobs_mod._last_cockpit = None
        panel = RfpPartsPanel()
        self.app.processEvents()
        self.assertTrue(hasattr(panel, "monitor"))
        self.assertTrue(hasattr(panel, "splitter"))
        from PySide6.QtWidgets import QPushButton

        texts = [btn.text() for btn in panel.findChildren(QPushButton)]
        self.assertIn("Собрать свод только из ДС", texts)
        self.assertIn("Наложить RFP на группы", texts)
        self.assertIn("Только имена и покрытие", texts)
        self.assertIn("Проверить реестр", texts)
        self.assertIn("Собрать свод частей RFP", texts)
        self.assertIn("Подставить копию роботу", texts)
        self.assertIn("Папка реестра", texts)
        panel.show()
        self.app.processEvents()
        use_btns = [
            btn
            for btn in panel.findChildren(QPushButton)
            if btn.text() == "Подставить копию роботу"
        ]
        self.assertEqual(len(use_btns), 1)
        btn_use = use_btns[0]
        self.assertFalse(btn_use.isVisible())
        snapshot = DsCockpitSnapshot(
            kind="coverage",
            summary="Покрытие: групп=1, файлов ДС=1, RFP=0, без RFP=1, ERROR=0, WARN=1",
            registry_path=Path("registry.xlsx"),
            registry_format="new",
            error_count=0,
            warn_count=1,
            match_count=1,
            empty_code=1,
            migrated_registry_path=Path("migrated.xlsx"),
            next_step="Канон UNC не заменён",
            registry_rows=[
                CockpitRow(
                    cells=("13", "Активен", "ДС13", "13", "согл УЛ ДС13", "Вся ДС", ""),
                    tone="match",
                )
            ],
            group_rows=[
                CockpitRow(
                    cells=("ДС13", "13", "13", "согл УЛ ДС13", "нет", "MATCH"),
                    tone="match",
                )
            ],
            file_rows=[
                CockpitRow(cells=("ДС", "spec.xlsx", "13", "OK"), tone="ok"),
            ],
            coverage_rows=[
                CockpitRow(
                    cells=("ДС13", "spec.xlsx", "нет RFP", "согл УЛ ДС13", "нет файла RFP"),
                    tone="warn",
                )
            ],
        )
        panel.fill_from_snapshot(snapshot)
        self.app.processEvents()
        self.assertTrue(btn_use.isVisible())
        self.assertEqual(panel._table_registry.rowCount(), 1)
        self.assertEqual(panel._table_coverage.rowCount(), 1)
        self.assertIn("ERROR=0", panel._cockpit_summary.text())
        self.assertIn("MATCH=1", panel._cockpit_summary.text())
        bg = panel._table_coverage.item(0, 0).background().color().name()
        self.assertEqual(bg.lower(), "#fff8c5")
        match_bg = panel._table_registry.item(0, 0).background().color().name()
        self.assertEqual(match_bg.lower(), "#e6f4ea")
        panel.close()
        panel.deleteLater()
        self.app.processEvents()


if __name__ == "__main__":
    unittest.main(verbosity=2)
