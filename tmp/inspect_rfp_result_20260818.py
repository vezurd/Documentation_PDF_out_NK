"""Inspect 2026.08.18.16.55 result vs rfp_parts_net layout (read-only)."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tmp" / "inspect_20260818"
TMP.mkdir(parents=True, exist_ok=True)

sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")


def main() -> int:
    from openpyxl import load_workbook

    print("=== copied files ===")
    for p in sorted(TMP.iterdir()):
        print(f"{p.name}\t{p.stat().st_size}")

    reports = Path(r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP\RFP сводный файл")
    print(f"\n=== RFP сводный файл exists={reports.exists()} ===")
    if reports.exists():
        dirs = [d for d in reports.iterdir() if d.is_dir()]
        dirs.sort(key=lambda d: d.name, reverse=True)
        for d in dirs[:8]:
            net = d / "rfp_parts_net.xlsx"
            print(f"  {d.name}\tnet_exists={net.exists()}\tsize={net.stat().st_size if net.exists() else '-'}")

    net_candidates = list(TMP.glob("rfp_parts_net.xlsx"))
    if not net_candidates:
        print("\nNo local rfp_parts_net.xlsx copy")
    else:
        net = net_candidates[0]
        wb = load_workbook(net, read_only=True, data_only=True)
        print(f"\n=== net sheets: {wb.sheetnames} ===")
        ws = wb[wb.sheetnames[0]]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            rows.append(row)
            if i >= 4:
                break
        wb.close()
        for i, row in enumerate(rows, start=1):
            print(f"net row {i}:")
            for j, v in enumerate(row[:18]):
                print(f"  [{j}] {v!r}")

    title_files = list(TMP.glob("*title_system*.xlsx")) + list(TMP.glob("*сравнение_title*"))
    if not title_files:
        title_files = [p for p in TMP.iterdir() if "title" in p.name.lower() or p.stat().st_size < 20000 and p.suffix == ".xlsx"]
    print(f"\n=== title/small xlsx ===")
    for p in TMP.iterdir():
        if p.suffix.lower() == ".xlsx" and p.stat().st_size < 100_000:
            wb = load_workbook(p, read_only=True, data_only=True)
            print(f"\nFILE {p.name} sheets={wb.sheetnames}")
            ws = wb.active
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                print(" ", row[:12])
                if i >= 15:
                    break
            wb.close()

    match_files = [p for p in TMP.iterdir() if p.suffix.lower() == ".xlsx" and p.stat().st_size > 1_000_000]
    if match_files:
        match = match_files[0]
        print(f"\n=== match file {match.name} ===")
        wb = load_workbook(match, read_only=True, data_only=True)
        print("sheets:", wb.sheetnames)
        ws = wb[wb.sheetnames[0]]
        headers = None
        n = 0
        empty_code = 0
        empty_name = 0
        empty_values = 0
        has_code = 0
        has_mto = 0
        sample_empty = []
        sample_filled = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 1:
                headers = list(row)
                print("headers (first 25):")
                for j, h in enumerate(headers[:25]):
                    print(f"  [{j}] {h!r}")
                continue
            n += 1
            # Heuristic: col A often DS_NAME / title; find CODE/NAME/VALUES by header
            code = name = values = code_mto = None
            if headers:
                def col(name_part: str):
                    for j, h in enumerate(headers):
                        if h and name_part.lower() in str(h).lower():
                            return row[j] if j < len(row) else None
                    return None
            # counts using first-row headers
            def by_exact(substrs):
                for j, h in enumerate(headers or []):
                    hs = str(h or "")
                    if any(s in hs for s in substrs):
                        return row[j] if j < len(row) else None
                return None

            code = by_exact(["Код РД", "CODE"])
            name = by_exact(["Наименование"])
            values = by_exact(["Кол-во RFP", "VALUES", "Кол."])
            code_mto = by_exact(["Код МТО", "CODE_MTO"])
            if code not in (None, ""):
                has_code += 1
                if len(sample_filled) < 5:
                    sample_filled.append(row[:12])
            else:
                empty_code += 1
                if len(sample_empty) < 5:
                    sample_empty.append(row[:12])
            if name in (None, ""):
                empty_name += 1
            if values in (None, ""):
                empty_values += 1
            if code_mto not in (None, ""):
                has_mto += 1
            if n >= 80000:
                break
        wb.close()
        print(f"data rows scanned: {n}")
        print(f"has_code={has_code} empty_code={empty_code} empty_name={empty_name} empty_values={empty_values} has_mto={has_mto}")
        print("sample filled CODE:")
        for s in sample_filled:
            print(" ", s)
        print("sample empty CODE:")
        for s in sample_empty:
            print(" ", s)

    for p in TMP.glob("ul_compare_report*.txt"):
        print(f"\n=== {p.name} (head) ===")
        text = p.read_text(encoding="utf-8", errors="replace")
        print("\n".join(text.splitlines()[:80]))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
