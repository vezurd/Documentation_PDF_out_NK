"""Add a Russian format-hint sheet to СписокДС.xlsx (does not change Лист2 data)."""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.tags_rfp_compare.ds_manager_matrix import load_ds_manager_matrix

LOCAL = ROOT / "tmp" / "СписокДС_copy.xlsx"
UNC = Path(r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP\Список ДС - Фамилии МП.xlsx")
HINT_TITLE = "Формат заполнения"

_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_FILL_TITLE = PatternFill("solid", fgColor="1F4E79")
_FILL_OK = PatternFill("solid", fgColor="C6EFCE")
_FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
_FILL_WARN = PatternFill("solid", fgColor="FFF2CC")
_FILL_HEAD = PatternFill("solid", fgColor="D6DCE4")
_FILL_NOTE = PatternFill("solid", fgColor="DEEBF7")
_FONT_WHITE = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
_FONT_H = Font(name="Calibri", size=11, bold=True)
_FONT_N = Font(name="Calibri", size=11)
_WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")


def _comment(text: str, *, width: int = 420, height: int = 180) -> Comment:
    note = Comment(text, "Подсказка")
    note.width = width
    note.height = height
    return note


def _write_cell(
    sheet: Worksheet,
    row: int,
    col: int,
    value: object,
    *,
    fill: PatternFill | None = None,
    font: Font | None = None,
    comment: str | None = None,
    comment_size: tuple[int, int] | None = None,
) -> None:
    cell = sheet.cell(row, col, value)
    cell.alignment = _WRAP
    cell.border = _THIN
    cell.font = font or _FONT_N
    if fill is not None:
        cell.fill = fill
    if comment:
        width, height = comment_size or (420, 180)
        cell.comment = _comment(comment, width=width, height=height)


def _build_hint(sheet: Worksheet) -> None:
    sheet.sheet_properties.tabColor = "ED7D31"
    widths = {1: 28, 2: 16, 3: 42, 4: 28, 5: 22, 6: 62}
    for idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(idx)].width = width

    sheet.merge_cells("A1:F1")
    _write_cell(
        sheet,
        1,
        1,
        "Как заполнять столбец A «Имя ДС» (лист Лист2)",
        fill=_FILL_TITLE,
        font=_FONT_WHITE,
    )
    sheet.row_dimensions[1].height = 24
    for col in range(2, 7):
        sheet.cell(1, col).fill = _FILL_TITLE

    intro = (
        "Этот файл — матрица менеджеров ДС, не справочник фактических номеров для УЛ.\n"
        "Программа читает только лист, где в шапке есть «Имя ДС» (или «№ ДС») и «Фамилия» "
        "(или «МП»). Лист позиций с «№ позиции» не читается. Этот лист-подсказку робот "
        "игнорирует — не ставьте в первой строке одновременно заголовки «Имя ДС» и «Фамилия».\n\n"
        "Столбец A нужен, чтобы по префиксу имени файла RFP подставить фамилию в Step4 "
        "(колонка «Менеджер ДС») и путь к файлу. Вкладка «RFP · ДС ↔ УЛ» эту книгу не читает: "
        "она сравнивает имена файлов в RFP_Зиновьев с папками «согл УЛ ДС…».\n\n"
        "Правило номера: число сразу после «ДС» — фактический ДС. Оно должно совпадать "
        "с папкой УЛ и с началом имени файла. Число после «_» — порядковый/информационный "
        "(корректировка), в ключ посадки УЛ не входит. Буква ревизии (А/Б) в конце допустима: "
        "ДС24_92Б совпадёт с файлом ДС24_92.\n\n"
        "На Лист2 сейчас много составных имён записано наоборот относительно файлов "
        "(в таблице ДС92_24Б, а файл называется ДС24_92. …). Из-за этого менеджер в Step4 "
        "не находится, хотя вкладка ДС↔УЛ при этом зелёная: файлы и папки УЛ согласованы между собой."
    )
    sheet.merge_cells("A2:F8")
    _write_cell(sheet, 2, 1, intro, fill=_FILL_NOTE)
    sheet.row_dimensions[2].height = 48
    for row in range(2, 9):
        sheet.row_dimensions[row].height = 26
        for col in range(1, 7):
            sheet.cell(row, col).fill = _FILL_NOTE
            sheet.cell(row, col).border = _THIN

    headers = (
        "Пример для столбца A",
        "Верно?",
        "Что видит программа",
        "Папка УЛ / файл RFP",
        "Где используется",
        "Пояснение",
    )
    header_row = 10
    for col, title in enumerate(headers, start=1):
        _write_cell(sheet, header_row, col, title, fill=_FILL_HEAD, font=_FONT_H)
    sheet.row_dimensions[header_row].height = 28
    sheet.auto_filter.ref = "A10:F22"
    sheet.freeze_panes = "A11"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    examples: list[tuple] = [
        (
            "ДС24",
            "Да",
            "фактический = 24, порядковый = 24",
            "согл УЛ ДС24; файл ДС24. AGCC…",
            "A → МП; имена → вкладка ДС↔УЛ",
            "Обычная ДС без корректировки. Пишите с префиксом ДС, без пробелов.",
            _FILL_OK,
            (
                "Верный простой формат.\n"
                "Фактический номер = 24.\n"
                "Папка УЛ: согл УЛ ДС24.\n"
                "В Step4 «Фактический ДС» = ДС24, «Порядковый ДС» = ДС24.\n"
                "Вкладка RFP · ДС ↔ УЛ эту ячейку не читает — она смотрит имя файла."
            ),
        ),
        (
            "ДС24_92",
            "Да",
            "фактический = 24, порядковый = 92 (корректировка)",
            "согл УЛ ДС24; файл ДС24_92. AGCC…",
            "A → МП по префиксу файла",
            "Составное имя: сначала фактический (как в УЛ), затем порядковый.",
            _FILL_OK,
            (
                "Верный составной формат.\n"
                "Шаблон: ДС{фактический}_{порядковый}{буква}.\n"
                "Пример живого файла: ДС24_92. AGCC.287-…xlsx\n"
                "Программа возьмёт фактический 24 и сопоставит с папкой «согл УЛ ДС24».\n"
                "Число 92 информационное, в ключ посадки УЛ не входит."
            ),
        ),
        (
            "ДС24_92Б",
            "Да",
            "фактический = 24; буква Б снимается при поиске МП",
            "файл ДС24_92. … всё равно найдёт эту строку",
            "A → МП (алиас без буквы)",
            "Буква ревизии в конце допустима и не ломает поиск менеджера.",
            _FILL_OK,
            (
                "Буква А/Б в конце — это ревизия, не часть номера.\n"
                "Строка ДС24_92Б совпадёт с файлом ДС24_92 и с ДС24_92Б.\n"
                "Не пишите букву в середине и не отделяйте её пробелом."
            ),
        ),
        (
            "ДС01",
            "Да",
            "нули снимаются → ключ ДС1",
            "файл ДС1_Госфин_… или ДС01. …",
            "A → МП",
            "Ведущие нули можно оставлять: ДС01 = ДС1.",
            _FILL_OK,
            "Ведущие нули программа отбрасывает. ДС01 и ДС1 — одна ДС для поиска менеджера.",
        ),
        (
            "4905",
            "Да, лучше ДС4905",
            "к числу дописывается префикс ДС → ДС4905",
            "согл УЛ ДС4905; файл ДС4905.xlsx",
            "A → МП",
            "Голое число допустимо. Для единообразия пишите ДС4905.",
            _FILL_WARN,
            (
                "Голое число 4905 программа превратит в ДС4905.\n"
                "На вкладке ДС↔УЛ эта ячейка всё равно не участвует.\n"
                "Предпочтительная запись в столбце A: ДС4905 (с префиксом)."
            ),
        ),
        (
            "ДС92_24Б",
            "Нет (переворот)",
            "фактический = 92, порядковый = 24 — наоборот к файлу ДС24_92",
            "файла ДС92_24Б нет; УЛ ДС92 нет; есть согл УЛ ДС24",
            "A не совпадёт с файлом → «Нет в матрице МП»",
            "Так сейчас на Лист2. Нужно ДС24_92 или ДС24_92Б.",
            _FILL_BAD,
            (
                "ОШИБКА ПОРЯДКА ЧИСЕЛ.\n"
                "В таблице записано ДС92_24Б, а файл называется ДС24_92. AGCC…\n"
                "Папка УЛ: согл УЛ ДС24 (не ДС92).\n"
                "Программа считает фактическим ПЕРВОЕ число после ДС.\n"
                "Исправьте на ДС24_92Б (или ДС24_92).\n"
                "Иначе Step4 не найдёт менеджера, хотя вкладка ДС↔УЛ будет зелёной."
            ),
        ),
        (
            "ДС14_48",
            "Да (как в файле)",
            "фактический = 14",
            "согл УЛ ДС14; файл ДС14_48. AGCC…",
            "A → МП",
            "На Лист2 сейчас стоит перевёрнутое ДС48_14 — замените на этот вид.",
            _FILL_OK,
            (
                "Живой файл: ДС14_48. AGCC.287-0000-12.4.1-RFP-0007_0_RU.xlsx\n"
                "В таблице ошибочно ДС48_14.\n"
                "Правило: копируйте префикс имени файла до точки+пробела."
            ),
        ),
        (
            "47/13А",
            "Нет",
            "слэш станет подчёркиванием → ДС47_13А",
            "не совпадёт с файлами ДС13_Приложение…",
            "устаревший слэш",
            "Слэш больше не используйте. Пишите ДС{фактический}_{порядковый}.",
            _FILL_BAD,
            (
                "Старый формат чеклиста «47/13А» программа превратит в ДС47_13А.\n"
                "Это не имя файла и не папка УЛ.\n"
                "Пишите так же, как начинается файл в RFP_Зиновьев."
            ),
        ),
        (
            "ДС13_Приложение",
            "Осторожно",
            "фактический всё же 13; ключ поиска длинный",
            "файлы ДС13_Приложение №1…",
            "A → МП только если префикс совпал",
            "Для приложений достаточно ДС13 — сработает startswith при сверке папки, "
            "но поиск менеджера по точному ярлыку файла надёжнее с общим ДС13.",
            _FILL_WARN,
            (
                "Имя файла: ДС13_Приложение №1 – «Детализированная спецификация №3»_6600.xlsx\n"
                "Ярлык, который видит робот: ДС13_ПРИЛОЖЕНИЕ (обрезается по первым «словам»).\n"
                "Строка ДС13 в столбце A надёжнее, чем ДС13_ или ДС47_ПРИЛОЖЕНИЕ.\n"
                "Не ставьте чужой номер (ДС47) — фактический у этих файлов 13."
            ),
        ),
        (
            "ДС4905_1",
            "Нет (суффикс копии)",
            "без «. » в имени это не корректировка; с «. » робот решит sequential=1",
            "файлы ДС4905.xlsx / ДС4905_(2).xlsx",
            "путаница с копией файла",
            "Пишите ДС4905. Не используйте _1/_2 как номер ДС.",
            _FILL_BAD,
            (
                "ДС4905_1.xlsx — это копия файла, а не «фактический 1».\n"
                "Корректировка должна иметь точку и пробел: «ДС14_48. AGCC…».\n"
                "В столбце A для этой ДС достаточно ДС4905."
            ),
        ),
        (
            "ДС7_ГФ",
            "Нет",
            "ключ ДС7_ГФ не равен ярлыку файла ДС7_ГОСФИН",
            "файл ДС7_Госфин.xlsx; папка согл УЛ ДС7 ГФ …",
            "A не совпадёт",
            "Для госфина пишите ДС7 (номер). Пометку ГФ в столбец A не копируйте.",
            _FILL_BAD,
            (
                "Папка УЛ «согл УЛ ДС7 ГФ 2 тит»: фактический = 7, пометка ГФ папку не отменяет.\n"
                "Файл: ДС7_Госфин.xlsx → ярлык ДС7_ГОСФИН.\n"
                "Строка ДС7 в столбце A совпадёт (простой номер).\n"
                "Строка ДС7_ГФ — нет, это другой ключ."
            ),
        ),
        (
            "ДС 24 _ 92 Б",
            "Нет",
            "пробелы схлопнутся, но лучше не полагаться",
            "—",
            "A",
            "Без пробелов: ДС24_92Б. Кириллическая С в «ДС», не латинская C.",
            _FILL_BAD,
            "Пробелы внутри имени лучше не ставить. Буква ДС — кириллица. Латинская DS/C не распознается.",
        ),
    ]

    start = 11
    for offset, item in enumerate(examples):
        (
            example,
            ok,
            sees,
            paths,
            where,
            explain,
            fill,
            note,
        ) = item
        row = start + offset
        sheet.row_dimensions[row].height = 48
        _write_cell(
            sheet,
            row,
            1,
            example,
            fill=fill,
            font=_FONT_H,
            comment=note,
            comment_size=(480, 220),
        )
        _write_cell(sheet, row, 2, ok, fill=fill)
        _write_cell(sheet, row, 3, sees, fill=fill)
        _write_cell(sheet, row, 4, paths, fill=fill)
        _write_cell(sheet, row, 5, where, fill=fill)
        _write_cell(sheet, row, 6, explain, fill=fill)

    map_row = start + len(examples) + 1
    sheet.merge_cells(start_row=map_row, start_column=1, end_row=map_row, end_column=6)
    _write_cell(
        sheet,
        map_row,
        1,
        "Что переписать на Лист2 (снимок 03.09.2026). Данные на Лист2 этот лист не меняет — правьте вручную.",
        fill=_FILL_HEAD,
        font=_FONT_H,
    )
    sheet.row_dimensions[map_row].height = 28

    map_headers = (
        "Сейчас в столбце A",
        "Верно?",
        "Как должно быть (как файл)",
        "Фактический ДС / папка УЛ",
        "Порядковый (хвост)",
        "Почему",
    )
    map_header_row = map_row + 1
    for col, title in enumerate(map_headers, start=1):
        _write_cell(sheet, map_header_row, col, title, fill=_FILL_HEAD, font=_FONT_H)

    inversions = [
        ("ДС48_14", "ДС14_48", "14", "48"),
        ("ДС61_15", "ДС15_61", "15", "61"),
        ("ДС75_16А", "ДС16_75 или ДС16_75А", "16", "75"),
        ("ДС76_17А", "ДС17_76 или ДС17_76А", "17", "76"),
        ("ДС64_20А", "ДС20_64 или ДС20_64А", "20", "64"),
        ("ДС68_21", "ДС21_68", "21", "68"),
        ("ДС89_22А", "ДС22_89 или ДС22_89А", "22", "89"),
        ("ДС92_24Б", "ДС24_92 или ДС24_92Б", "24", "92"),
        ("ДС72_26А", "ДС26_72 или ДС26_72А", "26", "72"),
        ("ДС87_30А", "ДС30_87 или ДС30_87А", "30", "87"),
        ("ДС98_35Б", "ДС35_98 или ДС35_98Б", "35", "98"),
        ("ДС73_36", "ДС36_73", "36", "73"),
        ("ДС88_37А", "ДС37_88 или ДС37_88А", "37", "88"),
        ("ДС97_41А", "ДС41_97 или ДС41_97А", "41", "97"),
        ("ДС100_45А", "ДС45_100 или ДС45_100А", "45", "100"),
        ("ДС90_51А", "ДС51_90 или ДС51_90А", "51", "90"),
        ("ДС103_66А", "ДС66_103 или ДС66_103А", "66", "103"),
        ("ДС96_70А", "ДС70_96 или ДС70_96А", "70", "96"),
    ]
    first_data = map_header_row + 1
    for offset, (wrong, right, actual, sequential) in enumerate(inversions):
        row = first_data + offset
        sheet.row_dimensions[row].height = 22
        _write_cell(
            sheet,
            row,
            1,
            wrong,
            fill=_FILL_BAD,
            comment=(
                f"Сейчас в столбце A: {wrong}.\n"
                f"Файл в RFP_Зиновьев начинается с {right.split()[0]}.\n"
                f"Папка УЛ: согл УЛ ДС{actual}.\n"
                f"Замените на {right}."
            ),
        )
        _write_cell(sheet, row, 2, "Нет, переворот", fill=_FILL_BAD)
        _write_cell(sheet, row, 3, right, fill=_FILL_OK)
        _write_cell(sheet, row, 4, f"ДС{actual} / согл УЛ ДС{actual}", fill=_FILL_OK)
        _write_cell(sheet, row, 5, sequential, fill=_FILL_WARN)
        _write_cell(
            sheet,
            row,
            6,
            "Первое число после ДС должно быть как в папке УЛ и в имени файла.",
            fill=_FILL_BAD,
        )

    footer_row = first_data + len(inversions) + 1
    sheet.merge_cells(
        start_row=footer_row, start_column=1, end_row=footer_row + 3, end_column=6
    )
    footer = (
        "Куда что идёт.\n"
        "1) Столбец A Лист2 → поиск менеджера и сверка «есть ли файл» с папкой RFP_Зиновьев "
        "(не fatal). Ключ должен совпасть с префиксом файла (ДС24_92), а не с «человеческим» "
        "порядком «сначала старый номер».\n"
        "2) Имя файла RFP → колонки Step4 «Порядковый ДС» (весь префикс) и «Фактический ДС» "
        "(первое число после ДС). Отсюда же ключ посадки УЛ (actual, титул, марка, код).\n"
        "3) Имя папки УЛ «согл УЛ ДС24» → фактический 24. Вкладка RFP · ДС ↔ УЛ стыкует "
        "файл и папку только по этому числу. СписокДС.xlsx туда не подмешивается.\n"
        "4) Не копируйте в A хвосты копии файла (_1, _(2)), «Приложение», «ГФ», «AGCC»."
    )
    _write_cell(sheet, footer_row, 1, footer, fill=_FILL_NOTE)
    sheet.row_dimensions[footer_row].height = 36
    for row in range(footer_row, footer_row + 4):
        for col in range(1, 7):
            sheet.cell(row, col).fill = _FILL_NOTE
            sheet.cell(row, col).border = _THIN


def main() -> int:
    src = LOCAL if LOCAL.exists() else UNC
    if not src.exists():
        print(f"ERROR: workbook not found: {src}")
        return 1

    workbook = load_workbook(src)
    if HINT_TITLE in workbook.sheetnames:
        del workbook[HINT_TITLE]
    sheet = workbook.create_sheet(HINT_TITLE)
    _build_hint(sheet)

    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    data = payload.getvalue()
    LOCAL.write_bytes(data)
    print(f"wrote local {LOCAL} bytes={len(data)}")

    matrix = load_ds_manager_matrix(LOCAL)
    if matrix.load_error:
        print(f"ERROR: matrix load after hint sheet: {matrix.load_error}")
        return 1
    print(f"matrix still reads Лист2: entries={len(matrix.entries)}")

    try:
        UNC.write_bytes(data)
        print(f"wrote UNC {UNC}")
    except OSError as exc:
        print(f"WARN: UNC not written ({exc}). Close Excel if the file is open, then copy {LOCAL}")
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
