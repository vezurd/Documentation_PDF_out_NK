"""Read-only: compare sheet shape, comments and fills of two Step4 xlsx files."""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import openpyxl

DIR = Path(__file__).resolve().parent / "step4_compare"
OLD = DIR / "Шаг4_Сопоставление_RFP_MTO_20260909_131931.xlsx"
NEW = DIR / "Шаг4_Сопоставление_RFP_MTO_20260909_134240.xlsx"


def describe(path: Path) -> dict[str, object]:
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        info: dict[str, object] = {"sheets": wb.sheetnames}
        for sheet in wb.worksheets:
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            comments = 0
            fills: Counter[str] = Counter()
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.comment is not None:
                        comments += 1
                    fill = cell.fill
                    if fill is not None and fill.fgColor is not None:
                        fills[str(fill.fgColor.rgb)] += 1
            info[sheet.title] = {
                "dims": f"{sheet.max_row}x{sheet.max_column}",
                "headers": headers,
                "comments": comments,
                "fills": dict(fills.most_common(12)),
            }
        return info
    finally:
        wb.close()


def main() -> None:
    old = describe(OLD)
    new = describe(NEW)
    print(f"sheets old: {old['sheets']}")
    print(f"sheets new: {new['sheets']}")
    for title in old["sheets"]:  # type: ignore[union-attr]
        o = old.get(title, {})
        n = new.get(title, {})
        print(f"\n=== лист «{title}» ===")
        print(f"  dims:     old={o.get('dims')}  new={n.get('dims')}")
        print(f"  comments: old={o.get('comments')}  new={n.get('comments')}")
        oh = o.get("headers") or []
        nh = n.get("headers") or []
        print(f"  columns:  old={len(oh)}  new={len(nh)}  equal={oh == nh}")
        if oh != nh:
            print(f"    only_old={[h for h in oh if h not in nh]}")
            print(f"    only_new={[h for h in nh if h not in oh]}")
        of = o.get("fills") or {}
        nf = n.get("fills") or {}
        if of != nf:
            print("  fills отличаются:")
            for key in sorted(set(of) | set(nf)):
                if of.get(key) != nf.get(key):
                    print(f"    {key}: old={of.get(key)} new={nf.get(key)}")
        else:
            print("  fills: совпадают")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    main()
