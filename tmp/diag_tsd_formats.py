"""Diagnose TSD column offsets and SO-PL layouts (read-only)."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl

from RFQ.ds_compare.tsd_packing_load import (
    DEFAULT_TSD_PACKING_ROOT,
    collect_tsd_files,
    classify_tsd_sheets,
    read_workbook_sheet_names,
)

OUT = ROOT / "tmp" / "tsd_format_diag.txt"

# Candidates: List1 companions + known empty Single after VENDOR fix
LIST1_TIPS = [
    "PL_2076961.4_2511.03",
    "PL_2076961.5_2512.01",
    "PL_2076961.12_2893.17",
    "PL_2076961.9_2709.16",
    "PL_2076961.10_2891.12",
    "PL-АГХК-ПЭППиЛАО-2076961.15.6",
    "PL_2076961.32_3705.04",
    "PL_2076961.12_2893.16",  # empty blank
    "PL_2076961.20_3220.10",  # empty blank
    "PL-АГХК-ПЭППиЛАО-2076961.15.34",  # Single1=0 in earlier report
]

# Header markers for standard packing list (col B = index 1)
STD_HEADER_B = "код рд"
# SO-PL style: CODE in G = index 6
SOPL_HINTS = ("name of zip", "agcc.323", "so - pl")


def _find_header_row(ws, max_row: int = 40, max_col: int = 80) -> tuple[int | None, list]:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True), start=1):
        texts = [str(c).strip().lower() if c is not None else "" for c in row]
        joined = " | ".join(texts)
        if "код рд" in joined or "code rd" in joined or "name of zip" in joined.lower() or "code" == texts[6] if len(texts) > 6 else False:
            return i, list(row)
        if any("наименование" in t and "name" in t for t in texts) and any("код" in t for t in texts):
            return i, list(row)
    return None, []


def _col_map_from_header(header_row: list) -> dict[str, int]:
    """Map semantic keys to 0-based indices by header substrings."""
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        t = str(cell).lower().replace("\n", " ")
        if idx not in mapping.values():
            pass
        if ("код рд" in t or "code rd" in t) and "CODE" not in mapping:
            mapping["CODE"] = idx
        elif ("спецификац" in t or "ол" in t[:20]) and "SPEC" not in mapping and "код рд" not in t:
            if "спецификац" in t or "№ спецификации" in t or "name of document" in t:
                mapping["SPEC"] = idx
        elif ("tag" in t or "лини" in t) and "TAGS" not in mapping:
            mapping["TAGS"] = idx
        elif ("наименован" in t or (t.strip() == "name") or "description of" in t) and "NAME" not in mapping:
            mapping["NAME"] = idx
        elif ("количество" in t or "quantity" in t) and "qty" not in t.lower()[20:] and "VALUES" not in mapping:
            if "км" not in t or "cable" not in t:
                mapping["VALUES"] = idx
        elif ("единица" in t or "unit of measurement" in t or t.strip() in ("uom", "units")) and "UNITS" not in mapping:
            mapping["UNITS"] = idx
        elif ("завод изготовитель" in t or "manufacturer" in t) and "VENDOR" not in mapping:
            mapping["VENDOR"] = idx
        elif ("title" in t and "system" not in t) and "TITLE" not in mapping:
            mapping["TITLE"] = idx
        elif ("system" in t or "система" in t) and "SYSTEM" not in mapping:
            mapping["SYSTEM"] = idx
    # second pass for Name of document
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        t = str(cell).lower().replace("\n", " ")
        if "name of document" in t:
            mapping["SPEC"] = idx
        if t.strip() in ("code",) and "CODE" not in mapping:
            mapping["CODE"] = idx
        if t.strip() in ("tag",) and "TAGS" not in mapping:
            mapping["TAGS"] = idx
        if t.strip() in ("name",) and "NAME" not in mapping:
            mapping["NAME"] = idx
        if t.strip() in ("qty", "quantity") and "VALUES" not in mapping:
            mapping["VALUES"] = idx
        if t.strip() in ("uom", "unit") and "UNITS" not in mapping:
            mapping["UNITS"] = idx
    return mapping


def _first_data_preview(ws, header_row_idx: int, cols: list[int], n: int = 3) -> list:
    out = []
    for row in ws.iter_rows(
        min_row=header_row_idx + 1,
        max_row=header_row_idx + 8,
        max_col=max(cols) + 1 if cols else 20,
        values_only=True,
    ):
        vals = [row[c] if c < len(row) else None for c in cols]
        if any(v not in (None, "") for v in vals):
            # skip numeric banner
            if all(isinstance(v, (int, float)) or (isinstance(v, str) and v.isdigit()) for v in vals if v not in (None, "")):
                continue
            out.append(vals)
            if len(out) >= n:
                break
    return out


def main() -> int:
    docs = collect_tsd_files(DEFAULT_TSD_PACKING_ROOT)
    lines: list[str] = []

    lines.append("=== A) List1 / suspicious files: column map vs expected STD (B=1 CODE) ===")
    expected_std = {"CODE": 1, "SPEC": 2, "TAGS": 3, "NAME": 7, "VALUES": 8, "UNITS": 9, "VENDOR": 12}
    offset_suspects: list[str] = []

    for tip in LIST1_TIPS:
        cand = next((d for d in docs if tip in d.file_name), None)
        if not cand:
            lines.append(f"MISSING {tip}")
            continue
        wb = openpyxl.load_workbook(cand.file_full_path, read_only=True, data_only=True)
        lines.append(f"\nFILE: {cand.file_name}")
        lines.append(f"  sheets: {wb.sheetnames}")
        for sheet in wb.sheetnames:
            if str(sheet).lower().startswith("master"):
                continue
            ws = wb[sheet]
            hidx, header = _find_header_row(ws)
            cmap = _col_map_from_header(header) if header else {}
            lines.append(f"  sheet {sheet!r} header_row={hidx} map={cmap}")
            if not cmap:
                # dump first non-empty rows
                for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, max_col=14, values_only=True), 1):
                    cells = [(j, v) for j, v in enumerate(row) if v not in (None, "")]
                    if cells:
                        lines.append(f"    r{i}: {cells[:8]}")
                continue
            # compare to std
            diffs = []
            for k, exp in expected_std.items():
                got = cmap.get(k)
                if got is None:
                    diffs.append(f"{k}=MISSING")
                elif got != exp:
                    diffs.append(f"{k}: got={got} expected={exp} (shift={got - exp})")
            if diffs:
                lines.append(f"    OFFSET/DIFF: {diffs}")
                # If CODE shifted left by 1 -> missing column A
                if cmap.get("CODE") == 0:
                    offset_suspects.append(f"{cand.file_name} | sheet={sheet} | CODE at A(0) instead of B(1)")
                    lines.append("    >>> SUSPECT: missing empty column A (everything shifted left by 1)")
            else:
                lines.append("    map matches STD B/C/D/H/I/J/M")
            cols = [cmap[k] for k in ("CODE", "SPEC", "NAME", "VALUES", "UNITS") if k in cmap]
            preview = _first_data_preview(ws, hidx or 1, cols)
            lines.append(f"    preview: {preview}")
        wb.close()

    lines.append("\n=== OFFSET SUSPECTS (for correction) ===")
    if not offset_suspects:
        lines.append("(none found by header map)")
    else:
        for s in offset_suspects:
            lines.append(s)

    lines.append("\n=== B) SO - PL / AGCC.323 files: column maps ===")
    sopl_docs = [
        d
        for d in docs
        if "ГФ" in d.file_full_path or "AGCC.323" in d.file_name or "2064203" in d.file_name
    ]
    maps_counter: dict[str, list[str]] = {}
    for d in sopl_docs:
        wb = openpyxl.load_workbook(d.file_full_path, read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            if str(sheet).lower().startswith("master"):
                continue
            ws = wb[sheet]
            hidx, header = _find_header_row(ws, max_col=100)
            cmap = _col_map_from_header(header) if header else {}
            key = str(sorted(cmap.items()))
            maps_counter.setdefault(key, []).append(f"{d.file_name}::{sheet}")
            # also show raw header cells at user-mentioned indices
            if header:
                idxs = [6, 9, 12, 13, 14, 56, 57]  # G J M N O BE BF (0-based: 6,9,12,13,14,56,57)
                snap = {i: (header[i] if i < len(header) else None) for i in idxs}
                lines.append(f"{d.file_name} | {sheet!r} header@{hidx} map={cmap}")
                lines.append(f"  snap G/J/M/N/O/BE/BF: {snap}")
        wb.close()

    lines.append("\n=== SO-PL map variants ===")
    for key, files in maps_counter.items():
        lines.append(f"map {key}")
        lines.append(f"  count={len(files)} sample={files[:5]}")

    # Name of document in pickle
    lines.append("\n=== C) 'Name of document' in cache ===")
    try:
        import pickle
        from base.tables_columns import SPECIFICATION_NAME, ANNOTATION, CODE, TITLE, NAME

        with open(ROOT / "RFQ/ds_compare/cache/tsd_packing/tsd_packing_rows.cache", "rb") as f:
            payload = pickle.load(f)
        hits = []
        for r in payload["data"]:
            spec = str(r.el[SPECIFICATION_NAME].value or "")
            code = str(r.el[CODE].value or "")
            name = str(r.el[NAME].value or "")
            if "Name of document" in spec or code == "Name of Zip" or "Name of document" in name:
                hits.append(
                    (
                        r.el[ANNOTATION].value,
                        r.el[TITLE].value,
                        code,
                        spec,
                        name[:60],
                    )
                )
        lines.append(f"hits={len(hits)}")
        for h in hits[:30]:
            lines.append(repr(h))
    except Exception as e:
        lines.append(f"cache read error: {e!r}")

    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
