"""Local checks for copying the newest RD/SQ MTO into the robot folder."""

from __future__ import annotations

import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rd_catalog.models import FileKind, FileRecord, ReviewState, SourceKind
from rd_catalog.robot_mto_sync import (
    RobotMtoSyncError,
    archive_folder_name,
    execute_robot_mto_sync,
    plan_robot_mto_sync,
    plan_summary_text,
    preview_robot_mto_compare,
)
from rd_catalog.scan import _candidate_kind


def _record(
    file_id: int,
    *,
    path: str,
    source: SourceKind,
    title: str = "2225",
    mark: str = "KSB",
    revision: str = "01",
    appendix: str | None = None,
    mtime_ns: int = 1,
    size: int = 10,
    present: bool = True,
    file_kind: str = FileKind.MTO_XLSX.value,
    discipline_block: str = "MTO-0001",
) -> FileRecord:
    name = Path(path).name
    return FileRecord(
        id=file_id,
        path=path,
        path_key=path.casefold(),
        source=source,
        present=present,
        review_state=ReviewState.ACKNOWLEDGED,
        first_seen_run_id=1,
        last_seen_run_id=1,
        data={
            "file_kind": file_kind,
            "name": name,
            "parse_status": "parsed",
            "title": title,
            "mark": mark,
            "title_system": f"{title}-{mark}",
            "discipline_block": discipline_block,
            "revision": revision,
            "appendix": appendix,
            "mtime_ns": mtime_ns,
            "size": size,
        },
    )


def _write(path: Path, payload: bytes = b"mto-source") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def main() -> None:
    """Plan and execute robot MTO add/replace without touching UNC."""

    assert (
        _candidate_kind("AGCC.287-2225-KSB.MTO-0001_01_RU.xlsx", sq_only=True)
        is FileKind.MTO_XLSX
    )
    assert _candidate_kind("foo.dwg", sq_only=True, allow_source=True) is None
    when = datetime(2026, 8, 27, 11, 17)
    assert (
        archive_folder_name("AGCC.287-2225-KSB.MTO-0001_01_RU", when)
        == "_old_AGCC.287-2225-KSB.MTO-0001_01_RU_2026.08.27_11.17"
    )

    with tempfile.TemporaryDirectory(prefix="rd_catalog_robot_mto_") as temp:
        root = Path(temp)
        rd_dir = root / "rd"
        sq_dir = root / "sq"
        robot_root = root / "robot"
        rd_name = "AGCC.287-2225-KSB.MTO-0001_02_RU.xlsx"
        sq_name = "AGCC.287-2225-KSB.MTO-0001_03_RU.xlsx"
        old_name = "AGCC.287-2225-KSB.MTO-0001_01_RU.xlsx"
        rd_path = rd_dir / rd_name
        sq_path = sq_dir / sq_name
        robot_dir = robot_root / "2225" / "KSB"
        old_robot = robot_dir / old_name
        _write(rd_path, b"rd-rev-02")
        _write(sq_path, b"sq-rev-03")
        _write(old_robot, b"robot-rev-01")

        rd_record = _record(
            1,
            path=str(rd_path),
            source=SourceKind.RD,
            revision="02",
            mtime_ns=20,
            size=rd_path.stat().st_size,
        )
        sq_record = _record(
            2,
            path=str(sq_path),
            source=SourceKind.SQ,
            revision="03",
            mtime_ns=10,
            size=sq_path.stat().st_size,
        )
        robot_record = _record(
            3,
            path=str(old_robot),
            source=SourceKind.ROBOT,
            revision="01",
            mtime_ns=5,
            size=old_robot.stat().st_size,
        )
        pdf_record = _record(
            4,
            path=str(rd_dir / "AGCC.287-2225-KSB.OD-0001_02_RU.pdf"),
            source=SourceKind.RD,
            revision="02",
            file_kind=FileKind.PDF.value,
            discipline_block="OD-0001",
        )

        plan = plan_robot_mto_sync(
            title="2225",
            mark="KSB",
            records=(rd_record, sq_record, robot_record, pdf_record),
            detected_current_ids={1, 4},
            robot_root=robot_root,
            robot_flat_structure=False,
            discover_siblings=False,
        )
        assert plan.source.path == str(sq_path)
        assert plan.source.source is SourceKind.SQ
        assert plan.is_replace
        assert plan.robot_files[0].path == str(old_robot)
        assert Path(plan.destination_dir) == robot_dir
        assert Path(plan.destination_path) == robot_dir / sq_name
        summary = plan_summary_text(plan, now=when)
        assert "Источник: SQ" in summary
        assert "Ревизия: 03" in summary
        assert "_old_AGCC.287-2225-KSB.MTO-0001_01_RU_2026.08.27_11.17" in summary

        result = execute_robot_mto_sync(plan, now=when)
        assert not result.added
        assert Path(result.destination_path).read_bytes() == b"sq-rev-03"
        assert not old_robot.exists()
        archived = Path(result.archived[0].archived_path)
        assert archived.read_bytes() == b"robot-rev-01"
        assert archived.parent.name == (
            "_old_AGCC.287-2225-KSB.MTO-0001_01_RU_2026.08.27_11.17"
        )
        assert archived.parent.parent == robot_dir

        rd_only = plan_robot_mto_sync(
            title="2225",
            mark="KSB",
            records=(rd_record, pdf_record),
            detected_current_ids={1, 4},
            robot_root=robot_root,
            robot_flat_structure=False,
            discover_siblings=False,
        )
        assert rd_only.source.path == str(rd_path)
        assert rd_only.source.source is SourceKind.RD

        other_title = _record(
            9,
            path=str(rd_dir / "AGCC.287-1513-POS.MTO-0001_04_RU.xlsx"),
            source=SourceKind.RD,
            title="1513",
            mark="POS",
            revision="04",
            mtime_ns=99,
        )
        try:
            plan_robot_mto_sync(
                title="1513",
                mark="POS",
                records=(other_title,),
                detected_current_ids=set(),
                robot_root=robot_root,
                robot_flat_structure=False,
                discover_siblings=False,
            )
        except RobotMtoSyncError as exc:
            assert "Нет файла MTO" in str(exc)
        else:
            raise AssertionError("expected missing overlay-current RD MTO")

        add_root = root / "robot_add"
        add_source = rd_dir / "AGCC.287-9192-POS.MTO-0001_01_RU.xlsx"
        _write(add_source, b"new-pos")
        add_record = _record(
            5,
            path=str(add_source),
            source=SourceKind.RD,
            title="9192",
            mark="POS",
            revision="01",
            mtime_ns=30,
            size=add_source.stat().st_size,
        )
        add_plan = plan_robot_mto_sync(
            title="9192",
            mark="POS",
            records=(add_record,),
            detected_current_ids={5},
            robot_root=add_root,
            robot_flat_structure=False,
            discover_siblings=False,
        )
        assert not add_plan.is_replace
        assert Path(add_plan.destination_dir) == add_root / "9192" / "POS"
        add_result = execute_robot_mto_sync(add_plan, now=when)
        assert add_result.added
        assert Path(add_result.destination_path).read_bytes() == b"new-pos"
        assert not add_result.archived

        sibling_dir = root / "sq_pkg"
        sibling = sibling_dir / "AGCC.287-5850-SKUD.MTO-0001_02_RU.xlsx"
        _write(sibling, b"sq-sibling")
        sibling_plan = plan_robot_mto_sync(
            title="5850",
            mark="SKUD",
            records=(),
            detected_current_ids=set(),
            robot_root=root / "robot_sibling",
            robot_flat_structure=True,
            sq_folders=(str(sibling_dir),),
            discover_siblings=True,
        )
        assert sibling_plan.source.path == str(sibling)
        assert sibling_plan.source.source is SourceKind.SQ
        assert sibling_plan.source.revision == "02"

        same_rev_rd = _record(
            6,
            path=str(rd_path),
            source=SourceKind.RD,
            revision="02",
            mtime_ns=100,
            size=1,
        )
        same_rev_sq = _record(
            7,
            path=str(sq_path),
            source=SourceKind.SQ,
            revision="02",
            mtime_ns=50,
            size=1,
        )
        newer_mtime = plan_robot_mto_sync(
            title="2225",
            mark="KSB",
            records=(same_rev_rd, same_rev_sq),
            detected_current_ids={6},
            robot_root=robot_root,
            robot_flat_structure=False,
            discover_siblings=False,
        )
        assert newer_mtime.source.path == str(rd_path)

        outside = plan_robot_mto_sync(
            title="2225",
            mark="KSB",
            records=(rd_record,),
            detected_current_ids={1},
            robot_root=robot_root,
            robot_flat_structure=False,
            discover_siblings=False,
        )
        bad = outside.__class__(
            title=outside.title,
            mark=outside.mark,
            source=outside.source,
            robot_files=outside.robot_files,
            destination_dir=str(root / "not_robot"),
            destination_path=str(root / "not_robot" / rd_name),
            robot_root=str(robot_root),
        )
        try:
            execute_robot_mto_sync(bad, now=when)
        except RobotMtoSyncError as exc:
            assert "каталоге робота" in str(exc)
        else:
            raise AssertionError("expected refuse write outside robot_root")

        dup_a = robot_dir / "AGCC.287-2225-KSB.MTO-0001_00_RU.xlsx"
        dup_b = robot_dir / "AGCC.287-2225-KSB.MTO-0001_00-AN01_RU.xlsx"
        _write(dup_a, b"dup-a")
        _write(dup_b, b"dup-b")
        dup_plan = plan_robot_mto_sync(
            title="2225",
            mark="KSB",
            records=(
                rd_record,
                _record(
                    11,
                    path=str(dup_a),
                    source=SourceKind.ROBOT,
                    revision="0",
                    mtime_ns=1,
                    size=1,
                ),
                _record(
                    12,
                    path=str(dup_b),
                    source=SourceKind.ROBOT,
                    revision="0",
                    appendix="01",
                    mtime_ns=2,
                    size=1,
                ),
            ),
            detected_current_ids={1},
            robot_root=robot_root,
            robot_flat_structure=False,
            discover_siblings=False,
        )
        assert len(dup_plan.robot_files) == 2
        dup_result = execute_robot_mto_sync(dup_plan, now=when)
        assert len(dup_result.archived) == 2
        assert not dup_a.exists()
        assert not dup_b.exists()
        assert Path(dup_result.destination_path).read_bytes() == b"rd-rev-02"

        cmp_dir = root / "compare_live"
        src_xlsx = cmp_dir / "rd" / "AGCC.287-2225-KSB.MTO-0001_02_RU.xlsx"
        bot_xlsx = (
            cmp_dir / "robot" / "2225" / "KSB" / "AGCC.287-2225-KSB.MTO-0001_01_RU.xlsx"
        )
        _write(src_xlsx, b"src-xlsx")
        _write(bot_xlsx, b"bot-xlsx")
        src_rows = [
            {
                "CODE": "BCC0001",
                "UNITS": "шт",
                "VALUES": 1,
                "TAGS": ["T1"],
                "NAME": "Прибор A",
                "VENDOR": "V",
                "TYPE_MARK": "M",
            },
            {
                "CODE": "BCC0002",
                "UNITS": "шт",
                "VALUES": 1,
                "TAGS": ["T2"],
                "NAME": "Прибор B",
                "VENDOR": "V",
                "TYPE_MARK": "M",
            },
        ]
        bot_rows = [
            {
                "CODE": "BCC0001",
                "UNITS": "шт",
                "VALUES": 2,
                "TAGS": ["T1"],
                "NAME": "Прибор A",
                "VENDOR": "V",
                "TYPE_MARK": "M",
            },
            {
                "CODE": "BCC0003",
                "UNITS": "шт",
                "VALUES": 1,
                "TAGS": ["T3"],
                "NAME": "Прибор C",
                "VENDOR": "V",
                "TYPE_MARK": "M",
            },
        ]
        cmp_rd = _record(
            21,
            path=str(src_xlsx),
            source=SourceKind.RD,
            revision="02",
            mtime_ns=40,
            size=1,
        )
        cmp_robot = _record(
            22,
            path=str(bot_xlsx),
            source=SourceKind.ROBOT,
            revision="01",
            mtime_ns=10,
            size=1,
        )
        cmp_plan = plan_robot_mto_sync(
            title="2225",
            mark="KSB",
            records=(cmp_rd, cmp_robot),
            detected_current_ids={21},
            robot_root=cmp_dir / "robot",
            robot_flat_structure=False,
            discover_siblings=False,
        )

        def _loader(path):
            text = str(path)
            if text == str(src_xlsx):
                return src_rows
            if text == str(bot_xlsx):
                return bot_rows
            return []

        preview = preview_robot_mto_compare(
            cmp_plan,
            cmp_dir / "out",
            loader=_loader,
            now=when,
        )
        assert preview.appear == 1
        assert preview.disappear == 1
        assert preview.changed == 1
        assert "не кэш скана" in preview.summary
        assert "появятся строк: 1" in preview.summary
        assert "пропадут строк: 1" in preview.summary
        assert "изменятся строк: 1" in preview.summary
        assert preview.report_path
        assert Path(preview.report_path).is_file()
        import openpyxl

        workbook = openpyxl.load_workbook(preview.report_path)
        assert "Сводка" in workbook.sheetnames
        assert "Появятся" in workbook.sheetnames
        assert "Пропадут" in workbook.sheetnames
        assert "Изменятся" in workbook.sheetnames
        workbook.close()

        add_preview = preview_robot_mto_compare(
            add_plan,
            cmp_dir / "out_add",
            loader=lambda _path: src_rows,
            now=when,
        )
        assert add_preview.report_path is None
        assert add_preview.source_rows == 2
        assert "сравнивать не с чем" in add_preview.summary

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication, QPlainTextEdit

        from rd_catalog.robot_mto_dialog import RobotMtoSyncDialog

        app = QApplication.instance() or QApplication([])
        dialog = RobotMtoSyncDialog(plan, now=when, runtime_dir=root / "runtime")
        dialog.show()
        app.processEvents()
        assert "Заменить MTO у робота" in dialog.windowTitle()
        assert dialog._compare_button.text() == "Сравнить"
        assert dialog._open_report_button.text() == "Открыть файл сравнения"
        assert not dialog._open_report_button.isEnabled()
        texts = [
            widget.toPlainText()
            for widget in dialog.findChildren(QPlainTextEdit)
        ]
        assert any("Источник: SQ" in text for text in texts)
        dialog.close()
        app.processEvents()

    print("RD catalog robot MTO sync: OK")


if __name__ == "__main__":
    main()
