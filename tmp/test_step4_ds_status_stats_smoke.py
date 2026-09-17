"""Smoke tests for Step4 «Статистика ДС» quality sheet."""

from __future__ import annotations

import re
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from base.base_classes import RowStd, RowType, TableComments
from base.tables_columns import (
    DS_ACTUAL,
    DS_MANAGER,
    DS_NAME,
    ROW_TYPE,
    UL_COMPARE_STATUS,
)
from RFQ.tags_rfp_compare.step4.step4_6_quality_sheets import (
    build_ds_ul_status_stats,
    ds_stats_cell_tone,
    ds_stats_row_has_problems,
    ds_stats_status_from_header,
    write_quality_sheets,
)
from RFQ.tags_rfp_compare.step4.step4_packing_compare import (
    STATUS_COMPLETE,
    STATUS_OPEN_UPD,
    STATUS_PACKING_ONLY,
    STATUS_SHORTFALL,
)
from utils.colors import Color


def _position_row(
    *,
    ds_actual: str = "",
    ds_name: str = "",
    ds_manager: str = "",
    ul_status: str = "",
) -> RowStd:
    row = RowStd.get_std_check_row({}, TableComments())
    row.row_type = RowType.position_row
    row.el[ROW_TYPE].value = RowType.position_row
    row.el[DS_ACTUAL].value = ds_actual
    row.el[DS_NAME].value = ds_name
    row.el[DS_MANAGER].value = ds_manager
    row.el[UL_COMPARE_STATUS].value = ul_status
    return row


def _non_position_row(*, ul_status: str = STATUS_COMPLETE) -> RowStd:
    row = RowStd.get_std_check_row({}, TableComments())
    row.row_type = RowType.empty_row
    row.el[ROW_TYPE].value = RowType.empty_row
    row.el[DS_ACTUAL].value = "ДС1"
    row.el[DS_NAME].value = "ДС1"
    row.el[UL_COMPARE_STATUS].value = ul_status
    return row


def _header_total(header: str) -> int:
    match = re.search(r"\((\d+)\)\s*$", header)
    assert match is not None, header
    return int(match.group(1))


def _status_label(header: str) -> str:
    return re.sub(r"\s*\(\d+\)\s*$", "", header)


def _cell_fill_hex(cell) -> str:
    rgb = getattr(cell.fill.fgColor, "rgb", None)
    if rgb is None:
        return ""
    text = str(rgb)
    if len(text) >= 8 and text[:2].upper() in ("FF", "00"):
        return text[2:].upper()
    return text.upper()


class Step4DsStatusStatsSmoke(unittest.TestCase):
    def test_two_ds_mixed_statuses_and_header_totals(self) -> None:
        rows = [
            _position_row(
                ds_actual="ДС24",
                ds_name="ДС24",
                ds_manager="Иванов",
                ul_status=STATUS_SHORTFALL,
            ),
            _position_row(
                ds_actual="ДС24",
                ds_name="ДС24",
                ds_manager="Иванов",
                ul_status=STATUS_COMPLETE,
            ),
            _position_row(
                ds_actual="ДС82",
                ds_name="ДС82",
                ds_manager="Петров",
                ul_status=STATUS_OPEN_UPD,
            ),
        ]
        records, headers, totals = build_ds_ul_status_stats(rows)

        self.assertEqual(headers[:3], ["Фактический ДС", "Порядковый ДС", "Фамилия менеджера"])
        self.assertEqual(records[0][:3], ["ДС24", "ДС24", "Иванов"])
        self.assertEqual(records[1][:3], ["ДС82", "ДС82", "Петров"])

        shortfall_idx = headers.index(f"{STATUS_SHORTFALL} (1)")
        complete_idx = headers.index(f"{STATUS_COMPLETE} (1)")
        open_idx = headers.index(f"{STATUS_OPEN_UPD} (1)")

        self.assertEqual(records[0][shortfall_idx], 1)
        self.assertEqual(records[0][complete_idx], 1)
        self.assertEqual(records[0][open_idx], 0)
        self.assertEqual(records[1][open_idx], 1)

        for header in headers[3:]:
            label = _status_label(header)
            self.assertEqual(_header_total(header), totals[label])
            self.assertEqual(
                _header_total(header),
                sum(record[headers.index(header)] for record in records),
            )

    def test_empty_manager_row_emitted(self) -> None:
        rows = [
            _position_row(
                ds_actual="ДС1",
                ds_name="ДС1",
                ds_manager="",
                ul_status=STATUS_COMPLETE,
            ),
        ]
        records, headers, _totals = build_ds_ul_status_stats(rows)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0][:3], ["ДС1", "ДС1", ""])

    def test_empty_actual_and_sequential_emitted(self) -> None:
        rows = [
            _position_row(
                ds_actual="",
                ds_name="",
                ds_manager="Сидоров",
                ul_status=STATUS_OPEN_UPD,
            ),
        ]
        records, _headers, totals = build_ds_ul_status_stats(rows)
        self.assertEqual(records[0][:3], ["", "", "Сидоров"])
        self.assertEqual(totals[STATUS_OPEN_UPD], 1)

    def test_unexpected_status_becomes_extra_column(self) -> None:
        unexpected = "Неожиданный статус"
        rows = [
            _position_row(
                ds_actual="ДС5",
                ds_name="ДС5",
                ds_manager="Козлов",
                ul_status=unexpected,
            ),
        ]
        _records, headers, totals = build_ds_ul_status_stats(rows)
        status_labels = [_status_label(header) for header in headers[3:]]
        self.assertIn(unexpected, status_labels)
        self.assertEqual(totals[unexpected], 1)
        self.assertIn(f"{unexpected} (1)", headers)

    def test_only_position_rows_counted(self) -> None:
        rows = [
            _position_row(
                ds_actual="ДС1",
                ds_name="ДС1",
                ds_manager="A",
                ul_status=STATUS_SHORTFALL,
            ),
            _non_position_row(ul_status=STATUS_COMPLETE),
        ]
        records, headers, totals = build_ds_ul_status_stats(rows)
        self.assertEqual(len(records), 1)
        self.assertEqual(totals[STATUS_SHORTFALL], 1)
        self.assertEqual(totals.get(STATUS_COMPLETE, 0), 0)
        shortfall_idx = headers.index(f"{STATUS_SHORTFALL} (1)")
        self.assertEqual(records[0][shortfall_idx], 1)

    def test_packing_only_and_junk_label_sort(self) -> None:
        rows = [
            _position_row(
                ds_actual="ДС10",
                ds_name="ДС10",
                ds_manager="ГЭМ",
                ul_status=STATUS_PACKING_ONLY,
            ),
            _position_row(
                ds_actual="",
                ds_name="",
                ds_manager="",
                ul_status=STATUS_PACKING_ONLY,
            ),
            _position_row(
                ds_actual="без-номера",
                ds_name="без-номера",
                ds_manager="Яковлев",
                ul_status=STATUS_SHORTFALL,
            ),
        ]
        records, headers, totals = build_ds_ul_status_stats(rows)
        self.assertEqual(totals[STATUS_PACKING_ONLY], 2)
        self.assertEqual(records[0][:3], ["ДС10", "ДС10", "ГЭМ"])
        self.assertEqual(records[-1][:3], ["", "", ""])
        packing_idx = headers.index(f"{STATUS_PACKING_ONLY} (2)")
        self.assertEqual(records[0][packing_idx], 1)
        self.assertEqual(records[-1][packing_idx], 1)

    def test_leftover_sequential_marker_is_own_group(self) -> None:
        rows = [
            _position_row(
                ds_actual="ДС92",
                ds_name="ДС92_24Б",
                ds_manager="Михеев Д.",
                ul_status=STATUS_COMPLETE,
            ),
            _position_row(
                ds_actual="ДС92",
                ds_name="ДС92_RFP_не_найден",
                ds_manager="",
                ul_status=STATUS_PACKING_ONLY,
            ),
        ]
        records, headers, totals = build_ds_ul_status_stats(rows)
        self.assertEqual(len(records), 2)
        sequentials = [record[1] for record in records]
        self.assertIn("ДС92_24Б", sequentials)
        self.assertIn("ДС92_RFP_не_найден", sequentials)
        self.assertEqual(totals[STATUS_PACKING_ONLY], 1)
        packing_idx = headers.index(f"{STATUS_PACKING_ONLY} (1)")
        leftover = next(
            record for record in records if record[1] == "ДС92_RFP_не_найден"
        )
        self.assertEqual(leftover[0], "ДС92")
        self.assertEqual(leftover[packing_idx], 1)

    def test_cell_tone_zero_ok_complete_always_ok(self) -> None:
        self.assertEqual(ds_stats_cell_tone(STATUS_SHORTFALL, 0), "ok")
        self.assertEqual(ds_stats_cell_tone(STATUS_SHORTFALL, 1), "problem")
        self.assertEqual(ds_stats_cell_tone(STATUS_PACKING_ONLY, 4), "problem")
        self.assertEqual(ds_stats_cell_tone(STATUS_COMPLETE, 0), "ok")
        self.assertEqual(ds_stats_cell_tone(STATUS_COMPLETE, 30850), "ok")
        self.assertEqual(
            ds_stats_status_from_header(f"{STATUS_COMPLETE} (30850)"),
            STATUS_COMPLETE,
        )

    def test_row_has_problems_only_when_yellow_counts(self) -> None:
        rows = [
            _position_row(
                ds_actual="ДС3",
                ds_name="ДС3",
                ds_manager="Ок",
                ul_status=STATUS_COMPLETE,
            ),
            _position_row(
                ds_actual="ДС4",
                ds_name="ДС4",
                ds_manager="Проблема",
                ul_status=STATUS_SHORTFALL,
            ),
        ]
        records, headers, _totals = build_ds_ul_status_stats(rows)
        complete_row = next(record for record in records if record[1] == "ДС3")
        problem_row = next(record for record in records if record[1] == "ДС4")
        self.assertFalse(ds_stats_row_has_problems(complete_row, headers))
        self.assertTrue(ds_stats_row_has_problems(problem_row, headers))

    def test_workbook_sheet_and_header_totals(self) -> None:
        import io
        import tempfile

        import xlsxwriter
        import openpyxl

        rows = [
            _position_row(
                ds_actual="ДС92",
                ds_name="ДС92_24Б",
                ds_manager="Михеев Д.",
                ul_status=STATUS_SHORTFALL,
            ),
            _position_row(
                ds_actual="ДС92",
                ds_name="ДС92_24Б",
                ds_manager="Михеев Д.",
                ul_status=STATUS_PACKING_ONLY,
            ),
            _position_row(
                ds_actual="ДС92",
                ds_name="ДС92_24Б",
                ds_manager="Михеев Д.",
                ul_status=STATUS_COMPLETE,
            ),
            _position_row(
                ds_actual="ДС1",
                ds_name="ДС1_RFP_не_найден",
                ds_manager="",
                ul_status=STATUS_OPEN_UPD,
            ),
            _position_row(
                ds_actual="ДС3",
                ds_name="ДС3",
                ds_manager="Ок",
                ul_status=STATUS_COMPLETE,
            ),
        ]
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "stats.xlsx"
            workbook = xlsxwriter.Workbook(str(path))
            try:
                write_quality_sheets(workbook, rows)
            finally:
                workbook.close()
            loaded = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), data_only=True)
            try:
                self.assertIn("Статистика ДС", loaded.sheetnames)
                sheet = loaded["Статистика ДС"]
                headers = [cell.value for cell in sheet[1]]
                self.assertEqual(headers[0], "Фактический ДС")
                self.assertIn(f"{STATUS_SHORTFALL} (1)", headers)
                self.assertIn(f"{STATUS_PACKING_ONLY} (1)", headers)
                self.assertIn(f"{STATUS_COMPLETE} (2)", headers)
                self.assertGreaterEqual(sheet.column_dimensions["A"].width or 0, 12)
                self.assertGreaterEqual(sheet.column_dimensions["B"].width or 0, 14)

                ok_hex = Color.match_matched.upper()
                problem_hex = Color.yellow.upper()
                shortfall_col = headers.index(f"{STATUS_SHORTFALL} (1)") + 1
                packing_col = headers.index(f"{STATUS_PACKING_ONLY} (1)") + 1
                complete_col = headers.index(f"{STATUS_COMPLETE} (2)") + 1
                open_col = headers.index(f"{STATUS_OPEN_UPD} (1)") + 1
                by_seq = {
                    excel_row[1].value: excel_row
                    for excel_row in sheet.iter_rows(min_row=2, max_row=sheet.max_row)
                }
                ds92 = by_seq["ДС92_24Б"]
                leftover = by_seq["ДС1_RFP_не_найден"]
                complete_ds = by_seq["ДС3"]
                self.assertEqual(ds92[shortfall_col - 1].value, 1)
                self.assertEqual(_cell_fill_hex(ds92[shortfall_col - 1]), problem_hex)
                self.assertEqual(_cell_fill_hex(ds92[packing_col - 1]), problem_hex)
                self.assertEqual(ds92[complete_col - 1].value, 1)
                self.assertEqual(_cell_fill_hex(ds92[complete_col - 1]), ok_hex)
                self.assertEqual(_cell_fill_hex(ds92[open_col - 1]), ok_hex)
                self.assertEqual(_cell_fill_hex(leftover[open_col - 1]), problem_hex)
                self.assertEqual(_cell_fill_hex(leftover[complete_col - 1]), ok_hex)
                self.assertIsNone(ds92[0].fill.patternType)
                self.assertIsNone(leftover[0].fill.patternType)
                for identity_col in range(3):
                    self.assertEqual(_cell_fill_hex(complete_ds[identity_col]), ok_hex)
                self.assertEqual(complete_ds[complete_col - 1].value, 1)
                self.assertEqual(_cell_fill_hex(complete_ds[complete_col - 1]), ok_hex)
            finally:
                loaded.close()


if __name__ == "__main__":
    unittest.main()
