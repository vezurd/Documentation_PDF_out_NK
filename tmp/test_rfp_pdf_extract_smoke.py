"""Smoke tests for RFQ.rfp_parts.pdf_rfp_extract (synthetic PDF, no UNC)."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.rfp_parts.analyze_rfp_parts import REQUIRED_FIELDS, _match_header, _parse_decimal
from RFQ.rfp_parts.pdf_rfp_extract import (
    extract_rfp_pdf,
    parse_ds_label_from_title_text,
    pdf_rfp_stamp_dir,
    preview_ds_label,
)

_HEADERS = [
    "№ п/п",
    "Титул",
    "Спецификация",
    "TAG",
    "Код 1С",
    "Код РД",
    "Наименование МТР",
    "Технические характеристики",
    "Кол-во",
    "Ед. изм",
    "Кол-во",
    "Ед. изм",
]
_COL_WIDTHS = (42, 72, 78, 88, 62, 78, 118, 128, 52, 50, 52, 50)
_CYR_FONT = next(
    str(candidate)
    for candidate in (
        Path(r"C:\Windows\Fonts\calibri.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\tahoma.ttf"),
    )
    if candidate.is_file()
)


def _insert_cyr(page: fitz.Page, point: tuple[float, float], text: str, *, fontsize: float) -> None:
    page.insert_text(
        point,
        text,
        fontsize=fontsize,
        fontname="cyr",
        fontfile=_CYR_FONT,
    )


def _insert_ascii(page: fitz.Page, point: tuple[float, float], text: str, *, fontsize: float) -> None:
    page.insert_text(point, text, fontsize=fontsize)


def _build_synthetic_pdf(path: Path, *, diadoc: bool = False) -> None:
    doc = fitz.open()
    title = doc.new_page(width=1191, height=842)
    _insert_cyr(
        title,
        (48, 80),
        "Приложение №2 к Дополнительному соглашению №98/35Б от 21.07.2026",
        fontsize=12,
    )

    page = doc.new_page(width=1191, height=842)
    origin_x, origin_y = 18.0, 36.0
    header_h = 28.0
    data_h = 36.0
    row_heights = (header_h, data_h, data_h)
    xs = [origin_x]
    for width in _COL_WIDTHS:
        xs.append(xs[-1] + width)
    ys = [origin_y]
    for height in row_heights:
        ys.append(ys[-1] + height)

    for y in ys:
        page.draw_line(fitz.Point(xs[0], y), fitz.Point(xs[-1], y), color=(0, 0, 0), width=0.8)
    for x in xs:
        page.draw_line(fitz.Point(x, ys[0]), fitz.Point(x, ys[-1]), color=(0, 0, 0), width=0.8)
    for row_i in range(len(row_heights)):
        for col_i in range(len(_COL_WIDTHS)):
            page.draw_rect(
                fitz.Rect(xs[col_i], ys[row_i], xs[col_i + 1], ys[row_i + 1]),
                color=(0, 0, 0),
                width=0.4,
            )

    def _put(
        col: int,
        row: int,
        text: str,
        *,
        dy: float = 11.0,
        fontsize: float = 7.0,
    ) -> None:
        point = (xs[col] + 2.0, ys[row] + dy)
        if any(ord(ch) > 127 for ch in text):
            _insert_cyr(page, point, text, fontsize=fontsize)
        else:
            _insert_ascii(page, point, text, fontsize=fontsize)

    for col, label in enumerate(_HEADERS):
        _put(col, 0, label, dy=16.0, fontsize=6.0)

    _put(0, 1, "1", dy=14.0)
    _put(1, 1, "8950-", dy=12.0)
    _put(1, 1, "SOT4", dy=24.0)
    _put(2, 1, "SPEC1")
    _put(3, 1, "TAG-1")
    _put(4, 1, "1640127")
    _put(5, 1, "BCC0002121")
    _put(6, 1, "Test item")
    _put(7, 1, "mark")
    _put(8, 1, "1,000")
    _put(9, 1, "шт")
    _put(10, 1, "1,000")
    _put(11, 1, "шт")

    _put(0, 2, "2", dy=14.0)
    _put(1, 2, "8950-SOT4")
    _put(2, 2, "SPEC2")
    _put(3, 2, "TAG-A; TAG-B")
    _put(4, 2, "1640128")
    _put(5, 2, "BCC0002122")
    _put(6, 2, "Second item")
    _put(7, 2, "mark")
    _put(8, 2, "1,000")
    _put(9, 2, "шт")
    _put(10, 2, "1,000")
    _put(11, 2, "шт")

    _insert_ascii(page, (36.0, ys[-1] + 48.0), "ORPHAN", fontsize=12)
    if diadoc:
        _insert_cyr(
            page,
            (700.0, 790.0),
            "Передан через Диадок 28.08.2025 17:34 GMT+03:00",
            fontsize=8,
        )
        _insert_ascii(
            page,
            (700.0, 802.0),
            "395e88f1-d12c-465a-a159-43ea777f1945",
            fontsize=7,
        )
        _insert_cyr(page, (700.0, 814.0), "Страница 2 из 9", fontsize=8)
        _insert_cyr(page, (980.0, 24.0), "Страница 1 из 8", fontsize=8)
    doc.save(path)
    doc.close()


class RfpPdfExtractSmokeTest(unittest.TestCase):
    def test_parse_ds_label_from_title_text(self) -> None:
        self.assertEqual(parse_ds_label_from_title_text(""), "")
        label = parse_ds_label_from_title_text("№98/35Б")
        self.assertEqual(label, "ДС98_35Б")
        self.assertNotIn("/", label)
        self.assertEqual(
            parse_ds_label_from_title_text(
                "Приложение №2 к Дополнительному соглашению №98/35Б от 21.07.2026"
            ),
            "ДС98_35Б",
        )
        self.assertEqual(
            parse_ds_label_from_title_text(
                "Дополнительному соглашению №98 от 21.07.2026"
            ),
            "ДС98",
        )

    def test_pdf_rfp_stamp_dir_unique_under_local_base(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            first = pdf_rfp_stamp_dir(base / "spec.pdf")
            first.mkdir(parents=True)
            second = pdf_rfp_stamp_dir(base / "spec.pdf")
            self.assertNotEqual(first, second)
            self.assertEqual(first.parent, base.resolve())
            self.assertTrue(first.name.startswith("результат распознавания PDF_"))
            self.assertTrue(second.name.endswith("_2"))

    def test_extract_synthetic_pdf(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            pdf_path = tmp_path / "AGCC.287-test-RFP-0098.pdf"
            out_dir = tmp_path / "out"
            _build_synthetic_pdf(pdf_path)

            self.assertEqual(preview_ds_label(pdf_path), "ДС98_35Б")

            result = extract_rfp_pdf(pdf_path, out_dir=out_dir)
            self.assertTrue(result.xlsx_path.name.startswith("ДС98_35Б"))
            self.assertEqual(result.contract_errors, 0)
            self.assertGreaterEqual(result.outside_words, 1)
            report = result.report_path.read_text(encoding="utf-8")
            self.assertIn("ORPHAN", report)

            wb = load_workbook(result.xlsx_path, read_only=True, data_only=True)
            try:
                self.assertIn("Перечень материалов", wb.sheetnames)
                ws = wb["Перечень материалов"]
                rows = [list(row) for row in ws.iter_rows(values_only=True)]
            finally:
                wb.close()

            self.assertGreaterEqual(len(rows), 4)
            header = ["" if cell is None else str(cell) for cell in rows[1]]
            columns = _match_header(header, prefer_lot_qty=True)
            for field in REQUIRED_FIELDS:
                self.assertIn(field, columns, msg=f"missing {field} in {header!r}")

            title_idx = columns["DS_TITLE"]
            tags_idx = columns["TAGS"]
            qty_idx = columns["VALUES"]
            first_data = ["" if cell is None else str(cell) for cell in rows[2]]
            second_data = ["" if cell is None else str(cell) for cell in rows[3]]
            title = first_data[title_idx]
            self.assertEqual(title, "8950-SOT4")
            self.assertNotIn(" ", title)
            self.assertEqual(_parse_decimal(first_data[qty_idx]), _parse_decimal("1"))
            self.assertNotIn(" ", first_data[qty_idx])
            self.assertEqual(_parse_decimal(second_data[qty_idx]), _parse_decimal("1"))
            tags = second_data[tags_idx]
            self.assertIn("TAG-A", tags)
            self.assertIn("TAG-B", tags)
            self.assertNotIn("TAG-ATAG-B", tags)

    def test_strip_diadoc_keeps_header_page_number(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            pdf_path = tmp_path / "stamp.pdf"
            out_dir = tmp_path / "out"
            _build_synthetic_pdf(pdf_path, diadoc=True)
            original = fitz.open(pdf_path)
            try:
                self.assertIn("Диадок", original[1].get_text())
            finally:
                original.close()

            result = extract_rfp_pdf(pdf_path, out_dir=out_dir, strip_stamp=True)
            self.assertIsNotNone(result.cleaned_pdf_path)
            self.assertTrue(result.cleaned_pdf_path.is_file())
            self.assertEqual(result.contract_errors, 0)
            cleaned = fitz.open(result.cleaned_pdf_path)
            try:
                text = cleaned[1].get_text()
            finally:
                cleaned.close()
            self.assertNotIn("Диадок", text)
            self.assertNotIn("395e88f1", text)
            self.assertNotIn("из\xa09", text)
            self.assertIn("из\xa08", text)
            source = fitz.open(pdf_path)
            try:
                self.assertIn("Диадок", source[1].get_text())
            finally:
                source.close()


if __name__ == "__main__":
    raise SystemExit(unittest.main(verbosity=2))
