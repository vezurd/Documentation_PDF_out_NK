"""HTTP smoke for rd_catalog_web FastAPI app."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient

from rd_catalog.config import load_config
from rd_catalog.kits_table_layout import (
    KitsTableLayout,
    resolve_kits_table_layout,
    write_kits_table_layout,
)
from rd_catalog.monitor_views import KITS_HEADERS
from rd_catalog_web.app import create_app


def _empty_config(root: Path):
    runtime = root / "runtime"
    runtime.mkdir()
    override = root / "config.json"
    override.write_text(
        json.dumps(
            {
                "rd_root": str(root / "rd"),
                "sq_root": str(root / "sq"),
                "robot_root": str(root / "robot"),
                "runtime_dir": str(runtime),
                "db_path": str(runtime / "catalog.sqlite"),
                "robot_flat_structure": True,
                "skip_dirs": ["old"],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return load_config(override)


class CatalogWebTests(unittest.TestCase):
    def test_meta_and_kits(self) -> None:
        with tempfile.TemporaryDirectory(prefix="rd_web_") as temp:
            root = Path(temp)
            config = _empty_config(root)
            app = create_app(root / "config.json")
            client = TestClient(app)
            meta = client.get("/api/meta")
            self.assertEqual(meta.status_code, 200)
            body = meta.json()
            self.assertIn("kits_headers", body)
            self.assertIn("generated_at", body)
            self.assertIn("palette", body)
            self.assertIn("Титул", body["kits_headers"])
            self.assertIn("kits_header_widths", body)
            self.assertIn("kits_paint_legend", body)
            self.assertEqual(
                body["kits_paint_legend_button"],
                "Показать легенду",
            )
            self.assertTrue(body["kits_paint_legend"])
            self.assertIn("samples", body["kits_paint_legend"][0])
            self.assertEqual(
                body["kits_progress_text"],
                "Ок: 0 / 0 · Совпадает: 0 / 0 · Код A: 0 · ТДО: 0 · "
                "Нет в РД: 0 · Проблемы MTO: 0",
            )
            self.assertNotIn("kits_progress_tooltip", body)
            self.assertEqual(body["kits_progress_visible_sep"], " · видно: ")
            self.assertEqual(body["kits_xlsx_button"], "Сохранить Excel")
            self.assertEqual(set(body["kits_headers"]), set(KITS_HEADERS))
            expected = resolve_kits_table_layout(config.runtime_dir, KITS_HEADERS)
            self.assertEqual(body["kits_headers"], list(expected.order))
            self.assertEqual(body["kits_header_widths"], expected.widths)
            kits = client.get("/api/kits")
            self.assertEqual(kits.status_code, 200)
            self.assertIsInstance(kits.json(), list)
            home = client.get("/")
            self.assertEqual(home.status_code, 200)
            html = home.text
            self.assertIn('data-tab="kits"', html)
            self.assertIn('data-tab="an"', html)
            self.assertNotIn('data-tab="heatmap"', html)
            self.assertNotIn('data-tab="worklist"', html)
            self.assertNotIn('data-tab="journal"', html)
            self.assertNotIn('data-tab="collisions"', html)
            self.assertNotIn('data-tab="readiness"', html)
            self.assertEqual(client.get("/api/heatmap").status_code, 404)
            self.assertEqual(client.get("/api/an").status_code, 200)

    def test_kits_headers_follow_runtime_layout_json(self) -> None:
        with tempfile.TemporaryDirectory(prefix="rd_web_layout_") as temp:
            root = Path(temp)
            config = _empty_config(root)
            write_kits_table_layout(
                Path(config.runtime_dir) / "kits_table_layout.json",
                KitsTableLayout(
                    order=("Марка", "Титул", "РД · рев."),
                    widths={"Титул": 41, "Марка": 42},
                ),
            )
            app = create_app(root / "config.json")
            body = TestClient(app).get("/api/meta").json()
            self.assertEqual(body["kits_headers"][:4], ["Марка", "Ок", "Титул", "РД · рев."])
            self.assertEqual(set(body["kits_headers"]), set(KITS_HEADERS))
            self.assertEqual(body["kits_header_widths"]["Титул"], 41)
            self.assertEqual(body["kits_header_widths"]["Марка"], 42)

    def test_kits_xlsx_uses_layout_and_skips_unknown_keys(self) -> None:
        import io

        import openpyxl

        with tempfile.TemporaryDirectory(prefix="rd_web_xlsx_") as temp:
            root = Path(temp)
            config = _empty_config(root)
            write_kits_table_layout(
                Path(config.runtime_dir) / "kits_table_layout.json",
                KitsTableLayout(
                    order=("Марка", "Титул", "Ок"),
                    widths={"Титул": 41, "Марка": 42, "Ок": 40},
                ),
            )
            client = TestClient(create_app(root / "config.json"))
            empty = client.post("/api/kits/xlsx", json={"keys": []})
            self.assertEqual(empty.status_code, 200)
            self.assertIn(
                "spreadsheetml.sheet",
                empty.headers.get("content-type", ""),
            )
            self.assertIn(
                "filename*=UTF-8''",
                empty.headers.get("content-disposition", ""),
            )
            workbook = openpyxl.load_workbook(io.BytesIO(empty.content))
            try:
                sheet = workbook.active
                self.assertEqual(sheet.title, "Комплекты")
                headers = [
                    sheet.cell(1, col).value
                    for col in range(1, len(KITS_HEADERS) + 1)
                ]
                self.assertEqual(headers[:3], ["Марка", "Титул", "Ок"])
                self.assertEqual(set(headers), set(KITS_HEADERS))
                self.assertIsNone(sheet.cell(2, 1).value)
                self.assertTrue(sheet.auto_filter.ref.startswith("A1:"))
            finally:
                workbook.close()
            painted = client.post(
                "/api/kits/xlsx",
                json={
                    "keys": [{"title": "9110", "mark": "KSB"}],
                    "columns": [
                        {"header": "Марка", "width_px": 70},
                        {"header": "Титул", "width_px": 57},
                    ],
                },
            )
            self.assertEqual(painted.status_code, 200)
            workbook = openpyxl.load_workbook(io.BytesIO(painted.content))
            try:
                sheet = workbook.active
                self.assertEqual(
                    [sheet.cell(1, col).value for col in range(1, 3)],
                    ["Марка", "Титул"],
                )
                self.assertIsNone(sheet.cell(2, 1).value)
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, "A1:B1")
            finally:
                workbook.close()
            bad = client.post("/api/kits/xlsx", json={"keys": "nope"})
            self.assertEqual(bad.status_code, 400)

    def test_listen_urls_never_return_bind_any(self) -> None:
        from unittest.mock import patch

        from rd_catalog_web.urls import lan_share_url, listen_urls, local_open_url

        with patch("rd_catalog_web.urls.lan_ipv4", return_value=["10.1.2.3"]):
            urls = listen_urls("0.0.0.0", 8765)
            self.assertEqual(urls[0], "http://127.0.0.1:8765")
            self.assertEqual(urls[1], "http://10.1.2.3:8765")
            self.assertTrue(all("0.0.0.0" not in url for url in urls))
            self.assertEqual(local_open_url("0.0.0.0", 8765), "http://127.0.0.1:8765")
            self.assertEqual(lan_share_url("0.0.0.0", 8765), "http://10.1.2.3:8765")
        with patch("rd_catalog_web.urls.lan_ipv4", return_value=[]):
            self.assertEqual(lan_share_url("0.0.0.0", 8765), "http://127.0.0.1:8765")
        self.assertEqual(listen_urls("127.0.0.1", 9000), ["http://127.0.0.1:9000"])

    def test_lan_ipv4_prefers_ethernet_over_vpn(self) -> None:
        from unittest.mock import patch

        from rd_catalog_web import urls as urlmod

        fake_info = [
            (0, 0, 0, "", ("10.95.175.58", 0)),
            (0, 0, 0, "", ("172.16.3.235", 0)),
            (0, 0, 0, "", ("169.254.1.1", 0)),
        ]
        aliases = {
            "10.95.175.58": "NetherlandsAmsterdamS4",
            "172.16.3.235": "Ethernet",
        }
        with (
            patch.object(urlmod.socket, "gethostname", return_value="pc"),
            patch.object(urlmod.socket, "getaddrinfo", return_value=fake_info),
            patch.object(urlmod, "_windows_ip_aliases", return_value=aliases),
        ):
            ips = urlmod.lan_ipv4()
        self.assertEqual(ips, ["172.16.3.235", "10.95.175.58"])
        with patch.object(urlmod, "lan_ipv4", return_value=ips):
            self.assertEqual(
                urlmod.lan_share_url("0.0.0.0", 8765),
                "http://172.16.3.235:8765",
            )


if __name__ == "__main__":
    unittest.main()
