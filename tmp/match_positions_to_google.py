"""Match Excel positions without BCC codes to the Google code base (СПО_1).

Reads ``Позиции без кодо.xlsx``, loads Google via ``base.base_google.load_base``,
and writes an Excel with suggested BCC, quality, Google name and type/mark.

Run from repo root:

  set PYTHONUTF8=1
  python tmp/match_positions_to_google.py
"""

from __future__ import annotations

import os
import re
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

from base.base_google import load_base  # noqa: E402
from base.tables_columns import CODE, NAME, TYPE_MARK, VENDOR  # noqa: E402

INPUT_XLSX = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP\_temp"
    r"\Позиции без кодо.xlsx"
)
OUTPUT_NAME = "Позиции без кодов_подбор_BCC.xlsx"

SRC_HEADERS = [
    "Группа",
    "№",
    "Наименование",
    "Производитель",
    "Код производителя",
    "КОД 1С ВСС",
    "Количество",
    "Col H",
    "Col I",
    "Col J",
    "Col K",
    "Col L",
    "Col M",
    "Col N",
    "Col O",
    "Col P",
]
MATCH_HEADERS = [
    "Предполагаемый код BCC",
    "Оценка качества",
    "Балл 0-100",
    "Наименование Google",
    "Тип / марка Google",
    "Вендор Google",
    "Причина подбора",
    "Альт.2 код",
    "Альт.2 балл",
    "Альт.2 наименование",
    "Альт.3 код",
    "Альт.3 балл",
    "Альт.3 наименование",
]

FILL_EXACT = PatternFill("solid", fgColor="C6EFCE")
FILL_HIGH = PatternFill("solid", fgColor="D9EAD3")
FILL_MED = PatternFill("solid", fgColor="FFF2CC")
FILL_WEAK = PatternFill("solid", fgColor="FCE4D6")
FILL_KIT = PatternFill("solid", fgColor="D0E0E3")
FILL_NONE = PatternFill("solid", fgColor="F4CCCC")
FILL_HEADER = PatternFill("solid", fgColor="305496")
FILL_MATCH_HEADER = PatternFill("solid", fgColor="548235")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")


# Visually identical Latin/Cyrillic letters — needed for SKU compact compare (А/A, Р/P).
_LOOKALIKE = str.maketrans(
    {
        "а": "a",
        "в": "b",
        "с": "c",
        "е": "e",
        "н": "h",
        "к": "k",
        "м": "m",
        "о": "o",
        "р": "p",
        "т": "t",
        "х": "x",
        "у": "y",
    }
)


def _norm(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip().lower().replace("ё", "е")
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("×", "x").replace("*", "x")
    text = re.sub(r"(?<=\d)\s*[xх]\s*(?=\d)", "x", text)
    text = text.replace("«", "").replace("»", "").replace('"', "")
    text = re.sub(r"\s+", " ", text)
    return text


def _fold(text: str) -> str:
    return text.translate(_LOOKALIKE)


def _compact(value: object) -> str:
    return re.sub(r"[^0-9a-zа-я]+", "", _fold(_norm(value)))


def _tokens(value: object) -> set[str]:
    return set(re.findall(r"[0-9a-zа-я]{2,}", _fold(_norm(value))))


_SPEC_PATTERNS = (
    re.compile(r"\d+[.,]?\d*\s*x\s*\d+[.,]?\d*(?:\s*x\s*\d+[.,]?\d*)?"),
    re.compile(r"\bm\s*\d+(?:\s*x\s*\d+[.,]?\d*)?"),
    re.compile(r"\b\d+\s*a\b"),
    re.compile(r"\b\d\s*p\b"),
    re.compile(r"\b\d+p\b"),
    re.compile(r"\bip\s*\d{2}"),
    re.compile(r"\bdin\s*\d+"),
    re.compile(r"\bva-?\d+\w*"),
)

_MODEL_RE = re.compile(r"[a-zа-я0-9]+(?:-[a-zа-я0-9]+){2,}")


def _specs(value: object) -> set[str]:
    text = _fold(_norm(value))
    found: set[str] = set()
    for pat in _SPEC_PATTERNS:
        for match in pat.findall(text):
            found.add(re.sub(r"\s+", "", match).replace(",", "."))
    return found


def _models(value: object) -> set[str]:
    text = _fold(_norm(value)).replace("_", "-")
    return set(_MODEL_RE.findall(text))


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


@dataclass
class GoogleItem:
    code: str
    name: str
    type_mark: str
    vendor: str
    name_n: str = ""
    mark_n: str = ""
    blob_n: str = ""
    compact_mark: str = ""
    compact_blob: str = ""
    compact_name: str = ""
    tokens: set[str] = field(default_factory=set)
    specs: set[str] = field(default_factory=set)
    models: set[str] = field(default_factory=set)
    is_kit: bool = False


@dataclass
class ExcelRow:
    excel_row: int
    values: list[object]
    name: str
    vendor: str
    mfr: str
    code_1c: str
    name_n: str = ""
    blob_n: str = ""
    compact_mfr: str = ""
    compact_name: str = ""
    tokens: set[str] = field(default_factory=set)
    specs: set[str] = field(default_factory=set)
    models: set[str] = field(default_factory=set)


@dataclass
class Candidate:
    item: GoogleItem
    score: float
    reasons: list[str]
    kit_only: bool = False


def _prepare_google(rows: list) -> list[GoogleItem]:
    items: list[GoogleItem] = []
    for row in rows:
        code = _cell_str(row.el[CODE].value)
        name = _cell_str(row.el[NAME].value)
        mark = _cell_str(row.el[TYPE_MARK].value)
        vendor = _cell_str(row.el.get(VENDOR).value if row.el.get(VENDOR) else "")
        blob = f"{name} {mark} {vendor}".strip()
        name_n = _norm(name)
        item = GoogleItem(
            code=code,
            name=name,
            type_mark=mark,
            vendor=vendor,
            name_n=name_n,
            mark_n=_norm(mark),
            blob_n=_norm(blob),
            compact_mark=_compact(mark),
            compact_blob=_compact(blob),
            compact_name=_compact(name),
            tokens=_tokens(blob),
            specs=_specs(blob),
            models=_models(blob),
            is_kit=("в составе" in name_n) or (len(name_n) > 280),
        )
        if item.code:
            items.append(item)
    return items


def _load_excel(path: Path) -> list[ExcelRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out: list[ExcelRow] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        values = list(row[:16])
        while len(values) < 16:
            values.append(None)
        name = _cell_str(values[2])
        if not name:
            continue
        vendor = _cell_str(values[3])
        mfr = _cell_str(values[4])
        blob = f"{name} {vendor} {mfr}"
        out.append(
            ExcelRow(
                excel_row=i,
                values=values,
                name=name,
                vendor=vendor,
                mfr=mfr,
                code_1c=_cell_str(values[5]),
                name_n=_norm(name),
                blob_n=_norm(blob),
                compact_mfr=_compact(mfr),
                compact_name=_compact(name),
                tokens=_tokens(blob),
                specs=_specs(blob),
                models=_models(blob),
            )
        )
    wb.close()
    return out


def _mark_hit(src: ExcelRow, g: GoogleItem) -> float:
    """SKU-like hits only. Short marks (M5, К-2, DIN125) must not match as substrings."""
    mfr = src.compact_mfr
    if mfr:
        if mfr == g.compact_mark and len(mfr) >= 4:
            return 1.0
        if len(mfr) >= 5 and mfr in g.compact_blob:
            return 0.9
    if src.compact_name and len(src.compact_name) >= 10 and g.compact_mark:
        mark = g.compact_mark
        if len(mark) >= 10 and mark == src.compact_name:
            return 0.95
        if len(mark) >= 10 and mark in src.compact_name:
            if len(mark) >= 0.35 * len(src.compact_name):
                return 0.9
    return 0.0


def _cheap_score(src: ExcelRow, g: GoogleItem) -> float:
    inter = src.tokens & g.tokens
    if not src.tokens:
        token_recall = 0.0
        jacc = 0.0
    else:
        token_recall = len(inter) / len(src.tokens)
        union = src.tokens | g.tokens
        jacc = len(inter) / len(union) if union else 0.0
    spec_recall = 0.0
    if src.specs:
        spec_recall = len(src.specs & g.specs) / len(src.specs)
    return 3.2 * _mark_hit(src, g) + 1.4 * token_recall + 1.0 * jacc + 1.2 * spec_recall


def _ratio(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    return SequenceMatcher(None, left[:240], right[:240]).ratio()


def _metric_sizes(specs: set[str]) -> set[str]:
    return {s for s in specs if re.fullmatch(r"m\d+(?:x\d+(?:\.\d+)?)?", s)}


def _full_score(src: ExcelRow, g: GoogleItem) -> Candidate:
    reasons: list[str] = []
    name_sim = _ratio(src.name_n, g.name_n)
    blob_sim = _ratio(src.blob_n, g.blob_n)
    inter = src.tokens & g.tokens
    union = src.tokens | g.tokens
    jacc = (len(inter) / len(union)) if union else 0.0
    recall = (len(inter) / len(src.tokens)) if src.tokens else 0.0

    mark_score = _mark_hit(src, g)
    kit_only = False
    if mark_score >= 0.95 and src.compact_mfr and src.compact_mfr == g.compact_mark:
        reasons.append("код производителя = тип/марка Google")
    elif mark_score >= 0.9 and src.compact_mfr and src.compact_mfr in g.compact_blob:
        reasons.append("код производителя найден в Google")
        if g.is_kit:
            kit_only = True
            reasons.append("совпадение внутри комплекта Google")
    elif mark_score >= 0.88:
        reasons.append("наименование совпало с типом/маркой Google")

    spec_penalty = 0.0
    if src.specs:
        missing = src.specs - g.specs
        if missing:
            spec_penalty = 0.18 * (len(missing) / len(src.specs))
            reasons.append("не совпала спецификация: " + ", ".join(sorted(missing)[:4]))
        else:
            reasons.append("спецификация совпала")

    src_m = _metric_sizes(src.specs)
    g_m = _metric_sizes(g.specs)
    if src_m and g_m and src_m.isdisjoint(g_m):
        spec_penalty += 0.28
        reasons.append("другой типоразмер (M): " + ",".join(sorted(src_m)) + " vs " + ",".join(sorted(g_m)))

    if src.models and g.models and src.models.isdisjoint(g.models):
        spec_penalty += 0.30
        reasons.append("модель отличается")

    if name_sim >= 0.92:
        reasons.append("наименование почти совпало")
    elif name_sim >= 0.75:
        reasons.append("наименование похоже")

    if g.is_kit and src.name_n and src.name_n in g.name_n and len(g.name_n) > len(src.name_n) * 1.8:
        kit_only = True
        reasons.append("позиция входит в состав комплекта Google")
    elif mark_score >= 0.9 and g.compact_mark and src.compact_mfr and src.compact_mfr != g.compact_mark:
        if len(g.blob_n) > max(80, len(src.blob_n) * 2.2) or len(g.models) >= 2:
            kit_only = True
            reasons.append("код найден в составной позиции Google")

    score = (
        0.34 * name_sim
        + 0.16 * blob_sim
        + 0.12 * jacc
        + 0.10 * recall
        + 0.48 * mark_score
        - spec_penalty
    )
    if kit_only:
        score *= 0.42
    elif mark_score >= 0.95 and spec_penalty < 0.1:
        score = max(score, 0.92)
    elif mark_score >= 0.9 and spec_penalty < 0.1 and not kit_only:
        score = max(score, 0.88)
    elif src.specs and not (src.specs - g.specs) and spec_penalty < 0.1:
        if recall >= 0.45 or name_sim >= 0.55:
            score = max(score, 0.60)
    score = max(0.0, min(1.0, score))
    if not reasons:
        reasons.append("лучшее нечёткое совпадение по тексту")
    return Candidate(item=g, score=score, reasons=reasons, kit_only=kit_only)


def _quality(top: Candidate | None, gap: float) -> tuple[str, int]:
    if top is None:
        return "не найден", 0
    points = int(round(top.score * 100))
    markish = any("код производителя" in r or "типом/маркой" in r for r in top.reasons)
    if top.kit_only:
        return "только в комплекте", points
    if top.score < 0.36 and not markish:
        return "не найден", points
    if top.score >= 0.88 and markish:
        label = "точный"
    elif top.score >= 0.78:
        label = "высокий"
    elif top.score >= 0.56:
        label = "средний"
    else:
        label = "слабый"
    if gap < 0.045 and top.score >= 0.50 and not markish:
        label = "спорный (" + label + ")"
    return label, points


def _fill_for_quality(label: str) -> PatternFill:
    if label.startswith("точный") or label.startswith("высокий"):
        return FILL_EXACT if label.startswith("точный") else FILL_HIGH
    if label.startswith("средний") or "спорный" in label:
        return FILL_MED
    if label.startswith("слабый"):
        return FILL_WEAK
    if "комплект" in label:
        return FILL_KIT
    return FILL_NONE


def _forced_candidates(src: ExcelRow, google_items: list[GoogleItem]) -> list[GoogleItem]:
    out: list[GoogleItem] = []
    dims = {s for s in src.specs if "x" in s or s.startswith("din")}
    for g in google_items:
        if _mark_hit(src, g) >= 0.88:
            out.append(g)
            continue
        if dims and dims <= g.specs:
            out.append(g)
    return out


def match_rows(excel_rows: list[ExcelRow], google_items: list[GoogleItem]) -> list[dict]:
    results: list[dict] = []
    for src in excel_rows:
        ranked = sorted(
            google_items,
            key=lambda g: _cheap_score(src, g),
            reverse=True,
        )[:48]
        seen = {id(g) for g in ranked}
        for g in _forced_candidates(src, google_items):
            if id(g) not in seen:
                ranked.append(g)
                seen.add(id(g))
        cands = sorted(
            (_full_score(src, g) for g in ranked),
            key=lambda c: c.score,
            reverse=True,
        )
        if (
            len(cands) >= 2
            and not str(cands[0].item.code).upper().startswith("BCC")
            and str(cands[1].item.code).upper().startswith("BCC")
            and (cands[0].score - cands[1].score) <= 0.08
        ):
            cands[0], cands[1] = cands[1], cands[0]
        top = cands[0] if cands else None
        second = cands[1] if len(cands) > 1 else None
        third = cands[2] if len(cands) > 2 else None
        gap = (top.score - second.score) if top and second else 1.0
        label, points = _quality(top, gap)
        if label == "не найден":
            code = ""
            g_name = top.item.name if top else ""
            g_mark = top.item.type_mark if top else ""
            g_vendor = top.item.vendor if top else ""
            reason = (
                "близкого кода в Google нет"
                + ("; лучший слабый: " + "; ".join(top.reasons[:2]) if top else "")
            )
        else:
            code = top.item.code
            g_name = top.item.name
            g_mark = top.item.type_mark
            g_vendor = top.item.vendor
            reason = "; ".join(top.reasons[:4])
        results.append(
            {
                "src": src,
                "code": code,
                "label": label,
                "points": points,
                "g_name": g_name,
                "g_mark": g_mark,
                "g_vendor": g_vendor,
                "reason": reason,
                "alt2_code": second.item.code if second else "",
                "alt2_points": int(round(second.score * 100)) if second else "",
                "alt2_name": second.item.name if second else "",
                "alt3_code": third.item.code if third else "",
                "alt3_points": int(round(third.score * 100)) if third else "",
                "alt3_name": third.item.name if third else "",
                "fill": _fill_for_quality(label),
            }
        )
    return results


def _write_xlsx(results: list[dict], out_path: Path, google_count: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Подбор BCC"
    headers = SRC_HEADERS + MATCH_HEADERS
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILL_MATCH_HEADER if col > len(SRC_HEADERS) else FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"

    for r_idx, rec in enumerate(results, 2):
        src: ExcelRow = rec["src"]
        for c_idx, val in enumerate(src.values, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.border = THIN
            cell.alignment = WRAP
        extra = [
            rec["code"],
            rec["label"],
            rec["points"],
            rec["g_name"],
            rec["g_mark"],
            rec["g_vendor"],
            rec["reason"],
            rec["alt2_code"],
            rec["alt2_points"],
            rec["alt2_name"],
            rec["alt3_code"],
            rec["alt3_points"],
            rec["alt3_name"],
        ]
        for offset, val in enumerate(extra):
            cell = ws.cell(r_idx, len(SRC_HEADERS) + 1 + offset, val)
            cell.border = THIN
            cell.alignment = WRAP
            cell.fill = rec["fill"]
        ws.row_dimensions[r_idx].height = 36

    widths = {
        1: 10,
        2: 8,
        3: 42,
        4: 12,
        5: 22,
        6: 14,
        7: 12,
        17: 18,
        18: 20,
        19: 12,
        20: 42,
        21: 28,
        22: 14,
        23: 36,
        24: 14,
        25: 10,
        26: 36,
        27: 14,
        28: 10,
        29: 36,
    }
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 12)

    summary = wb.create_sheet("Сводка")
    counts = Counter(rec["label"] for rec in results)
    summary["A1"] = "Подбор позиций без BCC к Google-базе СПО_1"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Источник: {INPUT_XLSX}"
    summary["A3"] = f"Строк Excel (без пустых/итога): {len(results)}"
    summary["A4"] = f"Записей Google после фильтра кода: {google_count}"
    summary["A5"] = f"Сформировано: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    summary["A7"] = "Оценка"
    summary["B7"] = "Кол-во"
    summary["A7"].font = Font(bold=True)
    summary["B7"].font = Font(bold=True)
    for i, (label, n) in enumerate(sorted(counts.items(), key=lambda x: -x[1]), 8):
        summary.cell(i, 1, label)
        summary.cell(i, 2, n)
        summary.cell(i, 1).fill = _fill_for_quality(label)
    note_row = 8 + len(counts) + 2
    summary.cell(
        note_row,
        1,
        "Тип / марка Google — столбец type_mark листа СПО_1 (не титул РД вида 8445-SOT). "
        "Код 1С ВСС в Google-базе не хранится; подбор идёт по наименованию, "
        "коду производителя и спецификации (сечение, DIN, ток и т.п.).",
    )
    summary.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=6)
    summary.cell(note_row, 1).alignment = WRAP
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 12
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> int:
    if not INPUT_XLSX.is_file():
        print(f"INPUT missing: {INPUT_XLSX}", flush=True)
        return 1
    print("loading google code base via load_base()…", flush=True)
    google_rows = load_base() or []
    google_items = _prepare_google(google_rows)
    print(f"google items: {len(google_items)}", flush=True)
    excel_rows = _load_excel(INPUT_XLSX)
    print(f"excel rows: {len(excel_rows)}", flush=True)
    print("matching…", flush=True)
    results = match_rows(excel_rows, google_items)
    counts = Counter(rec["label"] for rec in results)
    for label, n in sorted(counts.items(), key=lambda x: -x[1]):
        print(f"  {label}: {n}", flush=True)

    unc_out = INPUT_XLSX.with_name(OUTPUT_NAME)
    local_out = ROOT / "tmp" / OUTPUT_NAME
    saved: list[Path] = []
    for path in (unc_out, local_out):
        try:
            _write_xlsx(results, path, len(google_items))
            saved.append(path)
            print(f"saved: {path}", flush=True)
        except OSError as exc:
            print(f"WARN cannot write {path}: {exc}", flush=True)
    if not saved:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
