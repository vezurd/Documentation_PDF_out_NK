"""Offline smoke for DS audit / quality gate / Step1 baseline (no UNC writes)."""

from __future__ import annotations

import sys
import tempfile
import time
import unittest
from collections import Counter
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import base.t_comm_initial_classes as t_com_init_cls
from RFQ.rfp_parts.ds_baseline import (
    BASELINE_XLSX_NAME,
    DUPLICATE_TAGS_REPORT_PREFIX,
    EMPTY_CODE_REPORT_PREFIX,
    ISSUE_EMPTY_CODE,
    ISSUE_EXACT_DUPLICATE,
    ISSUE_GROUP_FALLBACK,
    ISSUE_INTERNAL_SHIFT,
    ISSUE_MULTI_DATA_SHEET,
    ISSUE_NO_HEADER,
    ISSUE_QTY_EMPTY,
    ISSUE_QTY_FORMULA,
    ISSUE_QTY_NEGATIVE,
    ISSUE_QTY_ZERO,
    ISSUE_TAG_DUPLICATE,
    ISSUE_TAG_MISMATCH,
    ISSUE_UNKNOWN_GOOGLE,
    ISSUE_UNRESOLVED_ID,
    QUALITY_REPORT_PREFIX,
    STRUCTURE_REPORT_PREFIX,
    TAG_MISMATCH_REPORT_PREFIX,
    IdentityDsUnitsConverter,
    _path_uri,
    build_ds_baseline,
    collect_ds_workbooks,
    resolve_ds_source_id,
)
from RFQ.rfp_parts.ds_registry import (
    MODE_FILTER,
    MODE_NEEDS_SPLIT,
    MODE_NO_UL,
    MODE_WHOLE,
    DsRegistryRelation,
    DsRegistryRow,
    STATUS_ACTIVE,
    load_registry,
    write_registry_workbook,
)
from RFQ.units_convert.models import GoogleUnitsIndex, STATUS_IDENTITY
from base.base_classes import RowType, TableComments
from base.base_mto import get_std_from_excel_file
from base.tables_columns import CODE, DS_NAME, DS_TITLE, TAGS, UNITS, VALUES, VALUES_2

DS_HEADER = [
    "№ п/п",
    "Титул",
    "Раздел",
    "Спецификация",
    "RFQ",
    "Наименование Позиций Товара по РД",
    "Код 1С СОУ",
    "Код РД",
    "Наименование Позиций Товара Поставщика",
    "Технические требования (ГОСТ/ ТУ и др.)",
    "Ед. изм.",
    "Кол-во",
]

STAMP = "20260102_030405"


def _load_xlsx(path: Path):
    return load_workbook(BytesIO(path.read_bytes()))


def _ds_row(
    *,
    npp: object = 1,
    title: str = "8529",
    system: str = "SOS",
    spec: str = "spec-1",
    rfq: str = "RFQ-1",
    name: str = "Кабель",
    code_1c: str = "1C-1",
    code: str = "BCC0000001",
    supplier: str = "Vendor",
    type_mark: str = "NYM",
    units: str = "шт",
    qty: object = 2,
    extra: list[object] | None = None,
) -> list[object]:
    row = [
        npp,
        title,
        system,
        spec,
        rfq,
        name,
        code_1c,
        code,
        supplier,
        type_mark,
        units,
        qty,
    ]
    if extra:
        row.extend(extra)
    return row


def _write_xlsx(
    path: Path,
    rows: list[list[object]],
    *,
    sheet: str = "Спека",
    extra_sheets: dict[str, list[list[object]]] | None = None,
    leading_empty: int = 0,
    insert_internal_at: int | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet

    def _adjust(row: list[object]) -> list[object]:
        data = list(row)
        if insert_internal_at is not None:
            data.insert(insert_internal_at, "служебная")
        if leading_empty:
            data = [None] * leading_empty + data
        return data

    for row in rows:
        ws.append(_adjust(row))
    if extra_sheets:
        for name, sheet_rows in extra_sheets.items():
            extra = wb.create_sheet(name)
            for row in sheet_rows:
                extra.append(_adjust(row))
    wb.save(path)
    wb.close()


def _google_index() -> GoogleUnitsIndex:
    return GoogleUnitsIndex(
        units_by_code={"BCC0000001": "шт"},
        display_units_by_code={"BCC0000001": "шт"},
        display_code_by_code={"BCC0000001": "BCC0000001"},
    )


def _rel(
    *,
    group_id: str,
    rfp_key: str,
    ul_folder: str,
    mode: str,
    title_filter: str = "",
    mark_filter: str = "",
    block_index: int = 1,
) -> DsRegistryRelation:
    return DsRegistryRelation(
        group_id=group_id,
        rfp_key=rfp_key,
        ul_folder=ul_folder,
        mode=mode,
        title_filter=title_filter,
        mark_filter=mark_filter,
        block_index=block_index,
    )


def _write_registry(path: Path) -> None:
    rows = [
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="13",
            relations=(_rel(group_id="ДС13", rfp_key="13", ul_folder="согл УЛ ДС13", mode=MODE_WHOLE),),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="47",
            relations=(_rel(group_id="ДС47", rfp_key="47", ul_folder="согл УЛ ДС47", mode=MODE_WHOLE),),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="88",
            relations=(_rel(group_id="ДС88", rfp_key="88", ul_folder="согл УЛ ДС88", mode=MODE_WHOLE),),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="37",
            relations=(_rel(group_id="ДС37", rfp_key="37", ul_folder="согл УЛ ДС37", mode=MODE_WHOLE),),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="4905",
            relations=(_rel(group_id="ДС4905", rfp_key="4905", ul_folder="согл УЛ 4905", mode=MODE_WHOLE),),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="4905_1",
            relations=(
                _rel(group_id="ДС4905_1", rfp_key="4905_1", ul_folder="согл УЛ 4905_1", mode=MODE_WHOLE),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="8",
            relations=(
                _rel(
                    group_id="ДС8",
                    rfp_key="8",
                    ul_folder="согл УЛ ДС8",
                    mode=MODE_NEEDS_SPLIT,
                    block_index=1,
                ),
                _rel(
                    group_id="ДС81",
                    rfp_key="81",
                    ul_folder="согл УЛ ДС81",
                    mode=MODE_NEEDS_SPLIT,
                    block_index=2,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="101",
            relations=(
                _rel(
                    group_id="NO_UL:101",
                    rfp_key="101",
                    ul_folder="",
                    mode=MODE_NO_UL,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="15",
            relations=(
                _rel(
                    group_id="ДС15",
                    rfp_key="15",
                    ul_folder="согл УЛ ДС15",
                    mode=MODE_FILTER,
                    title_filter="1600",
                    mark_filter="POS",
                    block_index=1,
                ),
                _rel(
                    group_id="ДС61",
                    rfp_key="61",
                    ul_folder="согл УЛ ДС61",
                    mode=MODE_FILTER,
                    title_filter="1711",
                    mark_filter="KIP",
                    block_index=2,
                ),
            ),
        ),
        DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="99",
            relations=(_rel(group_id="ДС99", rfp_key="99", ul_folder="согл УЛ ДС99", mode=MODE_WHOLE),),
        ),
    ]
    write_registry_workbook(path, rows)


def _run_baseline(ds_root: Path, registry_path: Path, out_dir: Path):
    registry = load_registry(registry_path)
    return build_ds_baseline(
        ds_root,
        registry,
        out_dir,
        write_baseline=True,
        converter=IdentityDsUnitsConverter(),
        google_index=_google_index(),
        stamp=STAMP,
    )


class DsSourceIdResolveSmokeTest(unittest.TestCase):
    def test_ancestor_filename_prefix_and_traps(self) -> None:
        active = ("13", "37", "45", "88", "100", "4905", "4905_1")
        ancestor = resolve_ds_source_id("ДС_88/ДС 37 extra.xlsx", active)
        self.assertEqual(ancestor.source_id, "88")
        self.assertEqual(ancestor.method, "ancestor")

        prefix = resolve_ds_source_id("ДС13.xlsx", active)
        self.assertEqual(prefix.source_id, "13")
        self.assertEqual(prefix.method, "filename_prefix")

        prefix_second = resolve_ds_source_id("ДС_88_pack_ДС 37 extra.xlsx", active)
        self.assertEqual(prefix_second.source_id, "88")
        self.assertEqual(prefix_second.method, "filename_prefix")

        prefix_100 = resolve_ds_source_id("ДС_100_foo_ДС 45.xlsx", active)
        self.assertEqual(prefix_100.source_id, "100")
        self.assertEqual(prefix_100.method, "filename_prefix")

        sequential = resolve_ds_source_id("ДС_88_24Б.xlsx", active)
        self.assertEqual(sequential.source_id, "88")
        self.assertEqual(sequential.method, "filename_prefix")

        long_id = resolve_ds_source_id("4905_1.xlsx", active)
        self.assertEqual(long_id.source_id, "4905_1")
        self.assertNotEqual(long_id.source_id, "4905")

        only_short = resolve_ds_source_id("ДС4905_1.xlsx", ("4905",))
        self.assertEqual(only_short.source_id, "4905")
        self.assertEqual(only_short.method, "filename_prefix")

        buried = resolve_ds_source_id("archive/copy_ДС88_ДС37.xlsx", active)
        self.assertEqual(buried.method, "unresolved")
        self.assertIsNone(buried.source_id)

        ordinal = resolve_ds_source_id(
            "ДС_50_1._Спецификация №_65_ДС 50v1.xlsx",
            active + ("50", "65"),
        )
        self.assertEqual(ordinal.source_id, "50")
        self.assertEqual(ordinal.method, "filename_prefix")

        later_number = resolve_ds_source_id(
            "ДС_96_1._Спецификация №75_ДС 70 на уменьшение.xlsx",
            active + ("75", "96", "70"),
        )
        self.assertEqual(later_number.source_id, "96")

        appendix = resolve_ds_source_id(
            "ДС_75_Приложение 1 к ДС 75_16а Спецификация№41.xlsx",
            active + ("75", "41"),
        )
        self.assertEqual(appendix.source_id, "75")


class DsBaselineSmokeTest(unittest.TestCase):
    def test_collect_skips_lock_registry_svod_and_reports(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            ds_root.mkdir()
            registry_path = root / "Реестр_ДС_new.xlsx"
            _write_registry(registry_path)
            _write_xlsx(ds_root / "ДС13.xlsx", [DS_HEADER, _ds_row()])
            _write_xlsx(ds_root / "ДС_47" / "spec.xlsx", [DS_HEADER, _ds_row(npp=2)])
            (ds_root / "~$lock.xlsx").write_bytes(b"PK\x03\x04lock")
            _write_xlsx(ds_root / "Реестр_ДС_УЛ.xlsx", [["Актуальный ДС"]])
            _write_xlsx(ds_root / "СВОД_ДС_ручной.xlsx", [DS_HEADER, _ds_row()])
            _write_xlsx(ds_root / BASELINE_XLSX_NAME, [DS_HEADER])
            _write_xlsx(
                ds_root / f"{STRUCTURE_REPORT_PREFIX}_{STAMP}.xlsx",
                [["Сводка"]],
            )
            files, skipped = collect_ds_workbooks(ds_root, registry_path=registry_path)
            rels = [item.relpath for item in files]
            self.assertEqual(rels, ["ДС13.xlsx", "ДС_47/spec.xlsx"])
            reasons = {item.reason for item in skipped}
            self.assertTrue(reasons & {"excel_lock", "own_or_manual_or_registry_name", "own_output"})

    def test_blocking_set_writes_reports_without_baseline(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)

            _write_xlsx(
                ds_root / "ДС99_shift.xlsx",
                [DS_HEADER[:7] + ["служебная"] + DS_HEADER[7:], _ds_row()],
            )
            _write_xlsx(
                ds_root / "ДС99_sheets.xlsx",
                [DS_HEADER, _ds_row()],
                extra_sheets={"Копия": [DS_HEADER, _ds_row(npp=2)]},
            )
            formula_rows = [DS_HEADER, _ds_row(qty="=2+2")]
            _write_xlsx(ds_root / "ДС99_formula.xlsx", formula_rows)
            _write_xlsx(ds_root / "ДС99_empty_qty.xlsx", [DS_HEADER, _ds_row(qty=None)])
            _write_xlsx(ds_root / "ДС99_neg.xlsx", [DS_HEADER, _ds_row(qty=-3)])
            _write_xlsx(ds_root / "ДС99_zero.xlsx", [DS_HEADER, _ds_row(qty=0)])
            _write_xlsx(ds_root / "ДС99_empty_code.xlsx", [DS_HEADER, _ds_row(code="")])
            _write_xlsx(
                ds_root / "ДС99_dup.xlsx",
                [DS_HEADER, _ds_row(), _ds_row()],
            )
            _write_xlsx(
                ds_root / "ДС99_tags.xlsx",
                [
                    DS_HEADER + ["Tag"],
                    _ds_row(extra=["T-DUP;T-DUP"]),
                    _ds_row(npp=2, code="BCC9999999", extra=["T-ONLY"]),
                ],
            )
            _write_xlsx(ds_root / "unmapped.xlsx", [DS_HEADER, _ds_row()])

            result = _run_baseline(ds_root, registry_path, out_dir)
            self.assertTrue(result.blocking)
            self.assertIsNone(result.baseline_path)
            self.assertFalse((out_dir / BASELINE_XLSX_NAME).exists())
            self.assertIsNotNone(result.structure_report_path)
            self.assertTrue(result.structure_report_path.is_file())
            self.assertTrue(result.quality_report_path.is_file())
            self.assertTrue(result.empty_code_report_path.is_file())
            self.assertTrue(result.duplicate_tags_report_path.is_file())
            self.assertTrue(result.tag_mismatch_report_path.is_file())
            self.assertTrue(
                result.structure_report_path.name.startswith(STRUCTURE_REPORT_PREFIX)
            )
            self.assertTrue(
                result.quality_report_path.name.startswith(QUALITY_REPORT_PREFIX)
            )
            self.assertIn(STAMP, result.structure_report_path.name)
            self.assertIn(STAMP, result.empty_code_report_path.name)
            self.assertTrue(
                result.empty_code_report_path.name.startswith(EMPTY_CODE_REPORT_PREFIX)
            )
            self.assertTrue(
                result.duplicate_tags_report_path.name.startswith(
                    DUPLICATE_TAGS_REPORT_PREFIX
                )
            )
            self.assertTrue(
                result.tag_mismatch_report_path.name.startswith(
                    TAG_MISMATCH_REPORT_PREFIX
                )
            )

            codes = {item.code for item in result.issues}
            self.assertIn(ISSUE_INTERNAL_SHIFT, codes)
            self.assertIn(ISSUE_MULTI_DATA_SHEET, codes)
            self.assertIn(ISSUE_QTY_FORMULA, codes)
            self.assertIn(ISSUE_QTY_EMPTY, codes)
            self.assertIn(ISSUE_QTY_NEGATIVE, codes)
            self.assertIn(ISSUE_QTY_ZERO, codes)
            self.assertIn(ISSUE_EMPTY_CODE, codes)
            self.assertIn(ISSUE_EXACT_DUPLICATE, codes)
            self.assertIn(ISSUE_UNRESOLVED_ID, codes)
            self.assertIn(ISSUE_UNKNOWN_GOOGLE, codes)
            self.assertIn(ISSUE_TAG_DUPLICATE, codes)
            self.assertIn(ISSUE_TAG_MISMATCH, codes)
            self.assertTrue(any(not item.blocking for item in result.issues if item.code == ISSUE_QTY_ZERO))
            self.assertTrue(
                any(item.blocking for item in result.issues if item.code == ISSUE_QTY_FORMULA)
            )

            shift_issues = [
                item
                for item in result.issues
                if Path(item.relpath).name == "ДС99_shift.xlsx"
            ]
            shift_layout = [
                item for item in shift_issues if item.code == ISSUE_INTERNAL_SHIFT
            ]
            self.assertEqual(len(shift_layout), 1)
            self.assertIn("Раскладка не канон", shift_layout[0].message)
            self.assertFalse(
                any(item.code == "missing_core_role" for item in shift_issues)
            )
            self.assertFalse(
                any(item.file_name == "ДС99_shift.xlsx" for item in result.positions)
            )

            structure = _load_xlsx(result.structure_report_path)
            try:
                self.assertEqual(structure.sheetnames[0], "Проблемы строк")
                self.assertEqual(
                    set(structure.sheetnames),
                    {"Сводка", "Файлы ДС", "Колонки", "Листы", "Проблемы строк"},
                )
                files_ws = structure["Файлы ДС"]
                linked = False
                for row in files_ws.iter_rows(min_row=2, max_col=2):
                    for cell in row:
                        if cell.hyperlink is not None:
                            linked = True
                            self.assertTrue(
                                str(cell.hyperlink.target).startswith("file:")
                            )
                self.assertTrue(linked)
                problems = structure["Проблемы строк"]
                jump = None
                for row in problems.iter_rows(min_row=2, max_col=6):
                    link = row[2].hyperlink
                    sheet = str(row[4].value or "")
                    excel_row = row[5].value
                    if (
                        link is not None
                        and sheet
                        and isinstance(excel_row, int)
                    ):
                        jump = (link, sheet, excel_row)
                        break
                self.assertIsNotNone(jump)
                link, sheet, excel_row = jump
                self.assertIn(".xlsx", str(link.target).casefold())
                self.assertIn(sheet, str(link.location or ""))
                self.assertIn(f"A{excel_row}", str(link.location or ""))
            finally:
                structure.close()

            quality = _load_xlsx(result.quality_report_path)
            try:
                self.assertEqual(
                    set(quality.sheetnames),
                    {
                        "Сводка",
                        "Позиции без кода",
                        "Количества",
                        "Теги",
                        "Дубли",
                        "Коды вне Google",
                        "Источники",
                    },
                )
                self.assertNotIn("Единицы и конвертация", quality.sheetnames)
                empty_ws = quality["Позиции без кода"]
                self.assertGreaterEqual(empty_ws.max_row, 2)
                empty_link = empty_ws.cell(row=2, column=1).hyperlink
                self.assertIsNotNone(empty_link)
                self.assertIsNone(empty_ws.cell(row=2, column=2).hyperlink)
                empty_sheet = str(empty_ws.cell(row=2, column=3).value or "")
                empty_row = empty_ws.cell(row=2, column=4).value
                self.assertIn(empty_sheet, str(empty_link.location or ""))
                self.assertIn(f"A{empty_row}", str(empty_link.location or ""))
                self.assertIn(".xlsx", str(empty_link.target).casefold())
                src_ws = quality["Источники"]
                self.assertIsNotNone(src_ws.cell(row=2, column=2).hyperlink)
                qty_ws = quality["Количества"]
                self.assertGreaterEqual(qty_ws.max_row, 2)
                qty_link = qty_ws.cell(row=2, column=3).hyperlink
                self.assertIsNotNone(qty_link)
                qty_sheet = str(qty_ws.cell(row=2, column=5).value or "")
                qty_row = qty_ws.cell(row=2, column=6).value
                self.assertIn(qty_sheet, str(qty_link.location or ""))
                self.assertIn(f"A{qty_row}", str(qty_link.location or ""))
                self.assertIn(".xlsx", str(qty_link.target).casefold())
                self.assertIsNone(qty_ws.cell(row=2, column=4).hyperlink)
            finally:
                quality.close()

            empty_sidecar = _load_xlsx(result.empty_code_report_path)
            try:
                ws = empty_sidecar.active
                self.assertEqual(ws.title, "Без кода")
                sidecar_link = ws.cell(row=2, column=1).hyperlink
                self.assertIsNotNone(sidecar_link)
                sidecar_sheet = str(ws.cell(row=2, column=3).value or "")
                sidecar_row = ws.cell(row=2, column=4).value
                self.assertIn(sidecar_sheet, str(sidecar_link.location or ""))
                self.assertIn(f"A{sidecar_row}", str(sidecar_link.location or ""))
                self.assertIn(".xlsx", str(sidecar_link.target).casefold())
            finally:
                empty_sidecar.close()

            # Positions with empty CODE still present in memory.
            self.assertTrue(any(not item.code_normalized for item in result.positions))

    def test_valid_set_writes_step1_baseline_and_groups(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)

            _write_xlsx(ds_root / "ДС13.xlsx", [DS_HEADER, _ds_row(title="8529", system="SOS")])
            _write_xlsx(
                ds_root / "ДС_47" / "spec.xlsx",
                [DS_HEADER, _ds_row(npp=3, title="8529", system="SOS")],
                leading_empty=2,
            )
            _write_xlsx(
                ds_root / "4905_1.xlsx",
                [DS_HEADER, _ds_row(npp=4, title="4905", system="KSB")],
            )
            _write_xlsx(
                ds_root / "ДС8.xlsx",
                [DS_HEADER, _ds_row(npp=5, title="8000", system="AAA")],
            )
            _write_xlsx(
                ds_root / "ДС15.xlsx",
                [
                    DS_HEADER,
                    _ds_row(npp=6, title="1600", system="POS"),
                    _ds_row(npp=7, title="9999", system="ZZZ"),
                ],
            )
            _write_xlsx(
                ds_root / "ДС101.xlsx",
                [DS_HEADER, _ds_row(npp=8, title="1010", system="NO")],
            )

            result = _run_baseline(ds_root, registry_path, out_dir)
            self.assertFalse(result.blocking, result.summary_line())
            self.assertIsNotNone(result.baseline_path)
            self.assertTrue(result.baseline_path.is_file())
            self.assertEqual(result.baseline_path.name, BASELINE_XLSX_NAME)
            self.assertEqual(result.baseline_path.parent, out_dir)
            self.assertEqual(result.structure_report_path.parent.parent, out_dir)
            self.assertRegex(
                result.structure_report_path.parent.name,
                r"^\d{4}\.\d{2}\.\d{2}_\d{2}\.\d{2}(?:_\d+)?$",
            )
            self.assertTrue(result.structure_report_path.is_file())
            self.assertTrue(result.quality_report_path.is_file())
            self.assertIsNone(result.empty_code_report_path)

            by_source = {item.source_id: item for item in result.positions}
            self.assertEqual(by_source["13"].group_label, "ДС13")
            self.assertEqual(by_source["47"].group_label, "ДС47")
            self.assertFalse(by_source["13"].grouping_fallback)
            self.assertEqual(by_source["4905_1"].source_id, "4905_1")
            self.assertEqual(by_source["4905_1"].group_label, "ДС4905_1")
            self.assertEqual(by_source["101"].group_label, "ДС101")

            pos8 = [item for item in result.positions if item.source_id == "8"]
            self.assertEqual(len(pos8), 1)
            self.assertEqual(pos8[0].group_label, "ДС8")
            self.assertFalse(pos8[0].grouping_fallback)
            self.assertFalse(pos8[0].overlay_blocked)

            pos15 = [item for item in result.positions if item.source_id == "15"]
            matched = next(item for item in pos15 if item.title == "1600")
            unmatched = next(item for item in pos15 if item.title == "9999")
            self.assertEqual(matched.group_label, "ДС15")
            self.assertFalse(matched.grouping_fallback)
            self.assertEqual(unmatched.group_label, "ДС15")
            self.assertFalse(unmatched.grouping_fallback)
            self.assertEqual(unmatched.group_id, "ДС15")

            offset_row = next(item for item in result.positions if item.source_id == "47")
            self.assertEqual(offset_row.leading_empty, 2)
            self.assertEqual(offset_row.code_normalized, "BCC0000001")
            self.assertEqual(offset_row.conversion_status, STATUS_IDENTITY)
            self.assertEqual(offset_row.qty, Decimal("2"))
            self.assertEqual(offset_row.qty_source, Decimal("2"))
            self.assertEqual(offset_row.conversion_coefficient, Decimal("1"))

            wb = load_workbook(result.baseline_path, data_only=True)
            try:
                ws = wb.active
                headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                data_rows = list(ws.iter_rows(min_row=2, values_only=True))
                self.assertEqual(headers[4], "Теги / Tag")
                names = {row[0] for row in data_rows if row and row[0]}
                self.assertIn("ДС13", names)
                self.assertIn("ДС4905_1", names)
                self.assertIn("ДС101", names)
                self.assertIn("ДС8", names)
                for row in data_rows:
                    if not row or row[0] is None:
                        continue
                    self.assertTrue(row[4] in {"", None})
                    self.assertEqual(row[9], row[17])
            finally:
                wb.close()

            t_com = TableComments(
                file_full_path=str(result.baseline_path),
                dir_path="-1",
                tabel_class=t_com_init_cls.RFP_AGGREGATED,
            )
            loaded = get_std_from_excel_file(t_com)
            types = Counter(row.row_type for row in loaded)
            self.assertGreaterEqual(types[RowType.position_row], 1)
            pos = [row for row in loaded if row.row_type == RowType.position_row]
            self.assertTrue(pos)
            sample = next(row for row in pos if str(row.el[CODE].value) == "BCC0000001")
            self.assertEqual(str(sample.el[TAGS].value or ""), "")
            self.assertIn("-", str(sample.el[DS_TITLE].value))
            self.assertEqual(int(sample.el[VALUES].value), int(sample.el[VALUES_2].value))
            self.assertEqual(str(sample.el[UNITS].value), "шт")
            ds_names = {str(row.el[DS_NAME].value) for row in pos}
            self.assertIn("ДС13", ds_names)
            self.assertEqual(ds_names.intersection({"13", "47"}), set())

            reused = [row for row in pos if str(row.el[DS_NAME].value) == "ДС13"]
            self.assertGreaterEqual(len(reused), 1)

            group_ids = {item.group_id for item in result.groups}
            self.assertIn("ДС13", group_ids)
            reused_group = next(item for item in result.groups if item.group_id == "ДС13")
            self.assertEqual(set(reused_group.source_ids), {"13"})
            self.assertFalse(reused_group.fallback)
            own_47 = next(item for item in result.groups if item.group_id == "ДС47")
            self.assertEqual(set(own_47.source_ids), {"47"})

            bag8 = next(item for item in result.groups if item.group_id == "ДС8")
            self.assertFalse(bag8.fallback)
            self.assertFalse(bag8.overlay_blocked)

            quality = _load_xlsx(result.quality_report_path)
            try:
                self.assertNotIn("Единицы и конвертация", quality.sheetnames)
                self.assertEqual(
                    set(quality.sheetnames),
                    {
                        "Сводка",
                        "Позиции без кода",
                        "Количества",
                        "Теги",
                        "Дубли",
                        "Коды вне Google",
                        "Источники",
                    },
                )
                src_ws = quality["Источники"]
                self.assertIsNotNone(src_ws.cell(row=2, column=2).hyperlink)
            finally:
                quality.close()

    def test_spec_header_with_tag_word_is_not_tag_column(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            header = list(DS_HEADER)
            header[3] = "Линия/ TAG-Номер/ Спецификация"
            spec = "AGCC.287-8630-KSB4.MTO-0001"
            _write_xlsx(
                ds_root / "ДС13.xlsx",
                [
                    header,
                    _ds_row(spec=spec, qty=2),
                    _ds_row(npp=2, spec=spec, qty=4),
                ],
            )
            result = _run_baseline(ds_root, registry_path, out_dir)
            self.assertFalse(result.blocking, result.summary_line())
            self.assertEqual(len(result.positions), 2)
            self.assertEqual(
                {item.qty for item in result.positions},
                {Decimal("2"), Decimal("4")},
            )
            self.assertTrue(all(not item.diagnostic_tags for item in result.positions))
            self.assertFalse(
                any(item.code == ISSUE_TAG_DUPLICATE for item in result.issues)
            )
            self.assertFalse(
                any(item.code == ISSUE_TAG_MISMATCH for item in result.issues)
            )

    def test_non_canon_shift_emits_one_layout_issue_and_no_positions(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            _write_xlsx(
                ds_root / "ДС99_shift.xlsx",
                [DS_HEADER[:7] + ["служебная"] + DS_HEADER[7:], _ds_row()],
            )
            result = _run_baseline(ds_root, registry_path, out_dir)
            layout = [
                item for item in result.issues if item.code == ISSUE_INTERNAL_SHIFT
            ]
            self.assertEqual(len(layout), 1)
            self.assertIn("Раскладка не канон", layout[0].message)
            self.assertEqual(layout[0].excel_row, 1)
            self.assertFalse(
                any(item.code == "missing_core_role" for item in result.issues)
            )
            self.assertEqual(result.positions, [])
            self.assertTrue(result.blocking)

    def test_no_ds_header_says_active_sheet_failed_header_check(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            header = list(DS_HEADER)
            header[5] = "Прочее"
            _write_xlsx(ds_root / "ДС99_upd.xlsx", [header, _ds_row()])
            result = _run_baseline(ds_root, registry_path, out_dir)
            missing = [item for item in result.issues if item.code == ISSUE_NO_HEADER]
            self.assertEqual(len(missing), 1, result.issues)
            message = missing[0].message
            self.assertIn("не прошёл проверку шапки ДС", message)
            self.assertIn("нет обязательных ролей: Наименование", message)
            self.assertNotIn("единственным data-sheet", message)
            self.assertEqual(result.positions, [])
            self.assertTrue(result.blocking)

    def test_canon_footer_itogo_sum_is_not_a_position(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            _write_xlsx(
                ds_root / "ДС13.xlsx",
                [
                    DS_HEADER,
                    _ds_row(npp=1, qty=2),
                    _ds_row(npp="ИТОГО:", qty="=SUM(L2:L2)", code=""),
                ],
            )
            result = _run_baseline(ds_root, registry_path, out_dir)
            self.assertEqual(len(result.positions), 1)
            self.assertEqual(result.positions[0].qty, Decimal("2"))
            self.assertFalse(
                any(item.code == ISSUE_QTY_FORMULA for item in result.issues)
            )
            self.assertFalse(
                any(item.code == ISSUE_EMPTY_CODE for item in result.issues)
            )
            self.assertFalse(result.blocking, result.summary_line())

    def test_contract_footer_without_material_fields_is_not_a_position(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            blank = [""] * 12

            def clause(title: str = "", name: str = "") -> list[object]:
                row = list(blank)
                row[1] = title
                row[5] = name
                return row

            _write_xlsx(
                ds_root / "ДС13.xlsx",
                [
                    DS_HEADER,
                    _ds_row(npp=1, qty=2),
                    clause(title="2. Условия уже отсечены отдельной фразой нет"),
                    clause(title="Покупатель перечисляет денежные средства"),
                    clause(name='ООО "Би.Си.Си."'),
                    clause(title="к/счет: 30101810200000000823"),
                    clause(name="___________________ / Сергеев С.Д. /"),
                    _ds_row(
                        npp=2,
                        code="",
                        code_1c="",
                        name="Металлодетектор",
                        units="шт",
                        qty=2,
                    ),
                    _ds_row(
                        npp=3,
                        code="",
                        code_1c="",
                        name="Кабель без количества",
                        units="м",
                        qty=None,
                    ),
                ],
            )
            result = _run_baseline(ds_root, registry_path, out_dir)
            self.assertEqual(len(result.positions), 3)
            self.assertEqual(
                [item.excel_row for item in result.positions],
                [2, 8, 9],
            )
            empty_codes = [
                item for item in result.issues if item.code == ISSUE_EMPTY_CODE
            ]
            self.assertEqual([item.excel_row for item in empty_codes], [8, 9])
            qty_empty = [
                item for item in result.issues if item.code == ISSUE_QTY_EMPTY
            ]
            self.assertEqual([item.excel_row for item in qty_empty], [9])

    def test_signature_and_ref_footer_are_not_positions(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            blank = [""] * 12

            def put(column: int, text: str) -> list[object]:
                row = list(blank)
                row[column] = text
                return row

            ref_row = ["#REF!"] * 12
            ref_row[2] = ""
            _write_xlsx(
                ds_root / "ДС13.xlsx",
                [
                    DS_HEADER,
                    _ds_row(npp=1, qty=2),
                    put(6, "/Савва С.В./"),
                    put(6, "/Никифоров И.С./"),
                    put(6, "_________________________________________/Савва С."),
                    put(11, "____________________________________/Григорьев М.С./"),
                    ref_row,
                    _ds_row(
                        npp=2,
                        code="",
                        code_1c="002601499",
                        name="Кабель без кода РД",
                        units="",
                        qty=None,
                    ),
                    _ds_row(
                        npp=3,
                        code="",
                        code_1c="",
                        name="Металлодетектор",
                        units="шт",
                        qty=2,
                    ),
                ],
            )
            result = _run_baseline(ds_root, registry_path, out_dir)
            self.assertEqual(
                [item.excel_row for item in result.positions],
                [2, 8, 9],
            )
            empty_codes = [
                item.excel_row
                for item in result.issues
                if item.code == ISSUE_EMPTY_CODE
            ]
            self.assertEqual(empty_codes, [8, 9])
            self.assertFalse(
                any(item.code == ISSUE_QTY_FORMULA for item in result.issues)
            )
            self.assertFalse(
                any(item.code == "qty_non_numeric" for item in result.issues)
            )
            qty_empty = [
                item.excel_row
                for item in result.issues
                if item.code == ISSUE_QTY_EMPTY
            ]
            self.assertEqual(qty_empty, [8])

    def test_empty_code_row_stays_a_position(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            _write_xlsx(
                ds_root / "ДС13.xlsx",
                [
                    DS_HEADER,
                    _ds_row(
                        code="",
                        title="8529",
                        name="Кабель",
                        units="шт",
                        qty=2,
                    ),
                ],
            )
            result = _run_baseline(ds_root, registry_path, out_dir)
            self.assertEqual(len(result.positions), 1)
            self.assertFalse(result.positions[0].code_normalized)
            self.assertEqual(result.positions[0].qty, Decimal("2"))
            self.assertTrue(
                any(item.code == ISSUE_EMPTY_CODE for item in result.issues)
            )
            self.assertTrue(result.blocking)


class DsBaselineInflatedSheetSmokeTest(unittest.TestCase):
    def test_inflated_dimensions_keeps_positions_and_progress(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)

            path = ds_root / "ДС13.xlsx"
            path.parent.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            ws = wb.active
            assert ws is not None
            ws.title = "Спека"
            ws.append(DS_HEADER + ["Tag"])
            ws.append(_ds_row(npp=1, extra=["T-A"]))
            ws.append(_ds_row(npp=2, extra=["T-B"]))
            far = ws.cell(row=100_000, column=16_384)
            far.value = None
            far.font = Font(bold=True)
            wb.save(path)
            wb.close()

            probe = load_workbook(path, read_only=True, data_only=False)
            try:
                dims = probe.active
                self.assertGreaterEqual(dims.max_row or 0, 100_000)
                self.assertGreaterEqual(dims.max_column or 0, 16_384)
            finally:
                probe.close()

            progress: list[tuple[int, int, str, float | None]] = []
            registry = load_registry(registry_path)
            started = time.perf_counter()
            result = build_ds_baseline(
                ds_root,
                registry,
                out_dir,
                write_baseline=True,
                converter=IdentityDsUnitsConverter(),
                google_index=_google_index(),
                stamp=STAMP,
                progress_callback=lambda index, total, relpath, elapsed: progress.append(
                    (index, total, relpath, elapsed)
                ),
            )
            elapsed = time.perf_counter() - started
            self.assertLess(elapsed, 20.0, f"inflated sheet parse took {elapsed:.1f}s")
            self.assertFalse(result.blocking, result.summary_line())
            self.assertEqual(len(result.positions), 2)
            self.assertEqual(
                {item.ds_number for item in result.positions},
                {"1", "2"},
            )
            self.assertEqual(result.positions[0].diagnostic_tags, ("T-A",))
            self.assertEqual(result.positions[1].diagnostic_tags, ("T-B",))
            self.assertEqual(
                [(item[0], item[1], item[2]) for item in progress],
                [(1, 1, "ДС13.xlsx"), (1, 1, "ДС13.xlsx")],
            )
            self.assertIsNone(progress[0][3])
            self.assertIsNotNone(progress[1][3])


class DsBaselineTagDuplicateSmokeTest(unittest.TestCase):
    def test_shared_tag_emits_one_issue_and_phase_elapsed(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_root = root / "ds"
            out_dir = root / "out"
            registry_path = root / "registry.xlsx"
            _write_registry(registry_path)
            _write_xlsx(
                ds_root / "ДС13.xlsx",
                [DS_HEADER + ["Tag"], _ds_row(extra=["T-SHARE"])],
            )
            _write_xlsx(
                ds_root / "ДС_47" / "spec.xlsx",
                [DS_HEADER + ["Tag"], _ds_row(npp=2, extra=["T-SHARE"])],
            )
            phases: list[str] = []
            registry = load_registry(registry_path)
            result = build_ds_baseline(
                ds_root,
                registry,
                out_dir,
                write_baseline=True,
                converter=IdentityDsUnitsConverter(),
                google_index=_google_index(),
                stamp=STAMP,
                phase_callback=phases.append,
            )
            cross = [
                item
                for item in result.issues
                if item.code == ISSUE_TAG_DUPLICATE and "T-SHARE" in item.message
            ]
            self.assertEqual(len(cross), 1, cross)
            self.assertIn("в 2 строках", cross[0].message)
            self.assertTrue(
                any("конвертация единиц измерения" in line for line in phases),
                phases,
            )
            self.assertTrue(
                any("дубли тегов между строками" in line for line in phases),
                phases,
            )
            self.assertTrue(
                any(" — " in line and " с" in line for line in phases),
                phases,
            )


class DsBaselinePathUriSmokeTest(unittest.TestCase):
    def test_absolute_uri_does_not_call_resolve(self) -> None:
        _path_uri.cache_clear()
        with patch.object(Path, "resolve", side_effect=AssertionError("resolve")):
            uri = _path_uri(r"C:\tmp\ds.xlsx")
        self.assertTrue(str(uri).startswith("file:"))
        unc = _path_uri(r"\\bcc\eng\ds.xlsx")
        self.assertTrue(str(unc).startswith("file:"))


if __name__ == "__main__":
    unittest.main()
