"""Offline regression tests for automatic RFP DS checklist updates."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.rfp_parts import build_ds_checklist_xlsx as updater
from RFQ.rfp_parts.ds_checklist import load_ds_checklist


def _write_summary(path: Path, rows: list[tuple[object, object]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Би.Си.Си."
    sheet.cell(1, 3, "№ ДС")
    sheet.cell(1, 4, "Описание")
    for row_index, (number, description) in enumerate(rows, start=2):
        sheet.cell(row_index, 3, number)
        sheet.cell(row_index, 4, description)
    workbook.save(path)
    workbook.close()


class ChecklistUpdaterSmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.summary = self.root / "summary.xlsx"
        self.checklist = self.root / "checklist" / "ДС.xlsx"
        self.increase = self.root / "increase"
        self.decrease = self.root / "decrease"
        self.increase.mkdir()
        self.decrease.mkdir()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _update(self) -> updater.ChecklistUpdateResult:
        return updater.ensure_ds_checklist_current(
            self.summary,
            self.checklist,
            self.increase,
            self.decrease,
        )

    def test_missing_checklist_is_created_then_unchanged_without_rewrite(self) -> None:
        _write_summary(self.summary, [(10, "Увеличение"), ("20/1А", "Уменьшение")])

        first = self._update()
        self.assertEqual(first.status, "updated")
        self.assertEqual(first.changes_count, 2)
        self.assertTrue(self.checklist.is_file())
        first_mtime = self.checklist.stat().st_mtime_ns
        first_bytes = self.checklist.read_bytes()

        second = self._update()
        self.assertEqual(second.status, "unchanged")
        self.assertEqual(second.changes_count, 0)
        self.assertEqual(self.checklist.stat().st_mtime_ns, first_mtime)
        self.assertEqual(self.checklist.read_bytes(), first_bytes)

    def test_semantic_source_change_rewrites_checklist(self) -> None:
        _write_summary(self.summary, [(10, "Увеличение")])
        self.assertEqual(self._update().status, "updated")

        _write_summary(self.summary, [(10, "Увеличение"), (11, "Увеличение")])
        result = self._update()

        self.assertEqual(result.status, "updated")
        self.assertEqual(result.changes_count, 1)
        checklist = load_ds_checklist(self.checklist)
        self.assertEqual([entry.ds for entry in checklist.increase], ["10", "11"])

    def test_not_required_survives_normalized_identity_and_category_move(self) -> None:
        _write_summary(self.summary, [(14, "Увеличение")])
        self.assertEqual(self._update().status, "updated")

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "ДС"
        sheet["A1"] = "ДС на увеличение"
        sheet["B1"] = "В папке увеличение"
        sheet["A2"] = " ДС 14 "
        sheet["B2"] = "  не ТРЕБУЕТСЯ "
        workbook.save(self.checklist)
        workbook.close()

        self.assertEqual(self._update().status, "unchanged")
        _write_summary(self.summary, [(14, "Уменьшение по договору")])
        moved = self._update()

        self.assertEqual(moved.status, "updated")
        self.assertEqual(moved.changes_count, 2)
        workbook = openpyxl.load_workbook(
            io.BytesIO(self.checklist.read_bytes()), data_only=True
        )
        try:
            sheet = workbook["ДС"]
            self.assertEqual(sheet["D2"].value, "14")
            self.assertEqual(sheet["E2"].value, "Не требуется")
        finally:
            workbook.close()

    def test_missing_source_returns_error(self) -> None:
        result = self._update()

        self.assertEqual(result.status, "error")
        self.assertIsNotNone(result.error)
        self.assertIn("ERROR:", result.summary_line)
        self.assertFalse(self.checklist.exists())

    def test_missing_source_keeps_existing_checklist(self) -> None:
        self.checklist.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "ДС"
        sheet["A1"] = "ДС на увеличение"
        sheet["B1"] = "В папке увеличение"
        sheet["A2"] = "10"
        sheet["B2"] = "ЕСТЬ"
        workbook.save(self.checklist)
        workbook.close()
        (self.increase / "ДС10.xlsx").write_bytes(b"")
        original = self.checklist.read_bytes()

        result = self._update()

        self.assertEqual(result.status, "unchanged")
        self.assertIsNone(result.error)
        self.assertIn("источник списка ДС недоступен", result.summary_line)
        self.assertIn("WARN:", result.summary_line)
        self.assertIsNotNone(result.compare_result)
        self.assertIsNone(result.compare_result.load_error)
        self.assertEqual(result.compare_result.error_count, 0)
        self.assertEqual(self.checklist.read_bytes(), original)

    def test_missing_exact_path_uses_glob_match(self) -> None:
        glob_summary = self.root / "Сводная таблица ДС по вед.договорам extra.xlsx"
        _write_summary(glob_summary, [(10, "Увеличение")])
        (self.increase / "ДС10.xlsx").write_bytes(b"")
        missing_exact = self.root / "Сводная таблица ДС по вед.договорам.xlsx"

        result = updater.ensure_ds_checklist_current(
            missing_exact,
            self.checklist,
            self.increase,
            self.decrease,
        )

        self.assertEqual(result.status, "updated")
        self.assertEqual(result.summary_path, glob_summary)
        self.assertTrue(self.checklist.is_file())
        checklist = load_ds_checklist(self.checklist)
        self.assertEqual([entry.ds for entry in checklist.increase], ["10"])

    def test_replace_failure_is_error_and_leaves_no_temporary_file(self) -> None:
        _write_summary(self.summary, [(10, "Увеличение")])

        with mock.patch.object(updater.os, "replace", side_effect=OSError("replace denied")):
            result = self._update()

        self.assertEqual(result.status, "error")
        self.assertIn("replace denied", result.error or "")
        self.assertFalse(self.checklist.exists())
        self.assertEqual(list(self.checklist.parent.glob(f".{self.checklist.stem}.*.xlsx")), [])

    def test_folder_mismatch_does_not_override_update_status(self) -> None:
        _write_summary(self.summary, [(99, "Увеличение")])

        updated = self._update()
        unchanged = self._update()

        self.assertEqual(updated.status, "updated")
        self.assertEqual(unchanged.status, "unchanged")
        for result in (updated, unchanged):
            self.assertIsNotNone(result.compare_result)
            self.assertGreater(result.compare_result.error_count, 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
