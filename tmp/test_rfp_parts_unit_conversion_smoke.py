"""Offline smoke: parts unit conversion before counters/coarse collisions."""

from __future__ import annotations

import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from RFQ.units_convert.models import PAIR_COEF_PREFIX

from RFQ.rfp_parts.analyze_rfp_parts import (
    NetSummaryRow,
    RfpRecord,
    _apply_parts_conversion_plan,
    _assert_unified_target_units,
    _build_parts_conversion_requests,
    _build_summary_rows,
    _build_unit_counters,
    _compute_coarse_collisions,
    _write_net_xlsx,
    build_parts_build_deps_snapshot,
    step1_net_extra_columns,
)
import RFQ.units_convert.matrix as matrix_module
from RFQ.units_convert import (
    ConversionInvariant,
    ConversionPlan,
    ConversionRequest,
    GoogleUnitsIndex,
    UnitsConversionError,
    build_conversion_plan,
    build_google_units_index,
)
from base.base_classes import RowStd, RowType, TableComments
from base.base_mto import get_std_from_excel_file
from base.tables_columns import (
    CODE,
    ColNames,
    UNITS,
    UNITS_CHECK_STATUS,
    UNITS_CONVERSION_TRACE,
    VALUES,
)
import base.t_comm_initial_classes as t_com_init_cls


def _google_index(*pairs: tuple[str, str]) -> GoogleUnitsIndex:
    rows: list[RowStd] = []
    for code, units in pairs:
        row = RowStd()
        row.el[CODE].value = code
        row.el[UNITS].value = units
        rows.append(row)
    return build_google_units_index(rows)


def _record(
    *,
    code: str = "BCC0000516",
    units: str = "шт",
    values: Decimal = Decimal("3"),
    tags: str = "",
    file_name: str = "DS10.xlsx",
    excel_row: int = 24,
    ds_name: str = "ДС10",
) -> RfpRecord:
    return RfpRecord(
        kind="parts",
        file_name=file_name,
        sheet="Перечень материалов",
        excel_row=excel_row,
        ds_name=ds_name,
        ds_number="10",
        ds_title="8529-SOS",
        ds_specification="8529-SOS",
        tags=tags,
        ds_code_1c="",
        code=code,
        name="Cable",
        type_mark="",
        values=values,
        units=units,
    )


class PartsUnitConversionSmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.matrix_path = Path(self.temp_dir.name) / "units_matrix.xlsx"
        self.google = _google_index(("BCC0000516", "шт"))
        matrix_module._before_replace_hook = None

    def tearDown(self) -> None:
        matrix_module._before_replace_hook = None
        self.temp_dir.cleanup()

    def _seed_nonidentity(self, *, source_unit: str, google_unit: str = "м") -> None:
        google = _google_index(("BCC0000516", google_unit))
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [
                    ConversionRequest(
                        request_id="seed",
                        contour="parts",
                        code="BCC0000516",
                        source_unit=source_unit,
                        quantity="1",
                        tags_count=0,
                        invariant=ConversionInvariant.NONE,
                        location="seed",
                    )
                ],
                google,
                self.matrix_path,
            )

    def _set_matrix_coef(self, coef: object) -> None:
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            coef_col = None
            for col_idx in range(1, ws.max_column + 1):
                title = str(ws.cell(row=1, column=col_idx).value or "")
                if title.startswith(PAIR_COEF_PREFIX):
                    coef_col = col_idx
                    break
            assert coef_col is not None
            ws.cell(row=2, column=coef_col, value=coef)
            wb.save(self.matrix_path)
        finally:
            wb.close()

    def _plan(self, records: list[RfpRecord]) -> ConversionPlan:
        requests = _build_parts_conversion_requests(records)
        return build_conversion_plan(requests, self.google, self.matrix_path)

    def test_conversion_before_counters_and_coarse(self) -> None:
        records = [_record(units="шт."), _record(units="шт", file_name="DS11.xlsx")]
        plan = self._plan(records)
        converted = _apply_parts_conversion_plan(records, plan)
        _assert_unified_target_units(converted)
        coarse = _compute_coarse_collisions(converted)
        self.assertEqual(len(coarse.ambiguous_meta), 0)
        aggregation = _build_unit_counters(converted, [])
        self.assertEqual(sum(aggregation.units.values()), 6)

    def test_tagged_invariant_fatal(self) -> None:
        self._seed_nonidentity(source_unit="шт", google_unit="м")
        self._set_matrix_coef("0.5")
        tagged = _record(
            tags="8529-SS-01-S-UZ-0711;8529-SS-01-S-UZ-0712",
            values=Decimal("2"),
            units="шт",
        )
        requests = _build_parts_conversion_requests([tagged])
        self.assertEqual(requests[0].invariant, ConversionInvariant.TAGS_EQUAL)
        self.assertEqual(requests[0].tags_count, 2)
        google = _google_index(("BCC0000516", "м"))
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(requests, google, self.matrix_path)

    def test_no_partial_records_on_fatal(self) -> None:
        self._seed_nonidentity(source_unit="шт", google_unit="м")
        self._set_matrix_coef("0.5")
        tagged = _record(
            tags="8529-SS-01-S-UZ-0711;8529-SS-01-S-UZ-0712",
            values=Decimal("2"),
            units="шт",
        )
        google = _google_index(("BCC0000516", "м"))
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                _build_parts_conversion_requests([tagged]),
                google,
                self.matrix_path,
            )
        self.assertEqual(tagged.units_check_status, "")
        self.assertEqual(tagged.units_conversion_trace, "")

    def test_normalized_spelling_not_collision(self) -> None:
        records = [_record(units="шт."), _record(units="шт", file_name="DS11.xlsx")]
        converted = _apply_parts_conversion_plan(records, self._plan(records))
        coarse = _compute_coarse_collisions(converted)
        self.assertEqual(coarse.ambiguous_meta, [])

    def test_all_traces_merged_into_summary(self) -> None:
        first = _record(
            units="шт.",
            file_name="DS10.xlsx",
            excel_row=10,
        )
        second = _record(
            units="шт",
            file_name="DS11.xlsx",
            excel_row=11,
        )
        converted = _apply_parts_conversion_plan(
            [first, second],
            self._plan([first, second]),
        )
        aggregation = _build_unit_counters(converted, [])
        key = next(iter(aggregation.units))
        self.assertIn("identity", aggregation.units_status_by_key[key])
        self.assertIn("code=", aggregation.units_trace_by_key[key])
        self.assertIn(";", aggregation.units_trace_by_key[key])
        rows = _build_summary_rows(aggregation, _compute_coarse_collisions(converted))
        self.assertEqual(len(rows), 1)
        self.assertIn("identity", rows[0].record.units_check_status)
        self.assertIn("code=", rows[0].record.units_conversion_trace)

    def test_net_status_trace_and_filters_layout(self) -> None:
        item = NetSummaryRow(
            record=RfpRecord(
                kind="net",
                file_name="DS10.xlsx",
                sheet="Перечень материалов",
                excel_row=24,
                ds_name="ДС10",
                ds_number="1",
                ds_title="8529-SOS",
                ds_specification="8529-SOS",
                tags="",
                ds_code_1c="",
                code="BCC0000516",
                name="Cable",
                type_mark="",
                values=Decimal("3"),
                units="шт",
                units_check_status="identity",
                units_conversion_trace="code=BCC0000516; status=identity",
            ),
            source="Сумма частей",
        )
        path = Path(self.temp_dir.name) / "net_layout.xlsx"
        self.assertEqual(_write_net_xlsx(path, [item]), 1)
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Перечень материалов"]
        rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
        wb.close()
        header, data = rows[0], rows[1]

        extra = step1_net_extra_columns()
        self.assertEqual(extra[18], UNITS_CHECK_STATUS)
        self.assertEqual(extra[19], UNITS_CONVERSION_TRACE)
        self.assertEqual(header[18], "Статус ед. изм.")
        self.assertEqual(header[19], "Trace ед. изм.")
        self.assertEqual(header[20], "Источник строки")
        self.assertEqual(header[25], "Через замену кода")
        self.assertEqual(data[18], "identity")
        self.assertIn("code=", str(data[19]))

        t_com = TableComments(
            file_full_path=str(path),
            dir_path="-1",
            tabel_class=t_com_init_cls.RFP_AGGREGATED,
        )
        loaded = get_std_from_excel_file(t_com)
        pos = next(row for row in loaded if row.row_type == RowType.position_row)
        self.assertEqual(pos.el[UNITS_CHECK_STATUS].value, "identity")
        self.assertIn("code=", str(pos.el[UNITS_CONVERSION_TRACE].value))
        self.assertEqual(int(pos.el[VALUES].value), 3)
        self.assertEqual(pos.el[UNITS].value, "шт")
        self.assertEqual(pos.el[CODE].value, "BCC0000516")

    def test_build_deps_snapshot_roundtrip(self) -> None:
        records = [_record()]
        plan = self._plan(records)
        snapshot = build_parts_build_deps_snapshot(
            plan,
            google_index=self.google,
            matrix_path=self.matrix_path,
        )
        self.assertEqual(snapshot["ALGORITHM_VERSION"], "v5")
        self.assertEqual(snapshot["AGGREGATION_KEY"], "ds_name+title+code+tag")
        self.assertTrue(snapshot["dependencies"])
        self.assertEqual(snapshot["google_units_by_code"]["BCC0000516"], "шт")

    def test_same_position_different_ds_not_merged(self) -> None:
        first = _record(
            ds_name="ДС43",
            file_name="ДС43.xlsx",
            values=Decimal("80"),
        )
        second = _record(
            ds_name="ДС47_13",
            file_name="ДС47_13.xlsx",
            values=Decimal("80"),
        )
        aggregation = _build_unit_counters([first, second], [], emit_warnings=False)
        self.assertEqual(len(aggregation.units), 2)
        self.assertEqual(sum(aggregation.units.values()), 160)
        rows = _build_summary_rows(
            aggregation, _compute_coarse_collisions([first, second])
        )
        self.assertEqual({row.record.ds_name for row in rows}, {"ДС43", "ДС47_13"})
        self.assertEqual(
            [row.record.values for row in rows], [Decimal("80"), Decimal("80")]
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
