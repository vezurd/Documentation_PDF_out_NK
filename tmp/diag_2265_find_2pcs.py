"""Find where 2 pcs of 2265-KSB / DС67 went in yesterday Step4."""
from __future__ import annotations

import sys

from openpyxl import load_workbook

XLSX = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.09.02.13.49"
    r"\Шаг4_Сопоставление_RFP_MTO_20260902_140222.xlsx"
)


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def _norm(v: object) -> str:
    return "" if v is None else str(v).strip()


def main() -> None:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [_norm(v) for v in next(rows)]
    idx = {n: i for i, n in enumerate(header) if n}

    def cell(row, name):
        i = idx.get(name)
        if i is None or i >= len(row):
            return None
        return row[i]

    _safe_print("--- all 2265-KSB rows with tag FV-0711 or code 2772 or DS67 coupling ---")
    n = 0
    sums: dict[str, float] = {}
    for rn, row in enumerate(rows, start=2):
        title = _norm(cell(row, "Титул/Марка"))
        if title != "2265-KSB":
            continue
        ds = _norm(cell(row, "Порядковый ДС"))
        actual = _norm(cell(row, "Фактический ДС"))
        code_rfp = _norm(cell(row, "Код RFP")).upper()
        code_mto = _norm(cell(row, "Код MTO")).upper()
        code_ul = _norm(cell(row, "Код УЛ")).upper()
        tags = _norm(cell(row, "Теги RFP"))
        name = _norm(cell(row, "Наименование RFP"))
        ul_name = _norm(cell(row, "Наименование УЛ"))
        pos = _norm(cell(row, "№ позиции"))
        hit = (
            "2265-S-FV-0711" in tags
            or "BCC0002772" in (code_rfp, code_mto, code_ul)
            or (
                actual == "ДС67"
                and "BCC0000642" in (code_rfp, code_mto, code_ul)
            )
            or (
                ds == "ДС67"
                and ("Муфта вводная" in name or "Муфта вводная" in ul_name)
            )
        )
        if not hit:
            continue
        n += 1
        qty = cell(row, "Кол-во RFP")
        _safe_print(
            f"r{rn} ds={ds!r} actual={actual!r} pos={pos!r} "
            f"rfp={code_rfp!r} qty={qty!r} tags={tags!r} "
            f"mto={code_mto!r} mto_qty={cell(row, 'Количество MTO')!r} "
            f"ul={code_ul!r} ul_qty={cell(row, 'Кол-во по УЛ')!r} "
            f"ul_status={cell(row, 'Статус УЛ')!r} "
            f"pos_status={cell(row, 'MTO, Статус позиции')!r} "
            f"name={name[:60]!r}"
        )
        try:
            q = float(qty) if qty not in (None, "") else 0.0
        except (TypeError, ValueError):
            q = 0.0
        key = f"{ds}|{code_rfp or code_mto or code_ul}"
        sums[key] = sums.get(key, 0.0) + q
    _safe_print(f"rows={n}")
    _safe_print("rfp_qty_sums=" + str(sums))
    wb.close()


if __name__ == "__main__":
    main()
