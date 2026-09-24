"""Offline smoke tests for RFQ.units_convert matrix/plan workflow."""

from __future__ import annotations

import copy
import os
import sys
import tempfile
import time
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

import RFQ.units_convert.matrix as matrix_module
from RFQ.units_convert import (
    ALGORITHM_VERSION,
    ConversionInvariant,
    ConversionRequest,
    RowStdBinding,
    UnitsConversionError,
    action_by_id,
    apply_conversion_plan,
    build_conversion_plan,
    build_google_units_index,
    converted_quantity,
    converted_unit,
    normalize_code,
)
from RFQ.units_convert.matrix import (
    DATA_ALIGNMENT,
    HEADER_ALIGNMENT,
    HEADER_ROW_HEIGHT,
    MATRIX_FIXED_COLUMN_WIDTHS,
    PAIR_1_COLUMN_WIDTHS,
    PAIR_2_PLUS_COLUMN_WIDTHS,
    PAIR_COLUMN_STRIDE,
    PAIR_START_COL,
    _parse_coefficient,
    expected_matrix_column_width,
    load_matrix,
    normalize_reciprocal_coefficient,
)
from RFQ.units_convert.models import (
    MATRIX_HEADER,
    PAIR_COEF_PREFIX,
    PAIR_FORMULA_PREFIX,
    STATUS_IDENTITY,
    STATUS_NO_GOOGLE,
    ConversionDependency,
)
from base.base_classes import CheckElement, RowStd
from base.tables_columns import CODE, NAME, UNITS, UNITS_CHECK_STATUS, UNITS_CONVERSION_TRACE, VALUES, VALUES_2


def _google_row(code: str, units: str, *, name: str = "") -> RowStd:
    row = RowStd()
    row.el[CODE].value = code
    row.el[UNITS].value = units
    if name:
        row.el[NAME].value = name
    return row


def _header_column(ws, title: str) -> int:
    for col_idx in range(1, ws.max_column + 1):
        if str(ws.cell(row=1, column=col_idx).value or "").strip() == title:
            return col_idx
    raise AssertionError(f"Header column {title!r} not found")


def _apply_persisted_matrix_formatting(ws, *, pair_count: int, last_row: int) -> None:
    last_col = 4 + pair_count * PAIR_COLUMN_STRIDE
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT
    for col in range(1, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = expected_matrix_column_width(
            col,
            pair_count=pair_count,
        )
        ws.cell(row=1, column=col).alignment = HEADER_ALIGNMENT
    for row_index in range(2, last_row + 1):
        for col in range(1, last_col + 1):
            ws.cell(row=row_index, column=col).alignment = DATA_ALIGNMENT


def _assert_persisted_matrix_formatting(
    test_case: unittest.TestCase,
    ws,
    *,
    pair_count: int,
    last_row: int,
) -> None:
    last_col = 4 + pair_count * PAIR_COLUMN_STRIDE
    test_case.assertAlmostEqual(
        ws.row_dimensions[1].height,
        HEADER_ROW_HEIGHT,
        places=6,
    )
    for col in range(1, last_col + 1):
        letter = get_column_letter(col)
        test_case.assertAlmostEqual(
            ws.column_dimensions[letter].width,
            expected_matrix_column_width(col, pair_count=pair_count),
            places=6,
            msg=f"column {letter}",
        )
        header_cell = ws.cell(row=1, column=col)
        test_case.assertEqual(header_cell.alignment.horizontal, "left")
        test_case.assertEqual(header_cell.alignment.vertical, "center")
        test_case.assertTrue(header_cell.alignment.wrap_text)
    for row_index in range(2, last_row + 1):
        for col in range(1, last_col + 1):
            cell = ws.cell(row=row_index, column=col)
            test_case.assertEqual(cell.alignment.horizontal, "left")
            test_case.assertIn(cell.alignment.vertical, (None, "bottom"))
            test_case.assertTrue(cell.alignment.wrap_text)


def _coef_column(ws, pair_index: int) -> int:
    return _header_column(ws, f"{PAIR_COEF_PREFIX}{pair_index}")


def _data_row(
    *,
    code: str,
    units: str,
    qty: object,
    status: str = "",
    trace: str = "",
    slim: bool = False,
) -> RowStd:
    row = RowStd()
    row.el[CODE].value = code
    row.el[UNITS].value = units
    row.el[VALUES].value = qty
    if not slim:
        row.el[UNITS_CHECK_STATUS].value = status or None
        row.el[UNITS_CONVERSION_TRACE].value = trace or None
    else:
        row.el.pop(UNITS_CHECK_STATUS, None)
        row.el.pop(UNITS_CONVERSION_TRACE, None)
    return row


def _request(
    request_id: str,
    *,
    code: str,
    units: str,
    qty: object,
    tags: int = 0,
    invariant: ConversionInvariant = ConversionInvariant.NONE,
    contour: str = "rfp",
    location: str = "",
    name: str = "",
) -> ConversionRequest:
    return ConversionRequest(
        request_id=request_id,
        contour=contour,
        code=code,
        source_unit=units,
        quantity=qty,
        tags_count=tags,
        invariant=invariant,
        location=location or request_id,
        item_name=name,
    )


def _binding(request_id: str, row: RowStd) -> RowStdBinding:
    return RowStdBinding(
        request_id=request_id,
        row=row,
        quantity_column=VALUES,
        units_column=UNITS,
        status_column=UNITS_CHECK_STATUS,
        trace_column=UNITS_CONVERSION_TRACE,
    )


class UnitsConvertMatrixSmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.matrix_path = Path(self._tmpdir.name) / "units_matrix.xlsx"
        matrix_module._before_replace_hook = None

    def tearDown(self) -> None:
        matrix_module._before_replace_hook = None
        self._tmpdir.cleanup()
        import gc

        gc.collect()

    def _seed_identity_matrix(self, code: str, google_unit: str) -> None:
        google = build_google_units_index([_google_row(code, google_unit)])
        build_conversion_plan(
            [_request("seed", code=code, units=google_unit, qty="1")],
            google,
            self.matrix_path,
        )

    def _seed_nonidentity_placeholder(self, code: str, google_unit: str, source_unit: str) -> None:
        google = build_google_units_index([_google_row(code, google_unit)])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("seed", code=code, units=source_unit, qty="1")],
                google,
                self.matrix_path,
            )

    def _set_matrix_coef(self, row_index: int, pair_index: int, coef: object) -> None:
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            coef_col = _coef_column(ws, pair_index)
            ws.cell(row=row_index, column=coef_col, value=coef)
            wb.save(self.matrix_path)
        finally:
            wb.close()

    def _set_matrix_google_unit(self, row_index: int, unit: str) -> None:
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            ws.cell(row=row_index, column=4, value=unit)
            wb.save(self.matrix_path)
        finally:
            wb.close()

    def _formula_gray_rgb_suffix(self) -> str:
        return "D9D9D9"

    def _find_question_mark_coef_cell(self) -> str:
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            for row_idx in range(2, ws.max_row + 1):
                for col_idx in range(1, ws.max_column + 1):
                    title = str(ws.cell(row=1, column=col_idx).value or "")
                    if not title.startswith(PAIR_COEF_PREFIX):
                        continue
                    if ws.cell(row=row_idx, column=col_idx).value == "?":
                        return ws.cell(row=row_idx, column=col_idx).coordinate
            self.fail("Expected '?' coefficient cell was not written")
        finally:
            wb.close()

    def test_algorithm_version(self) -> None:
        self.assertEqual(ALGORITHM_VERSION, "v5")

    def test_normalize_code(self) -> None:
        self.assertEqual(normalize_code(" bcc 0001 "), "BCC0001")

    def test_google_name_captured_first_non_empty(self) -> None:
        rows = [
            _google_row("BCC0025", "шт", name=""),
            _google_row("BCC0025", "шт", name="Первая позиция"),
            _google_row("bcc0025", "шт", name="Вторая позиция"),
        ]
        index = build_google_units_index(rows)
        self.assertEqual(index.google_name("BCC0025"), "Первая позиция")

    def test_duplicate_same_unit_different_name_nonfatal(self) -> None:
        rows = [
            _google_row("BCC0026", "шт", name="Имя A"),
            _google_row("BCC0026", "шт", name="Имя B"),
        ]
        index = build_google_units_index(rows)
        self.assertEqual(index.google_unit("BCC0026"), "шт")
        self.assertEqual(index.google_name("BCC0026"), "Имя A")

    def test_conflicting_google_units_fatal(self) -> None:
        rows = [
            _google_row("BCC0001", "шт"),
            _google_row("bcc0001", "м"),
        ]
        with self.assertRaises(UnitsConversionError) as ctx:
            build_google_units_index(rows)
        self.assertIn("Конфликт единиц измерения Google", str(ctx.exception))
        self.assertNotIn("Conflicting Google units", str(ctx.exception))

    def test_identity_conversion(self) -> None:
        google = build_google_units_index([_google_row("BCC0001", "шт")])
        row = _data_row(code="BCC0001", units="шт", qty="10")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0001", units="шт", qty="10")],
            google,
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(action.status, STATUS_IDENTITY)
        self.assertFalse(action.quantity_changed)
        self.assertFalse(action.units_changed)
        self.assertEqual(converted_quantity(action), Decimal("10"))
        self.assertEqual(len(plan.dependencies), 1)
        apply_conversion_plan(plan, [_binding("r1", row)])
        self.assertEqual(row.el[VALUES].value, "10")
        self.assertEqual(row.el[UNITS].value, "шт")

    def test_missing_google_single_unit_warning_no_change(self) -> None:
        google = build_google_units_index([])
        row = _data_row(code="BCC0009", units="шт", qty="3")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0009", units="шт", qty="3")],
            google,
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(action.status, STATUS_NO_GOOGLE)
        self.assertTrue(plan.warnings)
        self.assertIn("Нет единицы Google", plan.warnings[0])
        self.assertIn("количество не изменено", plan.warnings[0])
        self.assertEqual(plan.dependencies[0].target_unit, "шт")
        apply_conversion_plan(plan, [_binding("r1", row)])
        self.assertEqual(row.el[VALUES].value, "3")
        self.assertEqual(row.el[UNITS].value, "шт")

    def test_no_google_source_unit_cell_gets_name_comment(self) -> None:
        google = build_google_units_index([])
        build_conversion_plan(
            [_request("r1", code="BCC0031", units="шт", qty="1", name="Кабель силовой 1кВ")],
            google,
            self.matrix_path,
        )
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            self.assertEqual(ws["B2"].value, "Кабель силовой 1кВ")
            self.assertEqual(ws["E2"].value, "шт")
            comment = ws["E2"].comment
            self.assertIsNotNone(comment)
            self.assertIn("Кабель силовой 1кВ", comment.text)
        finally:
            wb.close()

    def test_google_code_source_unit_has_no_name_comment(self) -> None:
        google = build_google_units_index([_google_row("BCC0032", "шт")])
        build_conversion_plan(
            [_request("r1", code="BCC0032", units="шт", qty="1", name="Не должно попасть в комментарий")],
            google,
            self.matrix_path,
        )
        wb = load_workbook(self.matrix_path)
        try:
            self.assertIsNone(wb.active["E2"].comment)
        finally:
            wb.close()

    def test_missing_google_multi_unit_fatal(self) -> None:
        google = build_google_units_index([])
        requests = [
            _request("r1", code="BCC0009", units="шт", qty="1"),
            _request("r2", code="BCC0009", units="м", qty="2"),
        ]
        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(requests, google, self.matrix_path)
        self.assertIn("Несколько активных исходных ЕИ", str(ctx.exception))
        self.assertIn("Ед. изм. Google", str(ctx.exception))
        document = load_matrix(self.matrix_path)
        row = document.rows[0]
        self.assertEqual(row.code_status, "collision")
        self.assertEqual(row.google_normalized, "")
        self.assertEqual(row.google_name, "")
        self.assertEqual(
            {pair.source_normalized for pair in row.pairs},
            {"шт", "м"},
        )

    def test_collision_without_google_writes_both_names_to_column_b(self) -> None:
        google = build_google_units_index([])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [
                    _request("r1", code="BCC0033", units="компл", qty="1", name="Щит распределительный"),
                    _request("r2", code="BCC0033", units="шт", qty="1", name="Шкаф силовой"),
                ],
                google,
                self.matrix_path,
            )
        document = load_matrix(self.matrix_path)
        self.assertEqual(
            document.rows[0].google_name,
            "Щит распределительный\nШкаф силовой",
        )
        wb = load_workbook(self.matrix_path)
        try:
            self.assertEqual(
                wb.active["B2"].value,
                "Щит распределительный\nШкаф силовой",
            )
        finally:
            wb.close()

    def test_matrix_google_column_resolves_collision_without_google_base(self) -> None:
        google = build_google_units_index([])
        requests = [
            _request("r1", code="BCC0029", units="шт", qty="1"),
            _request("r2", code="BCC0029", units="м", qty="2"),
        ]
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(requests, google, self.matrix_path)
        self._set_matrix_google_unit(2, "шт")
        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(requests, google, self.matrix_path)
        self.assertIn("Требуется коэффициент «?»", str(ctx.exception))
        document = load_matrix(self.matrix_path)
        row = document.rows[0]
        self.assertEqual(row.google_normalized, "шт")
        self.assertEqual(row.code_status, "Требует уточнения")
        meter_pair = next(pair for pair in row.pairs if pair.source_normalized == "м")
        self.assertEqual(meter_pair.coefficient_raw, "?")
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            meter_col = None
            for col in range(5, ws.max_column + 1, 3):
                if str(ws.cell(row=2, column=col).value or "").strip() == "м":
                    meter_col = col
                    break
            self.assertIsNotNone(meter_col)
            ws.cell(row=2, column=meter_col + 1, value="10")
            wb.save(self.matrix_path)
        finally:
            wb.close()
        plan = build_conversion_plan(requests, google, self.matrix_path)
        piece = action_by_id(plan, "r1")
        meter = action_by_id(plan, "r2")
        assert piece is not None and meter is not None
        self.assertEqual(converted_unit(piece), "шт")
        self.assertEqual(converted_quantity(piece), Decimal("1"))
        self.assertEqual(converted_unit(meter), "шт")
        self.assertEqual(converted_quantity(meter), Decimal("20"))

    def test_google_base_overrides_matrix_d_local_target(self) -> None:
        google_empty = build_google_units_index([])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [
                    _request("r1", code="BCC0030", units="шт", qty="1"),
                    _request("r2", code="BCC0030", units="м", qty="1"),
                ],
                google_empty,
                self.matrix_path,
            )
        self._set_matrix_google_unit(2, "шт")
        google = build_google_units_index([_google_row("BCC0030", "кг")])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("r1", code="BCC0030", units="шт", qty="1")],
                google,
                self.matrix_path,
            )
        document = load_matrix(self.matrix_path)
        self.assertEqual(document.rows[0].google_normalized, "кг")

    def test_required_question_mark_writes_yellow_and_fatal(self) -> None:
        google = build_google_units_index([_google_row("BCC0002", "шт")])
        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(
                [_request("r1", code="BCC0002", units="м", qty="4")],
                google,
                self.matrix_path,
            )
        message = str(ctx.exception)
        self.assertIn("Требуется коэффициент «?»", message)
        self.assertNotIn("Required coefficient", message)
        self.assertIn("Что нужно сделать:", message)
        self.assertIn(str(self.matrix_path), message)
        self.assertIn("количество в ЕИ Google = исходное количество × коэффициент", message)
        self.assertIn("Сохраните и закройте файл Excel", message)
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            self.assertEqual(ws["D1"].value, "Ед. изм. Google")
            self.assertEqual(ws["E1"].value, "Найденная ЕИ 1")
            self.assertEqual(ws["F1"].value, "Коэффициент 1")
            self.assertEqual(ws["G1"].value, f"{PAIR_FORMULA_PREFIX}1")
            question_cell = self._find_question_mark_coef_cell()
            cell = ws[question_cell]
            self.assertEqual(cell.value, "?")
            rgb = str(cell.fill.fgColor.rgb or "")
            self.assertTrue(rgb.endswith("FFFF00"))
            formula_cell = ws["G2"]
            self.assertTrue(str(formula_cell.value or "").startswith("="))
            self.assertIn("F2", str(formula_cell.value))
            self.assertIn("$D2", str(formula_cell.value))
            self.assertIn("E2", str(formula_cell.value))
            formula_rgb = str(formula_cell.fill.fgColor.rgb or "")
            self.assertTrue(formula_rgb.endswith(self._formula_gray_rgb_suffix()))
            header_formula_rgb = str(ws["G1"].fill.fgColor.rgb or "")
            self.assertTrue(header_formula_rgb.endswith(self._formula_gray_rgb_suffix()))
            self.assertEqual(ws.auto_filter.ref, "A1:J2")
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws["C2"].value, "Требует уточнения")
            self.assertTrue(str(ws["C2"].fill.fgColor.rgb or "").endswith("FFFF00"))
        finally:
            wb.close()

    def test_blank_active_coefficient_is_rewritten_as_yellow_question_mark(self) -> None:
        path = Path(self._tmpdir.name) / "blank_active_coefficient.xlsx"
        self._write_current_layout_matrix(
            path,
            rows=[("BCC0003", "шт", [("шт", "1"), ("м", "")])],
        )
        google = build_google_units_index([_google_row("BCC0003", "шт")])

        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(
                [_request("r1", code="BCC0003", units="м", qty="4")],
                google,
                path,
            )

        self.assertIn("Требуется коэффициент «?»", str(ctx.exception))
        document = load_matrix(path)
        meter_pair = next(
            pair for pair in document.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair.coefficient_raw, "?")
        wb = load_workbook(path)
        try:
            ws = wb.active
            self.assertEqual(ws["I2"].value, "?")
            self.assertTrue(str(ws["I2"].fill.fgColor.rgb or "").endswith("FFFF00"))
            self.assertEqual(ws["C2"].value, "Требует уточнения")
        finally:
            wb.close()

    def test_google_name_written_to_matrix(self) -> None:
        google = build_google_units_index(
            [_google_row("BCC0027", "шт", name="Кабель силовой")]
        )
        build_conversion_plan(
            [_request("r1", code="BCC0027", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            self.assertEqual(ws["B1"].value, "Наименование Google")
            self.assertEqual(ws["B2"].value, "Кабель силовой")
            self.assertIn(ws["C2"].value, ("", None))
            self.assertEqual(ws["D2"].value, "шт")
        finally:
            wb.close()

    def test_inactive_historical_question_mark_not_fatal(self) -> None:
        self._seed_nonidentity_placeholder("BCC0003", "шт", "м")
        plan = build_conversion_plan(
            [_request("r2", code="BCC0003", units="шт", qty="5")],
            build_google_units_index([_google_row("BCC0003", "шт")]),
            self.matrix_path,
        )
        action = action_by_id(plan, "r2")
        assert action is not None
        self.assertEqual(action.status, STATUS_IDENTITY)
        self.assertEqual(load_matrix(self.matrix_path).rows[0].code_status, "Требует уточнения")

    def test_filled_coefficient_clears_needs_clarification_status(self) -> None:
        self._seed_nonidentity_placeholder("BCC0028", "шт", "м")
        self.assertEqual(load_matrix(self.matrix_path).rows[0].code_status, "Требует уточнения")
        self._set_matrix_coef(2, 2, "10")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0028", units="м", qty="1")],
            build_google_units_index([_google_row("BCC0028", "шт")]),
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(converted_quantity(action), Decimal("10"))
        self.assertEqual(load_matrix(self.matrix_path).rows[0].code_status, "")

    def test_valid_integer_conversion(self) -> None:
        self._seed_nonidentity_placeholder("BCC0004", "шт", "м")
        self._set_matrix_coef(2, 2, 100)
        row = _data_row(code="BCC0004", units="м", qty="10")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0004", units="м", qty="10")],
            build_google_units_index([_google_row("BCC0004", "шт")]),
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(converted_quantity(action), Decimal("1000"))
        apply_conversion_plan(plan, [_binding("r1", row)])
        self.assertIsInstance(row.el[VALUES].value, int)
        self.assertEqual(row.el[VALUES].value, 1000)
        self.assertEqual(converted_unit(action), "шт")

    def test_machine_residual_adjustment(self) -> None:
        self._seed_nonidentity_placeholder("BCC0005", "шт", "м")
        self._set_matrix_coef(2, 2, "3.0000000000001")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0005", units="м", qty="1")],
            build_google_units_index([_google_row("BCC0005", "шт")]),
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertTrue(action.machine_residual_adjusted)
        self.assertEqual(converted_quantity(action), Decimal("3"))

    def test_real_fractional_result_is_logged_and_applied(self) -> None:
        self._seed_nonidentity_placeholder("BCC0006", "шт", "м")
        self._set_matrix_coef(2, 2, "3.5")
        row = _data_row(code="BCC0006", units="м", qty="1")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0006", units="м", qty="1")],
            build_google_units_index([_google_row("BCC0006", "шт")]),
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertTrue(action.fractional_result)
        self.assertEqual(converted_quantity(action), Decimal("3.5"))
        self.assertEqual(len(plan.fractional_issues), 1)
        self.assertIn("дробный результат", "".join(plan.warnings).lower())
        apply_conversion_plan(plan, [_binding("r1", row)])
        self.assertEqual(row.el[VALUES].value, "3.5")
        self.assertEqual(row.el[UNITS].value, "шт")

    def test_fractional_result_skips_tags_equal_and_writes_xlsx(self) -> None:
        self._seed_nonidentity_placeholder("BCC0006B", "шт", "м")
        self._set_matrix_coef(2, 2, "3.5")
        plan = build_conversion_plan(
            [
                _request(
                    "r1",
                    code="BCC0006B",
                    units="м",
                    qty="1",
                    tags=3,
                    invariant=ConversionInvariant.TAGS_EQUAL,
                    contour="rfp",
                )
            ],
            build_google_units_index([_google_row("BCC0006B", "шт")]),
            self.matrix_path,
        )
        self.assertEqual(len(plan.fractional_issues), 1)
        self.assertTrue(plan.fractional_issues[0].invariant_skipped)
        log_dir = Path(self._tmpdir.name) / "frac_logs"
        from RFQ.units_convert import write_fractional_conversion_logs

        written = write_fractional_conversion_logs(plan, log_dir)
        self.assertEqual(len(written), 1)
        self.assertEqual(written[0].name, "дробные значения после конвертации RFP.xlsx")
        wb = load_workbook(written[0])
        try:
            ws = wb.active
            self.assertEqual(ws["C2"].value, "BCC0006B")
            self.assertEqual(str(ws["H2"].value), "3.5")
            self.assertEqual(ws["L2"].value, "да")
        finally:
            wb.close()

    def test_coef_one_fractional_quantity_preserved(self) -> None:
        google = build_google_units_index([_google_row("BCC0007", "шт")])
        row = _data_row(code="BCC0007", units="шт.", qty="2,5")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0007", units="шт.", qty="2,5")],
            google,
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(converted_quantity(action), Decimal("2.5"))
        apply_conversion_plan(plan, [_binding("r1", row)])
        self.assertEqual(row.el[VALUES].value, "2,5")
        self.assertEqual(row.el[UNITS].value, "шт")

    def test_identity_canonical_text_change_enforces_tags_invariant(self) -> None:
        google = build_google_units_index([_google_row("BCC0015", "шт")])
        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(
                [
                    _request(
                        "r1",
                        code="BCC0015",
                        units="шт.",
                        qty="2",
                        tags=3,
                        invariant=ConversionInvariant.TAGS_EQUAL,
                    )
                ],
                google,
                self.matrix_path,
            )
        message = str(ctx.exception)
        self.assertIn("TAGS_EQUAL", message)
        self.assertIn("Нарушен инвариант TAGS_EQUAL", message)

    def test_tags_equal_invariant(self) -> None:
        self._seed_nonidentity_placeholder("BCC0008", "шт", "м")
        self._set_matrix_coef(2, 2, "2")
        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(
                [
                    _request(
                        "r1",
                        code="BCC0008",
                        units="м",
                        qty="5",
                        tags=4,
                        invariant=ConversionInvariant.TAGS_EQUAL,
                    )
                ],
                build_google_units_index([_google_row("BCC0008", "шт")]),
                self.matrix_path,
            )
        message = str(ctx.exception)
        self.assertIn("TAGS_EQUAL", message)
        self.assertIn("Нарушен инвариант TAGS_EQUAL", message)

    def test_tags_not_exceed_quantity_invariant(self) -> None:
        self._seed_nonidentity_placeholder("BCC0010", "шт", "м")
        self._set_matrix_coef(2, 2, "2")
        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(
                [
                    _request(
                        "r1",
                        code="BCC0010",
                        units="м",
                        qty="2.5",
                        tags=6,
                        invariant=ConversionInvariant.TAGS_NOT_EXCEED_QUANTITY,
                    )
                ],
                build_google_units_index([_google_row("BCC0010", "шт")]),
                self.matrix_path,
            )
        message = str(ctx.exception)
        self.assertIn("TAGS_NOT_EXCEED_QUANTITY", message)
        self.assertIn("Нарушен инвариант TAGS_NOT_EXCEED_QUANTITY", message)

    def test_empty_code_or_source_unit_fatal(self) -> None:
        google = build_google_units_index([_google_row("BCC0016", "шт")])
        with self.assertRaises(UnitsConversionError) as ctx:
            build_conversion_plan(
                [_request("r1", code="   ", units="шт", qty="1", location="loc-empty-code")],
                google,
                self.matrix_path,
            )
        self.assertIn("Пустой нормализованный код", str(ctx.exception))
        self.assertIn("loc-empty-code", str(ctx.exception))
        with self.assertRaises(UnitsConversionError) as ctx2:
            build_conversion_plan(
                [_request("r2", code="BCC0016", units="  ", qty="1", location="loc-empty-unit")],
                google,
                self.matrix_path,
            )
        self.assertIn("Пустая нормализованная исходная ЕИ", str(ctx2.exception))
        self.assertIn("loc-empty-unit", str(ctx2.exception))

    def test_no_partial_mutation_on_binding_validation_error(self) -> None:
        google = build_google_units_index([_google_row("BCC0011", "шт")])
        row_ok = _data_row(code="BCC0011", units="шт", qty="1", status="keep", trace="keep")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0011", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        before_ok = copy.deepcopy(
            {
                VALUES: row_ok.el[VALUES].value,
                UNITS: row_ok.el[UNITS].value,
                UNITS_CHECK_STATUS: row_ok.el[UNITS_CHECK_STATUS].value,
                UNITS_CONVERSION_TRACE: row_ok.el[UNITS_CONVERSION_TRACE].value,
            }
        )
        with self.assertRaises(UnitsConversionError):
            apply_conversion_plan(plan, [])
        self.assertEqual(row_ok.el[VALUES].value, before_ok[VALUES])
        self.assertEqual(row_ok.el[UNITS].value, before_ok[UNITS])
        self.assertEqual(row_ok.el[UNITS_CHECK_STATUS].value, before_ok[UNITS_CHECK_STATUS])
        self.assertEqual(row_ok.el[UNITS_CONVERSION_TRACE].value, before_ok[UNITS_CONVERSION_TRACE])

    def test_apply_merges_existing_status_and_trace(self) -> None:
        google = build_google_units_index([_google_row("BCC0012", "шт")])
        row = _data_row(code="BCC0012", units="шт", qty="1", status="rfp-old", trace="mto-old")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0012", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        apply_conversion_plan(plan, [_binding("r1", row)])
        self.assertIn("rfp-old", str(row.el[UNITS_CHECK_STATUS].value))
        self.assertIn("identity", str(row.el[UNITS_CHECK_STATUS].value))
        self.assertIn("mto-old", str(row.el[UNITS_CONVERSION_TRACE].value))

    def test_apply_hydrates_missing_status_trace_columns(self) -> None:
        google = build_google_units_index([_google_row("BCC0017", "шт")])
        row = _data_row(code="BCC0017", units="шт.", qty="1", slim=True)
        self.assertNotIn(UNITS_CHECK_STATUS, row.el)
        plan = build_conversion_plan(
            [_request("r1", code="BCC0017", units="шт.", qty="1")],
            google,
            self.matrix_path,
        )
        apply_conversion_plan(plan, [_binding("r1", row)])
        self.assertIn(UNITS_CHECK_STATUS, row.el)
        self.assertIn(UNITS_CONVERSION_TRACE, row.el)
        self.assertIn("converted", str(row.el[UNITS_CHECK_STATUS].value))

    def test_multiple_bindings_same_row_allowed(self) -> None:
        google = build_google_units_index([_google_row("BCC0018", "шт")])
        row = _data_row(code="BCC0018", units="м", qty="10")
        row.el[VALUES_2].value = "10"
        self._seed_nonidentity_placeholder("BCC0018", "шт", "м")
        self._set_matrix_coef(2, 2, 2)
        plan = build_conversion_plan(
            [
                _request("r1", code="BCC0018", units="м", qty="10"),
                _request("r2", code="BCC0018", units="м", qty="10"),
            ],
            google,
            self.matrix_path,
        )
        apply_conversion_plan(
            plan,
            [
                _binding("r1", row),
                RowStdBinding(
                    request_id="r2",
                    row=row,
                    quantity_column=VALUES_2,
                    units_column=UNITS,
                    status_column=UNITS_CHECK_STATUS,
                    trace_column=UNITS_CONVERSION_TRACE,
                ),
            ],
        )
        self.assertEqual(row.el[VALUES].value, 20)
        self.assertEqual(row.el[VALUES_2].value, 20)

    def test_binding_rejects_duplicate_request_id(self) -> None:
        google = build_google_units_index([_google_row("BCC0019", "шт")])
        row = _data_row(code="BCC0019", units="шт", qty="1")
        plan = build_conversion_plan(
            [_request("r1", code="BCC0019", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        with self.assertRaises(UnitsConversionError):
            apply_conversion_plan(plan, [_binding("r1", row), _binding("r1", row)])

    def test_absent_code_later_appears_in_google_reorders_and_resets(self) -> None:
        google_empty = build_google_units_index([])
        build_conversion_plan(
            [_request("r1", code="BCC0013", units="шт", qty="1")],
            google_empty,
            self.matrix_path,
        )
        doc_before = load_matrix(self.matrix_path)
        self.assertEqual(doc_before.rows[0].code_status, "Нет в базе")

        google = build_google_units_index([_google_row("BCC0013", "м")])
        build_conversion_plan([], google, self.matrix_path)
        doc_after = load_matrix(self.matrix_path)
        self.assertEqual(doc_after.rows[0].google_normalized, "м")
        self.assertEqual(doc_after.rows[0].pairs[0].source_normalized, "м")
        self.assertEqual(doc_after.rows[0].pairs[0].coefficient_raw, "1")
        inactive_pair = next(
            pair for pair in doc_after.rows[0].pairs if pair.source_normalized == "шт"
        )
        self.assertEqual(inactive_pair.coefficient_raw, "?")
        self.assertEqual(doc_after.rows[0].code_status, "Требует уточнения")

    def test_google_target_change_resets_nonidentity_and_retains_old_target(self) -> None:
        self._seed_nonidentity_placeholder("BCC0014", "шт", "м")
        self._set_matrix_coef(2, 2, "10")
        google_new = build_google_units_index([_google_row("BCC0014", "кг")])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("r1", code="BCC0014", units="м", qty="1")],
                google_new,
                self.matrix_path,
            )
        doc = load_matrix(self.matrix_path)
        row = doc.rows[0]
        self.assertEqual(row.google_normalized, "кг")
        kg_pair = next(pair for pair in row.pairs if pair.source_normalized == "кг")
        self.assertEqual(kg_pair.coefficient_raw, "1")
        self.assertEqual(row.pairs[0].source_normalized, "кг")
        meter_pair = next(pair for pair in row.pairs if pair.source_normalized == "м")
        self.assertEqual(meter_pair.coefficient_raw, "?")
        old_target_pair = next(pair for pair in row.pairs if pair.source_normalized == "шт")
        self.assertEqual(old_target_pair.coefficient_raw, "?")

    def test_unrelated_google_code_not_added_to_matrix(self) -> None:
        google = build_google_units_index(
            [
                _google_row("BCC0001", "шт"),
                _google_row("BCC9999", "кг"),
            ]
        )
        build_conversion_plan(
            [_request("r1", code="BCC0001", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        doc = load_matrix(self.matrix_path)
        codes = {row.code_normalized for row in doc.rows}
        self.assertEqual(codes, {"BCC0001"})

    def test_dependencies_include_identity_and_no_google(self) -> None:
        google = build_google_units_index([_google_row("BCC0021", "шт")])
        plan_google = build_conversion_plan(
            [_request("g", code="BCC0021", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        self.assertEqual(
            plan_google.dependencies,
            (
                ConversionDependency(
                    code="BCC0021",
                    source_unit="шт",
                    target_unit="шт",
                    coefficient=Decimal("1"),
                ),
            ),
        )
        plan_no_google = build_conversion_plan(
            [_request("n", code="BCC0022", units="м", qty="2")],
            build_google_units_index([]),
            Path(self._tmpdir.name) / "units_matrix_ng.xlsx",
        )
        dep = plan_no_google.dependencies[0]
        self.assertEqual(dep.code, "BCC0022")
        self.assertEqual(dep.source_unit, "м")
        self.assertEqual(dep.target_unit, "м")
        self.assertEqual(dep.coefficient, Decimal("1"))

    def test_duplicate_matrix_code_rows_fatal(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Код"
        ws["B1"] = "Статус кода"
        ws["C1"] = "Ед. изм. Google"
        ws["D1"] = "Найденная ЕИ 1"
        ws["E1"] = "Коэффициент 1"
        ws["A2"] = "BCC0030"
        ws["A3"] = "bcc0030"
        path = Path(self._tmpdir.name) / "dup_matrix.xlsx"
        wb.save(path)
        wb.close()
        with self.assertRaises(UnitsConversionError) as ctx:
            load_matrix(path)
        self.assertIn("Дублирующийся нормализованный код матрицы", str(ctx.exception))

    def test_concurrent_destination_change_retries_without_blind_overwrite(self) -> None:
        google = build_google_units_index([_google_row("BCC0020", "шт")])
        self._seed_identity_matrix("BCC0020", "шт")
        hook_calls = {"count": 0}

        def _touch_destination(_tmp: Path, dest: Path) -> None:
            if hook_calls["count"]:
                return
            hook_calls["count"] += 1
            os.utime(dest, (time.time() + 5, time.time() + 5))

        matrix_module._before_replace_hook = _touch_destination
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("r1", code="BCC0020", units="м", qty="1")],
                google,
                self.matrix_path,
            )
        self.assertEqual(hook_calls["count"], 1)
        doc = load_matrix(self.matrix_path)
        meter_pair = next(
            (pair for pair in doc.rows[0].pairs if pair.source_normalized == "м"),
            None,
        )
        self.assertIsNotNone(meter_pair)
        self.assertEqual(meter_pair.coefficient_raw, "?")

    def test_permission_error_on_replace_is_explicit_units_error(self) -> None:
        google = build_google_units_index([_google_row("BCC0021", "шт")])
        self._seed_identity_matrix("BCC0021", "шт")
        original = self.matrix_path.read_bytes()

        def _busy_replace(_src: str | os.PathLike[str], _dst: str | os.PathLike[str]) -> None:
            raise PermissionError(13, "Permission denied")

        with patch("RFQ.units_convert.matrix.os.replace", side_effect=_busy_replace):
            with patch("RFQ.units_convert.matrix.time.sleep"):
                with self.assertRaises(UnitsConversionError) as ctx:
                    build_conversion_plan(
                        [_request("r1", code="BCC0021", units="м", qty="1")],
                        google,
                        self.matrix_path,
                    )
        message = str(ctx.exception)
        self.assertIn("файл занят", message)
        self.assertIn(self.matrix_path.name, message)
        self.assertIn("Закройте файл матрицы", message)
        self.assertIn("Повторите сбор частей или запуск RFP", message)
        self.assertIsInstance(ctx.exception.__cause__, PermissionError)
        self.assertEqual(self.matrix_path.read_bytes(), original)

    def test_permission_error_on_load_is_explicit_units_error(self) -> None:
        self.matrix_path.write_bytes(b"locked-matrix")

        def _busy_read(self_path: Path) -> bytes:
            raise PermissionError(13, "Permission denied")

        with patch.object(Path, "read_bytes", _busy_read):
            with self.assertRaises(UnitsConversionError) as ctx:
                load_matrix(self.matrix_path)
        message = str(ctx.exception)
        self.assertIn("файл занят", message)
        self.assertIn("прочитать", message)
        self.assertIn(self.matrix_path.name, message)
        self.assertIsInstance(ctx.exception.__cause__, PermissionError)

    def test_matrix_rows_sorted_by_normalized_code(self) -> None:
        google = build_google_units_index(
            [
                _google_row("BCC0002", "шт"),
                _google_row("BCC0001", "шт"),
            ]
        )
        build_conversion_plan(
            [
                _request("r1", code="BCC0002", units="шт", qty="1"),
                _request("r2", code="BCC0001", units="шт", qty="1"),
            ],
            google,
            self.matrix_path,
        )
        doc = load_matrix(self.matrix_path)
        codes = [row.code_normalized for row in doc.rows]
        self.assertEqual(codes, sorted(codes))

    def test_two_pair_triple_layout_autofilter_span(self) -> None:
        google = build_google_units_index([_google_row("BCC0100", "шт")])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("r1", code="BCC0100", units="м", qty="1")],
                google,
                self.matrix_path,
            )
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            self.assertEqual(ws["H1"].value, "Найденная ЕИ 2")
            self.assertEqual(ws["I1"].value, "Коэффициент 2")
            self.assertEqual(ws["J1"].value, f"{PAIR_FORMULA_PREFIX}2")
            self.assertEqual(ws.auto_filter.ref, "A1:J2")
        finally:
            wb.close()

    def test_legacy_two_column_matrix_migrates_to_triple_layout(self) -> None:
        legacy_path = Path(self._tmpdir.name) / "legacy_matrix.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Код"
        ws["B1"] = "Статус кода"
        ws["C1"] = "Ед. изм. Google"
        ws["D1"] = "Найденная ЕИ 1"
        ws["E1"] = "Коэффициент 1"
        ws["A2"] = "BCC0101"
        ws["C2"] = "шт"
        ws["D2"] = "упак"
        ws["E2"] = "100"
        wb.save(legacy_path)
        wb.close()

        google = build_google_units_index(
            [_google_row("BCC0101", "шт", name="Упаковка товара")]
        )
        plan = build_conversion_plan(
            [_request("r1", code="BCC0101", units="упак", qty="2")],
            google,
            legacy_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(converted_quantity(action), Decimal("200"))

        doc = load_matrix(legacy_path)
        self.assertTrue(doc.layout_current)
        pack_pair = next(
            pair for pair in doc.rows[0].pairs if pair.source_normalized == "упак"
        )
        self.assertEqual(pack_pair.coefficient_raw, "100")
        self.assertEqual(pack_pair.coefficient, Decimal("100"))

        wb = load_workbook(legacy_path)
        try:
            ws = wb.active
            self.assertEqual(ws["B2"].value, "Упаковка товара")
            self.assertEqual(ws["G1"].value, f"{PAIR_FORMULA_PREFIX}1")
            self.assertTrue(str(ws["G2"].value or "").startswith("="))
            self.assertIn("$D2", str(ws["G2"].value))
            self.assertEqual(ws["I2"].value, "100")
            self.assertEqual(ws.auto_filter.ref, "A1:J2")
        finally:
            wb.close()

    def test_legacy_multi_pair_two_column_matrix_migrates(self) -> None:
        legacy_path = Path(self._tmpdir.name) / "legacy_multi_matrix.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Код"
        ws["B1"] = "Статус кода"
        ws["C1"] = "Ед. изм. Google"
        ws["D1"] = "Найденная ЕИ 1"
        ws["E1"] = "Коэффициент 1"
        ws["F1"] = "Найденная ЕИ 2"
        ws["G1"] = "Коэффициент 2"
        ws["A2"] = "BCC0102"
        ws["C2"] = "шт"
        ws["D2"] = "шт"
        ws["E2"] = "1"
        ws["F2"] = "м"
        ws["G2"] = "10"
        wb.save(legacy_path)
        wb.close()

        doc_before = load_matrix(legacy_path)
        self.assertFalse(doc_before.layout_current)
        meter_pair = next(
            pair for pair in doc_before.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair.coefficient_raw, "10")

        google = build_google_units_index([_google_row("BCC0102", "шт")])
        build_conversion_plan(
            [_request("r1", code="BCC0102", units="м", qty="1")],
            google,
            legacy_path,
        )

        doc_after = load_matrix(legacy_path)
        self.assertTrue(doc_after.layout_current)
        meter_pair_after = next(
            pair for pair in doc_after.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair_after.coefficient_raw, "10")
        self.assertEqual(meter_pair_after.coefficient, Decimal("10"))

        wb = load_workbook(legacy_path)
        try:
            ws = wb.active
            self.assertEqual(ws["H1"].value, "Найденная ЕИ 2")
            self.assertEqual(ws["I1"].value, "Коэффициент 2")
            self.assertEqual(ws["J1"].value, f"{PAIR_FORMULA_PREFIX}2")
            self.assertEqual(ws.auto_filter.ref, "A1:J2")
        finally:
            wb.close()

    def test_legacy_triple_matrix_migrates_with_name_and_coef(self) -> None:
        legacy_path = Path(self._tmpdir.name) / "legacy_triple_matrix.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Код"
        ws["B1"] = "Статус кода"
        ws["C1"] = "Ед. изм. Google"
        ws["D1"] = "Найденная ЕИ 1"
        ws["E1"] = "Коэффициент 1"
        ws["F1"] = f"{PAIR_FORMULA_PREFIX}1"
        ws["A2"] = "BCC0104"
        ws["C2"] = "шт"
        ws["D2"] = "упак"
        ws["E2"] = "50"
        ws["F2"] = '=IF(OR(D2="",E2="",E2="?",$C2=""),"",E2&"/"&$C2&" = 1/"&D2)'
        wb.save(legacy_path)
        wb.close()

        doc_before = load_matrix(legacy_path)
        self.assertFalse(doc_before.layout_current)
        self.assertEqual(doc_before.rows[0].pairs[0].coefficient_raw, "50")

        google = build_google_units_index(
            [_google_row("BCC0104", "шт", name="Legacy triple item")]
        )
        build_conversion_plan(
            [_request("r1", code="BCC0104", units="упак", qty="2")],
            google,
            legacy_path,
        )

        doc_after = load_matrix(legacy_path)
        self.assertTrue(doc_after.layout_current)
        self.assertEqual(doc_after.rows[0].google_name, "Legacy triple item")
        pack_pair = next(
            pair for pair in doc_after.rows[0].pairs if pair.source_normalized == "упак"
        )
        self.assertEqual(pack_pair.coefficient_raw, "50")
        self.assertEqual(pack_pair.coefficient, Decimal("50"))

        wb = load_workbook(legacy_path)
        try:
            ws = wb.active
            self.assertEqual(ws["B2"].value, "Legacy triple item")
            self.assertEqual(_coef_column(ws, 2), 9)
            self.assertEqual(ws.cell(row=2, column=_coef_column(ws, 2)).value, "50")
            self.assertIn("$D2", str(ws["G2"].value))
        finally:
            wb.close()

    def test_new_layout_multi_pair_load_ignores_formula(self) -> None:
        triple_path = Path(self._tmpdir.name) / "triple_matrix.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Код"
        ws["B1"] = "Наименование Google"
        ws["C1"] = "Статус кода"
        ws["D1"] = "Ед. изм. Google"
        ws["E1"] = "Найденная ЕИ 1"
        ws["F1"] = "Коэффициент 1"
        ws["G1"] = f"{PAIR_FORMULA_PREFIX}1"
        ws["H1"] = "Найденная ЕИ 2"
        ws["I1"] = "Коэффициент 2"
        ws["J1"] = f"{PAIR_FORMULA_PREFIX}2"
        ws["A2"] = "BCC0103"
        ws["B2"] = "Item name"
        ws["D2"] = "шт"
        ws["E2"] = "шт"
        ws["F2"] = "1"
        ws["G2"] = '=IF(OR(E2="",F2="",F2="?",$D2=""),"",F2&"/"&$D2&" = 1/"&E2)'
        ws["H2"] = "м"
        ws["I2"] = "5"
        ws["J2"] = '=IF(OR(H2="",I2="",I2="?",$D2=""),"",I2&"/"&$D2&" = 1/"&H2)'
        for coord in ("F2", "I2"):
            ws[coord].number_format = "@"
        _apply_persisted_matrix_formatting(ws, pair_count=2, last_row=2)
        wb.save(triple_path)
        wb.close()

        doc = load_matrix(triple_path)
        self.assertTrue(doc.layout_current)
        self.assertEqual(len(doc.rows[0].pairs), 2)
        meter_pair = next(
            pair for pair in doc.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair.coefficient_raw, "5")
        self.assertEqual(meter_pair.coefficient, Decimal("5"))

        google = build_google_units_index([_google_row("BCC0103", "шт")])
        plan = build_conversion_plan(
            [_request("r1", code="BCC0103", units="м", qty="2")],
            google,
            triple_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(converted_quantity(action), Decimal("10"))

    def _write_current_layout_matrix(
        self,
        path: Path,
        *,
        rows: list[tuple[str, str, list[tuple[str, object]]]],
        coef_text_format: bool = True,
    ) -> None:
        """Write a current-layout matrix; ``rows`` are (code, google_unit, [(source, coef), ...])."""
        wb = Workbook()
        ws = wb.active
        for col_index, title in enumerate(MATRIX_HEADER, start=1):
            ws.cell(row=1, column=col_index, value=title)
        pair_count = max((len(pairs) for _, _, pairs in rows), default=1)
        for pair_index in range(1, pair_count + 1):
            source_col = PAIR_START_COL + (pair_index - 1) * PAIR_COLUMN_STRIDE
            ws.cell(row=1, column=source_col, value=f"Найденная ЕИ {pair_index}")
            ws.cell(row=1, column=source_col + 1, value=f"{PAIR_COEF_PREFIX}{pair_index}")
            ws.cell(
                row=1,
                column=source_col + 2,
                value=f"{PAIR_FORMULA_PREFIX}{pair_index}",
            )
        for row_index, (code, google_unit, pairs) in enumerate(rows, start=2):
            ws.cell(row=row_index, column=1, value=code)
            ws.cell(row=row_index, column=4, value=google_unit)
            for pair_index, (source, coef) in enumerate(pairs, start=1):
                source_col = PAIR_START_COL + (pair_index - 1) * PAIR_COLUMN_STRIDE
                coef_col = source_col + 1
                ws.cell(row=row_index, column=source_col, value=source)
                coef_cell = ws.cell(row=row_index, column=coef_col, value=coef)
                if coef_text_format:
                    coef_cell.number_format = "@"
            for pair_index in range(len(pairs) + 1, pair_count + 1):
                coef_col = PAIR_START_COL + (pair_index - 1) * PAIR_COLUMN_STRIDE + 1
                blank_coef_cell = ws.cell(row=row_index, column=coef_col)
                if coef_text_format:
                    blank_coef_cell.number_format = "@"
        last_row = len(rows) + 1
        _apply_persisted_matrix_formatting(ws, pair_count=pair_count, last_row=last_row)
        wb.save(path)
        wb.close()

    def test_non_text_coef_triggers_text_format_migration(self) -> None:
        path = Path(self._tmpdir.name) / "non_text_coef.xlsx"
        self._write_current_layout_matrix(
            path,
            rows=[("BCC0400", "шт", [("шт", "1"), ("м", "100")])],
            coef_text_format=False,
        )
        doc_before = load_matrix(path)
        self.assertFalse(doc_before.layout_current)

        google = build_google_units_index([_google_row("BCC0400", "шт")])
        plan = build_conversion_plan(
            [_request("r1", code="BCC0400", units="м", qty="2")],
            google,
            path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(converted_quantity(action), Decimal("200"))

        doc_after = load_matrix(path)
        self.assertTrue(doc_after.layout_current)
        wb = load_workbook(path)
        try:
            ws = wb.active
            for pair_index in (1, 2):
                coef_col = _coef_column(ws, pair_index)
                for row_idx in range(2, ws.max_row + 1):
                    self.assertEqual(
                        ws.cell(row=row_idx, column=coef_col).number_format,
                        "@",
                    )
        finally:
            wb.close()

    def test_text_formatted_current_matrix_skips_rewrite(self) -> None:
        google = build_google_units_index([_google_row("BCC0401", "шт")])
        requests = [_request("r1", code="BCC0401", units="шт", qty="1")]
        build_conversion_plan(requests, google, self.matrix_path)
        self.assertTrue(load_matrix(self.matrix_path).layout_current)

        hook_calls: list[tuple[Path, Path]] = []

        def _record_hook(tmp_path: Path, dest_path: Path) -> None:
            hook_calls.append((tmp_path, dest_path))

        matrix_module._before_replace_hook = _record_hook
        build_conversion_plan(requests, google, self.matrix_path)
        self.assertEqual(hook_calls, [])

    def test_multi_row_multi_pair_sequential_load_values(self) -> None:
        path = Path(self._tmpdir.name) / "multi_row_matrix.xlsx"
        self._write_current_layout_matrix(
            path,
            rows=[
                ("BCC0402", "шт", [("шт", "1"), ("м", "10"), ("упак", "50")]),
                ("BCC0403", "кг", [("кг", "1"), ("т", "1000")]),
            ],
        )
        doc = load_matrix(path)
        self.assertTrue(doc.layout_current)
        self.assertEqual([row.code_normalized for row in doc.rows], ["BCC0402", "BCC0403"])
        row_a = doc.rows[0]
        pack_pair = next(pair for pair in row_a.pairs if pair.source_normalized == "упак")
        self.assertEqual(pack_pair.coefficient_raw, "50")
        self.assertEqual(pack_pair.coefficient, Decimal("50"))
        row_b = doc.rows[1]
        ton_pair = next(pair for pair in row_b.pairs if pair.source_normalized == "т")
        self.assertEqual(ton_pair.coefficient_raw, "1000")
        self.assertEqual(ton_pair.coefficient, Decimal("1000"))

    def test_save_hook_sees_temp_before_replace(self) -> None:
        google = build_google_units_index([_google_row("BCC0404", "шт")])
        observed: dict[str, object] = {}

        def _inspect_temp(tmp_path: Path, dest_path: Path) -> None:
            observed["tmp_exists"] = tmp_path.exists()
            observed["tmp_size"] = tmp_path.stat().st_size if tmp_path.exists() else 0
            observed["dest"] = dest_path

        matrix_module._before_replace_hook = _inspect_temp
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("r1", code="BCC0404", units="м", qty="1")],
                google,
                self.matrix_path,
            )
        self.assertTrue(observed.get("tmp_exists"))
        self.assertGreater(int(observed.get("tmp_size", 0)), 0)
        self.assertEqual(observed.get("dest"), self.matrix_path)

    def test_build_conversion_plan_uses_indexed_lookup(self) -> None:
        path = Path(self._tmpdir.name) / "indexed_lookup_matrix.xlsx"
        rows_data = [
            (f"BCC{5000 + index:04d}", "шт", [("шт", "1"), ("м", str(index + 1))])
            for index in range(120)
        ]
        self._write_current_layout_matrix(path, rows=rows_data)
        google_rows = [_google_row(code, "шт") for code, _, _ in rows_data]
        google = build_google_units_index(google_rows)
        requests = [
            _request(
                f"r{index}",
                code=code,
                units="м",
                qty="1",
            )
            for index, (code, _, pairs) in enumerate(rows_data)
        ]

        def _forbidden_linear_lookup(*_args, **_kwargs):
            raise AssertionError("linear matrix lookup must not run in build_conversion_plan")

        with patch(
            "RFQ.units_convert.matrix.lookup_coefficient",
            side_effect=_forbidden_linear_lookup,
        ), patch(
            "RFQ.units_convert.matrix.matrix_row_for_code",
            side_effect=_forbidden_linear_lookup,
        ):
            plan = build_conversion_plan(requests, google, path)

        self.assertEqual(len(plan.actions), len(requests))
        sample = action_by_id(plan, "r0")
        assert sample is not None
        self.assertEqual(converted_quantity(sample), Decimal("1"))

    def test_load_matrix_benchmark_two_thousand_rows(self) -> None:
        path = Path(self._tmpdir.name) / "bench_matrix.xlsx"
        rows_data = [
            (f"BCC{6000 + index:04d}", "шт", [("шт", "1"), ("м", "2")])
            for index in range(2000)
        ]
        self._write_current_layout_matrix(path, rows=rows_data)
        started = time.perf_counter()
        doc = load_matrix(path)
        elapsed = time.perf_counter() - started
        self.assertEqual(len(doc.rows), 2000)
        print(f"[matrix benchmark] load_matrix 2000 rows: {elapsed:.3f}s")

    def test_persisted_formatting_constants(self) -> None:
        self.assertEqual(
            MATRIX_FIXED_COLUMN_WIDTHS,
            {1: 30.42578125, 2: 127.0, 3: 13.28515625, 4: 9.140625},
        )
        self.assertEqual(PAIR_1_COLUMN_WIDTHS, (13.0, 13.0, 23.0))
        self.assertEqual(PAIR_2_PLUS_COLUMN_WIDTHS, (9.140625, 13.0, 23.28515625))

    def test_generated_two_pair_workbook_persisted_formatting(self) -> None:
        google = build_google_units_index([_google_row("BCC0500", "шт")])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("r1", code="BCC0500", units="м", qty="1")],
                google,
                self.matrix_path,
            )
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            _assert_persisted_matrix_formatting(self, ws, pair_count=2, last_row=2)
            self.assertEqual(ws["F2"].number_format, "@")
            self.assertEqual(ws["I2"].number_format, "@")
            self.assertEqual(ws["F2"].value, "1")
            self.assertEqual(ws["I2"].value, "?")
            formula_rgb = str(ws["G2"].fill.fgColor.rgb or "")
            self.assertTrue(formula_rgb.endswith(self._formula_gray_rgb_suffix()))
            yellow_rgb = str(ws["I2"].fill.fgColor.rgb or "")
            self.assertTrue(yellow_rgb.endswith("FFFF00"))
            self.assertTrue(str(ws["G2"].value or "").startswith("="))
        finally:
            wb.close()
        self.assertTrue(load_matrix(self.matrix_path).layout_current)

    def test_generated_three_pair_workbook_pair3_fallback_widths(self) -> None:
        path = Path(self._tmpdir.name) / "three_pair_matrix.xlsx"
        self._write_current_layout_matrix(
            path,
            rows=[("BCC0501", "шт", [("шт", "1"), ("м", "10"), ("упак", "50")])],
        )
        wb = load_workbook(path)
        try:
            ws = wb.active
            for col, letter in ((11, "K"), (12, "L"), (13, "M")):
                self.assertAlmostEqual(
                    ws.column_dimensions[letter].width,
                    expected_matrix_column_width(col, pair_count=3),
                    places=6,
                )
                self.assertAlmostEqual(
                    expected_matrix_column_width(col, pair_count=3),
                    PAIR_2_PLUS_COLUMN_WIDTHS[col - 11],
                    places=6,
                )
        finally:
            wb.close()

    def test_formatting_mismatch_triggers_migration_and_restore(self) -> None:
        google = build_google_units_index([_google_row("BCC0502", "шт")])
        build_conversion_plan(
            [_request("r1", code="BCC0502", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        self.assertTrue(load_matrix(self.matrix_path).layout_current)

        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            ws.column_dimensions["A"].width = 12.0
            ws.row_dimensions[1].height = 15.0
            ws["B2"].alignment = Alignment(horizontal="center", wrap_text=True)
            wb.save(self.matrix_path)
        finally:
            wb.close()

        self.assertFalse(load_matrix(self.matrix_path).layout_current)
        build_conversion_plan(
            [_request("r1", code="BCC0502", units="шт", qty="1")],
            google,
            self.matrix_path,
        )
        self.assertTrue(load_matrix(self.matrix_path).layout_current)
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            _assert_persisted_matrix_formatting(self, ws, pair_count=1, last_row=2)
        finally:
            wb.close()

    def test_sparse_unstyled_cells_mark_layout_stale_without_crashing(self) -> None:
        path = Path(self._tmpdir.name) / "sparse_unstyled_matrix.xlsx"
        wb = Workbook()
        ws = wb.active
        header = (
            *MATRIX_HEADER,
            f"{matrix_module.PAIR_SOURCE_PREFIX}1",
            f"{PAIR_COEF_PREFIX}1",
            f"{PAIR_FORMULA_PREFIX}1",
        )
        ws.append(header)
        ws.append(["BCC0504", None, "", "шт", "шт", "1"])
        _apply_persisted_matrix_formatting(ws, pair_count=1, last_row=1)
        ws["F2"].number_format = "@"
        wb.save(path)
        wb.close()

        document = load_matrix(path)

        self.assertEqual(len(document.rows), 1)
        self.assertEqual(document.rows[0].code_normalized, "BCC0504")
        self.assertFalse(document.layout_current)

    def test_correctly_formatted_semantically_unchanged_skips_rewrite(self) -> None:
        google = build_google_units_index([_google_row("BCC0503", "шт")])
        requests = [_request("r1", code="BCC0503", units="шт", qty="1")]
        build_conversion_plan(requests, google, self.matrix_path)
        doc = load_matrix(self.matrix_path)
        self.assertTrue(doc.layout_current)

        hook_calls: list[tuple[Path, Path]] = []

        def _record_hook(tmp_path: Path, dest_path: Path) -> None:
            hook_calls.append((tmp_path, dest_path))

        matrix_module._before_replace_hook = _record_hook
        build_conversion_plan(requests, google, self.matrix_path)
        self.assertEqual(hook_calls, [])
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            _assert_persisted_matrix_formatting(self, ws, pair_count=1, last_row=2)
        finally:
            wb.close()


class ReciprocalCoefficientSmokeTest(unittest.TestCase):
    def test_normalize_reciprocal_coefficient_helper(self) -> None:
        third = Decimal("1") / Decimal("3")
        seventh = Decimal("1") / Decimal("7")
        self.assertEqual(normalize_reciprocal_coefficient(Decimal("0.33")), third)
        self.assertEqual(normalize_reciprocal_coefficient(Decimal("0.14")), seventh)
        self.assertEqual(normalize_reciprocal_coefficient(Decimal("0.5")), Decimal("0.5"))
        self.assertEqual(normalize_reciprocal_coefficient(Decimal("0.25")), Decimal("0.25"))
        self.assertEqual(normalize_reciprocal_coefficient(Decimal("0.32")), Decimal("0.32"))
        self.assertEqual(normalize_reciprocal_coefficient(Decimal("2")), Decimal("2"))
        self.assertEqual(normalize_reciprocal_coefficient(Decimal("100")), Decimal("100"))


class ReciprocalMatrixConversionSmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.matrix_path = Path(self._tmpdir.name) / "units_matrix.xlsx"
        matrix_module._before_replace_hook = None

    def tearDown(self) -> None:
        matrix_module._before_replace_hook = None
        self._tmpdir.cleanup()
        import gc

        gc.collect()

    def _seed_nonidentity_placeholder(self, code: str, google_unit: str, source_unit: str) -> None:
        google = build_google_units_index([_google_row(code, google_unit)])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("seed", code=code, units=source_unit, qty="1")],
                google,
                self.matrix_path,
            )

    def _set_matrix_coef(self, row_index: int, pair_index: int, coef: object) -> None:
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            coef_col = _coef_column(ws, pair_index)
            ws.cell(row=row_index, column=coef_col, value=coef)
            wb.save(self.matrix_path)
        finally:
            wb.close()

    def test_raw_033_meters_to_pieces(self) -> None:
        code = "BCC0200"
        self._seed_nonidentity_placeholder(code, "шт", "м")
        self._set_matrix_coef(2, 2, "0,33")
        google = build_google_units_index([_google_row(code, "шт")])

        plan_qty3 = build_conversion_plan(
            [_request("r3", code=code, units="м", qty="3")],
            google,
            self.matrix_path,
        )
        action3 = action_by_id(plan_qty3, "r3")
        assert action3 is not None
        third = Decimal("1") / Decimal("3")
        self.assertEqual(converted_quantity(action3), Decimal("1"))
        self.assertEqual(action3.coefficient, third)
        self.assertIn(f"coef={third}", action3.trace)
        dep3 = next(
            dep
            for dep in plan_qty3.dependencies
            if dep.source_unit == "м" and dep.target_unit == "шт"
        )
        self.assertEqual(dep3.coefficient, third)

        plan_qty300 = build_conversion_plan(
            [_request("r300", code=code, units="м", qty="300")],
            google,
            self.matrix_path,
        )
        action300 = action_by_id(plan_qty300, "r300")
        assert action300 is not None
        self.assertEqual(converted_quantity(action300), Decimal("100"))
        self.assertEqual(action300.coefficient, third)

        doc = load_matrix(self.matrix_path)
        meter_pair = next(
            pair for pair in doc.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair.coefficient_raw, "0,33")
        self.assertEqual(meter_pair.coefficient, third)

        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            coef_col = _coef_column(ws, 2)
            self.assertEqual(ws.cell(row=2, column=coef_col).value, "0,33")
        finally:
            wb.close()

    def test_raw_032_does_not_snap_and_stays_fractional(self) -> None:
        code = "BCC0201"
        self._seed_nonidentity_placeholder(code, "шт", "м")
        self._set_matrix_coef(2, 2, "0,32")
        google = build_google_units_index([_google_row(code, "шт")])
        plan = build_conversion_plan(
            [_request("r1", code=code, units="м", qty="3")],
            google,
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertTrue(action.fractional_result)
        self.assertEqual(converted_quantity(action), Decimal("0.96"))
        self.assertEqual(len(plan.fractional_issues), 1)
        self.assertNotIn("Fractional conversion result exceeds tolerance", str(plan.warnings))
        doc = load_matrix(self.matrix_path)
        meter_pair = next(
            pair for pair in doc.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair.coefficient_raw, "0,32")
        self.assertEqual(meter_pair.coefficient, Decimal("0.32"))


class ExactFractionCoefficientSmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.matrix_path = Path(self._tmpdir.name) / "units_matrix.xlsx"
        matrix_module._before_replace_hook = None

    def tearDown(self) -> None:
        matrix_module._before_replace_hook = None
        self._tmpdir.cleanup()
        import gc

        gc.collect()

    def _seed_nonidentity_placeholder(self, code: str, google_unit: str, source_unit: str) -> None:
        google = build_google_units_index([_google_row(code, google_unit)])
        with self.assertRaises(UnitsConversionError):
            build_conversion_plan(
                [_request("seed", code=code, units=source_unit, qty="1")],
                google,
                self.matrix_path,
            )

    def _set_matrix_coef(self, row_index: int, pair_index: int, coef: object) -> None:
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            coef_col = _coef_column(ws, pair_index)
            ws.cell(row=row_index, column=coef_col, value=coef)
            wb.save(self.matrix_path)
        finally:
            wb.close()

    def _assert_coef_cells_text_format(self) -> None:
        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            for col_idx in range(1, ws.max_column + 1):
                title = str(ws.cell(row=1, column=col_idx).value or "")
                if not title.startswith(PAIR_COEF_PREFIX):
                    continue
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    self.assertEqual(
                        cell.number_format,
                        "@",
                        msg=f"{cell.coordinate} must stay text-formatted",
                    )
        finally:
            wb.close()

    def test_parse_fraction_syntax_direct(self) -> None:
        third = Decimal("1") / Decimal("3")
        two_thirds = Decimal("2") / Decimal("3")
        half = Decimal("1") / Decimal("2")

        for raw, expected in (
            ("1/3", third),
            ("=1/3", third),
            ("= 1/3", third),
            ("1,5/3", Decimal("0.5")),
            ("2/3", two_thirds),
        ):
            with self.subTest(raw=raw):
                coef_raw, coef_value, is_placeholder = _parse_coefficient(raw)
                self.assertEqual(coef_raw, raw.strip())
                self.assertEqual(coef_value, expected)
                self.assertFalse(is_placeholder)

        for raw in ("1/0", "1/2/3", "a/b", "1/", "/3", "=1/0"):
            with self.subTest(raw=raw):
                coef_raw, coef_value, is_placeholder = _parse_coefficient(raw)
                self.assertEqual(coef_raw, raw.strip())
                self.assertIsNone(coef_value)
                self.assertTrue(is_placeholder)

    def test_fraction_one_third_integration(self) -> None:
        code = "BCC0300"
        self._seed_nonidentity_placeholder(code, "шт", "м")
        self._set_matrix_coef(2, 2, "1/3")
        google = build_google_units_index([_google_row(code, "шт")])
        third = Decimal("1") / Decimal("3")

        plan_qty3 = build_conversion_plan(
            [_request("r3", code=code, units="м", qty="3")],
            google,
            self.matrix_path,
        )
        action3 = action_by_id(plan_qty3, "r3")
        assert action3 is not None
        self.assertEqual(converted_quantity(action3), Decimal("1"))
        self.assertEqual(action3.coefficient, third)

        plan_qty300 = build_conversion_plan(
            [_request("r300", code=code, units="м", qty="300")],
            google,
            self.matrix_path,
        )
        action300 = action_by_id(plan_qty300, "r300")
        assert action300 is not None
        self.assertEqual(converted_quantity(action300), Decimal("100"))
        self.assertEqual(action300.coefficient, third)

        doc = load_matrix(self.matrix_path)
        meter_pair = next(
            pair for pair in doc.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair.coefficient_raw, "1/3")
        self.assertEqual(meter_pair.coefficient, third)

        wb = load_workbook(self.matrix_path)
        try:
            ws = wb.active
            coef_col = _coef_column(ws, 2)
            coef_cell = ws.cell(row=2, column=coef_col)
            self.assertEqual(coef_cell.value, "1/3")
            self.assertEqual(coef_cell.number_format, "@")
            formula_cell = ws.cell(row=2, column=coef_col + 1)
            formula_text = str(formula_cell.value or "")
            self.assertIn(f"{get_column_letter(coef_col)}2", formula_text)
            self.assertIn('&"/"&', formula_text)
            self.assertIn("$D2", formula_text)
        finally:
            wb.close()

        self._assert_coef_cells_text_format()

    def test_fraction_two_thirds_integration(self) -> None:
        code = "BCC0301"
        self._seed_nonidentity_placeholder(code, "шт", "м")
        self._set_matrix_coef(2, 2, "2/3")
        google = build_google_units_index([_google_row(code, "шт")])
        two_thirds = Decimal("2") / Decimal("3")

        plan = build_conversion_plan(
            [_request("r1", code=code, units="м", qty="3")],
            google,
            self.matrix_path,
        )
        action = action_by_id(plan, "r1")
        assert action is not None
        self.assertEqual(converted_quantity(action), Decimal("2"))
        self.assertEqual(action.coefficient, two_thirds)

        doc = load_matrix(self.matrix_path)
        meter_pair = next(
            pair for pair in doc.rows[0].pairs if pair.source_normalized == "м"
        )
        self.assertEqual(meter_pair.coefficient_raw, "2/3")
        self.assertEqual(meter_pair.coefficient, two_thirds)

    def test_writer_text_format_on_blank_coef_cells(self) -> None:
        code = "BCC0302"
        self._seed_nonidentity_placeholder(code, "шт", "м")
        self._assert_coef_cells_text_format()


if __name__ == "__main__":
    unittest.main()
