"""Readonly: compare checklist xlsx statuses vs actual folders."""

from __future__ import annotations

from openpyxl import load_workbook

from RFQ.rfp_parts.ds_checklist import (
    DEFAULT_CHECKLIST_FILE,
    DEFAULT_DECREASE_DIR,
    DEFAULT_INCREASE_DIR,
    collect_folder_ds,
    load_ds_checklist,
    present_in,
    run_checklist_compare,
)


def main() -> None:
    wb = load_workbook(DEFAULT_CHECKLIST_FILE, read_only=True, data_only=True)
    try:
        ws = wb["ДС"] if "ДС" in wb.sheetnames else wb.active
        print("Sheet:", ws.title)
        print("--- rows with НЕТ or suspect DS ---")
        for row in ws.iter_rows(min_row=1, max_col=8, values_only=True):
            vals = [("" if v is None else str(v).strip()) for v in row]
            while len(vals) < 8:
                vals.append("")
            inc, inc_st, dec, dec_st, other = vals[0], vals[1], vals[3], vals[4], vals[5]
            blob = " ".join(vals).upper()
            suspect = any(
                t in blob
                for t in (
                    "58",
                    "59",
                    "85",
                    "97",
                    "98",
                    "96/70",
                    "11/4",
                    "12/6",
                )
            )
            if "НЕТ" in (inc_st.upper(), dec_st.upper()) or suspect:
                print(
                    f"A={inc!r:12} B={inc_st!r:16} "
                    f"D={dec!r:12} E={dec_st!r:16} F={other!r}"
                )
    finally:
        wb.close()

    print("--- loose filename search ---")
    needles = ("ДС58", "ДС59", "ДС85", "ДС97", "ДС98", "ДС96", "ДС11", "ДС12")
    for label, folder in (
        ("INC", DEFAULT_INCREASE_DIR),
        ("DEC", DEFAULT_DECREASE_DIR),
    ):
        for path in sorted(folder.iterdir()):
            if not path.is_file():
                continue
            name_u = path.name.upper().replace(" ", "")
            if any(n in name_u for n in needles):
                print(label, path.name)

    print("--- checklist comments for ERROR candidates ---")
    cl = load_ds_checklist()
    candidates = {"58", "59", "85", "97", "98", "96/70А", "11/4", "12/6"}
    for entry in cl.increase + cl.decrease:
        if entry.ds in candidates or entry.ds.upper() in {c.upper() for c in candidates}:
            print(
                entry.kind,
                entry.ds,
                "comment=",
                repr(entry.comment),
                "not_required=",
                entry.not_required,
            )

    print("--- extras in folders not covered by checklist? ---")
    result = run_checklist_compare(progress=False)
    extras = [i for i in result.issues if i.code == "extra"]
    print("extra count", len(extras))
    for issue in extras:
        print(issue.level, issue.kind, issue.ds, issue.message)

    print("--- folder vs expected present matrix for decrease list ---")
    dec_ds = collect_folder_ds(DEFAULT_DECREASE_DIR)
    print("folder keys:", sorted(dec_ds.keys()))
    for entry in cl.decrease:
        print(
            f"{entry.ds:12} not_req={entry.not_required} "
            f"present={present_in(dec_ds, entry.ds)} comment={entry.comment!r}"
        )


if __name__ == "__main__":
    main()
