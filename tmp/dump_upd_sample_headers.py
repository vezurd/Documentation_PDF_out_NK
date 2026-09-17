"""Dump first rows of representative UPD workbooks (UTF-8)."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(
    r"\\bcc\root\CurProjects\di_manegers\Пятницкий_П"
    r"\Амурский ГХК\Поставки\Файл закачки УПД по всем ДС"
)
OUT = Path(__file__).resolve().parents[1] / "tmp" / "upd_sample_headers.txt"

SAMPLES = [
    ROOT / "ГФ 2" / "УПД ГФ.xlsx",
    ROOT / "ГФ 2" / "УПДГФ.xlsx",
    ROOT / "2087882" / "УПД шлагбаум.xlsx",
    ROOT / "ДС 13" / "PL_2076961.1_2424.04.xlsx",
    ROOT / "ДС 13" / "PL_2076961.1_2424.08.xlsx",
    ROOT / "8350" / "Общая.xlsx",
    ROOT / "ДС 15" / "PL_2076961.3_2510.13.xlsx",
    ROOT / "ДС 31" / "PL_2076961.15_2896.01.xlsx",
]


def _cell(v: object) -> str:
    if v is None:
        return ""
    return str(v).replace("\r\n", " / ").replace("\n", " / ").strip()


def dump_wb(path: Path, lines: list[str], max_row: int = 8) -> None:
    rel = path.relative_to(ROOT)
    lines.append("=" * 88)
    lines.append(f"FILE {rel}  exists={path.exists()}  size={path.stat().st_size if path.exists() else 0}")
    if not path.exists():
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        lines.append(f"sheets: {[ws.title for ws in wb.worksheets]}")
        for ws in wb.worksheets:
            lines.append(f"-- sheet={ws.title!r} max_row={ws.max_row} max_col={ws.max_column}")
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, values_only=True), start=1):
                cells = [_cell(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                parts = [f"{get_column_letter(idx)}={val}" for idx, val in enumerate(cells, start=1) if val]
                lines.append(f"  r{i}: " + " | ".join(parts) if parts else f"  r{i}: <empty>")
    finally:
        wb.close()


def main() -> int:
    lines: list[str] = []
    # If named samples missing, pick first xlsx from a few folders.
    missing = [p for p in SAMPLES if not p.exists()]
    extras: list[Path] = []
    if missing:
        lines.append("missing named samples:")
        for p in missing:
            lines.append(f"  {p}")
        for folder in ["ДС 31", "ДС 23", "ДС 45", "ДС 29"]:
            d = ROOT / folder
            if d.is_dir():
                found = sorted(d.glob("*.xlsx"))[:2]
                extras.extend(found)

    for p in SAMPLES + extras:
        if p.exists() and p.suffix.lower() == ".xlsx":
            dump_wb(p, lines)
            lines.append("")

    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
