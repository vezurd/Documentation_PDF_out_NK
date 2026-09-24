"""Offline smoke tests for RFP Step4 packing-list integration."""

from __future__ import annotations

import inspect
import io
import json
import pickle
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _configure_utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")
            except Exception:
                pass


_configure_utf8_stdio()

from base.base_classes import RowStd, RowType, TableComments
from base.tables_columns import (
    ANNOTATION,
    CODE,
    CODE_MTO,
    CODE_VO,
    DS_ACTUAL,
    DS_MANAGER,
    DS_NAME,
    DS_NUMBER,
    DS_SYSTEM,
    DS_TITLE,
    IN_CABINET,
    MATCH_STATUS,
    MATCH_STATUS_VO,
    MTO_CODE_STRUCK,
    NAME,
    PATH_MTO,
    PATH_RFP,
    ROW_TYPE,
    TAG_MTO,
    TAG_VO,
    TAGS,
    TITLE,
    TYPE_MARK,
    UL_CODE,
    UL_COMPARE_STATUS,
    UL_DATA_STATUS,
    UL_NAME,
    UL_ORDERED_VALUES,
    UL_REMAINING_VALUES,
    UL_SOURCE_FILES,
    UL_TAG_MATCH_STATUS,
    UL_TAGS,
    UL_TYPE_MARK,
    UL_UNITS,
    UL_VALUES,
    UL_VENDOR,
    UNITS,
    UNITS_CHECK_STATUS,
    UNITS_CONVERSION_TRACE,
    UNITS_MTO,
    VALUES,
    VALUES_MTO,
    VALUES_MTO_RFP_DIFF,
    VALUES_MTO_UL_DIFF,
    VALUES_VO,
    VENDOR,
)
from RFQ.packing_list_provider import (
    PACKING_CACHE_VERSION,
    PackingCacheMeta,
    PackingDataset,
    PackingIssue,
    PackingQualityLevel,
    RegistryPackingIndex,
    build_registry_planting_index,
    load_packing_dataset,
    normalize_packing_key_part,
    packing_match_key,
    registry_packing_scope,
    rfp_packing_match_key,
    rfp_row_actual_ds,
    rfp_row_supply_key,
    slim_packing_rows,
)
from RFQ.tags_rfp_compare.gem_supply_codes import (
    apply_gem_supply_status,
    load_gem_supply_codes,
    paint_gem_supply_rows,
)
from RFQ.tags_rfp_compare.rfp_tags_utils import (
    get_default_asbuild_config,
    get_default_config,
)
from RFQ.tags_rfp_compare.step4.step4_2_match_rfp_with_mto import add_mto_data_to_rfp_row
from RFQ.tags_rfp_compare.step4.step4_4_add_unmatched_mto_rows import _copy_mto_el_to_row
from RFQ.tags_rfp_compare.step4.step4_postmerge_ops import (
    apply_mto_rfp_quantity_diff,
    apply_mto_ul_quantity_diff,
    converted_units_comment_text,
)
from RFQ.tags_rfp_compare.step4.step4_6_save_match_result_to_excel import (
    excel_path_link,
    parse_single_ul_source,
    save_match_result_to_excel,
    should_write_excel_comment,
)
from RFQ.tags_rfp_compare.step4.step4_quantity_balance import (
    format_input_ul_detail,
    run_output_ul_quantity_balance_check,
    sum_ul_values_by_title,
)
from RFQ.tags_rfp_compare.step4.step4_packing_compare import (
    FALLBACK_COMMENT_MARKER,
    DATA_PARTIAL,
    DATA_ROW_PROBLEM,
    DATA_UNAVAILABLE,
    STATUS_COMPLETE,
    STATUS_DATA_PROBLEM,
    STATUS_GEM_SUPPLY,
    STATUS_ZIP_SMR_PNR,
    STATUS_MTO_DELIVERED,
    STATUS_MTO_DELIVERED_SHORT,
    STATUS_MTO_ONLY,
    STATUS_OPEN_UPD,
    STATUS_OVERDELIVERY,
    STATUS_PACKING_ONLY,
    STATUS_RFP_ONLY_EXTRA,
    STATUS_SHORTFALL,
    STATUS_TAG_MISMATCH_COLLECTED,
    STATUS_TAG_MATCHED_VIA_MTO,
    STATUS_UNAVAILABLE,
    STATUS_VO_DELIVERED,
    STATUS_VO_DELIVERED_SHORT,
    STATUS_VO_ONLY,
    RfpPackingFatalError,
    _PackingUnit,
    _RowAllocationState,
    _delivered_qty,
    _format_ul_source_files,
    _merge_packing_units_trace,
    _status_for_remaining,
    build_ordered_snapshot,
    apply_zip_smr_pnr_status,
    compare_rfp_rows_with_packing,
    packing_queue_input_by_title,
    parse_composite_title,
    save_rfp_packing_report,
    ul_name_is_zip_smr_pnr,
)
from RFQ.tags_rfp_compare.step4_analyze_and_match import (
    _collapse_rfp_rows_before_export,
    step4_analyze_and_match,
)
from utils.colors import Color

UL_COLUMNS = (
    UL_ORDERED_VALUES,
    UL_VALUES,
    UL_UNITS,
    UL_REMAINING_VALUES,
    UL_COMPARE_STATUS,
    UL_DATA_STATUS,
    UL_TAG_MATCH_STATUS,
    UL_CODE,
    UL_SOURCE_FILES,
    UL_NAME,
    UL_TYPE_MARK,
    UL_VENDOR,
    UL_TAGS,
)

UNITS_COLUMNS = (
    UNITS,
    UNITS_MTO,
    UNITS_CHECK_STATUS,
    UNITS_CONVERSION_TRACE,
)


def _ul_source(file_name: str, folder: str = "согл УЛ ДС1") -> str:
    return f"{folder}/{file_name}"


def _row(
    *,
    title: str = "8950-SOO1",
    system: str = "",
    code: str = "CODE",
    quantity: object = 1,
    tags: object = "",
    tag_mto: object = "",
    units: str = "шт",
    code_mto: str = "",
    code_vo: str = "",
    ds_name: str = "ДС1",
    ds_number: str = "",
    source: str = "согл УЛ ДС1/packing-a.xlsx",
    sheet: str = "Single 1",
    excel_row: int = 10,
    match_status: str = "",
    match_status_vo: str = "",
    values_mto: object = None,
    values_vo: object = None,
    tag_vo: object = "",
) -> RowStd:
    row = RowStd.get_std_check_row({}, TableComments())
    row.row_type = RowType.position_row
    row.el[ROW_TYPE].value = RowType.position_row
    values = {
        DS_NAME: ds_name,
        DS_TITLE: title,
        DS_SYSTEM: system,
        DS_NUMBER: ds_number,
        CODE: code,
        CODE_MTO: code_mto,
        CODE_VO: code_vo,
        VALUES: quantity,
        TAGS: tags,
        TAG_MTO: tag_mto,
        TAG_VO: tag_vo,
        UNITS: units,
        NAME: f"Item {code}",
        TYPE_MARK: "TM",
        VENDOR: "Vendor",
        ANNOTATION: source,
        TITLE: sheet,
        MATCH_STATUS: match_status,
        MATCH_STATUS_VO: match_status_vo,
    }
    for column, value in values.items():
        row.el[column].value = value
    if values_mto is not None:
        row.el[VALUES_MTO].value = values_mto
    if values_vo is not None:
        row.el[VALUES_VO].value = values_vo
    row._packing_source_row = excel_row
    return row


def _dataset(
    rows: list[RowStd],
    *,
    quality: PackingQualityLevel = PackingQualityLevel.OK,
    issues: list[PackingIssue] | None = None,
) -> PackingDataset:
    return PackingDataset(
        rows=rows,
        meta=PackingCacheMeta(
            version=PACKING_CACHE_VERSION,
            created_at="2026-07-29T12:00:00+00:00",
            root="C:/packing",
            fingerprint="offline-smoke",
            counters={"position_rows": len(rows)},
            critical_ok=quality == PackingQualityLevel.OK,
            report_paths={"critical": "tsd_packing_critical.txt"},
        ),
        quality=quality,
        issues=list(issues or []),
        cache_path="memory.cache",
    )


def _values(rows: list[RowStd], status: str) -> list[RowStd]:
    return [row for row in rows if row.get_value(UL_COMPARE_STATUS) == status]


class RfpPackingMatcherSmokeTest(unittest.TestCase):
    def test_ul_source_files_grouped_by_file_with_newlines(self) -> None:
        first = r"согл УЛ ДС20\Packing list PL_2076961.8_2890.05 ред.3.12.xlsx"
        second = r"согл УЛ ДС20\Packing list PL_2076961.8_2890.12_редакция1_211125.xlsx"
        self.assertEqual(
            _format_ul_source_files(
                [
                    f"{first} · Single 1 · строка 18",
                    f"{first} · Single 2 · строка 16",
                    f"{first} · Single 3 · строка 16",
                    f"{first} · Single 4 · строка 16",
                    f"{first} · Single 5 · строка 16",
                    f"{second} · Single 2 · строка 16",
                ]
            ),
            "\n".join(
                [
                    first,
                    "  Single 1 · строка 18",
                    "  Single 2 · строка 16",
                    "  Single 3 · строка 16",
                    "  Single 4 · строка 16",
                    "  Single 5 · строка 16",
                    second,
                    "  Single 2 · строка 16",
                ]
            ),
        )
        self.assertEqual(
            _format_ul_source_files([f"{first} · Single 1 · строка 18"]),
            f"{first} · Single 1 · строка 18",
        )

        target = _row(quantity=3, ds_name="ДС20")
        packing = [
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=first,
                sheet="Single 1",
                excel_row=18,
            ),
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=first,
                sheet="Single 2",
                excel_row=16,
            ),
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=second,
                sheet="Single 2",
                excel_row=16,
            ),
        ]
        rows, _audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(20, "8950", "SOO1", "CODE"): 3.0},
            _dataset(packing),
        )
        self.assertEqual(
            rows[0].get_value(UL_SOURCE_FILES),
            "\n".join(
                [
                    first,
                    "  Single 1 · строка 18",
                    "  Single 2 · строка 16",
                    second,
                    "  Single 2 · строка 16",
                ]
            ),
        )

    def test_strict_composite_title_and_original_snapshot(self) -> None:
        self.assertEqual(parse_composite_title(" 8950-SOO1 "), ("8950", "SOO1"))
        self.assertEqual(parse_composite_title("7421-SKUD.1"), ("7421", "SKUD.1"))
        self.assertEqual(parse_composite_title("8350-KSBMT"), ("8350", "KSBMT"))
        self.assertEqual(parse_composite_title("8950-POS3"), ("8950", "POS3"))
        for malformed in (
            "",
            "8950",
            "SOO1",
            "x8950-SOO1",
            "8950-SOO1-extra",
            "8950 SOO1",
            "8950-",
        ):
            self.assertIsNone(parse_composite_title(malformed), malformed)

        source = _row(quantity=3)
        ignored = _row(title="bad title", code="IGNORED", quantity=9)
        snapshot = build_ordered_snapshot([source, ignored])
        self.assertEqual(snapshot, {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 3.0})

        dotted = _row(title="7421-SKUD.1", code="SKUD-CODE", quantity=1)
        dotted_snapshot = build_ordered_snapshot([dotted])
        self.assertEqual(
            dotted_snapshot,
            {rfp_packing_match_key(1, "7421", "SKUD.1", "SKUD-CODE"): 1.0},
        )
        packing = _row(
            title="7421",
            system="SKUD.1",
            code="SKUD-CODE",
            quantity=1,
        )
        rows, audit = compare_rfp_rows_with_packing(
            [dotted],
            dotted_snapshot,
            _dataset([packing]),
        )
        self.assertEqual(dotted.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(dotted.get_value(UL_VALUES), 1.0)
        self.assertEqual(audit.stats.not_in_packing, 0)
        self.assertFalse(
            any(issue.code == "rfp_composite_title" for issue in audit.issues)
        )

    def test_quantity_three_allocates_to_unit_and_collapsed_rows_once(self) -> None:
        unit = _row(quantity=1, tags="T-2")
        collapsed = _row(quantity=2)
        packing = [
            _row(
                title="8950",
                system="SOO1",
                quantity=2,
                tags=["T-1", "T-2"],
                source=_ul_source("packing-a.xlsx"),
            ),
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                tags=["T-3"],
                source=_ul_source("packing-b.xlsx"),
            ),
        ]
        rows, audit = compare_rfp_rows_with_packing(
            [unit, collapsed],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 3.0},
            _dataset(packing),
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(unit.get_value(UL_VALUES), 1.0)
        self.assertEqual(unit.get_value(UL_TAGS), "T-2")
        self.assertEqual(collapsed.get_value(UL_VALUES), 2.0)
        self.assertEqual(
            set(collapsed.get_value(UL_TAGS).split("; ")),
            {"T-1", "T-3"},
        )
        collapsed_sources = collapsed.get_value(UL_SOURCE_FILES)
        self.assertIn("packing-a.xlsx", collapsed_sources)
        self.assertIn("packing-b.xlsx", collapsed_sources)
        self.assertIn("\n", collapsed_sources)
        self.assertNotIn("; ", collapsed_sources)
        self.assertEqual(unit.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(collapsed.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(audit.stats.packing_units, 3)
        self.assertEqual(audit.stats.allocated_units, 3)
        self.assertEqual(audit.stats.packing_only_added, 0)

    def test_tagged_and_untagged_leftovers_collapse_separately(self) -> None:
        target = _row(quantity=1, tags="A")
        packing = _row(
            title="8950",
            system="SOO1",
            quantity=5,
            tags=["A", "B"],
        )
        rows, audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
        )

        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 2)
        by_tag = {row.get_value(UL_TAGS) or "": row for row in leftovers}
        self.assertEqual(by_tag["B"].get_value(UL_VALUES), 1.0)
        self.assertEqual(by_tag[""].get_value(UL_VALUES), 3.0)
        self.assertEqual(by_tag["B"].get_value(DS_TITLE), "8950-SOO1")
        self.assertEqual(by_tag[""].get_value(DS_TITLE), "8950-SOO1")
        self.assertEqual(audit.stats.allocated_units, 1)
        self.assertEqual(audit.stats.packing_only_added, 2)

    def test_code_fallbacks_and_struck_mto_exclusion(self) -> None:
        mto = _row(code="RFP-MISS", code_mto="MTO-OK")
        vo = _row(code="RFP-MISS-2", code_vo="VO-OK")
        struck = _row(
            code="RFP-MISS-3",
            code_mto="MTO-STRUCK",
            code_vo="VO-AFTER-STRUCK",
        )
        struck.el[CODE_MTO].struck_value = "MTO-STRUCK"
        struck.el[MTO_CODE_STRUCK].value = "MTO-STRUCK"
        packing = [
            _row(title="8950", system="SOO1", code="MTO-OK", source=_ul_source("mto.xlsx")),
            _row(title="8950", system="SOO1", code="VO-OK", source=_ul_source("vo.xlsx")),
            _row(
                title="8950",
                system="SOO1",
                code="MTO-STRUCK",
                source=_ul_source("struck.xlsx"),
            ),
            _row(
                title="8950",
                system="SOO1",
                code="VO-AFTER-STRUCK",
                source=_ul_source("vo-fallback.xlsx"),
            ),
        ]
        ordered = {
            rfp_packing_match_key(1, "8950", "SOO1", mto.get_value(CODE)): 1.0,
            rfp_packing_match_key(1, "8950", "SOO1", vo.get_value(CODE)): 1.0,
            rfp_packing_match_key(1, "8950", "SOO1", struck.get_value(CODE)): 1.0,
        }
        rows, audit = compare_rfp_rows_with_packing(
            [mto, vo, struck],
            ordered,
            _dataset(packing),
        )

        self.assertIn("CODE_MTO", mto.el[UL_COMPARE_STATUS].comment)
        self.assertIn("CODE_VO", vo.el[UL_COMPARE_STATUS].comment)
        self.assertEqual(struck.get_value(UL_CODE), "VO-AFTER-STRUCK")
        self.assertIn("CODE_VO", struck.el[UL_COMPARE_STATUS].comment)
        self.assertEqual(audit.stats.fallback_code_matches, 3)
        self.assertEqual(len(_values(rows, STATUS_PACKING_ONLY)), 1)
        self.assertEqual(
            _values(rows, STATUS_PACKING_ONLY)[0].get_value(UL_CODE),
            "MTO-STRUCK",
        )

    def test_complete_short_no_ul_overdelivery_and_ul_only(self) -> None:
        complete = _row(code="COMPLETE", quantity=1)
        short = _row(code="SHORT", quantity=2)
        missing = _row(code="MISSING", quantity=1)
        over = _row(code="OVER", quantity=1)
        packing = [
            _row(title="8950", system="SOO1", code="COMPLETE", quantity=1),
            _row(title="8950", system="SOO1", code="SHORT", quantity=1),
            _row(title="8950", system="SOO1", code="OVER", quantity=3),
            _row(title="8950", system="SOO1", code="ONLY", quantity=2),
        ]
        ordered = {
            rfp_packing_match_key(1, "8950", "SOO1", row.get_value(CODE)): row.get_value(VALUES)
            for row in (complete, short, missing, over)
        }
        rows, _audit = compare_rfp_rows_with_packing(
            [complete, short, missing, over],
            ordered,
            _dataset(packing),
        )

        self.assertEqual(complete.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(short.get_value(UL_COMPARE_STATUS), STATUS_SHORTFALL)
        self.assertEqual(short.get_value(UL_REMAINING_VALUES), 1.0)
        self.assertEqual(missing.get_value(UL_COMPARE_STATUS), STATUS_RFP_ONLY_EXTRA)
        self.assertIsNone(missing.get_value(UL_VALUES))
        self.assertEqual(over.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(sum(row.get_value(UL_VALUES) for row in leftovers), 4.0)
        self.assertTrue(any(row.get_value(UL_CODE) == "OVER" for row in leftovers))
        self.assertTrue(any(row.get_value(UL_CODE) == "ONLY" for row in leftovers))
        self.assertEqual(_status_for_remaining(-1.0)[0], STATUS_OVERDELIVERY)

    def test_cross_tag_same_code_sets_tag_match_status(self) -> None:
        target = _row(quantity=1, tags="A")
        packing = _row(title="8950", system="SOO1", quantity=1, tags=["B"])
        rows, _audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(target.get_value(UL_VALUES), 1.0)
        self.assertEqual(
            target.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MISMATCH_COLLECTED,
        )
        self.assertEqual(target.el[UL_TAG_MATCH_STATUS].color, Color.yellow)
        self.assertEqual(target.get_value(UL_TAGS), "B")
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_sibling_rfp_tag_not_stolen_when_wrong_title_has_other_tag(self) -> None:
        row_1001 = _row(
            title="8525-SOT",
            code="BCC0000152",
            quantity=1,
            tags="8525-SS-01-S-AVI-1001",
            code_mto="BCC0000152",
            values_mto=1,
        )
        row_1002 = _row(
            title="8525-SOT",
            code="BCC0000152",
            quantity=1,
            tags="8525-SS-01-S-AVI-1002",
            code_mto="BCC0000152",
            values_mto=1,
        )
        packing = [
            _row(
                title="8525",
                system="SOT",
                code="BCC0000152",
                quantity=1,
                tags=["8525-SS-01-S-AVI-1002"],
            ),
            _row(
                title="6550",
                system="SOT",
                code="BCC0000152",
                quantity=1,
                tags=["8525-SS-01-S-AVI-1001"],
            ),
        ]
        rows, _audit = compare_rfp_rows_with_packing(
            [row_1001, row_1002],
            {rfp_packing_match_key(1, "8525", "SOT", "BCC0000152"): 2.0},
            _dataset(packing),
        )
        self.assertEqual(row_1002.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(row_1002.get_value(UL_VALUES), 1.0)
        self.assertEqual(row_1002.get_value(UL_TAGS), "8525-SS-01-S-AVI-1002")
        self.assertFalse(row_1002.get_value(UL_TAG_MATCH_STATUS))
        self.assertEqual(row_1001.get_value(UL_COMPARE_STATUS), STATUS_OPEN_UPD)
        self.assertIsNone(row_1001.get_value(UL_VALUES))
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        self.assertEqual(leftovers[0].get_value(DS_TITLE), "6550-SOT")
        self.assertEqual(leftovers[0].get_value(UL_TAGS), "8525-SS-01-S-AVI-1001")

    def test_sibling_on_same_title_does_not_steal_other_rfp_tag(self) -> None:
        row_a = _row(
            quantity=1,
            tags="A",
            code_mto="CODE",
            values_mto=1,
        )
        row_b = _row(
            quantity=1,
            tags="B",
            code_mto="CODE",
            values_mto=1,
        )
        packing = _row(title="8950", system="SOO1", quantity=1, tags=["B"])
        rows, _audit = compare_rfp_rows_with_packing(
            [row_a, row_b],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 2.0},
            _dataset([packing]),
        )
        self.assertEqual(row_a.get_value(UL_COMPARE_STATUS), STATUS_OPEN_UPD)
        self.assertIsNone(row_a.get_value(UL_VALUES))
        self.assertEqual(row_b.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(row_b.get_value(UL_VALUES), 1.0)
        self.assertEqual(row_b.get_value(UL_TAGS), "B")
        self.assertFalse(row_b.get_value(UL_TAG_MATCH_STATUS))
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_mto_tag_match_before_blind_fill(self) -> None:
        target = _row(quantity=1, tags="A", tag_mto="B")
        packing = _row(title="8950", system="SOO1", quantity=1, tags=["B"])
        rows, _audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
            use_mto_tags=True,
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(target.get_value(UL_VALUES), 1.0)
        self.assertEqual(
            target.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MATCHED_VIA_MTO,
        )
        self.assertEqual(target.el[UL_TAG_MATCH_STATUS].color, Color.soft_yellow)
        self.assertEqual(target.get_value(UL_TAGS), "B")
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_mto_tag_match_disabled_falls_back_to_mismatch(self) -> None:
        target = _row(quantity=1, tags="A", tag_mto="B")
        packing = _row(title="8950", system="SOO1", quantity=1, tags=["B"])
        compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
            use_mto_tags=False,
        )
        self.assertEqual(
            target.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MISMATCH_COLLECTED,
        )
        self.assertEqual(target.el[UL_TAG_MATCH_STATUS].color, Color.yellow)

    def test_empty_mto_tag_still_mismatch_on_cross_tag(self) -> None:
        target = _row(quantity=1, tags="A", tag_mto="")
        packing = _row(title="8950", system="SOO1", quantity=1, tags=["B"])
        compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
        )
        self.assertEqual(
            target.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MISMATCH_COLLECTED,
        )

    def test_mismatch_priority_over_via_mto_in_same_allocation(self) -> None:
        # RFP misses A; MTO takes B (via_mto); blind takes C (mismatch) → mismatch wins.
        target = _row(quantity=2, tags="A", tag_mto="B")
        packing = [
            _row(title="8950", system="SOO1", quantity=1, tags=["B"], excel_row=10),
            _row(title="8950", system="SOO1", quantity=1, tags=["C"], excel_row=11),
        ]
        compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 2.0},
            _dataset(packing),
            use_mto_tags=True,
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(target.get_value(UL_VALUES), 2.0)
        self.assertEqual(target.get_value(UL_TAGS), "B; C")
        self.assertEqual(
            target.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MISMATCH_COLLECTED,
        )
        self.assertEqual(target.el[UL_TAG_MATCH_STATUS].color, Color.yellow)

    def test_rfp_then_blind_cross_tag_is_mismatch_not_via_mto(self) -> None:
        target = _row(quantity=2, tags="A", tag_mto="B")
        packing = [
            _row(title="8950", system="SOO1", quantity=1, tags=["A"], excel_row=10),
            _row(title="8950", system="SOO1", quantity=1, tags=["C"], excel_row=11),
        ]
        compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 2.0},
            _dataset(packing),
            use_mto_tags=True,
        )
        self.assertEqual(
            target.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MISMATCH_COLLECTED,
        )

    def test_empty_allocation_does_not_write_zero_or_shortfall(self) -> None:
        target = _row(code="MISSING", quantity=2, code_mto="MTO-MISS", values_mto=2)
        rows, _audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "MISSING"): 2.0},
            _dataset([]),
        )
        self.assertIs(rows[0], target)
        self.assertIsNone(target.get_value(UL_VALUES))
        self.assertNotEqual(target.get_value(UL_COMPARE_STATUS), STATUS_SHORTFALL)
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_OPEN_UPD)
        self.assertEqual(target.el[UL_COMPARE_STATUS].color, Color.yellow)

    def test_presence_statuses_without_packing(self) -> None:
        rfp_only = _row(code="RFP-ONLY", quantity=1)
        mto_only = _row(
            code="",
            quantity="",
            code_mto="MTO-ONLY",
            values_mto=1,
            match_status="Добавлен из МТО",
        )
        vo_only = _row(
            code="",
            quantity="",
            code_vo="VO-ONLY",
            values_vo=1,
            match_status_vo="Добавлен из VO",
        )
        compare_rfp_rows_with_packing(
            [rfp_only, mto_only, vo_only],
            {rfp_packing_match_key(1, "8950", "SOO1", "RFP-ONLY"): 1.0},
            _dataset([]),
        )
        self.assertEqual(rfp_only.get_value(UL_COMPARE_STATUS), STATUS_RFP_ONLY_EXTRA)
        self.assertEqual(mto_only.get_value(UL_COMPARE_STATUS), STATUS_MTO_ONLY)
        self.assertEqual(mto_only.el[UL_COMPARE_STATUS].color, Color.yellow)
        self.assertEqual(vo_only.get_value(UL_COMPARE_STATUS), STATUS_VO_ONLY)
        self.assertIsNone(rfp_only.get_value(UL_VALUES))
        self.assertIsNone(mto_only.get_value(UL_VALUES))

    def test_mto_only_fallback_consumes_packing_queue(self) -> None:
        target = _row(
            code="",
            quantity="",
            code_mto="MTO-OK",
            values_mto=2,
            match_status="Добавлен из МТО",
        )
        packing = _row(title="8950", system="SOO1", code="MTO-OK", quantity=2)
        rows, audit = compare_rfp_rows_with_packing(
            [target],
            {},
            _dataset([packing]),
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)
        self.assertEqual(target.el[UL_COMPARE_STATUS].color, Color.match_matched)
        self.assertEqual(target.get_value(UL_VALUES), 2.0)
        self.assertEqual(target.get_value(UL_CODE), "MTO-OK")
        self.assertIn("CODE_MTO", target.el[UL_COMPARE_STATUS].comment)
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))
        self.assertEqual(audit.stats.allocated_units, 2)
        self.assertEqual(audit.stats.complete_zero_allocated, 0)

    def test_slim_cache_rows_match_rfp_compare(self) -> None:
        """v6 slim packing rows must satisfy the RFP matcher without KeyError."""
        target = _row(code="SLIM1", quantity=2)
        packing = slim_packing_rows(
            [
                _row(
                    title="8950",
                    system="SOO1",
                    code="SLIM1",
                    quantity=2,
                    tags=["A", "B"],
                )
            ]
        )
        ordered = {rfp_packing_match_key(1, "8950", "SOO1", "SLIM1"): 2.0}
        rows, audit = compare_rfp_rows_with_packing(
            [target], ordered, _dataset(packing)
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(audit.stats.invalid_quantity_rows, 0)
        self.assertEqual(len(rows), 1)

    def test_all_invalid_quantities_are_explicit_problem_rows(self) -> None:
        invalid = [None, "", "text", float("nan"), float("inf"), -1]
        packing = [
            _row(
                title="8950",
                system="SOO1",
                code=f"BAD-{index}",
                quantity=value,
                excel_row=20 + index,
            )
            for index, value in enumerate(invalid)
        ]
        rows, audit = compare_rfp_rows_with_packing([], {}, _dataset(packing))

        self.assertEqual(audit.stats.invalid_quantity_rows, len(invalid))
        self.assertEqual(len(rows), len(invalid))
        self.assertTrue(
            all(row.get_value(UL_COMPARE_STATUS) == STATUS_DATA_PROBLEM for row in rows)
        )
        self.assertTrue(all(row.get_value(UL_DATA_STATUS) == DATA_ROW_PROBLEM for row in rows))
        self.assertTrue(all(row.get_value(UL_VALUES) is None for row in rows))
        self.assertTrue(all(row.el[UL_VALUES].color == Color.red for row in rows))

    def test_fractional_ul_quantity_stays_in_queue(self) -> None:
        """Converted remainder (35 m × 0.01 = 0.35) must still match RFP, not drop out."""
        target = _row(quantity=0.35, units="км")
        packing = _row(title="8950", system="SOO1", quantity=0.35, units="км")
        ordered = {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 0.35}
        rows, audit = compare_rfp_rows_with_packing(
            [target], ordered, _dataset([packing])
        )
        self.assertEqual(audit.stats.invalid_quantity_rows, 0)
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertAlmostEqual(float(target.get_value(UL_VALUES)), 0.35)
        self.assertAlmostEqual(float(target.get_value(UL_REMAINING_VALUES)), 0.0)
        self.assertFalse(any(row.get_value(UL_COMPARE_STATUS) == STATUS_PACKING_ONLY for row in rows))
        qty_in, rows_in, _by_title = packing_queue_input_by_title(_dataset([packing]))
        self.assertAlmostEqual(qty_in, 0.35)
        self.assertEqual(rows_in, 1)

    def test_fractional_ul_splits_leftover_without_dropping_qty(self) -> None:
        target = _row(quantity=0.2, units="км")
        packing = _row(title="8950", system="SOO1", quantity=0.35, units="км")
        ordered = {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 0.2}
        rows, audit = compare_rfp_rows_with_packing(
            [target], ordered, _dataset([packing])
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertAlmostEqual(float(target.get_value(UL_VALUES)), 0.2)
        leftovers = [
            row for row in rows if row.get_value(UL_COMPARE_STATUS) == STATUS_PACKING_ONLY
        ]
        self.assertEqual(len(leftovers), 1)
        self.assertAlmostEqual(float(leftovers[0].get_value(UL_VALUES)), 0.15)
        self.assertEqual(audit.stats.invalid_quantity_rows, 0)

    def test_tags_exceed_quantity_is_fatal_and_auditable(self) -> None:
        packing = [
            _row(
                title="8950",
                system="SOO1",
                code="FATAL-1",
                quantity=1,
                tags=["A", "B"],
                source=_ul_source("fatal-a.xlsx"),
                excel_row=41,
            ),
            _row(
                title="8950",
                system="SOO1",
                code="FATAL-2",
                quantity=0,
                tags=["C"],
                source=_ul_source("fatal-b.xlsx"),
                excel_row=42,
            ),
        ]
        with self.assertRaises(RfpPackingFatalError) as raised:
            compare_rfp_rows_with_packing([], {}, _dataset(packing))

        fatal_issues = [
            issue
            for issue in raised.exception.audit.issues
            if issue.code == "packing_tags_exceed_quantity"
        ]
        self.assertEqual(len(fatal_issues), 2)
        self.assertEqual({issue.excel_row for issue in fatal_issues}, {41, 42})
        with tempfile.TemporaryDirectory() as tmp:
            report = Path(save_rfp_packing_report(raised.exception.audit, tmp))
            text = report.read_text(encoding="utf-8")
            self.assertIn("fatal-a.xlsx", text)
            self.assertIn("fatal-b.xlsx", text)
            self.assertIn("packing_tags_exceed_quantity", text)

    def test_unit_mismatch_is_row_problem(self) -> None:
        target = _row(code="UNIT", quantity=1, units="компл")
        packing = _row(
            title="8950",
            system="SOO1",
            code="UNIT",
            quantity=1,
            units="шт",
        )
        _rows, audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "UNIT"): 1.0},
            _dataset([packing]),
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_DATA_PROBLEM)
        self.assertEqual(target.get_value(UL_DATA_STATUS), DATA_ROW_PROBLEM)
        self.assertEqual(target.get_value(UL_REMAINING_VALUES), "")
        self.assertEqual(target.el[UL_UNITS].color, Color.yellow)
        self.assertEqual(audit.stats.unit_mismatches, 1)
        self.assertTrue(any(issue.code == "packing_units_mismatch" for issue in audit.issues))

    def test_unavailable_and_partial_datasets_are_visible(self) -> None:
        unavailable_issue = PackingIssue("cache_missing", "offline cache absent")
        unavailable = PackingDataset(
            rows=[],
            meta=None,
            quality=PackingQualityLevel.UNAVAILABLE,
            issues=[unavailable_issue],
            cache_path="missing.cache",
        )
        target = _row(code="UNAVAILABLE")
        compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "UNAVAILABLE"): 1.0},
            unavailable,
        )
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_UNAVAILABLE)
        self.assertEqual(target.get_value(UL_DATA_STATUS), DATA_UNAVAILABLE)
        self.assertEqual(target.el[UL_COMPARE_STATUS].color, Color.red)

        partial_issue = PackingIssue("loader_warning", "one source skipped")
        partial_target = _row(code="PARTIAL")
        partial_packing = _row(
            title="8950",
            system="SOO1",
            code="PARTIAL",
        )
        _rows, audit = compare_rfp_rows_with_packing(
            [partial_target],
            {rfp_packing_match_key(1, "8950", "SOO1", "PARTIAL"): 1.0},
            _dataset(
                [partial_packing],
                quality=PackingQualityLevel.PARTIAL,
                issues=[partial_issue],
            ),
        )
        self.assertEqual(partial_target.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(partial_target.get_value(UL_DATA_STATUS), DATA_PARTIAL)
        self.assertEqual(partial_target.el[UL_DATA_STATUS].color, Color.red)
        self.assertEqual(audit.quality, PackingQualityLevel.PARTIAL)

    def test_malformed_result_title_is_explicit(self) -> None:
        target = _row(title="prefix-8950-SOO1", code="BAD-TITLE")
        rows, audit = compare_rfp_rows_with_packing([target], {}, _dataset([]))
        self.assertIs(rows[0], target)
        self.assertEqual(target.get_value(UL_COMPARE_STATUS), STATUS_DATA_PROBLEM)
        self.assertEqual(target.get_value(UL_DATA_STATUS), DATA_ROW_PROBLEM)
        self.assertTrue(any(issue.code == "rfp_composite_title" for issue in audit.issues))

    def test_two_ds_same_title_code_do_not_share_ul_queue(self) -> None:
        row_ds29 = _row(
            ds_name="ДС29",
            title="8445-SOT",
            code="BCC0003052",
            quantity=1,
        )
        row_ds82 = _row(
            ds_name="ДС82",
            title="8445-SOT",
            code="BCC0003052",
            quantity=1,
        )
        packing = _row(
            title="8445",
            system="SOT",
            code="BCC0003052",
            quantity=1,
            source=_ul_source("file.xlsx", "согл УЛ ДС82"),
        )
        ordered = {
            rfp_packing_match_key(29, "8445", "SOT", "BCC0003052"): 1.0,
            rfp_packing_match_key(82, "8445", "SOT", "BCC0003052"): 1.0,
        }
        rows, _audit = compare_rfp_rows_with_packing(
            [row_ds29, row_ds82],
            ordered,
            _dataset([packing]),
        )
        self.assertIsNone(row_ds29.get_value(UL_VALUES))
        self.assertEqual(row_ds29.get_value(UL_COMPARE_STATUS), STATUS_RFP_ONLY_EXTRA)
        self.assertEqual(row_ds82.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(row_ds82.get_value(UL_VALUES), 1.0)
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertFalse(leftovers)
        self.assertNotEqual(row_ds29.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)

    def test_correction_ds_matches_actual_folder_not_sequential(self) -> None:
        rfp_match = _row(ds_name="ДС92_24Б", title="8950-SOO1", code="CODE")
        packing_actual = _row(
            title="8950",
            system="SOO1",
            code="CODE",
            quantity=1,
            source=_ul_source("file.xlsx", "согл УЛ ДС92"),
        )
        compare_rfp_rows_with_packing(
            [rfp_match],
            {rfp_packing_match_key(92, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing_actual]),
        )
        self.assertEqual(rfp_match.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(rfp_match.get_value(UL_VALUES), 1.0)
        self.assertEqual(rfp_match.get_value(DS_ACTUAL), "ДС92")

        rfp_seq = _row(ds_name="ДС92_24Б", title="8950-SOO1", code="CODE")
        packing_sequential = _row(
            title="8950",
            system="SOO1",
            code="CODE",
            quantity=1,
            source=_ul_source("file.xlsx", "согл УЛ ДС24"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [rfp_seq],
            {rfp_packing_match_key(92, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing_sequential]),
        )
        self.assertIsNone(rfp_seq.get_value(UL_VALUES))
        self.assertEqual(rfp_seq.get_value(UL_COMPARE_STATUS), STATUS_RFP_ONLY_EXTRA)
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        self.assertEqual(leftovers[0].get_value(UL_COMPARE_STATUS), STATUS_PACKING_ONLY)
        self.assertEqual(leftovers[0].get_value(DS_NAME), "ДС24_RFP_не_найден")
        self.assertEqual(leftovers[0].get_value(DS_ACTUAL), "ДС24")

    def test_shared_folder_two_tags_registry_both_complete(self) -> None:
        folder = "согл УЛ ДС15"
        folder_key = normalize_packing_key_part(folder)
        index = RegistryPackingIndex(
            known_source_ids=frozenset({15, 61}),
            ds_to_folders={
                15: frozenset({folder_key}),
                61: frozenset({folder_key}),
            },
            folder_owners={folder_key: frozenset({15, 61})},
        )
        row_15 = _row(ds_name="ДС15", quantity=1, tags="T15")
        row_61 = _row(ds_name="ДС61", quantity=1, tags="T61")
        packing = [
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                tags=["T15"],
                source=_ul_source("t15.xlsx", folder),
                excel_row=10,
            ),
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                tags=["T61"],
                source=_ul_source("t61.xlsx", folder),
                excel_row=11,
            ),
        ]
        ordered = {
            rfp_packing_match_key(15, "8950", "SOO1", "CODE"): 1.0,
            rfp_packing_match_key(61, "8950", "SOO1", "CODE"): 1.0,
        }
        with registry_packing_scope(index):
            rows, _audit = compare_rfp_rows_with_packing(
                [row_15, row_61],
                ordered,
                _dataset(packing),
            )
        self.assertEqual(row_15.get_value(UL_VALUES), 1)
        self.assertEqual(row_61.get_value(UL_VALUES), 1)
        self.assertEqual(row_15.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(row_61.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_shared_folder_one_untagged_slot_not_cloned(self) -> None:
        folder = "согл УЛ ДС15"
        folder_key = normalize_packing_key_part(folder)
        index = RegistryPackingIndex(
            known_source_ids=frozenset({15, 61}),
            ds_to_folders={
                15: frozenset({folder_key}),
                61: frozenset({folder_key}),
            },
            folder_owners={folder_key: frozenset({15, 61})},
        )
        row_15 = _row(ds_name="ДС15", quantity=1)
        row_61 = _row(ds_name="ДС61", quantity=1)
        packing = [
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=_ul_source("shared.xlsx", folder),
            )
        ]
        ordered = {
            rfp_packing_match_key(15, "8950", "SOO1", "CODE"): 1.0,
            rfp_packing_match_key(61, "8950", "SOO1", "CODE"): 1.0,
        }
        with registry_packing_scope(index):
            rows, _audit = compare_rfp_rows_with_packing(
                [row_15, row_61],
                ordered,
                _dataset(packing),
            )
        values = [row_15.get_value(UL_VALUES), row_61.get_value(UL_VALUES)]
        numeric = [float(value) for value in values if value is not None]
        self.assertEqual(len(numeric), 1)
        self.assertEqual(sum(numeric), 1.0)
        self.assertTrue(any(value is None for value in values))
        self.assertEqual(
            sum(1 for row in (row_15, row_61) if row.get_value(UL_VALUES) == 1),
            1,
        )
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_one_ds_two_registry_folders_walks_both(self) -> None:
        folder_a = "согл УЛ ДС1 ГФ 5титулов"
        folder_b = "согл УЛ ДС7 ГФ 2 тит"
        key_a = normalize_packing_key_part(folder_a)
        key_b = normalize_packing_key_part(folder_b)
        index = RegistryPackingIndex(
            known_source_ids=frozenset({8}),
            ds_to_folders={8: frozenset({key_a, key_b})},
            folder_owners={key_a: frozenset({8}), key_b: frozenset({8})},
        )
        packing = [
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=_ul_source("a.xlsx", folder_a),
                excel_row=10,
            ),
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=_ul_source("b.xlsx", folder_b),
                excel_row=11,
            ),
        ]
        rfp_one = _row(ds_name="ДС8", quantity=1)
        with registry_packing_scope(index):
            rows_one, _audit = compare_rfp_rows_with_packing(
                [rfp_one],
                {rfp_packing_match_key(8, "8950", "SOO1", "CODE"): 1.0},
                _dataset(packing),
            )
        self.assertEqual(rfp_one.get_value(UL_VALUES), 1)
        self.assertEqual(rfp_one.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        leftovers_one = _values(rows_one, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers_one), 1)

        rfp_two = _row(ds_name="ДС8", quantity=2)
        with registry_packing_scope(index):
            rows_two, _audit = compare_rfp_rows_with_packing(
                [rfp_two],
                {rfp_packing_match_key(8, "8950", "SOO1", "CODE"): 2.0},
                _dataset(packing),
            )
        self.assertEqual(rfp_two.get_value(UL_VALUES), 2)
        self.assertEqual(rfp_two.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertFalse(_values(rows_two, STATUS_PACKING_ONLY))

    def test_registry_isolation_ds29_does_not_take_ds82_folder(self) -> None:
        folder_29 = "согл УЛ ДС29"
        folder_82 = "согл УЛ ДС82"
        key_29 = normalize_packing_key_part(folder_29)
        key_82 = normalize_packing_key_part(folder_82)
        index = RegistryPackingIndex(
            known_source_ids=frozenset({29, 82}),
            ds_to_folders={
                29: frozenset({key_29}),
                82: frozenset({key_82}),
            },
            folder_owners={
                key_29: frozenset({29}),
                key_82: frozenset({82}),
            },
        )
        row_29 = _row(ds_name="ДС29", title="8445-SOT", code="BCC0003052", quantity=1)
        packing = [
            _row(
                title="8445",
                system="SOT",
                code="BCC0003052",
                quantity=1,
                source=_ul_source("file.xlsx", folder_82),
            )
        ]
        with registry_packing_scope(index):
            rows, _audit = compare_rfp_rows_with_packing(
                [row_29],
                {rfp_packing_match_key(29, "8445", "SOT", "BCC0003052"): 1.0},
                _dataset(packing),
            )
        self.assertIsNone(row_29.get_value(UL_VALUES))
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        self.assertNotEqual(row_29.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)

    def test_unknown_source_still_name_matches_own_folder(self) -> None:
        folder_92 = "согл УЛ ДС92"
        folder_15 = "согл УЛ ДС15"
        key_15 = normalize_packing_key_part(folder_15)
        index = RegistryPackingIndex(
            known_source_ids=frozenset({15}),
            ds_to_folders={15: frozenset({key_15})},
            folder_owners={key_15: frozenset({15})},
        )
        rfp_match = _row(ds_name="ДС92_24Б", title="8950-SOO1", code="CODE")
        packing_actual = _row(
            title="8950",
            system="SOO1",
            code="CODE",
            quantity=1,
            source=_ul_source("file.xlsx", folder_92),
        )
        with registry_packing_scope(index):
            compare_rfp_rows_with_packing(
                [rfp_match],
                {rfp_packing_match_key(92, "8950", "SOO1", "CODE"): 1.0},
                _dataset([packing_actual]),
            )
        self.assertEqual(rfp_match.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(rfp_match.get_value(UL_VALUES), 1.0)

    def test_known_source_empty_folders_does_not_name_match(self) -> None:
        index = RegistryPackingIndex(
            known_source_ids=frozenset({15}),
            ds_to_folders={},
        )
        rfp = _row(ds_name="ДС15", title="8950-SOO1", code="CODE", quantity=1)
        packing = _row(
            title="8950",
            system="SOO1",
            code="CODE",
            quantity=1,
            source=_ul_source("file.xlsx", "согл УЛ ДС15"),
        )
        with registry_packing_scope(index):
            rows, _audit = compare_rfp_rows_with_packing(
                [rfp],
                {rfp_packing_match_key(15, "8950", "SOO1", "CODE"): 1.0},
                _dataset([packing]),
            )
        self.assertIsNone(rfp.get_value(UL_VALUES))
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        self.assertEqual(leftovers[0].get_value(UL_COMPARE_STATUS), STATUS_PACKING_ONLY)

    def test_mixed_slash_supply_key_uses_own_folders_not_actual_fifteen(self) -> None:
        from RFQ.rfp_parts.ds_registry import (
            DsRegistryRelation,
            DsRegistryRow,
            STATUS_ACTIVE,
        )

        folder_mixed = "согл УЛ смесь 15-61"
        folder_15 = "согл УЛ ДС15"
        key_mixed = normalize_packing_key_part(folder_mixed)
        key_15 = normalize_packing_key_part(folder_15)

        class _Doc:
            def __init__(self, rows: list) -> None:
                self.active_rows = rows

        document = _Doc(
            [
                DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id="15/61",
                    relations=(
                        DsRegistryRelation(
                            group_id="ДС15/61",
                            rfp_key="99",
                            ul_folder=folder_mixed,
                        ),
                    ),
                ),
                DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id="15",
                    relations=(
                        DsRegistryRelation(
                            group_id="ДС15",
                            rfp_key="15",
                            ul_folder=folder_15,
                        ),
                    ),
                ),
            ]
        )
        index = build_registry_planting_index(document)
        self.assertEqual(index.mixed_ds_to_folders["15/61"], frozenset({key_mixed}))
        self.assertEqual(index.ds_to_folders[15], frozenset({key_15}))
        self.assertNotIn(99, index.rfp_number_to_actual)

        mixed_rfp = _row(ds_name="ДС15/61", quantity=1)
        simple_rfp = _row(ds_name="ДС15", quantity=1)
        underscore = _row(ds_name="ДС15_61", quantity=1)
        self.assertIsNone(rfp_row_actual_ds(mixed_rfp))
        self.assertEqual(rfp_row_supply_key(mixed_rfp), "15/61")
        self.assertEqual(rfp_row_actual_ds(underscore), 15)
        self.assertEqual(rfp_row_supply_key(underscore), "15")

        snapshot = build_ordered_snapshot([mixed_rfp, simple_rfp])
        mixed_key = rfp_packing_match_key("15/61", "8950", "SOO1", "CODE")
        simple_key = rfp_packing_match_key(15, "8950", "SOO1", "CODE")
        self.assertEqual(snapshot[mixed_key], 1.0)
        self.assertEqual(snapshot[simple_key], 1.0)
        self.assertNotEqual(mixed_key, simple_key)

        packing = [
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=_ul_source("mixed.xlsx", folder_mixed),
                excel_row=10,
            ),
            _row(
                title="8950",
                system="SOO1",
                quantity=1,
                source=_ul_source("only15.xlsx", folder_15),
                excel_row=11,
            ),
        ]
        with registry_packing_scope(index):
            compare_rfp_rows_with_packing(
                [mixed_rfp],
                {mixed_key: 1.0},
                _dataset(packing),
            )
        self.assertEqual(mixed_rfp.get_value(UL_VALUES), 1.0)
        self.assertEqual(mixed_rfp.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertIn("смесь", str(mixed_rfp.get_value(UL_SOURCE_FILES)))
        self.assertNotIn("only15", str(mixed_rfp.get_value(UL_SOURCE_FILES)))
        self.assertEqual(mixed_rfp.get_value(DS_ACTUAL), "ДС15/61")

        empty_mixed = _row(ds_name="ДС15/61", quantity=1)
        empty_index = RegistryPackingIndex(
            known_source_ids=frozenset({15}),
            ds_to_folders={15: frozenset({key_15})},
        )
        with registry_packing_scope(empty_index):
            rows, _audit = compare_rfp_rows_with_packing(
                [empty_mixed],
                {mixed_key: 1.0},
                _dataset(
                    [
                        _row(
                            title="8950",
                            system="SOO1",
                            quantity=1,
                            source=_ul_source("only15.xlsx", folder_15),
                        )
                    ]
                ),
            )
        self.assertIsNone(empty_mixed.get_value(UL_VALUES))
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)

        underscore_rfp = _row(ds_name="ДС15_61", quantity=1)
        with registry_packing_scope(index):
            compare_rfp_rows_with_packing(
                [underscore_rfp],
                {simple_key: 1.0},
                _dataset(
                    [
                        _row(
                            title="8950",
                            system="SOO1",
                            quantity=1,
                            source=_ul_source("only15.xlsx", folder_15),
                        )
                    ]
                ),
            )
        self.assertEqual(underscore_rfp.get_value(UL_VALUES), 1.0)
        self.assertEqual(underscore_rfp.get_value(DS_ACTUAL), "ДС15")

    def test_mixed_manager_same_surname_and_conflict_comment(self) -> None:
        from openpyxl import Workbook

        from RFQ.tags_rfp_compare.ds_manager_matrix import (
            apply_ds_source_columns,
            load_ds_manager_matrix,
        )

        matrix_path = Path(tempfile.mkdtemp()) / "mp.xlsx"
        wb = Workbook()
        sheet = wb.active
        sheet.title = "Лист2"
        sheet.append(["Имя ДС", "Фамилия"])
        sheet.append(["ДС15", "Иванов"])
        sheet.append(["ДС61", "Иванов"])
        sheet.append(["ДС13", "Петров"])
        sheet.append(["ДС47", "Сидоров"])
        wb.save(matrix_path)
        wb.close()
        matrix = load_ds_manager_matrix(matrix_path)

        same = _row(ds_name="ДС15/61", quantity=1)
        same.el[DS_ACTUAL].value = ""
        apply_ds_source_columns(
            [same],
            matrix=matrix,
            rfp_paths_by_key={},
            mto_paths_by_title={},
        )
        self.assertEqual(same.get_value(DS_ACTUAL), "ДС15/61")
        self.assertEqual(same.get_value(DS_MANAGER), "Иванов")

        conflict = _row(ds_name="ДС13/47", quantity=1)
        conflict.el[DS_ACTUAL].value = ""
        apply_ds_source_columns(
            [conflict],
            matrix=matrix,
            rfp_paths_by_key={},
            mto_paths_by_title={},
        )
        self.assertEqual(conflict.get_value(DS_ACTUAL), "ДС13/47")
        self.assertEqual(conflict.get_value(DS_MANAGER), "")
        self.assertEqual(conflict.el[DS_ACTUAL].comment, "смешанный ДС13/47")

    def test_gf_folder_without_actual_does_not_land(self) -> None:
        rfp = _row(ds_name="ДС1", title="8950-SOO1", code="CODE", quantity=1)
        packing = _row(
            title="8950",
            system="SOO1",
            code="CODE",
            quantity=1,
            source=_ul_source("file.xlsx", "согл УЛ ГФ 5титулов"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [rfp],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
        )
        self.assertIsNone(rfp.get_value(UL_VALUES))
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        self.assertEqual(leftovers[0].get_value(UL_COMPARE_STATUS), STATUS_PACKING_ONLY)
        self.assertEqual(leftovers[0].get_value(DS_NAME), "RFP_не_найден")
        self.assertEqual(leftovers[0].get_value(DS_ACTUAL), "")

    def test_packing_only_leftover_unique_matrix_manager(self) -> None:
        from openpyxl import Workbook

        from RFQ.tags_rfp_compare.ds_manager_matrix import (
            apply_ds_source_columns,
            load_ds_manager_matrix,
        )

        rfp = _row(ds_name="ДС99", title="8950-SOO1", code="CODE", quantity=1)
        packing = _row(
            title="8950",
            system="SOO1",
            code="CODE",
            quantity=1,
            source=_ul_source("file.xlsx", "согл УЛ ДС01"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [rfp],
            {rfp_packing_match_key(99, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
        )
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        leftover = leftovers[0]
        self.assertEqual(leftover.get_value(DS_NAME), "ДС1_RFP_не_найден")
        self.assertEqual(leftover.get_value(DS_ACTUAL), "ДС1")

        with tempfile.TemporaryDirectory() as tmpdir:
            matrix_path = Path(tmpdir) / "matrix.xlsx"
            wb = Workbook()
            positions = wb.active
            positions.title = "Лист1"
            positions.append(["Имя ДС", "№ позиции", "Титул/Марка", "Код RFP", "Наименование RFP"])
            positions.append(["ДС6", "1", "8310-SOS", "BCC0000341", "dummy"])
            managers = wb.create_sheet("Лист2")
            managers.append(["Имя ДС", "Кол-во RFP", "Фамилия"])
            managers.append(["ДС01", None, "ГЭМ"])
            wb.save(matrix_path)
            wb.close()
            matrix = load_ds_manager_matrix(matrix_path)
            apply_ds_source_columns([leftover], matrix=matrix, rfp_paths_by_key={}, mto_paths_by_title={})
        self.assertEqual(leftover.get_value(DS_MANAGER), "ГЭМ")

    def test_packing_only_leftover_ds92_identity_columns(self) -> None:
        rfp = _row(ds_name="ДС92_24Б", title="8950-SOO1", code="CODE", quantity=1)
        packing = _row(
            title="8950",
            system="SOO1",
            code="OTHER",
            quantity=1,
            source=_ul_source("file.xlsx", "согл УЛ ДС92"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [rfp],
            {rfp_packing_match_key(92, "8950", "SOO1", "CODE"): 1.0},
            _dataset([packing]),
        )
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        self.assertEqual(leftovers[0].get_value(DS_NAME), "ДС92_RFP_не_найден")
        self.assertEqual(leftovers[0].get_value(DS_ACTUAL), "ДС92")

    def test_incident_mto_only_empty_ds_consumes_leftover(self) -> None:
        mto = _row(
            title="2319-KSB",
            code="",
            quantity="",
            ds_name="",
            code_mto="BCC0000538",
            values_mto=2,
            match_status="Добавлен из МТО",
        )
        packing = _row(
            title="2319",
            system="KSB",
            code="BCC0000538",
            quantity=2,
            source=_ul_source("PL.xlsx", "согл УЛ ДС71"),
        )
        rows, audit = compare_rfp_rows_with_packing(
            [mto],
            {},
            _dataset([packing]),
        )
        self.assertEqual(mto.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)
        self.assertEqual(mto.el[UL_COMPARE_STATUS].color, Color.match_matched)
        self.assertEqual(mto.get_value(UL_VALUES), 2.0)
        self.assertEqual(mto.get_value(DS_ACTUAL), "ДС71")
        self.assertNotIn("RFP_не_найден", str(mto.get_value(DS_NAME) or ""))
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))
        self.assertEqual(audit.stats.allocated_units, 2)

    def test_mto_only_without_ul_stays_mto_only(self) -> None:
        mto = _row(
            title="2319-KSB",
            code="",
            quantity="",
            ds_name="",
            code_mto="BCC0000538",
            values_mto=2,
            match_status="Добавлен из МТО",
        )
        compare_rfp_rows_with_packing([mto], {}, _dataset([]))
        self.assertEqual(mto.get_value(UL_COMPARE_STATUS), STATUS_MTO_ONLY)
        self.assertIsNone(mto.get_value(UL_VALUES))
        self.assertFalse(mto.el[DS_ACTUAL].comment)

    def test_leftover_other_ds_sits_on_mto_not_rfp_8445(self) -> None:
        rfp_ds29 = _row(
            ds_name="ДС29",
            title="8445-SOT",
            code="BCC0003052",
            quantity=1,
        )
        mto_only = _row(
            ds_name="",
            title="8445-SOT",
            code="",
            quantity="",
            code_mto="BCC0003052",
            values_mto=1,
            match_status="Добавлен из МТО",
        )
        packing = [
            _row(
                title="8445",
                system="SOT",
                code="BCC0003052",
                quantity=1,
                source=_ul_source("ds29.xlsx", "согл УЛ ДС29"),
                excel_row=10,
            ),
            _row(
                title="8445",
                system="SOT",
                code="BCC0003052",
                quantity=1,
                source=_ul_source("ds82.xlsx", "согл УЛ ДС82"),
                excel_row=11,
            ),
        ]
        ordered = {rfp_packing_match_key(29, "8445", "SOT", "BCC0003052"): 1.0}
        rows, _audit = compare_rfp_rows_with_packing(
            [rfp_ds29, mto_only],
            ordered,
            _dataset(packing),
        )
        self.assertEqual(rfp_ds29.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(rfp_ds29.get_value(UL_VALUES), 1.0)
        self.assertEqual(mto_only.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)
        self.assertEqual(mto_only.get_value(UL_VALUES), 1.0)
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_vo_only_empty_ds_consumes_leftover(self) -> None:
        vo = _row(
            code="",
            quantity="",
            ds_name="",
            code_vo="VO-OK",
            values_vo=1,
            match_status_vo="Добавлен из VO",
            tag_vo="VO-TAG",
        )
        packing = _row(
            title="8950",
            system="SOO1",
            code="VO-OK",
            quantity=1,
            tags=["VO-TAG"],
            source=_ul_source("vo.xlsx", "согл УЛ ДС5"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [vo],
            {},
            _dataset([packing]),
        )
        self.assertEqual(vo.get_value(UL_COMPARE_STATUS), STATUS_VO_DELIVERED)
        self.assertEqual(vo.el[UL_COMPARE_STATUS].color, Color.match_matched)
        self.assertEqual(vo.get_value(UL_VALUES), 1.0)
        self.assertEqual(vo.get_value(DS_ACTUAL), "ДС5")
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_mto_plus_vo_same_row_is_block_b(self) -> None:
        row = _row(
            code="",
            quantity="",
            ds_name="",
            code_mto="MTO-VO",
            values_mto=1,
            match_status="Добавлен из МТО",
            match_status_vo="Тег сопоставлен",
            code_vo="VO-SIDE",
            values_vo=1,
            tag_vo="KEEP-VO",
        )
        packing = _row(
            title="8950",
            system="SOO1",
            code="MTO-VO",
            quantity=1,
            source=_ul_source("b.xlsx", "согл УЛ ДС8"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [row],
            {},
            _dataset([packing]),
        )
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)
        self.assertEqual(row.get_value(CODE_VO), "VO-SIDE")
        self.assertEqual(row.get_value(TAG_VO), "KEEP-VO")
        self.assertEqual(row.get_value(MATCH_STATUS_VO), "Тег сопоставлен")
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_mto_only_partial_leftover_is_short_no_dump(self) -> None:
        mto = _row(
            code="",
            quantity="",
            ds_name="",
            code_mto="MTO-PART",
            values_mto=2,
            match_status="Добавлен из МТО",
        )
        packing = _row(
            title="8950",
            system="SOO1",
            code="MTO-PART",
            quantity=1,
            source=_ul_source("part.xlsx", "согл УЛ ДС3"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [mto],
            {},
            _dataset([packing]),
        )
        self.assertEqual(mto.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED_SHORT)
        self.assertEqual(mto.el[UL_COMPARE_STATUS].color, Color.yellow)
        self.assertEqual(mto.get_value(UL_VALUES), 1.0)
        self.assertFalse(_values(rows, STATUS_PACKING_ONLY))

    def test_mto_only_excess_leftover_dumps_tail(self) -> None:
        mto = _row(
            code="",
            quantity="",
            ds_name="",
            code_mto="MTO-EX",
            values_mto=2,
            match_status="Добавлен из МТО",
        )
        packing = _row(
            title="8950",
            system="SOO1",
            code="MTO-EX",
            quantity=3,
            source=_ul_source("ex.xlsx", "согл УЛ ДС4"),
        )
        rows, _audit = compare_rfp_rows_with_packing(
            [mto],
            {},
            _dataset([packing]),
        )
        self.assertEqual(mto.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)
        self.assertEqual(mto.get_value(UL_VALUES), 2.0)
        leftovers = _values(rows, STATUS_PACKING_ONLY)
        self.assertEqual(len(leftovers), 1)
        self.assertEqual(leftovers[0].get_value(UL_VALUES), 1.0)

    def test_two_mto_only_same_code_keep_own_tags_in_b1(self) -> None:
        first = _row(
            code="",
            quantity="",
            ds_name="",
            code_mto="SHARE",
            values_mto=1,
            tag_mto="TAG-A",
            match_status="Добавлен из МТО",
            ds_number="1",
        )
        second = _row(
            code="",
            quantity="",
            ds_name="",
            code_mto="SHARE",
            values_mto=1,
            tag_mto="TAG-B",
            match_status="Добавлен из МТО",
            ds_number="2",
        )
        packing = [
            _row(
                title="8950",
                system="SOO1",
                code="SHARE",
                quantity=1,
                tags=["TAG-B"],
                source=_ul_source("b.xlsx", "согл УЛ ДС9"),
                excel_row=20,
            ),
            _row(
                title="8950",
                system="SOO1",
                code="SHARE",
                quantity=1,
                tags=["TAG-A"],
                source=_ul_source("a.xlsx", "согл УЛ ДС9"),
                excel_row=21,
            ),
        ]
        compare_rfp_rows_with_packing(
            [first, second],
            {},
            _dataset(packing),
        )
        self.assertEqual(first.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)
        self.assertEqual(second.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)
        self.assertEqual(first.get_value(UL_TAGS), "TAG-A")
        self.assertEqual(second.get_value(UL_TAGS), "TAG-B")
        self.assertNotEqual(
            first.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MISMATCH_COLLECTED,
        )
        self.assertNotEqual(
            second.get_value(UL_TAG_MATCH_STATUS),
            STATUS_TAG_MISMATCH_COLLECTED,
        )

    def test_running_delivered_matches_full_sum(self) -> None:
        source_row = _row()
        key = ("1", "8445", "sot", "code")
        state = _RowAllocationState(
            row=_row(),
            key=key,
            fallback="",
            ordered=4.0,
        )

        def unit(qty: float) -> _PackingUnit:
            return _PackingUnit(
                key=key,
                source_row=source_row,
                tag="",
                units="шт",
                source="test",
                qty=qty,
            )

        takes = 0
        for qty in (1.0, 1.0, 0.3, 0.7):
            state.take(unit(qty))
            takes += 1
            self.assertEqual(
                state.remaining_qty(),
                state.ordered - _delivered_qty(state.allocated),
            )
        self.assertEqual(state.delivered, _delivered_qty(state.allocated))
        self.assertEqual(len(state.allocated), takes)


class RfpPackingIntegrationSmokeTest(unittest.TestCase):
    def test_existing_rfp_collapse_characterization_is_unchanged(self) -> None:
        untagged_a = _row(quantity=1, ds_number="10")
        untagged_b = _row(quantity=2, ds_number="10")
        tagged_a = _row(quantity=1, tags="T1", ds_number="11")
        tagged_b = _row(quantity=1, tags="T2", ds_number="11")
        collapsed = _collapse_rfp_rows_before_export(
            [untagged_a, untagged_b, tagged_a, tagged_b]
        )
        self.assertEqual(len(collapsed), 3)
        self.assertEqual(collapsed[0].get_value(VALUES), 3.0)
        self.assertIs(collapsed[1], tagged_a)
        self.assertIs(collapsed[2], tagged_b)

    def test_ul_check_elements_are_independent_in_light_copies(self) -> None:
        source = _row()
        first = RowStd.get_row_copy_light(source)
        second = RowStd.get_row_copy_light(source)
        batch_first, batch_second = RowStd.batch_copy_light(source, 2)

        for left, right in ((first, second), (batch_first, batch_second)):
            for column in UL_COLUMNS + UNITS_COLUMNS:
                self.assertIsNot(left.el[column], right.el[column], column)
            right_value = right.el[UL_VALUES].value
            left.el[UL_VALUES].value = 99
            left.el[UL_VALUES].color = Color.red
            left.el[UL_VALUES].comment = "left only"
            self.assertEqual(right.el[UL_VALUES].value, right_value)
            self.assertNotEqual(right.el[UL_VALUES].value, 99)
            self.assertEqual(right.el[UL_VALUES].color, Color.no)
            self.assertEqual(right.el[UL_VALUES].comment, "")

    def test_packing_units_trace_merge_order_and_dedupe(self) -> None:
        source_row = _row()
        key = ("1", "8950", "soo1", "CODE")

        def unit(*, status: str = "", trace: str = "") -> _PackingUnit:
            return _PackingUnit(
                key=key,
                source_row=source_row,
                tag="",
                units="шт",
                source="test",
                units_check_status=status,
                units_conversion_trace=trace,
            )

        ordered = _row()
        ordered.el[UNITS_CHECK_STATUS].value = "identity"
        ordered.el[UNITS_CONVERSION_TRACE].value = "rfp-trace"
        _merge_packing_units_trace(
            ordered,
            [
                unit(status="converted", trace="ul-trace-a"),
                unit(status="identity", trace="ul-trace-b"),
                unit(status="converted", trace="ul-trace-a"),
            ],
        )
        self.assertEqual(ordered.get_value(UNITS_CHECK_STATUS), "identity; converted")
        self.assertEqual(
            ordered.get_value(UNITS_CONVERSION_TRACE),
            "rfp-trace; ul-trace-a; ul-trace-b",
        )

        empty_target = _row()
        empty_target.el[UNITS_CHECK_STATUS].value = ""
        empty_target.el[UNITS_CONVERSION_TRACE].value = ""
        _merge_packing_units_trace(empty_target, [unit(), unit(status="", trace="")])
        self.assertEqual(empty_target.el[UNITS_CHECK_STATUS].value, "")
        self.assertEqual(empty_target.el[UNITS_CONVERSION_TRACE].value, "")
        self.assertIsNotNone(empty_target.el[UNITS_CHECK_STATUS].value)
        self.assertIsNotNone(empty_target.el[UNITS_CONVERSION_TRACE].value)

        trace_only = _row()
        trace_only.el[UNITS_CHECK_STATUS].value = "identity"
        trace_only.el[UNITS_CONVERSION_TRACE].value = None
        _merge_packing_units_trace(
            trace_only,
            [unit(status="", trace="ul-only-trace")],
        )
        self.assertEqual(trace_only.get_value(UNITS_CHECK_STATUS), "identity")
        self.assertEqual(trace_only.get_value(UNITS_CONVERSION_TRACE), "ul-only-trace")

    def test_mto_units_and_trace_transfer(self) -> None:
        rfp_row = _row(units="шт", quantity=2)
        rfp_row.el[UNITS_CHECK_STATUS].value = "RFP status"
        rfp_row.el[UNITS_CONVERSION_TRACE].value = "RFP trace"
        mto_row = _row(units="шт.", quantity=3)
        mto_row.el[UNITS_CHECK_STATUS].value = "MTO status"
        mto_row.el[UNITS_CONVERSION_TRACE].value = "MTO trace"
        add_mto_data_to_rfp_row(rfp_row, mto_row)
        self.assertEqual(rfp_row.get_value(UNITS_MTO), "шт")
        self.assertIn("RFP status", rfp_row.get_value(UNITS_CHECK_STATUS))
        self.assertIn("MTO status", rfp_row.get_value(UNITS_CHECK_STATUS))
        self.assertIn("RFP trace", rfp_row.get_value(UNITS_CONVERSION_TRACE))
        self.assertIn("MTO trace", rfp_row.get_value(UNITS_CONVERSION_TRACE))

    def test_collapse_different_units_are_not_merged(self) -> None:
        row_a = _row(quantity=1, units="шт", ds_number="10")
        row_a.el[UNITS_MTO].value = "шт"
        row_b = _row(quantity=2, units="шт", ds_number="10")
        row_b.el[UNITS_MTO].value = "м"
        collapsed = _collapse_rfp_rows_before_export([row_a, row_b])
        self.assertEqual(len(collapsed), 2)

    def test_mto_rfp_diff_numeric_and_unit_mismatch(self) -> None:
        matched = _row(units="шт", quantity=2)
        matched.el[UNITS_MTO].value = "шт"
        matched.el[VALUES_MTO].value = 2
        diff_row = _row(units="шт", quantity=1)
        diff_row.el[UNITS_MTO].value = "шт"
        diff_row.el[VALUES_MTO].value = 3
        mismatch = _row(units="шт", quantity=2)
        mismatch.el[UNITS_MTO].value = "м"
        mismatch.el[VALUES_MTO].value = 2
        apply_mto_rfp_quantity_diff([matched, diff_row, mismatch])
        self.assertEqual(matched.get_value(VALUES_MTO_RFP_DIFF), 0)
        self.assertEqual(matched.el[VALUES_MTO_RFP_DIFF].color, Color.match_matched)
        self.assertEqual(diff_row.get_value(VALUES_MTO_RFP_DIFF), 2)
        self.assertEqual(diff_row.el[VALUES_MTO_RFP_DIFF].color, Color.yellow)
        self.assertEqual(mismatch.get_value(VALUES_MTO_RFP_DIFF), "")
        self.assertEqual(mismatch.el[UNITS].color, Color.yellow)
        self.assertEqual(mismatch.el[UNITS_MTO].color, Color.yellow)
        self.assertIn("Несовместимые", mismatch.el[UNITS].comment)

    def test_mto_ul_diff_numeric_empty_and_unit_mismatch(self) -> None:
        matched = _row(units="шт", quantity=2, values_mto=2)
        matched.el[UNITS_MTO].value = "шт"
        matched.el[UL_UNITS].value = "шт"
        matched.el[UL_VALUES].value = 2
        shortfall = _row(units="шт", quantity=1, values_mto=3)
        shortfall.el[UNITS_MTO].value = "шт"
        shortfall.el[UL_UNITS].value = "шт"
        shortfall.el[UL_VALUES].value = 1
        empty_ul = _row(units="шт", quantity=1, values_mto=2)
        empty_ul.el[UNITS_MTO].value = "шт"
        mismatch = _row(units="шт", quantity=2, values_mto=2)
        mismatch.el[UNITS_MTO].value = "шт"
        mismatch.el[UL_UNITS].value = "м"
        mismatch.el[UL_VALUES].value = 2
        apply_mto_ul_quantity_diff([matched, shortfall, empty_ul, mismatch])
        self.assertEqual(matched.get_value(VALUES_MTO_UL_DIFF), 0)
        self.assertEqual(matched.el[VALUES_MTO_UL_DIFF].color, Color.match_matched)
        self.assertEqual(shortfall.get_value(VALUES_MTO_UL_DIFF), 2)
        self.assertEqual(shortfall.el[VALUES_MTO_UL_DIFF].color, Color.yellow)
        self.assertEqual(empty_ul.get_value(VALUES_MTO_UL_DIFF), "")
        self.assertEqual(mismatch.get_value(VALUES_MTO_UL_DIFF), "")
        self.assertEqual(mismatch.el[UL_UNITS].color, Color.yellow)
        self.assertIn("Несовместимые", mismatch.el[UL_UNITS].comment)

    def test_mto_rfp_diff_fills_empty_mto_units_from_rfp(self) -> None:
        row = _row(units="шт", quantity=1)
        row.el[UNITS_MTO].value = None
        row.el[VALUES_MTO].value = 1
        row.el[UNITS_CHECK_STATUS].value = "identity"
        row.el[UNITS_CONVERSION_TRACE].value = (
            "code=BCC000181; src=шт; tgt=шт; qty=1; coef=1; result=1; status=identity"
        )
        apply_mto_rfp_quantity_diff([row])
        self.assertEqual(row.get_value(UNITS_MTO), "шт")
        self.assertEqual(row.get_value(VALUES_MTO_RFP_DIFF), 0)
        self.assertEqual(row.el[VALUES_MTO_RFP_DIFF].color, Color.match_matched)
        self.assertFalse(str(row.el[UNITS_MTO].comment or "").strip())

    def test_converted_units_comment_only_for_matrix_conversion(self) -> None:
        identity = converted_units_comment_text(
            "identity",
            "code=BCC1; src=шт; tgt=шт; qty=1; coef=1; result=1; status=identity",
        )
        spelling = converted_units_comment_text(
            "converted",
            "code=BCC1; src=шт.; tgt=шт; qty=1; coef=1; result=1; status=converted",
        )
        converted = converted_units_comment_text(
            "identity; converted",
            "code=BCC1; src=шт; tgt=шт; qty=1; coef=1; result=1; status=identity; "
            "code=BCC1; src=м; tgt=шт; qty=10; coef=0.001; result=0.01; status=converted",
        )
        self.assertEqual(identity, "")
        self.assertEqual(spelling, "")
        self.assertIn("Конвертация по матрице", converted)
        self.assertIn("10 м × 0.001 → 0.01 шт", converted)
        self.assertNotIn("status=identity", converted)
        row = _row(units="шт", quantity=1)
        row.el[UNITS_MTO].value = "шт"
        row.el[VALUES_MTO].value = 1
        row.el[UNITS_CHECK_STATUS].value = "converted"
        row.el[UNITS_CONVERSION_TRACE].value = (
            "code=BCC1; src=м; tgt=шт; qty=10; coef=0.001; result=0.01; status=converted"
        )
        apply_mto_rfp_quantity_diff([row])
        self.assertIn("Конвертация по матрице", row.el[UNITS_MTO].comment)
        self.assertEqual(row.get_value(VALUES_MTO_RFP_DIFF), 0)

    def test_copy_mto_el_keeps_units_fields(self) -> None:
        src = _row(units="шт", quantity=1)
        src.el[CODE_MTO].value = "BCC1"
        src.el[VALUES_MTO].value = 2
        src.el[UNITS_MTO].value = "шт"
        src.el[UNITS_CHECK_STATUS].value = "converted"
        src.el[UNITS_CONVERSION_TRACE].value = "code=BCC1; status=converted"
        dest = _row(units="шт", quantity=1)
        dest.el.pop(UNITS_MTO, None)
        _copy_mto_el_to_row(src, dest)
        self.assertEqual(dest.get_value(UNITS_MTO), "шт")
        self.assertEqual(dest.get_value(VALUES_MTO), 2)
        self.assertEqual(dest.get_value(UNITS_CHECK_STATUS), "converted")

    def test_excel_save_survives_empty_units_trace(self) -> None:
        empty = _row(code="EMPTY-TRACE", quantity=1, units="шт")
        missing = _row(code="NONE-TRACE", quantity=1, units="шт")
        missing.el[UNITS_CONVERSION_TRACE].value = None
        with tempfile.TemporaryDirectory() as tmp:
            path = save_match_result_to_excel([empty, missing], tmp)
            self.assertTrue(path)
            self.assertTrue(Path(path).is_file())

    def test_excel_ul_headers_color_numeric_source_and_comments(self) -> None:
        target = _row(code="XLSX", quantity=2, units="шт")
        target.el[UNITS_MTO].value = "шт"
        target.el[VALUES_MTO].value = 2
        target.el[UNITS_CHECK_STATUS].value = "OK"
        target.el[UNITS_CONVERSION_TRACE].value = "trace line"
        apply_mto_rfp_quantity_diff([target])
        target.el[PATH_RFP].value = (
            r"\\bcc\eng\RFP_Зиновьев\ДС75_16А. "
            r"AGCC.287-0000-12.4.1-RFP-0010_01_RU.xlsx"
        )
        target.el[PATH_MTO].value = (
            r"\\bcc\eng\PrDoc\РД\6550\SOT"
            r"\AGCC.287-6550-SOT.MTO-0001_0-AN01_RU.xlsx"
        )
        packing = _row(
            title="8950",
            system="SOO1",
            code="XLSX",
            quantity=2,
            source=_ul_source("source.xlsx"),
            sheet="Packing",
            excel_row=77,
        )
        rows, audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "XLSX"): 2.0},
            _dataset([packing]),
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = save_match_result_to_excel(rows, tmp, packing_audit=audit)
            self.assertIsNotNone(path)
            workbook = openpyxl.load_workbook(
                io.BytesIO(Path(path).read_bytes()),
                data_only=True,
            )
            try:
                self.assertIn("Сводка", workbook.sheetnames)
                self.assertIn("Статистика ДС", workbook.sheetnames)
                self.assertIn("Посадка", workbook.sheetnames)
                self.assertIn("Статусы", workbook.sheetnames)
                summary = workbook["Сводка"]
                metrics = {
                    summary.cell(row_idx, 1).value: summary.cell(row_idx, 2).value
                    for row_idx in range(2, summary.max_row + 1)
                    if summary.cell(row_idx, 1).value
                }
                self.assertEqual(metrics["Поставка комплектна"], 1)
                landing = workbook["Посадка"]
                self.assertEqual(landing.cell(2, 1).value, "RFP")
                sheet = workbook["Сопоставление RFP и MTO"]
                headers = {
                    cell.value: cell.column for cell in sheet[1] if cell.value is not None
                }
                header_labels = list(headers.keys())
                self.assertLess(
                    header_labels.index("Ед. изм. RFP"),
                    header_labels.index("Тег (VO/MTO/RFP)"),
                )
                self.assertEqual(
                    header_labels[
                        header_labels.index("Количество MTO") + 1 :
                        header_labels.index("Номера MTO") + 1
                    ],
                    ["Ед. изм. MTO", "Номера MTO"],
                )
                self.assertEqual(
                    header_labels[header_labels.index("МТО − УЛ") - 1],
                    "МТО − RFP",
                )
                self.assertEqual(
                    header_labels[header_labels.index("Кол-во по УЛ") - 1],
                    "МТО − УЛ",
                )
                self.assertEqual(
                    header_labels[header_labels.index("Ед. изм. RFP") + 1],
                    "Теги RFP",
                )
                self.assertEqual(
                    header_labels[header_labels.index("Теги RFP") + 1],
                    "Тег (VO/MTO/RFP)",
                )
                self.assertNotIn("Лот", headers)
                self.assertNotIn("Качество данных УЛ", headers)
                self.assertNotIn("Статус ед. изм.", headers)
                self.assertNotIn("Trace ед. изм.", headers)
                self.assertEqual(sheet.cell(2, headers["Ед. изм. MTO"]).value, "шт")
                self.assertIsNone(sheet.cell(2, headers["Ед. изм. MTO"]).comment)
                expected_headers = {
                    "Фактический ДС",
                    "Ед. изм. RFP",
                    "Теги RFP",
                    "МТО − RFP",
                    "МТО − УЛ",
                    "Кол-во по УЛ",
                    "Ед. изм. УЛ",
                    "RFP − УЛ",
                    "Статус УЛ",
                    "Статус тегов УЛ",
                    "Тег УЛ",
                    "Наименование УЛ",
                    "Код УЛ",
                    "Источник УЛ (файл · вкладка · строка)",
                    "Тип марки УЛ",
                    "Поставщик УЛ",
                }
                self.assertTrue(expected_headers.issubset(headers))
                self.assertIn("Порядковый ДС", headers)
                self.assertNotIn("Имя ДС", headers)
                self.assertEqual(
                    header_labels.index("Фактический ДС"),
                    header_labels.index("Порядковый ДС") + 1,
                )
                self.assertEqual(
                    header_labels.index("№ позиции"),
                    header_labels.index("Фактический ДС") + 1,
                )
                self.assertEqual(
                    header_labels.index("Менеджер ДС"),
                    header_labels.index("№ позиции") + 1,
                )
                self.assertEqual(sheet.cell(2, headers["Порядковый ДС"]).value, "ДС1")
                self.assertEqual(sheet.cell(2, headers["Фактический ДС"]).value, "ДС1")
                self.assertEqual(target.get_value(DS_NAME), "ДС1")
                self.assertEqual(target.get_value(DS_ACTUAL), "ДС1")
                self.assertNotIn("Теги УЛ", headers)
                self.assertEqual(
                    header_labels[header_labels.index("Статус УЛ") + 1],
                    "Статус тегов УЛ",
                )
                self.assertEqual(
                    header_labels[header_labels.index("Статус тегов УЛ") + 1],
                    "Тег УЛ",
                )
                self.assertEqual(
                    header_labels[header_labels.index("Тег УЛ") + 1],
                    "Наименование УЛ",
                )
                self.assertEqual(
                    header_labels.index("Менеджер ДС"),
                    header_labels.index("№ позиции") + 1,
                )
                self.assertEqual(header_labels[-2], "Путь к RFP")
                self.assertEqual(header_labels[-1], "Путь к МТО")
                self.assertEqual(
                    header_labels.index("Путь к МТО"),
                    header_labels.index("Путь к RFP") + 1,
                )
                self.assertNotIn("Заказано по RFP", headers)
                for label in ("Кол-во RFP", "Кол-во по УЛ", "RFP − УЛ", "МТО − RFP"):
                    self.assertIsInstance(sheet.cell(2, headers[label]).value, (int, float))
                status = sheet.cell(2, headers["Статус УЛ"])
                self.assertEqual(status.value, STATUS_COMPLETE)
                self.assertEqual(status.fill.fgColor.rgb[-6:].lower(), Color.green.lower())
                self.assertEqual(
                    sheet.cell(1, headers["Статус УЛ"]).fill.fgColor.rgb[-6:].lower(),
                    "c6efce",
                )
                self.assertEqual(
                    sheet.cell(1, headers["Ед. изм. УЛ"]).fill.fgColor.rgb[-6:].lower(),
                    "c6efce",
                )
                source = sheet.cell(
                    2,
                    headers["Источник УЛ (файл · вкладка · строка)"],
                )
                self.assertEqual(
                    str(source.value).replace("\\", "/"),
                    "согл УЛ ДС1/source.xlsx · Packing · строка 77",
                )
                self.assertIsNotNone(source.hyperlink)
                self.assertIn("source.xlsx", str(source.hyperlink.target))
                self.assertIn("Packing", str(source.hyperlink.location or ""))
                self.assertIn("A77", str(source.hyperlink.location or ""))
                path_rfp = sheet.cell(2, headers["Путь к RFP"])
                self.assertIn("RFP_Зиновьев", str(path_rfp.value))
                self.assertIn(
                    "ДС75_16А. AGCC.287-0000-12.4.1-RFP-0010_01_RU.xlsx",
                    str(path_rfp.value),
                )
                self.assertIsNone(path_rfp.hyperlink)
                path_mto = sheet.cell(2, headers["Путь к МТО"])
                self.assertIn(
                    r"6550\SOT\AGCC.287-6550-SOT.MTO-0001_0-AN01_RU.xlsx",
                    str(path_mto.value).replace("/", "\\"),
                )
                self.assertIsNone(path_mto.hyperlink)
                self.assertIsNone(sheet.cell(2, headers["Статус УЛ"]).comment)
                self.assertIn("source.xlsx", target.el[UL_COMPARE_STATUS].comment)
            finally:
                workbook.close()

    def test_should_write_excel_comment_ul_policy(self) -> None:
        row = _row(code="POLICY")
        row.el[UL_COMPARE_STATUS].value = STATUS_COMPLETE
        row.el[UL_COMPARE_STATUS].comment = "Ключ УЛ: ('8950', 'soo1', 'POLICY')"
        self.assertFalse(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_SHORTFALL
        self.assertTrue(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_DATA_PROBLEM
        self.assertTrue(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_COMPLETE
        row.el[UL_COMPARE_STATUS].comment = (
            f"{FALLBACK_COMMENT_MARKER} CODE_MTO='MTO-OK'"
        )
        self.assertTrue(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_PACKING_ONLY
        row.el[UL_COMPARE_STATUS].comment = "Ключ УЛ: leftover"
        self.assertFalse(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_MTO_DELIVERED
        row.el[UL_COMPARE_STATUS].comment = "Ключ УЛ leftover MTO"
        self.assertFalse(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_MTO_DELIVERED_SHORT
        self.assertTrue(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_VO_DELIVERED_SHORT
        self.assertTrue(
            should_write_excel_comment(
                UL_COMPARE_STATUS, row, row.el[UL_COMPARE_STATUS].comment
            )
        )
        row.el[UNITS_MTO].comment = "Конвертация по матрице:\n1 шт × 1 → 1 м"
        self.assertTrue(
            should_write_excel_comment(UNITS_MTO, row, row.el[UNITS_MTO].comment)
        )
        self.assertFalse(should_write_excel_comment(UL_COMPARE_STATUS, row, ""))
        self.assertFalse(should_write_excel_comment(UL_COMPARE_STATUS, row, None))

    def test_excel_ul_comments_only_for_problem_and_fallback(self) -> None:
        complete = _row(code="OK", quantity=1)
        short = _row(code="SHORT", quantity=2)
        mismatch = _row(code="UNIT", quantity=1, units="компл")
        fallback = _row(code="RFP-MISS", code_mto="MTO-OK")
        packing = [
            _row(title="8950", system="SOO1", code="OK", quantity=1, source=_ul_source("ok.xlsx")),
            _row(
                title="8950",
                system="SOO1",
                code="SHORT",
                quantity=1,
                source=_ul_source("short.xlsx"),
            ),
            _row(
                title="8950",
                system="SOO1",
                code="UNIT",
                quantity=1,
                units="шт",
                source=_ul_source("unit.xlsx"),
            ),
            _row(
                title="8950",
                system="SOO1",
                code="MTO-OK",
                source=_ul_source("fallback.xlsx"),
            ),
            _row(
                title="8950",
                system="SOO1",
                code="LEFTOVER",
                quantity=1,
                source=_ul_source("left.xlsx"),
            ),
        ]
        ordered = {
            rfp_packing_match_key(1, "8950", "SOO1", "OK"): 1.0,
            rfp_packing_match_key(1, "8950", "SOO1", "SHORT"): 2.0,
            rfp_packing_match_key(1, "8950", "SOO1", "UNIT"): 1.0,
            rfp_packing_match_key(1, "8950", "SOO1", "RFP-MISS"): 1.0,
        }
        rows, audit = compare_rfp_rows_with_packing(
            [complete, short, mismatch, fallback],
            ordered,
            _dataset(packing),
        )
        self.assertEqual(complete.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(short.get_value(UL_COMPARE_STATUS), STATUS_SHORTFALL)
        self.assertEqual(mismatch.get_value(UL_COMPARE_STATUS), STATUS_DATA_PROBLEM)
        self.assertIn(FALLBACK_COMMENT_MARKER, fallback.el[UL_COMPARE_STATUS].comment)
        leftovers = [
            row
            for row in rows
            if row.get_value(UL_COMPARE_STATUS) == STATUS_PACKING_ONLY
        ]
        self.assertTrue(leftovers)
        self.assertTrue(str(complete.el[UL_COMPARE_STATUS].comment or "").strip())
        with tempfile.TemporaryDirectory() as tmp:
            path = save_match_result_to_excel(rows, tmp, packing_audit=audit)
            workbook = openpyxl.load_workbook(
                io.BytesIO(Path(path).read_bytes()),
                data_only=True,
            )
            try:
                sheet = workbook["Сопоставление RFP и MTO"]
                headers = {
                    cell.value: cell.column for cell in sheet[1] if cell.value is not None
                }
                status_col = headers["Статус УЛ"]
                code_col = headers["Код RFP"]
                ul_code_col = headers["Код УЛ"]
                by_rfp_code: dict[str, object] = {}
                leftover_comments = []
                for excel_row in range(2, sheet.max_row + 1):
                    status_cell = sheet.cell(excel_row, status_col)
                    rfp_code = str(sheet.cell(excel_row, code_col).value or "").strip()
                    if rfp_code:
                        by_rfp_code[rfp_code] = status_cell.comment
                    if status_cell.value == STATUS_PACKING_ONLY:
                        leftover_comments.append(status_cell.comment)
                self.assertIsNone(by_rfp_code["OK"])
                self.assertIsNotNone(by_rfp_code["SHORT"])
                self.assertIn("short.xlsx", by_rfp_code["SHORT"].text)
                self.assertIsNotNone(by_rfp_code["UNIT"])
                self.assertIn("packing_units_mismatch", by_rfp_code["UNIT"].text)
                self.assertIsNotNone(by_rfp_code["RFP-MISS"])
                self.assertIn(FALLBACK_COMMENT_MARKER, by_rfp_code["RFP-MISS"].text)
                self.assertTrue(leftover_comments)
                self.assertTrue(all(comment is None for comment in leftover_comments))
                leftover_codes = [
                    str(sheet.cell(excel_row, ul_code_col).value or "")
                    for excel_row in range(2, sheet.max_row + 1)
                    if sheet.cell(excel_row, status_col).value == STATUS_PACKING_ONLY
                ]
                self.assertIn("LEFTOVER", leftover_codes)
            finally:
                workbook.close()

    def test_config_toggles_and_pipeline_order(self) -> None:
        default_config = get_default_config()
        default_asbuild_config = get_default_asbuild_config()
        self.assertIn("use_multiprocessing", default_config["step4"])
        self.assertIn("use_multiprocessing", default_asbuild_config["step4"])
        self.assertTrue(default_config["step4"]["include_packing_lists"])
        self.assertTrue(default_config["step4"]["ul_match_use_mto_tags"])
        self.assertTrue(
            default_config["paths"]["gem_supply_codes"].endswith(
                "Коды_Поставок_ГЭМ_8950.xlsx"
            )
        )
        self.assertFalse(default_asbuild_config["step4"]["include_packing_lists"])
        self.assertTrue(default_asbuild_config["step4"]["ul_match_use_mto_tags"])
        self.assertTrue(default_config["rfp_parts"]["auto_update_checklist"])
        self.assertTrue(default_config["rfp_parts"]["use_latest_net"])
        self.assertFalse(default_asbuild_config["rfp_parts"]["auto_update_checklist"])
        self.assertFalse(default_asbuild_config["rfp_parts"]["use_latest_net"])
        main_config = json.loads(
            (ROOT / "RFQ/tags_rfp_compare/rfp_tags_compare_config.json").read_text(
                encoding="utf-8"
            )
        )
        asbuild_config = json.loads(
            (
                ROOT
                / "RFQ/tags_rfp_compare/rfp_tags_compare_config_asbuild.json"
            ).read_text(encoding="utf-8")
        )
        self.assertTrue(main_config["step4"]["include_packing_lists"])
        self.assertTrue(main_config["step4"]["ul_match_use_mto_tags"])
        self.assertFalse(asbuild_config["step4"]["include_packing_lists"])
        aggregate_source = (
            ROOT / "RFQ/tags_rfp_compare/agregate_tags.py"
        ).read_text(encoding="utf-8")
        include_once = aggregate_source.index(
            'include_packing_lists = bool(step4_cfg.get("include_packing_lists", True))'
        )
        self.assertIn(
            'ul_match_use_mto_tags = bool(step4_cfg.get("ul_match_use_mto_tags", True))',
            aggregate_source,
        )
        self.assertIn("ul_match_use_mto_tags=ul_match_use_mto_tags", aggregate_source)
        packing_load = aggregate_source.index("packing_dataset = load_packing_dataset()")
        ul_preflight = aggregate_source.index("ensure_tsd_packing_cache_current()")
        gate_call = aggregate_source.index("run_units_gate(")
        finalize_call = aggregate_source.index("step1_finalize_rfp_data(")
        step4_call = aggregate_source.index("step4_analyze_and_match(")
        self.assertLess(include_once, ul_preflight)
        self.assertLess(ul_preflight, packing_load)
        self.assertLess(packing_load, gate_call)
        self.assertLess(gate_call, finalize_call)
        self.assertLess(finalize_call, step4_call)
        self.assertIn("if include_packing_lists:", aggregate_source)
        self.assertIn("packing_dataset=packing_dataset", aggregate_source)

        source = inspect.getsource(step4_analyze_and_match)
        balance = source.index("run_output_quantity_balance_check(")
        legacy_check = source.index("check_values_sum(", balance)
        packing = source.index("compare_rfp_rows_with_packing(", legacy_check)
        ul_input = source.index("print_input_ul_quantity_balance(", legacy_check)
        ul_balance = source.index("run_output_ul_quantity_balance_check(", packing)
        gem_overlay = source.index("apply_gem_supply_status(", packing)
        excel = source.index("save_match_result_to_excel(", ul_balance)
        ds_source = source.index("apply_ds_source_columns(", ul_balance)
        gem_row_fill = source.index("paint_gem_supply_rows(", ds_source)
        self.assertLess(balance, legacy_check)
        self.assertLess(legacy_check, ul_input)
        self.assertLess(ul_input, packing)
        self.assertLess(packing, gem_overlay)
        self.assertLess(gem_overlay, ul_balance)
        self.assertLess(ul_balance, ds_source)
        self.assertLess(ds_source, gem_row_fill)
        self.assertLess(gem_row_fill, excel)
        self.assertIn("use_mto_tags=ul_match_use_mto_tags", source)
        self.assertIn(
            'emit_milestone("global_checks", "Done", balance_report.summary)',
            source,
        )
        self.assertIn(
            'emit_milestone("packing_lists", "Done", ul_balance_report.summary)',
            source,
        )
        self.assertIn("format_input_rfp_mto_vo_detail", source)
        self.assertIn("format_input_ul_detail", source)

        fatal_handler = source.index("except RfpPackingFatalError", packing)
        fatal_raise = source.index("raise", fatal_handler)
        self.assertLess(fatal_raise, excel)
        self.assertIn("if include_packing_lists:", source)

    def test_provider_missing_corrupt_old_and_partial_cache(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            cache = Path(tmp) / "packing.cache"
            self.assertEqual(
                load_packing_dataset(cache).quality,
                PackingQualityLevel.UNAVAILABLE,
            )

            cache.write_bytes(b"not a pickle")
            self.assertEqual(
                load_packing_dataset(cache).quality,
                PackingQualityLevel.UNAVAILABLE,
            )

            row = _row(title="8950", system="SOO1")
            payload = {
                "meta": {
                    "version": "old",
                    "created_at": "2026-07-29T12:00:00+00:00",
                    "root": "C:/packing",
                    "fingerprint": "abc",
                    "counters": {"position_rows": 1},
                    "critical": {"ok": True, "remarks": []},
                    "report_paths": {},
                    "loader_issues": [],
                },
                "data": [row],
            }
            with cache.open("wb") as stream:
                pickle.dump(payload, stream)
            self.assertEqual(
                load_packing_dataset(cache).quality,
                PackingQualityLevel.UNAVAILABLE,
            )

            payload["meta"]["version"] = PACKING_CACHE_VERSION
            payload["meta"]["critical"]["ok"] = False
            payload["meta"]["critical"]["remarks"] = ["offline partial"]
            with cache.open("wb") as stream:
                pickle.dump(payload, stream)
            loaded = load_packing_dataset(cache)
            self.assertEqual(loaded.quality, PackingQualityLevel.PARTIAL)
            self.assertTrue(loaded.issues)

    def test_ul_quantity_balance_matches_queued_units(self) -> None:
        target = _row(quantity=1, tags="A")
        packing = _row(
            title="8950",
            system="SOO1",
            quantity=5,
            tags=["A", "B"],
        )
        invalid = _row(
            title="8950",
            system="SOO1",
            quantity="n/a",
        )
        dataset = _dataset([packing, invalid])
        qty_in, rows_in, by_title = packing_queue_input_by_title(dataset)
        self.assertEqual(qty_in, 5.0)
        self.assertEqual(rows_in, 1)
        self.assertEqual(by_title["8950-SOO1"], 5.0)

        result_rows, audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            dataset,
        )
        self.assertEqual(audit.stats.packing_units, 5)
        qty_out, _rows_out, by_title_out = sum_ul_values_by_title(result_rows)
        self.assertEqual(qty_out, 5.0)
        self.assertEqual(by_title_out["8950-SOO1"], 5.0)
        report = run_output_ul_quantity_balance_check(
            qty_in,
            by_title,
            result_rows,
            "",
            available=True,
        )
        self.assertFalse(report.fatal)
        self.assertFalse(report.errors)
        self.assertTrue(report.summary.startswith("OK УЛ "))

    def test_ul_quantity_balance_mismatch_is_fatal(self) -> None:
        target = _row(quantity=1)
        packing = _row(title="8950", system="SOO1", quantity=1)
        dataset = _dataset([packing])
        qty_in, _rows_in, by_title = packing_queue_input_by_title(dataset)
        result_rows, _audit = compare_rfp_rows_with_packing(
            [target],
            {rfp_packing_match_key(1, "8950", "SOO1", "CODE"): 1.0},
            dataset,
        )
        result_rows[0].el[UL_VALUES].value = 0
        with tempfile.TemporaryDirectory() as tmp:
            report = run_output_ul_quantity_balance_check(
                qty_in,
                by_title,
                result_rows,
                tmp,
                available=True,
            )
        self.assertTrue(report.fatal)
        self.assertEqual(len(report.errors), 1)
        self.assertEqual(report.errors[0].part, "УЛ")
        self.assertEqual(report.errors[0].sum_in, 1.0)
        self.assertEqual(report.errors[0].sum_out, 0.0)
        self.assertTrue(report.summary.startswith("ошибка УЛ "))

    def test_ul_quantity_balance_unavailable_is_skipped(self) -> None:
        dataset = _dataset([], quality=PackingQualityLevel.UNAVAILABLE)
        qty_in, rows_in, by_title = packing_queue_input_by_title(dataset)
        self.assertEqual((qty_in, rows_in, by_title), (0.0, 0, {}))
        report = run_output_ul_quantity_balance_check(
            qty_in,
            by_title,
            [],
            "",
            available=False,
        )
        self.assertFalse(report.fatal)
        self.assertFalse(report.errors)
        self.assertEqual(report.summary, "УЛ недоступны")
        self.assertEqual(format_input_ul_detail(available=False), "вход УЛ недоступны")
        self.assertEqual(
            format_input_ul_detail(available=True, qty=12.0, rows=3),
            "вход УЛ 12.00 (3 стр.)",
        )


class Step4ExcelPathLinkSmokeTest(unittest.TestCase):
    """UL source hyperlinks only; PATH_RFP / PATH_MTO stay plain text."""

    def test_ul_single_source_gets_url_multiple_does_not(self) -> None:
        rfp = r"\\bcc\eng\parts\ДС1.xlsx"
        display, url = excel_path_link(PATH_RFP, rfp)
        self.assertEqual(display, rfp)
        self.assertIsNone(url)

        ul = "folder\\pl.xlsx · Single 1 · строка 12"
        display, url = excel_path_link(
            UL_SOURCE_FILES, ul, packing_root=r"C:\tsd"
        )
        self.assertEqual(display, ul)
        self.assertIn("pl.xlsx", str(url))
        self.assertIn("Single 1", str(url))
        self.assertIn("A12", str(url))

        multi_ul = "a.xlsx · S1 · строка 1\nb.xlsx · S2 · строка 2"
        display, url = excel_path_link(
            UL_SOURCE_FILES, multi_ul, packing_root=r"C:\tsd"
        )
        self.assertEqual(display, multi_ul)
        self.assertIsNone(url)
        self.assertIsNone(parse_single_ul_source(multi_ul))


class GemSupplyOverlaySmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.codes_path = Path(self._tmp.name) / "Коды_Поставок_ГЭМ_8950.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Код"])
        sheet.append(["BCC0001233"])
        sheet.append(["BCC0000290"])
        workbook.save(self.codes_path)
        workbook.close()
        self.gem = load_gem_supply_codes(self.codes_path)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _open_upd_row(
        self,
        *,
        title: str = "8950-SOT4",
        code: str = "BCC0001233",
        comment: str = "",
        **kwargs,
    ) -> RowStd:
        row = _row(title=title, code=code, **kwargs)
        row.el[UL_COMPARE_STATUS].value = STATUS_OPEN_UPD
        if comment:
            row.el[UL_COMPARE_STATUS].comment = comment
        return row

    def _assert_full_row_match_matched(self, row: RowStd) -> None:
        sample_cols = (
            DS_NAME,
            DS_TITLE,
            CODE,
            NAME,
            VALUES,
            UNITS,
            TAGS,
            CODE_MTO,
            VALUES_MTO,
            UL_COMPARE_STATUS,
            UL_VALUES,
            PATH_RFP,
            DS_MANAGER,
        )
        for column in sample_cols:
            self.assertEqual(
                row.el[column].color,
                Color.match_matched,
                msg=f"{column} should be match_matched on a GEM row",
            )

    def test_open_upd_8950_gem_code_becomes_gem_supply(self) -> None:
        row = self._open_upd_row(comment="keep-me")
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 1)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_GEM_SUPPLY)
        self.assertEqual(row.el[UL_COMPARE_STATUS].comment, "keep-me")
        self._assert_full_row_match_matched(row)

    def test_open_upd_8950_gem_code_in_cabinet_unchanged(self) -> None:
        row = self._open_upd_row(comment="keep-me")
        row.el[IN_CABINET].value = "2612-S-SK-1501"
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_OPEN_UPD)
        self.assertEqual(row.el[UL_COMPARE_STATUS].comment, "keep-me")
        self.assertNotEqual(row.el[UL_COMPARE_STATUS].color, Color.match_matched)

    def test_mto_only_8950_gem_code_in_cabinet_unchanged(self) -> None:
        row = _row(title="8950-SOT4", code="", code_mto="BCC0001233")
        row.el[UL_COMPARE_STATUS].value = STATUS_MTO_ONLY
        row.el[IN_CABINET].value = "SX-01"
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_MTO_ONLY)
        self.assertNotEqual(row.el[UL_COMPARE_STATUS].color, Color.match_matched)

    def test_open_upd_other_title_stays_open_upd(self) -> None:
        row = self._open_upd_row(title="8445-SOT1")
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_OPEN_UPD)
        self.assertNotEqual(row.el[UL_COMPARE_STATUS].color, Color.match_matched)

    def test_shortfall_8950_gem_code_unchanged(self) -> None:
        row = _row(title="8950-SOT4", code="BCC0001233")
        row.el[UL_COMPARE_STATUS].value = STATUS_SHORTFALL
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_SHORTFALL)

    def test_complete_8950_gem_code_unchanged(self) -> None:
        row = _row(title="8950-SOT4", code="BCC0001233")
        row.el[UL_COMPARE_STATUS].value = STATUS_COMPLETE
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)

    def test_mto_delivered_8950_gem_code_unchanged(self) -> None:
        row = _row(title="8950-SOT4", code="", code_mto="BCC0001233")
        row.el[UL_COMPARE_STATUS].value = STATUS_MTO_DELIVERED
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_MTO_DELIVERED)

    def test_mto_only_8950_gem_code_becomes_gem_supply(self) -> None:
        row = _row(title="8950-SOT4", code="", code_mto="BCC0001233")
        row.el[UL_COMPARE_STATUS].value = STATUS_MTO_ONLY
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 1)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_GEM_SUPPLY)
        self._assert_full_row_match_matched(row)

    def test_gem_row_fill_restored_after_later_colouring(self) -> None:
        row = self._open_upd_row()
        apply_gem_supply_status([row], self.gem)
        row.el[VALUES_MTO_UL_DIFF].color = Color.no
        row.el[DS_MANAGER].color = Color.yellow
        row.el[PATH_RFP].color = Color.yellow
        painted = paint_gem_supply_rows([row])
        self.assertEqual(painted, 1)
        self._assert_full_row_match_matched(row)
        self.assertEqual(row.el[VALUES_MTO_UL_DIFF].color, Color.match_matched)

    def test_mto_only_other_title_stays_mto_only(self) -> None:
        row = _row(title="8445-SOT1", code="", code_mto="BCC0001233")
        row.el[UL_COMPARE_STATUS].value = STATUS_MTO_ONLY
        changed = apply_gem_supply_status([row], self.gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_MTO_ONLY)

    def test_missing_path_load_error_apply_noop(self) -> None:
        missing = Path(self._tmp.name) / "missing_gem_codes.xlsx"
        gem = load_gem_supply_codes(missing)
        self.assertTrue(gem.load_error)
        self.assertFalse(gem.codes)
        row = self._open_upd_row()
        changed = apply_gem_supply_status([row], gem)
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_OPEN_UPD)

    def test_load_fixture_xlsx_normalizes_codes(self) -> None:
        self.assertFalse(self.gem.load_error)
        self.assertIn(
            normalize_packing_key_part("BCC0001233"),
            self.gem.codes,
        )
        self.assertIn(
            normalize_packing_key_part("BCC0000290"),
            self.gem.codes,
        )


class ZipSmrPnrOverlaySmokeTest(unittest.TestCase):
    def test_customer_phrase_packing_only_becomes_zip_green(self) -> None:
        row = _row()
        row.el[UL_NAME].value = (
            "Комплект ЗИП для СМР, ПНР 1600 SOTМуфта вводная..."
        )
        row.el[UL_COMPARE_STATUS].value = STATUS_PACKING_ONLY
        row.el[UL_COMPARE_STATUS].color = Color.soft_cyan
        row.el[UL_COMPARE_STATUS].comment = "keep-zip"
        row.el[DS_TITLE].color = Color.soft_cyan
        changed = apply_zip_smr_pnr_status([row])
        self.assertEqual(changed, 1)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_ZIP_SMR_PNR)
        self.assertEqual(row.el[UL_COMPARE_STATUS].color, Color.match_matched)
        self.assertEqual(row.el[UL_COMPARE_STATUS].comment, "keep-zip")
        self.assertEqual(row.el[DS_TITLE].color, Color.match_matched)

    def test_variants_match(self) -> None:
        names = [
            "Комплект ЗИП для СМР,ППНР,ПНР Контроллер",
            "ЗИП для СМР, ПНР. Аккумулятор",
            "ЗИП, для проведения СМР, ПредПНР, ПНР: Коробка",
        ]
        for name in names:
            with self.subTest(name=name):
                self.assertTrue(ul_name_is_zip_smr_pnr(name))
                row = _row()
                row.el[UL_NAME].value = name
                row.el[UL_COMPARE_STATUS].value = STATUS_PACKING_ONLY
                changed = apply_zip_smr_pnr_status([row])
                self.assertEqual(changed, 1)
                self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_ZIP_SMR_PNR)

    def test_uzip_and_camera_kit_ignored(self) -> None:
        for name in (
            "УЗИП оборудование защиты",
            "Комплект ЗИП: Карта памяти MicroSD 64G",
        ):
            with self.subTest(name=name):
                self.assertFalse(ul_name_is_zip_smr_pnr(name))
                row = _row()
                row.el[UL_NAME].value = name
                row.el[UL_COMPARE_STATUS].value = STATUS_PACKING_ONLY
                changed = apply_zip_smr_pnr_status([row])
                self.assertEqual(changed, 0)
                self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_PACKING_ONLY)

    def test_data_problem_not_overlaid(self) -> None:
        row = _row()
        row.el[UL_NAME].value = "Комплект ЗИП для СМР, ПНР тест"
        row.el[UL_COMPARE_STATUS].value = STATUS_DATA_PROBLEM
        row.el[UL_COMPARE_STATUS].color = Color.red
        changed = apply_zip_smr_pnr_status([row])
        self.assertEqual(changed, 0)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_DATA_PROBLEM)
        self.assertEqual(row.el[UL_COMPARE_STATUS].color, Color.red)

    def test_complete_row_only_status_cell(self) -> None:
        row = _row()
        row.el[UL_NAME].value = "ЗИП для СМР, ПНР. Аккумулятор"
        row.el[UL_COMPARE_STATUS].value = STATUS_COMPLETE
        row.el[UL_COMPARE_STATUS].color = Color.green
        row.el[DS_TITLE].color = Color.no
        changed = apply_zip_smr_pnr_status([row])
        self.assertEqual(changed, 1)
        self.assertEqual(row.get_value(UL_COMPARE_STATUS), STATUS_ZIP_SMR_PNR)
        self.assertEqual(row.el[UL_COMPARE_STATUS].color, Color.match_matched)
        self.assertEqual(row.el[DS_TITLE].color, Color.no)


if __name__ == "__main__":
    unittest.main(verbosity=2)
