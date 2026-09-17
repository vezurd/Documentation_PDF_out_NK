"""Read-only: check whether Step4 row order is stable across runs, and compare comments as multisets."""

from __future__ import annotations

import hashlib
import sys
from collections import Counter
from pathlib import Path

import openpyxl

DIR = Path(__file__).resolve().parent / "step4_compare"
OLD = DIR / "Шаг4_Сопоставление_RFP_MTO_20260909_131931.xlsx"
NEW = DIR / "Шаг4_Сопоставление_RFP_MTO_20260909_134240.xlsx"

RESULTS_ROOT = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_РЕЗУЛЬТАТА_ПРОВЕРКИ"
)
SHEET = "Сопоставление RFP и MTO"


def comment_multiset(path: Path) -> Counter[str]:
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        sheet = wb[SHEET]
        found: Counter[str] = Counter()
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    found[cell.comment.text] += 1
        return found
    finally:
        wb.close()


def order_signature(path: Path) -> tuple[int, list[str], str]:
    """Return (row count, first DS values, hash of the DS/code order)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb[SHEET]
        keys: list[str] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index == 0:
                continue
            keys.append(f"{row[0]}|{row[2]}|{row[5]}")
        digest = hashlib.sha1("\n".join(keys).encode("utf-8")).hexdigest()[:12]
        first = [k.split("|")[0] for k in keys[:12]]
        return len(keys), first, digest
    finally:
        wb.close()


def main() -> None:
    print("=== комментарии как мультимножество ===")
    old_comments = comment_multiset(OLD)
    new_comments = comment_multiset(NEW)
    only_old = sum((old_comments - new_comments).values())
    only_new = sum((new_comments - old_comments).values())
    print(f"  всего old={sum(old_comments.values())} new={sum(new_comments.values())}")
    print(f"  only_old={only_old}  only_new={only_new}")
    for text, count in list((old_comments - new_comments).items())[:5]:
        print(f"    ТОЛЬКО OLD x{count}: {text[:160]!r}")
    for text, count in list((new_comments - old_comments).items())[:5]:
        print(f"    ТОЛЬКО NEW x{count}: {text[:160]!r}")

    print("\n=== порядок строк по прогонам ===")
    runs = sorted(
        (p for p in RESULTS_ROOT.iterdir() if p.is_dir() and p.name.startswith("_результат")),
        key=lambda p: p.name,
    )
    for run in runs:
        files = sorted(run.glob("Шаг4_Сопоставление_RFP_MTO_*.xlsx"))
        if not files:
            continue
        try:
            count, first, digest = order_signature(files[-1])
        except Exception as error:  # noqa: BLE001 - диагностика, любые ошибки чтения важны
            print(f"  {run.name}: ошибка чтения — {type(error).__name__}: {error}")
            continue
        print(f"  {run.name}")
        print(f"    строк={count}  order_hash={digest}  первые ДС={first}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    main()
