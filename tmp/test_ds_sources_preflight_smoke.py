"""Smoke tests for grouped source preflight and workbook source sheets."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import xlsxwriter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from base.base_classes import RowType
from RFQ.ds_compare.ds_merge_folder import merge_ds_folder
from RFQ.ds_compare.ds_sources_preflight import (
    SourceFileRecord,
    SourceManifest,
    refresh_sources_before_grouped,
)
from RFQ.ds_compare.ds_vs_mto_excel_xlsxwriter import _write_source_sheets
from RFQ.ds_compare.tsd_packing_load import load_and_cache_tsd_packing
from RFQ.packing_list_provider import PackingQualityLevel


class DsSourcesPreflightSmokeTest(unittest.TestCase):
    """Check stale DS merge and source-sheet structure."""

    def test_stale_ds_summary_is_replaced(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            merge_dir = root / "ds"
            merge_dir.mkdir()
            source = merge_dir / "ДС1.xlsx"
            source.write_bytes(b"source")
            old_summary = root / "old_summary.xlsx"
            old_summary.write_bytes(b"old")
            os.utime(old_summary, (1, 1))
            os.utime(source, (2, 2))
            new_summary = root / "new_summary.xlsx"

            def fake_merge(*args, **kwargs) -> str:
                del args, kwargs
                new_summary.write_bytes(b"new")
                return str(new_summary)

            cfg = {
                "gui_paths": {
                    "last_ds_file": str(old_summary),
                    "last_rfq_file": "",
                    "last_merge_ds_folder": str(merge_dir),
                    "last_tsd_packing_folder": str(root / "missing-packing"),
                    "last_tsd_summary_file": "",
                }
            }
            with (
                patch(
                    "RFQ.ds_compare.ds_sources_preflight.merge_ds_folder",
                    side_effect=fake_merge,
                ),
                patch(
                    "RFQ.ds_compare.ds_sources_preflight.save_ds_compare_config",
                    return_value=True,
                ),
            ):
                result = refresh_sources_before_grouped(
                    ds_path=str(old_summary),
                    rfq_only=False,
                    cfg=cfg,
                )

            self.assertTrue(result.ds_refreshed)
            self.assertEqual(result.effective_ds_path, str(new_summary))
            self.assertEqual(cfg["gui_paths"]["last_ds_file"], str(new_summary))
            self.assertEqual(len(result.manifest.ds), 2)

    def test_three_source_sheets_are_written(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "sources.xlsx"
            record = SourceFileRecord(
                key="1",
                role="исходный",
                path=r"C:\source.xlsx",
                modified_at="2026-08-12 12:00:00",
                status="использован",
            )
            manifest = SourceManifest(
                ds=[record],
                mto=[record],
                packing=[record],
                packing_root=r"C:\packing",
                packing_fingerprint="abc",
                packing_cache_path=r"C:\packing.cache",
            )
            workbook = xlsxwriter.Workbook(str(output))
            _write_source_sheets(workbook, manifest)
            workbook.close()

            workbook_read = openpyxl.load_workbook(output, read_only=True)
            try:
                self.assertEqual(
                    workbook_read.sheetnames,
                    ["Источники ДС", "Источники МТО", "Источники УЛ"],
                )
            finally:
                workbook_read.close()

    def test_packing_cache_hit_reuses_existing_summary(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            summary = root / "existing_summary.xlsx"
            summary.write_bytes(b"existing")
            dataset = SimpleNamespace(
                available=True,
                rows=[object()],
                meta=SimpleNamespace(
                    report_paths={"summary": str(summary)},
                    counters={"files_total": 1, "position_rows": 1},
                ),
                quality=PackingQualityLevel.OK,
                issues=[],
            )
            with (
                patch(
                    "RFQ.ds_compare.tsd_packing_load.collect_tsd_files",
                    return_value=[],
                ),
                patch(
                    "RFQ.ds_compare.tsd_packing_load._files_fingerprint",
                    return_value="fingerprint",
                ),
                patch(
                    "RFQ.ds_compare.tsd_packing_load.load_packing_dataset",
                    return_value=dataset,
                ),
                patch(
                    "RFQ.ds_compare.tsd_packing_load.save_tsd_packing_summary_xlsx"
                ) as save_summary,
            ):
                result = load_and_cache_tsd_packing(str(root), force=False)

            self.assertEqual(result.summary_path, str(summary))
            self.assertTrue(result.from_cache)
            save_summary.assert_not_called()

    def test_ensure_tsd_packing_cache_current_uses_force_false(self) -> None:
        from RFQ.ds_compare.tsd_packing_load import (
            TsdLoadResult,
            TsdLoadStats,
            ensure_tsd_packing_cache_current,
            format_packing_cache_freshness_detail,
        )

        fake = TsdLoadResult(
            "summary.xlsx",
            TsdLoadStats(position_rows=3),
            True,
            from_cache=True,
        )
        with patch(
            "RFQ.ds_compare.tsd_packing_load.load_and_cache_tsd_packing",
            return_value=fake,
        ) as loader:
            result = ensure_tsd_packing_cache_current(root=r"C:\tsd")

        loader.assert_called_once_with(r"C:\tsd", force=False)
        self.assertIs(result, fake)
        self.assertIn("unchanged", format_packing_cache_freshness_detail(result))
        rebuilt = TsdLoadResult(
            "summary.xlsx",
            TsdLoadStats(position_rows=3),
            False,
            from_cache=False,
        )
        rebuilt_detail = format_packing_cache_freshness_detail(rebuilt)
        self.assertIn("rebuilt", rebuilt_detail)
        self.assertIn("критичные замечания", rebuilt_detail)

    def test_ds_merge_stops_on_wrong_quantity_column(self) -> None:
        class BadRow:
            row_type = RowType.position_row

            @staticmethod
            def get_value(key: str) -> object:
                del key
                return "шт"

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "ДС99_wrong_layout.xlsx"
            source.write_bytes(b"source")
            doc = SimpleNamespace(
                file_full_path=str(source),
                file_name=source.name,
            )
            with (
                patch(
                    "RFQ.ds_compare.ds_merge_folder.utils.path.get_files_single",
                    return_value=[doc],
                ),
                patch(
                    "RFQ.ds_compare.ds_merge_folder.load_ds_data",
                    return_value=[BadRow()],
                ),
                patch(
                    "RFQ.ds_compare.ds_merge_folder.utils.path.get_path_out_dir",
                    return_value=str(root / "out"),
                ),
            ):
                result = merge_ds_folder(str(root), open_folder=False)

            self.assertIsNone(result)
            report = root / "out" / "DS_merge_load_report.txt"
            self.assertIn("Ошибка формата", report.read_text(encoding="utf-8"))

if __name__ == "__main__":
    unittest.main()
