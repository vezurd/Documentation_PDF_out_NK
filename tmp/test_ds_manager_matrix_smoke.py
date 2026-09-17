"""Smoke tests for DS manager matrix loader and column application."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from base.base_classes import RowStd, RowType, TableComments
from base.tables_columns import DS_ACTUAL, DS_MANAGER, DS_NAME, DS_TITLE, PATH_MTO, PATH_RFP
from RFQ.tags_rfp_compare.ds_manager_matrix import (
    _fill_leftover_ds_managers,
    apply_ds_source_columns,
    build_mto_paths_by_title,
    collect_parts_ds_paths,
    compare_matrix_to_parts,
    load_ds_manager_matrix,
    lookup_manager,
)
from RFQ.tags_rfp_compare.ds_manager_roster import (
    ROSTER_SHEET_TITLE,
    build_ds_roster_rows,
    sync_ds_roster_sheet,
)
from utils.colors import Color


def _write_matrix_xlsx(path: Path) -> None:
    """Write СписокДС-like workbook: position dump first, manager sheet second."""
    wb = Workbook()
    positions = wb.active
    positions.title = "Лист1"
    positions.append(
        ["Имя ДС", "№ позиции", "Титул/Марка", "Код RFP", "Наименование RFP"]
    )
    positions.append(["ДС6", "1", "8310-SOS", "BCC0000341", "dummy"])

    managers = wb.create_sheet("Лист2")
    managers.append(["Имя ДС", "Кол-во RFP", "Фамилия"])
    managers.append(["ДС01", None, "ГЭМ"])
    managers.append(["ДС47_13", 10, "Михеев Д."])
    managers.append(["ДС73_36", None, ""])
    wb.save(path)
    wb.close()


def _write_legacy_matrix_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ДС"
    ws.append(["№ ДС", "МП", "Комментарий"])
    ws.append(["10", "Михеев Д.", ""])
    ws.append(["", "", "ДС61 nested comment — ignored"])
    ws.append(["47/13А", "Михеев Д.", ""])
    wb.save(path)
    wb.close()


def _write_empty_workbook(path: Path) -> None:
    wb = Workbook()
    wb.save(path)
    wb.close()


class DsManagerMatrixSmokeTest(unittest.TestCase):
    """End-to-end smoke for manager matrix load, compare, and column fill."""

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.tmp = Path(self._tmpdir.name)

    def tearDown(self) -> None:
        self._tmpdir.cleanup()

    def test_load_matrix_keys_and_manager(self) -> None:
        matrix_path = self.tmp / "matrix.xlsx"
        _write_matrix_xlsx(matrix_path)

        matrix = load_ds_manager_matrix(matrix_path)
        self.assertIsNone(matrix.load_error)
        self.assertEqual(len(matrix.entries), 3)

        by_raw = {entry.raw_ds: entry for entry in matrix.entries}
        self.assertIn("ДС01", by_raw)
        self.assertIn("ДС47_13", by_raw)
        self.assertIn("ДС73_36", by_raw)
        self.assertNotIn("ДС6", by_raw)

        self.assertIn("ДС01", matrix.by_key)
        self.assertIn("ДС1", matrix.by_key)
        self.assertIn("ДС47_13", matrix.by_key)
        self.assertEqual(matrix.by_key["ДС01"].manager, "ГЭМ")
        self.assertEqual(lookup_manager(matrix, "ДС1").manager, "ГЭМ")
        self.assertEqual(lookup_manager(matrix, "ДС47_13А").manager, "Михеев Д.")
        self.assertEqual(lookup_manager(matrix, "ДС73_36").manager, "")

    def test_compare_matrix_to_parts(self) -> None:
        matrix_path = self.tmp / "matrix.xlsx"
        _write_matrix_xlsx(matrix_path)
        matrix = load_ds_manager_matrix(matrix_path)

        parts_dir = self.tmp / "parts"
        parts_dir.mkdir()
        _write_empty_workbook(parts_dir / "ДС1_test.xlsx")
        _write_empty_workbook(parts_dir / "ДС99_x.xlsx")

        paths = collect_parts_ds_paths(parts_dir)
        self.assertIn("ДС1_TEST", paths)
        self.assertIn("ДС99_X", paths)

        result = compare_matrix_to_parts(matrix, parts_dir)
        self.assertIsNone(result.load_error)
        self.assertEqual(result.matrix_count, 3)
        self.assertEqual(result.folder_ds_count, 2)
        self.assertIn("ДС47_13", result.missing_file)
        self.assertIn("ДС73_36", result.missing_file)
        self.assertNotIn("ДС01", result.missing_file)
        self.assertIn("ДС99_X", result.extra_files)
        self.assertIn("без файла=2", result.summary_line)
        self.assertIn("файлы без матрицы=1", result.summary_line)

    def test_apply_ds_source_columns(self) -> None:
        matrix_path = self.tmp / "matrix.xlsx"
        _write_matrix_xlsx(matrix_path)
        matrix = load_ds_manager_matrix(matrix_path)

        row_ok = RowStd()
        row_ok.row_type = RowType.position_row
        row_ok.el[DS_NAME].value = "ДС1"
        row_ok.el[DS_TITLE].value = "2265-KSB"

        row_missing = RowStd()
        row_missing.row_type = RowType.position_row
        row_missing.el[DS_NAME].value = "ДС47_13А"

        row_empty = RowStd()
        row_empty.row_type = RowType.position_row
        row_empty.el[DS_NAME].value = None

        rfp_paths = {"ДС1": [str(self.tmp / "parts" / "ДС1_test.xlsx")]}
        mto_paths = build_mto_paths_by_title(
            {
                "2265-KSB": [
                    _row_with_mto_path("/data/mto/file1.xlsx"),
                    _row_with_mto_path("/data/mto/file1.xlsx"),
                    _row_with_mto_path("/data/mto/file2.xlsx"),
                ],
            }
        )

        apply_ds_source_columns(
            [row_ok, row_missing, row_empty],
            matrix=matrix,
            rfp_paths_by_key=rfp_paths,
            mto_paths_by_title=mto_paths,
        )

        self.assertEqual(row_ok.el[DS_MANAGER].value, "ГЭМ")
        self.assertEqual(row_ok.el[DS_MANAGER].color, Color.no)
        self.assertIn("ДС1_test.xlsx", str(row_ok.el[PATH_RFP].value))
        self.assertEqual(
            row_ok.el[PATH_MTO].value,
            "/data/mto/file1.xlsx; /data/mto/file2.xlsx",
        )

        self.assertEqual(row_missing.el[DS_MANAGER].value, "Михеев Д.")
        self.assertEqual(row_missing.el[DS_MANAGER].color, Color.no)
        self.assertEqual(row_missing.el[PATH_RFP].color, Color.yellow)
        self.assertEqual(row_missing.el[PATH_RFP].comment, "")
        self.assertEqual(row_missing.el[PATH_RFP].value, "")

        self.assertEqual(row_empty.el[DS_MANAGER].value, "")
        self.assertEqual(row_empty.el[DS_MANAGER].color, Color.no)
        self.assertEqual(row_empty.el[PATH_RFP].value, "")
        self.assertEqual(row_empty.el[PATH_MTO].value, "")

    def test_leftover_sequential_unique_matrix_manager(self) -> None:
        matrix_path = self.tmp / "matrix.xlsx"
        _write_matrix_xlsx(matrix_path)
        matrix = load_ds_manager_matrix(matrix_path)

        leftover = RowStd()
        leftover.row_type = RowType.position_row
        leftover.el[DS_NAME].value = "ДС1_RFP_не_найден"
        leftover.el[DS_ACTUAL].value = "ДС1"
        leftover.el[DS_TITLE].value = "2265-KSB"

        apply_ds_source_columns(
            [leftover],
            matrix=matrix,
            rfp_paths_by_key={"ДС1": [str(self.tmp / "parts" / "ДС1_test.xlsx")]},
            mto_paths_by_title={},
        )
        self.assertEqual(leftover.el[DS_MANAGER].value, "ГЭМ")
        self.assertEqual(leftover.el[DS_MANAGER].color, Color.no)
        self.assertIn("ДС1_test.xlsx", str(leftover.el[PATH_RFP].value))

    def test_leftover_inherits_unique_sibling_manager(self) -> None:
        leftover = RowStd()
        leftover.row_type = RowType.position_row
        leftover.el[DS_NAME].value = "ДС10_RFP_не_найден"
        leftover.el[DS_ACTUAL].value = "ДС10"

        sibling = RowStd()
        sibling.row_type = RowType.position_row
        sibling.el[DS_NAME].value = "ДС10"
        sibling.el[DS_ACTUAL].value = "ДС10"
        sibling.el[DS_MANAGER].value = "Михеев Д."

        _fill_leftover_ds_managers([leftover, sibling], None)
        self.assertEqual(leftover.el[DS_MANAGER].value, "Михеев Д.")

    def test_leftover_empty_when_actual_has_several_surnames(self) -> None:
        leftover = RowStd()
        leftover.row_type = RowType.position_row
        leftover.el[DS_NAME].value = "ДС101_RFP_не_найден"
        leftover.el[DS_ACTUAL].value = "ДС101"

        first = RowStd()
        first.row_type = RowType.position_row
        first.el[DS_NAME].value = "ДС101"
        first.el[DS_ACTUAL].value = "ДС101"
        first.el[DS_MANAGER].value = "Исаков К."

        second = RowStd()
        second.row_type = RowType.position_row
        second.el[DS_NAME].value = "ДС101"
        second.el[DS_ACTUAL].value = "ДС101"
        second.el[DS_MANAGER].value = "Беляева Е."

        _fill_leftover_ds_managers([leftover, first, second], None)
        self.assertEqual(leftover.el[DS_MANAGER].value, "")
        self.assertEqual(leftover.el[DS_MANAGER].color, Color.no)

    def test_leftover_empty_when_matrix_has_several_surnames(self) -> None:
        matrix_path = self.tmp / "conflict.xlsx"
        wb = Workbook()
        managers = wb.active
        managers.title = "Лист2"
        managers.append(["Имя ДС", "Кол-во RFP", "Фамилия"])
        managers.append(["ДС101", None, "Исаков К."])
        managers.append(["ДС101", None, "Беляева Е."])
        wb.save(matrix_path)
        wb.close()
        matrix = load_ds_manager_matrix(matrix_path)

        leftover = RowStd()
        leftover.row_type = RowType.position_row
        leftover.el[DS_NAME].value = "ДС101_RFP_не_найден"
        leftover.el[DS_ACTUAL].value = "ДС101"

        apply_ds_source_columns(
            [leftover],
            matrix=matrix,
            rfp_paths_by_key={},
            mto_paths_by_title={},
        )
        self.assertEqual(leftover.el[DS_MANAGER].value, "")

    def test_gf_leftover_without_actual_has_no_manager(self) -> None:
        leftover = RowStd()
        leftover.row_type = RowType.position_row
        leftover.el[DS_NAME].value = "RFP_не_найден"
        leftover.el[DS_ACTUAL].value = ""

        apply_ds_source_columns(
            [leftover],
            matrix=None,
            rfp_paths_by_key={},
            mto_paths_by_title={},
        )
        self.assertEqual(leftover.el[DS_MANAGER].value, "")

    def test_gosfin_label_looks_up_manager_by_actual(self) -> None:
        matrix_path = self.tmp / "matrix.xlsx"
        _write_matrix_xlsx(matrix_path)
        matrix = load_ds_manager_matrix(matrix_path)
        self.assertEqual(lookup_manager(matrix, "ДС1_ГОСФИН_AGCC").manager, "ГЭМ")
        self.assertEqual(lookup_manager(matrix, "ДС1_ПРИЛОЖЕНИЕ").manager, "ГЭМ")

        gosfin = RowStd()
        gosfin.row_type = RowType.position_row
        gosfin.el[DS_NAME].value = "ДС1_ГОСФИН_AGCC"
        gosfin.el[DS_ACTUAL].value = "ДС1"
        gosfin.el[DS_TITLE].value = "2265-KSB"

        apply_ds_source_columns(
            [gosfin],
            matrix=matrix,
            rfp_paths_by_key={"ДС1": [str(self.tmp / "parts" / "ДС1_test.xlsx")]},
            mto_paths_by_title={},
        )
        self.assertEqual(gosfin.el[DS_MANAGER].value, "ГЭМ")
        self.assertIn("ДС1_test.xlsx", str(gosfin.el[PATH_RFP].value))

    def test_load_error_missing_file(self) -> None:
        matrix = load_ds_manager_matrix(self.tmp / "missing.xlsx")
        self.assertIsNotNone(matrix.load_error)
        self.assertEqual(matrix.entries, [])
        result = compare_matrix_to_parts(matrix, self.tmp)
        self.assertEqual(result.missing_file, [])
        self.assertIn("ошибка", result.summary_line.lower())

    def test_load_legacy_ds_mp_headers(self) -> None:
        matrix_path = self.tmp / "legacy.xlsx"
        _write_legacy_matrix_xlsx(matrix_path)
        matrix = load_ds_manager_matrix(matrix_path)
        self.assertIsNone(matrix.load_error)
        self.assertEqual(lookup_manager(matrix, "ДС10").manager, "Михеев Д.")
        self.assertEqual(lookup_manager(matrix, "ДС47_13А").manager, "Михеев Д.")

    def test_roster_inverted_name_and_sync(self) -> None:
        matrix_path = self.tmp / "roster.xlsx"
        wb = Workbook()
        dump = wb.active
        dump.title = "Лист1"
        dump.append(["Имя ДС", "№ позиции", "Код RFP"])
        dump.append(["ДС6", "1", "dummy"])
        managers = wb.create_sheet("Лист2")
        managers.append(["Имя ДС", "Кол-во RFP", "Фамилия"])
        managers.append(["ДС01", None, "ГЭМ"])
        managers.append(["ДС92_24Б", None, "Михеев Д."])
        managers.append(["ДС73_36", None, ""])
        managers.append(["ДС101", None, "Исаков К."])
        managers.append(["ДС101", None, "Беляева Е."])
        wb.save(matrix_path)
        wb.close()

        matrix = load_ds_manager_matrix(matrix_path)
        parts_dir = self.tmp / "parts_roster"
        parts_dir.mkdir()
        _write_empty_workbook(parts_dir / "ДС1. test.xlsx")
        _write_empty_workbook(parts_dir / "ДС24_92. AGCC.xlsx")
        _write_empty_workbook(parts_dir / "ДС101. extra.xlsx")

        rows = build_ds_roster_rows(matrix, parts_dir)
        by_copy = {row.copy_name: row for row in rows}
        self.assertEqual(by_copy["ДС1"].manager, "ГЭМ")
        self.assertEqual(by_copy["ДС1"].actual_label, "ДС1")
        self.assertEqual(by_copy["ДС1"].sequential_label, "ДС1")
        inverted = by_copy["ДС24_92Б"]
        self.assertEqual(inverted.actual_label, "ДС24")
        self.assertEqual(inverted.sequential_label, "ДС92")
        self.assertEqual(inverted.manager, "Михеев Д.")
        self.assertEqual(inverted.was_on_sheet1, "ДС92_24Б")
        self.assertIn("перевёрнут", inverted.status)
        ds101_rows = [row for row in rows if row.actual_label == "ДС101"]
        self.assertEqual(len(ds101_rows), 2)
        matched_101 = [row for row in ds101_rows if "дубликат" not in row.status]
        dup_101 = [row for row in ds101_rows if "дубликат" in row.status]
        self.assertEqual(matched_101[0].manager, "Исаков К.")
        self.assertEqual(dup_101[0].manager, "Беляева Е.")
        self.assertTrue(dup_101[0].duplicate_label.startswith("да:"))

        first = sync_ds_roster_sheet(matrix, parts_dir)
        self.assertTrue(first.wrote)
        self.assertIsNone(first.write_error)
        self.assertIn("Столбцы статуса записаны", first.summary_line)
        self.assertTrue(first.summary_line.startswith("Фамилии по файлам RFP заполнены."))
        self.assertIn("Дубликаты:", first.summary_line)

        workbook = load_workbook(matrix_path)
        self.assertIn(ROSTER_SHEET_TITLE, workbook.sheetnames)
        self.assertEqual(workbook.sheetnames[0], "Лист1")
        self.assertIn("Лист2", workbook.sheetnames)
        roster = workbook[ROSTER_SHEET_TITLE]
        self.assertEqual(roster["A1"].value, "Имя ДС")
        self.assertEqual(roster["B1"].value, "Кол-во RFP")
        self.assertEqual(roster["C1"].value, "Фамилия")
        self.assertEqual(roster["D1"].value, "Фактический ДС")
        list2 = workbook["Лист2"]
        self.assertEqual(list2["A1"].value, "Имя ДС")
        self.assertEqual(list2["C1"].value, "Фамилия")
        self.assertEqual(list2["A3"].value, "ДС92_24Б")
        self.assertEqual(list2["C3"].value, "Михеев Д.")
        self.assertIsNone(list2["A3"].comment)
        self.assertIsNone(list2["C4"].comment)
        self.assertEqual(list2["D1"].value, "Статус")
        self.assertEqual(list2["E1"].value, "Дубликат")
        self.assertEqual(list2["F1"].value, "Что делать")
        self.assertIn("перевёрнут", str(list2["D3"].value))
        self.assertEqual(str(list2["E3"].value), "нет")
        self.assertIn("только на Лист2", str(list2["D4"].value))
        self.assertIn("дубликат", str(list2["D6"].value).lower())
        self.assertTrue(str(list2["E6"].value).startswith("да:"))
        self.assertIn("есть лишняя строка", str(list2["E5"].value))
        self.assertIn("Дозаполнять не нужно", str(list2["F4"].value))
        workbook.close()

        pasted_path = self.tmp / "pasted_list2.xlsx"
        src = load_workbook(matrix_path)
        roster_src = src[ROSTER_SHEET_TITLE]
        pasted = Workbook()
        pasted_sheet = pasted.active
        pasted_sheet.title = "Лист2"
        for row in roster_src.iter_rows(values_only=True):
            pasted_sheet.append(list(row))
        src.close()
        pasted.save(pasted_path)
        pasted.close()
        from_paste = load_ds_manager_matrix(pasted_path)
        self.assertIsNone(from_paste.load_error)
        self.assertEqual(lookup_manager(from_paste, "ДС24_92").manager, "Михеев Д.")
        self.assertEqual(lookup_manager(from_paste, "ДС1").manager, "ГЭМ")

        reloaded = load_ds_manager_matrix(matrix_path)
        self.assertIsNone(reloaded.load_error)
        self.assertEqual(len(reloaded.entries), 5)
        self.assertEqual(lookup_manager(reloaded, "ДС01").manager, "ГЭМ")
        self.assertEqual(lookup_manager(reloaded, "ДС101").manager, "Исаков К.")

        second = sync_ds_roster_sheet(reloaded, parts_dir)
        self.assertFalse(second.wrote)
        self.assertIn("лишние на Лист2", second.summary_line)

        _write_empty_workbook(parts_dir / "ДС99. extra.xlsx")
        third = sync_ds_roster_sheet(reloaded, parts_dir)
        self.assertTrue(third.wrote)
        extra_rows = build_ds_roster_rows(reloaded, parts_dir)
        extra_by_copy = {row.copy_name: row for row in extra_rows}
        self.assertEqual(extra_by_copy["ДС99"].manager, "")
        self.assertIn("нет на Лист2", extra_by_copy["ДС99"].status)

        reloaded_after = load_ds_manager_matrix(matrix_path)
        _write_empty_workbook(parts_dir / "ДС88. lock.xlsx")
        with patch("RFQ.tags_rfp_compare.ds_manager_roster.os.replace", side_effect=PermissionError("locked")):
            locked = sync_ds_roster_sheet(reloaded_after, parts_dir)
        self.assertTrue(locked.wrote)
        self.assertTrue(locked.used_fallback)
        self.assertIsNone(locked.write_error)
        fallback = Path(locked.saved_path or "")
        self.assertTrue(fallback.exists())
        self.assertNotEqual(fallback, matrix_path)
        self.assertTrue(fallback.stem.startswith(matrix_path.stem + "_"))
        self.assertIn("копия", locked.summary_line)
        fallback.unlink()

    def test_roster_compare_job_prints_and_stores(self) -> None:
        matrix_path = self.tmp / "job.xlsx"
        wb = Workbook()
        managers = wb.active
        managers.title = "Лист2"
        managers.append(["Имя ДС", "Кол-во RFP", "Фамилия"])
        managers.append(["ДС01", None, "ГЭМ"])
        wb.save(matrix_path)
        wb.close()

        parts_dir = self.tmp / "parts_job"
        parts_dir.mkdir()
        _write_empty_workbook(parts_dir / "ДС1. test.xlsx")

        from RFQ.tags_rfp_compare.ds_manager_roster import (
            get_last_ds_roster_compare,
            run_ds_roster_compare_job,
        )

        job = run_ds_roster_compare_job(matrix_path, parts_dir)
        self.assertTrue(job.success)
        self.assertTrue(job.message.startswith("ОК:"))
        stored = get_last_ds_roster_compare()
        self.assertIsNotNone(stored)
        assert stored is not None
        self.assertEqual(stored.matrix_path, matrix_path)
        by_copy = {row.copy_name: row for row in stored.rows}
        self.assertEqual(by_copy["ДС1"].manager, "ГЭМ")
        self.assertTrue(stored.is_ok)


def _row_with_mto_path(path: str) -> RowStd:
    row = RowStd(t_com=TableComments())
    row.t_com.file_full_path = path
    return row


if __name__ == "__main__":
    unittest.main()
