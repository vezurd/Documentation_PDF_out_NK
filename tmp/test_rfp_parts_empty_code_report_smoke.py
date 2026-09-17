"""Smoke tests for the empty-CODE sidecar xlsx (collect + stamp + Launch copy)."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.rfp_parts.analyze_rfp_parts import (
    PARTS_KIND,
    _extract_records,
    _write_empty_code_xlsx,
)
from RFQ.rfp_parts.file_status import (
    EMPTY_CODE_XLSX_NAME,
    copy_empty_code_report_to_result_dir,
)

_HEADER = [
    "№ п/п",
    "Титул",
    "Спецификация",
    "Tag",
    "Код 1С",
    "Код РД",
    "Наименование МТР",
    "Технические характеристики",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
    "Кол-во",
    "Ед. изм",
]
_DATA = [
    1,
    "8525-SKUD",
    "spec",
    "T-1",
    "",
    "BCC0000001",
    "Кабель",
    "mark",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
    2,
    "шт",
]


def _write_book(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    wb = Workbook()
    default = wb.active
    names = list(sheets)
    default.title = names[0]
    for title, rows in sheets.items():
        ws = default if title == names[0] else wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(path)
    wb.close()


class RfpPartsEmptyCodeReportSmokeTest(unittest.TestCase):
    def test_coded_and_empty_code_rows(self) -> None:
        empty_row = list(_DATA)
        empty_row[0] = 2
        empty_row[3] = "T-EMPTY"
        empty_row[5] = ""
        empty_row[6] = "Без кода"
        empty_row[16] = 5
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "ДС7_ГФ.xlsx"
            _write_book(path, {"Перечень материалов": [_HEADER, _DATA, empty_row]})
            warnings: list[tuple[str, str, str]] = []
            skipped: list = []
            records, _stats = _extract_records(
                path, PARTS_KIND, warnings, skipped_empty_code=skipped
            )
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].code, "BCC0000001")
            self.assertEqual(len(skipped), 1)
            snap = skipped[0]
            self.assertEqual(snap.file_name, path.name)
            self.assertEqual(snap.name, "Без кода")
            self.assertEqual(snap.values_text, "5")
            self.assertEqual(snap.excel_row, 3)
            warn_msgs = [msg for _level, _fname, msg in warnings]
            self.assertTrue(
                any(
                    "пропущена, нет полей" in msg and "Код РД" in msg
                    for msg in warn_msgs
                )
            )

            out_path = Path(raw) / EMPTY_CODE_XLSX_NAME
            written = _write_empty_code_xlsx(out_path, skipped)
            self.assertEqual(written, 1)
            self.assertTrue(out_path.is_file())
            wb = load_workbook(out_path, read_only=True, data_only=True)
            try:
                ws = wb["Без кода"]
                data_rows = list(ws.iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(data_rows), 1)
                blob = " ".join(
                    "" if cell is None else str(cell)
                    for row in data_rows
                    for cell in row
                )
                self.assertNotIn("BCC0000001", blob)
            finally:
                wb.close()

    def test_missing_name_with_code_is_not_collected(self) -> None:
        missing_name = list(_DATA)
        missing_name[6] = ""
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "ДС8.xlsx"
            _write_book(path, {"Перечень материалов": [_HEADER, missing_name]})
            warnings: list[tuple[str, str, str]] = []
            skipped: list = []
            records, _stats = _extract_records(
                path, PARTS_KIND, warnings, skipped_empty_code=skipped
            )
            self.assertEqual(records, [])
            self.assertEqual(skipped, [])
            self.assertTrue(
                any("пропущена, нет полей" in msg for _, _, msg in warnings)
            )

    def test_copy_helper_copies_and_missing_is_none(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            stamp = Path(raw) / "stamp"
            result = Path(raw) / "result"
            stamp.mkdir()
            src = stamp / EMPTY_CODE_XLSX_NAME
            wb = Workbook()
            wb.active.title = "Без кода"
            wb.save(src)
            wb.close()
            dest = copy_empty_code_report_to_result_dir(stamp, result)
            self.assertIsNotNone(dest)
            assert dest is not None
            self.assertTrue(dest.is_file())
            self.assertEqual(dest.name, EMPTY_CODE_XLSX_NAME)
            self.assertEqual(dest.parent, result)
            missing = copy_empty_code_report_to_result_dir(
                Path(raw) / "no_stamp", result
            )
            self.assertIsNone(missing)

    def test_empty_list_not_written_and_missing_copy_is_none(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            out_path = Path(raw) / EMPTY_CODE_XLSX_NAME
            written = _write_empty_code_xlsx(out_path, [])
            self.assertEqual(written, 0)
            self.assertFalse(out_path.exists())
            self.assertIsNone(
                copy_empty_code_report_to_result_dir(Path(raw), Path(raw) / "dest")
            )


if __name__ == "__main__":
    unittest.main(verbosity=2)
