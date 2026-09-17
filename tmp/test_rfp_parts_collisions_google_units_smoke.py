"""Offline smoke: Google UNITS column and source-row fix list in collisions xlsx."""

from __future__ import annotations

import sys
import unittest
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from RFQ.rfp_parts.analyze_rfp_parts import (
    CoarseCollisionResult,
    RfpRecord,
    _build_google_code_lookup,
    _collect_units_fix_rows,
    _compute_coarse_collisions,
    _google_units_for_code,
    _google_units_matches_variants,
    _write_collisions_xlsx,
)
from base.tables_columns import CODE, EQUIPMENT_CODE, UNITS


def _row(**values: object) -> SimpleNamespace:
    return SimpleNamespace(
        el={key: SimpleNamespace(value=value) for key, value in values.items()}
    )


def _record(
    *,
    code: str,
    units: str,
    name: str = "Cable",
    file_name: str = "DS10.xlsx",
    sheet: str = "Перечень материалов",
    excel_row: int = 24,
    ds_name: str = "ДС10",
    ds_title: str = "8529-SOS",
) -> RfpRecord:
    return RfpRecord(
        kind="parts",
        file_name=file_name,
        sheet=sheet,
        excel_row=excel_row,
        ds_name=ds_name,
        ds_number="10",
        ds_title=ds_title,
        ds_specification=ds_title,
        tags="",
        ds_code_1c="",
        code=code,
        name=name,
        type_mark="",
        values=Decimal("1"),
        units=units,
    )


class GoogleUnitsCollisionsSmokeTest(unittest.TestCase):
    def test_lookup_indexes_equipment_and_units_by_bcc(self) -> None:
        rows = [
            _row(**{CODE: "BCC0000516", EQUIPMENT_CODE: "KA", UNITS: "м"}),
            _row(**{CODE: " bcc0000516 ", EQUIPMENT_CODE: "ignored", UNITS: "шт"}),
            _row(**{CODE: "BCC0000999", EQUIPMENT_CODE: "XB", UNITS: "шт."}),
        ]
        lookup = _build_google_code_lookup(rows)
        self.assertEqual(lookup.equipment_by_code["BCC0000516"], "KA")
        self.assertEqual(lookup.units_by_code["BCC0000516"], "м")
        self.assertEqual(lookup.units_by_code["BCC0000999"], "шт.")

    def test_units_compare_ignores_trailing_dot(self) -> None:
        self.assertTrue(_google_units_matches_variants("шт.", ("шт", "м")))
        self.assertFalse(_google_units_matches_variants("кг", ("шт", "м")))

    def test_fix_rows_keep_only_source_lines_not_matching_google(self) -> None:
        wrong = _record(
            code="BCC0000516",
            units="шт",
            file_name="ДС17.xlsx",
            excel_row=41,
            ds_name="ДС17",
        )
        correct = _record(
            code="BCC0000516",
            units="м",
            file_name="ДС76.xlsx",
            excel_row=18,
            ds_name="ДС76",
        )
        dotted = _record(
            code="BCC0000516",
            units="м.",
            file_name="ДС17.xlsx",
            excel_row=90,
            ds_name="ДС17",
        )
        coarse = _compute_coarse_collisions([wrong, correct, dotted])
        self.assertEqual(len(coarse.ambiguous_meta), 1)
        fix_rows = _collect_units_fix_rows(coarse, {"BCC0000516": "м"})
        locations = [
            (item.record.file_name, item.record.excel_row, item.file_units)
            for item in fix_rows
        ]
        self.assertEqual(
            locations,
            [("ДС17.xlsx", 41, "шт"), ("ДС17.xlsx", 90, "м.")],
        )

    def test_collisions_xlsx_lists_file_sheet_row(self) -> None:
        wrong = _record(
            code="BCC0000516",
            units="шт",
            file_name="ДС17.xlsx",
            sheet="Перечень материалов",
            excel_row=41,
            ds_name="ДС17",
        )
        correct = _record(
            code="BCC0000516",
            units="м",
            file_name="ДС76.xlsx",
            excel_row=18,
            ds_name="ДС76",
        )
        coarse = CoarseCollisionResult(
            qty_by_key={("8529-SOS", "BCC0000516"): Decimal("2")},
            examples={("8529-SOS", "BCC0000516"): wrong},
            ambiguous_meta=[
                (("8529-SOS", "BCC0000516"), wrong, ("м", "шт")),
            ],
            records_by_key={
                ("8529-SOS", "BCC0000516"): [wrong, correct],
            },
        )
        path = ROOT / "tmp" / "_smoke_rfp_parts_collisions.xlsx"
        try:
            rows = _write_collisions_xlsx(
                path, coarse, units_by_code={"BCC0000516": "м"}
            )
            self.assertEqual(rows, 1)
            wb = load_workbook(path)
            try:
                ws = wb["К исправлению"]
                self.assertEqual(
                    [ws.cell(1, col).value for col in range(1, 9)],
                    [
                        "Файл",
                        "Лист",
                        "Строка",
                        "Имя ДС",
                        "Титул-система",
                        "Код",
                        "Ед. изм. (в файле)",
                        "Ед. изм. (Google)",
                    ],
                )
                self.assertEqual(ws.cell(2, 1).value, "ДС17.xlsx")
                self.assertEqual(ws.cell(2, 2).value, "Перечень материалов")
                self.assertEqual(ws.cell(2, 3).value, 41)
                self.assertEqual(ws.cell(2, 7).value, "шт")
                self.assertEqual(ws.cell(2, 8).value, "м")
                self.assertEqual(ws.max_row, 2)

                keys = wb["Ключи"]
                self.assertEqual(keys.cell(1, 3).value, "Ед. изм. (варианты)")
                self.assertEqual(keys.cell(2, 3).value, "м | шт")
                self.assertEqual(keys.cell(2, 5).value, 1)

                summary = {
                    row[0]: row[1]
                    for row in wb["Сводка"].iter_rows(min_row=2, values_only=True)
                    if row[0]
                }
                self.assertEqual(summary["Строк к исправлению"], 1)
                self.assertEqual(summary["Файлов с ошибкой"], 1)
            finally:
                wb.close()
        finally:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass

    def test_google_units_lookup_uses_normalized_code(self) -> None:
        self.assertEqual(
            _google_units_for_code(" bcc0000516 ", {"BCC0000516": "м"}),
            "м",
        )


if __name__ == "__main__":
    unittest.main()
