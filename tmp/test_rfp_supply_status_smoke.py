"""Smoke: optional parts «Исключен из поставки» → net slot 15 → Step4 skip."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import base.t_comm_initial_classes as t_com_init_cls
from RFQ.packing_list_provider import (
    PACKING_CACHE_VERSION,
    PackingCacheMeta,
    PackingDataset,
    PackingQualityLevel,
)
from RFQ.rfp_parts.analyze_rfp_parts import (
    PARTS_KIND,
    _build_summary_rows,
    _build_unit_counters,
    _compute_coarse_collisions,
    _extract_records,
    _write_net_xlsx,
)
from RFQ.tags_rfp_compare.rfp_supply_status import (
    CANONICAL_EXCLUDED_FROM_SUPPLY,
    KIND_EXCLUDED,
    KIND_TAG_REPLACED,
    KIND_UNKNOWN,
    RfpSupplyStatusError,
    RfpSupplyStatusIssue,
    apply_rfp_supply_status_on_rows,
    format_tag_replaced_status,
    is_excluded_from_supply,
    parse_rfp_supply_status,
    raise_if_rfp_supply_status_issues,
)
from RFQ.tags_rfp_compare.step4.step4_2_match_rfp_with_mto import match_rfp_with_mto
from RFQ.tags_rfp_compare.step4.step4_6_save_match_result_to_excel import (
    OUTPUT_COLUMNS_CONFIG,
)
from RFQ.tags_rfp_compare.step4.step4_packing_compare import (
    STATUS_COMPLETE,
    STATUS_EXCLUDED_FROM_SUPPLY,
    build_ordered_snapshot,
    compare_rfp_rows_with_packing,
    rfp_packing_match_key,
)
from RFQ.tags_rfp_compare.step4.step4_postmerge_ops import apply_mto_ul_quantity_diff
from RFQ.tags_rfp_compare.step4_analyze_and_match import (
    _collapse_rfp_rows_before_export,
)
from base.base_classes import RowStd, RowType, TableComments
from base.base_mto import get_std_from_excel_file
from base.tables_columns import (
    ANNOTATION,
    CODE,
    CODE_MTO,
    ColNames,
    DS_NAME,
    DS_NUMBER,
    DS_SYSTEM,
    DS_TITLE,
    MATCH_STATUS,
    NAME,
    RFP_SUPPLY_STATUS,
    TAGS,
    TITLE,
    TYPE_MARK,
    UL_COMPARE_STATUS,
    UL_VALUES,
    UNITS,
    VALUES,
    VALUES_2,
    VALUES_MTO,
    VALUES_MTO_UL_DIFF,
    VENDOR,
)

_CODE = "BCC0002946"
_TITLE = "8950-POS5"
_TAG_EXCL = "8950-FE-05-S-FV-0201"
_TAG_LIVE = "8950-FE-05-S-FV-0202"


def _header_row(*, status_label: str | None) -> list[object]:
    row = [""] * 20
    row[0] = "№ п/п"
    row[1] = "Титул"
    row[2] = "Спецификация"
    row[3] = "Линия, Tag-номер"
    row[4] = "Код 1С"
    row[5] = "Код РД"
    row[6] = "Наименование МТР"
    row[7] = "Технические характеристики"
    row[16] = "Кол-во"
    row[17] = "Ед. изм"
    if status_label is not None:
        row[19] = status_label
    return row


def _data_row(
    *,
    number: int,
    tag: str,
    status: object = "",
    price: object = "",
) -> list[object]:
    row = [""] * 20
    row[0] = number
    row[1] = _TITLE
    row[2] = "AGCC.287-8950-POS5.MTO-0001"
    row[3] = tag
    row[5] = _CODE
    row[6] = "ГИР 4/250M (LT)"
    row[7] = "ГИР 4/250M (LT)"
    row[16] = 1
    row[17] = "шт"
    row[19] = status if status != "" else price
    return row


def _write_parts_xlsx(
    path: Path,
    *,
    status_label: str | None,
    rows: list[list[object]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Перечень материалов"
    ws.append(["титульный"])
    ws.append(["шапка 2"])
    ws.append(["шапка 3"])
    ws.append(_header_row(status_label=status_label))
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def _rfp_row(
    *,
    tag: str,
    excluded: bool,
    ds_number: str = "1",
    quantity: object = 1,
) -> RowStd:
    row = RowStd.get_std_check_row({}, TableComments())
    row.row_type = RowType.position_row
    row.el[DS_NAME].value = "ДС29"
    row.el[DS_NUMBER].value = ds_number
    row.el[DS_TITLE].value = _TITLE
    row.el[CODE].value = _CODE
    row.el[TAGS].value = tag
    row.el[VALUES].value = quantity
    row.el[UNITS].value = "шт"
    row.el[NAME].value = "ГИР 4/250M (LT)"
    row.el[TYPE_MARK].value = "TM"
    row.el[VENDOR].value = "Vendor"
    if excluded:
        row.el[RFP_SUPPLY_STATUS].value = CANONICAL_EXCLUDED_FROM_SUPPLY
    return row


def _mto_row(*, tag: str, quantity: object = 1) -> RowStd:
    row = RowStd.get_std_check_row({}, TableComments())
    row.row_type = RowType.position_row
    row.el[DS_TITLE].value = _TITLE
    row.el[CODE].value = _CODE
    row.el[TAGS].value = tag
    row.el[VALUES].value = quantity
    row.el[UNITS].value = "шт"
    row.el[NAME].value = "MTO GIR"
    return row


def _packing_row(*, tag: str) -> RowStd:
    row = RowStd.get_std_check_row({}, TableComments())
    row.row_type = RowType.position_row
    row.el[DS_TITLE].value = "8950"
    row.el[DS_SYSTEM].value = "POS5"
    row.el[TITLE].value = "Single 1"
    row.el[CODE].value = _CODE
    row.el[TAGS].value = tag
    row.el[VALUES].value = 1
    row.el[UNITS].value = "шт"
    row.el[NAME].value = "UL GIR"
    row.el[ANNOTATION].value = "согл УЛ ДС29/packing-a.xlsx"
    row._packing_source_row = 10
    return row


def _dataset(rows: list[RowStd]) -> PackingDataset:
    return PackingDataset(
        rows=rows,
        meta=PackingCacheMeta(
            version=PACKING_CACHE_VERSION,
            created_at="2026-09-03T12:00:00+00:00",
            root="C:/packing",
            fingerprint="supply-status-smoke",
            counters={"position_rows": len(rows)},
            critical_ok=True,
            report_paths={"critical": "tsd_packing_critical.txt"},
        ),
        quality=PackingQualityLevel.OK,
        issues=[],
        cache_path="memory.cache",
    )


class RfpSupplyStatusSmokeTest(unittest.TestCase):
    def test_column_is_in_rfp_aggregated_slot_15_and_excel_config(self) -> None:
        extra = {
            int(idx): name
            for idx, name in ColNames.RFP_AGGREAGATED.column_dict.items()
            if isinstance(idx, int) and idx > 10
        }
        self.assertEqual(extra[15], RFP_SUPPLY_STATUS)
        self.assertEqual(extra[17], VALUES_2)
        visible = [
            item.col_name
            for item in OUTPUT_COLUMNS_CONFIG
            if item.output and item.col_name == RFP_SUPPLY_STATUS
        ]
        self.assertEqual(visible, [RFP_SUPPLY_STATUS])
        label = next(
            item.header_label
            for item in OUTPUT_COLUMNS_CONFIG
            if item.col_name == RFP_SUPPLY_STATUS
        )
        self.assertEqual(label, "Статус поставки")
        self.assertNotEqual(label, "MTO, Статус позиции")

    def test_parts_header_writes_slot_15_and_skips_price_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            parts_dir = Path(tmp)
            ds29 = parts_dir / "ДС29. AGCC.287-0000-12.4.1-RFP-0018_01_RU.xlsx"
            _write_parts_xlsx(
                ds29,
                status_label="MTO, Статус позиции",
                rows=[
                    _data_row(number=1, tag=_TAG_EXCL, status=CANONICAL_EXCLUDED_FROM_SUPPLY),
                    _data_row(number=2, tag=_TAG_LIVE, status=""),
                ],
            )
            price_book = parts_dir / "ДС10. price.xlsx"
            _write_parts_xlsx(
                price_book,
                status_label="Цена",
                rows=[_data_row(number=1, tag="8950-FE-01-S-FV-0001", price=12345)],
            )

            warnings: list[tuple[str, str, str]] = []
            ds29_records, ds29_stats = _extract_records(ds29, PARTS_KIND, warnings)
            self.assertEqual(ds29_stats.rows, 2)
            by_tag = {item.tags: item for item in ds29_records}
            self.assertEqual(
                by_tag[_TAG_EXCL].rfp_supply_status,
                CANONICAL_EXCLUDED_FROM_SUPPLY,
            )
            self.assertEqual(by_tag[_TAG_LIVE].rfp_supply_status, "")

            price_records, price_stats = _extract_records(price_book, PARTS_KIND, warnings)
            self.assertEqual(price_stats.rows, 1)
            self.assertEqual(price_records[0].rfp_supply_status, "")
            self.assertFalse(
                any("некорректный статус позиции" in item[2] for item in warnings)
            )

            aggregation = _build_unit_counters(ds29_records, warnings)
            net_rows = _build_summary_rows(
                aggregation, _compute_coarse_collisions(ds29_records)
            )
            net_path = parts_dir / "rfp_parts_net.xlsx"
            self.assertEqual(_write_net_xlsx(net_path, net_rows), 2)

            wb = load_workbook(net_path, read_only=True, data_only=True)
            ws = wb["Перечень материалов"]
            header, *data = list(ws.iter_rows(values_only=True))
            wb.close()
            self.assertEqual(header[15], "Статус поставки")
            self.assertEqual(header[20], "Источник строки")
            by_net_tag = {row[4]: row for row in data}
            self.assertEqual(by_net_tag[_TAG_EXCL][15], CANONICAL_EXCLUDED_FROM_SUPPLY)
            self.assertIn(by_net_tag[_TAG_LIVE][15], (None, ""))

            t_com = TableComments(
                file_full_path=str(net_path),
                dir_path="-1",
                tabel_class=t_com_init_cls.RFP_AGGREGATED,
            )
            loaded = get_std_from_excel_file(t_com)
            pos = [
                row for row in loaded if row.row_type == RowType.position_row
            ]
            loaded_by_tag = {
                (row.get_tags_list() or [""])[0]: row for row in pos
            }
            self.assertTrue(is_excluded_from_supply(loaded_by_tag[_TAG_EXCL]))
            self.assertFalse(is_excluded_from_supply(loaded_by_tag[_TAG_LIVE]))

    def test_step4_match_and_packing_skip_excluded_but_keep_live(self) -> None:
        excluded = _rfp_row(tag=_TAG_EXCL, excluded=True)
        live = _rfp_row(tag=_TAG_LIVE, excluded=False, ds_number="2")
        mto_excl = _mto_row(tag=_TAG_EXCL)
        mto_live = _mto_row(tag=_TAG_LIVE)

        _matched, unmatched_mto = match_rfp_with_mto(
            [excluded, live],
            {_TITLE: [mto_excl, mto_live]},
        )
        self.assertEqual(live.get_value(MATCH_STATUS), "Тег сопоставлен")
        self.assertEqual(live.get_value(CODE_MTO), _CODE)
        self.assertEqual(live.get_value(VALUES_MTO), 1)
        self.assertNotEqual(excluded.get_value(MATCH_STATUS), "Тег сопоставлен")
        self.assertFalse(excluded.get_value(CODE_MTO))
        self.assertTrue(is_excluded_from_supply(excluded))
        self.assertIn(mto_excl, unmatched_mto)
        self.assertNotIn(mto_live, unmatched_mto)

        snapshot = build_ordered_snapshot([excluded, live])
        self.assertEqual(
            snapshot,
            {rfp_packing_match_key("29", "8950", "POS5", _CODE): 1.0},
        )
        packing = [
            _packing_row(tag=_TAG_EXCL),
            _packing_row(tag=_TAG_LIVE),
        ]
        rows, _audit = compare_rfp_rows_with_packing(
            [excluded, live],
            snapshot,
            _dataset(packing),
        )
        self.assertEqual(excluded.get_value(UL_COMPARE_STATUS), STATUS_EXCLUDED_FROM_SUPPLY)
        self.assertEqual(
            excluded.get_value(UL_COMPARE_STATUS),
            CANONICAL_EXCLUDED_FROM_SUPPLY,
        )
        self.assertIn(excluded.get_value(UL_VALUES), (None, ""))
        self.assertEqual(live.get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(live.get_value(UL_VALUES), 1.0)
        leftover = [
            row
            for row in rows
            if row is not live and row is not excluded
        ]
        self.assertTrue(leftover)

        apply_mto_ul_quantity_diff([excluded, live])
        self.assertIn(excluded.get_value(VALUES_MTO_UL_DIFF), (None, ""))
        self.assertEqual(live.get_value(VALUES_MTO_UL_DIFF), 0)

    def test_collapse_does_not_mix_excluded_no_tag_with_live(self) -> None:
        excluded = _rfp_row(tag="", excluded=True, quantity=1)
        live = _rfp_row(tag="", excluded=False, quantity=1)
        collapsed = _collapse_rfp_rows_before_export([excluded, live])
        self.assertEqual(len(collapsed), 2)
        flags = {is_excluded_from_supply(row) for row in collapsed}
        self.assertEqual(flags, {True, False})

    def test_parse_known_and_unknown_supply_status(self) -> None:
        self.assertEqual(parse_rfp_supply_status("").kind, "empty")
        self.assertEqual(parse_rfp_supply_status("  ").kind, "empty")
        excluded = parse_rfp_supply_status("исключен из поставки")
        self.assertEqual(excluded.kind, KIND_EXCLUDED)
        replaced = parse_rfp_supply_status(
            "8950-FE-05-S-BZ-8189 тег заменен на 8950-FE-05-S-BZ-1811"
        )
        self.assertEqual(replaced.kind, KIND_TAG_REPLACED)
        self.assertEqual(replaced.old_tag, "8950-FE-05-S-BZ-8189")
        self.assertEqual(replaced.new_tag, "8950-FE-05-S-BZ-1811")
        yo = parse_rfp_supply_status(
            "8950-FE-05-S-BZ-8189 тег заменён на 8950-FE-05-S-BZ-1811"
        )
        self.assertEqual(yo.kind, KIND_TAG_REPLACED)
        self.assertEqual(yo.new_tag, "8950-FE-05-S-BZ-1811")
        unknown = parse_rfp_supply_status("снято с поставки")
        self.assertEqual(unknown.kind, KIND_UNKNOWN)
        self.assertEqual(unknown.text, "снято с поставки")

    def test_parts_tag_replacement_rewrites_tag_before_net(self) -> None:
        old_tag = "8950-FE-05-S-BZ-8189"
        new_tag = "8950-FE-05-S-BZ-1811"
        status = f"{old_tag} тег заменен на {new_tag}"
        with tempfile.TemporaryDirectory() as tmp:
            parts_dir = Path(tmp)
            ds29 = parts_dir / "ДС29. AGCC.287-0000-12.4.1-RFP-0018_01_RU.xlsx"
            _write_parts_xlsx(
                ds29,
                status_label="MTO, Статус позиции",
                rows=[
                    _data_row(number=1, tag=old_tag, status=status),
                    _data_row(number=2, tag=_TAG_LIVE, status=""),
                ],
            )
            warnings: list[tuple[str, str, str]] = []
            issues: list[RfpSupplyStatusIssue] = []
            records, stats = _extract_records(
                ds29, PARTS_KIND, warnings, status_issues=issues
            )
            self.assertEqual(stats.rows, 2)
            self.assertEqual(issues, [])
            by_tag = {item.tags: item for item in records}
            self.assertIn(new_tag, by_tag)
            self.assertNotIn(old_tag, by_tag)
            self.assertEqual(
                by_tag[new_tag].rfp_supply_status,
                format_tag_replaced_status(old_tag, new_tag),
            )

            aggregation = _build_unit_counters(records, warnings)
            net_rows = _build_summary_rows(
                aggregation, _compute_coarse_collisions(records)
            )
            net_path = parts_dir / "rfp_parts_net.xlsx"
            self.assertEqual(_write_net_xlsx(net_path, net_rows), 2)
            wb = load_workbook(net_path, read_only=True, data_only=True)
            ws = wb["Перечень материалов"]
            _header, *data = list(ws.iter_rows(values_only=True))
            wb.close()
            by_net_tag = {row[4]: row for row in data}
            self.assertEqual(
                by_net_tag[new_tag][15],
                format_tag_replaced_status(old_tag, new_tag),
            )
            self.assertIn(by_net_tag[_TAG_LIVE][15], (None, ""))

    def test_unknown_status_is_error_and_stops(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            parts_dir = Path(tmp)
            ds29 = parts_dir / "ДС29. AGCC.287-0000-12.4.1-RFP-0018_01_RU.xlsx"
            _write_parts_xlsx(
                ds29,
                status_label="MTO, Статус позиции",
                rows=[
                    _data_row(number=1, tag=_TAG_LIVE, status="снято с поставки"),
                    _data_row(number=2, tag=_TAG_EXCL, status=""),
                ],
            )
            warnings: list[tuple[str, str, str]] = []
            issues: list[RfpSupplyStatusIssue] = []
            records, _stats = _extract_records(
                ds29, PARTS_KIND, warnings, status_issues=issues
            )
            self.assertEqual([item.tags for item in records], [_TAG_EXCL])
            self.assertEqual(len(issues), 1)
            self.assertEqual(issues[0].code, _CODE)
            self.assertEqual(issues[0].raw, "снято с поставки")
            self.assertEqual(issues[0].file_name, ds29.name)
            self.assertTrue(
                any("некорректный статус позиции" in item[2] for item in warnings)
            )
            with self.assertRaises(RfpSupplyStatusError) as ctx:
                raise_if_rfp_supply_status_issues(issues)
            message = str(ctx.exception)
            self.assertIn(ds29.name, message)
            self.assertIn("строка", message)
            self.assertIn(_CODE, message)
            self.assertIn("снято с поставки", message)

    def test_tag_replacement_mismatch_is_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            ds29 = Path(tmp) / "ДС29. AGCC.xlsx"
            _write_parts_xlsx(
                ds29,
                status_label="MTO, Статус позиции",
                rows=[
                    _data_row(
                        number=1,
                        tag="8950-FE-05-S-BZ-0001",
                        status=(
                            "8950-FE-05-S-BZ-8189 тег заменен на "
                            "8950-FE-05-S-BZ-1811"
                        ),
                    ),
                ],
            )
            warnings: list[tuple[str, str, str]] = []
            issues: list[RfpSupplyStatusIssue] = []
            records, _stats = _extract_records(
                ds29, PARTS_KIND, warnings, status_issues=issues
            )
            self.assertEqual(records, [])
            self.assertEqual(len(issues), 1)
            self.assertIn("не содержит заменяемый тег", issues[0].detail)

    def test_step1_apply_replaces_tag_and_matches_mto(self) -> None:
        old_tag = "8950-FE-05-S-BZ-8189"
        new_tag = "8950-FE-05-S-BZ-1811"
        row = _rfp_row(tag=old_tag, excluded=False)
        row._xlsx_source_row = 15
        row.el[RFP_SUPPLY_STATUS].value = f"{old_tag} тег заменен на {new_tag}"
        apply_rfp_supply_status_on_rows([row], file_name="ДС29.xlsx")
        self.assertEqual(row.el[TAGS].value, new_tag)
        self.assertEqual(
            row.el[RFP_SUPPLY_STATUS].value,
            format_tag_replaced_status(old_tag, new_tag),
        )
        self.assertFalse(is_excluded_from_supply(row))
        mto = _mto_row(tag=new_tag)
        _matched, unmatched = match_rfp_with_mto([row], {_TITLE: [mto]})
        self.assertEqual(row.get_value(MATCH_STATUS), "Тег сопоставлен")
        self.assertNotIn(mto, unmatched)

        bad = _rfp_row(tag=_TAG_LIVE, excluded=False)
        bad._xlsx_source_row = 22
        bad.el[RFP_SUPPLY_STATUS].value = "непонятный статус"
        with self.assertRaises(RfpSupplyStatusError) as ctx:
            apply_rfp_supply_status_on_rows([bad], file_name="ДС10.xlsx")
        message = str(ctx.exception)
        self.assertIn("ДС10.xlsx", message)
        self.assertIn("22", message)
        self.assertIn(_CODE, message)
        self.assertIn("непонятный статус", message)


if __name__ == "__main__":
    unittest.main()
