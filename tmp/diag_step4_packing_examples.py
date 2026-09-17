"""Read-only Step4 packing examples (ДС82/19048, ДС75/17250-17251).

Finds newest Step4 xlsx under result_dir_base from config, extracts target rows
and related «Только в УЛ» leftovers. Does not modify production code.

Run from repo root:
  set PYTHONUTF8=1
  python tmp/diag_step4_packing_examples.py
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "RFQ" / "tags_rfp_compare" / "rfp_tags_compare_config.json"

STEP4_GLOB = "Шаг4_Сопоставление_RFP_MTO_*.xlsx"
FOLDER_PREFIX = "_результат_проверки_"

WANTED_FIELDS = (
    "excel_row",
    "Имя ДС",
    "№ позиции",
    "Титул/Марка",
    "Код RFP",
    "Код MTO",
    "Кол-во RFP",
    "Статус УЛ",
    "Статус тегов УЛ",
    "Теги УЛ",
    "Код УЛ",
    "Источник УЛ (файл · вкладка · строка)",
    "Теги RFP",
    "Наименование RFP",
    "Наименование УЛ",
    "MTO, Статус позиции",
)

LEFTOVER_CODES = frozenset({"BCC0003052", "BCC0000152"})
LEFTOVER_TAG_PARTS = (
    "1001",
    "1002",
    "8445-GA-01-S-FV-0011",
)


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def _find_latest_step4(base: Path) -> tuple[Path | None, list[str]]:
    notes: list[str] = []
    try:
        if not base.exists():
            notes.append(f"result_dir_base missing: {base}")
            return None, notes
        if not base.is_dir():
            notes.append(f"result_dir_base not a directory: {base}")
            return None, notes
    except OSError as exc:
        notes.append(f"result_dir_base OSError: {exc!r}")
        return None, notes

    candidates: list[tuple[float, Path, Path]] = []
    try:
        for child in base.iterdir():
            if not child.is_dir():
                continue
            name = child.name
            if not name.lower().startswith(FOLDER_PREFIX):
                continue
            for xlsx in child.glob(STEP4_GLOB):
                try:
                    mtime = xlsx.stat().st_mtime
                except OSError:
                    continue
                candidates.append((mtime, child, xlsx))
    except OSError as exc:
        notes.append(f"iterdir OSError: {exc!r}")
        return None, notes

    if not candidates:
        notes.append(f"no folders matching {FOLDER_PREFIX}* with {STEP4_GLOB}")
        return None, notes

    candidates.sort(key=lambda t: t[0], reverse=True)
    _, folder, xlsx = candidates[0]
    notes.append(f"newest folder: {folder.name}")
    notes.append(f"newest xlsx: {xlsx.name} mtime={candidates[0][0]:.0f}")
    if len(candidates) > 1:
        notes.append(f"total Step4 xlsx candidates: {len(candidates)}")
    return xlsx, notes


def _match_example1(row: dict[str, str]) -> bool:
    ds = row.get("Имя ДС", "")
    pos = _cell_text(row.get("№ позиции", ""))
    tags = row.get("Теги RFP", "")
    return ds.strip() == "ДС82" and pos == "19048" and "8445-GA-01-S-FV-0011" in tags


def _match_example2(row: dict[str, str]) -> bool:
    ds = row.get("Имя ДС", "")
    pos = _cell_text(row.get("№ позиции", ""))
    code = _cell_text(row.get("Код RFP", "")).upper()
    if code != "BCC0000152":
        return False
    if pos not in {"17250", "17251"}:
        return False
    return "ДС75" in ds


def _match_leftover(row: dict[str, str]) -> bool:
    status = row.get("Статус УЛ", "")
    if status != "Только в УЛ":
        return False
    code_ul = _cell_text(row.get("Код УЛ", "")).upper()
    if code_ul in LEFTOVER_CODES:
        return True
    blob = " | ".join(
        row.get(k, "") for k in ("Теги УЛ", "Теги RFP", "Тег MTO", "Тег (VO/MTO/RFP)")
    )
    return any(part in blob for part in LEFTOVER_TAG_PARTS)


def _ul_comment(cell) -> str:
    c = getattr(cell, "comment", None)
    if c is None:
        return ""
    if isinstance(c, Comment):
        return _cell_text(c.text)
    return _cell_text(c)


def _print_row(label: str, row: dict[str, str]) -> None:
    _safe_print(f"\n--- {label} (excel_row={row.get('excel_row', '?')}) ---")
    for field in WANTED_FIELDS:
        if field == "excel_row":
            continue
        val = row.get(field, "")
        if field == "Источник УЛ (файл · вкладка · строка)":
            key = field
        elif field == "Источник УЛ":
            key = "Источник УЛ (файл · вкладка · строка)"
        else:
            key = field
        if key in row:
            _safe_print(f"  {field}: {row.get(key, '')}")
    comment = row.get("_ul_comment", "")
    if comment:
        _safe_print(f"  Статус УЛ comment: {comment}")


def main() -> int:
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
    base = Path(cfg["paths"]["result_dir_base"])
    _safe_print(f"result_dir_base: {base}")

    xlsx, find_notes = _find_latest_step4(base)
    for note in find_notes:
        _safe_print(note)
    if xlsx is None:
        return 1

    try:
        if not xlsx.exists():
            _safe_print(f"xlsx missing: {xlsx}")
            return 1
    except OSError as exc:
        _safe_print(f"xlsx exists() OSError: {exc!r}")
        return 1

    _safe_print(f"\nloading: {xlsx}")
    wb = load_workbook(xlsx, read_only=False, data_only=True)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]
    _safe_print(f"sheet: {sheet_name}")

    col_of: dict[str, int] = {}
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    for idx, cell in enumerate(header_row):
        name = _cell_text(cell.value)
        if name and name not in col_of:
            col_of[name] = idx

    src_col = col_of.get("Источник УЛ (файл · вкладка · строка)")
    ul_status_col = col_of.get("Статус УЛ")

    ex1: list[dict[str, str]] = []
    ex2: list[dict[str, str]] = []
    leftovers: list[dict[str, str]] = []

    for excel_row, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row: dict[str, str] = {"excel_row": str(excel_row)}
        for name, idx in col_of.items():
            if idx < len(cells):
                row[name] = _preview_cell(cells[idx].value)
        if ul_status_col is not None and ul_status_col < len(cells):
            row["_ul_comment"] = _ul_comment(cells[ul_status_col])

        if _match_example1(row):
            ex1.append(row)
        if _match_example2(row):
            ex2.append(row)
        if _match_leftover(row):
            leftovers.append(row)

    wb.close()

    _safe_print("\n========== EXAMPLE 1: ДС82 pos 19048 tag 8445-GA-01-S-FV-0011 ==========")
    if not ex1:
        _safe_print("  (no matching rows)")
    for i, row in enumerate(ex1, 1):
        _print_row(f"example1 #{i}", row)

    _safe_print("\n========== EXAMPLE 2: ДС75 pos 17250/17251 code BCC0000152 ==========")
    if not ex2:
        _safe_print("  (no matching rows)")
    for i, row in enumerate(sorted(ex2, key=lambda r: r.get("№ позиции", "")), 1):
        _print_row(f"example2 #{i}", row)

    _safe_print("\n========== LEFTOVER «Только в УЛ» (watch codes/tags) ==========")
    _safe_print(f"count: {len(leftovers)}")
    for i, row in enumerate(leftovers, 1):
        _print_row(f"leftover #{i}", row)

    _safe_print("\n========== INTERPRETATION HINTS ==========")
    if ex1:
        r = ex1[0]
        _safe_print(
            f"Ex1: title={r.get('Титул/Марка')!r} status_ul={r.get('Статус УЛ')!r} "
            f"code_ul={r.get('Код УЛ')!r}"
        )
    if ex2:
        for r in ex2:
            _safe_print(
                f"Ex2 pos {r.get('№ позиции')}: title={r.get('Титул/Марка')!r} "
                f"status_ul={r.get('Статус УЛ')!r} ul_tags={r.get('Теги УЛ')!r}"
            )

    return 0


def _preview_cell(value: object, limit: int = 200) -> str:
    text = _cell_text(value)
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


if __name__ == "__main__":
    raise SystemExit(main())
