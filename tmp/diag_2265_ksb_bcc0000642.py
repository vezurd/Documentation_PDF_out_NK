"""Read-only diagnosis of 2265-KSB / BCC0000642 leftover vs RFP qty."""
from __future__ import annotations

import os
import sys
from pathlib import Path

from openpyxl import load_workbook

RESULT_DIR = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.09.03.10.39"
)
YESTERDAY_MATCH = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.09.02.13.49"
    r"\Шаг4_Сопоставление_RFP_MTO_20260902_140222.xlsx"
)
TITLE_NEEDLE = "2265-KSB"
CODE_NEEDLE = "BCC0000642"


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def _norm(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _upper(value: object) -> str:
    return _norm(value).upper()


def _num(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def list_dir() -> list[str]:
    _safe_print(f"RESULT_DIR exists={os.path.exists(RESULT_DIR)}")
    names = os.listdir(RESULT_DIR)
    _safe_print(f"file_count={len(names)}")
    out: list[str] = []
    for name in sorted(names):
        full = os.path.join(RESULT_DIR, name)
        try:
            size = os.path.getsize(full)
        except OSError as exc:
            size = -1
            _safe_print(f"size_error {ascii(name)} {exc!r}")
        _safe_print(f"FILE {ascii(name)} size={size}")
        out.append(full)
    return out


def dump_values_check(path: str) -> None:
    _safe_print(f"\n===== VALUES CHECK {ascii(os.path.basename(path))} =====")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            _safe_print(f"-- sheet {ascii(sheet_name)}")
            rows = ws.iter_rows(values_only=True)
            header = [_norm(v) for v in next(rows, [])]
            _safe_print(f"header={header}")
            for i, row in enumerate(rows, start=2):
                joined = " | ".join(_norm(v) for v in row)
                if TITLE_NEEDLE.lower() in joined.lower() or "ошиб" in joined.lower():
                    _safe_print(f"r{i}: {joined}")
    finally:
        wb.close()


def _cell(idx: dict[str, int], row: tuple, name: str) -> object:
    i = idx.get(name)
    if i is None or i >= len(row):
        return None
    return row[i]


def dump_match_file(path: str) -> None:
    _safe_print(f"\n===== MATCH FILE {ascii(os.path.basename(path))} =====")
    # Comments are needed; read_only mode does not load them.
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        _safe_print(f"sheets={wb.sheetnames}")
        ws = wb[wb.sheetnames[0]]
        header = [_norm(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {name: i for i, name in enumerate(header) if name}
        _safe_print(f"n_headers={len(header)}")

        interesting: list[tuple[int, tuple, str]] = []
        for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
            row = tuple(c.value for c in cells)
            title_raw = _upper(_cell(idx, row, "Титул/Марка"))
            if "2265" not in title_raw or "KSB" not in title_raw:
                continue
            all_text = " ".join(_upper(v) for v in row if v is not None)
            if CODE_NEEDLE not in all_text and "2265-S-FV-0711" not in all_text:
                continue
            ul_comment = ""
            ul_i = idx.get("Статус УЛ")
            if ul_i is not None and ul_i < len(cells) and cells[ul_i].comment:
                ul_comment = str(cells[ul_i].comment.text).replace("\n", " | ")
            interesting.append((rn, row, ul_comment))

        _safe_print(f"matched_rows={len(interesting)}")

        cols_print = [
            "Тип строки",
            "Порядковый ДС",
            "Фактический ДС",
            "№ позиции",
            "Титул/Марка",
            "Код RFP",
            "Кол-во RFP",
            "Ед. изм. RFP",
            "Теги RFP",
            "Тег (VO/MTO/RFP)",
            "MTO, Статус позиции",
            "Наименование RFP",
            "Тип марки RFP",
            "Количество MTO",
            "Ед. изм. MTO",
            "Код MTO",
            "Тег MTO",
            "МТО − RFP",
            "МТО − УЛ",
            "Кол-во по УЛ",
            "Ед. изм. УЛ",
            "RFP − УЛ",
            "Статус УЛ",
            "Статус тегов УЛ",
            "Тег УЛ",
            "Наименование УЛ",
            "Код УЛ",
            "Источник УЛ (файл · вкладка · строка)",
            "Статус сопоставления",
            "Путь к RFP",
            "Путь к МТО",
        ]
        present = [c for c in cols_print if c in idx]
        missing = [c for c in cols_print if c not in idx]
        if missing:
            _safe_print(f"missing_cols={missing}")

        sums: dict[str, float] = {}
        status_counts: dict[str, int] = {}
        for rn, row, ul_comment in interesting:
            status = _norm(_cell(idx, row, "Статус УЛ"))
            status_counts[status] = status_counts.get(status, 0) + 1
            parts = [f"excel_row={rn}"]
            for col in present:
                val = _cell(idx, row, col)
                parts.append(f"{col}={val!r}")
            if ul_comment:
                parts.append(f"UL_COMMENT={ul_comment!r}")
            _safe_print("ROW " + " || ".join(parts))
            for qty_col in (
                "Кол-во RFP",
                "Количество MTO",
                "Кол-во по УЛ",
                "RFP − УЛ",
                "МТО − RFP",
                "МТО − УЛ",
            ):
                n = _num(_cell(idx, row, qty_col))
                if n is None:
                    continue
                key = f"{qty_col}|{status or '(empty status)'}"
                sums[key] = sums.get(key, 0.0) + n
                sums[qty_col] = sums.get(qty_col, 0.0) + n

        _safe_print("\n--- status counts ---")
        for k, v in sorted(status_counts.items(), key=lambda kv: (-kv[1], kv[0])):
            _safe_print(f"  {k!r}: {v}")
        _safe_print("\n--- qty sums among selected rows ---")
        for k, v in sorted(sums.items()):
            _safe_print(f"  {k} = {v:g}")
    finally:
        wb.close()


def dump_quality_if_any(path: str) -> None:
    _safe_print(f"\n===== OTHER SHEETS IN {ascii(os.path.basename(path))} =====")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames[1:]:
            ws = wb[sheet_name]
            _safe_print(f"\n-- sheet {ascii(sheet_name)} --")
            n = 0
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                joined = " | ".join(_norm(v) for v in row)
                if not joined:
                    continue
                hit = (
                    TITLE_NEEDLE.lower() in joined.lower()
                    or CODE_NEEDLE.lower() in joined.lower()
                    or (i <= 8)
                )
                if hit:
                    _safe_print(f"r{i}: {joined[:500]}")
                    n += 1
                if n >= 80:
                    _safe_print("... truncated sheet hits ...")
                    break
    finally:
        wb.close()


def dump_small_xlsx_hits(path: str) -> None:
    _safe_print(f"\n===== SMALL XLSX {ascii(os.path.basename(path))} =====")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            _safe_print(f"-- sheet {ascii(sheet_name)}")
            hits = 0
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                joined = " | ".join(_norm(v) for v in row)
                if i <= 3:
                    _safe_print(f"r{i}: {joined[:400]}")
                needle_hit = (
                    TITLE_NEEDLE.lower() in joined.lower()
                    or CODE_NEEDLE.lower() in joined.lower()
                    or "2265" in joined
                )
                if i > 3 and needle_hit:
                    _safe_print(f"HIT r{i}: {joined[:500]}")
                    hits += 1
                    if hits >= 40:
                        _safe_print("... truncated hits ...")
                        break
            _safe_print(f"hits={hits}")
    finally:
        wb.close()


def search_nearby_match_files() -> None:
    parent = os.path.dirname(RESULT_DIR)
    _safe_print(f"\n===== PARENT DIR {ascii(parent)} =====")
    try:
        names = os.listdir(parent)
    except OSError as exc:
        _safe_print(f"listdir_error {exc!r}")
        return
    names_sorted = sorted(names, reverse=True)
    _safe_print(f"parent_count={len(names_sorted)}")
    for name in names_sorted[:15]:
        full = os.path.join(parent, name)
        try:
            is_dir = os.path.isdir(full)
        except OSError:
            is_dir = False
        _safe_print(f"{'DIR' if is_dir else 'FILE'} {ascii(name)}")
        if not is_dir:
            continue
        try:
            children = os.listdir(full)
        except OSError as exc:
            _safe_print(f"  listdir_error {exc!r}")
            continue
        for child in sorted(children):
            if "сопоставлен" in child.lower() or child.lower().startswith("step4_"):
                cfull = os.path.join(full, child)
                try:
                    size = os.path.getsize(cfull)
                except OSError:
                    size = -1
                _safe_print(f"  MATCHCAND {ascii(child)} size={size}")

    temp_dir = os.environ.get("TEMP") or os.environ.get("TMP") or ""
    _safe_print(f"\n===== TEMP {ascii(temp_dir)} =====")
    if not temp_dir:
        return
    try:
        for name in os.listdir(temp_dir):
            if not name.lower().startswith("step4_"):
                continue
            full = os.path.join(temp_dir, name)
            try:
                size = os.path.getsize(full)
            except OSError:
                size = -1
            _safe_print(f"TEMPFILE {ascii(name)} size={size}")
    except OSError as exc:
        _safe_print(f"temp_error {exc!r}")


def main() -> int:
    files = list_dir()
    match_files: list[str] = []
    values_files: list[str] = []
    other_xlsx: list[str] = []
    for full in files:
        name = os.path.basename(full)
        lower = name.lower()
        if not lower.endswith(".xlsx"):
            continue
        if "проверк" in lower and "сумм" in lower:
            values_files.append(full)
        elif "сопоставлен" in lower:
            match_files.append(full)
        else:
            other_xlsx.append(full)

    # Fallback by size: biggest xlsx is almost always the match file.
    if not match_files:
        xlsx = [f for f in files if f.lower().endswith(".xlsx")]
        xlsx.sort(key=lambda p: os.path.getsize(p), reverse=True)
        _safe_print("\nNo name-match for Step4 match file; largest xlsx candidates:")
        for full in xlsx[:8]:
            _safe_print(f"  {ascii(os.path.basename(full))} size={os.path.getsize(full)}")
        if xlsx and os.path.getsize(xlsx[0]) > 200_000:
            match_files = [xlsx[0]]

    for path in values_files:
        dump_values_check(path)
    for path in match_files:
        dump_match_file(path)
        dump_quality_if_any(path)
    for path in other_xlsx:
        dump_small_xlsx_hits(path)

    reports = [
        os.path.join(RESULT_DIR, name)
        for name in os.listdir(RESULT_DIR)
        if "ul_compare" in name.lower() or name.lower().endswith(".txt")
    ]
    _safe_print(f"\ntext_reports={len(reports)}")
    for path in reports:
        _safe_print(f"TXT {ascii(os.path.basename(path))} size={os.path.getsize(path)}")
        if os.path.getsize(path) < 200_000:
            with open(path, encoding="utf-8", errors="replace") as fh:
                text = fh.read()
            if TITLE_NEEDLE in text or CODE_NEEDLE in text:
                _safe_print("contains title/code")

    search_nearby_match_files()
    _safe_print("\n===== YESTERDAY MATCH (likely the open Excel) =====")
    if os.path.exists(YESTERDAY_MATCH):
        dump_match_file(YESTERDAY_MATCH)
        dump_quality_if_any(YESTERDAY_MATCH)
    else:
        _safe_print(f"missing {ascii(YESTERDAY_MATCH)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
