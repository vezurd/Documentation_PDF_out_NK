"""Smoke: compact RFP tag-count mismatch Excel (same layout as MTO)."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from base.base_classes import RowStd, RowType
from base.tables_columns import CODE, DS_NAME, DS_TITLE, NAME, TAGS, VALUES
from RFQ.rfp_parts.file_status import DUPLICATE_TAGS_XLSX_NAME
from RFQ.tags_rfp_compare.step1_load_rfp import _check_rfp_tags_and_values
from RFQ.tags_rfp_compare.tag_count_mismatch_excel import (
    load_parts_lot_tag_mismatches,
    save_tag_count_mismatch_excel,
)


def _rfp_row(*, code: str, name: str, ds_name: str, ds_title: str, tags: list[str], qty: object) -> RowStd:
    row = RowStd()
    row.row_type = RowType.position_row
    row.el[CODE].value = code
    row.el[NAME].value = name
    row.el[DS_NAME].value = ds_name
    row.el[DS_TITLE].value = ds_title
    row.el[TAGS].value = tags
    row.el[VALUES].value = qty
    return row


class RfpTagCountMismatchSmokeTest(unittest.TestCase):
    def test_save_rfp_layout_matches_mto_shape(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            path = save_tag_count_mismatch_excel(
                [
                    {
                        "title_system": "8445-SOT",
                        "code": "BCC0001",
                        "name": "Клапан",
                        "ds_name": "ДС23",
                        "tags_count": 2,
                        "values": 5,
                        "tags": "T-1, T-2",
                    }
                ],
                tmp,
                source_label="RFP",
                name_header="NAME",
                name_key="name",
                ds_header="DS_NAME",
                ds_key="ds_name",
                timestamp="20260904_105837",
            )
            self.assertIsNotNone(path)
            self.assertEqual(
                Path(path).name,
                "Отчет по несоответствию тегов - RFP_20260904_105837.xlsx",
            )
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb.active
                headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                self.assertEqual(
                    headers,
                    [
                        "№",
                        "title_system",
                        "CODE",
                        "NAME",
                        "DS_NAME",
                        "Кол-во тегов",
                        "VALUES",
                        "Теги",
                    ],
                )
                data = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
                self.assertEqual(data[1], "8445-SOT")
                self.assertEqual(data[3], "Клапан")
                self.assertEqual(data[4], "ДС23")
                self.assertEqual(data[5], 2)
                self.assertEqual(data[6], 5)
                self.assertEqual(data[7], "T-1\nT-2")
            finally:
                wb.close()

    def test_load_parts_keeps_only_lot_vs_tags(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            src = Path(tmp) / DUPLICATE_TAGS_XLSX_NAME
            wb = Workbook()
            ws = wb.active
            ws.append(
                [
                    "Тип",
                    "Файл",
                    "Лист",
                    "Строка",
                    "Тег",
                    "Кол-во лота",
                    "Число тегов",
                    "Код",
                    "Наименование",
                    "Имя ДС",
                    "Титул",
                    "Теги в ячейке",
                    "Где ещё встречается",
                ]
            )
            ws.append(
                [
                    "Дубль тега",
                    "a.xlsx",
                    "Лист",
                    10,
                    "DUP-1",
                    1,
                    1,
                    "BCC0000",
                    "Dup",
                    "ДС1",
                    "1000-SOT",
                    "DUP-1",
                    "b.xlsx",
                ]
            )
            ws.append(
                [
                    "Лот > тегов",
                    "b.xlsx",
                    "Лист",
                    20,
                    "T-1; T-2",
                    5,
                    2,
                    "BCC0002",
                    "Насос",
                    "ДС29",
                    "8950-POS5",
                    "T-1, T-2",
                    "",
                ]
            )
            ws.append(
                [
                    "Тегов > лота",
                    "c.xlsx",
                    "Лист",
                    30,
                    "X-1; X-2; X-3",
                    1,
                    3,
                    "BCC0003",
                    "Муфта",
                    "ДС67",
                    "2265-KSB",
                    "X-1, X-2, X-3",
                    "",
                ]
            )
            wb.save(src)
            wb.close()
            rows = load_parts_lot_tag_mismatches(src)
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["code"], "BCC0002")
            self.assertEqual(rows[0]["values"], 5)
            self.assertEqual(rows[0]["tags_count"], 2)
            self.assertEqual(rows[1]["code"], "BCC0003")
            self.assertEqual(rows[1]["ds_name"], "ДС67")

    def test_step1_uses_parts_report_when_net_has_no_mismatch(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            parts = Path(tmp) / DUPLICATE_TAGS_XLSX_NAME
            wb = Workbook()
            ws = wb.active
            ws.append(
                [
                    "Тип",
                    "Код",
                    "Наименование",
                    "Имя ДС",
                    "Титул",
                    "Число тегов",
                    "Кол-во лота",
                    "Теги в ячейке",
                ]
            )
            ws.append(
                [
                    "Лот > тегов",
                    "BCC0009",
                    "Позиция",
                    "ДС8",
                    "8445-SOT",
                    1,
                    3,
                    "ONLY-1",
                ]
            )
            wb.save(parts)
            wb.close()
            net_row = _rfp_row(
                code="BCC0009",
                name="Позиция",
                ds_name="ДС8",
                ds_title="8445-SOT",
                tags=["ONLY-1"],
                qty=1,
            )
            _check_rfp_tags_and_values(
                [net_row],
                tmp,
                parts_tag_report_path=str(parts),
            )
            reports = list(Path(tmp).glob("Отчет по несоответствию тегов - RFP_*.xlsx"))
            self.assertEqual(len(reports), 1)
            out = load_workbook(reports[0], read_only=True, data_only=True)
            try:
                ws = out.active
                data = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
                self.assertEqual(data[2], "BCC0009")
                self.assertEqual(data[5], 1)
                self.assertEqual(data[6], 3)
            finally:
                out.close()

    def test_step1_writes_net_mismatch_without_parts_report(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            row = _rfp_row(
                code="BCC0010",
                name="Кран",
                ds_name="ДС2",
                ds_title="8525-AVI",
                tags=["A-1", "A-2"],
                qty=5,
            )
            _check_rfp_tags_and_values([row], tmp)
            reports = list(Path(tmp).glob("Отчет по несоответствию тегов - RFP_*.xlsx"))
            self.assertEqual(len(reports), 1)
            out = load_workbook(reports[0], read_only=True, data_only=True)
            try:
                ws = out.active
                data = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
                self.assertEqual(data[2], "BCC0010")
                self.assertEqual(data[5], 2)
                self.assertEqual(data[6], 5)
                self.assertEqual(data[3], "Кран")
            finally:
                out.close()


if __name__ == "__main__":
    unittest.main()
