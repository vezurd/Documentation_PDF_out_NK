"""Build numbered ID handoff list from tsd_packing_critical.txt."""

from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl

from RFQ.ds_compare.ds_compare_config import (
    load_ds_compare_config,
    normalize_gui_paths,
)
from RFQ.ds_compare.tsd_packing_load import (
    TSD_PACKING_CRITICAL_REPORT_NAME,
    is_sopl_sheet,
    tsd_packing_cache_dir,
)

OUT = ROOT / "tmp" / "tsd_id_fix_list.txt"

_ISSUE_RE = re.compile(
    r"(?P<n>\d+)\. ERROR (?P<code>\w+) \[(?P<path>.+?) · (?P<sheet>.+?) · "
    r"строка (?P<row>\d+)\](?P<rest>.*)"
)
_FIELD_RE = re.compile(r"field=([^:\s]+)")
_VALUE_RE = re.compile(r"value=(\{[^}]*\})")

_WANTED = {
    "source_required_fields",
    "packing_tags_exceed_quantity",
}

_SOPL_COLS = {
    "CODE": 6,
    "NAME": 12,
    "VALUES": 13,
    "UNITS": 14,
    "TAGS": 9,
    "TITLE": 56,
    "MARK": 57,
}
_STD_COLS = {
    "CODE": 1,
    "NAME": 7,
    "VALUES": 8,
    "UNITS": 9,
    "TAGS": 3,
    "SPEC": 2,
}


def _short(value: object, limit: int = 80) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\n", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _snap(ws: object, excel_row: int, *, sopl: bool) -> dict[str, str]:
    colmap = _SOPL_COLS if sopl else _STD_COLS
    max_col = max(colmap.values()) + 1
    cells = next(
        ws.iter_rows(
            min_row=excel_row,
            max_row=excel_row,
            max_col=max_col,
            values_only=True,
        ),
        (),
    )
    out: dict[str, str] = {}
    for key, idx0 in colmap.items():
        value = cells[idx0] if idx0 < len(cells) else None
        out[key] = _short(value)
    return out


def _action(code: str, field: str, *, sopl: bool) -> str:
    if code == "packing_key_missing":
        if sopl:
            return (
                "Заполнить Титул и Марку в столбцах BE («Титул») и BF («Марка») "
                "для каждой позиции (сейчас колонки пустые / отсутствуют в файле)."
            )
        return "Заполнить титул и марку позиции (ключ сопоставления)."
    if code == "packing_tags_exceed_quantity":
        return (
            "Исправить рассинхрон: число тегов в столбце TAG не должно превышать "
            "количество (Quantity). Либо убрать лишние теги, либо увеличить количество."
        )
    if code == "source_required_fields":
        if field == "code":
            return (
                "Заполнить «Код РД» (столбец B) — сейчас пусто или только пробелы."
            )
        if field == "units":
            return "Заполнить «Единица измерения» (столбец J)."
        if field == "values":
            return "Заполнить «Количество» (столбец I)."
        if "values" in field and "units" in field:
            return (
                "Заполнить «Количество» (столбец I) и «Единица измерения» (столбец J) "
                "— сейчас пусто при заполненных коде/наименовании."
            )
        if "name" in field:
            return (
                "Заполнить обязательные поля позиции "
                f"({field}): CODE, NAME, VALUES, UNITS."
            )
        return f"Заполнить обязательные поля позиции: {field or 'CODE/NAME/VALUES/UNITS'}."
    return "Проверить строку по отчёту робота."


def main() -> int:
    critical = tsd_packing_cache_dir() / TSD_PACKING_CRITICAL_REPORT_NAME
    text = critical.read_text(encoding="utf-8")
    root = ""
    for line in text.splitlines()[:5]:
        if line.lower().startswith("root:"):
            root = line.split(":", 1)[1].strip()
            break
    if not root:
        gui = normalize_gui_paths(load_ds_compare_config().get("gui_paths"))
        root = gui.get("last_tsd_packing_folder") or ""
    root_path = Path(root)

    by_file: dict[str, list[dict[str, object]]] = defaultdict(list)
    for match in _ISSUE_RE.finditer(text):
        code = match.group("code")
        if code not in _WANTED:
            continue
        rest = match.group("rest")
        field_m = _FIELD_RE.search(rest)
        value_m = _VALUE_RE.search(rest)
        by_file[match.group("path")].append(
            {
                "code": code,
                "sheet": match.group("sheet"),
                "row": int(match.group("row")),
                "field": field_m.group(1) if field_m else "",
                "value": value_m.group(1) if value_m else "",
            }
        )

    lines: list[str] = [
        "Список замечаний по исходным УЛ для исправления (ИД)",
        "Только: пустые CODE/NAME/VALUES/UNITS и теги > количества.",
        "(Отдельно: SO-PL без Титула/Марки BE/BF — файл "
        "согл УЛ ГФ 5титулов\\AGCC.323-2000-20642031-M15-AN004833FR_01_EN.xlsx, "
        "строки позиций с 17 — заполнить BE/BF.)",
        f"Источник: {critical}",
        f"Корень ТСД: {root}",
        "",
        "Типы ошибок:",
        "- source_required_fields — не заполнены обязательные поля позиции "
        "(код / наименование / количество / ед.изм.).",
        "- packing_tags_exceed_quantity — тегов больше, чем количество.",
        "",
    ]

    for file_no, rel in enumerate(sorted(by_file.keys(), key=str.lower), start=1):
        items = sorted(
            by_file[rel],
            key=lambda it: (str(it["sheet"]), int(it["row"]), str(it["code"])),
        )
        lines.append(f"{file_no}. Файл: {rel}")
        abs_path = root_path / rel
        lines.append(f"   Полный путь: {abs_path}")
        if not abs_path.is_file():
            lines.append("   Файл не найден по указанному пути — проверить имя/папку.")
            for it in items:
                lines.append(
                    f"   - лист «{it['sheet']}», строка {it['row']}: "
                    f"тип {it['code']}"
                    + (f" (поля: {it['field']})" if it["field"] else "")
                    + f". Что сделать: {_action(str(it['code']), str(it['field']), sopl=is_sopl_sheet(str(it['sheet'])))}"
                )
            lines.append("")
            continue

        wb = openpyxl.load_workbook(abs_path, read_only=True, data_only=True)
        try:
            for it in items:
                sheet = str(it["sheet"])
                sopl = is_sopl_sheet(sheet)
                action = _action(str(it["code"]), str(it["field"]), sopl=sopl)
                line = (
                    f"   - лист «{sheet}», строка {it['row']}: тип {it['code']}"
                )
                if it["field"]:
                    line += f" (поля: {it['field']})"
                if it["value"]:
                    line += f"; {it['value']}"
                line += f". Что сделать: {action}"
                lines.append(line)

                if sheet in wb.sheetnames:
                    snap = _snap(wb[sheet], int(it["row"]), sopl=sopl)
                    if sopl:
                        lines.append(
                            "     сейчас: "
                            f"CODE={snap['CODE']!r}; NAME={snap['NAME']!r}; "
                            f"QTY={snap['VALUES']!r}; UNITS={snap['UNITS']!r}; "
                            f"TITLE(BE)={snap['TITLE']!r}; MARK(BF)={snap['MARK']!r}"
                        )
                    else:
                        lines.append(
                            "     сейчас: "
                            f"CODE={snap['CODE']!r}; NAME={snap['NAME']!r}; "
                            f"QTY={snap['VALUES']!r}; UNITS={snap['UNITS']!r}; "
                            f"TAGS={snap['TAGS']!r}"
                        )
        finally:
            wb.close()
        lines.append("")

    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {OUT}")
    print(f"files={len(by_file)} issues={sum(len(v) for v in by_file.values())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
