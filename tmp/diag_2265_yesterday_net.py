"""Yesterday VALUES check + step1 duplicates for 2265-KSB / BCC0000642."""
from __future__ import annotations

import os
import sys

from openpyxl import load_workbook

DIR = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.09.02.13.49"
)
NET_BASE = (
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP\RFP сводный файл"
)


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def dump_xlsx_hits(path: str, needles: tuple[str, ...]) -> None:
    _safe_print(f"\nFILE {ascii(os.path.basename(path))} size={os.path.getsize(path)}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                joined = " | ".join("" if v is None else str(v).strip() for v in row)
                if i <= 2 or any(n.lower() in joined.lower() for n in needles):
                    _safe_print(f"{sheet_name} r{i}: {joined[:500]}")
    finally:
        wb.close()


def main() -> None:
    _safe_print("=== yesterday result dir ===")
    for name in sorted(os.listdir(DIR)):
        _safe_print(ascii(name))
        full = os.path.join(DIR, name)
        low = name.lower()
        if not low.endswith(".xlsx"):
            continue
        if any(k in low for k in ("сумм", "дубл", "несовпад", "tag_effective")):
            dump_xlsx_hits(full, ("2265-KSB", "BCC0000642", "2265-S-FV-0711"))

    _safe_print("\n=== latest rfp_parts_net ===")
    try:
        stamps = sorted(os.listdir(NET_BASE), reverse=True)
    except OSError as exc:
        _safe_print(repr(exc))
        return
    for stamp in stamps[:8]:
        _safe_print(f"stamp {ascii(stamp)}")
        net = os.path.join(NET_BASE, stamp, "rfp_parts_net.xlsx")
        if os.path.isfile(net):
            _safe_print(f"NET {ascii(net)} size={os.path.getsize(net)}")
            dump_xlsx_hits(net, ("2265-KSB", "BCC0000642"))
            tags = os.path.join(NET_BASE, stamp, "rfp_parts_duplicate_tags.xlsx")
            if os.path.isfile(tags):
                dump_xlsx_hits(tags, ("2265-S-FV-0711", "BCC0000642", "ДС67"))
            break


if __name__ == "__main__":
    main()
