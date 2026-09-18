"""Smoke tests for launch units gate ordering and fractional safety."""

from __future__ import annotations

import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _ok_ds_id_coverage() -> mock.Mock:
    """Avoid UNC scans of RFP_Зиновьев / TSD during units-gate order tests."""
    fake = mock.Mock()
    fake.load_error = None
    fake.summary_line.return_value = "ОК: пар 0, замечаний нет"
    return fake

from base.base_classes import RowStd, RowType
from base.tables_columns import (
    CODE,
    DS_SYSTEM,
    DS_TITLE,
    TAGS,
    UNITS,
    UNITS_CHECK_STATUS,
    UNITS_CONVERSION_TRACE,
    VALUES,
    VALUES_2,
)
from RFQ.packing_list_provider import PackingCacheMeta, PackingDataset, PackingQualityLevel
from RFQ.tags_rfp_compare.step1_load_rfp import (
    split_rfp_rows,
    step1_finalize_rfp_data,
    try_exact_int_quantity,
)
from RFQ.tags_rfp_compare.step4.step4_1_check_mto_data import check_mto_data
from RFQ.tags_rfp_compare.step4_analyze_and_match import (
    _resolve_packing_dataset_for_step4,
    step4_analyze_and_match,
)
from RFQ.tags_rfp_compare.units_gate import (
    build_units_gate_plan,
    format_units_gate_progress_summary,
    run_units_gate,
)
from RFQ.units_convert.models import (
    ConversionAction,
    ConversionInvariant,
    ConversionPlan,
    ConversionRequest,
    PAIR_COEF_PREFIX,
    STATUS_CONVERTED,
    STATUS_IDENTITY,
    STATUS_NO_GOOGLE,
    UnitsConversionError,
)
from RFQ.units_convert import build_conversion_plan, build_google_units_index
from openpyxl import load_workbook
from utils.cache_utils import cache_manager


def _row_gate_snapshot(row: RowStd) -> dict[str, object]:
    def _cell(column: str) -> object:
        return row.el[column].value if column in row.el else None

    return {
        VALUES: _cell(VALUES),
        UNITS: _cell(UNITS),
        UNITS_CHECK_STATUS: _cell(UNITS_CHECK_STATUS),
        UNITS_CONVERSION_TRACE: _cell(UNITS_CONVERSION_TRACE),
    }


def _seed_matrix_placeholder(
    matrix_path: Path,
    *,
    code: str,
    google_unit: str,
    source_unit: str,
) -> None:
    google_index = build_google_units_index([_google_row(code, google_unit)])
    try:
        build_conversion_plan(
            [
                ConversionRequest(
                    request_id="seed",
                    contour="seed",
                    code=code,
                    source_unit=source_unit,
                    quantity="1",
                    tags_count=0,
                    invariant=ConversionInvariant.NONE,
                    location="seed",
                )
            ],
            google_index,
            matrix_path,
        )
    except UnitsConversionError:
        pass
    else:
        raise AssertionError("expected UnitsConversionError for placeholder seed")


def _coef_column(ws, pair_index: int) -> int:
    title = f"{PAIR_COEF_PREFIX}{pair_index}"
    for col_idx in range(1, ws.max_column + 1):
        if str(ws.cell(row=1, column=col_idx).value or "").strip() == title:
            return col_idx
    raise AssertionError(f"Header column {title!r} not found")


def _set_matrix_coef(matrix_path: Path, row_index: int, pair_index: int, coef: object) -> None:
    wb = load_workbook(matrix_path)
    try:
        ws = wb.active
        coef_col = _coef_column(ws, pair_index)
        ws.cell(row=row_index, column=coef_col, value=coef)
        wb.save(matrix_path)
    finally:
        wb.close()


def _find_question_mark_coef_cell(matrix_path: Path) -> tuple[int, int]:
    wb = load_workbook(matrix_path)
    try:
        ws = wb.active
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                title = str(ws.cell(row=1, column=col_idx).value or "")
                if not title.startswith(PAIR_COEF_PREFIX):
                    continue
                if ws.cell(row=row_idx, column=col_idx).value == "?":
                    pair_index = int(title.removeprefix(PAIR_COEF_PREFIX))
                    return row_idx, pair_index
        raise AssertionError("Expected '?' coefficient cell was not written")
    finally:
        wb.close()


def _prepare_matrix_coef(
    matrix_path: Path,
    *,
    code: str,
    google_unit: str,
    source_unit: str,
    coef: object,
) -> None:
    _seed_matrix_placeholder(
        matrix_path,
        code=code,
        google_unit=google_unit,
        source_unit=source_unit,
    )
    row_index, pair_index = _find_question_mark_coef_cell(matrix_path)
    _set_matrix_coef(matrix_path, row_index, pair_index, coef)


def _position_row(*, code: str, units: str, qty: object, tags: list[str] | None = None) -> RowStd:
    row = RowStd()
    row.row_type = RowType.position_row
    row.el[CODE].value = code
    row.el[UNITS].value = units
    row.el[VALUES].value = qty
    if tags is not None:
        row.el[TAGS].value = tags
    return row


def _google_row(code: str, units: str) -> RowStd:
    row = RowStd()
    row.el[CODE].value = code
    row.el[UNITS].value = units
    return row


def _conversion_action(*, request_id: str, status: str) -> ConversionAction:
    return ConversionAction(
        request_id=request_id,
        code_raw="BCC0001",
        code_normalized="BCC0001",
        original_unit="шт",
        source_unit="шт",
        target_unit="шт",
        original_quantity=Decimal("1"),
        coefficient=Decimal("1"),
        result_quantity=Decimal("1"),
        status=status,
        trace="",
        quantity_changed=status == STATUS_CONVERTED,
        units_changed=status == STATUS_CONVERTED,
        machine_residual_adjusted=False,
    )


class Step4UnitsGateOrderSmokeTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.result_dir = self._tmpdir.name

    def tearDown(self) -> None:
        self._tmpdir.cleanup()

    def test_raw_fractional_rfp_value_preserved_before_finalize(self) -> None:
        row = _position_row(code="BCC0001", units="шт", qty="1.5")
        before = row.el[VALUES].value
        with mock.patch(
            "RFQ.tags_rfp_compare.step1_load_rfp._check_rfp_tags_and_values",
            return_value=[],
        ):
            finalized, _errors = step1_finalize_rfp_data(
                [row],
                self.result_dir,
                debug=False,
                skip_split=True,
                units_ban=[],
            )
        self.assertEqual(before, "1.5")
        self.assertEqual(finalized[0].el[VALUES].value, before)

    def test_finalize_does_not_split_fractional_rfp(self) -> None:
        row = _position_row(code="BCC0001", units="шт", qty="1.5")
        with mock.patch(
            "RFQ.tags_rfp_compare.step1_load_rfp._check_rfp_tags_and_values",
            return_value=[],
        ):
            finalized, _errors = step1_finalize_rfp_data(
                [row],
                self.result_dir,
                debug=False,
                skip_split=False,
                units_ban=[],
            )
        self.assertEqual(len(finalized), 1)
        self.assertAlmostEqual(float(finalized[0].el[VALUES].value), 1.5)

    def test_gate_runs_before_finalize_and_step4_snapshots(self) -> None:
        call_order: list[str] = []

        def _gate(*_args, **_kwargs):
            call_order.append("gate")
            return mock.Mock(
                plan=mock.Mock(warnings=()),
                google_index=mock.Mock(),
                warnings=(),
                summary="RFP: проверено 1, преобразовано 0, без Google 0, без изменений 1; MTO: проверено 0, преобразовано 0, без Google 0; УЛ: проверено 0, преобразовано 0, без Google 0",
            )

        def _finalize(*_args, **_kwargs):
            call_order.append("finalize")
            return [], []

        def _step4(*_args, **_kwargs):
            call_order.append("step4")
            return []

        config = {
            "paths": {
                "rfp_path": "",
                "mto_path": "",
                "vo_path": "",
                "code_ban_file": "",
                "replacement_table_file": "",
                "result_dir_base": self.result_dir,
                "units_convert_matrix": str(Path(self.result_dir) / "matrix.xlsx"),
            },
            "step1": {"rfp_pipeline_mode": "standard", "skip_split": False},
            "step2": {"export_positions_database_excel": False},
            "step3": {},
            "step4": {"include_packing_lists": False},
            "rfp_parts": {"use_latest_net": False},
            "rfp_tags_utils": {"save_input_fingerprints": False},
            "memory_log": False,
        }

        with (
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.load_config", return_value=config),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.resolve_rfp_pipeline_mode", return_value="standard"),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.resolve_effective_rfp_path", return_value=mock.Mock(path="rfp.xlsx", detail="test")),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step1_load_rfp_raw", return_value=[_position_row(code="BCC0001", units="шт", qty=1)]),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step2_load_mto_data", return_value={}),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step3_load_vo_data", return_value={}),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.load_base", return_value=[_google_row("BCC0001", "шт")]),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.run_units_gate", side_effect=_gate),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.emit_milestone") as emit_milestone,
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step1_finalize_rfp_data", side_effect=_finalize),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step4_analyze_and_match", side_effect=_step4),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.ensure_rfp_parts_net_current"),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.check_ds_id_coverage",
                return_value=_ok_ds_id_coverage(),
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.open_dir"),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.finalize_timing_log"),
        ):
            from RFQ.tags_rfp_compare.agregate_tags import main

            main(config_override=config)

        self.assertEqual(call_order, ["gate", "finalize", "step4"])
        milestone_calls = [(call.args[0], call.args[1]) for call in emit_milestone.call_args_list]
        running_positions = [
            index
            for index, (milestone_id, state) in enumerate(milestone_calls)
            if milestone_id == "units_gate" and state == "Running"
        ]
        done_positions = [
            index
            for index, (milestone_id, state) in enumerate(milestone_calls)
            if milestone_id == "units_gate" and state == "Done"
        ]
        self.assertEqual(len(running_positions), 1)
        self.assertEqual(len(done_positions), 1)
        self.assertLess(running_positions[0], done_positions[0])
        self.assertIn(("ds_mp_check", "Skipped"), milestone_calls)
        self.assertIn(("ul_preflight", "Skipped"), milestone_calls)
        ds_mp_at = next(
            index
            for index, (milestone_id, _state) in enumerate(milestone_calls)
            if milestone_id == "ds_mp_check"
        )
        parts_at = next(
            index
            for index, (milestone_id, _state) in enumerate(milestone_calls)
            if milestone_id == "parts_preflight"
        )
        ul_at = next(
            index
            for index, (milestone_id, _state) in enumerate(milestone_calls)
            if milestone_id == "ul_preflight"
        )
        self.assertLess(ds_mp_at, parts_at)
        self.assertLess(parts_at, ul_at)

    def test_orchestrator_ul_preflight_runs_before_rfp_load(self) -> None:
        from RFQ.ds_compare.tsd_packing_load import TsdLoadResult, TsdLoadStats

        packing_result = TsdLoadResult(
            summary_path="summary.xlsx",
            stats=TsdLoadStats(position_rows=12),
            success=True,
            from_cache=True,
        )
        config = {
            "paths": {
                "rfp_path": "",
                "mto_path": "",
                "vo_path": "",
                "code_ban_file": "",
                "replacement_table_file": "",
                "result_dir_base": self.result_dir,
                "units_convert_matrix": str(Path(self.result_dir) / "matrix.xlsx"),
            },
            "step1": {"rfp_pipeline_mode": "standard", "skip_split": False},
            "step2": {"export_positions_database_excel": False},
            "step3": {},
            "step4": {"include_packing_lists": True},
            "rfp_parts": {"use_latest_net": False},
            "rfp_tags_utils": {"save_input_fingerprints": False},
            "memory_log": False,
        }

        with (
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.load_config", return_value=config),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.resolve_rfp_pipeline_mode", return_value="standard"),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.resolve_effective_rfp_path",
                return_value=mock.Mock(path="rfp.xlsx", detail="test"),
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.step1_load_rfp_raw",
                return_value=[_position_row(code="BCC0001", units="шт", qty=1)],
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step2_load_mto_data", return_value={}),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step3_load_vo_data", return_value={}),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.load_base",
                return_value=[_google_row("BCC0001", "шт")],
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.run_units_gate",
                return_value=mock.Mock(
                    plan=mock.Mock(warnings=()),
                    google_index=mock.Mock(),
                    warnings=(),
                    summary="RFP: проверено 1, преобразовано 0, без Google 0, без изменений 1; MTO: проверено 0, преобразовано 0, без Google 0; УЛ: проверено 0, преобразовано 0, без Google 0",
                ),
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.emit_milestone") as emit_milestone,
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.step1_finalize_rfp_data",
                return_value=([], []),
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step4_analyze_and_match", return_value=[]),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.ensure_rfp_parts_net_current"),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.ensure_tsd_packing_cache_current",
                return_value=packing_result,
            ) as refresh,
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.load_packing_dataset",
                return_value=PackingDataset(
                    rows=[],
                    meta=None,
                    quality=PackingQualityLevel.OK,
                    issues=[],
                    cache_path="cache",
                ),
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.check_ds_id_coverage",
                return_value=_ok_ds_id_coverage(),
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.open_dir"),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.finalize_timing_log"),
        ):
            from RFQ.tags_rfp_compare.agregate_tags import main

            main(config_override=config)

        refresh.assert_called_once_with()
        milestone_calls = [
            (call.args[0], call.args[1]) for call in emit_milestone.call_args_list
        ]
        ul_done = next(
            index
            for index, (milestone_id, state) in enumerate(milestone_calls)
            if milestone_id == "ul_preflight" and state == "Done"
        )
        rfp_running = next(
            index
            for index, (milestone_id, state) in enumerate(milestone_calls)
            if milestone_id == "rfp_load" and state == "Running"
        )
        self.assertLess(ul_done, rfp_running)
        self.assertIn("unchanged", emit_milestone.call_args_list[ul_done].args[2])

    def test_progress_summary_counts_per_contour(self) -> None:
        plan = ConversionPlan(
            actions=(
                _conversion_action(request_id="rfp:0:values", status=STATUS_IDENTITY),
                _conversion_action(request_id="rfp:1:values", status=STATUS_CONVERTED),
                _conversion_action(request_id="mto:1:values", status=STATUS_NO_GOOGLE),
                _conversion_action(request_id="ul:0:values", status=STATUS_CONVERTED),
                _conversion_action(request_id="seed:0:values", status=STATUS_CONVERTED),
            ),
            warnings=(),
            dependencies=(),
        )
        summary = format_units_gate_progress_summary(plan)
        self.assertIn("RFP: проверено 2, преобразовано 1, без Google 0, без изменений 1", summary)
        self.assertIn("MTO: проверено 1, преобразовано 0, без Google 1", summary)
        self.assertIn("УЛ: проверено 1, преобразовано 1, без Google 0", summary)

    def test_units_gate_milestone_error_targets_units_gate_not_global_checks(self) -> None:
        config = {
            "paths": {
                "rfp_path": "",
                "mto_path": "",
                "vo_path": "",
                "code_ban_file": "",
                "replacement_table_file": "",
                "result_dir_base": self.result_dir,
                "units_convert_matrix": str(Path(self.result_dir) / "matrix.xlsx"),
            },
            "step1": {"rfp_pipeline_mode": "standard", "skip_split": False},
            "step2": {"export_positions_database_excel": False},
            "step3": {},
            "step4": {"include_packing_lists": False},
            "rfp_parts": {"use_latest_net": False},
            "rfp_tags_utils": {"save_input_fingerprints": False},
            "memory_log": False,
        }

        def _raise_units_error(*_args, **_kwargs):
            raise UnitsConversionError("В матрице отсутствует коэффициент")

        with (
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.load_config", return_value=config),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.resolve_rfp_pipeline_mode", return_value="standard"),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.resolve_effective_rfp_path",
                return_value=mock.Mock(path="rfp.xlsx", detail="test"),
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.step1_load_rfp_raw",
                return_value=[_position_row(code="BCC0001", units="шт", qty=1)],
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step2_load_mto_data", return_value={}),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step3_load_vo_data", return_value={}),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.load_base", return_value=[_google_row("BCC0001", "шт")]),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.run_units_gate", side_effect=_raise_units_error),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.emit_milestone") as emit_milestone,
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.ensure_rfp_parts_net_current"),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.check_ds_id_coverage",
                return_value=_ok_ds_id_coverage(),
            ),
        ):
            from RFQ.tags_rfp_compare.agregate_tags import main

            with self.assertRaises(SystemExit):
                main(config_override=config)

        error_milestones = [
            (call.args[0], call.args[1], call.args[2] if len(call.args) > 2 else "")
            for call in emit_milestone.call_args_list
            if call.args[1] == "Error"
        ]
        self.assertIn(("units_gate", "Error", "В матрице отсутствует коэффициент"), error_milestones)
        self.assertNotIn(
            ("global_checks", "Error", "конвертация ед. изм.: В матрице отсутствует коэффициент"),
            error_milestones,
        )
        self.assertIn(
            ("complete", "Error", "конвертация ед. изм.: В матрице отсутствует коэффициент"),
            error_milestones,
        )

    def test_step4_internal_order_records_gate_before_snapshots(self) -> None:
        call_order: list[str] = []

        def _record(name: str):
            call_order.append(name)
            return _Snap()

        class _Snap:
            rfp = {}
            mto = {}
            vo = {}
            rfp_rows = {}

        rfp = [_position_row(code="8950-SOO1", units="шт", qty=1)]
        rfp[0].el[DS_TITLE].value = "8950-SOO1"
        mto = {"8950-SOO1": []}
        vo: dict[str, list[RowStd]] = {}

        with (
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.build_input_snapshot",
                side_effect=lambda *_a, **_k: _record("build_input_snapshot"),
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.build_ordered_snapshot",
                side_effect=lambda *_a, **_k: (_record("build_ordered_snapshot"), {})[1],
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match._compare_title_systems_rfp_mto_vo",
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.print_input_quantity_balance",
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.load_replacement_table",
                return_value={},
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.apply_post_split_input_to_snapshot",
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match._detect_and_save_effective_tag_duplicates",
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.process_title_mark",
                side_effect=lambda *_a, **_k: (
                    _record("tm_loop"),
                    {
                        "title_mark": "8950-SOO1",
                        "result_rows": [],
                        "mto_summary": {},
                        "vo_summary": {},
                        "mto_diagnostics": {},
                        "vo_diagnostics": {},
                        "post_split_mto_sum": 0.0,
                        "post_split_mto_rows": 0,
                        "post_split_vo_sum": 0.0,
                        "post_split_vo_rows": 0,
                    },
                )[1],
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.run_output_quantity_balance_check",
                return_value=mock.Mock(fatal=False, summary="ok"),
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.compare_rfp_rows_with_packing",
                return_value=(
                    [],
                    mock.Mock(
                        format_short=lambda: "ok",
                        summary="ok",
                        phase_timings={},
                    ),
                ),
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.run_output_ul_quantity_balance_check",
                return_value=mock.Mock(fatal=False, summary="ok"),
            ),
            mock.patch(
                "RFQ.tags_rfp_compare.step4_analyze_and_match.save_rfp_packing_report",
            ),
            mock.patch("RFQ.tags_rfp_compare.step4_analyze_and_match.check_values_sum"),
            mock.patch("RFQ.tags_rfp_compare.step4_analyze_and_match.save_match_result_to_excel", return_value="out.xlsx"),
            mock.patch("RFQ.tags_rfp_compare.step4_analyze_and_match.emit_milestone"),
        ):
            step4_analyze_and_match(
                rfp,
                mto,
                vo,
                self.result_dir,
                debug=False,
                debug_step4_1=False,
                debug_step4_2=False,
                debug_step4_3=False,
                debug_step4_4=False,
                include_packing_lists=True,
                packing_dataset=PackingDataset(
                    rows=[],
                    meta=None,
                    quality=PackingQualityLevel.UNAVAILABLE,
                    issues=[],
                    cache_path="test.cache",
                ),
                timing_log_enabled=False,
            )

        self.assertEqual(
            call_order[:3],
            ["build_input_snapshot", "build_ordered_snapshot", "tm_loop"],
        )

    def test_rfp_values_and_values_2_convert_and_merge_traces(self) -> None:
        matrix_path = Path(self.result_dir) / "matrix_rfp.xlsx"
        code = "BCC0018"
        google = [_google_row(code, "шт")]
        _prepare_matrix_coef(
            matrix_path,
            code=code,
            google_unit="шт",
            source_unit="м",
            coef="2",
        )
        row = _position_row(code=code, units="м", qty="10")
        row.el[VALUES_2].value = "5"
        result = run_units_gate(
            rfp_rows=[row],
            mto_data={},
            packing_dataset=None,
            google_rows=google,
            matrix_path=matrix_path,
        )
        self.assertEqual(row.el[VALUES].value, 20)
        self.assertIn("RFP: проверено 2, преобразовано 2, без Google 0", result.summary)
        self.assertEqual(row.el[VALUES_2].value, 10)
        self.assertEqual(row.el[UNITS].value, "шт")
        self.assertIn("converted", str(row.el[UNITS_CHECK_STATUS].value))
        trace = str(row.el[UNITS_CONVERSION_TRACE].value)
        self.assertIn("qty=10", trace)
        self.assertIn("qty=5", trace)
        self.assertIn("result=20", trace)
        self.assertIn("result=10", trace)

    def test_ul_tags_not_exceed_fatal_is_transactional(self) -> None:
        matrix_path = Path(self.result_dir) / "matrix_ul.xlsx"
        code = "BCC0100"
        google = [_google_row(code, "шт")]
        _prepare_matrix_coef(
            matrix_path,
            code=code,
            google_unit="шт",
            source_unit="компл",
            coef="0.5",
        )
        rfp_row = _position_row(code=code, units="шт", qty=5)
        mto_row = _position_row(code=code, units="шт", qty=5)
        ul_row = _position_row(code=code, units="компл", qty=2, tags=["T1", "T2"])
        ul_row.el[DS_TITLE].value = "8950"
        ul_row.el[DS_SYSTEM].value = "SOO1"
        ul_row._packing_source_row = 10
        dataset = PackingDataset(
            rows=[ul_row],
            meta=PackingCacheMeta(
                version="v6",
                created_at="now",
                root=".",
                fingerprint="fp",
                critical_ok=True,
            ),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path=str(matrix_path),
        )
        before_rfp = _row_gate_snapshot(rfp_row)
        before_mto = _row_gate_snapshot(mto_row)
        before_ul = _row_gate_snapshot(ul_row)
        with self.assertRaises(UnitsConversionError) as ctx:
            run_units_gate(
                rfp_rows=[rfp_row],
                mto_data={"TS": [mto_row]},
                packing_dataset=dataset,
                google_rows=google,
                matrix_path=matrix_path,
            )
        message = str(ctx.exception)
        self.assertIn("TAGS_NOT_EXCEED_QUANTITY", message)
        self.assertIn("Нарушен инвариант TAGS_NOT_EXCEED_QUANTITY", message)
        self.assertEqual(_row_gate_snapshot(rfp_row), before_rfp)
        self.assertEqual(_row_gate_snapshot(mto_row), before_mto)
        self.assertEqual(_row_gate_snapshot(ul_row), before_ul)

    def test_mto_and_ul_convert_in_common_gate(self) -> None:
        matrix_path = Path(self.result_dir) / "matrix_success.xlsx"
        mto_code = "BCC0300"
        ul_code = "BCC0301"
        google = [_google_row(mto_code, "шт"), _google_row(ul_code, "шт")]
        _prepare_matrix_coef(
            matrix_path,
            code=mto_code,
            google_unit="шт",
            source_unit="м",
            coef="2",
        )
        _prepare_matrix_coef(
            matrix_path,
            code=ul_code,
            google_unit="шт",
            source_unit="компл",
            coef="2",
        )
        mto_row = _position_row(code=mto_code, units="м", qty=4)
        ul_row = _position_row(code=ul_code, units="компл", qty=2, tags=["T1"])
        ul_row.el[DS_TITLE].value = "8950"
        ul_row.el[DS_SYSTEM].value = "SOO1"
        ul_row._packing_source_row = 11
        dataset = PackingDataset(
            rows=[ul_row],
            meta=PackingCacheMeta(
                version="v6",
                created_at="now",
                root=".",
                fingerprint="fp",
                critical_ok=True,
            ),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path=str(matrix_path),
        )
        run_units_gate(
            rfp_rows=[],
            mto_data={"TS1": [mto_row]},
            packing_dataset=dataset,
            google_rows=google,
            matrix_path=matrix_path,
        )
        self.assertEqual(mto_row.el[VALUES].value, 8)
        self.assertEqual(mto_row.el[UNITS].value, "шт")
        self.assertIn("converted", str(mto_row.el[UNITS_CHECK_STATUS].value))
        self.assertEqual(ul_row.el[VALUES].value, 4)
        self.assertEqual(ul_row.el[UNITS].value, "шт")
        self.assertIn("converted", str(ul_row.el[UNITS_CHECK_STATUS].value))

    def test_empty_code_mto_rows_skipped_not_fatal(self) -> None:
        matrix_path = Path(self.result_dir) / "matrix_empty_code.xlsx"
        google = [_google_row("BCC0500", "шт")]
        coded = _position_row(code="BCC0500", units="шт", qty=2)
        empty = _position_row(code="", units="шт", qty=3)
        whitespace = _position_row(code="   ", units="компл", qty=1)
        result = run_units_gate(
            rfp_rows=[],
            mto_data={"8630-KSB3": [coded, empty, whitespace]},
            packing_dataset=None,
            google_rows=google,
            matrix_path=matrix_path,
        )
        self.assertEqual(coded.el[VALUES].value, 2)
        self.assertEqual(empty.el[VALUES].value, 3)
        self.assertEqual(empty.el[UNITS].value, "шт")
        self.assertIsNone(
            empty.el[UNITS_CHECK_STATUS].value if UNITS_CHECK_STATUS in empty.el else None
        )
        self.assertEqual(whitespace.el[VALUES].value, 1)
        self.assertEqual(whitespace.el[UNITS].value, "компл")
        self.assertIn("MTO: проверено 1", result.summary)
        plan, _, bindings = build_units_gate_plan(
            rfp_rows=[],
            mto_data={"8630-KSB3": [coded, empty, whitespace]},
            packing_dataset=None,
            google_rows=google,
            matrix_path=matrix_path,
        )
        self.assertEqual([item.request_id for item in plan.actions], ["mto:1:values"])
        self.assertEqual([item.request_id for item in bindings], ["mto:1:values"])

    def test_mto_request_ids_globally_unique(self) -> None:
        matrix_path = Path(self.result_dir) / "matrix_mto_ids.xlsx"
        google = [_google_row("BCC0400", "шт"), _google_row("BCC0401", "шт")]
        plan, _, bindings = build_units_gate_plan(
            rfp_rows=[],
            mto_data={
                "a:1": [_position_row(code="BCC0400", units="шт", qty=1)],
                "a": [_position_row(code="BCC0401", units="шт", qty=1)],
            },
            packing_dataset=None,
            google_rows=google,
            matrix_path=matrix_path,
        )
        mto_ids = [binding.request_id for binding in bindings if binding.request_id.startswith("mto:")]
        self.assertEqual(mto_ids, ["mto:1:values", "mto:2:values"])
        self.assertEqual(len(mto_ids), len(set(mto_ids)))

    def test_supplied_packing_dataset_used_without_reload(self) -> None:
        supplied = PackingDataset(
            rows=[],
            meta=None,
            quality=PackingQualityLevel.UNAVAILABLE,
            issues=[],
            cache_path="supplied.cache",
        )
        with mock.patch(
            "RFQ.tags_rfp_compare.step4_analyze_and_match.load_packing_dataset",
            side_effect=AssertionError("load_packing_dataset must not be called"),
        ):
            resolved = _resolve_packing_dataset_for_step4(
                include_packing_lists=True,
                packing_dataset=supplied,
            )
        self.assertIs(resolved, supplied)

    def test_direct_step4_fallback_loads_packing_dataset(self) -> None:
        fallback = PackingDataset(
            rows=[],
            meta=None,
            quality=PackingQualityLevel.UNAVAILABLE,
            issues=[],
            cache_path="fallback.cache",
        )
        with mock.patch(
            "RFQ.tags_rfp_compare.step4_analyze_and_match.load_packing_dataset",
            return_value=fallback,
        ) as loader:
            resolved = _resolve_packing_dataset_for_step4(
                include_packing_lists=True,
                packing_dataset=None,
            )
        loader.assert_called_once()
        self.assertIs(resolved, fallback)

    def test_orchestrator_does_not_write_converted_pickle_caches(self) -> None:
        config = {
            "paths": {
                "rfp_path": "",
                "mto_path": "",
                "vo_path": "",
                "code_ban_file": "",
                "replacement_table_file": "",
                "result_dir_base": self.result_dir,
                "units_convert_matrix": str(Path(self.result_dir) / "matrix.xlsx"),
            },
            "step1": {"rfp_pipeline_mode": "standard", "skip_split": False},
            "step2": {"export_positions_database_excel": False},
            "step3": {},
            "step4": {"include_packing_lists": False},
            "rfp_parts": {"use_latest_net": False},
            "rfp_tags_utils": {"save_input_fingerprints": False},
            "memory_log": False,
        }
        with (
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.load_config", return_value=config),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.resolve_rfp_pipeline_mode", return_value="standard"),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.resolve_effective_rfp_path", return_value=mock.Mock(path="rfp.xlsx", detail="test")),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step1_load_rfp_raw", return_value=[_position_row(code="BCC0001", units="шт", qty=1)]),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step2_load_mto_data", return_value={}),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step3_load_vo_data", return_value={}),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.load_base", return_value=[_google_row("BCC0001", "шт")]),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.run_units_gate",
                return_value=mock.Mock(
                    plan=mock.Mock(warnings=()),
                    google_index=mock.Mock(),
                    warnings=(),
                    summary="RFP: проверено 1, преобразовано 0, без Google 0, без изменений 1; MTO: проверено 0, преобразовано 0, без Google 0; УЛ: проверено 0, преобразовано 0, без Google 0",
                ),
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step1_finalize_rfp_data", return_value=([], [])),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.step4_analyze_and_match", return_value=[]),
            mock.patch(
                "RFQ.tags_rfp_compare.agregate_tags.check_ds_id_coverage",
                return_value=_ok_ds_id_coverage(),
            ),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.open_dir"),
            mock.patch("RFQ.tags_rfp_compare.agregate_tags.finalize_timing_log"),
            mock.patch.object(cache_manager, "save_to_cache") as save_cache,
            mock.patch.object(cache_manager, "save_rfp_step1_to_cache") as save_step1,
        ):
            from RFQ.tags_rfp_compare.agregate_tags import main

            main(config_override=config)
        save_cache.assert_not_called()
        save_step1.assert_not_called()

    def test_mto_fractional_value_remains_unsplit(self) -> None:
        row = _position_row(code="BCC0200", units="шт", qty="2.5")
        mto_data = {"TS1": [row]}
        updated, diagnostics = check_mto_data(
            mto_data,
            self.result_dir,
            debug=False,
            defer_reports=True,
            units_ban=[],
        )
        self.assertEqual(len(updated["TS1"]), 1)
        self.assertAlmostEqual(float(updated["TS1"][0].el[VALUES].value), 2.5)
        self.assertEqual(diagnostics["count_mismatches"], [])

    def test_try_exact_int_quantity(self) -> None:
        self.assertEqual(try_exact_int_quantity("2.0"), 2)
        self.assertIsNone(try_exact_int_quantity("1.5"))
        self.assertIsNone(try_exact_int_quantity(Decimal("1.5")))

    def test_split_rfp_rows_skips_fractional_values(self) -> None:
        row = _position_row(code="BCC0001", units="шт", qty="1.5")
        expanded, split_tags, split_no_tags = split_rfp_rows([row], units_ban=[], code_ban=set())
        self.assertEqual(len(expanded), 1)
        self.assertAlmostEqual(float(expanded[0].el[VALUES].value), 1.5)
        self.assertEqual(split_tags, 0)
        self.assertEqual(split_no_tags, 0)


if __name__ == "__main__":
    unittest.main()
