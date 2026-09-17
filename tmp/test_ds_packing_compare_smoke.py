"""Synthetic smoke tests for packing provider and grouped DS enrichment."""

from __future__ import annotations

import io
import pickle
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from base.base_classes import RowStd, RowType, TableComments
from base.tables_columns import (
    CODE,
    CODE_2,
    ColNames,
    DS_SYSTEM,
    DS_TITLE,
    NAME,
    RFQ_CODE,
    RFQ_COMPARE_STATUS,
    RFQ_VALUES,
    ROW_TYPE,
    TAGS,
    TITLE,
    TYPE_MARK,
    UL_COMPARE_STATUS,
    UL_DATA_STATUS,
    UL_ORDERED_VALUES,
    UL_REMAINING_VALUES,
    UL_UNITS,
    UL_VALUES,
    UNITS,
    UNITS_2,
    VALUES,
    VENDOR,
    ANNOTATION,
)
from RFQ.ds_compare.ds_packing_grouped_compare import (
    STATUS_COMPLETE,
    STATUS_DATA_PROBLEM,
    STATUS_ACCOUNTED_ABOVE,
    STATUS_PACKING_ONLY,
    STATUS_SHORTFALL,
    STATUS_UNAVAILABLE,
    compare_grouped_rows_with_packing,
)
from RFQ.ds_compare.ds_quantity_parse import try_parse_quantity
from RFQ.ds_compare.tsd_packing_load import load_tsd_file_rows
from RFQ.ds_compare.ds_units_normalize import highlight_compared_units_columns
from RFQ.ds_compare.ds_vs_mto_excel_columns import DS_VS_MTO_OUTPUT_COLUMNS_CONFIG
from RFQ.ds_compare.ds_vs_mto_excel_xlsxwriter import save_ds_vs_mto_excel_xlsxwriter
from RFQ.packing_list_provider import (
    PACKING_CACHE_COLUMNS,
    PACKING_CACHE_VERSION,
    PackingCacheMeta,
    PackingDataset,
    PackingQualityLevel,
    load_packing_dataset,
    packing_row_key,
    packing_row_source,
    slim_packing_row,
    slim_packing_rows,
)
from utils.colors import Color


def _row(
    *,
    title: str,
    system: str,
    code: str,
    quantity: object,
    units: str = "шт",
    code_2: str = "",
    rfq_quantity: object = 0,
    source: str = "packing.xlsx",
    sheet: str = "Single 1",
    excel_row: int = 10,
) -> RowStd:
    row = RowStd.get_std_check_row({}, TableComments())
    row.row_type = RowType.position_row
    row.el[ROW_TYPE].value = RowType.position_row
    values = {
        DS_TITLE: title,
        DS_SYSTEM: system,
        CODE: code,
        CODE_2: code_2,
        VALUES: quantity,
        RFQ_VALUES: rfq_quantity,
        UNITS: units,
        UNITS_2: units,
        NAME: f"Item {code or code_2}",
        TYPE_MARK: "TM",
        VENDOR: "Vendor",
        ANNOTATION: source,
        TITLE: sheet,
    }
    for column, value in values.items():
        row.el[column].value = value
    row._packing_source_row = excel_row
    return row


def _meta(*, critical_ok: bool = True) -> PackingCacheMeta:
    return PackingCacheMeta(
        version=PACKING_CACHE_VERSION,
        created_at="2026-07-29T12:00:00+00:00",
        root="C:/packing",
        fingerprint="abc",
        counters={"files_total": 1, "position_rows": 1},
        critical_ok=critical_ok,
        report_paths={"critical": "tsd_packing_critical.txt"},
    )


class PackingCompareSmokeTest(unittest.TestCase):
    def test_grouped_formulas_code2_errors_and_ul_only(self) -> None:
        compared = [
            _row(
                title="8529",
                system="SOS",
                code="BCC1",
                quantity=5,
                rfq_quantity=2,
            ),
            _row(
                title="8529",
                system="SOS",
                code="OLD",
                code_2="NEW",
                quantity=2,
            ),
        ]
        packing = [
            _row(title="8529", system="SOS", code="BCC1", quantity=3, excel_row=11),
            _row(title="8529", system="SOS", code="BCC1", quantity=4, excel_row=12),
            _row(title="8529", system="SOS", code="NEW", quantity=1, excel_row=13),
            _row(title="8529", system="SOS", code="UL_ONLY", quantity=8, excel_row=14),
            _row(title="8529", system="SOS", code="BAD", quantity="wrong", excel_row=15),
        ]
        dataset = PackingDataset(
            rows=packing,
            meta=_meta(),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path="memory.cache",
        )

        rows, audit = compare_grouped_rows_with_packing(compared, dataset)

        self.assertEqual(compared[0].get_value(UL_ORDERED_VALUES), 7.0)
        self.assertEqual(compared[0].get_value(UL_VALUES), 7.0)
        self.assertEqual(compared[0].get_value(UL_REMAINING_VALUES), 0.0)
        self.assertEqual(compared[0].get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(compared[1].get_value(UL_ORDERED_VALUES), 2.0)
        self.assertEqual(compared[1].get_value(UL_REMAINING_VALUES), 1.0)
        self.assertEqual(compared[1].get_value(UL_COMPARE_STATUS), STATUS_SHORTFALL)
        self.assertIn("MTO code", compared[1].el[UL_ORDERED_VALUES].comment)
        self.assertTrue(
            any(row.get_value(UL_COMPARE_STATUS) == STATUS_PACKING_ONLY for row in rows)
        )
        self.assertTrue(
            any(row.get_value(UL_COMPARE_STATUS) == STATUS_DATA_PROBLEM for row in rows)
        )
        self.assertEqual(audit.stats.invalid_quantity_rows, 1)
        self.assertGreaterEqual(audit.stats.packing_only_added, 2)

    def test_units_include_packing(self) -> None:
        compared = [
            _row(
                title="8529",
                system="SOS",
                code="BCC2",
                quantity=1,
                units="шт",
            )
        ]
        compared[0].el[UNITS_2].value = "компл"
        dataset = PackingDataset(
            rows=[_row(title="8529", system="SOS", code="BCC2", quantity=1, units="шт")],
            meta=_meta(),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path="memory.cache",
        )
        rows, _audit = compare_grouped_rows_with_packing(compared, dataset)
        stats = highlight_compared_units_columns(rows)
        self.assertEqual(stats.mismatch, 1)
        self.assertEqual(rows[0].el[UL_UNITS].color, Color.yellow)

    def test_rfq_only_matches_packing(self) -> None:
        rfq_only = _row(
            title="8529",
            system="SOS",
            code="",
            quantity=0,
            rfq_quantity=3,
        )
        rfq_only.el[RFQ_CODE].value = "RFQ_ONLY"
        rfq_only.el[RFQ_COMPARE_STATUS].value = "Новая"
        packing = _row(
            title="8529",
            system="SOS",
            code="RFQ_ONLY",
            quantity=3,
        )
        dataset = PackingDataset(
            rows=[packing],
            meta=_meta(),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path="memory.cache",
        )
        rows, audit = compare_grouped_rows_with_packing([rfq_only], dataset)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(audit.stats.packing_only_added, 0)

    def test_duplicate_targets_are_not_double_subtracted(self) -> None:
        compared = [
            _row(title="8529", system="SOS", code="DUP", quantity=2),
            _row(title="8529", system="SOS", code="DUP", quantity=3),
        ]
        dataset = PackingDataset(
            rows=[_row(title="8529", system="SOS", code="DUP", quantity=5)],
            meta=_meta(),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path="memory.cache",
        )
        rows, audit = compare_grouped_rows_with_packing(compared, dataset)
        self.assertEqual(rows[0].get_value(UL_ORDERED_VALUES), 5.0)
        self.assertEqual(rows[0].get_value(UL_REMAINING_VALUES), 0.0)
        self.assertEqual(rows[1].get_value(UL_COMPARE_STATUS), STATUS_ACCOUNTED_ABOVE)
        self.assertEqual(audit.stats.duplicate_target_rows, 1)

    def test_invalid_and_unavailable_are_explicit(self) -> None:
        compared = [_row(title="8529", system="SOS", code="BAD", quantity=1)]
        bad_dataset = PackingDataset(
            rows=[
                _row(title="8529", system="SOS", code="BAD", quantity=float("inf")),
                _row(title="8529", system="SOS", code="EMPTY", quantity=""),
            ],
            meta=_meta(),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path="memory.cache",
        )
        rows, audit = compare_grouped_rows_with_packing(compared, bad_dataset)
        self.assertEqual(rows[0].get_value(UL_COMPARE_STATUS), STATUS_DATA_PROBLEM)
        self.assertIsNone(rows[0].get_value(UL_VALUES))
        self.assertIsNone(rows[0].get_value(UL_REMAINING_VALUES))
        self.assertEqual(rows[0].el[UL_VALUES].color, Color.red)
        self.assertEqual(audit.stats.invalid_quantity_rows, 2)

        unavailable = PackingDataset(
            rows=[],
            meta=None,
            quality=PackingQualityLevel.UNAVAILABLE,
            issues=[],
            cache_path="missing.cache",
        )
        rows, _audit = compare_grouped_rows_with_packing(
            [_row(title="8529", system="SOS", code="MISS", quantity=1)],
            unavailable,
        )
        self.assertEqual(rows[0].get_value(UL_COMPARE_STATUS), STATUS_UNAVAILABLE)
        self.assertEqual(rows[0].el[UL_DATA_STATUS].color, Color.red)

    def test_common_quantity_parser_rejects_non_finite(self) -> None:
        self.assertFalse(try_parse_quantity(float("nan"))[0])
        self.assertFalse(try_parse_quantity(float("inf"))[0])
        self.assertFalse(try_parse_quantity("NaN")[0])

    def test_schema_invariant(self) -> None:
        expected = set(ColNames.DsVsMto.column_dict.values())
        configured = {column.col_name for column in DS_VS_MTO_OUTPUT_COLUMNS_CONFIG}
        self.assertEqual(expected, configured)

    def test_tsd_loader_keeps_source_row_after_long_empty_gap(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "packing_source.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Single 1"
            sheet.cell(310, 2, "BCC-LATE")
            sheet.cell(310, 3, "8529-SOS")
            sheet.cell(310, 8, "Late item;TM")
            sheet.cell(310, 9, 1)
            sheet.cell(310, 10, "шт")
            workbook.save(path)
            workbook.close()

            positions, _plan, _all_rows, _hits, issues = load_tsd_file_rows(
                str(path),
                root=tmp,
            )
            self.assertEqual(len(positions), 1)
            self.assertEqual(positions[0]._packing_source_row, 310)
            self.assertFalse(issues)

    def test_tsd_loader_ignores_header_and_signature_required_fields(self) -> None:
        """Header/signature rows must not emit source_required_fields noise."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "packing_header_signature.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Single 7"
            # Document header / banner area (old heuristic: NAME + UNITS → false ERROR).
            sheet.cell(7, 8, "Упаковочный лист")
            sheet.cell(7, 10, "____")
            sheet.cell(8, 4, "Директор Московского филиала")
            sheet.cell(8, 8, "К.Э. Иванов")
            sheet.cell(8, 10, "____")
            # STD letterhead labels mapped into CODE/VALUES (common critical noise).
            sheet.cell(9, 2, "Грузоотправитель / Shipper")
            sheet.cell(9, 9, "Продавец / Seller")
            sheet.cell(10, 2, "Упаковочный лист / Packing list No.")
            sheet.cell(10, 8, "2026-02-13")
            sheet.cell(10, 9, "Место назначения / Destination")
            # Real positions (table band).
            sheet.cell(12, 2, "BCC-OK-1")
            sheet.cell(12, 3, "8529-SOS")
            sheet.cell(12, 8, "Good item;TM")
            sheet.cell(12, 9, 1)
            sheet.cell(12, 10, "шт")
            sheet.cell(14, 2, "BCC-OK-2")
            sheet.cell(14, 3, "8529-SOS")
            sheet.cell(14, 8, "Another;TM")
            sheet.cell(14, 9, 2)
            sheet.cell(14, 10, "шт")
            # Signature / footer after the table (Single 7 · строка 22 pattern).
            sheet.cell(22, 4, "Директор Московского филиала")
            sheet.cell(22, 8, "К.Э. Колесников")
            sheet.cell(22, 10, "____")
            workbook.save(path)
            workbook.close()

            positions, _plan, _all_rows, _hits, issues = load_tsd_file_rows(
                str(path),
                root=tmp,
            )
            self.assertEqual(len(positions), 2)
            required_field_issues = [
                issue for issue in issues if issue.code == "source_required_fields"
            ]
            self.assertEqual(required_field_issues, [])

    def test_tsd_loader_ignores_sopl_banner_and_legend(self) -> None:
        """SO - PL letterhead / package legend must not flood critical remarks."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sopl_banner.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "SO - PL"
            # Letterhead (cols G=7, N=14, O=15 in 1-based Excel).
            sheet.cell(2, 7, "For delivery to / Условия поставки:")
            sheet.cell(2, 14, "Packing list No. / Упаковочный лист №:")
            sheet.cell(2, 15, "PLAN006580FR")
            sheet.cell(3, 7, "Consignee /Получатель:")
            sheet.cell(3, 14, "Date / Дата:")
            sheet.cell(3, 15, "2025-07-22")
            # Column header row.
            sheet.cell(5, 3, "Name of document")
            sheet.cell(5, 6, "Vendor")
            sheet.cell(5, 7, "PO item")
            sheet.cell(5, 13, "Russian translation")
            sheet.cell(5, 14, "Quantity")
            sheet.cell(5, 15, "Units")
            # Real positions.
            sheet.cell(6, 3, "8529-SOS")
            sheet.cell(6, 7, "BCC0000155")
            sheet.cell(6, 13, "Камера")
            sheet.cell(6, 14, 2)
            sheet.cell(6, 15, "шт")
            sheet.cell(6, 57, "8529")
            sheet.cell(6, 58, "SOS")
            sheet.cell(7, 3, "8529-SOS")
            sheet.cell(7, 7, "BCC0000152")
            sheet.cell(7, 13, "Камера 2")
            sheet.cell(7, 14, 1)
            sheet.cell(7, 15, "шт")
            sheet.cell(7, 57, "8529")
            sheet.cell(7, 58, "SOS")
            # Footer legend.
            sheet.cell(10, 7, "Type of package / Вид упаковки:")
            sheet.cell(10, 13, "Stackability / Штабелируемость:")
            sheet.cell(11, 7, "001 = case / Ящик")
            sheet.cell(11, 13, "00 = None / Не допускается")
            workbook.save(path)
            workbook.close()

            positions, _plan, _all_rows, _hits, issues = load_tsd_file_rows(
                str(path),
                root=tmp,
            )
            self.assertEqual(len(positions), 2)
            required_field_issues = [
                issue for issue in issues if issue.code == "source_required_fields"
            ]
            self.assertEqual(required_field_issues, [])

    def test_tsd_loader_flags_incomplete_outside_band_product_code(self) -> None:
        """Product-like CODE just outside the table band is still reported."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "outside_band_bcc.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Single 1"
            # Incomplete BCC row above the table (should still warn).
            sheet.cell(8, 2, "BCC-ABOVE")
            sheet.cell(8, 8, "Almost position")
            sheet.cell(10, 2, "BCC-OK")
            sheet.cell(10, 3, "8529-SOS")
            sheet.cell(10, 8, "Good;TM")
            sheet.cell(10, 9, 1)
            sheet.cell(10, 10, "шт")
            workbook.save(path)
            workbook.close()

            _positions, _plan, _all_rows, _hits, issues = load_tsd_file_rows(
                str(path),
                root=tmp,
            )
            required_field_issues = [
                issue for issue in issues if issue.code == "source_required_fields"
            ]
            self.assertEqual(len(required_field_issues), 1)
            self.assertEqual(required_field_issues[0].excel_row, 8)

    def test_tsd_loader_flags_incomplete_real_position(self) -> None:
        """Incomplete mid-table row with CODE must still raise source_required_fields."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "packing_incomplete.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Single 1"
            sheet.cell(10, 2, "BCC-OK")
            sheet.cell(10, 3, "8529-SOS")
            sheet.cell(10, 8, "Good;TM")
            sheet.cell(10, 9, 1)
            sheet.cell(10, 10, "шт")
            # Damaged position inside the table band: CODE + NAME, no VALUES/UNITS.
            sheet.cell(11, 2, "BCC-BROKEN")
            sheet.cell(11, 3, "8529-SOS")
            sheet.cell(11, 8, "Broken item;TM")
            sheet.cell(12, 2, "BCC-OK-2")
            sheet.cell(12, 3, "8529-SOS")
            sheet.cell(12, 8, "Good 2;TM")
            sheet.cell(12, 9, 1)
            sheet.cell(12, 10, "шт")
            workbook.save(path)
            workbook.close()

            positions, _plan, _all_rows, _hits, issues = load_tsd_file_rows(
                str(path),
                root=tmp,
            )
            self.assertEqual(len(positions), 2)
            required_field_issues = [
                issue for issue in issues if issue.code == "source_required_fields"
            ]
            self.assertEqual(len(required_field_issues), 1)
            self.assertEqual(required_field_issues[0].excel_row, 11)
            self.assertIn("values", required_field_issues[0].field)
            self.assertIn("units", required_field_issues[0].field)

    def test_excel_has_numeric_ul_values_color_and_comment(self) -> None:
        compared = [_row(title="8529", system="SOS", code="XLSX", quantity=2)]
        dataset = PackingDataset(
            rows=[_row(title="8529", system="SOS", code="XLSX", quantity=2)],
            meta=_meta(),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path="memory.cache",
        )
        rows, _audit = compare_grouped_rows_with_packing(compared, dataset)
        with tempfile.TemporaryDirectory() as tmp:
            path = save_ds_vs_mto_excel_xlsxwriter(
                rows,
                tmp,
                "test",
                output_file_name="packing.xlsx",
            )
            self.assertIsNotNone(path)
            workbook = openpyxl.load_workbook(
                io.BytesIO(Path(path).read_bytes()),
                data_only=True,
            )
            try:
                sheet = workbook["DS vs MTO"]
                headers = {
                    cell.value: cell.column
                    for cell in sheet[1]
                    if cell.value is not None
                }
                ordered = sheet.cell(2, headers["Заказано (ДС + RFQ)"])
                status = sheet.cell(2, headers["Статус УЛ"])
                self.assertIsInstance(ordered.value, (int, float))
                self.assertIsNotNone(ordered.comment)
                self.assertEqual(status.value, STATUS_COMPLETE)
                self.assertEqual(status.fill.fgColor.rgb[-6:].lower(), Color.green.lower())
                header = sheet.cell(1, headers["Статус УЛ"])
                self.assertEqual(header.fill.fgColor.rgb[-6:].lower(), "c9daf8")
            finally:
                workbook.close()


class PackingSlimCacheTest(unittest.TestCase):
    """v6 slim pickle: fewer columns, same position keys as full RowStd."""

    def test_slim_row_keeps_compare_fields_and_drops_extra_el(self) -> None:
        full = _row(
            title="8529",
            system="SOS",
            code="BCC-SLIM",
            quantity=3,
            source="folder/file.xlsx",
            sheet="Single 2",
            excel_row=42,
        )
        full.el[NAME].value = "Cable"
        full.el[TYPE_MARK].value = "TM1"
        full.el[TAGS].value = ["T1", "T2"]
        # Full RowStd has the entire ColNames set.
        self.assertGreater(len(full.el), len(PACKING_CACHE_COLUMNS))

        slim = slim_packing_row(full)
        self.assertEqual(set(slim.el), set(PACKING_CACHE_COLUMNS))
        self.assertEqual(slim.row_type, RowType.position_row)
        self.assertEqual(packing_row_key(slim), packing_row_key(full))
        self.assertEqual(packing_row_source(slim), ("folder/file.xlsx", "Single 2", 42))
        self.assertEqual(slim.get_tags_list(), ["T1", "T2"])
        self.assertEqual(slim.get_value(NAME), "Cable")
        self.assertEqual(slim.get_value(VALUES), 3.0)

    def test_slim_pickle_smaller_same_count_and_loadable(self) -> None:
        rows = [
            _row(title="8529", system="SOS", code=f"BCC{i}", quantity=1, excel_row=10 + i)
            for i in range(20)
        ]
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            slim = slim_packing_rows(rows)
            full_path = root / "full.cache"
            slim_path = root / "slim.cache"
            with full_path.open("wb") as stream:
                pickle.dump(
                    {
                        "meta": {
                            "version": PACKING_CACHE_VERSION,
                            "created_at": "2026-07-30T12:00:00+00:00",
                            "root": "C:/packing",
                            "fingerprint": "fp",
                            "counters": {"position_rows": len(rows)},
                            "critical": {"ok": True, "remarks": []},
                            "report_paths": {},
                            "loader_issues": [],
                        },
                        "data": rows,
                    },
                    stream,
                    protocol=pickle.HIGHEST_PROTOCOL,
                )
            with slim_path.open("wb") as stream:
                pickle.dump(
                    {
                        "meta": {
                            "version": PACKING_CACHE_VERSION,
                            "created_at": "2026-07-30T12:00:00+00:00",
                            "root": "C:/packing",
                            "fingerprint": "fp",
                            "counters": {"position_rows": len(slim)},
                            "critical": {"ok": True, "remarks": []},
                            "report_paths": {},
                            "loader_issues": [],
                        },
                        "data": slim,
                    },
                    stream,
                    protocol=pickle.HIGHEST_PROTOCOL,
                )

            self.assertEqual(len(slim), len(rows))
            self.assertLess(slim_path.stat().st_size, full_path.stat().st_size // 2)

            dataset = load_packing_dataset(slim_path)
            self.assertEqual(dataset.quality, PackingQualityLevel.OK)
            self.assertEqual(len(dataset.rows), 20)
            self.assertEqual(packing_row_key(dataset.rows[0]), packing_row_key(rows[0]))
            self.assertEqual(dataset.meta.version, PACKING_CACHE_VERSION)


    def test_slim_rows_work_in_grouped_compare(self) -> None:
        compared = [_row(title="8529", system="SOS", code="BCC1", quantity=5)]
        packing = slim_packing_rows(
            [_row(title="8529", system="SOS", code="BCC1", quantity=5, excel_row=11)]
        )
        dataset = PackingDataset(
            rows=packing,
            meta=_meta(),
            quality=PackingQualityLevel.OK,
            issues=[],
            cache_path="memory.cache",
        )
        rows, audit = compare_grouped_rows_with_packing(compared, dataset)
        self.assertEqual(rows[0].get_value(UL_COMPARE_STATUS), STATUS_COMPLETE)
        self.assertEqual(audit.stats.invalid_quantity_rows, 0)


class PackingProviderSmokeTest(unittest.TestCase):
    def _payload(self, row: RowStd, *, version: str = PACKING_CACHE_VERSION, critical_ok: bool = True) -> dict:
        return {
            "meta": {
                "version": version,
                "created_at": "2026-07-29T12:00:00+00:00",
                "root": "C:/packing",
                "fingerprint": "abc",
                "counters": {"files_total": 1, "position_rows": 1},
                "critical": {"ok": critical_ok, "remarks": []},
                "report_paths": {},
                "loader_issues": [],
            },
            "data": [row],
        }

    def test_ok_partial_missing_corrupt_and_old_cache(self) -> None:
        row = _row(title="8529", system="SOS", code="BCC1", quantity=1)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            cache = root / "packing.cache"

            with cache.open("wb") as stream:
                pickle.dump(self._payload(row), stream)
            self.assertEqual(load_packing_dataset(cache).quality, PackingQualityLevel.OK)

            with cache.open("wb") as stream:
                pickle.dump(self._payload(row, critical_ok=False), stream)
            self.assertEqual(
                load_packing_dataset(cache).quality,
                PackingQualityLevel.PARTIAL,
            )

            with cache.open("wb") as stream:
                pickle.dump(self._payload(row, version="old"), stream)
            self.assertEqual(
                load_packing_dataset(cache).quality,
                PackingQualityLevel.UNAVAILABLE,
            )

            cache.write_bytes(b"not a pickle")
            self.assertEqual(
                load_packing_dataset(cache).quality,
                PackingQualityLevel.UNAVAILABLE,
            )
            cache.unlink()
            self.assertEqual(
                load_packing_dataset(cache).quality,
                PackingQualityLevel.UNAVAILABLE,
            )


if __name__ == "__main__":
    unittest.main()
