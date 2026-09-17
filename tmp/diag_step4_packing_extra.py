"""Extra queries for 8445-SOT1 / BCC0003052 leftovers."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.08.27.13.58"
    r"\Шаг4_Сопоставление_RFP_MTO_20260827_140410.xlsx"
)
FIELDS = (
    "Имя ДС", "№ позиции", "Титул/Марка", "Код RFP", "Код MTO", "Кол-во RFP",
    "Статус УЛ", "Статус тегов УЛ", "Теги УЛ", "Код УЛ",
    "Источник УЛ (файл · вкладка · строка)", "Теги RFP",
)


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def dump_row(rn: int, row: list, idx: dict[str, int], cells) -> None:
    def cell(name: str) -> str:
        i = idx.get(name)
        if i is None or i >= len(row):
            return ""
        v = row[i]
        return "" if v is None else str(v).strip()

    _safe_print(f"\n--- row {rn} ---")
    for f in FIELDS:
        _safe_print(f"  {f}: {cell(f)}")
    i = idx.get("Статус УЛ")
    if i is not None and i < len(cells):
        c = cells[i].comment
        if c:
            _safe_print(f"  Статус УЛ comment: {str(c.text).replace(chr(10), ' ')}")


def main() -> None:
    wb = load_workbook(XLSX, read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {n: i for i, n in enumerate(header) if n}

    _safe_print("=== row 4245 (8445-SOT1) and nearby BCC0003052 complete ===")
    for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row = [c.value for c in cells]
        title = str(row[idx["Титул/Марка"]] or "").strip() if "Титул/Марка" in idx else ""
        code_rfp = str(row[idx["Код RFP"]] or "").strip().upper() if "Код RFP" in idx else ""
        code_ul = str(row[idx["Код УЛ"]] or "").strip().upper() if "Код УЛ" in idx else ""
        tags_ul = str(row[idx["Теги УЛ"]] or "") if "Теги УЛ" in idx else ""
        status = str(row[idx["Статус УЛ"]] or "").strip() if "Статус УЛ" in idx else ""
        if rn in {4245, 2778}:
            dump_row(rn, row, idx, cells)
        if title.startswith("8445") and code_ul == "BCC0003052" and status == "Поставка комплектна":
            dump_row(rn, row, idx, cells)
        if "8445-GA-01-S-FV-0011" in tags_ul:
            dump_row(rn, row, idx, cells)

    _safe_print("\n=== leftovers code BCC0003052 title 8445* ===")
    for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row = [c.value for c in cells]
        status = str(row[idx["Статус УЛ"]] or "").strip()
        if status != "Только в УЛ":
            continue
        title = str(row[idx["Титул/Марка"]] or "").strip()
        code = str(row[idx["Код УЛ"]] or "").strip().upper()
        if not title.startswith("8445") or code != "BCC0003052":
            continue
        dump_row(rn, row, idx, cells)

    _safe_print("\n=== leftovers tag 8445-GA-01-S-FV-0011 ===")
    for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row = [c.value for c in cells]
        if str(row[idx["Статус УЛ"]] or "").strip() != "Только в УЛ":
            continue
        tags = str(row[idx["Теги УЛ"]] or "")
        if "8445-GA-01-S-FV-0011" not in tags:
            continue
        dump_row(rn, row, idx, cells)

    _safe_print("\n=== 8525-SOT BCC0000152 all result rows ===")
    for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row = [c.value for c in cells]
        title = str(row[idx["Титул/Марка"]] or "").strip()
        if title != "8525-SOT":
            continue
        code_rfp = str(row[idx["Код RFP"]] or "").strip().upper()
        code_ul = str(row[idx["Код УЛ"]] or "").strip().upper()
        if code_rfp != "BCC0000152" and code_ul != "BCC0000152":
            continue
        dump_row(rn, row, idx, cells)

    wb.close()


if __name__ == "__main__":
    main()
