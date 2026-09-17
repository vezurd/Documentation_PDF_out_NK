"""Smoke: rfp_parts_net.xlsx core columns match Step1 RFP_AGGREAGATED indices."""

from __future__ import annotations

import sys
import tempfile
import unittest
from collections import Counter
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

import base.t_comm_initial_classes as t_com_init_cls
from RFQ.rfp_parts.analyze_rfp_parts import (
    NetSummaryRow,
    RfpRecord,
    _ensure_net_layout_matches_step1,
    _write_net_xlsx,
    step1_net_core_columns,
    step1_net_extra_columns,
)
from base.base_classes import RowType, TableComments
from base.base_mto import get_std_from_excel_file
from base.tables_columns import (
    CODE,
    ColNames,
    NAME,
    RFP_SUPPLY_STATUS,
    UNITS,
    UNITS_CHECK_STATUS,
    UNITS_CONVERSION_TRACE,
    VALUES,
    VALUES_2,
)


def _record() -> RfpRecord:
    return RfpRecord(
        kind="net",
        file_name="DS10.xlsx",
        sheet="Перечень материалов",
        excel_row=24,
        ds_name="ДС10",
        ds_number="1",
        ds_title="8529-SOS",
        ds_specification="AGCC.287-8529-SOS.MTO-0001",
        tags="8529-SS-01-S-UZ-0711",
        ds_code_1c="",
        code="BCC0000516",
        name="Кабель",
        type_mark="NYM",
        values=Decimal("3"),
        units="шт",
    )


class NetStep1LayoutSmokeTest(unittest.TestCase):
    def test_core_map_follows_rfp_aggregated_integer_keys(self) -> None:
        core = step1_net_core_columns()
        extra = step1_net_extra_columns()
        source = {
            int(idx): name
            for idx, name in ColNames.RFP_AGGREAGATED.column_dict.items()
            if isinstance(idx, int) and 0 <= idx <= 10
        }
        self.assertEqual(core, source)
        self.assertEqual(core[6], CODE)
        self.assertEqual(core[9], VALUES)
        self.assertEqual(core[10], UNITS)
        self.assertEqual(extra[15], RFP_SUPPLY_STATUS)
        self.assertEqual(extra[17], VALUES_2)
        self.assertEqual(extra[18], UNITS_CHECK_STATUS)
        self.assertEqual(extra[19], UNITS_CONVERSION_TRACE)
        self.assertNotIn(8.5, core)
        _ensure_net_layout_matches_step1()

    def test_written_net_is_step1_position_row(self) -> None:
        item = NetSummaryRow(record=_record(), source="Сумма частей")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "rfp_parts_net.xlsx"
            self.assertEqual(_write_net_xlsx(path, [item]), 1)
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb["Перечень материалов"]
            rows = list(ws.iter_rows(values_only=True))
            wb.close()

            header, data = rows[0], rows[1]
            self.assertEqual(header[8], "Технические характеристики")
            self.assertEqual(header[9], "Кол-во")
            self.assertEqual(header[10], "Ед. изм.")
            self.assertNotEqual(header[9], "Поставщик")
            self.assertEqual(header[15], "Статус поставки")
            self.assertEqual(header[17], "Закупка по Лоту, кол-во")
            self.assertEqual(header[18], "Статус ед. изм.")
            self.assertEqual(header[19], "Trace ед. изм.")
            self.assertEqual(header[20], "Источник строки")
            self.assertEqual(header[25], "Через замену кода")
            self.assertEqual(data[6], "BCC0000516")
            self.assertEqual(data[9], 3)
            self.assertEqual(data[10], "шт")
            self.assertEqual(data[17], 3)

            t_com = TableComments(
                file_full_path=str(path),
                dir_path="-1",
                tabel_class=t_com_init_cls.RFP_AGGREGATED,
            )
            loaded = get_std_from_excel_file(t_com)
        types = Counter(row.row_type for row in loaded)
        self.assertGreaterEqual(types[RowType.position_row], 1)
        pos = next(row for row in loaded if row.row_type == RowType.position_row)
        self.assertEqual(pos.el[CODE].value, "BCC0000516")
        self.assertEqual(pos.el[NAME].value, "Кабель")
        self.assertEqual(int(pos.el[VALUES].value), 3)
        self.assertEqual(pos.el[UNITS].value, "шт")
        supply_cell = pos.el.get(RFP_SUPPLY_STATUS)
        if supply_cell is not None:
            self.assertIn(supply_cell.value, (None, ""))


if __name__ == "__main__":
    unittest.main()
