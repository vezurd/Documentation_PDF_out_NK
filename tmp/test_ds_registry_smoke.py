"""Offline smoke for the canonical DS registry (no UNC writes)."""

from __future__ import annotations

import os
import re
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.rfp_parts.ds_registry import (
    HEADER_CONDITIONAL_MARK,
    HEADER_REQUIRED_MARK,
    HDR_SOURCE_ID,
    HDR_STATUS,
    ISSUE_DUPLICATE_LINK,
    ISSUE_IDENTITY_SPLIT,
    ISSUE_ORPHAN,
    ISSUE_RFP_FILE_MISSING,
    ISSUE_SHARED_RFP,
    ISSUE_SHARED_UL,
    MIGRATION_REPORT_PREFIX,
    MODE_NEEDS_SPLIT,
    MODE_NO_UL,
    MODE_WHOLE,
    RFP_STATUS_MATCH,
    RFP_STATUS_MISS,
    REGISTRY_SHEET_NAME,
    REGISTRY_TABLE_NAME,
    STATUS_ACTIVE,
    STATUS_DISABLED,
    STATUS_HISTORY,
    RegistryLinks,
    DsRegistryRelation,
    DsRegistryReplaceError,
    DsRegistryRow,
    backup_and_replace_registry,
    canonical_supply_group_id,
    DsRegistryDocument,
    collect_orphan_rows,
    collapse_registry_rows,
    copy_registry_snapshot,
    detect_registry_format,
    legal_rfp_file_clusters,
    load_registry,
    migrate_registry,
    no_ul_group_id,
    registry_packing_maps,
    validate_registry_rows,
    write_registry_workbook,
)


def _load_xlsx(path: Path):
    return load_workbook(BytesIO(path.read_bytes()))


def _fill_rgb(cell) -> str:
    color = getattr(cell.fill, "fgColor", None)
    rgb = getattr(color, "rgb", None) if color is not None else None
    return str(rgb or "").upper()


def _write_legacy(path: Path, rows: list[tuple[object, object, object, object]]) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "Актуальный ДС"
    ws["B1"] = "Старый ДС"
    ws["C1"] = "УЛ"
    ws["D1"] = "Примечание"
    for excel_row, (current, old, ul_text, note) in enumerate(rows, start=2):
        ws.cell(row=excel_row, column=1, value=current)
        ws.cell(row=excel_row, column=2, value=old)
        ws.cell(row=excel_row, column=3, value=ul_text)
        ws.cell(row=excel_row, column=4, value=note)
    wb.save(path)
    wb.close()


def _legacy_53_style_rows() -> list[tuple[object, object, object, object]]:
    return [
        (13, "старый 13", "согл УЛ ДС13", ""),
        (47, "", "согл УЛ ДС13", ""),
        (8, "", "согл УЛ ДС8;согл УЛ ДС81", "два УЛ"),
        (101, "", "", "без упаковки"),
        ("4905_1", "4905", "согл УЛ 4905", ""),
        (4905, "", "согл УЛ 4905", ""),
        ("", "99", "", "архивная запись без актуального ДС"),
    ]


class DsRegistrySmokeTest(unittest.TestCase):
    def test_legacy_migration_53_style_patterns(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            legacy_path = root / "Реестр_ДС_УЛ.xlsx"
            output_path = root / "Реестр_ДС_новый.xlsx"
            report_path = root / f"{MIGRATION_REPORT_PREFIX}_smoke.xlsx"
            _write_legacy(legacy_path, _legacy_53_style_rows())
            self.assertEqual(detect_registry_format(legacy_path), "legacy")

            result = migrate_registry(
                legacy_path, output_path, report_path=report_path
            )
            self.assertTrue(output_path.is_file())
            self.assertTrue(report_path.is_file())
            self.assertEqual(result.source_path, legacy_path)
            self.assertEqual(result.output_path, output_path)
            self.assertEqual(result.report_path, report_path)
            self.assertEqual(result.active_count, 6)
            self.assertEqual(result.history_count, 1)
            self.assertFalse(result.validation.is_ok)
            self.assertTrue(result.validation.blocks_overlay)
            codes = {item.code for item in result.validation.issues}
            self.assertNotIn(ISSUE_SHARED_UL, codes)
            self.assertIn(ISSUE_SHARED_RFP, codes)

            by_id = {row.source_id: row for row in result.rows if row.source_id}
            self.assertEqual(by_id["13"].status, STATUS_ACTIVE)
            self.assertEqual(by_id["13"].previous_ds, "старый 13")
            self.assertEqual(by_id["13"].revision, "")
            self.assertEqual(by_id["13"].relations[0].group_id, "ДС13")
            self.assertEqual(by_id["13"].relations[0].rfp_key, "13")
            self.assertEqual(by_id["13"].relations[0].ul_folder, "согл УЛ ДС13")

            self.assertEqual(by_id["47"].relations[0].group_id, "ДС47")
            self.assertEqual(by_id["47"].relations[0].rfp_key, "13")
            self.assertEqual(by_id["47"].relations[0].ul_folder, "согл УЛ ДС13")

            row8 = by_id["8"]
            self.assertEqual(len(row8.relations), 2)
            self.assertEqual(row8.relations[0].group_id, "ДС8")
            self.assertEqual(row8.relations[0].rfp_key, "8")
            self.assertEqual(row8.relations[0].ul_folder, "согл УЛ ДС8")
            self.assertEqual(row8.relations[1].group_id, "ДС8")
            self.assertEqual(row8.relations[1].rfp_key, "81")
            self.assertEqual(row8.relations[1].ul_folder, "согл УЛ ДС81")

            row101 = by_id["101"]
            self.assertEqual(row101.relations[0].group_id, "ДС101")
            self.assertEqual(row101.relations[0].rfp_key, "101")
            self.assertEqual(row101.relations[0].ul_folder, "")

            self.assertIn("4905_1", by_id)
            self.assertEqual(by_id["4905_1"].source_id, "4905_1")
            self.assertEqual(by_id["4905_1"].relations[0].group_id, "ДС4905_1")
            self.assertEqual(by_id["4905"].relations[0].group_id, "ДС4905")
            self.assertEqual(by_id["4905"].relations[0].rfp_key, "4905")
            self.assertEqual(by_id["4905"].relations[0].ul_folder, "согл УЛ 4905")

            history = [row for row in result.rows if row.is_history]
            self.assertEqual(len(history), 1)
            self.assertEqual(history[0].status, STATUS_HISTORY)
            self.assertEqual(history[0].source_id, "")
            self.assertEqual(history[0].previous_ds, "99")
            self.assertEqual(history[0].note, "архивная запись без актуального ДС")
            self.assertEqual(history[0].relations, ())
            self.assertEqual(history[0].original_excel_row, 8)
            self.assertEqual(history[0].revision, "")

            report_wb = _load_xlsx(report_path)
            try:
                self.assertEqual(
                    set(report_wb.sheetnames),
                    {"Сводка", "Соответствие строк", "Проблемы"},
                )
                mapping = report_wb["Соответствие строк"]
                mapped_ids = [
                    str(row[3] or "")
                    for row in mapping.iter_rows(min_row=2, values_only=True)
                ]
                self.assertIn("13", mapped_ids)
                self.assertIn("4905_1", mapped_ids)
                self.assertIn("", mapped_ids)
            finally:
                report_wb.close()

    def test_writer_headers_table_validations_colors_comments(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "registry.xlsx"
            write_registry_workbook(
                path,
                [
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="23",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС23",
                                rfp_key="23",
                                ul_folder="согл УЛ ДС23",
                                mode=MODE_WHOLE,
                                block_index=1,
                            ),
                        ),
                        original_excel_row=2,
                    )
                ],
            )
            wb = _load_xlsx(path)
            try:
                self.assertEqual(
                    wb.sheetnames,
                    [REGISTRY_SHEET_NAME, "Как заполнять", "Исходный ДС"],
                )
                legend = wb["Как заполнять"]
                self.assertIn("Как заполнять реестр", str(legend["A1"].value))
                names = [
                    str(legend.cell(row, 1).value or "")
                    for row in range(4, 20)
                ]
                self.assertIn("Несколько строк", names)
                self.assertIn("Сверка RFP", names)
                ws = wb[REGISTRY_SHEET_NAME]
                self.assertEqual(ws.freeze_panes, "A2")
                self.assertIn(REGISTRY_TABLE_NAME, ws.tables)
                table = ws.tables[REGISTRY_TABLE_NAME]
                self.assertTrue(str(table.ref).startswith("A1:"))
                self.assertTrue(ws["A1"].alignment.wrap_text)
                self.assertEqual(ws["A1"].value, HDR_STATUS)
                self.assertEqual(ws["B1"].value, HDR_SOURCE_ID)
                self.assertEqual(ws["B2"].value, "ДС23")
                self.assertEqual(ws["C1"].value, "Исторический номер ДС")
                self.assertIn(HEADER_REQUIRED_MARK, str(ws["A1"].value))
                self.assertIn(HEADER_CONDITIONAL_MARK, str(ws["E1"].value))
                self.assertTrue(_fill_rgb(ws["A1"]).endswith("1B4F72"))
                self.assertTrue(_fill_rgb(ws["B1"]).endswith("1B4F72"))
                self.assertTrue(_fill_rgb(ws["E1"]).endswith("FFC000"))
                self.assertTrue(_fill_rgb(ws["D1"]).endswith("D9D9D9"))
                self.assertIsNotNone(ws["A1"].comment)
                self.assertIn("Активен", ws["A1"].comment.text)
                self.assertIn("*", ws["A1"].comment.text)
                self.assertIsNotNone(ws["E1"].comment)
                self.assertIn("Одна папка", ws["E1"].comment.text)
                info = wb["Исходный ДС"]
                self.assertEqual(info["A2"].value, "ДС23")
                self.assertEqual(info["E2"].value, 2)
                self.assertIsNone(ws["D2"].value)
                formulas = [
                    str(dv.formula1) for dv in ws.data_validations.dataValidation
                ]
                joined = " ".join(formulas)
                self.assertIn(STATUS_ACTIVE, joined)
                self.assertIn(STATUS_HISTORY, joined)
                self.assertTrue(ws.conditional_formatting._cf_rules)
            finally:
                wb.close()

    def test_parser_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "roundtrip.xlsx"
            rows = [
                DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id="13",
                    previous_ds="old",
                    note="keep",
                    relations=(
                        DsRegistryRelation(
                            group_id="ДС13",
                            rfp_key="13",
                            ul_folder="согл УЛ ДС13",
                            mode=MODE_WHOLE,
                            block_index=1,
                        ),
                    ),
                    original_excel_row=10,
                ),
                DsRegistryRow(
                    status=STATUS_HISTORY,
                    source_id="",
                    previous_ds="88",
                    note="архив",
                    relations=(),
                    original_excel_row=11,
                ),
                DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id="4905_1",
                    relations=(
                        DsRegistryRelation(
                            group_id="ДС4905",
                            rfp_key="4905",
                            ul_folder="согл УЛ 4905",
                            mode=MODE_WHOLE,
                            block_index=1,
                        ),
                    ),
                    original_excel_row=12,
                ),
            ]
            write_registry_workbook(path, rows)
            loaded = load_registry(path)
            self.assertEqual(detect_registry_format(path), "new")
            self.assertTrue(loaded.validation.is_ok)
            self.assertEqual(len(loaded.rows), 3)
            self.assertEqual(loaded.rows[0].source_id, "13")
            self.assertEqual(loaded.rows[0].previous_ds, "old")
            self.assertEqual(loaded.rows[0].note, "keep")
            self.assertEqual(loaded.rows[0].original_excel_row, 10)
            self.assertEqual(loaded.rows[0].relations[0].group_id, "ДС13")
            self.assertEqual(loaded.rows[0].relations[0].rfp_key, "13")
            self.assertEqual(loaded.rows[1].status, STATUS_HISTORY)
            self.assertEqual(loaded.rows[1].relations, ())
            self.assertEqual(loaded.rows[2].source_id, "4905_1")
            self.assertEqual(loaded.rows[2].relations[0].group_id, "ДС4905_1")
            self.assertEqual(loaded.rows[2].relations[0].rfp_key, "4905")

    def test_validate_duplicate_partial_gap_mismatch(self) -> None:
        duplicate = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            ),
        ]
        dup_codes = {item.code for item in validate_registry_rows(duplicate).issues}
        self.assertIn(ISSUE_DUPLICATE_LINK, dup_codes)

        partial = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="1",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="",
                        rfp_key="",
                        ul_folder="согл УЛ ДС1",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            )
        ]
        partial_codes = {item.code for item in validate_registry_rows(partial).issues}
        self.assertEqual(partial_codes, set())

        gap = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="8",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="",
                        rfp_key="",
                        ul_folder="",
                        mode="",
                        block_index=1,
                    ),
                    DsRegistryRelation(
                        group_id="ДС81",
                        rfp_key="81",
                        ul_folder="согл УЛ ДС81",
                        mode=MODE_WHOLE,
                        block_index=2,
                    ),
                ),
            )
        ]
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "gap.xlsx"
            write_registry_workbook(path, gap)
            loaded = load_registry(path)
            self.assertTrue(loaded.validation.is_ok, loaded.validation.issues)
            self.assertEqual(len(loaded.rows[0].relations), 1)
            self.assertEqual(loaded.rows[0].relations[0].ul_folder, "согл УЛ ДС81")

        mismatch = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="47",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13 доп",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            ),
        ]
        mismatch_codes = {item.code for item in validate_registry_rows(mismatch).issues}
        self.assertIn(ISSUE_SHARED_RFP, mismatch_codes)

    def test_ul_identity_roundtrip_and_mismatch(self) -> None:
        self.assertEqual(canonical_supply_group_id(13), "ДС13")
        self.assertEqual(canonical_supply_group_id(4905), "ДС4905")
        self.assertEqual(canonical_supply_group_id("4905"), "ДС4905")
        self.assertEqual(no_ul_group_id("101"), "NO_UL:101")

        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "identity_roundtrip.xlsx"
            write_registry_workbook(
                path,
                [
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="13",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС13",
                                rfp_key="13",
                                ul_folder="согл УЛ ДС13",
                                mode=MODE_WHOLE,
                                block_index=1,
                            ),
                        ),
                    ),
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="4905_1",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС4905",
                                rfp_key="4905",
                                ul_folder="согл УЛ 4905",
                                mode=MODE_WHOLE,
                                block_index=1,
                            ),
                        ),
                    ),
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="101",
                        relations=(
                            DsRegistryRelation(
                                group_id=no_ul_group_id("101"),
                                rfp_key="101",
                                ul_folder="",
                                mode=MODE_NO_UL,
                                block_index=1,
                            ),
                        ),
                    ),
                ],
            )
            loaded = load_registry(path)
            self.assertTrue(loaded.validation.is_ok)
            self.assertEqual(loaded.rows[0].relations[0].group_id, "ДС13")
            self.assertEqual(loaded.rows[0].relations[0].rfp_key, "13")
            self.assertEqual(loaded.rows[1].relations[0].group_id, "ДС4905_1")
            self.assertEqual(loaded.rows[1].relations[0].rfp_key, "4905")
            self.assertEqual(loaded.rows[2].relations[0].group_id, "ДС101")

        bare_group = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            )
        ]
        bare_codes = {item.code for item in validate_registry_rows(bare_group).issues}
        self.assertEqual(bare_codes, set())

        wrong_key = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="47",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            )
        ]
        key_codes = {item.code for item in validate_registry_rows(wrong_key).issues}
        self.assertEqual(key_codes, set())

        unparsed_filled = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="папка без номера ДС",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            )
        ]
        unparsed_codes = {
            item.code for item in validate_registry_rows(unparsed_filled).issues
        }
        self.assertEqual(unparsed_codes, set())

        no_ul = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="101",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id=no_ul_group_id("101"),
                        rfp_key="101",
                        ul_folder="",
                        mode=MODE_NO_UL,
                        block_index=1,
                    ),
                ),
            )
        ]
        no_ul_codes = {item.code for item in validate_registry_rows(no_ul).issues}
        self.assertEqual(no_ul_codes, set())

    def test_backup_atomic_replace_and_confirm_false(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            target = root / "Реестр_ДС_УЛ.xlsx"
            incoming = root / "new.xlsx"
            write_registry_workbook(
                target,
                [
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="1",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС1",
                                rfp_key="1",
                                ul_folder="согл УЛ ДС1",
                                mode=MODE_WHOLE,
                                block_index=1,
                            ),
                        ),
                    )
                ],
            )
            write_registry_workbook(
                incoming,
                [
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="2",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС2",
                                rfp_key="2",
                                ul_folder="согл УЛ ДС2",
                                mode=MODE_WHOLE,
                                block_index=1,
                            ),
                        ),
                    )
                ],
            )
            before = target.read_bytes()
            with self.assertRaises(DsRegistryReplaceError) as refused:
                backup_and_replace_registry(target, incoming, confirm=False)
            self.assertIn("confirm=True", str(refused.exception))
            self.assertEqual(target.read_bytes(), before)
            self.assertEqual(list(root.glob("*_backup_*")), [])

            replaced = backup_and_replace_registry(target, incoming, confirm=True)
            self.assertTrue(replaced.replaced)
            self.assertTrue(replaced.backup_path.is_file())
            self.assertEqual(load_registry(target).rows[0].source_id, "2")
            self.assertEqual(load_registry(replaced.backup_path).rows[0].source_id, "1")

            original_after = target.read_bytes()
            with patch(
                "RFQ.rfp_parts.ds_registry.os.replace",
                side_effect=PermissionError("locked"),
            ):
                with self.assertRaises(DsRegistryReplaceError) as locked:
                    backup_and_replace_registry(target, incoming, confirm=True)
            self.assertIn("занят", str(locked.exception))
            self.assertEqual(target.read_bytes(), original_after)
            self.assertEqual(load_registry(target).rows[0].source_id, "2")


class LockedRegistryCopySmokeTest(unittest.TestCase):
    def test_open_excel_writes_sibling(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            target = Path(raw) / "Реестр_ДС_УЛ_migrated.xlsx"
            target.write_bytes(b"locked")
            real_replace = os.replace

            def fake_replace(src, dst):
                if Path(dst) == target:
                    raise PermissionError("locked")
                return real_replace(src, dst)

            row = DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        mode=MODE_WHOLE,
                        block_index=1,
                    ),
                ),
            )
            with patch(
                "RFQ.rfp_parts.ds_registry.os.replace", side_effect=fake_replace
            ):
                written = write_registry_workbook(target, [row])
            self.assertNotEqual(written.resolve(), target.resolve())
            self.assertIn("_новый_", written.name)
            self.assertEqual(target.read_bytes(), b"locked")
            self.assertEqual(load_registry(written).rows[0].source_id, "13")


class RegistrySheetLinksSmokeTest(unittest.TestCase):
    def test_ds_number_file_link_and_yellow_split(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds_file = root / "ДС11_spec.xlsx"
            ds_file.write_bytes(b"ds")
            rfp_file = root / "ДС13. sample.xlsx"
            rfp_file.write_bytes(b"rfp")
            path = root / "registry.xlsx"
            rows = [
                DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id="11",
                    previous_ds="2",
                    relations=(
                        DsRegistryRelation(
                            group_id="ДС11",
                            rfp_key="13",
                            ul_folder="согл УЛ ДС11",
                            mode=MODE_WHOLE,
                            block_index=1,
                        ),
                    ),
                ),
                DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id="8",
                    relations=(
                        DsRegistryRelation(
                            group_id="ДС1",
                            rfp_key="99",
                            ul_folder="согл УЛ ДС1",
                            mode=MODE_NEEDS_SPLIT,
                            block_index=1,
                        ),
                    ),
                ),
            ]
            write_registry_workbook(
                path,
                rows,
                links=RegistryLinks(
                    ds_files={"11": (ds_file,)},
                    rfp_files={"13": (rfp_file,)},
                    scanned_ds=True,
                    scanned_rfp=True,
                    reconcile_by_source={"11": RFP_STATUS_MATCH, "8": RFP_STATUS_MISS},
                ),
            )
            loaded = load_registry(path)
            self.assertEqual(loaded.rows[0].source_id, "11")
            self.assertEqual(loaded.rows[0].previous_ds, "2")
            self.assertEqual(loaded.rows[0].relations[0].rfp_key, "13")
            self.assertTrue(loaded.validation.is_ok, loaded.validation.issues)
            wb = _load_xlsx(path)
            try:
                ws = wb[REGISTRY_SHEET_NAME]
                self.assertEqual(ws["B2"].value, "ДС11")
                self.assertEqual(ws["C2"].value, "2")
                self.assertEqual(ws["D2"].value, ds_file.name)
                self.assertNotIn(str(root), str(ws["D2"].value))
                self.assertEqual(ws["G2"].value, rfp_file.name)
                self.assertEqual(ws["H2"].value, RFP_STATUS_MATCH)
                self.assertEqual(ws["H3"].value, RFP_STATUS_MISS)
                self.assertTrue(_fill_rgb(ws["F3"]).endswith("FFFF00"))
                self.assertFalse(_fill_rgb(ws["E2"]).endswith("FFFF00"))
            finally:
                wb.close()

    def test_missing_ul_folder_is_blue_and_shared_folder_is_not_yellow(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "registry.xlsx"

            def row(source_id: str, folder: str) -> DsRegistryRow:
                return DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id=source_id,
                    relations=(
                        DsRegistryRelation(
                            group_id=f"ДС{source_id}",
                            rfp_key=source_id,
                            ul_folder=folder,
                            mode=MODE_WHOLE,
                            block_index=1,
                        ),
                    ),
                )

            present = "согл УЛ ДС11"
            missing = "согл УЛ нет на диске"
            shared = "согл УЛ общая"
            write_registry_workbook(
                path,
                [row("11", present), row("12", missing), row("13", shared), row("47", shared)],
                links=RegistryLinks(
                    scanned_ul=True,
                    ul_names=frozenset({"согл ул дс11"}),
                ),
            )
            loaded = load_registry(path)
            codes = {issue.code for issue in loaded.validation.issues}
            self.assertNotIn("ul_folder_missing", codes)
            self.assertNotIn(ISSUE_SHARED_UL, codes)
            self.assertTrue(loaded.validation.is_ok, loaded.validation.issues)
            wb = _load_xlsx(path)
            try:
                ws = wb[REGISTRY_SHEET_NAME]
                self.assertFalse(_fill_rgb(ws["E2"]).endswith("BDD7EE"))
                self.assertTrue(_fill_rgb(ws["E3"]).endswith("BDD7EE"))
                self.assertFalse(_fill_rgb(ws["E4"]).endswith("FFFF00"))
                self.assertFalse(_fill_rgb(ws["E5"]).endswith("FFFF00"))
                self.assertTrue(_fill_rgb(ws["E4"]).endswith("BDD7EE"))
                self.assertTrue(_fill_rgb(ws["E5"]).endswith("BDD7EE"))
            finally:
                wb.close()

    def test_rfp_file_cell_is_one_name_per_line(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            present = root / "ДС13._ (5)_7230-SOT.xlsx"
            present.write_bytes(b"a")
            other = root / "ДС13._ (6)_7412-SKUD.xlsx"
            other.write_bytes(b"b")
            missing = "нет такого.xlsx"

            def row(source_id: str, rfp_file: str, rfp_key: str) -> DsRegistryRow:
                return DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id=source_id,
                    relations=(
                        DsRegistryRelation(
                            group_id=f"ДС{source_id}",
                            rfp_key=rfp_key,
                            ul_folder="",
                            mode=MODE_WHOLE,
                            rfp_file=rfp_file,
                            block_index=1,
                        ),
                    ),
                )

            mixed = f"{present.name}\n{missing}"
            checked = validate_registry_rows(
                [row("13", mixed, "13")], rfp_root=root
            )
            self.assertTrue(checked.is_ok, checked.issues)
            missing_issues = [
                issue
                for issue in checked.issues
                if issue.code == ISSUE_RFP_FILE_MISSING
            ]
            self.assertEqual(len(missing_issues), 1)
            self.assertIn(missing, missing_issues[0].message)
            self.assertNotIn(present.name, missing_issues[0].message)

            both = f"{present.name}\n{other.name}"
            complete = validate_registry_rows(
                [row("13", both, "13")], rfp_root=root
            )
            self.assertFalse(
                any(issue.code == ISSUE_RFP_FILE_MISSING for issue in complete.issues)
            )

            path = root / "registry.xlsx"
            write_registry_workbook(
                path,
                [row("13", both, "13")],
                links=RegistryLinks(
                    scanned_rfp=True,
                    rfp_files={"13": (present, other)},
                ),
            )
            wb = _load_xlsx(path)
            try:
                ws = wb[REGISTRY_SHEET_NAME]
                self.assertIn(present.name, str(ws["G2"].value))
                self.assertIn(other.name, str(ws["G2"].value))
                self.assertFalse(_fill_rgb(ws["G2"]).endswith("FFFF00"))
                self.assertIsNone(ws["G2"].comment)
            finally:
                wb.close()

            write_registry_workbook(
                path,
                [row("13", mixed, "13")],
                links=RegistryLinks(
                    scanned_rfp=True,
                    rfp_files={"13": (present,)},
                ),
            )
            wb = _load_xlsx(path)
            try:
                ws = wb[REGISTRY_SHEET_NAME]
                self.assertTrue(_fill_rgb(ws["G2"]).endswith("FFFF00"))
                comment = "" if ws["G2"].comment is None else str(ws["G2"].comment.text)
                self.assertIn(missing, comment)
                self.assertNotIn(present.name, comment)
                self.assertNotIn("\\n", comment)
            finally:
                wb.close()

    def test_no_ul_warns_only_when_folder_exists(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            (root / "согл УЛ ДС101").mkdir()
            present = DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="101",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="NO_UL:101",
                        rfp_key="101",
                        ul_folder="",
                        mode=MODE_NO_UL,
                        block_index=1,
                    ),
                ),
            )
            absent = DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="102",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="NO_UL:102",
                        rfp_key="102",
                        ul_folder="",
                        mode=MODE_NO_UL,
                        block_index=1,
                    ),
                ),
            )
            named_missing = DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="103",
                excel_row=4,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС103",
                        rfp_key="",
                        ul_folder="согл УЛ нет на диске",
                        mode="",
                        block_index=1,
                    ),
                ),
            )
            warned = validate_registry_rows(
                [present, absent, named_missing], ul_root=root
            )
            self.assertTrue(warned.is_ok, warned.issues)
            orphans = collect_orphan_rows([present, absent], ul_root=root)
            self.assertTrue(
                any(rel.ul_folder == "согл УЛ ДС101" for row in orphans for rel in row.relations)
            )

    def test_legacy_header_name_still_loads(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "old.xlsx"
            wb = Workbook()
            ws = wb.active
            assert ws is not None
            ws.title = REGISTRY_SHEET_NAME
            headers = [
                "Статус *",
                "ID ДС источника *",
                "Предыдущий / старый ДС",
                "Ревизия",
                "Примечание",
                "Исходная строка реестра",
                "ID группы поставки 1 *",
                "Фактический ДС / ключ RFP 1 *",
                "Папка УЛ 1 †",
                "Режим распределения 1 *",
                "Фильтр титула 1 †",
                "Фильтр марки 1 †",
            ]
            for col, header in enumerate(headers, start=1):
                ws.cell(row=1, column=col, value=header)
            values = [
                "Активен",
                "ДС11",
                "2",
                "",
                "заметка",
                4,
                "ДС11",
                "11",
                "согл УЛ ДС11",
                "Вся ДС",
                "",
                "",
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=2, column=col, value=value)
            wb.save(path)
            wb.close()
            loaded = load_registry(path)
            self.assertEqual(loaded.rows[0].source_id, "11")
            self.assertEqual(loaded.rows[0].previous_ds, "2")
            self.assertEqual(loaded.rows[0].note, "заметка")
            self.assertEqual(loaded.rows[0].relations[0].rfp_key, "11")

    def test_one_bag_shared_folder_and_packing_key(self) -> None:
        folder_1 = "согл УЛ ДС1 ГФ 5титулов"
        folder_7 = "согл УЛ ДС7 ГФ 2 тит"
        ds8 = DsRegistryRow(
            status=STATUS_ACTIVE,
            source_id="8",
            relations=(
                DsRegistryRelation(
                    group_id="ДС1",
                    rfp_key="8",
                    ul_folder=folder_1,
                    mode="",
                    block_index=1,
                ),
                DsRegistryRelation(
                    group_id="ДС7",
                    rfp_key="1",
                    ul_folder=folder_7,
                    mode="",
                    block_index=2,
                ),
            ),
        )
        collapsed, collapse_issues = collapse_registry_rows([ds8])
        self.assertEqual(collapse_issues, [])
        self.assertEqual(len(collapsed), 1)
        self.assertEqual(
            {rel.group_id for rel in collapsed[0].relations},
            {"ДС8"},
        )
        clean = validate_registry_rows(collapsed)
        self.assertTrue(clean.is_ok, clean.issues)
        document = DsRegistryDocument(
            path=Path("registry.xlsx"),
            rows=collapsed,
            validation=clean,
            max_relation_blocks=1,
        )
        maps = registry_packing_maps(document)
        self.assertIsNotNone(maps)
        assert maps is not None
        self.assertEqual(maps[0][folder_1.casefold()], 8)
        self.assertEqual(maps[0][folder_7.casefold()], 8)

        from base.base_classes import RowStd
        from base.tables_columns import ANNOTATION, DS_NAME
        from RFQ.packing_list_provider import (
            build_registry_planting_index,
            packing_row_actual_ds,
            registry_packing_scope,
            rfp_row_actual_ds,
        )

        index = build_registry_planting_index(document)
        packing = RowStd()
        packing.el[ANNOTATION].value = folder_1 + "/list.xlsx"
        rfp = RowStd()
        rfp.el[DS_NAME].value = "ДС1"
        self.assertEqual(packing_row_actual_ds(packing), 1)
        self.assertEqual(rfp_row_actual_ds(rfp), 1)
        with registry_packing_scope(index):
            self.assertEqual(packing_row_actual_ds(packing), 8)
            self.assertEqual(rfp_row_actual_ds(rfp), 8)

        shared = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="14",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС14",
                        rfp_key="14",
                        ul_folder="согл УЛ ДС14",
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="48",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС48",
                        rfp_key="14",
                        ul_folder="согл УЛ ДС14",
                        block_index=1,
                    ),
                ),
            ),
        ]
        blocked = validate_registry_rows(shared)
        codes = {item.code for item in blocked.issues}
        self.assertNotIn(ISSUE_SHARED_UL, codes)
        self.assertIn(ISSUE_SHARED_RFP, codes)
        self.assertTrue(blocked.blocks_overlay)
        dirty = DsRegistryDocument(
            path=Path("shared.xlsx"),
            rows=shared,
            validation=blocked,
            max_relation_blocks=1,
        )
        self.assertIsNone(registry_packing_maps(dirty))

    def test_closed_rfp_file_cluster_is_legal(self) -> None:
        shared_file = "ДС15_61. AGCC.287-0000-12.4.1-RFP-0009_02_RU.xlsx"

        def row(source_id: str, rfp_key: str) -> DsRegistryRow:
            return DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id=source_id,
                relations=(
                    DsRegistryRelation(
                        group_id=f"ДС{source_id}",
                        rfp_key=rfp_key,
                        ul_folder=f"согл УЛ ДС{source_id}",
                        rfp_file=shared_file,
                        block_index=1,
                    ),
                ),
            )

        different_numbers = [row("15", "15"), row("61", "61")]
        checked = validate_registry_rows(different_numbers)
        self.assertTrue(checked.is_ok, checked.issues)
        self.assertNotIn(
            ISSUE_SHARED_RFP, {item.code for item in checked.issues}
        )
        self.assertEqual(
            legal_rfp_file_clusters(different_numbers),
            (frozenset({"15", "61"}),),
        )

        same_number = [row("15", "15"), row("61", "15")]
        checked_same = validate_registry_rows(same_number)
        self.assertTrue(checked_same.is_ok, checked_same.issues)
        self.assertNotIn(
            ISSUE_SHARED_RFP, {item.code for item in checked_same.issues}
        )

        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "cluster.xlsx"
            write_registry_workbook(path, same_number)
            loaded = load_registry(path)
            self.assertTrue(loaded.validation.is_ok, loaded.validation.issues)
            self.assertNotIn(
                ISSUE_SHARED_RFP,
                {item.code for item in loaded.validation.issues},
            )
            ids = {item.source_id for item in loaded.rows if item.source_id}
            self.assertEqual({"15", "61"}, ids)
            book = _load_xlsx(path)
            try:
                sheet = book[REGISTRY_SHEET_NAME]
                self.assertFalse(_fill_rgb(sheet["F2"]).endswith("FFFF00"))
                self.assertFalse(_fill_rgb(sheet["F3"]).endswith("FFFF00"))
                self.assertFalse(_fill_rgb(sheet["G2"]).endswith("FFFF00"))
                self.assertFalse(_fill_rgb(sheet["G3"]).endswith("FFFF00"))
            finally:
                book.close()

    def test_partial_rfp_file_overlap_is_shared(self) -> None:
        rows = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="15",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС15",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС15",
                        rfp_file="A.xlsx\nB.xlsx",
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="61",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС61",
                        rfp_key="61",
                        ul_folder="согл УЛ ДС61",
                        rfp_file="A.xlsx",
                        block_index=1,
                    ),
                ),
            ),
        ]
        checked = validate_registry_rows(rows)
        self.assertFalse(checked.is_ok)
        self.assertIn(ISSUE_SHARED_RFP, {item.code for item in checked.issues})

    def test_shared_rfp_number_without_closed_files_is_error(self) -> None:
        different_files = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="15",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС15",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС15",
                        rfp_file="A.xlsx",
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="61",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС61",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС61",
                        rfp_file="B.xlsx",
                        block_index=1,
                    ),
                ),
            ),
        ]
        different = validate_registry_rows(different_files)
        self.assertFalse(different.is_ok)
        self.assertIn(
            ISSUE_SHARED_RFP, {item.code for item in different.issues}
        )

        one_empty = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="15",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС15",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС15",
                        rfp_file="A.xlsx",
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="61",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС61",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС61",
                        rfp_file="",
                        block_index=1,
                    ),
                ),
            ),
        ]
        empty_side = validate_registry_rows(one_empty)
        self.assertFalse(empty_side.is_ok)
        self.assertIn(
            ISSUE_SHARED_RFP, {item.code for item in empty_side.issues}
        )

    def test_history_same_rfp_file_is_not_shared(self) -> None:
        shared_file = "shared-rfp.xlsx"

        def pair(second_status: str, second_id: str) -> list[DsRegistryRow]:
            return [
                DsRegistryRow(
                    status=STATUS_ACTIVE,
                    source_id="15",
                    excel_row=2,
                    relations=(
                        DsRegistryRelation(
                            group_id="ДС15",
                            rfp_key="15",
                            ul_folder="согл УЛ ДС15",
                            rfp_file=shared_file,
                            block_index=1,
                        ),
                    ),
                ),
                DsRegistryRow(
                    status=second_status,
                    source_id=second_id,
                    excel_row=3,
                    relations=(
                        DsRegistryRelation(
                            group_id=f"ДС{second_id}",
                            rfp_key=second_id,
                            ul_folder=f"согл УЛ ДС{second_id}",
                            rfp_file=shared_file,
                            block_index=1,
                        ),
                    ),
                ),
            ]

        for status, source_id in (
            (STATUS_HISTORY, "61"),
            (STATUS_DISABLED, "70"),
        ):
            rows = pair(status, source_id)
            checked = validate_registry_rows(rows)
            self.assertTrue(checked.is_ok, checked.issues)
            self.assertNotIn(
                ISSUE_SHARED_RFP, {item.code for item in checked.issues}
            )
            self.assertEqual(legal_rfp_file_clusters(rows), ())

    def test_duplicate_rfp_file_inside_one_source_is_error(self) -> None:
        rows = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="15",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС15",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС15",
                        rfp_file="A.xlsx\nA.xlsx",
                        block_index=1,
                    ),
                ),
            ),
        ]
        checked = validate_registry_rows(rows)
        codes = {item.code for item in checked.issues}
        self.assertIn(ISSUE_DUPLICATE_LINK, codes)
        self.assertNotIn(ISSUE_SHARED_RFP, codes)
        self.assertFalse(checked.is_ok)

    def test_packing_maps_omit_cluster_shared_rfp_number(self) -> None:
        shared_file = "ДС15_61. cluster.xlsx"
        rows = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="15",
                relations=(
                    DsRegistryRelation(
                        group_id="ДС15",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС15",
                        rfp_file=shared_file,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="61",
                relations=(
                    DsRegistryRelation(
                        group_id="ДС61",
                        rfp_key="15",
                        ul_folder="согл УЛ ДС61",
                        rfp_file=shared_file,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="99",
                relations=(
                    DsRegistryRelation(
                        group_id="ДС99",
                        rfp_key="99",
                        ul_folder="согл УЛ ДС99",
                        rfp_file="ДС99.xlsx",
                        block_index=1,
                    ),
                ),
            ),
        ]
        checked = validate_registry_rows(rows)
        self.assertTrue(checked.is_ok, checked.issues)
        document = DsRegistryDocument(
            path=Path("cluster-pack.xlsx"),
            rows=rows,
            validation=checked,
            max_relation_blocks=1,
        )
        maps = registry_packing_maps(document)
        self.assertIsNotNone(maps)
        assert maps is not None
        folders, numbers = maps
        self.assertNotIn(15, numbers)
        self.assertEqual(numbers[99], 99)
        self.assertEqual(folders["согл ул дс15"], 15)
        self.assertEqual(folders["согл ул дс61"], 61)
        self.assertEqual(folders["согл ул дс99"], 99)

    def test_shared_ul_folder_is_valid_and_indexed(self) -> None:
        folder = "согл УЛ общая"
        rows = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="15",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС15",
                        rfp_key="15",
                        ul_folder=folder,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="61",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС61",
                        rfp_key="61",
                        ul_folder=folder,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="99",
                excel_row=4,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС99",
                        rfp_key="99",
                        ul_folder="",
                        block_index=1,
                    ),
                ),
            ),
        ]
        checked = validate_registry_rows(rows)
        self.assertTrue(checked.is_ok, checked.issues)
        codes = {item.code for item in checked.issues}
        self.assertNotIn(ISSUE_SHARED_UL, codes)
        document = DsRegistryDocument(
            path=Path("shared-folder.xlsx"),
            rows=rows,
            validation=checked,
            max_relation_blocks=1,
        )
        from RFQ.packing_list_provider import (
            build_registry_planting_index,
            normalize_packing_key_part,
            registry_packing_scope,
            supply_ul_folders,
        )

        index = build_registry_planting_index(document)
        folder_key = normalize_packing_key_part(folder)
        self.assertEqual(index.folder_owners[folder_key], frozenset({15, 61}))
        self.assertIn(folder_key, index.ds_to_folders[15])
        self.assertIn(folder_key, index.ds_to_folders[61])
        self.assertIn(99, index.known_source_ids)
        self.assertNotIn(99, index.ds_to_folders)
        with registry_packing_scope(index):
            self.assertEqual(supply_ul_folders(15), frozenset({folder_key}))
            self.assertEqual(supply_ul_folders(61), frozenset({folder_key}))
            self.assertEqual(supply_ul_folders(99), frozenset())
            self.assertIsNone(supply_ul_folders(12345))
            self.assertIsNone(supply_ul_folders(None))

    def test_duplicate_ul_folder_inside_one_number_is_error(self) -> None:
        folder = "согл УЛ общая"
        rows = [
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="15",
                excel_row=2,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС15",
                        rfp_key="15",
                        ul_folder=folder,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="61",
                excel_row=3,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС61",
                        rfp_key="61",
                        ul_folder=folder,
                        block_index=1,
                    ),
                ),
            ),
            DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="61",
                excel_row=4,
                relations=(
                    DsRegistryRelation(
                        group_id="ДС61",
                        rfp_key="62",
                        ul_folder=folder,
                        block_index=1,
                    ),
                ),
            ),
        ]
        checked = validate_registry_rows(rows)
        codes = {item.code for item in checked.issues}
        self.assertIn(ISSUE_DUPLICATE_LINK, codes)
        self.assertNotIn(ISSUE_SHARED_UL, codes)
        self.assertFalse(checked.is_ok)

    def test_copy_registry_snapshot_preserves_name_and_swallows_oserror(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            source = root / "Реестр_ДС_УЛ.xlsx"
            source.write_bytes(b"registry-bytes")
            dest_dir = root / "result"
            dest_dir.mkdir()
            dest = copy_registry_snapshot(source, dest_dir)
            self.assertEqual(dest, dest_dir / source.name)
            assert dest is not None
            self.assertEqual(dest.read_bytes(), b"registry-bytes")

            dated = root / "Реестр_ДС_УЛ_20260923_120000.xlsx"
            dated.write_bytes(b"dated-bytes")
            dated_dest = copy_registry_snapshot(dated, dest_dir)
            self.assertEqual(dated_dest, dest_dir / dated.name)
            assert dated_dest is not None
            self.assertEqual(dated_dest.read_bytes(), b"dated-bytes")

            not_a_dir = root / "not_a_dir.txt"
            not_a_dir.write_bytes(b"x")
            self.assertIsNone(copy_registry_snapshot(source, not_a_dir))


class RegistryRescanSmokeTest(unittest.TestCase):
    def test_check_drops_yellow_rows_and_rewrites_ds_file(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            root = Path(raw)
            ds = root / "ds"
            (ds / "ДС_13").mkdir(parents=True)
            (ds / "ДС_75_Приложение 1.xlsx").write_bytes(b"a")
            (ds / "ДС_96_1._Спецификация №75.xlsx").write_bytes(b"b")
            (ds / "ДС_13" / "10_Спецификация.xlsx").write_bytes(b"c")
            (ds / "заметка.xlsx").write_bytes(b"x")
            (ds / "Реестр_ДС_УЛ.xlsx").write_bytes(b"reg")
            (ds / "СВОД_ДС_ручной.xlsx").write_bytes(b"svod")
            path = root / "registry.xlsx"
            stale = (
                "ДС_75_Приложение 1.xlsx\n"
                "ДС_96_1._Спецификация №75.xlsx"
            )
            write_registry_workbook(
                path,
                [
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="75",
                        ds_file=stale,
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС75",
                                rfp_key="16",
                                ul_folder="согл УЛ ДС16",
                                block_index=1,
                            ),
                        ),
                    ),
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="75",
                        ds_file="другой.xlsx",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС75",
                                rfp_key="",
                                ul_folder="согл УЛ ДС75б",
                                block_index=1,
                            ),
                        ),
                    ),
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="96",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС96",
                                rfp_key="70",
                                ul_folder="согл УЛ ДС70",
                                block_index=1,
                            ),
                        ),
                    ),
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="13",
                        ds_file="чужой.xlsx",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС13",
                                rfp_key="13",
                                ul_folder="согл УЛ ДС13",
                                block_index=1,
                            ),
                        ),
                    ),
                    DsRegistryRow(
                        status="",
                        source_id="",
                        ds_file="ДС_75_Приложение 1.xlsx",
                        note="сирота: файл ДС не записан в строке номера ДС75",
                    ),
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="",
                        note="черновик человека",
                        relations=(
                            DsRegistryRelation(
                                group_id="",
                                rfp_key="1",
                                ul_folder="папка человека",
                                block_index=1,
                            ),
                        ),
                    ),
                ],
            )
            split_rows = [
                DsRegistryRow(status=STATUS_ACTIVE, source_id="75", ds_file="a.xlsx"),
                DsRegistryRow(status=STATUS_ACTIVE, source_id="75", ds_file="b.xlsx"),
            ]
            _merged, split_issues = collapse_registry_rows(
                split_rows, compare_ds_file=False
            )
            self.assertEqual(split_issues, [])

            laps: list[str] = []
            doc = load_registry(
                path,
                ds_root=ds,
                on_lap=lambda label, _seconds: laps.append(label),
            )
            by_id = {row.source_id: row for row in doc.rows if row.source_id}
            self.assertEqual(laps.count("имена файлов ДС"), 1)
            self.assertIn("чтение листа", laps)
            self.assertIn("сироты ДС", laps)
            self.assertEqual(by_id["75"].ds_file, "ДС_75_Приложение 1.xlsx")
            self.assertEqual(
                by_id["96"].ds_file, "ДС_96_1._Спецификация №75.xlsx"
            )
            self.assertEqual(by_id["13"].ds_file, "10_Спецификация.xlsx")
            notes = [row.note for row in doc.rows if not row.source_id]
            self.assertIn("черновик человека", notes)
            self.assertEqual(
                [note for note in notes if note.startswith("сирота:")],
                ["сирота: файл ДС не назван ни на одной строке"],
            )
            self.assertEqual(
                [row.ds_file for row in doc.rows if not row.source_id and row.ds_file],
                ["заметка.xlsx"],
            )
            self.assertNotIn(
                ISSUE_IDENTITY_SPLIT,
                {item.code for item in doc.validation.issues},
            )

            again_path = root / "again.xlsx"
            write_registry_workbook(again_path, doc.rows)
            again = load_registry(again_path, ds_root=ds)
            again_notes = [
                row.note
                for row in again.rows
                if not row.source_id and row.note.startswith("сирота:")
            ]
            self.assertEqual(again_notes, notes[-1:])


class RegistryExcelFilterSmokeTest(unittest.TestCase):
    def test_space_custom_filter_does_not_block_read(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "registry.xlsx"
            write_registry_workbook(
                path,
                [
                    DsRegistryRow(
                        status=STATUS_ACTIVE,
                        source_id="11",
                        relations=(
                            DsRegistryRelation(
                                group_id="ДС11",
                                rfp_key="11",
                                ul_folder="согл УЛ ДС11",
                                block_index=1,
                            ),
                        ),
                    )
                ],
            )
            source = ZipFile(BytesIO(path.read_bytes()))
            out = BytesIO()
            with ZipFile(out, "w") as target:
                for info in source.infolist():
                    blob = source.read(info.filename)
                    if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                        text = blob.decode("utf-8")
                        text2, count = re.subn(
                            r"<autoFilter ref=\"A1:E\d+\"/>",
                            (
                                '<autoFilter ref="A1:E3"><filterColumn colId="0">'
                                '<customFilters><customFilter operator="notEqual" val=" "/>'
                                "</customFilters></filterColumn></autoFilter>"
                            ),
                            text,
                            count=1,
                        )
                        if count:
                            blob = text2.encode("utf-8")
                    target.writestr(info.filename, blob)
            path.write_bytes(out.getvalue())
            self.assertEqual(detect_registry_format(path), "new")
            self.assertEqual(load_registry(path).rows[0].source_id, "11")


class RegistrySheetLayoutSmokeTest(unittest.TestCase):
    def test_keeps_widths_and_fits_multiline_height(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as raw:
            path = Path(raw) / "registry.xlsx"
            row = DsRegistryRow(
                status=STATUS_ACTIVE,
                source_id="13",
                ds_file="один.xlsx\nдва.xlsx\nтри.xlsx",
                relations=(
                    DsRegistryRelation(
                        group_id="ДС13",
                        rfp_key="13",
                        ul_folder="согл УЛ ДС13",
                        block_index=1,
                    ),
                ),
            )
            write_registry_workbook(path, [row])
            book = _load_xlsx(path)
            sheet = book["Реестр ДС"]
            self.assertAlmostEqual(sheet.column_dimensions["D"].width, 93.57, places=2)
            self.assertGreaterEqual(sheet.row_dimensions[1].height, 45)
            self.assertEqual(sheet["A1"].alignment.horizontal, "center")
            self.assertTrue(sheet["A1"].alignment.wrap_text)
            self.assertEqual(sheet["D2"].alignment.horizontal, "left")
            self.assertEqual(sheet["D2"].alignment.vertical, "center")
            self.assertTrue(sheet["D2"].alignment.wrap_text)
            self.assertGreaterEqual(sheet.row_dimensions[2].height, 45)
            sheet.column_dimensions["D"].width = 40
            book.save(path)
            book.close()
            write_registry_workbook(path, [row])
            again = _load_xlsx(path)
            self.assertEqual(again["Реестр ДС"].column_dimensions["D"].width, 40)
            again.close()


if __name__ == "__main__":
    unittest.main()
