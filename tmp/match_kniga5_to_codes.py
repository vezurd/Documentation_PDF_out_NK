"""Match Книга5 positions to purchase codes via Приложение 1.1 + Google СПО_1.

Appendix columns: Код 1 / Код 2 (BCC), Код 1С (two cols), партномер, описание.
Google is loaded with ``base.base_google.load_base`` like the RFP pipeline.

Run from repo root:

  set PYTHONUTF8=1
  python tmp/match_kniga5_to_codes.py
"""

from __future__ import annotations

import os
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
sys.path.insert(0, str(ROOT / "tmp"))

import match_positions_to_google as mpg  # noqa: E402

from base.base_google import load_base  # noqa: E402
from base.base_utils import check_code  # noqa: E402
from base.tables_columns import CODE, NAME, TYPE_MARK, VENDOR  # noqa: E402

TEMP_DIR = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP\_temp"
)
INPUT_XLSX = TEMP_DIR / "Книга5.xlsx"
APPENDIX_XLSX = TEMP_DIR / "Приложение 1.1.xlsx"
OUTPUT_NAME = "Книга5_подбор_BCC.xlsx"

SRC_HEADERS = [
    "Группа",
    "№",
    "Наименование",
    "Производитель",
    "Код производителя",
    "КОД 1С ВСС",
    "Количество",
    "Правильный код BCC",
]
MATCH_HEADERS = [
    "Предполагаемый код BCC",
    "Оценка качества",
    "Балл 0-100",
    "Источник",
    "Наименование каталога",
    "Тип / марка / партномер",
    "Вендор",
    "Код 2 приложения",
    "Код 1С приложения",
    "Строка приложения",
    "Причина подбора",
    "Альт.2 код",
    "Альт.2 источник",
    "Альт.2 балл",
    "Альт.2 наименование",
    "Альт.3 код",
    "Альт.3 источник",
    "Альт.3 балл",
    "Альт.3 наименование",
]

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
FILL_HEADER = PatternFill("solid", fgColor="305496")
FILL_MATCH_HEADER = PatternFill("solid", fgColor="548235")

_BCC_RE = re.compile(r"BCC\d{7}", re.I)
_CODE_1C_RE = re.compile(r"\d{8}")


@dataclass
class CatalogItem:
    code: str
    name: str
    type_mark: str
    vendor: str
    source: str
    code2: str = ""
    codes_1c: tuple[str, ...] = ()
    articles: tuple[str, ...] = ()
    appendix_row: int = 0
    name_n: str = ""
    mark_n: str = ""
    blob_n: str = ""
    compact_mark: str = ""
    compact_blob: str = ""
    compact_name: str = ""
    tokens: set[str] = field(default_factory=set)
    specs: set[str] = field(default_factory=set)
    models: set[str] = field(default_factory=set)
    is_kit: bool = False
    article_compacts: tuple[str, ...] = ()


def _is_purchase_code(value: str) -> bool:
    text = (value or "").strip()
    if not text or "/" in text:
        return False
    if _BCC_RE.fullmatch(text):
        return True
    return check_code(text) == 1 and not text.isdigit()


def _split_articles(raw: object) -> list[str]:
    parts: list[str] = []
    for piece in re.split(r"[,;+\n]", mpg._cell_str(raw)):
        piece = re.sub(r"\s*-\s*\d+\s*$", "", piece).strip()
        if piece:
            parts.append(piece)
    return parts


def _extract_1c(*values: object) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        for part in re.split(r"[\n;,]+", mpg._cell_str(value)):
            part = part.strip()
            if _CODE_1C_RE.fullmatch(part) and part not in seen:
                seen.add(part)
                out.append(part)
    return out


def _finish_item(item: CatalogItem) -> CatalogItem:
    blob = f"{item.name} {item.type_mark} {item.vendor} {' '.join(item.articles)}"
    item.name_n = mpg._norm(item.name)
    item.mark_n = mpg._norm(item.type_mark)
    item.blob_n = mpg._norm(blob)
    item.compact_mark = mpg._compact(item.type_mark)
    item.compact_blob = mpg._compact(blob)
    item.compact_name = mpg._compact(item.name)
    item.tokens = mpg._tokens(blob)
    item.specs = mpg._specs(blob)
    item.models = mpg._models(blob)
    item.article_compacts = tuple(
        mpg._compact(a) for a in item.articles if mpg._compact(a)
    )
    item.is_kit = (
        "в составе" in item.name_n
        or len(item.name_n) > 280
        or len(item.articles) >= 3
    )
    return item


def _load_appendix(path: Path) -> list[CatalogItem]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    items: list[CatalogItem] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= 5:
            continue
        vals = list(row) + [None] * 11
        code1 = mpg._cell_str(vals[1])
        code2 = mpg._cell_str(vals[2])
        name = mpg._cell_str(vals[6]) or mpg._cell_str(vals[7])
        articles = _split_articles(vals[8])
        vendor = mpg._cell_str(vals[9])
        codes_1c = _extract_1c(vals[3], vals[4])
        primary = code1 if _is_purchase_code(code1) else ""
        alt = code2 if _is_purchase_code(code2) else ""
        if not primary and not alt and not articles and not codes_1c:
            if not name or name.isupper() and len(name) < 80:
                continue
        item = CatalogItem(
            code=primary,
            name=name,
            type_mark=", ".join(articles),
            vendor=vendor,
            source="приложение 1.1",
            code2=alt,
            codes_1c=tuple(codes_1c),
            articles=tuple(articles),
            appendix_row=i,
        )
        items.append(_finish_item(item))
    wb.close()
    return items


def _load_google_catalog() -> list[CatalogItem]:
    rows = load_base() or []
    items: list[CatalogItem] = []
    for row in rows:
        code = mpg._cell_str(row.el[CODE].value)
        if not code:
            continue
        name = mpg._cell_str(row.el[NAME].value)
        mark = mpg._cell_str(row.el[TYPE_MARK].value)
        vendor = mpg._cell_str(row.el.get(VENDOR).value if row.el.get(VENDOR) else "")
        articles = _split_articles(mark)
        item = CatalogItem(
            code=code,
            name=name,
            type_mark=mark,
            vendor=vendor,
            source="Google СПО_1",
            articles=tuple(articles) if articles else tuple(filter(None, [mark])),
        )
        items.append(_finish_item(item))
    return items


def _load_kniga(path: Path) -> list[mpg.ExcelRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out: list[mpg.ExcelRow] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        values = list(row[:8])
        while len(values) < 8:
            values.append(None)
        name = mpg._cell_str(values[2])
        if not name:
            continue
        vendor = mpg._cell_str(values[3])
        mfr = mpg._cell_str(values[4])
        blob = f"{name} {vendor} {mfr}"
        out.append(
            mpg.ExcelRow(
                excel_row=i,
                values=values,
                name=name,
                vendor=vendor,
                mfr=mfr,
                code_1c=mpg._cell_str(values[5]),
                name_n=mpg._norm(name),
                blob_n=mpg._norm(blob),
                compact_mfr=mpg._compact(mfr),
                compact_name=mpg._compact(name),
                tokens=mpg._tokens(blob),
                specs=mpg._specs(blob),
                models=mpg._models(blob),
            )
        )
    wb.close()
    return out


def _article_exact(src: mpg.ExcelRow, item: CatalogItem) -> bool:
    mfr = src.compact_mfr
    if not mfr or len(mfr) < 4:
        return False
    if mfr in item.article_compacts or mfr == item.compact_mark:
        return True
    for art in item.articles:
        last = art.replace("«", "").replace("»", "").strip().split()
        if last and len(mfr) >= 5 and mfr == mpg._compact(last[-1]):
            return True
    return False


def _mark_hit(src: mpg.ExcelRow, item: CatalogItem) -> float:
    """Exact article/SKU only; compact-blob substring needs a long distinctive code."""
    mfr = src.compact_mfr
    if _article_exact(src, item):
        return 1.0
    if mfr and len(mfr) >= 10 and mfr in item.compact_blob:
        return 0.9
    if src.compact_name and len(src.compact_name) >= 10 and item.compact_mark:
        mark = item.compact_mark
        if len(mark) >= 10 and mark == src.compact_name:
            return 0.95
        if len(mark) >= 10 and mark in src.compact_name:
            if len(mark) >= 0.35 * len(src.compact_name):
                return 0.9
    return 0.0


def _source_tag(item: CatalogItem) -> str:
    return "приложение 1.1" if item.source.startswith("приложение") else "Google"


def _full_score(src: mpg.ExcelRow, item: CatalogItem) -> mpg.Candidate:
    orig = mpg._mark_hit
    mpg._mark_hit = _mark_hit
    try:
        cand = mpg._full_score(src, item)
    finally:
        mpg._mark_hit = orig
    tag = _source_tag(item)
    cand.reasons = [
        r.replace("Google", tag).replace("тип/марка Google", "партномер/марка")
        for r in cand.reasons
    ]
    if src.code_1c and src.code_1c in item.codes_1c:
        cand.score = max(cand.score, 0.97 if item.code else 0.72)
        cand.reasons = [f"совпал КОД 1С с {tag}"] + cand.reasons
        cand.kit_only = False
    mfr = src.compact_mfr
    if _article_exact(src, item):
        if item.is_kit or len(item.articles) >= 2:
            cand.kit_only = True
            if "комплект" not in " ".join(cand.reasons):
                cand.reasons.append("артикул входит в набор приложения 1.1")
        elif item.code and not cand.kit_only:
            cand.score = max(cand.score, 0.93)
            cand.reasons = [f"партномер = код производителя ({tag})"] + cand.reasons
    if (
        mfr
        and len(mfr) >= 10
        and mfr in item.compact_blob
        and mfr != item.compact_mark
        and (item.is_kit or len(item.name_n) > max(80, len(src.name_n) * 1.5))
    ):
        cand.kit_only = True
        if "комплект" not in " ".join(cand.reasons):
            cand.reasons.append("код найден в составной позиции каталога")
    cand.score = max(0.0, min(1.0, cand.score))
    return cand


def _cheap_score(src: mpg.ExcelRow, item: CatalogItem) -> float:
    orig = mpg._mark_hit
    mpg._mark_hit = _mark_hit
    try:
        base = mpg._cheap_score(src, item)
    finally:
        mpg._mark_hit = orig
    if src.code_1c and src.code_1c in item.codes_1c:
        base += 5.0
    if src.compact_mfr and _article_exact(src, item):
        base += 3.5
    return base


def _forced(src: mpg.ExcelRow, items: list[CatalogItem]) -> list[CatalogItem]:
    out: list[CatalogItem] = []
    dims = {s for s in src.specs if "x" in s or s.startswith("din")}
    for item in items:
        if src.code_1c and src.code_1c in item.codes_1c:
            out.append(item)
            continue
        if src.compact_mfr and _article_exact(src, item):
            out.append(item)
            continue
        if _mark_hit(src, item) >= 0.88:
            out.append(item)
            continue
        if dims and dims <= item.specs:
            out.append(item)
    return out


def _fill_for_quality(label: str) -> PatternFill:
    if "без BCC" in label:
        return mpg.FILL_KIT
    return mpg._fill_for_quality(label)


def _quality(top: mpg.Candidate | None, gap: float) -> tuple[str, int]:
    if top is None:
        return "не найден", 0
    item: CatalogItem = top.item
    points = int(round(top.score * 100))
    reasons = " ".join(top.reasons)
    if top.kit_only:
        return "только в комплекте", points
    if not item.code and top.score >= 0.70:
        return "в приложении без BCC", points
    if top.score < 0.36 and "КОД 1С" not in reasons and "партномер" not in reasons:
        return "не найден", points
    if "КОД 1С" in reasons and item.code:
        label = "точный (1С)"
    elif top.score >= 0.88 and (
        "партномер" in reasons or "код производителя" in reasons
    ):
        label = "точный"
    elif top.score >= 0.78:
        label = "высокий"
    elif top.score >= 0.56:
        label = "средний"
    else:
        label = "слабый"
    if gap < 0.045 and top.score >= 0.50 and "точный" not in label:
        label = "спорный (" + label + ")"
    return label, points


def match_rows(
    excel_rows: list[mpg.ExcelRow],
    catalog: list[CatalogItem],
) -> list[dict]:
    results: list[dict] = []
    for src in excel_rows:
        ranked = sorted(catalog, key=lambda g: _cheap_score(src, g), reverse=True)[:56]
        seen = {id(g) for g in ranked}
        for item in _forced(src, catalog):
            if id(item) not in seen:
                ranked.append(item)
                seen.add(id(item))
        cands = sorted(
            (_full_score(src, item) for item in ranked),
            key=lambda c: (c.score, 1 if c.item.code else 0),
            reverse=True,
        )
        if (
            len(cands) >= 2
            and not str(cands[0].item.code).upper().startswith("BCC")
            and str(cands[1].item.code).upper().startswith("BCC")
            and (cands[0].score - cands[1].score) <= 0.08
        ):
            cands[0], cands[1] = cands[1], cands[0]
        # If best row has no BCC but a close Google/appendix row does — keep both.
        top = cands[0] if cands else None
        if (
            top
            and not top.item.code
            and not top.kit_only
        ):
            for alt in cands[1:6]:
                if alt.item.code and alt.score >= 0.50:
                    top = alt
                    top.reasons = ["взят код из близкого кандидата с BCC"] + top.reasons
                    break
        second = next((c for c in cands if c is not top), None)
        third = next((c for c in cands if c is not top and c is not second), None)
        gap = (top.score - second.score) if top and second else 1.0
        label, points = _quality(top, gap)
        if label == "не найден":
            code = ""
            reason = "близкого кода нет"
            if top:
                reason += "; лучший слабый: " + "; ".join(top.reasons[:2])
        else:
            code = top.item.code if top else ""
            reason = "; ".join(top.reasons[:5]) if top else ""
        results.append(
            {
                "src": src,
                "code": code,
                "label": label,
                "points": points,
                "source": top.item.source if top else "",
                "g_name": top.item.name if top else "",
                "g_mark": top.item.type_mark if top else "",
                "g_vendor": top.item.vendor if top else "",
                "code2": top.item.code2 if top else "",
                "codes_1c": ", ".join(top.item.codes_1c) if top else "",
                "app_row": (top.item.appendix_row or "") if top else "",
                "reason": reason,
                "alt2_code": second.item.code if second else "",
                "alt2_src": second.item.source if second else "",
                "alt2_points": int(round(second.score * 100)) if second else "",
                "alt2_name": second.item.name if second else "",
                "alt3_code": third.item.code if third else "",
                "alt3_src": third.item.source if third else "",
                "alt3_points": int(round(third.score * 100)) if third else "",
                "alt3_name": third.item.name if third else "",
                "fill": _fill_for_quality(label),
            }
        )
    return results


def _write_xlsx(results: list[dict], out_path: Path, n_app: int, n_google: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Подбор BCC"
    headers = SRC_HEADERS + MATCH_HEADERS
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILL_MATCH_HEADER if col > len(SRC_HEADERS) else FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"
    for r_idx, rec in enumerate(results, 2):
        src: mpg.ExcelRow = rec["src"]
        for c_idx, val in enumerate(src.values[:8], 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.border = THIN
            cell.alignment = WRAP
        extra = [
            rec["code"],
            rec["label"],
            rec["points"],
            rec["source"],
            rec["g_name"],
            rec["g_mark"],
            rec["g_vendor"],
            rec["code2"],
            rec["codes_1c"],
            rec["app_row"],
            rec["reason"],
            rec["alt2_code"],
            rec["alt2_src"],
            rec["alt2_points"],
            rec["alt2_name"],
            rec["alt3_code"],
            rec["alt3_src"],
            rec["alt3_points"],
            rec["alt3_name"],
        ]
        for offset, val in enumerate(extra):
            cell = ws.cell(r_idx, len(SRC_HEADERS) + 1 + offset, val)
            cell.border = THIN
            cell.alignment = WRAP
            cell.fill = rec["fill"]
        ws.row_dimensions[r_idx].height = 36
    widths = {
        1: 10,
        2: 8,
        3: 42,
        4: 12,
        5: 22,
        6: 14,
        7: 12,
        8: 18,
        9: 18,
        10: 20,
        11: 12,
        12: 16,
        13: 42,
        14: 28,
        15: 14,
        16: 16,
        17: 16,
        18: 12,
        19: 36,
        20: 14,
        21: 16,
        22: 10,
        23: 36,
    }
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 14)

    summary = wb.create_sheet("Сводка")
    counts = Counter(rec["label"] for rec in results)
    src_counts = Counter(rec["source"] for rec in results if rec["code"])
    summary["A1"] = "Подбор BCC для Книга5: приложение 1.1 + Google СПО_1"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Книга5: {INPUT_XLSX}"
    summary["A3"] = f"Приложение 1.1: {APPENDIX_XLSX} ({n_app} строк каталога)"
    summary["A4"] = f"Google СПО_1: {n_google} кодов"
    summary["A5"] = f"Строк Книга5: {len(results)}"
    summary["A6"] = f"Сформировано: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    summary["A8"] = "Оценка"
    summary["B8"] = "Кол-во"
    summary["A8"].font = Font(bold=True)
    row_i = 9
    for label, n in sorted(counts.items(), key=lambda x: -x[1]):
        summary.cell(row_i, 1, label).fill = _fill_for_quality(label)
        summary.cell(row_i, 2, n)
        row_i += 1
    row_i += 1
    summary.cell(row_i, 1, "Источник предложенного кода")
    summary.cell(row_i, 1).font = Font(bold=True)
    row_i += 1
    for src, n in sorted(src_counts.items(), key=lambda x: -x[1]):
        summary.cell(row_i, 1, src or "(пусто)")
        summary.cell(row_i, 2, n)
        row_i += 1
    row_i += 1
    summary.cell(
        row_i,
        1,
        "Приложение 1.1: Код 1 и Код 2 — закупочные (BCC…); два столбца «Код 1С»; "
        "партномер может содержать список артикулов комплекта. "
        "Столбец «Правильный код BCC» не заполнялся — туда можно внести ручную правку. "
        "Предложение робота — в «Предполагаемый код BCC».",
    )
    summary.merge_cells(
        start_row=row_i, start_column=1, end_row=row_i + 3, end_column=6
    )
    summary.cell(row_i, 1).alignment = WRAP
    summary.column_dimensions["A"].width = 44
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> int:
    if not INPUT_XLSX.is_file():
        print(f"INPUT missing: {INPUT_XLSX}", flush=True)
        return 1
    if not APPENDIX_XLSX.is_file():
        print(f"APPENDIX missing: {APPENDIX_XLSX}", flush=True)
        return 1
    print("loading appendix 1.1…", flush=True)
    appendix = _load_appendix(APPENDIX_XLSX)
    print(f"appendix items: {len(appendix)}", flush=True)
    print("loading google code base…", flush=True)
    google = _load_google_catalog()
    print(f"google items: {len(google)}", flush=True)
    catalog = appendix + google
    excel_rows = _load_kniga(INPUT_XLSX)
    print(f"книга5 rows: {len(excel_rows)}", flush=True)
    print("matching…", flush=True)
    results = match_rows(excel_rows, catalog)
    counts = Counter(rec["label"] for rec in results)
    for label, n in sorted(counts.items(), key=lambda x: -x[1]):
        print(f"  {label}: {n}", flush=True)
    unc_out = TEMP_DIR / OUTPUT_NAME
    local_out = ROOT / "tmp" / OUTPUT_NAME
    saved = 0
    for path in (unc_out, local_out):
        try:
            _write_xlsx(results, path, len(appendix), len(google))
            print(f"saved: {path}", flush=True)
            saved += 1
        except OSError as exc:
            print(f"WARN cannot write {path}: {exc}", flush=True)
    return 0 if saved else 2


if __name__ == "__main__":
    raise SystemExit(main())
