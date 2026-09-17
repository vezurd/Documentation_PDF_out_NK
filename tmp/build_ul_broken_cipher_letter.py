"""Build letter-ready list of UL files with broken title/mark ciphers (*MT / SKUDM)."""

from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

STEP4 = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.08.25.11.56"
    r"\Шаг4_Сопоставление_RFP_MTO_20260825_120001.xlsx"
)
SUMMARY = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_RFP\УЛ сводный файл\tsd_packing_summary_2026.08.21_15.09.22.xlsx"
)
TSD_ROOT = Path(
    r"\\bcc\root\CurProjects\di_manegers\Пятницкий_П\Амурский ГХК\Поставки\ТСД по всем ДС"
)
OUT = ROOT / "tmp" / "ul_broken_title_cipher_letter.txt"

# Lost dots: KSB.MTO → KSBMT, SOS.MTO → SOSMT, SKUD.M → SKUDM, SOT.MTO → SOTMT
BAD_SYSTEM_RE = re.compile(r"^(?:KSBMT|SOSMT|SOTMT|SKUDM)$", re.IGNORECASE)
BAD_DISPLAY_RE = re.compile(
    r"^\d+-(?:KSBMT|SOSMT|SOTMT|SKUDM)$",
    re.IGNORECASE,
)


def _ranges(nums: list[int]) -> str:
    if not nums:
        return ""
    nums = sorted(set(nums))
    chunks: list[str] = []
    start = prev = nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
            continue
        chunks.append(f"{start}" if start == prev else f"{start}–{prev}")
        start = prev = n
    chunks.append(f"{start}" if start == prev else f"{start}–{prev}")
    return ", ".join(chunks)


def _parse_source(raw: object) -> list[tuple[str, str, int]]:
    """Parse Step4 «Источник УЛ» into (rel_path, sheet, excel_row)."""
    text = str(raw or "").strip()
    if not text:
        return []
    lines = [ln for ln in text.splitlines() if ln.strip()]
    out: list[tuple[str, str, int]] = []
    current_file = ""
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("…"):
            continue
        indented = line.startswith("  ") or line.startswith("\t")
        parts = [p.strip() for p in stripped.split(" · ")]
        if indented and current_file:
            sheet = parts[0] if parts else ""
            row_n = 0
            if len(parts) >= 2 and parts[1].startswith("строка "):
                try:
                    row_n = int(parts[1].replace("строка ", "").strip())
                except ValueError:
                    row_n = 0
            if row_n:
                out.append((current_file, sheet, row_n))
            continue
        if len(parts) >= 3 and parts[-1].startswith("строка "):
            try:
                row_n = int(parts[-1].replace("строка ", "").strip())
            except ValueError:
                row_n = 0
            current_file = parts[0]
            if row_n:
                out.append((current_file, parts[1], row_n))
        else:
            current_file = parts[0]
    return out


def _full_path(rel: str) -> Path:
    if rel.startswith("\\\\") or (len(rel) > 2 and rel[1] == ":"):
        return Path(rel)
    return TSD_ROOT / rel


def main() -> None:
    # file -> sheet -> rows
    by_file: dict[str, dict[str, list[int]]] = defaultdict(lambda: defaultdict(list))
    systems_by_file: dict[str, set[str]] = defaultdict(set)
    specs_by_file: dict[str, set[str]] = defaultdict(set)
    displays_by_file: dict[str, set[str]] = defaultdict(set)

    if not SUMMARY.is_file():
        raise SystemExit(f"summary not found: {SUMMARY}")

    print("reading summary…")
    sw = load_workbook(SUMMARY, read_only=True, data_only=True)
    ssheet = sw[sw.sheetnames[0]]
    sit = ssheet.iter_rows(values_only=True)
    sheader = [str(c or "").strip() for c in next(sit)]
    sidx = {n: i for i, n in enumerate(sheader) if n}

    def scell(row, name):
        i = sidx[name]
        return None if i >= len(row) else row[i]

    summary_rows = 0
    for row in sit:
        system = str(scell(row, "Марка") or "").strip()
        title = str(scell(row, "Титул") or "").strip()
        display = f"{title}-{system}" if title and system else ""
        if not (BAD_SYSTEM_RE.match(system) or BAD_DISPLAY_RE.match(display)):
            continue
        rel = str(scell(row, "Файл") or "").strip()
        sheet = str(scell(row, "Вкладка") or "").strip()
        try:
            excel_row = int(scell(row, "Строка в исходном УЛ"))
        except (TypeError, ValueError):
            continue
        if not rel or excel_row <= 0:
            continue
        summary_rows += 1
        by_file[rel][sheet].append(excel_row)
        systems_by_file[rel].add(system)
        if display:
            displays_by_file[rel].add(display)
        spec = str(scell(row, "Спецификация") or "").strip()
        if spec:
            specs_by_file[rel].add(spec)
    sw.close()
    print(f"summary bad rows: {summary_rows}, files: {len(by_file)}")

    # Cross-check leftover sources from Step4 (may add rows if summary skipped).
    print("reading Step4 leftover sources…")
    wb = load_workbook(STEP4, read_only=True, data_only=True)
    ws = wb["Сопоставление RFP и MTO"]
    it = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(it)]
    idx = {n: i for i, n in enumerate(header) if n}

    def cell(row, name):
        i = idx[name]
        return None if i >= len(row) else row[i]

    leftover_hits = 0
    for row in it:
        if str(cell(row, "Статус УЛ") or "").strip() != "Только в УЛ":
            continue
        title = str(cell(row, "Титул/Марка") or "").strip()
        if not BAD_DISPLAY_RE.match(title):
            continue
        leftover_hits += 1
        for rel, sheet, excel_row in _parse_source(
            cell(row, "Источник УЛ (файл · вкладка · строка)")
        ):
            by_file[rel][sheet].append(excel_row)
            displays_by_file[rel].add(title)
    wb.close()
    print(f"leftover bad rows: {leftover_hits}")

    expected = {
        "KSBMT": "KSB.MTO",
        "SOSMT": "SOS.MTO",
        "SOTMT": "SOT.MTO",
        "SKUDM": "SKUD.M",
    }

    lines: list[str] = []
    lines.append(
        "Список файлов упаковочных листов (УЛ), где в шифре титула/марки "
        "потеряны точки (или пробелы вместо точек)."
    )
    lines.append("")
    lines.append("Пример корректного шифра:")
    lines.append("  AGCC.3438-8350-KSB.MTO-0001")
    lines.append("Как сейчас попадает в робота (без точек между маркой и MTO/M):")
    lines.append("  8350-KSBMT  (ожидалось 8350-KSB.MTO → марка KSB.MTO)")
    lines.append("  8350-SOSMT / 8350-SOTMT / 8350-SKUDM — аналогично")
    lines.append("")
    lines.append("Ожидаемая замена марок:")
    for bad, good in expected.items():
        lines.append(f"  {bad} → {good}")
    lines.append("")
    lines.append(
        "Просьба в указанных файлах/вкладках/строках исправить колонку "
        "«Спецификация» (и при необходимости «Титул»/«Марка»), чтобы робот "
        "читал титул/марку как 8350-KSB.MTO (и аналоги), а не 8350-KSBMT."
    )
    lines.append("")
    lines.append(f"Корень папки ТСД: {TSD_ROOT}")
    lines.append(
        f"Свод робота: {SUMMARY.name} (строк с битой маркой: {summary_rows})"
    )
    lines.append("")

    for n, rel in enumerate(sorted(by_file, key=str.casefold), start=1):
        sheets = by_file[rel]
        full = _full_path(rel)
        lines.append(f"{n}. {full}")
        systems = sorted(systems_by_file.get(rel, set()), key=str.casefold)
        displays = sorted(displays_by_file.get(rel, set()), key=str.casefold)
        specs = sorted(specs_by_file.get(rel, set()), key=str.casefold)
        if systems:
            hint = ", ".join(
                f"{s}→{expected.get(s.upper(), '?')}" for s in systems
            )
            lines.append(f"   Марки сейчас: {hint}")
        if displays:
            lines.append(f"   Титул/Марка в роботе: {', '.join(displays)}")
        if specs:
            show = specs[:4]
            tail = f"; …(+{len(specs) - len(show)})" if len(specs) > len(show) else ""
            lines.append(f"   Примеры «Спецификация»: {'; '.join(show)}{tail}")
        total_rows = 0
        for sheet in sorted(sheets, key=str.casefold):
            rows = sheets[sheet]
            total_rows += len(set(rows))
            lines.append(
                f"   Вкладка «{sheet}»: строки {_ranges(rows)} "
                f"(всего {len(set(rows))} строк)"
            )
        lines.append(f"   Итого по файлу: {total_rows} строк")
        lines.append("")

    lines.append(f"Итого файлов: {len(by_file)}")
    text = "\n".join(lines)
    OUT.write_text(text, encoding="utf-8")
    print(text)
    print("\n--- written", OUT)


if __name__ == "__main__":
    main()
