"""Focused leftover query for packing examples."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.08.27.13.58"
    r"\Шаг4_Сопоставление_RFP_MTO_20260827_140410.xlsx"
)
SRC_COL = "Источник УЛ (файл · вкладка · строка)"

WATCH = {
    ("8445-SOT", "BCC0003052"),
    ("8445-SOT1", "BCC0003052"),
    ("8525-SOT", "BCC0000152"),
}
TAGS = {
    "8445-GA-01-S-FV-0011",
    "8525-SS-01-S-AVI-1001",
    "8525-SS-01-S-AVI-1002",
}


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def main() -> None:
    wb = load_workbook(XLSX, read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {n: i for i, n in enumerate(header) if n}

    def cell(row: list, name: str) -> str:
        i = idx.get(name)
        if i is None or i >= len(row):
            return ""
        v = row[i]
        return "" if v is None else str(v).strip()

    def ul_comment(cells, name: str) -> str:
        i = idx.get(name)
        if i is None or i >= len(cells):
            return ""
        c = cells[i].comment
        return str(c.text).replace("\n", " ") if c else ""

    _safe_print("=== leftovers by code+title or watch tag ===")
    n = 0
    for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row = [c.value for c in cells]
        if cell(row, "Статус УЛ") != "Только в УЛ":
            continue
        title = cell(row, "Титул/Марка")
        code = cell(row, "Код УЛ").upper()
        tags = cell(row, "Теги УЛ")
        if (title, code) not in WATCH and not any(t in tags for t in TAGS):
            continue
        n += 1
        _safe_print(
            f"row={rn} title={title!r} code_ul={code!r} tags_ul={tags!r} "
            f"src={cell(row, SRC_COL)!r}"
        )
        _safe_print(f"  comment={ul_comment(cells, 'Статус УЛ')[:160]!r}")
    _safe_print(f"count={n}")

    _safe_print("\n=== RFP rows title 8445* code BCC0003052 ===")
    for rn, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row = [c.value for c in cells]
        title = cell(row, "Титул/Марка")
        if not title.startswith("8445"):
            continue
        if cell(row, "Код RFP").upper() != "BCC0003052":
            continue
        _safe_print(
            f"row={rn} ds={cell(row, 'Имя ДС')!r} pos={cell(row, '№ позиции')!r} "
            f"title={title!r} status_ul={cell(row, 'Статус УЛ')!r} "
            f"tags_rfp={cell(row, 'Теги RFP')!r}"
        )

    wb.close()


if __name__ == "__main__":
    main()
