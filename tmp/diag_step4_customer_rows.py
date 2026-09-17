"""Read-only dump of customer-flagged Step4 Excel rows and the code replacement table.

Does not modify any file except the UTF-8 report under tmp/.
Run from repo root:

  set PYTHONUTF8=1
  python tmp/diag_step4_customer_rows.py
"""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX

STEP4_XLSX = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO"
    r"\_РЕЗУЛЬТАТА_ПРОВЕРКИ\_результат_проверки_2026.08.25.11.56"
    r"\Шаг4_Сопоставление_RFP_MTO_20260825_120001.xlsx"
)
REPL_XLSX = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\РД\Таблички_графики"
    r"\14_сравнение с ДС\_замены кодов\Code_Replacement_Table.xlsx"
)
REPORT_PATH = Path(__file__).resolve().parent / "diag_step4_customer_rows.md"

WATCH_CODES = (
    "BCC0002561",
    "BCC0000918",
    "BCC0000515",
    "BCC0000805",
    "BCCC0000805",
)

GROUPS: list[tuple[str, frozenset[int]]] = [
    ("п.3 цвет 4109-4142", frozenset(range(4109, 4143))),
    ("п.5 странные точечные", frozenset({5211, 5283, 5503, 7906, 7949, 7950})),
    ("п.6 9533-9539 Нет в МТО и недопоставка", frozenset(range(9533, 9540))),
    ("п.8/10 10845-11103 не совпали с RFP", frozenset(range(10845, 11104))),
    ("п.12 12677 и 12802 одно и то же", frozenset({12677, 12802})),
    ("п.14 15541-15811 не совпали с RFP", frozenset(range(15541, 15812))),
    ("п.15 18326-18391 не совпали с RFP", frozenset(range(18326, 18392))),
    ("п.16 18481-18491", frozenset(range(18481, 18492))),
    ("п.16 18514-18522", frozenset(range(18514, 18523))),
]

WANTED_ROWS = frozenset().union(*(rows for _, rows in GROUPS))

KEY_HEADERS = (
    "Имя ДС",
    "№ позиции",
    "Титул/Марка",
    "Код RFP",
    "Наименование RFP",
    "Поставщик RFP",
    "Кол-во RFP",
    "Ед. изм. RFP",
    "Теги RFP",
    "Тег (VO/MTO/RFP)",
    "MTO, Статус позиции",
    "Тип оборудования",
    "Тег MTO",
    "Код MTO",
    "Наименование MTO",
    "Количество MTO",
    "Ед. изм. MTO",
    "МТО − RFP",
    "Кол-во по УЛ",
    "Ед. изм. УЛ",
    "RFP − УЛ",
    "Статус УЛ",
    "Статус тегов УЛ",
    "Код УЛ",
    "Наименование УЛ",
    "Тег РКД",
    "Код РКД",
)

SAMPLE_LIMIT = 15
CELL_PREVIEW = 80
KSB_SAMPLE_LIMIT = 12
WATCH_HIT_LIMIT = 24


def _safe_print(text: str) -> None:
    enc = getattr(sys.stdout, "encoding", None) or "utf-8"
    sys.stdout.buffer.write((text + "\n").encode(enc, errors="backslashreplace"))


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").replace("|", "/").strip()


def _preview(text: str, limit: int = CELL_PREVIEW) -> str:
    text = _cell_text(text)
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def _fill_rgb(cell) -> str:
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fgColor is None:
        return ""
    color = fill.fgColor
    theme = getattr(color, "theme", None)
    tint = getattr(color, "tint", None)
    rgb = getattr(color, "rgb", None)
    indexed = getattr(color, "indexed", None)
    pattern = getattr(fill, "patternType", None)
    if pattern in (None, "none"):
        return ""
    parts: list[str] = []
    if rgb and str(rgb) not in ("00000000", "0"):
        parts.append(str(rgb).upper())
    if theme is not None:
        parts.append(f"theme{theme}")
        if tint:
            parts.append(f"tint{tint:.3f}")
    if indexed is not None:
        try:
            idx = int(indexed)
            if 0 <= idx < len(COLOR_INDEX):
                parts.append(f"idx{idx}:{COLOR_INDEX[idx]}")
            else:
                parts.append(f"idx{idx}")
        except (TypeError, ValueError):
            parts.append(f"idx{indexed}")
    return " ".join(parts)


def _exists_line(label: str, path: Path) -> str:
    try:
        exists = path.exists()
        size = path.stat().st_size if exists else 0
        return f"- **{label}** exists=`{exists}` size=`{size}` path=`{path}`"
    except OSError as exc:
        return f"- **{label}** exists=`OSError` error=`{exc!r}` path=`{path}`"


def _md_escape(text: str) -> str:
    return _preview(text, 120)


def _write_table(lines: list[str], rows: list[dict[str, str]], cols: list[str]) -> None:
    lines.append("| " + " | ".join(cols) + " |")
    lines.append("| " + " | ".join("---" for _ in cols) + " |")
    for row in rows:
        lines.append("| " + " | ".join(_md_escape(row.get(col, "")) for col in cols) + " |")
    lines.append("")


def _counter_md(counter: Counter[str], limit: int = 12) -> str:
    parts = [f"`{name}`={count}" for name, count in counter.most_common(limit)]
    extra = len(counter) - min(limit, len(counter))
    if extra > 0:
        parts.append(f"… ещё {extra}")
    return ", ".join(parts) if parts else "(нет)"


def _row_from_cells(
    excel_row: int,
    cells: tuple,
    col_of: dict[str, int],
) -> dict[str, str]:
    out: dict[str, str] = {"excel_row": str(excel_row)}
    for name, idx in col_of.items():
        if idx >= len(cells):
            continue
        cell = cells[idx]
        out[name] = _preview(cell.value)
        if name == "Титул/Марка":
            out["_fill_title"] = _fill_rgb(cell)
        elif name == "Код RFP":
            out["_fill_code"] = _fill_rgb(cell)
        elif name == "Код MTO":
            out["_fill_mto"] = _fill_rgb(cell)
        elif name == "Статус УЛ":
            out["_fill_ul"] = _fill_rgb(cell)
    return out


def _summarize(rows: list[dict[str, str]]) -> dict[str, object]:
    titles: Counter[str] = Counter()
    ul_status: Counter[str] = Counter()
    pos_status: Counter[str] = Counter()
    ds_names: Counter[str] = Counter()
    fills_title: Counter[str] = Counter()
    fills_ul: Counter[str] = Counter()
    rfp_empty = 0
    mto_empty = 0
    for row in rows:
        titles[row.get("Титул/Марка") or "(пусто)"] += 1
        ds_names[row.get("Имя ДС") or "(пусто)"] += 1
        ul_status[row.get("Статус УЛ") or "(пусто)"] += 1
        pos_status[row.get("MTO, Статус позиции") or "(пусто)"] += 1
        fills_title[row.get("_fill_title") or "(нет)"] += 1
        fills_ul[row.get("_fill_ul") or "(нет)"] += 1
        if not row.get("Код RFP"):
            rfp_empty += 1
        if not row.get("Код MTO"):
            mto_empty += 1
    return {
        "n": len(rows),
        "titles": titles,
        "ds_names": ds_names,
        "ul_status": ul_status,
        "pos_status": pos_status,
        "rfp_empty": rfp_empty,
        "mto_empty": mto_empty,
        "fills_title": fills_title,
        "fills_ul": fills_ul,
    }


def _sample(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    if len(rows) <= SAMPLE_LIMIT:
        return rows
    mid = len(rows) // 2
    picked = rows[:8] + rows[mid : mid + 4] + rows[-4:]
    seen: set[str] = set()
    out: list[dict[str, str]] = []
    for row in picked:
        key = row["excel_row"]
        if key in seen:
            continue
        seen.add(key)
        out.append(row)
    return out


def _load_replacement_hits(path: Path) -> tuple[list[str], list[tuple[str, str, str]]]:
    notes: list[str] = []
    hits: list[tuple[str, str, str]] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except OSError as exc:
        notes.append(f"не удалось открыть: `{exc!r}`")
        return notes, hits
    if "Коды_замен" not in wb.sheetnames:
        notes.append(f"нет листа `Коды_замен`; листы: {wb.sheetnames}")
        wb.close()
        return notes, hits
    ws = wb["Коды_замен"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    notes.append(f"заголовок: {header}")
    wanted = {code.upper() for code in WATCH_CODES}
    for row in rows:
        if not row:
            continue
        old = _cell_text(row[0] if len(row) > 0 else "").upper()
        new = _cell_text(row[1] if len(row) > 1 else "").upper()
        status = _cell_text(row[2] if len(row) > 2 else "")
        if old in wanted or new in wanted:
            hits.append((old, new, status))
    notes.append(f"всего совпадений по watch-кодам: {len(hits)}")
    wb.close()
    return notes, hits


TABLE_COLS = [
    "excel_row",
    "Имя ДС",
    "Титул/Марка",
    "Код RFP",
    "Код MTO",
    "Кол-во RFP",
    "Количество MTO",
    "Кол-во по УЛ",
    "Статус УЛ",
    "MTO, Статус позиции",
    "Наименование RFP",
    "Наименование MTO",
    "_fill_title",
    "_fill_ul",
]


def main() -> int:
    lines: list[str] = [
        "# Диагностика Step4 — замечания заказчика 2026.08.25",
        "",
        "Readonly. Пайплайн не менялся.",
        "",
        "## Файлы",
        "",
        _exists_line("Step4", STEP4_XLSX),
        _exists_line("Таблица замен", REPL_XLSX),
        "",
    ]
    _safe_print(lines[-4])
    _safe_print(lines[-3])

    try:
        step4_ok = STEP4_XLSX.exists()
    except OSError as exc:
        lines.append(f"Step4 exists() OSError: `{exc!r}`")
        step4_ok = False
    if not step4_ok:
        lines.append("Step4 xlsx недоступен — дальше не читаем.")
        REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
        _safe_print(f"wrote {REPORT_PATH}")
        return 1

    _safe_print("loading Step4 workbook (read_only)…")
    wb = load_workbook(STEP4_XLSX, read_only=True, data_only=False)
    lines.append(f"Листы Step4: `{wb.sheetnames}`")
    ws = wb[wb.sheetnames[0]]
    lines.append(f"Первый лист: `{ws.title}`")

    group_rows: dict[str, list[dict[str, str]]] = {title: [] for title, _ in GROUPS}
    ksb_rows: list[dict[str, str]] = []
    watch_hits: dict[str, list[dict[str, str]]] = {
        "BCC0002561 / BCC0000918": [],
        "BCC0000515 / BCC0000805 / BCCC0000805": [],
    }
    header_names: list[str] = []
    col_of: dict[str, int] = {}
    max_row_seen = 0
    pair_a = {"BCC0002561", "BCC0000918"}
    pair_b = {"BCC0000515", "BCC0000805", "BCCC0000805"}

    for excel_row, cells in enumerate(ws.iter_rows(), start=1):
        max_row_seen = excel_row
        if excel_row == 1:
            header_names = [_cell_text(c.value) for c in cells]
            for idx, name in enumerate(header_names):
                if name and name not in col_of:
                    col_of[name] = idx
            _safe_print(f"headers={len(col_of)}")
            continue

        if excel_row % 5000 == 0:
            _safe_print(f"row {excel_row}…")

        need_full = excel_row in WANTED_ROWS
        title_idx = col_of.get("Титул/Марка")
        rfp_idx = col_of.get("Код RFP")
        mto_idx = col_of.get("Код MTO")
        title_val = _cell_text(cells[title_idx].value) if title_idx is not None and title_idx < len(cells) else ""
        rfp_val = _cell_text(cells[rfp_idx].value).upper() if rfp_idx is not None and rfp_idx < len(cells) else ""
        mto_val = _cell_text(cells[mto_idx].value).upper() if mto_idx is not None and mto_idx < len(cells) else ""
        is_ksb = "9110-KSB" in title_val.upper() and len(ksb_rows) < KSB_SAMPLE_LIMIT
        hit_a = rfp_val in pair_a or mto_val in pair_a
        hit_b = rfp_val in pair_b or mto_val in pair_b
        if hit_a and len(watch_hits["BCC0002561 / BCC0000918"]) >= WATCH_HIT_LIMIT:
            hit_a = False
        if hit_b and len(watch_hits["BCC0000515 / BCC0000805 / BCCC0000805"]) >= WATCH_HIT_LIMIT:
            hit_b = False
        if not (need_full or is_ksb or hit_a or hit_b):
            continue
        parsed = _row_from_cells(excel_row, cells, col_of)
        if need_full:
            for title, rows_set in GROUPS:
                if excel_row in rows_set:
                    group_rows[title].append(parsed)
        if is_ksb:
            ksb_rows.append(parsed)
        if hit_a:
            watch_hits["BCC0002561 / BCC0000918"].append(parsed)
        if hit_b:
            watch_hits["BCC0000515 / BCC0000805 / BCCC0000805"].append(parsed)

    wb.close()
    lines.append(f"Прочитано строк (включая шапку): `{max_row_seen}`")
    lines.append(f"Заголовки ({len(col_of)}): " + ", ".join(f"`{h}`" for h in header_names if h))
    missing = [h for h in KEY_HEADERS if h not in col_of]
    if missing:
        lines.append("Нет на листе: " + ", ".join(f"`{h}`" for h in missing))
    lines.append("")

    for title, rows_set in GROUPS:
        rows = group_rows[title]
        missing_rows = sorted(r for r in rows_set if r > max_row_seen)
        lines.append(f"## {title}")
        lines.append("")
        if missing_rows:
            lines.append(f"Строк нет в файле (max_row={max_row_seen}): {missing_rows}")
        summary = _summarize(rows)
        lines.append(
            f"- строк: **{summary['n']}**; пустой Код RFP: **{summary['rfp_empty']}**; "
            f"пустой Код MTO: **{summary['mto_empty']}**"
        )
        lines.append(f"- титулы: {_counter_md(summary['titles'])}")
        lines.append(f"- Имя ДС: {_counter_md(summary['ds_names'])}")
        lines.append(f"- Статус УЛ: {_counter_md(summary['ul_status'])}")
        lines.append(f"- MTO статус позиции: {_counter_md(summary['pos_status'])}")
        lines.append(f"- заливка Титул/Марка: {_counter_md(summary['fills_title'])}")
        lines.append(f"- заливка Статус УЛ: {_counter_md(summary['fills_ul'])}")
        lines.append("")
        _write_table(lines, _sample(rows), TABLE_COLS)

    lines.append("## Титул 9110-KSB1")
    lines.append("")
    lines.append(f"Первые строки с `9110-KSB` в титуле: {len(ksb_rows)}")
    if ksb_rows:
        _write_table(lines, ksb_rows, TABLE_COLS)
    else:
        lines.append("")

    lines.append("## Watch-коды в Step4 (первые попадания)")
    lines.append("")
    for pair_label, hits in watch_hits.items():
        lines.append(f"### {pair_label}")
        lines.append(f"показов (cap {WATCH_HIT_LIMIT}): **{len(hits)}**")
        if hits:
            _write_table(lines, hits, TABLE_COLS)
        else:
            lines.append("")

    lines.append("## Таблица замен")
    lines.append("")
    notes, hits = _load_replacement_hits(REPL_XLSX)
    for note in notes:
        lines.append(f"- {note}")
    lines.append("")
    if hits:
        lines.append("| СТАРЫЙ_КОД | НОВЫЙ_КОД | СТАТУС |")
        lines.append("| --- | --- | --- |")
        for old, new, status in hits:
            lines.append(f"| {old} | {new} | {status} |")
        lines.append("")
    else:
        lines.append("Совпадений watch-кодов в таблице замен нет.")
        lines.append("")

    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    _safe_print(f"wrote {REPORT_PATH} ({REPORT_PATH.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
