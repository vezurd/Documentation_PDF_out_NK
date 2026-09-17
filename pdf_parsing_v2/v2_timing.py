"""Timing collector for the PDF v2 pipeline.

Collects:
- pipeline phases,
- detailed stage events,
- per-file extraction timings with breakdown columns.

Exports a multi-sheet Excel report suitable for performance analysis.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

import utils.excel as utils_excel

_EXTRACTION_DETAIL_THRESHOLD_SEC = 0.2
_EXTRACTION_BASE_COLUMNS = {"file_name", "elapsed_sec", "n_pages", "error"}


@dataclass
class TimingEntry:
    """One pipeline phase timing record."""

    label: str
    elapsed_sec: float
    detail: dict[str, Any] = field(default_factory=dict)


@dataclass
class StageTimingEntry:
    """One detailed stage timing record."""

    category: str
    stage: str
    elapsed_sec: float
    status: str = "ok"
    file_name: str = ""
    detail: dict[str, Any] = field(default_factory=dict)


class TimingCollector:
    """Accumulate timing data across pipeline phases and stages."""

    def __init__(self) -> None:
        self._entries: list[TimingEntry] = []
        self._stage_entries: list[StageTimingEntry] = []
        self._file_timings: dict[str, dict[str, Any]] = {}

    def record(self, label: str, elapsed_sec: float, **detail: Any) -> None:
        """Record a pipeline phase timing."""
        self._entries.append(TimingEntry(label, elapsed_sec, dict(detail)))

    def record_stage(
        self,
        category: str,
        stage: str,
        elapsed_sec: float,
        *,
        status: str = "ok",
        file_name: str = "",
        **detail: Any,
    ) -> None:
        """Record a detailed stage timing row."""
        self._stage_entries.append(
            StageTimingEntry(
                category=category,
                stage=stage,
                elapsed_sec=elapsed_sec,
                status=status,
                file_name=file_name,
                detail=dict(detail),
            )
        )

    def record_file(
        self, file_name: str, elapsed_sec: float, **detail: Any
    ) -> None:
        """Record per-file extraction timing and merge details."""
        current = dict(self._file_timings.get(file_name, {}))
        current["elapsed_sec"] = elapsed_sec
        current.update(detail)
        self._file_timings[file_name] = current

    @property
    def entries(self) -> list[TimingEntry]:
        return list(self._entries)

    @property
    def stage_entries(self) -> list[StageTimingEntry]:
        return list(self._stage_entries)

    @property
    def file_timings(self) -> dict[str, dict[str, Any]]:
        return dict(self._file_timings)

    def summary(self) -> dict[str, Any]:
        """Return JSON-serializable summary."""
        return {
            "phases": [
                {"label": e.label, "elapsed_sec": e.elapsed_sec, **e.detail}
                for e in self._entries
            ],
            "stages": [
                {
                    "category": e.category,
                    "stage": e.stage,
                    "elapsed_sec": e.elapsed_sec,
                    "status": e.status,
                    "file_name": e.file_name,
                    **e.detail,
                }
                for e in self._stage_entries
            ],
            "files": self._file_timings,
        }

    def _write_kv_sheet(
        self,
        ws: Any,
        title: str,
        rows: list[tuple[str, Any]],
        *,
        sanitize_illegal_chars: bool,
    ) -> None:
        ws.title = title
        ws.append(["Key", "Value"])
        for key, value in rows:
            ws.append(
                [
                    utils_excel.sanitize_cell_value_for_openpyxl(
                        key,
                        enabled=sanitize_illegal_chars,
                    ),
                    utils_excel.sanitize_cell_value_for_openpyxl(
                        value,
                        enabled=sanitize_illegal_chars,
                    ),
                ]
            )

    def _write_table_sheet(
        self,
        wb: Any,
        title: str,
        rows: list[dict[str, Any]],
        preferred_columns: list[str] | None = None,
        *,
        sanitize_illegal_chars: bool = True,
    ) -> None:
        ws = wb.create_sheet(title)
        if not rows:
            ws.append(["Info"])
            ws.append(["No data"])
            return
        present_columns: set[str] = set()
        for row in rows:
            present_columns.update(row.keys())
        columns: list[str] = []
        seen: set[str] = set()
        for column in preferred_columns or []:
            if column in present_columns and column not in seen:
                columns.append(column)
                seen.add(column)
        for row in rows:
            for column in row:
                if column not in seen:
                    columns.append(column)
                    seen.add(column)
        ws.append(
            [
                utils_excel.sanitize_cell_value_for_openpyxl(
                    column,
                    enabled=sanitize_illegal_chars,
                )
                for column in columns
            ]
        )
        for row in rows:
            ws.append(
                [
                    utils_excel.sanitize_cell_value_for_openpyxl(
                        row.get(column, ""),
                        enabled=sanitize_illegal_chars,
                    )
                    for column in columns
                ]
            )

    def _summary_rows(self) -> list[tuple[str, Any]]:
        files = self._file_timings
        stages = self._stage_entries
        total_errors = sum(1 for info in files.values() if info.get("error"))
        total_elapsed = 0.0
        for entry in self._entries:
            if entry.label == "total_pipeline":
                total_elapsed = entry.elapsed_sec
        rows: list[tuple[str, Any]] = [
            ("total_pipeline_sec", round(total_elapsed, 3)),
            ("n_phases", len(self._entries)),
            ("n_stage_events", len(stages)),
            ("n_files", len(files)),
            ("n_file_errors", total_errors),
        ]
        for category in ("extraction", "tags", "od", "rules", "report", "timing"):
            cat_total = sum(
                entry.elapsed_sec for entry in stages if entry.category == category
            )
            rows.append((f"{category}_total_sec", round(cat_total, 3)))
        return rows

    def _file_rows(self) -> list[dict[str, Any]]:
        visible_cols = self._visible_extraction_columns()
        rows: list[dict[str, Any]] = []
        for fname in sorted(self._file_timings):
            info = dict(self._file_timings[fname])
            row = {"file_name": fname}
            for key, value in info.items():
                if key not in visible_cols:
                    continue
                row[key] = self._round_value(value)
            rows.append(row)
        return rows

    def _stage_rows(self, category: str | None = None) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for entry in self._stage_entries:
            if category and entry.category != category:
                continue
            row = {
                "category": entry.category,
                "stage": entry.stage,
                "file_name": entry.file_name,
                "status": entry.status,
                "elapsed_sec": round(entry.elapsed_sec, 3),
            }
            for key, value in entry.detail.items():
                row[key] = self._round_value(value)
            rows.append(row)
        return rows

    def _visible_extraction_columns(self) -> set[str]:
        visible = set(_EXTRACTION_BASE_COLUMNS)
        max_by_key: dict[str, float] = {}
        for info in self._file_timings.values():
            for key, value in info.items():
                if key in _EXTRACTION_BASE_COLUMNS:
                    continue
                if isinstance(value, (int, float)):
                    max_by_key[key] = max(max_by_key.get(key, 0.0), float(value))
                elif value not in ("", None, {}):
                    visible.add(key)
        for key, value in max_by_key.items():
            if value >= _EXTRACTION_DETAIL_THRESHOLD_SEC:
                visible.add(key)
        return visible

    def _round_value(self, value: Any) -> Any:
        if isinstance(value, float):
            return round(value, 3)
        return value

    def to_excel(self, path: str, *, sanitize_illegal_chars: bool = True) -> None:
        """Write timing data to xlsx with summary and stage-specific sheets."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        self._write_kv_sheet(
            ws,
            "Summary",
            self._summary_rows(),
            sanitize_illegal_chars=sanitize_illegal_chars,
        )

        ws_phases = wb.create_sheet("Phases")
        ws_phases.append(["Phase", "Elapsed (s)", "Detail"])
        for e in self._entries:
            detail_text = str(e.detail or "")
            ws_phases.append(
                [
                    utils_excel.sanitize_cell_value_for_openpyxl(
                        e.label,
                        enabled=sanitize_illegal_chars,
                    ),
                    round(e.elapsed_sec, 3),
                    utils_excel.sanitize_cell_value_for_openpyxl(
                        detail_text,
                        enabled=sanitize_illegal_chars,
                    ),
                ]
            )

        self._write_table_sheet(
            wb,
            "Extraction",
            self._file_rows(),
            preferred_columns=[
                "file_name",
                "elapsed_sec",
                "n_pages",
                "error",
                "open_pdf",
                "page_loop_total",
                "char_index_build",
                "select_templates",
                "find_frame",
                "find_frame_filter_rects",
                "find_frame_span_scan",
                "find_frame_debug_top_rects",
                "find_frame_roi_build",
                "find_frame_roi_filter",
                "find_frame_union_build",
                "find_frame_total_drawings",
                "find_frame_after_edge_filter",
                "find_frame_border_band_dropped",
                "find_frame_after_border_band",
                "find_frame_border_band_enabled_pages",
                "find_tables",
                "grid_adapt",
                "extract_fields_text",
                "clean_fields",
                "template_scoring",
                "build_page_result",
            ],
            sanitize_illegal_chars=sanitize_illegal_chars,
        )
        self._write_table_sheet(
            wb,
            "Tags",
            self._stage_rows("tags"),
            preferred_columns=["stage", "elapsed_sec", "status", "file_name"],
            sanitize_illegal_chars=sanitize_illegal_chars,
        )
        self._write_table_sheet(
            wb,
            "OD",
            self._stage_rows("od"),
            preferred_columns=["stage", "elapsed_sec", "status", "file_name"],
            sanitize_illegal_chars=sanitize_illegal_chars,
        )
        self._write_table_sheet(
            wb,
            "Rules",
            self._stage_rows("rules"),
            preferred_columns=["stage", "elapsed_sec", "status", "file_name"],
            sanitize_illegal_chars=sanitize_illegal_chars,
        )
        self._write_table_sheet(
            wb,
            "Events",
            self._stage_rows(),
            preferred_columns=["category", "stage", "elapsed_sec", "status", "file_name"],
            sanitize_illegal_chars=sanitize_illegal_chars,
        )

        wb.save(path)
        wb.close()
