from datetime import datetime

import openpyxl
from openpyxl.styles import Border, Side, Alignment
from openpyxl.comments import Comment

import utils.path
from base.base_classes import *
from base.base_utils import tags_to_str
from base.tables_columns import (
    ColNames, UNITS, CODE, TAGS, VALUES, VENDOR, MASS, ANNOTATION, NUMBERS, TYPE_MARK, NAME,
    PROHIBITION, ROW_TYPE, TAGS_2, VALUES_2, NUMBERS_2, TYPE_MARK_2, NAME_2,
    DIFF_FLAG_TAG, DIFF_FLAG_VALUE, DIFF_FLAG_NAME, DIFF_FLAG_TYPE, DIFF_FLAG_SUMMARY, MTO_NAME,
    DS_NUMBER, BD_ID, BD_TABLE_NAME,
    DS_UNIT_PRICE_EXCL_VAT, DS_SHIPPING_PRICE_EXCL_VAT, DS_PACKAGING_PRICE_EXCL_VAT,
)
from utils.colors import Color
from utils.error_log import ErrorLog

excel_template_code_base = r"templates\out_template_google_code_base.xlsx"
excel_template_check_color = r"templates\out_template_check_mto_color.xlsx"
excel_template_check_rfq_color = r"templates\out_template_check_rfq_color.xlsx"
excel_template_nanocad_db = r"templates\out_template_NanocadDB.xlsx"
excel_template_ds_spec = r"templates\out_template_ds_spec.xlsx"
excel_template_ds_vs_mto = r"RFQ\ds_compare\templates\ds_vs_mto_template.xlsx"

def excel_base_out(in_list: list[RowStd],
                   in_dict=ColNames.column_list,
                   dir_base="base_check/",
                   file_name="excel_base_out"):
    # Загружаем шаблон Excel
    wb = openpyxl.load_workbook(excel_template_code_base)

    # grab the active worksheet
    ws = wb.active
    #
    title = []
    for att in in_dict:
        title.append(att)
    ws.append(title)
    # Добавляем к шаблону строчки с данными
    for row in in_list:
        #
        out_row = []
        for att in in_dict:
            out_row.append(str(row.el[att].value))
        #
        ws.append(out_row)

    # Save the file
    try:
        wb.save(dir_base + file_name + ".xlsx")
    except PermissionError as err:
        print(err)
        print("\tFAIL (excel_base_out)\n")


def check_color_out(base,
                    col_for_out_dict: dict,
                    out_dir,
                    file_prefix,
                    row_not_print_list=(
                            RowType.empty_row,
                    ),
                    excel_template_check_color=excel_template_check_color,
                    row_index=0,
                    comment_width=400,
                    comment_height=150,
                    show_row_type=False,
                    open_folder=False,
                    date_stamp=None,
                    output_middle_name=None):
    if not base:
        ErrorLog.add_error(f"ERROR: CHECK_COLOR_OUT: База для вывод в эксель пуста")
        return False
    if show_row_type:
        col_for_out_dict[99] = ROW_TYPE
    wb = None
    try:
        # Загружаем шаблон Excel
        wb = openpyxl.load_workbook(excel_template_check_color)
        wb.guess_types = True
        # grab the active worksheet
        ws = wb.active

        """
        Определяем форматирование в ячейках:
            horizontal (
            vertical
            wrap_text (True, False)
            
        HorizontalAlignmentsType
            "general", "left", "center", "right", "fill", "justify", "centerContinuous", "distributed"

        VerticalAlignmentsType
            "top", "center", "bottom", "justify", "distributed"
            
        """
        formating_dict = {
            TAGS: ["center", "center", True],
            TAGS_2: ["center", "center", True],
            NUMBERS: ["center", "center", True],
            NUMBERS_2: ["center", "center", True],
            NAME: ["left", "center", True],
            NAME_2: ["left", "center", True],
            TYPE_MARK: ["center", "center", True],
            TYPE_MARK_2: ["center", "center", True],
            CODE: ["center", "center", True],
            VENDOR: ["center", "center", True],
            UNITS: ["center", "center", True],
            VALUES: ["center", "center", True],
            VALUES_2: ["center", "center", True],
            MASS: ["center", "center", True],
            ANNOTATION: ["left", "center", True],
            PROHIBITION: ["center", "center", True],
            DIFF_FLAG_TAG: ["center", "center", True],
            DIFF_FLAG_VALUE: ["center", "center", True],
            DIFF_FLAG_NAME: ["center", "center", True],
            DIFF_FLAG_TYPE: ["center", "center", True],
            DIFF_FLAG_SUMMARY: ["center", "center", True],
            DS_NUMBER: ["center", "center", True],
            MTO_NAME: ["center", "center", True],
            BD_ID: ["center", "center", True],
            BD_TABLE_NAME: ["center", "center", True],
        }
        thin_border = Border(left=Side(style='thin'),
                             right=Side(style='thin'),
                             top=Side(style='thin'),
                             bottom=Side(style='thin'))

        # row_index = 0
        for row in base:
            if row.el[ROW_TYPE].value in row_not_print_list:
                continue
            row_index += 1
            column_index = 0
            for k, v in col_for_out_dict.items():
                column_index += 1
                cell = ws.cell(column=column_index, row=row_index)

                att = col_for_out_dict[k]  # Получаем имя столбца: "name" и т.п.

                value = row.el[att].value
                if isinstance(value, list):  # Если это список тегов (list)
                    value = tags_to_str(value)

                #Проверка типа данных перед записью (09/2025)
                # Для конвертации в числа для вывода в EXCEL
                float_format_list = [DS_UNIT_PRICE_EXCL_VAT,
                                     VALUES,
                                     DS_SHIPPING_PRICE_EXCL_VAT,
                                     DS_PACKAGING_PRICE_EXCL_VAT,]
                if att in float_format_list:
                    try:
                        value = float(value)
                    except (ValueError,TypeError):
                        pass
                # if att == DS_DELIVERY_TIME:
                #     value = datetime(value).date()

                #
                cell.value = value
                #

                # Добавляем комментарии в ячейку если они есть.
                cell_comment = row.el[att].comment
                # cell_comment += row.el[ROW_TYPE].value

                if len(cell_comment) > 0:
                    cell.comment = Comment(cell_comment, "Python_NK", width=comment_width, height=comment_height)

                if att in formating_dict:
                    horizontal = formating_dict[att][0]
                    vertical = formating_dict[att][1]
                    wrap_text = formating_dict[att][2]
                    cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text)
                    cell.border = thin_border
                else:
                    default_att = NUMBERS
                    horizontal = formating_dict[default_att][0]
                    vertical = formating_dict[default_att][1]
                    wrap_text = formating_dict[default_att][2]
                    cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text)
                    cell.border = thin_border

                color_index = row.el[att].color
                if color_index != Color.no:
                    try:
                        cell.fill = openpyxl.styles.PatternFill(start_color=color_index,
                                                                end_color=color_index,
                                                                fill_type="solid")
                    except ValueError:
                        print(f"color_index: <{column_index}>; value: <{value}>")

        # Save the file
        if output_middle_name is not None:
            base_file_name = str(output_middle_name).strip()
            if base_file_name.lower().endswith(".xlsx"):
                base_file_name = base_file_name[:-5]
        else:
            base_file_name = base[0].t_com.file_name.strip(".xlsx")
        if date_stamp:
            timestamp = datetime.now().strftime("%Y_%m_%d_%H_%M")
            base_file_name = base_file_name + "_" + timestamp
        file_name_out = file_prefix + "_" + base_file_name + ".xlsx"
        file_name = utils.path.file_path_plus_file_name(out_dir, file_name_out)
        try:
            utils.path.make_dir(out_dir)
            wb.save(file_name)
            print(f"\tЗавершен вывод в EXCEL (функция CHECK_MTO_COLOR)\n\t\t{file_name}\n")
            if open_folder:
                utils.path.open_dir(file_name)
            return file_name
        except PermissionError as err:
            print(err)
            print(f"\tОШИБКА вывода в ECELl функция CHECK_MTO_COLOR)\n\t{file_name}\n")
            return None
    finally:
        if wb is not None:
            wb.close()
