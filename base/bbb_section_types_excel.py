"""
Хранение конфигурации section_types в общем Excel файле на сетевом диске.

Файл используется несколькими пользователями. При конфликте записи (файл открыт
другим пользователем) выводится предупреждение, запись пропускается — некритично.
"""

import os
from datetime import datetime
from typing import Any, Dict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ─── Настройки файла ───────────────────────────────────────────────────────────
NETWORK_DIR = r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\РД\Таблички_графики\08_НК\11_BBB"
FILE_NAME = "BBB_типы_секций_MTO.xlsx"
SHEET_NAME = "Секции"

# Индексы столбцов (1-based)
COL_NAME      = 1   # Название секции
COL_TARGET    = 2   # Документ (BOE/BOM)
COL_VERIFIED  = 3   # Проверено (Да / Нет)
COL_SOURCE    = 4   # Источник файла
COL_ADDED_AT  = 5   # Дата добавления

HEADERS = [
    "Название секции",
    "Документ (BOE/BOM)",
    "Проверено",
    "Источник файла",
    "Дата добавления",
]

HEADER_COLOR   = "DBBCDB"
VERIFIED_COLOR = "D9EAD3"    # светло-зелёный — подтверждённые
NEW_COLOR      = "FFF2CC"    # светло-жёлтый  — новые, непроверенные

# ─── Базовые секции (используются при недоступности файла) ────────────────────
_BASE_SECTION_TYPES: Dict[str, Dict[str, Any]] = {
    "Аккумуляторные батареи":    {"target": "BOE", "verified": True},
    "Извещатели":                 {"target": "BOE", "verified": True},
    "Исполнительные устройства":  {"target": "BOE", "verified": True},
    "Кабеленесущие конструкции":  {"target": "BOE", "verified": True},
    "Кабельные изделия":          {"target": "BOE", "verified": True},
    "Коробки коммутационные":     {"target": "BOE", "verified": True},
    "Коробки коммутационные и блоки разветвительно-изолирующие": {"target": "BOE", "verified": True},
    "Материалы":                  {"target": "BOE", "verified": True},
    "Материалы для заземления":   {"target": "BOE", "verified": True},
    "Переферийное оборудование системы контроля и управления доступом": {"target": "BOE", "verified": True},
    "Приборы приемно-контрольные":{"target": "BOE", "verified": True},
    "Резервные источники питания":{"target": "BOE", "verified": True},
    "Шкафы и панели":             {"target": "BOE", "verified": True},
}


def get_excel_path() -> str:
    return os.path.join(NETWORK_DIR, FILE_NAME)


# ─── Загрузка ─────────────────────────────────────────────────────────────────

def load_section_types() -> Dict[str, Dict[str, Any]]:
    """
    Загружает section_types из Excel файла на сетевом диске.
    При недоступности файла возвращает базовые секции с предупреждением.
    """
    excel_path = get_excel_path()

    if not os.path.exists(excel_path):
        defaults = dict(_BASE_SECTION_TYPES)
        if _try_create_excel(excel_path, defaults):
            print(f"[INFO] Создан файл типов секций MTO: {excel_path}")
        else:
            print(
                f"[WARN] Файл типов секций MTO недоступен: {excel_path}\n"
                f"       Используются базовые секции (13 штук)."
            )
        return defaults

    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            wb.close()
            print(f"[WARN] Лист «{SHEET_NAME}» не найден в {FILE_NAME}. Используются базовые секции.")
            return dict(_BASE_SECTION_TYPES)

        ws = wb[SHEET_NAME]
        result: Dict[str, Dict[str, Any]] = {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            name = str(row[0]).strip()
            if not name:
                continue

            target_raw = str(row[1] or "BOE").strip().upper()
            target = target_raw if target_raw in ("BOE", "BOM") else "BOE"

            verified_raw = str(row[2] or "").strip().lower()
            verified = verified_raw in ("да", "yes", "true", "1")

            source_file = str(row[3]).strip() if len(row) > 3 and row[3] else ""
            added_at    = str(row[4]).strip() if len(row) > 4 and row[4] else ""

            entry: Dict[str, Any] = {"target": target, "verified": verified}
            if source_file:
                entry["source_file"] = source_file
            if added_at:
                entry["added_at"] = added_at

            result[name] = entry

        wb.close()
        return result

    except Exception as e:
        print(
            f"[WARN] Ошибка чтения файла типов секций MTO ({excel_path}): {e}\n"
            f"       Используются базовые секции."
        )
        return dict(_BASE_SECTION_TYPES)


# ─── Сохранение ───────────────────────────────────────────────────────────────

def save_section_types(section_types: Dict[str, Dict[str, Any]]) -> bool:
    """
    Сохраняет section_types в Excel на сетевом диске.
    При блокировке файла другим пользователем выводит предупреждение и возвращает False.
    """
    excel_path = get_excel_path()
    try:
        os.makedirs(os.path.dirname(excel_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        _write_sheet(ws, section_types)
        wb.save(excel_path)
        wb.close()
        return True

    except PermissionError:
        print(
            f"[WARN] Файл «{FILE_NAME}» заблокирован другим пользователем.\n"
            f"       Новая секция не сохранена в общий файл. Добавьте вручную:\n"
            f"       {excel_path}"
        )
        return False

    except Exception as e:
        print(f"[WARN] Ошибка сохранения файла типов секций MTO ({excel_path}): {e}")
        return False


def add_new_sections_if_missing(
    new_sections: Dict[str, Dict[str, Any]],
) -> bool:
    """
    Добавляет новые секции в Excel, если их там нет (читает → мёрджит → пишет).
    Возвращает True если запись прошла успешно.
    """
    if not new_sections:
        return True

    current = load_section_types()
    changed = False
    for name, entry in new_sections.items():
        if name not in current:
            current[name] = entry
            changed = True

    if not changed:
        return True

    return save_section_types(current)


# ─── Вспомогательные функции ──────────────────────────────────────────────────

def _try_create_excel(excel_path: str, section_types: Dict[str, Dict[str, Any]]) -> bool:
    try:
        os.makedirs(os.path.dirname(excel_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        _write_sheet(ws, section_types)
        wb.save(excel_path)
        wb.close()
        return True
    except Exception:
        return False


def _write_sheet(ws, section_types: Dict[str, Dict[str, Any]]) -> None:
    """Записывает заголовки и все строки на лист; verified — сверху по алфавиту."""

    # Заголовки
    header_fill = PatternFill(start_color=HEADER_COLOR, end_color=HEADER_COLOR, fill_type="solid")
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    # Сортировка: verified (алфавит) → unverified (алфавит)
    verified   = sorted((k, v) for k, v in section_types.items()
                         if isinstance(v, dict) and v.get("verified", False))
    unverified = sorted((k, v) for k, v in section_types.items()
                         if not (isinstance(v, dict) and v.get("verified", False)))

    verified_fill   = PatternFill(start_color=VERIFIED_COLOR, end_color=VERIFIED_COLOR, fill_type="solid")
    new_fill        = PatternFill(start_color=NEW_COLOR,      end_color=NEW_COLOR,      fill_type="solid")

    for row_idx, (name, entry) in enumerate(verified + unverified, start=2):
        is_verified = isinstance(entry, dict) and entry.get("verified", False)
        target      = entry.get("target", "BOE") if isinstance(entry, dict) else "BOE"
        verified_label = "Да" if is_verified else "Нет"
        source_file    = entry.get("source_file", "") if isinstance(entry, dict) else ""
        added_at       = entry.get("added_at",    "") if isinstance(entry, dict) else ""

        row_fill = verified_fill if is_verified else new_fill
        values = [name, target, verified_label, source_file, added_at]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="center")

    # Ширины столбцов
    col_widths = {COL_NAME: 58, COL_TARGET: 16, COL_VERIFIED: 12, COL_SOURCE: 45, COL_ADDED_AT: 18}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
