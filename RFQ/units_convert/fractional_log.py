"""Write per-contour xlsx logs for fractional unit-conversion results."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from RFQ.units_convert.models import ConversionPlan, FractionalConversionIssue

YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")

CONTOUR_LOG_FILENAMES: dict[str, str] = {
    "parts": "дробные значения после конвертации RFP.xlsx",
    "rfp": "дробные значения после конвертации RFP.xlsx",
    "mto": "дробные значения после конвертации MTO.xlsx",
    "ul": "дробные значения после конвертации УЛ.xlsx",
}

_HEADERS = (
    "Контур",
    "Источник",
    "Код",
    "Исходное кол-во",
    "Исходная ЕИ",
    "Целевая ЕИ",
    "Коэффициент",
    "Результат после конвертации",
    "Ближайшее целое",
    "Отклонение",
    "Теги",
    "Инвариант не проверялся",
)

_COLUMN_WIDTHS = (12, 70, 16, 16, 14, 14, 18, 26, 16, 16, 10, 24)


def _log_filename(contour: str) -> str:
    known = CONTOUR_LOG_FILENAMES.get(contour)
    if known:
        return known
    return f"дробные значения после конвертации {contour}.xlsx"


def write_fractional_conversion_logs(
    plan: ConversionPlan,
    out_dir: Path,
) -> list[Path]:
    """Write one xlsx per data type for fractional conversion rows.

    Args:
        plan: Conversion plan that may contain ``fractional_issues``.
        out_dir: Directory for log workbooks.

    Returns:
        Written workbook paths, one per contour group that had rows.
    """
    grouped: dict[str, list[FractionalConversionIssue]] = defaultdict(list)
    for issue in plan.fractional_issues:
        grouped[_log_filename(issue.contour)].append(issue)
    if not grouped:
        return []

    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for filename, issues in sorted(grouped.items()):
        path = out_dir / filename
        _write_issue_workbook(path, issues)
        written.append(path)
    return written


def _write_issue_workbook(
    path: Path,
    issues: list[FractionalConversionIssue],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Дробные значения"
    ws.append(list(_HEADERS))
    for issue in issues:
        ws.append(
            [
                issue.contour,
                issue.location,
                issue.code,
                str(issue.original_quantity),
                issue.source_unit,
                issue.target_unit,
                str(issue.coefficient),
                str(issue.result_quantity),
                str(issue.nearest_int),
                str(issue.residual),
                issue.tags_count,
                "да" if issue.invariant_skipped else "нет",
            ]
        )
        ws.cell(row=ws.max_row, column=8).fill = YELLOW_FILL
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{max(ws.max_row, 1)}"
    wb.save(path)
    wb.close()
