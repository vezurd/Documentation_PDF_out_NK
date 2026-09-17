"""Build and sync the informational actual/sequential DS sheet.

The manager sheet (``Имя ДС`` / ``Фамилия``) values are never rewritten.
Status columns to the right of those headers are refreshed on each sync.
The roster sheet is rebuilt when its data fingerprint differs.
"""

from __future__ import annotations

import os
import sys
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from RFQ.rfp_parts.ds_checklist import (
    DEFAULT_PARTS_DIR,
    parse_ds_name_from_file_name,
    table_ds_to_file_keys,
)
from RFQ.rfp_parts.ds_identity import DsIdentity, parse_rfp_ds_identity
from RFQ.tags_rfp_compare.ds_manager_matrix import (
    DsManagerEntry,
    DsManagerMatrix,
    ds_key_aliases,
    load_ds_manager_matrix,
    lookup_manager,
)

ROSTER_SHEET_TITLE = "Фактический и порядковый"
# Row 1 matches Лист2 so the whole table can be pasted there.
# Extra columns after «Фамилия» are informational; the loader ignores them.
_ROSTER_HEADERS = (
    "Имя ДС",
    "Кол-во RFP",
    "Фамилия",
    "Фактический ДС",
    "Порядковый ДС",
    "Статус",
    "Было на Лист2",
    "Дубликат",
    "Что делать",
)
# Written to the right of native Лист2 columns; never named Имя ДС / Фамилия.
_LIST2_EXTRA_HEADERS = (
    "Статус",
    "Дубликат",
    "Что делать",
)
_LIST2_NATIVE_HEADERS = frozenset({"Имя ДС", "№ ДС", "Кол-во RFP", "Фамилия", "МП"})

_KIND_EXACT = "exact"
_KIND_INVERTED = "inverted"
_KIND_ACTUAL = "actual_only"
_STATUS_OK = "фамилия с Лист2"
_STATUS_INVERTED = "фамилия с Лист2, имя на Лист2 перевёрнуто"
_STATUS_ACTUAL = "фамилия с Лист2 по фактическому номеру"
_STATUS_EMPTY = "на Лист2 нет фамилии — заполните"
_STATUS_MISSING = "нет на Лист2 — заполните"
_STATUS_LEFTOVER = "только на Лист2, файла RFP нет"
_STATUS_DUPLICATE = "дубликат на Лист2, файл уже закрыт другой строкой"
_REMARK_STATUSES = frozenset({_STATUS_LEFTOVER, _STATUS_DUPLICATE})
_FILL_STATUSES = frozenset({_STATUS_MISSING, _STATUS_EMPTY})

_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_FILL_HEAD = PatternFill("solid", fgColor="D6DCE4")
_FILL_OK = PatternFill("solid", fgColor="C6EFCE")
_FILL_WARN = PatternFill("solid", fgColor="FFF2CC")
_FILL_LEFT = PatternFill("solid", fgColor="F2F2F2")
_FONT_H = Font(name="Calibri", size=11, bold=True)
_FONT_N = Font(name="Calibri", size=11)
_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")


@dataclass(frozen=True)
class DsRosterRow:
    """One informational DS row for copy-paste onto the manager sheet."""

    actual_label: str
    sequential_label: str
    copy_name: str
    manager: str
    status: str
    was_on_sheet1: str
    note: str = ""
    source_row: int = 0
    duplicate_label: str = ""


@dataclass(frozen=True)
class DsRosterSyncResult:
    """Outcome of comparing and optionally rewriting the roster sheet."""

    wrote: bool
    changed: bool
    row_count: int
    matched_manager: int
    missing_manager: int
    leftovers: int
    summary_line: str
    write_error: str | None = None
    saved_path: str | None = None
    used_fallback: bool = False


@dataclass(frozen=True)
class DsRosterCompareResult:
    """GUI / CLI snapshot of RFP files vs the manager-matrix surnames."""

    matrix_path: Path
    parts_dir: Path
    rows: list[DsRosterRow]
    load_error: str | None
    sync: DsRosterSyncResult | None = None

    @property
    def needs_fill(self) -> bool:
        """True when an RFP file has no surname on Лист2."""
        return any(row.status in _FILL_STATUSES for row in self.rows)

    @property
    def is_ok(self) -> bool:
        """True when every RFP file has a surname and Лист2 has no leftovers."""
        if self.load_error:
            return False
        if self.sync is not None and self.sync.write_error:
            return False
        if self.needs_fill:
            return False
        if any(row.status in _REMARK_STATUSES for row in self.rows):
            return False
        if any(row.status == _STATUS_INVERTED for row in self.rows):
            return False
        return True

    def summary_line(self) -> str:
        """Russian one-liner for banners and the job monitor."""
        if self.sync is not None:
            return self.sync.summary_line
        return _summary(self.rows, wrote=False, error=self.load_error)


@dataclass(frozen=True, slots=True)
class DsRosterJobResult:
    """Duck-typed ``FunctionJobRunner`` result (no GUI import)."""

    success: bool
    message: str
    result_path: str | None = None


_last_roster: DsRosterCompareResult | None = None


def roster_status_tone(status: str) -> str:
    """Map a roster status string to a GUI color key.

    Args:
        status: Value of ``DsRosterRow.status``.

    Returns:
        One of ``ok``, ``warn``, ``leftover``, ``problem``.
    """
    if status == _STATUS_OK:
        return "ok"
    if status in {_STATUS_INVERTED, _STATUS_ACTUAL, _STATUS_DUPLICATE}:
        return "warn"
    if status == _STATUS_LEFTOVER:
        return "leftover"
    return "problem"


def resolve_ds_manager_matrix_path() -> Path:
    """Return ``paths.ds_manager_matrix`` from the RFP JSON config.

    Returns:
        Configured workbook path. Falls back to ``get_default_config`` when
        the live config cannot be loaded or the key is empty.
    """
    from RFQ.tags_rfp_compare.rfp_tags_utils import get_default_config, load_config

    try:
        cfg = load_config()
    except Exception:
        cfg = get_default_config()
    raw = str((cfg.get("paths") or {}).get("ds_manager_matrix") or "").strip()
    if not raw:
        raw = str(get_default_config()["paths"]["ds_manager_matrix"])
    return Path(raw)


def _identity_from_label(raw: str) -> DsIdentity:
    """Parse a matrix / prefix label; retry with ``. `` so compounds split."""
    text = str(raw or "").strip()
    ident = parse_rfp_ds_identity(text)
    if ident.kind != "compound":
        retry = parse_rfp_ds_identity(text + ". ")
        if retry.compound:
            return retry
    return ident


def _ds_label(number: int | None) -> str:
    return f"ДС{number}" if number is not None else ""


def _copy_name(ident: DsIdentity, *, letter: str = "") -> str:
    mark = (letter or ident.letter or "").upper()
    if ident.compound and ident.actual is not None and ident.sequential is not None:
        return f"ДС{ident.actual}_{ident.sequential}{mark}"
    if ident.actual is not None:
        return f"ДС{ident.actual}{mark}"
    keys = table_ds_to_file_keys(ident.raw)
    return keys[0] if keys else ident.label


def _entry_identity(entry: DsManagerEntry) -> DsIdentity:
    return _identity_from_label(entry.raw_ds)


def _pair(ident: DsIdentity) -> frozenset[int]:
    nums = [n for n in (ident.actual, ident.sequential) if n is not None]
    return frozenset(nums)


def _match_kind(entry: DsManagerEntry, ds_name: str, ident: DsIdentity) -> str | None:
    file_aliases = set(ds_key_aliases(ds_name))
    entry_aliases = set(entry.keys)
    if file_aliases & entry_aliases:
        return _KIND_EXACT
    eident = _entry_identity(entry)
    if (
        ident.compound
        and eident.compound
        and ident.actual is not None
        and ident.sequential is not None
        and ident.actual != ident.sequential
        and _pair(ident) == _pair(eident)
        and len(_pair(ident)) == 2
    ):
        if (eident.actual, eident.sequential) == (ident.sequential, ident.actual):
            return _KIND_INVERTED
        if (eident.actual, eident.sequential) == (ident.actual, ident.sequential):
            return _KIND_EXACT
    if (
        ident.actual is not None
        and not ident.compound
        and not eident.compound
        and eident.actual == ident.actual
    ):
        return _KIND_ACTUAL
    return None


def _pick_entry(
    matrix: DsManagerMatrix,
    ds_name: str,
    ident: DsIdentity,
) -> tuple[DsManagerEntry | None, str]:
    direct = lookup_manager(matrix, ds_name)
    if direct is not None:
        return direct, _KIND_EXACT

    ranked: list[tuple[tuple[int, int, int], DsManagerEntry, str]] = []
    for entry in matrix.entries:
        kind = _match_kind(entry, ds_name, ident)
        if kind is None:
            continue
        kind_rank = {_KIND_EXACT: 3, _KIND_INVERTED: 2, _KIND_ACTUAL: 1}[kind]
        has_mgr = 1 if entry.manager else 0
        ranked.append(((kind_rank, has_mgr, -len(entry.raw_ds)), entry, kind))
    if not ranked:
        return None, ""
    ranked.sort(key=lambda item: item[0], reverse=True)
    _score, entry, kind = ranked[0]
    return entry, kind


def _status_for(kind: str, manager: str, *, leftover: bool = False) -> str:
    if leftover:
        return _STATUS_LEFTOVER
    if not kind:
        return _STATUS_MISSING
    if not manager:
        return _STATUS_EMPTY
    if kind == _KIND_INVERTED:
        return _STATUS_INVERTED
    if kind == _KIND_ACTUAL:
        return _STATUS_ACTUAL
    return _STATUS_OK


def _collect_rfp_groups(parts_dir: Path) -> dict[tuple[int, int], tuple[str, DsIdentity]]:
    """Group RFP workbooks by (actual, sequential); keep a representative prefix."""
    grouped: dict[tuple[int, int], tuple[str, DsIdentity]] = {}
    if not parts_dir.exists():
        return grouped
    for path in sorted(parts_dir.iterdir(), key=lambda item: item.name.lower()):
        if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
            continue
        if path.name.startswith("~$"):
            continue
        ident = parse_rfp_ds_identity(path.name)
        if ident.actual is None:
            continue
        sequential = ident.sequential if ident.sequential is not None else ident.actual
        key = (ident.actual, sequential)
        ds_name = parse_ds_name_from_file_name(path.name)
        previous = grouped.get(key)
        if previous is None:
            grouped[key] = (ds_name, ident)
            continue
        prev_name, prev_ident = previous
        if ident.letter and not prev_ident.letter:
            grouped[key] = (ds_name, ident)
        elif len(ds_name) < len(prev_name):
            grouped[key] = (ds_name, ident)
    return grouped


def build_ds_roster_rows(
    matrix: DsManagerMatrix,
    parts_dir: str | Path | None = None,
) -> list[DsRosterRow]:
    """Build informational rows from RFP file names plus leftover matrix DS.

    Args:
        matrix: Loaded manager matrix (Лист2).
        parts_dir: ``RFP_Зиновьев``. Defaults to ``DEFAULT_PARTS_DIR``.

    Returns:
        Rows sorted by actual, then sequential; matrix leftovers last.
    """
    folder = Path(parts_dir or DEFAULT_PARTS_DIR)
    grouped = _collect_rfp_groups(folder)
    used: set[int] = set()
    rows: list[DsRosterRow] = []

    for (actual, sequential), (ds_name, ident) in sorted(grouped.items()):
        entry, kind = _pick_entry(matrix, ds_name, ident)
        letter = ident.letter
        if entry is not None:
            used.add(id(entry))
            eident = _entry_identity(entry)
            if not letter and eident.letter:
                letter = eident.letter
        manager = entry.manager if entry is not None else ""
        was = entry.raw_ds if entry is not None else ""
        rows.append(
            DsRosterRow(
                actual_label=_ds_label(actual),
                sequential_label=_ds_label(sequential),
                copy_name=_copy_name(ident, letter=letter),
                manager=manager,
                status=_status_for(kind, manager),
                was_on_sheet1=was,
                source_row=entry.source_row if entry is not None else 0,
            )
        )

    leftovers: list[DsRosterRow] = []
    for entry in matrix.entries:
        if id(entry) in used:
            continue
        ident = _entry_identity(entry)
        actual = ident.actual
        sequential = ident.sequential if ident.sequential is not None else ident.actual
        leftovers.append(
            DsRosterRow(
                actual_label=_ds_label(actual),
                sequential_label=_ds_label(sequential),
                copy_name=_copy_name(ident) or entry.raw_ds,
                manager=entry.manager,
                status=_STATUS_LEFTOVER,
                was_on_sheet1=entry.raw_ds,
                source_row=entry.source_row,
            )
        )
    leftovers.sort(key=lambda row: (row.actual_label, row.sequential_label, row.copy_name))
    rows.extend(leftovers)
    return _with_notes(rows)


def _row_note(row: DsRosterRow, *, duplicate_of: str = "") -> str:
    """Human verdict for the «Что делать» column and the GUI table."""
    if row.status == _STATUS_MISSING:
        return (
            f"Надо дозаполнить: добавьте строку на Лист2 с именем {row.copy_name} "
            f"и фамилией. Файл RFP есть, фактический {row.actual_label}."
        )
    if row.status == _STATUS_EMPTY:
        return (
            f"Надо дозаполнить фамилию. Файл RFP есть ({row.copy_name}), "
            f"фактический {row.actual_label}."
        )
    if row.status == _STATUS_DUPLICATE:
        name = row.was_on_sheet1 or row.copy_name
        extra = (
            f" Фактический {row.actual_label} уже закрыт ({duplicate_of})."
            if duplicate_of
            else ""
        )
        return (
            f"Дозаполнять не нужно. Дубликат: строка «{name}» не используется.{extra} "
            "Проверьте фамилию или удалите строку."
        )
    if row.status == _STATUS_LEFTOVER:
        name = row.was_on_sheet1 or row.copy_name
        bits = [
            "Дозаполнять не нужно.",
            f"Файла RFP с именем «{name}» нет.",
        ]
        if row.sequential_label and row.sequential_label != row.actual_label:
            bits.append(
                "Похоже на устаревшее/перевёрнутое имя или копию файла "
                "(например ДС4905_1). Можно удалить, если ДС уже есть другой строкой."
            )
        else:
            bits.append(
                "Лишняя строка: архив или ДС ещё без файла. Можно удалить, если не нужна."
            )
        return " ".join(bits)
    if row.status == _STATUS_INVERTED:
        return (
            f"Фамилию заполнять не нужно. Имя на Лист2 перевёрнуто "
            f"(было {row.was_on_sheet1}). Файл RFP: {row.copy_name}, "
            f"фактический {row.actual_label}. Исправьте «Имя ДС» на {row.copy_name}."
        )
    if row.status == _STATUS_ACTUAL:
        return (
            f"Фамилию заполнять не нужно. Сопоставлено по фактическому номеру "
            f"{row.actual_label}: строка Лист2 «{row.was_on_sheet1}», "
            f"файл RFP {row.copy_name}."
        )
    return (
        f"ОК, дозаполнять не нужно. Файл RFP {row.copy_name}, "
        f"фактический {row.actual_label}."
    )


def _duplicate_label(*, duplicate_of: str, duplicate_peer: str) -> str:
    if duplicate_of:
        return f"да: файл уже закрыт ({duplicate_of})"
    if duplicate_peer:
        return f"есть лишняя строка ({duplicate_peer})"
    return "нет"


def _with_notes(rows: list[DsRosterRow]) -> list[DsRosterRow]:
    covered: dict[str, str] = {}
    leftover_by_actual: dict[str, list[str]] = defaultdict(list)
    for row in rows:
        if row.status != _STATUS_LEFTOVER and row.actual_label:
            covered.setdefault(
                row.actual_label,
                f"{row.copy_name}" + (f", {row.manager}" if row.manager else ""),
            )
        elif row.status == _STATUS_LEFTOVER and row.actual_label:
            leftover_by_actual[row.actual_label].append(
                f"{row.was_on_sheet1}" + (f", {row.manager}" if row.manager else "")
            )
    out: list[DsRosterRow] = []
    for row in rows:
        dup_of = ""
        dup_peer = ""
        status = row.status
        if row.status == _STATUS_LEFTOVER and row.actual_label in covered:
            dup_of = covered[row.actual_label]
            status = _STATUS_DUPLICATE
        if row.status != _STATUS_LEFTOVER and row.actual_label in leftover_by_actual:
            dup_peer = "; ".join(leftover_by_actual[row.actual_label])
        updated = replace(row, status=status)
        out.append(
            replace(
                updated,
                note=_row_note(updated, duplicate_of=dup_of),
                duplicate_label=_duplicate_label(
                    duplicate_of=dup_of, duplicate_peer=dup_peer
                ),
            )
        )
    return out


def _fingerprint(rows: list[DsRosterRow]) -> tuple[tuple[str, ...], ...]:
    return tuple(
        (
            row.copy_name,
            "",
            row.manager,
            row.actual_label,
            row.sequential_label,
            row.status,
            row.was_on_sheet1,
            row.duplicate_label,
            row.note,
        )
        for row in rows
    )


def _read_stored_fingerprint(workbook) -> tuple[tuple[str, ...], ...] | None:
    if ROSTER_SHEET_TITLE not in workbook.sheetnames:
        return None
    sheet = workbook[ROSTER_SHEET_TITLE]
    header_row: int | None = None
    for idx, row in enumerate(sheet.iter_rows(max_row=8, values_only=True), start=1):
        texts = tuple(str(cell or "").strip() for cell in row[: len(_ROSTER_HEADERS)])
        if texts == _ROSTER_HEADERS:
            header_row = idx
            break
    if header_row is None:
        return None
    stored: list[tuple[str, ...]] = []
    width = len(_ROSTER_HEADERS)
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        padded = (row or ()) + (None,) * width
        values = tuple(str(cell or "").strip() for cell in padded[:width])
        if not any(values[:3]):
            continue
        stored.append(values)
    return tuple(stored)


def _paint(cell, *, fill: PatternFill, font: Font | None = None) -> None:
    cell.alignment = _WRAP
    cell.border = _THIN
    cell.font = font or _FONT_N
    cell.fill = fill


def _fill_for(row: DsRosterRow) -> PatternFill:
    if row.status == _STATUS_DUPLICATE:
        return _FILL_WARN
    if row.status == _STATUS_LEFTOVER:
        return _FILL_LEFT
    if row.manager:
        return _FILL_OK
    return _FILL_WARN


def _write_roster_sheet(sheet: Worksheet, rows: list[DsRosterRow]) -> None:
    sheet.sheet_properties.tabColor = "548235"
    widths = {1: 18, 2: 14, 3: 18, 4: 18, 5: 18, 6: 48, 7: 22, 8: 36, 9: 56}
    for idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(idx)].width = width

    for col, header in enumerate(_ROSTER_HEADERS, start=1):
        cell = sheet.cell(1, col, header)
        _paint(cell, fill=_FILL_HEAD, font=_FONT_H)
        cell.comment = None
    sheet.row_dimensions[1].height = 22

    for offset, row in enumerate(rows):
        excel_row = 2 + offset
        fill = _fill_for(row)
        values = (
            row.copy_name,
            "",
            row.manager,
            row.actual_label,
            row.sequential_label,
            row.status,
            row.was_on_sheet1,
            row.duplicate_label,
            row.note,
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, col, value)
            _paint(cell, fill=fill)
            cell.comment = None
        sheet.row_dimensions[excel_row].height = 20

    last_col = get_column_letter(len(_ROSTER_HEADERS))
    last = 1 + max(len(rows), 1)
    sheet.auto_filter.ref = f"A1:{last_col}{last}"
    sheet.freeze_panes = "A2"


def _summary(
    rows: list[DsRosterRow],
    *,
    wrote: bool,
    error: str | None,
    fallback_name: str | None = None,
) -> str:
    if error:
        return f"Ошибка: {error}"
    missing = [row for row in rows if row.status in _FILL_STATUSES]
    leftovers = [row for row in rows if row.status == _STATUS_LEFTOVER]
    duplicates = [row for row in rows if row.status == _STATUS_DUPLICATE]
    inverted = [row for row in rows if row.status == _STATUS_INVERTED]
    if wrote and fallback_name:
        write_bit = f" Записана копия (файл занят) {fallback_name}."
    elif wrote:
        write_bit = " Столбцы статуса записаны в книгу."
    else:
        write_bit = ""

    def _names(items: list[DsRosterRow], *, leftover: bool = False) -> str:
        labels: list[str] = []
        for row in items:
            label = (row.was_on_sheet1 if leftover else row.copy_name) or row.actual_label
            if label and label not in labels:
                labels.append(label)
        return ", ".join(labels)

    if missing:
        line = f"Надо дозаполнить: {_names(missing)}."
        if leftovers:
            line += f" Лишние на Лист2 (не дозаполнение): {_names(leftovers, leftover=True)}."
        if duplicates:
            line += f" Дубликаты: {_names(duplicates, leftover=True)}."
        return line + write_bit
    if leftovers or duplicates or inverted:
        bits = ["Фамилии по файлам RFP заполнены."]
        if leftovers:
            bits.append(f"Проверьте лишние на Лист2: {_names(leftovers, leftover=True)}.")
        if duplicates:
            bits.append(f"Дубликаты: {_names(duplicates, leftover=True)}.")
        if inverted:
            bits.append(f"Перевёрнутые имена: {_names(inverted, leftover=True)}.")
        return " ".join(bits) + write_bit
    return f"ОК: все файлы RFP с фамилией, лишних на Лист2 нет.{write_bit}"


def _sync_result(
    rows: list[DsRosterRow],
    *,
    wrote: bool,
    changed: bool,
    error: str | None = None,
    saved_path: str | None = None,
    used_fallback: bool = False,
) -> DsRosterSyncResult:
    fallback_name = Path(saved_path).name if used_fallback and saved_path else None
    return DsRosterSyncResult(
        wrote=wrote,
        changed=changed,
        row_count=len(rows),
        matched_manager=sum(
            1
            for row in rows
            if row.manager and row.status not in _REMARK_STATUSES
        ),
        missing_manager=sum(1 for row in rows if row.status in _FILL_STATUSES),
        leftovers=sum(1 for row in rows if row.status in _REMARK_STATUSES),
        summary_line=_summary(
            rows,
            wrote=wrote,
            error=error,
            fallback_name=fallback_name,
        ),
        write_error=error,
        saved_path=saved_path,
        used_fallback=used_fallback,
    )


def _stamped_copy_path(path: Path) -> Path:
    stamp = datetime.now().strftime("%Y.%m.%d_%H.%M.%S")
    candidate = path.with_name(f"{path.stem}_{stamp}{path.suffix}")
    if not candidate.exists():
        return candidate
    stamp = datetime.now().strftime("%Y.%m.%d_%H.%M.%S.%f")
    return path.with_name(f"{path.stem}_{stamp}{path.suffix}")


def _write_bytes_replace_or_stamp(path: Path, data: bytes) -> tuple[Path | None, str | None, bool]:
    """Replace ``path`` with ``data``; if locked, write a timestamped sibling.

    Returns:
        ``(saved_path, error, used_fallback)``.
    """
    tmp_path = path.with_name(f"{path.stem}._roster_tmp{path.suffix}")
    try:
        tmp_path.write_bytes(data)
        os.replace(tmp_path, path)
        return path, None, False
    except OSError:
        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            pass

    copy_path = _stamped_copy_path(path)
    try:
        copy_path.write_bytes(data)
        return copy_path, None, True
    except OSError as exc:
        error = str(exc)
        if isinstance(exc, PermissionError) or "Permission" in error:
            error = (
                f"не удалось записать ни «{path.name}», ни копию "
                f"«{copy_path.name}» — закройте Excel и повторите запуск RFP"
            )
        return None, error, False


def _header_map(sheet) -> dict[str, int]:
    """Return 1-based column index for each non-empty header on row 1."""
    found: dict[str, int] = {}
    for cell in sheet[1]:
        text = str(cell.value or "").strip()
        if text and text not in found:
            found[text] = int(cell.column)
    return found


def _ensure_list2_extra_headers(sheet) -> dict[str, int]:
    """Append status columns to the right of native Лист2 headers.

    Never inserts between ``Имя ДС`` / ``Фамилия``. Reuses a header if it
    already exists (for example after pasting the roster sheet).
    """
    found = _header_map(sheet)
    native_last = 1
    for header, col in found.items():
        if header in _LIST2_NATIVE_HEADERS:
            native_last = max(native_last, col)
    next_col = max([native_last, *found.values()], default=1) + 1
    widths = {"Статус": 48, "Дубликат": 36, "Что делать": 56}
    for header in _LIST2_EXTRA_HEADERS:
        if header in found:
            continue
        cell = sheet.cell(1, next_col, header)
        _paint(cell, fill=_FILL_HEAD, font=_FONT_H)
        cell.comment = None
        sheet.column_dimensions[get_column_letter(next_col)].width = widths.get(
            header, 24
        )
        found[header] = next_col
        next_col += 1
    return found


def _clear_sheet_comments(sheet, *, max_row: int, max_col: int) -> bool:
    changed = False
    if max_row < 1 or max_col < 1:
        return False
    for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.comment is not None:
                cell.comment = None
                changed = True
    return changed


def _list2_extra_values(row: DsRosterRow) -> dict[str, str]:
    return {
        "Статус": row.status,
        "Дубликат": row.duplicate_label,
        "Что делать": row.note,
    }


def _annotate_manager_sheet(
    workbook,
    matrix: DsManagerMatrix,
    rows: list[DsRosterRow],
) -> bool:
    """Write status columns on Лист2. Does not change names or surnames."""
    title = matrix.sheet_title
    if not title or title not in workbook.sheetnames:
        return False
    sheet = workbook[title]
    ds_col = matrix.ds_col + 1
    mp_col = matrix.mp_col + 1
    headers = _ensure_list2_extra_headers(sheet)
    extra_cols = [headers[name] for name in _LIST2_EXTRA_HEADERS if name in headers]
    max_col = max([ds_col, mp_col, *extra_cols], default=1)
    last_ds = 1
    for excel_row in range(2, (sheet.max_row or 1) + 1):
        if sheet.cell(excel_row, ds_col).value:
            last_ds = excel_row
    max_row = max(
        last_ds,
        *(row.source_row for row in rows if row.source_row),
        1,
    )
    changed = _clear_sheet_comments(sheet, max_row=max_row, max_col=max_col)
    by_row = {row.source_row: row for row in rows if row.source_row >= 2}

    for excel_row in range(2, max_row + 1):
        row = by_row.get(excel_row)
        payload = _list2_extra_values(row) if row is not None else None
        fill = _fill_for(row) if row is not None else None
        for header in _LIST2_EXTRA_HEADERS:
            col = headers.get(header)
            if col is None or col in {ds_col, mp_col}:
                continue
            cell = sheet.cell(excel_row, col)
            value = payload[header] if payload is not None else ""
            if str(cell.value or "") != value:
                cell.value = value
                changed = True
            if fill is not None:
                _paint(cell, fill=fill)
            cell.comment = None
    return changed


def _list2_status_match(
    workbook,
    matrix: DsManagerMatrix,
    rows: list[DsRosterRow],
) -> bool:
    title = matrix.sheet_title
    if not title or title not in workbook.sheetnames:
        return False
    sheet = workbook[title]
    headers = _header_map(sheet)
    if any(name not in headers for name in _LIST2_EXTRA_HEADERS):
        return False
    extra_cols = [headers[name] for name in _LIST2_EXTRA_HEADERS]
    ds_col = matrix.ds_col + 1
    mp_col = matrix.mp_col + 1
    max_col = max([ds_col, mp_col, *extra_cols])
    last_ds = 1
    for excel_row in range(2, (sheet.max_row or 1) + 1):
        if sheet.cell(excel_row, ds_col).value:
            last_ds = excel_row
    max_row = max(
        last_ds,
        *(row.source_row for row in rows if row.source_row),
        1,
    )
    for row_cells in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row_cells:
            if cell.comment is not None:
                return False
    by_row = {row.source_row: row for row in rows if row.source_row >= 2}
    for excel_row in range(2, max_row + 1):
        payload = _list2_extra_values(by_row[excel_row]) if excel_row in by_row else None
        for header in _LIST2_EXTRA_HEADERS:
            col = headers[header]
            actual = str(sheet.cell(excel_row, col).value or "").strip()
            expected = payload[header] if payload is not None else ""
            if actual != expected:
                return False
    return True


def sync_ds_roster_sheet(
    matrix: DsManagerMatrix,
    parts_dir: str | Path | None = None,
) -> DsRosterSyncResult:
    """Rebuild the roster sheet and refresh status columns on Лист2.

    Args:
        matrix: Loaded manager matrix. ``load_error`` skips the write.
        parts_dir: RFP parts folder. Defaults to ``DEFAULT_PARTS_DIR``.

    Returns:
        Sync result with a Russian ``summary_line``. Write errors are returned,
        not raised. If the target xlsx is locked, a sibling copy with date/time
        in the file name is written instead. Лист2 names and surnames are not
        rewritten; status / duplicate / «Что делать» columns are appended to
        the right of the native headers.
    """
    rows = build_ds_roster_rows(matrix, parts_dir=parts_dir)
    expected = _fingerprint(rows)
    if matrix.load_error:
        return _sync_result(rows, wrote=False, changed=False, error=matrix.load_error)

    path = matrix.path
    try:
        payload = path.read_bytes()
    except OSError as exc:
        return _sync_result(rows, wrote=False, changed=True, error=str(exc))

    try:
        workbook = load_workbook(BytesIO(payload))
    except Exception as exc:  # noqa: BLE001 — return write_error per contract
        return _sync_result(rows, wrote=False, changed=True, error=str(exc))

    try:
        stored = _read_stored_fingerprint(workbook)
        comments_ok = _list2_status_match(workbook, matrix, rows)
        if stored == expected and comments_ok:
            return _sync_result(rows, wrote=False, changed=False)

        if stored != expected:
            if ROSTER_SHEET_TITLE in workbook.sheetnames:
                del workbook[ROSTER_SHEET_TITLE]
            sheet = workbook.create_sheet(ROSTER_SHEET_TITLE)
            _write_roster_sheet(sheet, rows)
        _annotate_manager_sheet(workbook, matrix, rows)
        buffer = BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
    finally:
        workbook.close()

    saved_path, error, used_fallback = _write_bytes_replace_or_stamp(path, data)
    if error:
        return _sync_result(rows, wrote=False, changed=True, error=error)
    return _sync_result(
        rows,
        wrote=True,
        changed=True,
        saved_path=str(saved_path) if saved_path is not None else None,
        used_fallback=used_fallback,
    )


def build_ds_roster_compare(
    matrix_path: str | Path | None = None,
    parts_dir: str | Path | None = None,
    *,
    sync_workbook: bool = True,
) -> DsRosterCompareResult:
    """Load the manager matrix, build roster rows, optionally rewrite the sheet.

    Args:
        matrix_path: Workbook path. Defaults to ``paths.ds_manager_matrix``.
        parts_dir: RFP parts folder. Defaults to ``DEFAULT_PARTS_DIR``.
        sync_workbook: When true and the matrix loaded, rewrite the roster
            sheet on mismatch (same as a large RFP run). Лист2 names/surnames
            are not rewritten.

    Returns:
        Snapshot with rows, optional sync outcome, and ``load_error``.
    """
    path = Path(matrix_path) if matrix_path else resolve_ds_manager_matrix_path()
    parts = Path(parts_dir) if parts_dir else Path(DEFAULT_PARTS_DIR)
    matrix = load_ds_manager_matrix(path)
    rows = build_ds_roster_rows(matrix, parts)
    sync: DsRosterSyncResult | None = None
    if matrix.load_error:
        sync = _sync_result(rows, wrote=False, changed=False, error=matrix.load_error)
    elif sync_workbook:
        sync = sync_ds_roster_sheet(matrix, parts)
    return DsRosterCompareResult(
        matrix_path=path,
        parts_dir=parts,
        rows=rows,
        load_error=matrix.load_error,
        sync=sync,
    )


def get_last_ds_roster_compare() -> DsRosterCompareResult | None:
    """Return the last result from ``run_ds_roster_compare_job``, if any."""
    return _last_roster


def _safe_out(msg: str) -> None:
    """Write one line to stdout; tolerate Windows console encodings."""
    text = msg + "\n"
    stream = sys.stdout
    try:
        stream.write(text)
    except UnicodeEncodeError:
        encoding = getattr(stream, "encoding", None) or "utf-8"
        payload = text.encode(encoding, errors="backslashreplace")
        buf = getattr(stream, "buffer", None)
        if buf is not None:
            buf.write(payload)
        else:
            stream.write(payload.decode(encoding, errors="replace"))


def _print_roster(result: DsRosterCompareResult) -> None:
    """Write the summary and TSV table to stdout for the job monitor."""
    _safe_out(result.summary_line())
    _safe_out(f"Книга: {result.matrix_path}")
    _safe_out(f"RFP: {result.parts_dir}")
    if result.load_error:
        return
    _safe_out(
        "Имя ДС\tФактический ДС\tПорядковый ДС\tФамилия\tСтатус\tБыло на Лист2\tДубликат\tЧто делать"
    )
    for row in result.rows:
        _safe_out(
            f"{row.copy_name}\t{row.actual_label}\t{row.sequential_label}\t"
            f"{row.manager or '—'}\t{row.status}\t{row.was_on_sheet1 or '—'}\t"
            f"{row.duplicate_label or '—'}\t{row.note or '—'}"
        )


def run_ds_roster_compare_job(
    matrix_path: str | Path | None = None,
    parts_dir: str | Path | None = None,
) -> DsRosterJobResult:
    """Compare surnames vs RFP files, sync the roster sheet, print the table.

    Args:
        matrix_path: Optional workbook override.
        parts_dir: Optional RFP parts folder override.

    Returns:
        Job result. ``success`` is false on matrix load or write errors.
        Missing surnames still succeed so the GUI table can be shown.
        Does not raise.
    """
    global _last_roster
    os.environ.setdefault("PYTHONUTF8", "1")
    result = build_ds_roster_compare(matrix_path, parts_dir, sync_workbook=True)
    _last_roster = result
    _print_roster(result)
    write_error = result.sync.write_error if result.sync is not None else None
    saved = None
    if result.sync is not None and result.sync.saved_path:
        saved = result.sync.saved_path
    elif result.load_error is None:
        saved = str(result.matrix_path)
    return DsRosterJobResult(
        success=result.load_error is None and write_error is None,
        message=result.summary_line(),
        result_path=saved,
    )
