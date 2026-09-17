"""Load DS manager matrix xlsx and apply manager / source path columns in Step4."""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from RFQ.rfp_parts.ds_checklist import (
    matched_folder_keys,
    parse_ds_name_from_file_name,
    present_in,
    table_ds_to_file_keys,
)
from RFQ.rfp_parts.ds_identity import parse_rfp_ds_identity
from RFQ.tags_rfp_compare.step4.step4_packing_compare import is_packing_only_sequential
from base.base_classes import CheckElement, RowType
from base.tables_columns import (
    DS_ACTUAL,
    DS_MANAGER,
    DS_NAME,
    DS_TITLE,
    PATH_MTO,
    PATH_RFP,
)
from utils.colors import Color

_DS_HEADERS = frozenset({"№ ДС", "Имя ДС"})
_MP_HEADERS = frozenset({"МП", "Фамилия"})
_SKIP_MANAGER_SHEETS = frozenset({"Фактический и порядковый", "Формат заполнения"})
_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
_ZERO_PAD_RE = re.compile(r"^(ДС)0+(\d.*)$", re.IGNORECASE)
_LETTER_SUFFIX_RE = re.compile(r"^(ДС\d+(?:_\d+)?)[A-ZА-ЯЁ]+$", re.IGNORECASE)
# ``ДС1_ГОСФИН_AGCC`` / ``ДС13_ПРИЛОЖЕНИЕ`` / ``ДС13_`` → ``ДС1`` / ``ДС13``.
# Do not strip numeric corrections ``ДС92_24Б``.
_NONNUMERIC_TAIL_RE = re.compile(r"^(ДС\d+)(?:_[^\d].*|_+)$", re.IGNORECASE)


def _norm_ds_cell(value: object) -> str:
    """Normalize a DS number cell (same rules as ds_checklist)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


@dataclass(frozen=True)
class DsManagerEntry:
    """One DS row from the manager matrix workbook."""

    raw_ds: str
    keys: tuple[str, ...]
    manager: str
    source_row: int = 0


@dataclass
class DsManagerMatrix:
    """Parsed manager matrix and key lookup index."""

    path: Path
    entries: list[DsManagerEntry]
    by_key: dict[str, DsManagerEntry]
    load_error: str | None = None
    sheet_title: str = ""
    ds_col: int = 0
    mp_col: int = 0


@dataclass(frozen=True)
class DsManagerPresenceResult:
    """Compare matrix DS rows to RFP part files in a folder."""

    matrix_count: int
    folder_ds_count: int
    missing_file: list[str]
    extra_files: list[str]
    ok_count: int
    summary_line: str
    load_error: str | None = None


def _zero_stripped_ds_key(key: str) -> str:
    match = _ZERO_PAD_RE.match(key)
    if not match:
        return key
    return f"{match.group(1).upper()}{match.group(2)}"


def _letter_suffix_stripped_ds_key(key: str) -> str:
    match = _LETTER_SUFFIX_RE.match(key)
    if not match:
        return key
    return match.group(1).upper()


def _nonnumeric_tail_stripped_ds_key(key: str) -> str:
    match = _NONNUMERIC_TAIL_RE.match(key)
    if not match:
        return key
    return match.group(1).upper()


def ds_key_aliases(ds_name: str) -> tuple[str, ...]:
    """Return lookup keys for a DS label (zero-pad and trailing-letter variants)."""
    aliases: list[str] = []
    seen: set[str] = set()
    for base in table_ds_to_file_keys(ds_name):
        for variant in (
            base,
            _zero_stripped_ds_key(base),
            _letter_suffix_stripped_ds_key(base),
            _letter_suffix_stripped_ds_key(_zero_stripped_ds_key(base)),
            _nonnumeric_tail_stripped_ds_key(base),
            _nonnumeric_tail_stripped_ds_key(_zero_stripped_ds_key(base)),
        ):
            text = str(variant or "").strip()
            if not text or text in seen:
                continue
            seen.add(text)
            aliases.append(text)
    return tuple(aliases)


def _find_header_columns(header_row: tuple[object, ...]) -> tuple[int | None, int | None]:
    ds_col: int | None = None
    mp_col: int | None = None
    for idx, value in enumerate(header_row):
        text = str(value or "").strip()
        if text in _DS_HEADERS:
            ds_col = idx
        elif text in _MP_HEADERS:
            mp_col = idx
    return ds_col, mp_col


def _select_manager_sheet(workbook) -> tuple[object, int, int] | tuple[None, None, None]:
    """Pick the compact DS/manager sheet; skip the position dump (Лист1)."""
    for name in workbook.sheetnames:
        if name in _SKIP_MANAGER_SHEETS:
            continue
        worksheet = workbook[name]
        header_row = next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
            None,
        )
        if not header_row:
            continue
        ds_col, mp_col = _find_header_columns(header_row)
        if ds_col is not None and mp_col is not None:
            return worksheet, ds_col, mp_col
    return None, None, None


def _empty_matrix(path: Path, error: str) -> DsManagerMatrix:
    return DsManagerMatrix(path=path, entries=[], by_key={}, load_error=error)


def load_ds_manager_matrix(path: str | Path) -> DsManagerMatrix:
    """Load the DS manager matrix from an xlsx workbook.

    Args:
        path: Path to the manager matrix workbook (local or UNC).

    Returns:
        Parsed matrix with ``by_key`` index. On missing file or read errors,
        ``load_error`` is set and ``entries`` is empty (no exception raised).
    """
    matrix_path = Path(path)
    if not matrix_path.exists():
        return _empty_matrix(matrix_path, f"file not found: {matrix_path}")

    try:
        payload = matrix_path.read_bytes()
    except OSError as exc:
        return _empty_matrix(matrix_path, str(exc))

    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — return load_error per contract
        return _empty_matrix(matrix_path, str(exc))

    entries: list[DsManagerEntry] = []
    by_key: dict[str, DsManagerEntry] = {}

    try:
        sheet, ds_col, mp_col = _select_manager_sheet(workbook)
        if sheet is None or ds_col is None or mp_col is None:
            return _empty_matrix(
                matrix_path,
                "missing required headers "
                f"{sorted(_DS_HEADERS)!r} / {sorted(_MP_HEADERS)!r}",
            )

        seen_keys: set[str] = set()
        for excel_row, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            raw_ds = _norm_ds_cell(row[ds_col] if ds_col < len(row) else None)
            if not raw_ds:
                continue

            manager_value = row[mp_col] if mp_col < len(row) else None
            manager = str(manager_value or "").strip()

            file_keys = ds_key_aliases(raw_ds)
            entry = DsManagerEntry(
                raw_ds=raw_ds,
                keys=file_keys,
                manager=manager,
                source_row=excel_row,
            )
            entries.append(entry)

            for key in file_keys:
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                by_key.setdefault(key, entry)
    finally:
        workbook.close()

    return DsManagerMatrix(
        path=matrix_path,
        entries=entries,
        by_key=by_key,
        sheet_title=str(getattr(sheet, "title", "") or ""),
        ds_col=int(ds_col),
        mp_col=int(mp_col),
    )


def lookup_manager(matrix: DsManagerMatrix, ds_name: str) -> DsManagerEntry | None:
    """Return the matrix entry for a DS name or file key.

    Args:
        matrix: Loaded manager matrix.
        ds_name: DS label or file key (e.g. ``ДС10``, ``ДС01``, ``47/13А``).

    Returns:
        Matching entry, or ``None`` when not found in the matrix.
    """
    for key in ds_key_aliases(ds_name):
        entry = matrix.by_key.get(key)
        if entry is not None:
            return entry
    return None


def collect_parts_ds_paths(parts_dir: str | Path) -> dict[str, list[str]]:
    """Collect RFP part workbook paths keyed by parsed DS name (upper-case).

    Args:
        parts_dir: Folder with RFP part workbooks (non-recursive scan).

    Returns:
        Map ``UPPER_DS_KEY -> [path, ...]`` in discovery order. Missing
        directory yields an empty dict.
    """
    folder = Path(parts_dir)
    if not folder.exists():
        return {}

    found: dict[str, list[str]] = {}
    for path in folder.iterdir():
        if not path.is_file() or path.suffix.lower() not in _EXCEL_SUFFIXES:
            continue
        ds_key = parse_ds_name_from_file_name(path.name)
        if not ds_key:
            continue
        upper_key = ds_key.upper()
        found.setdefault(upper_key, []).append(str(path))
    return found


def compare_matrix_to_parts(
    matrix: DsManagerMatrix,
    parts_dir: str | Path,
) -> DsManagerPresenceResult:
    """Compare matrix DS rows to files in the RFP parts folder.

    Args:
        matrix: Loaded manager matrix.
        parts_dir: Folder with RFP part workbooks.

    Returns:
        Presence summary with Russian ``summary_line`` for milestones / console.
    """
    parts_paths = collect_parts_ds_paths(parts_dir)
    folder_ds = {key: key for key in parts_paths}
    folder_ds_count = len(parts_paths)

    if matrix.load_error:
        return DsManagerPresenceResult(
            matrix_count=len(matrix.entries),
            folder_ds_count=folder_ds_count,
            missing_file=[],
            extra_files=[],
            ok_count=0,
            summary_line=f"Матрица МП: ошибка — {matrix.load_error}",
            load_error=matrix.load_error,
        )

    missing_file: list[str] = []
    covered: set[str] = set()
    ok_count = 0

    for entry in matrix.entries:
        matched: set[str] = set()
        for key in entry.keys or ds_key_aliases(entry.raw_ds):
            matched |= matched_folder_keys(folder_ds, key)
            for folder_key in parts_paths:
                if folder_key not in matched and present_in(
                    {folder_key: folder_key}, key
                ):
                    matched.add(folder_key)
        if matched:
            ok_count += 1
            covered |= matched
        else:
            missing_file.append(entry.raw_ds)

    extra_files = sorted(key for key in parts_paths if key not in covered)

    summary_line = (
        f"Матрица МП: ДС={len(matrix.entries)}, "
        f"без файла={len(missing_file)}, "
        f"файлы без матрицы={len(extra_files)}"
    )

    return DsManagerPresenceResult(
        matrix_count=len(matrix.entries),
        folder_ds_count=folder_ds_count,
        missing_file=missing_file,
        extra_files=extra_files,
        ok_count=ok_count,
        summary_line=summary_line,
    )


def build_mto_paths_by_title(mto_data: dict) -> dict[str, list[str]]:
    """Build unique MTO source file paths grouped by title_system key.

    Args:
        mto_data: ``title_system -> list[RowStd]`` from Step2 load.

    Returns:
        Map ``title -> [file_full_path, ...]`` with stable first-seen order.
    """
    result: dict[str, list[str]] = {}
    for title, rows in mto_data.items():
        paths: list[str] = []
        seen: set[str] = set()
        for row in rows:
            t_com = getattr(row, "t_com", None)
            path = str(getattr(t_com, "file_full_path", "") or "").strip()
            if path and path not in seen:
                seen.add(path)
                paths.append(path)
        if paths:
            result[str(title)] = paths
    return result


def _ensure_column(row: object, column: str) -> None:
    el = getattr(row, "el", None)
    if el is None:
        return
    if column not in el:
        el[column] = CheckElement(None)


def _set_cell(
    row: object,
    column: str,
    value: object,
    *,
    color: str = Color.no,
    comment: str = "",
) -> None:
    _ensure_column(row, column)
    cell = row.el[column]
    cell.value = value
    Color.set_el_color(cell, color)
    cell.comment = comment


def _row_text(row: object, column: str) -> str:
    raw = getattr(row, "el", {}).get(column) if getattr(row, "el", None) else None
    return str(getattr(raw, "value", raw) or "").strip()


def ds_actual_number(label: str) -> int | None:
    """Return the actual DS number from a label, or None.

    ``ДС92_24Б`` is treated as actual 92 (same retry as RFP identity).
    """
    text = str(label or "").strip()
    if not text:
        return None
    ident = parse_rfp_ds_identity(text)
    if ident.kind != "compound":
        retry = parse_rfp_ds_identity(text + ". ")
        if retry.compound:
            return retry.actual
    return ident.actual


def _unique_manager_name(names: set[str]) -> str:
    cleaned = {name.strip() for name in names if str(name or "").strip()}
    if len(cleaned) != 1:
        return ""
    return next(iter(cleaned))


def _matrix_managers_for_actual(
    matrix: DsManagerMatrix | None,
    actual: int,
) -> set[str]:
    if matrix is None:
        return set()
    found: set[str] = set()
    for entry in matrix.entries:
        entry_actual = ds_actual_number(entry.raw_ds)
        if entry_actual != actual:
            continue
        manager = str(entry.manager or "").strip()
        if manager:
            found.add(manager)
    return found


def _fill_leftover_ds_managers(
    rows: list,
    matrix: DsManagerMatrix | None,
) -> None:
    """Fill leftover ``DS_MANAGER`` when the actual DS has one unique surname.

    Sibling RFP rows with the same actual number win. If there are no
    sibling surnames, use a unique matrix surname for that actual.
    Several different surnames leave the leftover empty (no yellow).
    """
    sibling_by_actual: dict[int, set[str]] = {}
    leftovers: list[tuple[object, int | None]] = []
    for row in rows:
        row_type = getattr(row, "row_type", None)
        if RowType is not None and row_type is not None and row_type != RowType.position_row:
            continue
        ds_name = _row_text(row, DS_NAME)
        ds_actual = _row_text(row, DS_ACTUAL)
        actual = ds_actual_number(ds_actual) or ds_actual_number(ds_name)
        if is_packing_only_sequential(ds_name):
            leftovers.append((row, actual))
            continue
        if actual is None:
            continue
        manager = _row_text(row, DS_MANAGER)
        if manager:
            sibling_by_actual.setdefault(actual, set()).add(manager)

    for row, actual in leftovers:
        if actual is None:
            continue
        sibling_names = sibling_by_actual.get(actual, set())
        if sibling_names:
            name = _unique_manager_name(sibling_names)
        else:
            name = _unique_manager_name(_matrix_managers_for_actual(matrix, actual))
        _set_cell(row, DS_MANAGER, name)


def _join_paths(paths: Sequence[str]) -> str:
    unique: list[str] = []
    seen: set[str] = set()
    for path in paths:
        text = str(path or "").strip()
        if text and text not in seen:
            seen.add(text)
            unique.append(text)
    return "; ".join(unique)


def _collect_rfp_paths(
    ds_name: str,
    rfp_paths_by_key: Mapping[str, Sequence[str]],
) -> list[str]:
    collected: list[str] = []
    seen: set[str] = set()
    for key in ds_key_aliases(ds_name):
        for path in rfp_paths_by_key.get(key, ()):
            text = str(path or "").strip()
            if text and text not in seen:
                seen.add(text)
                collected.append(text)
        for map_key, paths in rfp_paths_by_key.items():
            upper_map_key = str(map_key).upper()
            if upper_map_key == key or upper_map_key.startswith(key + "_"):
                for path in paths:
                    text = str(path or "").strip()
                    if text and text not in seen:
                        seen.add(text)
                        collected.append(text)
    return collected


def apply_ds_source_columns(
    rows: list,
    *,
    matrix: DsManagerMatrix | None,
    rfp_paths_by_key: Mapping[str, Sequence[str]],
    mto_paths_by_title: Mapping[str, Sequence[str]],
) -> None:
    """Fill DS manager and source path columns on result rows (in place).

    Leftover ``ДС{n}_RFP_не_найден`` rows get a surname only when that
    actual DS has one unique manager (from sibling RFP rows, else matrix).

    Args:
        rows: Step4 result rows (``RowStd`` or compatible objects with ``el``).
        matrix: Loaded manager matrix, or ``None`` to treat all DS as missing.
        rfp_paths_by_key: RFP part paths keyed by DS file key.
        mto_paths_by_title: MTO source paths keyed by ``DS_TITLE`` / title_system.
    """
    for row in rows:
        row_type = getattr(row, "row_type", None)
        if RowType is not None and row_type is not None and row_type != RowType.position_row:
            continue

        _ensure_column(row, DS_MANAGER)
        _ensure_column(row, PATH_RFP)
        _ensure_column(row, PATH_MTO)

        ds_name_raw = row.el.get(DS_NAME)
        ds_name = str(getattr(ds_name_raw, "value", ds_name_raw) or "").strip()
        ds_actual_raw = row.el.get(DS_ACTUAL)
        ds_actual = str(getattr(ds_actual_raw, "value", ds_actual_raw) or "").strip()
        leftover_seq = is_packing_only_sequential(ds_name)

        if leftover_seq:
            # Surname is filled in a second pass (unique actual → one MP).
            _set_cell(row, DS_MANAGER, "")
            path_lookup = ds_actual
        elif not ds_name and not ds_actual:
            _set_cell(row, DS_MANAGER, "")
            _set_cell(row, PATH_RFP, "")
            _set_cell(row, PATH_MTO, "")
            continue
        else:
            entry = None
            if matrix is not None:
                if ds_name:
                    entry = lookup_manager(matrix, ds_name)
                if entry is None and ds_actual:
                    entry = lookup_manager(matrix, ds_actual)
            if entry is None or not entry.manager:
                manager_comment = (
                    "Нет в матрице МП" if entry is None else "Пустой МП в матрице"
                )
                _set_cell(
                    row,
                    DS_MANAGER,
                    entry.manager if entry is not None else "",
                    color=Color.yellow,
                    comment=manager_comment,
                )
            else:
                _set_cell(row, DS_MANAGER, entry.manager)
            path_lookup = ds_name or ds_actual

        rfp_paths: list[str] = []
        if path_lookup:
            rfp_paths = _collect_rfp_paths(path_lookup, rfp_paths_by_key)
            if not rfp_paths and ds_actual and ds_actual != path_lookup:
                rfp_paths = _collect_rfp_paths(ds_actual, rfp_paths_by_key)
        if rfp_paths:
            _set_cell(row, PATH_RFP, _join_paths(rfp_paths))
        else:
            _set_cell(row, PATH_RFP, "", color=Color.yellow)

        ds_title_raw = row.el.get(DS_TITLE)
        ds_title = str(getattr(ds_title_raw, "value", ds_title_raw) or "").strip()
        mto_paths = list(mto_paths_by_title.get(ds_title, ()))
        _set_cell(row, PATH_MTO, _join_paths(mto_paths) if mto_paths else "")

    _fill_leftover_ds_managers(rows, matrix)
