"""
Корректное сравнение двух последних файлов Step4 по мультимножеству строк.
Не зависит от порядка строк и не использует неуникальные ключи.
"""

import argparse
import os
from collections import Counter
from datetime import datetime
from typing import Optional

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


FILE_PREFIX = "Шаг4_Сопоставление_RFP_MTO_"
FILE_SUFFIX = ".xlsx"
SHEET_NAME = "Сопоставление RFP и MTO"


def _find_step4_files(folder: str) -> list[str]:
    if not os.path.isdir(folder):
        return []
    files = [
        f for f in os.listdir(folder)
        if f.startswith(FILE_PREFIX) and f.endswith(FILE_SUFFIX)
    ]
    files.sort(reverse=True)
    return [os.path.join(folder, f) for f in files]


def _normalize_df(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    normalized = df[columns].fillna("").copy()
    for col in columns:
        normalized[col] = normalized[col].astype(str)
    return normalized


def _rows_as_tuples(df: pd.DataFrame) -> list[tuple]:
    return [tuple(row) for row in df.to_numpy()]


def _save_summary_excel(
    output_path: str,
    *,
    new_name: str,
    old_name: str,
    rows_new: int,
    rows_old: int,
    common_cols: int,
    only_new: int,
    only_old: int,
    result_label: str,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сравнение"

    headers = ["Параметр", "Значение"]
    ws.append(headers)
    rows = [
        ("Файл новый", new_name),
        ("Файл старый", old_name),
        ("Строк в новом", rows_new),
        ("Строк в старом", rows_old),
        ("Общих колонок", common_cols),
        ("Только в новом", only_new),
        ("Только в старом", only_old),
        ("Результат", result_label),
    ]
    for row in rows:
        ws.append(list(row))

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    align = Alignment(horizontal="left", vertical="center")
    for col in range(1, 3):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = align

    for row_idx in range(2, ws.max_row + 1):
        for col in range(1, 3):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = border
            cell.alignment = align

    ws.column_dimensions[get_column_letter(1)].width = 28
    ws.column_dimensions[get_column_letter(2)].width = 60
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{ws.max_row}"
    wb.save(output_path)


def compare_step4_files(folder: str, output_path: Optional[str] = None) -> Optional[str]:
    files = _find_step4_files(folder)
    if len(files) < 2:
        print(f"В папке найдено файлов с префиксом '{FILE_PREFIX}': {len(files)}. Нужно минимум 2.")
        return None

    new_path, old_path = files[0], files[1]
    new_name = os.path.basename(new_path)
    old_name = os.path.basename(old_path)
    print(f"Сравниваем:\n  1 (новый): {new_name}\n  2 (старый): {old_name}")

    df_new = pd.read_excel(new_path, sheet_name=SHEET_NAME)
    df_old = pd.read_excel(old_path, sheet_name=SHEET_NAME)

    common_cols = [c for c in df_new.columns if c in df_old.columns]
    if not common_cols:
        print("Ошибка: не найдено общих колонок для сравнения.")
        return None

    new_norm = _normalize_df(df_new, common_cols)
    old_norm = _normalize_df(df_old, common_cols)
    new_rows = _rows_as_tuples(new_norm)
    old_rows = _rows_as_tuples(old_norm)

    c_new = Counter(new_rows)
    c_old = Counter(old_rows)
    only_new = sum((c_new - c_old).values())
    only_old = sum((c_old - c_new).values())

    print(f"rows_new={len(new_rows)}")
    print(f"rows_old={len(old_rows)}")
    print(f"common_columns={len(common_cols)}")
    print(f"only_new={only_new}")
    print(f"only_old={only_old}")

    if output_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(folder, f"Шаг4_Сравнение_итог_{timestamp}.xlsx")

    result_label = "Идентично (multiset)" if only_new == 0 and only_old == 0 else "Есть различия (multiset)"
    _save_summary_excel(
        output_path,
        new_name=new_name,
        old_name=old_name,
        rows_new=len(new_rows),
        rows_old=len(old_rows),
        common_cols=len(common_cols),
        only_new=only_new,
        only_old=only_old,
        result_label=result_label,
    )

    print(f"Итог сравнения сохранен: {os.path.basename(output_path)}")
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Корректное сравнение двух последних Step4 файлов в папке (multiset)"
    )
    parser.add_argument("folder", type=str, help="Путь к папке с файлами")
    parser.add_argument(
        "-o",
        "--output",
        type=str,
        default=None,
        help="Путь к итоговому xlsx-файлу сводки",
    )
    args = parser.parse_args()

    folder = os.path.abspath(args.folder)
    print(f"Папка: {folder}")
    if not os.path.isdir(folder):
        print(f"Ошибка: папка не найдена: {folder}")
        return 1

    files = _find_step4_files(folder)
    print(f"Найдено файлов Шаг4_Сопоставление_RFP_MTO_*.xlsx: {len(files)}")
    for i, f in enumerate(files[:5], 1):
        print(f"  {i}. {os.path.basename(f)}")
    if len(files) > 5:
        print(f"  ... и ещё {len(files) - 5}")

    result = compare_step4_files(folder, args.output)
    return 0 if result else 1


if __name__ == "__main__":
    raise SystemExit(main())
