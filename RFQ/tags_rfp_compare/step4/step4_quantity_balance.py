"""Контроль сохранения сумм VALUES (RFP / MTO / VO) и количеств УЛ на входе step4 и перед Excel."""

from __future__ import annotations

import json
import os
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from prettytable import PrettyTable

from base.base_classes import RowStd, RowType
from base.tables_columns import (
    CODE_MTO,
    DS_TITLE,
    MATCH_STATUS,
    MATCH_STATUS_VO,
    UL_VALUES,
    VALUES,
    VALUES_MTO,
    VALUES_VO,
)
from RFQ.tags_rfp_compare.rfp_tags_utils import ensure_result_dir_exists
from RFQ.tags_rfp_compare.step4.step4_5_check_values_sum import (
    _calculate_result_mto_sums,
    _calculate_result_rfp_sums,
    _calculate_result_vo_sums,
)

EMPTY_TITLE_KEY = "<без title_system>"
_EPS = 0.01
_FATAL_MARKER = "quantity_balance_fatal.json"


class QuantityBalanceFatalError(RuntimeError):
    """Потеря количества: расхождение вход/выход больше допуска split (теги vs VALUES)."""

    def __init__(self, message: str, summary: str = "") -> None:
        super().__init__(message)
        self.summary = summary


def _float_dict() -> defaultdict[str, float]:
    return defaultdict(float)


def _int_dict() -> defaultdict[str, int]:
    return defaultdict(int)


@dataclass
class QuantityInputSnapshot:
    """Суммы VALUES на входе step4 (до split в worker)."""

    rfp: Dict[str, float] = field(default_factory=_float_dict)
    mto: Dict[str, float] = field(default_factory=_float_dict)
    vo: Dict[str, float] = field(default_factory=_float_dict)
    rfp_rows: Dict[str, int] = field(default_factory=_int_dict)
    mto_rows: Dict[str, int] = field(default_factory=_int_dict)
    vo_rows: Dict[str, int] = field(default_factory=_int_dict)


@dataclass
class QuantityBalanceRow:
    title_system: str
    part: str
    sum_in: float
    sum_out: float
    delta: float
    split_tolerance: float
    level: str  # ok | warn | error
    note: str = ""


@dataclass
class QuantityBalanceReport:
    phase: str
    rows: List[QuantityBalanceRow] = field(default_factory=list)
    warnings: List[QuantityBalanceRow] = field(default_factory=list)
    errors: List[QuantityBalanceRow] = field(default_factory=list)
    fatal: bool = False
    result_dir: str = ""
    excel_path: str = ""
    marker_path: str = ""
    summary: str = ""

    def emit_fatal(self) -> None:
        """Диалог и исключение (подробности уже в консоли из run_output_quantity_balance_check)."""
        show_quantity_balance_error_dialog(self._dialog_message())
        if self.fatal:
            raise QuantityBalanceFatalError(
                self._dialog_message(),
                summary=self.summary,
            )

    def _dialog_message(self) -> str:
        if self.phase == "ul_output":
            intro = (
                "Баланс количеств Step4 (УЛ): сумма UL_VALUES на выходе не равна "
                "валидным количествам упаковочных листов на входе."
            )
        else:
            intro = (
                "Баланс количеств Step4: расхождение вход/выход превышает допуск "
                "(несовпадения тегов/VALUES при раскрытии)."
            )
        lines = [
            intro,
            f"Критических строк: {len(self.errors)}.",
        ]
        for row in self.errors[:12]:
            lines.append(
                f"  {row.title_system} | {row.part}: вход={row.sum_in:.2f}, "
                f"выход={row.sum_out:.2f}, Δ={row.delta:+.2f}, допуск={row.split_tolerance:.2f}"
            )
        if len(self.errors) > 12:
            lines.append(f"  ... ещё {len(self.errors) - 12}")
        if self.excel_path:
            lines.append(f"\nПодробно: {self.excel_path}")
        return "\n".join(lines)


def _normalize_title(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _title_key(value) -> str:
    norm = _normalize_title(value)
    return norm if norm else EMPTY_TITLE_KEY


def _safe_float(value) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0.0


def sum_values_in_data_dict(
    data: Dict[str, List[RowStd]] | None,
) -> tuple[float, int]:
    """Сумма VALUES и число position_row (например MTO/VO после check_*_data в worker)."""
    total = 0.0
    count = 0
    for _title_mark, rows in (data or {}).items():
        for row in rows:
            if row.row_type != RowType.position_row:
                continue
            total += _safe_float(row.get_value(VALUES))
            count += 1
    return total, count


def apply_post_split_input_to_snapshot(
    snapshot: QuantityInputSnapshot,
    mto_by_title: Dict[str, float] | None = None,
    mto_rows_by_title: Dict[str, int] | None = None,
    vo_by_title: Dict[str, float] | None = None,
    vo_rows_by_title: Dict[str, int] | None = None,
) -> None:
    """Подменяет MTO/VO «вход» на суммы после split в worker (та же база, что и pipeline)."""
    if mto_by_title:
        for title, value in mto_by_title.items():
            snapshot.mto[_title_key(title)] = float(value)
    if mto_rows_by_title:
        for title, value in mto_rows_by_title.items():
            snapshot.mto_rows[_title_key(title)] = int(value)
    if vo_by_title:
        for title, value in vo_by_title.items():
            snapshot.vo[_title_key(title)] = float(value)
    if vo_rows_by_title:
        for title, value in vo_rows_by_title.items():
            snapshot.vo_rows[_title_key(title)] = int(value)


def build_input_snapshot(
    rfp_data: List[RowStd],
    mto_data: Dict[str, List[RowStd]],
    vo_data: Dict[str, List[RowStd]],
) -> QuantityInputSnapshot:
    """Суммы VALUES по title_system на входе step4 (сырые данные до split в worker)."""
    snap = QuantityInputSnapshot()
    for row in rfp_data:
        if row.row_type != RowType.position_row:
            continue
        key = _title_key(row.get_value(DS_TITLE))
        snap.rfp[key] += _safe_float(row.get_value(VALUES))
        snap.rfp_rows[key] += 1

    for title_mark, rows in (mto_data or {}).items():
        bucket = _title_key(title_mark)
        for row in rows:
            if row.row_type != RowType.position_row:
                continue
            snap.mto[bucket] += _safe_float(row.get_value(VALUES))
            snap.mto_rows[bucket] += 1

    for title_mark, rows in (vo_data or {}).items():
        bucket = _title_key(title_mark)
        for row in rows:
            if row.row_type != RowType.position_row:
                continue
            snap.vo[bucket] += _safe_float(row.get_value(VALUES))
            snap.vo_rows[bucket] += 1

    return snap


def build_output_snapshot(result_rows: List[RowStd]) -> Tuple[Dict[str, float], Dict[str, float], Dict[str, float]]:
    """Суммы RFP / MTO / VO в result_rows перед записью Excel."""
    return (
        _calculate_result_rfp_sums(result_rows),
        _calculate_result_mto_sums(result_rows),
        _calculate_result_vo_sums(result_rows),
    )


def build_split_tolerance_by_title(count_mismatches: List[dict]) -> Dict[str, float]:
    """Допуск по title_system: сумма |tags_count - values| из таблиц несовпадений split."""
    tol: Dict[str, float] = defaultdict(float)
    for item in count_mismatches or []:
        title = _title_key(item.get("title_system"))
        try:
            tags_n = int(item.get("tags_count", 0))
            values_n = int(item.get("values", 0))
        except (TypeError, ValueError):
            continue
        tol[title] += abs(tags_n - values_n)
    return dict(tol)


def _total_sum(sums: Dict[str, float]) -> float:
    return float(sum(sums.values()))


def _total_rows(row_counts: Dict[str, int]) -> int:
    return int(sum(row_counts.values()))


def _format_in_out_line(
    label: str,
    sum_in: float,
    sum_out: float,
    rows_in: int = 0,
    rows_out: int = 0,
) -> str:
    delta = sum_out - sum_in
    if rows_in or rows_out:
        return (
            f"{label} {sum_in:.2f}→{sum_out:.2f} (Δ{delta:+.2f}, "
            f"{rows_in}→{rows_out} стр.)"
        )
    return f"{label} {sum_in:.2f}→{sum_out:.2f} (Δ{delta:+.2f})"


def format_input_rfp_mto_vo_detail(snapshot: QuantityInputSnapshot) -> str:
    """Brief GUI/milestone line for Step4 RFP/MTO/VO input totals."""
    return (
        f"вход RFP {_total_sum(snapshot.rfp):.2f} | "
        f"MTO {_total_sum(snapshot.mto):.2f} | "
        f"VO {_total_sum(snapshot.vo):.2f}"
    )


def format_input_ul_detail(*, available: bool, qty: float = 0.0, rows: int = 0) -> str:
    """Brief GUI/milestone line for packing-list input totals."""
    if not available:
        return "вход УЛ недоступны"
    return f"вход УЛ {qty:.2f} ({rows} стр.)"


def print_input_quantity_balance(snapshot: QuantityInputSnapshot) -> None:
    """Краткая сводка сумм VALUES на входе step4 (без таблицы по title_system)."""
    print(
        "Баланс Step4 — вход (RFP сырой; MTO/VO уточняются после split в worker): "
        f"RFP {_total_sum(snapshot.rfp):.2f} ({_total_rows(snapshot.rfp_rows)} стр.) | "
        f"MTO {_total_sum(snapshot.mto):.2f} ({_total_rows(snapshot.mto_rows)} стр.) | "
        f"VO {_total_sum(snapshot.vo):.2f} ({_total_rows(snapshot.vo_rows)} стр.)"
    )


def run_output_quantity_balance_check(
    input_snapshot: QuantityInputSnapshot,
    result_rows: List[RowStd],
    count_mismatches: List[dict],
    result_dir: str,
) -> QuantityBalanceReport:
    """Сравнение входных и выходных сумм перед save_match_result_to_excel."""
    rfp_out, mto_out, vo_out = build_output_snapshot(result_rows)
    split_tol = build_split_tolerance_by_title(count_mismatches)
    report = QuantityBalanceReport(phase="output", result_dir=result_dir or "")

    all_titles = sorted(
        set(input_snapshot.rfp)
        | set(input_snapshot.mto)
        | set(input_snapshot.vo)
        | set(rfp_out)
        | set(mto_out)
        | set(vo_out),
        key=lambda t: (t == EMPTY_TITLE_KEY, t),
    )

    for title in all_titles:
        tol = split_tol.get(title, 0.0)
        for part, s_in, s_out in (
            ("RFP", input_snapshot.rfp.get(title, 0.0), rfp_out.get(title, 0.0)),
            ("MTO", input_snapshot.mto.get(title, 0.0), mto_out.get(title, 0.0)),
            ("VO", input_snapshot.vo.get(title, 0.0), vo_out.get(title, 0.0)),
        ):
            if s_in == 0.0 and s_out == 0.0:
                continue
            delta = s_out - s_in
            abs_delta = abs(delta)
            if abs_delta <= _EPS:
                level = "ok"
                note = ""
            elif abs_delta <= tol + _EPS:
                level = "warn"
                note = f"в допуске split (|Δ|≤{tol:.0f})"
            else:
                level = "error"
                note = f"превышен допуск split ({tol:.0f})"
            row = QuantityBalanceRow(
                title_system=title,
                part=part,
                sum_in=s_in,
                sum_out=s_out,
                delta=delta,
                split_tolerance=tol,
                level=level,
                note=note,
            )
            report.rows.append(row)
            if level == "warn":
                report.warnings.append(row)
            elif level == "error":
                report.errors.append(row)

    rfp_in = _total_sum(input_snapshot.rfp)
    mto_in = _total_sum(input_snapshot.mto)
    vo_in = _total_sum(input_snapshot.vo)
    rfp_out_t = _total_sum(rfp_out)
    mto_out_t = _total_sum(mto_out)
    vo_out_t = _total_sum(vo_out)
    in_out = (
        _format_in_out_line("RFP", rfp_in, rfp_out_t)
        + " | "
        + _format_in_out_line("MTO", mto_in, mto_out_t)
        + " | "
        + _format_in_out_line("VO", vo_in, vo_out_t)
    )

    if report.errors:
        report.fatal = True
        report.summary = f"ошибка {in_out}"
        if result_dir:
            report.excel_path = _save_balance_excel(report, result_dir) or ""
            report.marker_path = _write_fatal_marker(report, result_dir) or ""
        print("\n" + "=" * 80)
        print("Баланс Step4 — ОШИБКА: вход ≠ выход перед Excel")
        print("=" * 80)
        print("Сводка: " + in_out)
        print(
            f"Критических: {len(report.errors)}, "
            f"предупреждений (допуск split): {len(report.warnings)}"
        )
        _print_balance_compare_table(report.errors + report.warnings)
        if report.excel_path:
            print(f"Таблица -> {os.path.basename(report.excel_path)}")
    elif report.warnings:
        report.summary = f"OK split {in_out}"
        print(
            "Баланс Step4 — выход: OK с предупреждениями "
            f"({len(report.warnings)} title_system, допуск split). "
            + in_out
        )
    else:
        report.summary = f"OK {in_out}"
        print("Баланс Step4 — выход: OK. " + in_out)

    return report


def sum_ul_values_by_title(result_rows: List[RowStd]) -> Tuple[float, int, Dict[str, float]]:
    """Sum numeric ``UL_VALUES`` on position rows, grouped by ``DS_TITLE``.

    Empty / ``None`` cells are skipped (invalid packing leftovers, unavailable
    paint, unmatched RFP). Zero is counted.

    Returns:
        Total quantity, number of rows with a numeric UL_VALUES, per-title sums.
    """
    by_title: Dict[str, float] = defaultdict(float)
    total = 0.0
    count = 0
    for row in result_rows:
        if row.row_type != RowType.position_row:
            continue
        raw = row.get_value(UL_VALUES)
        if raw in (None, ""):
            continue
        qty = _safe_float(raw)
        key = _title_key(row.get_value(DS_TITLE))
        by_title[key] += qty
        total += qty
        count += 1
    return total, count, dict(by_title)


def print_input_ul_quantity_balance(*, available: bool, qty: float = 0.0, rows: int = 0) -> None:
    """Print UL input totals after the packing cache is loaded (and filtered)."""
    if not available:
        print("Баланс Step4 — вход УЛ: недоступны (проверка пропущена)")
        return
    print(f"Баланс Step4 — вход УЛ: {qty:.2f} ({rows} стр.)")


def run_output_ul_quantity_balance_check(
    qty_in: float,
    by_title_in: Dict[str, float],
    result_rows: List[RowStd],
    result_dir: str,
    *,
    available: bool,
) -> QuantityBalanceReport:
    """Compare queued packing quantities with sum of numeric UL_VALUES before Excel.

    УЛ is not part of the RFP/MTO/VO check. Unavailable cache skips the compare
    (Excel is still written with red «УЛ недоступны»). No split tolerance.
    """
    report = QuantityBalanceReport(phase="ul_output", result_dir=result_dir or "")
    if not available:
        report.summary = "УЛ недоступны"
        print("Баланс Step4 — выход УЛ: пропущен (УЛ недоступны)")
        return report

    qty_out, _rows_out, by_title_out = sum_ul_values_by_title(result_rows)
    all_titles = sorted(
        set(by_title_in) | set(by_title_out),
        key=lambda t: (t == EMPTY_TITLE_KEY, t),
    )
    for title in all_titles:
        s_in = by_title_in.get(title, 0.0)
        s_out = by_title_out.get(title, 0.0)
        if s_in == 0.0 and s_out == 0.0:
            continue
        delta = s_out - s_in
        abs_delta = abs(delta)
        if abs_delta <= _EPS:
            level = "ok"
            note = ""
        else:
            level = "error"
            note = "вход УЛ (очереди) ≠ сумма UL_VALUES"
        row = QuantityBalanceRow(
            title_system=title,
            part="УЛ",
            sum_in=s_in,
            sum_out=s_out,
            delta=delta,
            split_tolerance=0.0,
            level=level,
            note=note,
        )
        report.rows.append(row)
        if level == "error":
            report.errors.append(row)

    in_out = _format_in_out_line("УЛ", qty_in, qty_out)
    if report.errors:
        report.fatal = True
        report.summary = f"ошибка {in_out}"
        if result_dir:
            report.excel_path = _save_balance_excel(report, result_dir) or ""
            report.marker_path = _write_fatal_marker(report, result_dir) or ""
        print("\n" + "=" * 80)
        print("Баланс Step4 — ОШИБКА УЛ: вход ≠ выход перед Excel")
        print("=" * 80)
        print("Сводка: " + in_out)
        print(f"Критических: {len(report.errors)}")
        _print_balance_compare_table(report.errors)
        if report.excel_path:
            print(f"Таблица -> {os.path.basename(report.excel_path)}")
    else:
        report.summary = f"OK {in_out}"
        print("Баланс Step4 — выход УЛ: OK. " + in_out)
    return report


def show_quantity_balance_error_dialog(message: str, parent=None) -> None:
    try:
        import tkinter
        from tkinter import messagebox

        root = parent
        if root is None:
            root = tkinter.Tk()
            root.withdraw()
        messagebox.showerror("Баланс количеств Step4", message, parent=root)
        if parent is None:
            root.destroy()
    except Exception:
        print(message)


def find_latest_fatal_result_dir(result_dir_base: str) -> Optional[str]:
    """Последняя папка ``_результат_проверки_*`` с ``quantity_balance_fatal.json``."""
    if not result_dir_base or not os.path.isdir(result_dir_base):
        return None
    best_mtime = -1.0
    best_path: Optional[str] = None
    try:
        for name in os.listdir(result_dir_base):
            if not name.startswith("_результат_проверки_"):
                continue
            path = os.path.join(result_dir_base, name)
            if not os.path.isdir(path):
                continue
            if not os.path.isfile(os.path.join(path, _FATAL_MARKER)):
                continue
            mtime = os.path.getmtime(path)
            if mtime > best_mtime:
                best_mtime = mtime
                best_path = path
    except OSError:
        return None
    return best_path


def try_show_fatal_dialog_for_result_dir(result_dir: str) -> bool:
    """Читает marker после subprocess с ненулевым кодом. Возвращает True, если показан диалог."""
    if not result_dir:
        return False
    marker = os.path.join(result_dir, _FATAL_MARKER)
    if not os.path.isfile(marker):
        return False
    try:
        with open(marker, "r", encoding="utf-8") as f:
            data = json.load(f)
        show_quantity_balance_error_dialog(data.get("message", "Ошибка баланса количеств Step4."))
        return True
    except Exception:
        return False


def _print_balance_compare_table(rows: List[QuantityBalanceRow]) -> None:
    if not rows:
        return
    show = [r for r in rows if r.level != "ok"]
    if not show:
        print("\nСверка вход/выход: все title_system в пределах 0.01")
        return
    table = PrettyTable()
    table.field_names = [
        "title_system",
        "часть",
        "вход",
        "выход",
        "Δ",
        "допуск split",
        "уровень",
        "примечание",
    ]
    table.align = "l"
    for col in ("вход", "выход", "Δ", "допуск split"):
        table.align[col] = "r"
    for row in show:
        table.add_row(
            [
                row.title_system,
                row.part,
                f"{row.sum_in:.2f}",
                f"{row.sum_out:.2f}",
                f"{row.delta:+.2f}",
                f"{row.split_tolerance:.0f}",
                row.level.upper(),
                row.note,
            ],
        )
    print("\nСверка вход → выход (перед Excel):")
    print(table)


def _save_balance_excel(report: QuantityBalanceReport, result_dir: str) -> Optional[str]:
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        return None

    ensure_result_dir_exists(result_dir)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(result_dir, f"Шаг4_Баланс_количества_ошибки_{ts}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "balance_errors"
    headers = [
        "title_system",
        "часть",
        "сумма_вход",
        "сумма_выход",
        "дельта",
        "допуск_split",
        "уровень",
        "примечание",
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")
    for row in report.errors + report.warnings:
        ws.append(
            [
                row.title_system,
                row.part,
                row.sum_in,
                row.sum_out,
                row.delta,
                row.split_tolerance,
                row.level,
                row.note,
            ],
        )
    wb.save(path)
    return path


def _write_fatal_marker(report: QuantityBalanceReport, result_dir: str) -> Optional[str]:
    ensure_result_dir_exists(result_dir)
    path = os.path.join(result_dir, _FATAL_MARKER)
    payload = {
        "message": report._dialog_message(),
        "error_count": len(report.errors),
        "warning_count": len(report.warnings),
        "excel": report.excel_path,
        "errors": [
            {
                "title_system": r.title_system,
                "part": r.part,
                "sum_in": r.sum_in,
                "sum_out": r.sum_out,
                "delta": r.delta,
                "split_tolerance": r.split_tolerance,
            }
            for r in report.errors
        ],
    }
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        return path
    except OSError:
        return None
