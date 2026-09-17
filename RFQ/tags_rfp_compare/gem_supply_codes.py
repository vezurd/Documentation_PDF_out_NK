"""Load GEM-supply BCC codes and overlay UL status on Step4 rows."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from RFQ.packing_list_provider import normalize_packing_key_part
from RFQ.tags_rfp_compare.step4.step4_packing_compare import (
    STATUS_GEM_SUPPLY,
    STATUS_MTO_ONLY,
    STATUS_OPEN_UPD,
    parse_composite_title,
)
from base.base_classes import CheckElement, RowStd, RowType
from base.tables_columns import CODE, CODE_MTO, DS_TITLE, IN_CABINET, UL_COMPARE_STATUS
from utils.colors import Color

GEM_SUPPLY_TITLE = "8950"
_CODE_HEADER_TOKENS = ("код", "code", "bcc")
_GEM_OVERLAY_STATUSES = frozenset({STATUS_OPEN_UPD, STATUS_MTO_ONLY})


@dataclass
class GemSupplyCodes:
    """Parsed GEM-supply workbook: normalized codes and optional load error."""

    path: Path
    codes: frozenset[str]
    load_error: str | None = None


def _empty_codes(path: Path, error: str) -> GemSupplyCodes:
    return GemSupplyCodes(path=path, codes=frozenset(), load_error=error)


def _detect_code_column(header_row: tuple[object, ...]) -> int:
    for idx, value in enumerate(header_row):
        text = str(value or "").strip().lower()
        if any(token in text for token in _CODE_HEADER_TOKENS):
            return idx
    return 0


def load_gem_supply_codes(path: str | Path) -> GemSupplyCodes:
    """Load unique packing-normalized codes from the first worksheet of an xlsx.

    Row 1 is treated as a header. The code column is the first header cell
    whose stripped lowercased text contains ``код``, ``code``, or ``bcc``;
    otherwise column 0 is used.

    Args:
        path: Path to the GEM-supply workbook (local or UNC).

    Returns:
        Parsed codes as a ``frozenset``. On empty path, missing file, read
        ``OSError``, or workbook parse errors, ``load_error`` is set and
        ``codes`` is empty (no exception raised).
    """
    if isinstance(path, str) and not path.strip():
        return _empty_codes(Path(path), "empty path")

    gem_path = Path(path)
    try:
        if not gem_path.exists():
            return _empty_codes(gem_path, f"file not found: {gem_path}")
        payload = gem_path.read_bytes()
    except OSError as exc:
        return _empty_codes(gem_path, str(exc))

    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — return load_error per contract
        return _empty_codes(gem_path, str(exc))

    try:
        if not workbook.worksheets:
            return _empty_codes(gem_path, "workbook has no worksheets")

        worksheet = workbook.worksheets[0]
        rows_iter = worksheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return GemSupplyCodes(path=gem_path, codes=frozenset())

        code_col = _detect_code_column(header_row)
        collected: set[str] = set()
        for row in rows_iter:
            cell = row[code_col] if code_col < len(row) else None
            normalized = normalize_packing_key_part(cell)
            if not normalized:
                continue
            collected.add(normalized)
        return GemSupplyCodes(path=gem_path, codes=frozenset(collected))
    except Exception as exc:  # noqa: BLE001 — return load_error per contract
        return _empty_codes(gem_path, str(exc))
    finally:
        workbook.close()


def _paint_row_match_matched(row: RowStd) -> None:
    """Fill every CheckElement on *row* with ``Color.match_matched``."""
    for el in row.el.values():
        if isinstance(el, CheckElement):
            el.color = Color.match_matched


def paint_gem_supply_rows(rows: list[RowStd]) -> int:
    """Paint the full Excel row for every GEM-supply position.

    Call after later colouring steps (MTO−УЛ diff, manager/path columns)
    so they cannot leave unpainted cells on a GEM row.

    Args:
        rows: Step4 result rows to scan.

    Returns:
        Number of GEM-supply rows that were painted.
    """
    painted = 0
    for row in rows:
        if row.row_type != RowType.position_row:
            continue
        status = str(row.get_value(UL_COMPARE_STATUS) or "").strip()
        if status != STATUS_GEM_SUPPLY:
            continue
        _paint_row_match_matched(row)
        painted += 1
    return painted


def _is_in_cabinet_position(row: RowStd) -> bool:
    """True if the row belongs to a cabinet (non-empty ``IN_CABINET``)."""
    return bool(str(row.get_value(IN_CABINET) or "").strip())


def apply_gem_supply_status(
    rows: list[RowStd],
    gem: GemSupplyCodes | None,
    *,
    title: str = GEM_SUPPLY_TITLE,
) -> int:
    """Replace eligible UL statuses with GEM supply for matching 8950 rows.

    Mutates ``UL_COMPARE_STATUS`` (value + comment left as-is) and paints
    the whole row ``Color.match_matched`` (``#C6EFCE``), not only the
    status cell. Quantity values are left unchanged. Rows are updated only
    when they are position rows, currently ``STATUS_OPEN_UPD`` or
    ``STATUS_MTO_ONLY``, ``IN_CABINET`` is empty, the composite
    ``DS_TITLE`` title part matches ``title`` (default ``8950``), and
    ``CODE`` or ``CODE_MTO`` is in ``gem.codes``. Cabinet positions
    (non-empty ``IN_CABINET``) keep the original UL status and colour.

    Args:
        rows: Step4 result rows to scan.
        gem: Loaded GEM-supply codes, or ``None`` to skip.
        title: Composite-title prefix to match (default ``GEM_SUPPLY_TITLE``).

    Returns:
        Number of rows whose UL status was changed. ``0`` when ``gem`` is
        ``None`` or ``gem.codes`` is empty (including load-error payloads).
    """
    if gem is None or not gem.codes:
        return 0

    title_norm = normalize_packing_key_part(title)
    changed = 0
    for row in rows:
        if row.row_type != RowType.position_row:
            continue
        status = str(row.get_value(UL_COMPARE_STATUS) or "").strip()
        if status not in _GEM_OVERLAY_STATUSES:
            continue
        if _is_in_cabinet_position(row):
            continue
        parsed = parse_composite_title(row.get_value(DS_TITLE))
        if parsed is None:
            continue
        if normalize_packing_key_part(parsed[0]) != title_norm:
            continue
        code = normalize_packing_key_part(row.get_value(CODE))
        code_mto = normalize_packing_key_part(row.get_value(CODE_MTO))
        if code not in gem.codes and code_mto not in gem.codes:
            continue
        if UL_COMPARE_STATUS not in row.el:
            row.el[UL_COMPARE_STATUS] = CheckElement(None)
        el = row.el[UL_COMPARE_STATUS]
        el.value = STATUS_GEM_SUPPLY
        _paint_row_match_matched(row)
        changed += 1
    return changed
