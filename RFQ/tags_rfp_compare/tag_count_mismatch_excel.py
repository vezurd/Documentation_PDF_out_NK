"""Compact Excel report: tag count ≠ quantity (MTO Step4.1 / RFP Step1)."""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any, Iterable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from RFQ.tags_rfp_compare.rfp_tags_utils import ensure_result_dir_exists

PARTS_LOT_TAG_MISMATCH_LABELS = frozenset({"Лот > тегов", "Тегов > лота"})

_HEADER_FILL = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_MISMATCH_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def save_tag_count_mismatch_excel(
    count_mismatches: list[dict[str, Any]],
    result_dir: str,
    *,
    source_label: str,
    name_header: str,
    name_key: str,
    ds_header: str,
    ds_key: str,
    timestamp: str | None = None,
) -> str | None:
    """Write ``Отчет по несоответствию тегов - {source_label}_{ts}.xlsx``.

    Args:
        count_mismatches: Rows with keys ``title_system``, ``code``, ``tags_count``,
            ``values``, ``tags``, plus ``name_key`` / ``ds_key``.
        result_dir: Launch result folder.
        source_label: ``MTO`` or ``RFP`` (filename).
        name_header: Column D header (``MTO_NAME`` / ``NAME``).
        name_key: Dict key for column D.
        ds_header: Column E header (``DS_TITLE`` / ``DS_NAME``).
        ds_key: Dict key for column E.
        timestamp: Filename stamp; default ``YYYYMMDD_HHMMSS``.

    Returns:
        Saved path, or None on error / empty input.
    """
    if not count_mismatches:
        return None
    wb = None
    try:
        ensure_result_dir_exists(result_dir)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Несовпадения тегов и VALUES"
        headers = [
            "№",
            "title_system",
            "CODE",
            name_header,
            ds_header,
            "Кол-во тегов",
            "VALUES",
            "Теги",
        ]
        ws.append(headers)
        for col_num, _header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _HEADER_ALIGNMENT
            cell.border = _BORDER

        for idx, mismatch in enumerate(count_mismatches, 1):
            tags_str = str(mismatch.get("tags") or "")
            if ", " in tags_str:
                tags_with_newlines = tags_str.replace(", ", "\n")
            else:
                tags_with_newlines = tags_str
            ws.append(
                [
                    idx,
                    mismatch.get("title_system") or "не определен",
                    mismatch.get("code") or "",
                    mismatch.get(name_key) or "",
                    mismatch.get(ds_key) or "",
                    mismatch.get("tags_count"),
                    mismatch.get("values"),
                    tags_with_newlines,
                ]
            )
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=ws.max_row, column=col_num)
                cell.border = _BORDER
                if col_num == 8:
                    cell.alignment = Alignment(
                        horizontal="left", vertical="top", wrap_text=True
                    )
                elif col_num in (1, 6, 7):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                if col_num in (6, 7):
                    cell.fill = _MISMATCH_FILL

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 50
        ws.column_dimensions["E"].width = 20
        ws.column_dimensions["F"].width = 15
        ws.column_dimensions["G"].width = 12
        ws.column_dimensions["H"].width = 50
        last_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"
        ws.freeze_panes = "A2"

        stamp = timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"Отчет по несоответствию тегов - {source_label}_{stamp}.xlsx"
        file_path = os.path.join(result_dir, file_name)
        wb.save(file_path)
        return file_path
    except Exception as exc:
        print(f"Ошибка при сохранении таблицы несовпадений в Excel: {exc}")
        return None
    finally:
        if wb is not None:
            wb.close()


def load_parts_lot_tag_mismatches(path: str | os.PathLike) -> list[dict[str, Any]]:
    """Read «Лот > тегов» / «Тегов > лота» from ``Отчет по тегам - сбор частей.xlsx``.

    Args:
        path: Stamp-folder tag report.

    Returns:
        Compact mismatch dicts for ``save_tag_count_mismatch_excel`` (RFP keys).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: Iterable[tuple[Any, ...]] = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            return []
        col_index = {
            str(name).strip(): i
            for i, name in enumerate(header_row)
            if name is not None and str(name).strip()
        }
        type_i = col_index.get("Тип")
        if type_i is None:
            return []

        def _cell(row: tuple[Any, ...], header: str, default: str = "") -> Any:
            idx = col_index.get(header)
            if idx is None or idx >= len(row) or row[idx] is None:
                return default
            return row[idx]

        mismatches: list[dict[str, Any]] = []
        for row in rows:
            if not row or type_i >= len(row):
                continue
            kind = str(row[type_i] or "").strip()
            if kind not in PARTS_LOT_TAG_MISMATCH_LABELS:
                continue
            tags_text = _cell(row, "Теги в ячейке", "")
            if not tags_text:
                tags_text = _cell(row, "Тег", "")
            mismatches.append(
                {
                    "title_system": str(_cell(row, "Титул", "") or "не определен"),
                    "code": str(_cell(row, "Код", "") or ""),
                    "name": str(_cell(row, "Наименование", "") or ""),
                    "ds_name": str(_cell(row, "Имя ДС", "") or ""),
                    "tags_count": _cell(row, "Число тегов", None),
                    "values": _cell(row, "Кол-во лота", None),
                    "tags": str(tags_text or ""),
                }
            )
        return mismatches
    finally:
        wb.close()
