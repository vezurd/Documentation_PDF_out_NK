"""Load 1C UPD upload workbooks (файл закачки УПД) and merge position rows.

Two layouts of the same 1C report are supported:

* **compact** — ``Нпп`` + ``Позиции спецификации`` (no contract-structure levels)
* **wide** — extra ``Уровень 1..5`` before ``Позиция``

All mapped columns are kept, plus unknown extra headers. The merged xlsx adds
source file name, full path, relative path, sheet, and Excel row.
"""

from __future__ import annotations

import math
import os
import re
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import xlsxwriter
from openpyxl.utils import get_column_letter

from base.base_classes import RowType

DEFAULT_UPD_ROOT = (
    r"\\bcc\root\CurProjects\di_manegers\Пятницкий_П"
    r"\Амурский ГХК\Поставки\Файл закачки УПД по всем ДС"
)
DEFAULT_UPD_OUTPUT_DIR = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP\УПД сводный файл"
)

UPD_SUMMARY_STEM = "upd_summary"
UPD_SUMMARY_TIMESTAMP_FMT = "%Y.%m.%d_%H.%M.%S"
UPD_LOAD_REPORT_NAME = "upd_load_report.txt"

DEFAULT_UPD_LOAD_WORKERS = 8
_MIN_FILES_FOR_PARALLEL = 2
_HEADER_SCAN_ROWS = 20
_SUMMARY_HEADER_FILL = "#fce4d6"

LAYOUT_COMPACT = "compact"
LAYOUT_WIDE = "wide"

SOURCE_FILE = "source_file"
SOURCE_PATH = "source_path"
SOURCE_REL = "source_rel"
SOURCE_FOLDER = "source_folder"
SOURCE_SHEET = "source_sheet"
SOURCE_EXCEL_ROW = "source_excel_row"
SOURCE_LAYOUT = "source_layout"

_WS_RE = re.compile(r"[\s\u00a0]+")
_PRIMARY_SHEET_RE = re.compile(
    r"^(лист[_\s\-]?1(?:\s*\(\d+\))?|общая|общий)$",
    re.IGNORECASE,
)
_XLSX_SUFFIXES = {".xlsx", ".xlsm"}
_UNSUPPORTED_SUFFIXES = {".xls", ".xlsb"}

# Canonical data keys in export order (after source columns).
CANONICAL_COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("npp", "Нпп", 12),
    ("level_1", "Уровень 1", 28),
    ("level_2", "Уровень 2", 18),
    ("level_3", "Уровень 3", 18),
    ("level_4", "Уровень 4", 18),
    ("level_5", "Уровень 5", 18),
    ("spec_position", "Позиция (спецификация)", 40),
    ("spec_qty", "Кол-во (остаток)", 12),
    ("spec_units", "Ед.изм. (спецификация)", 12),
    ("spec_cost", "Стоимость (остаток)", 14),
    ("contract_stage", "Этап по договору", 28),
    ("vendor_code", "Код производителя", 16),
    ("name", "Наименование", 36),
    ("code_1c", "Код 1С", 12),
    ("goods_units", "Ед.изм. (товар)", 12),
    ("remainder", "Остаток", 10),
    ("price_excl_vat", "Цена (без НДС)", 12),
    ("amount_excl_vat", "Сумма (без НДС)", 14),
    ("price_bu", "Цена БУ (без НДС)", 12),
    ("amount_bu", "Сумма БУ (без НДС)", 14),
    ("sale_amount_incl_vat", "Сумма продажи (с НДС)", 16),
    ("section_code", "Раздел (код)", 12),
    ("section", "Раздел", 22),
    ("btp_row", "Номер строки БТП", 18),
    ("crm_guid", "CRM_GUID", 36),
    ("spec_currency", "Валюта спецификации", 12),
    ("ks2_flag", "По КС2 типа «Внутр.»", 14),
    ("ks2", "КС2 типа «Внутр.»", 16),
    ("bfirm", "БФирма", 18),
    ("spec_line_id", "Id строки спецификации", 36),
)
CANONICAL_KEYS: tuple[str, ...] = tuple(item[0] for item in CANONICAL_COLUMNS)
_CANONICAL_KEY_SET = frozenset(CANONICAL_KEYS)
_NUMERIC_KEYS = frozenset(
    {
        "spec_qty",
        "spec_cost",
        "remainder",
        "price_excl_vat",
        "amount_excl_vat",
        "price_bu",
        "amount_bu",
        "sale_amount_incl_vat",
    }
)
_SOURCE_COLUMNS: tuple[tuple[str, str, int], ...] = (
    (SOURCE_FILE, "Имя файла", 40),
    (SOURCE_PATH, "Путь к файлу", 70),
    (SOURCE_REL, "Относительный путь", 40),
    (SOURCE_FOLDER, "Папка ДС", 14),
    (SOURCE_SHEET, "Вкладка", 16),
    (SOURCE_EXCEL_ROW, "Строка в исходном УПД", 12),
    (SOURCE_LAYOUT, "Раскладка", 10),
)


@dataclass
class UpdHeader:
    """Detected 1C upload header on one worksheet."""

    detail_row: int
    layout: str
    column_keys: list[str | None]
    captions: list[str]


@dataclass
class UpdRow:
    """One classified UPD row (values keyed by canonical / extra columns)."""

    row_type: str
    values: dict[str, object]
    extra_captions: dict[str, str] = field(default_factory=dict)


@dataclass
class UpdSheetHit:
    """Per-sheet counters for the load report."""

    source_rel: str
    sheet_name: str
    layout: str
    position_rows: int
    status: str


@dataclass
class _UpdFileOutcome:
    source_rel: str
    source_path: str
    source_folder: str
    ok: bool
    error: str | None = None
    positions: list[UpdRow] = field(default_factory=list)
    empty_rows: int = 0
    other_rows: int = 0
    sheets_loaded: list[str] = field(default_factory=list)
    sheets_skipped: list[str] = field(default_factory=list)
    hits: list[UpdSheetHit] = field(default_factory=list)
    remarks: list[str] = field(default_factory=list)
    extra_captions: dict[str, str] = field(default_factory=dict)
    format_mismatch: bool = False


@dataclass
class UpdLoadStats:
    """Aggregate counters for one UPD merge run."""

    files_total: int = 0
    files_xlsx: int = 0
    files_ok: int = 0
    files_failed: int = 0
    files_unsupported: int = 0
    files_format_mismatch: int = 0
    files_without_position: int = 0
    sheets_loaded: int = 0
    sheets_skipped: int = 0
    position_rows: int = 0
    empty_rows: int = 0
    other_rows: int = 0
    layout_compact: int = 0
    layout_wide: int = 0
    folder_files: Counter[str] = field(default_factory=Counter)
    folder_positions: Counter[str] = field(default_factory=Counter)
    failures: list[str] = field(default_factory=list)
    remarks: list[str] = field(default_factory=list)
    sheet_hits: list[UpdSheetHit] = field(default_factory=list)

    def format_short(self) -> str:
        """One-line GUI / ActionResult summary."""
        return (
            f"файлов OK={self.files_ok}/{self.files_xlsx}, "
            f"ошибок={self.files_failed}, неподдерживаемых={self.files_unsupported}, "
            f"чужой формат={self.files_format_mismatch}, "
            f"position_row={self.position_rows} "
            f"(compact={self.layout_compact}, wide={self.layout_wide})"
        )


@dataclass
class UpdLoadResult:
    """Outcome of ``load_and_merge_upd``."""

    summary_path: str
    report_path: str
    stats: UpdLoadStats
    success: bool
    extra_captions: dict[str, str] = field(default_factory=dict)


_last_upd_result: UpdLoadResult | None = None


def upd_output_dir(path: str | Path | None = None) -> Path:
    """Return UPD summary directory; create if missing."""
    out = Path(path) if path else DEFAULT_UPD_OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    return out


def upd_summary_xlsx_name(*, when: datetime | None = None) -> str:
    """Return ``upd_summary_YYYY.MM.DD_HH.MM.SS.xlsx``."""
    stamp = (when or datetime.now()).strftime(UPD_SUMMARY_TIMESTAMP_FMT)
    return f"{UPD_SUMMARY_STEM}_{stamp}.xlsx"


def find_latest_upd_summary(output_dir: Path | None = None) -> Path | None:
    """Newest ``upd_summary_*.xlsx`` in the output dir, or ``None``."""
    directory = output_dir or DEFAULT_UPD_OUTPUT_DIR
    if not directory.is_dir():
        return None
    files = [
        path
        for path in directory.glob(f"{UPD_SUMMARY_STEM}*.xlsx")
        if path.is_file() and not path.name.startswith("~$")
    ]
    if not files:
        return None
    return max(files, key=lambda path: path.stat().st_mtime)


def get_last_upd_load_result() -> UpdLoadResult | None:
    """Return the last ``load_and_merge_upd`` result in this process."""
    return _last_upd_result


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value).replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    return _WS_RE.sub(" ", text).strip()


def _norm_header(value: object) -> str:
    text = _cell_text(value).casefold().replace("ё", "е")
    text = text.replace("*", " ")
    text = re.sub(r"[()\[\]«»\"'`/\\,.;:]+", " ", text)
    return _WS_RE.sub(" ", text).strip()


def _filled(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def is_primary_upd_sheet(name: str) -> bool:
    """True for Лист1 / Лист_1 / общая / общий (optional `` (2)`` suffix)."""
    return bool(_PRIMARY_SHEET_RE.match(str(name).strip()))


def _try_number(value: object) -> object:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return _cell_text(value)
        if value.is_integer():
            return int(value)
        return value
    text = _cell_text(value)
    if not text:
        return ""
    compact = text.replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        num = float(compact)
    except ValueError:
        return text
    if math.isnan(num) or math.isinf(num):
        return text
    if num.is_integer():
        return int(num)
    return num


def _map_header_label(norm: str, *, units_seen: int) -> str | None:
    """Map a normalized header cell to a canonical key, or None if group-only."""
    if not norm:
        return None
    if norm.startswith("нпп"):
        return "npp"
    level = re.fullmatch(r"уровень\s*([1-5])", norm)
    if level:
        return f"level_{level.group(1)}"
    if "кол-во" in norm or norm.startswith("количество"):
        return "spec_qty"
    if norm.startswith("стоимость"):
        return "spec_cost"
    if "этап по договору" in norm:
        return "contract_stage"
    if "код производителя" in norm:
        return "vendor_code"
    if "код 1с" in norm:
        return "code_1c"
    if norm.startswith("наименование"):
        return "name"
    if norm == "позиция" or norm.startswith("позиция "):
        return "spec_position"
    if norm.startswith("позиции спецификации") or "товарная позиция" in norm:
        return None
    if "структура договорной" in norm or "закупленные товарные" in norm:
        return None
    if "себестоимость остатков" in norm or norm == "в вал сделки":
        return None
    if "ед" in norm and "изм" in norm:
        return "spec_units" if units_seen == 0 else "goods_units"
    if norm == "остаток" or norm.startswith("остаток "):
        return "remainder"
    if "цена бу" in norm:
        return "price_bu"
    if "сумма бу" in norm:
        return "amount_bu"
    if "сумма продажи" in norm:
        return "sale_amount_incl_vat"
    if "цена" in norm and "ндс" in norm:
        return "price_excl_vat"
    if "сумма" in norm and "ндс" in norm:
        return "amount_excl_vat"
    if "раздел" in norm and "код" in norm:
        return "section_code"
    if "номер строки бтп" in norm or norm.endswith("бтп"):
        return "btp_row"
    if "crm" in norm and "guid" in norm:
        return "crm_guid"
    if "валюта спецификации" in norm:
        return "spec_currency"
    if norm.startswith("по кс2"):
        return "ks2_flag"
    if "кс2" in norm:
        return "ks2"
    if "бфирма" in norm:
        return "bfirm"
    if "id строки спецификации" in norm:
        return "spec_line_id"
    if norm == "раздел":
        return "section"
    return None


def _looks_like_detail_header(cells: list[object]) -> bool:
    joined = " ".join(_norm_header(c) for c in cells if _cell_text(c))
    if "наименование" not in joined:
        return False
    return "код 1с" in joined or "код производителя" in joined


def build_upd_header(
    detail_cells: list[object],
    group_cells: list[object] | None,
    *,
    detail_row: int,
) -> UpdHeader | None:
    """Build column map from the detail header row (optional group row above)."""
    width = len(detail_cells)
    if group_cells:
        width = max(width, len(group_cells))
    keys: list[str | None] = []
    captions: list[str] = []
    units_seen = 0
    used: dict[str, int] = {}
    extra_n = 0
    for idx in range(width):
        detail = detail_cells[idx] if idx < len(detail_cells) else None
        group = group_cells[idx] if group_cells and idx < len(group_cells) else None
        detail_text = _cell_text(detail)
        group_text = _cell_text(group)
        caption = detail_text or group_text
        norm = _norm_header(detail)
        key = _map_header_label(norm, units_seen=units_seen)
        if key is None and not detail_text:
            gkey = _map_header_label(_norm_header(group), units_seen=units_seen)
            if gkey in {"npp", "contract_stage"}:
                key = gkey
                caption = group_text
        if key in {"spec_units", "goods_units"}:
            units_seen += 1
        if key is None and caption:
            extra_n += 1
            key = f"extra:{extra_n}:{_norm_header(caption) or get_column_letter(idx + 1).lower()}"
        if key is not None and key in used and key in _CANONICAL_KEY_SET:
            extra_n += 1
            key = f"extra:{extra_n}:{key}"
        if key is not None:
            used[key] = used.get(key, 0) + 1
        keys.append(key)
        captions.append(caption)
    mapped = {k for k in keys if k in _CANONICAL_KEY_SET}
    if "name" not in mapped or "code_1c" not in mapped:
        return None
    if "spec_position" not in mapped and "spec_qty" not in mapped:
        return None
    layout = LAYOUT_WIDE if "level_1" in mapped else LAYOUT_COMPACT
    return UpdHeader(
        detail_row=detail_row,
        layout=layout,
        column_keys=keys,
        captions=captions,
    )


def detect_upd_header(preview_rows: list[list[object]]) -> UpdHeader | None:
    """Find the 1C upload detail header in the first scanned rows."""
    for idx, cells in enumerate(preview_rows):
        if not _looks_like_detail_header(cells):
            continue
        group = preview_rows[idx - 1] if idx > 0 else None
        header = build_upd_header(cells, group, detail_row=idx + 1)
        if header is not None:
            return header
    return None


def classify_upd_row(values: dict[str, object]) -> str:
    """Classify a data row after the header as position / empty / other."""
    data_values = [
        v
        for k, v in values.items()
        if k in _CANONICAL_KEY_SET or str(k).startswith("extra:")
    ]
    if not any(_filled(v) for v in data_values):
        return RowType.empty_row
    joined = " ".join(_norm_header(v) for v in data_values if _filled(v))
    if "наименование" in joined and ("код 1с" in joined or "позиции спецификации" in joined):
        return RowType.other_row
    if (
        _filled(values.get("spec_position"))
        or _filled(values.get("name"))
        or _filled(values.get("code_1c"))
        or _filled(values.get("npp"))
        or _filled(values.get("vendor_code"))
    ):
        return RowType.position_row
    return RowType.other_row


def _row_values_from_cells(cells: list[object], header: UpdHeader) -> tuple[dict[str, object], dict[str, str]]:
    values: dict[str, object] = {}
    extra_captions: dict[str, str] = {}
    for idx, key in enumerate(header.column_keys):
        if not key:
            continue
        raw = cells[idx] if idx < len(cells) else None
        if key in _NUMERIC_KEYS:
            values[key] = _try_number(raw)
        else:
            values[key] = _cell_text(raw)
        if key.startswith("extra:"):
            extra_captions[key] = header.captions[idx] or key
    return values, extra_captions


def _trim_row(cells: tuple[Any, ...] | list[Any]) -> list[object]:
    out = list(cells)
    while out and not _cell_text(out[-1]):
        out.pop()
    return out


def parse_upd_worksheet(ws: Any) -> tuple[UpdHeader | None, list[UpdRow], dict[str, str]]:
    """Read one worksheet: detect header, classify rows, return extras."""
    raw_rows: list[tuple[int, list[object]]] = []
    for excel_row, row in enumerate(ws.iter_rows(values_only=True), start=1):
        raw_rows.append((excel_row, _trim_row(row)))
    preview = [cells for _, cells in raw_rows[:_HEADER_SCAN_ROWS]]
    header = detect_upd_header(preview)
    if header is None:
        return None, [], {}
    rows: list[UpdRow] = []
    extra_captions: dict[str, str] = {}
    for excel_row, cells in raw_rows:
        if excel_row <= header.detail_row:
            continue
        values, extras = _row_values_from_cells(cells, header)
        extra_captions.update(extras)
        values["_excel_row"] = excel_row
        row_type = classify_upd_row(values)
        rows.append(UpdRow(row_type=row_type, values=values, extra_captions=extras))
    return header, rows, extra_captions


def _select_sheets_to_load(
    sheet_headers: dict[str, UpdHeader],
) -> tuple[list[str], list[str]]:
    """Prefer primary sheets when present so title-split copies are not doubled."""
    upd_names = list(sheet_headers)
    primary = [name for name in upd_names if is_primary_upd_sheet(name)]
    if primary:
        skipped = [name for name in upd_names if name not in primary]
        return primary, skipped
    return upd_names, []


def _rel_under_root(root: Path, file_path: Path) -> str:
    try:
        return str(file_path.resolve().relative_to(root.resolve()))
    except Exception:
        return file_path.name


def collect_upd_files(root: Path) -> tuple[list[Path], list[Path]]:
    """Recursive Excel files: ``(xlsx/xlsm, unsupported xls/xlsb)``."""
    supported: list[Path] = []
    unsupported: list[Path] = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [
            name
            for name in dirnames
            if "результат_проверки" not in name.casefold()
        ]
        for name in filenames:
            if name.startswith("~$"):
                continue
            path = Path(dirpath) / name
            suffix = path.suffix.lower()
            if suffix in _XLSX_SUFFIXES:
                supported.append(path)
            elif suffix in _UNSUPPORTED_SUFFIXES:
                unsupported.append(path)
    supported.sort(key=lambda p: str(p).casefold())
    unsupported.sort(key=lambda p: str(p).casefold())
    return supported, unsupported


def _load_one_file(path: Path, root: Path) -> _UpdFileOutcome:
    rel = _rel_under_root(root, path)
    folder = str(Path(rel).parent)
    if folder in {".", ""}:
        folder = "(root)"
    outcome = _UpdFileOutcome(
        source_rel=rel,
        source_path=str(path),
        source_folder=folder,
        ok=False,
    )
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        outcome.error = f"{type(exc).__name__}: {exc}"
        return outcome
    try:
        headers: dict[str, UpdHeader] = {}
        parsed: dict[str, tuple[list[UpdRow], dict[str, str]]] = {}
        for ws in wb.worksheets:
            header, rows, extras = parse_upd_worksheet(ws)
            if header is None:
                continue
            headers[ws.title] = header
            parsed[ws.title] = (rows, extras)
        if not headers:
            outcome.error = "нет вкладки формата закачки УПД (1С Остатки / Реализация 2_0)"
            outcome.format_mismatch = True
            outcome.remarks.append(f"{rel}: чужой формат (нет шапки УПД)")
            return outcome
        load_names, skip_names = _select_sheets_to_load(headers)
        outcome.sheets_loaded = list(load_names)
        outcome.sheets_skipped = list(skip_names)
        for name in skip_names:
            outcome.hits.append(
                UpdSheetHit(
                    source_rel=rel,
                    sheet_name=name,
                    layout=headers[name].layout,
                    position_rows=0,
                    status="skipped_split",
                )
            )
        for name in load_names:
            header = headers[name]
            rows, extras = parsed[name]
            outcome.extra_captions.update(extras)
            positions = 0
            for row in rows:
                excel_row = int(row.values.pop("_excel_row", 0) or 0)
                if row.row_type == RowType.empty_row:
                    outcome.empty_rows += 1
                    continue
                if row.row_type != RowType.position_row:
                    outcome.other_rows += 1
                    continue
                row.values[SOURCE_FILE] = path.name
                row.values[SOURCE_PATH] = str(path)
                row.values[SOURCE_REL] = rel
                row.values[SOURCE_FOLDER] = folder
                row.values[SOURCE_SHEET] = name
                row.values[SOURCE_EXCEL_ROW] = excel_row
                row.values[SOURCE_LAYOUT] = header.layout
                outcome.positions.append(row)
                positions += 1
            outcome.hits.append(
                UpdSheetHit(
                    source_rel=rel,
                    sheet_name=name,
                    layout=header.layout,
                    position_rows=positions,
                    status="loaded",
                )
            )
        outcome.ok = True
        if not outcome.positions:
            outcome.remarks.append(f"{rel}: шапка УПД есть, position_row=0")
        return outcome
    except Exception as exc:
        outcome.error = f"{type(exc).__name__}: {exc}"
        return outcome
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _resolve_workers(files_total: int, requested: int | None) -> int:
    if files_total < _MIN_FILES_FOR_PARALLEL:
        return 1
    n = DEFAULT_UPD_LOAD_WORKERS if requested is None else int(requested)
    return max(1, min(n, files_total))


def _source_export_columns() -> tuple[tuple[str, str, int], ...]:
    return _SOURCE_COLUMNS


def _data_export_columns(
    extra_captions: dict[str, str],
) -> list[tuple[str, str, int]]:
    cols = list(CANONICAL_COLUMNS)
    for key in sorted(extra_captions):
        cols.append((key, extra_captions[key], 18))
    return cols


def save_upd_summary_xlsx(
    rows: list[UpdRow],
    extra_captions: dict[str, str],
    out_path: Path,
) -> Path:
    """Write merged position rows (xlsxwriter)."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    columns = list(_source_export_columns()) + _data_export_columns(extra_captions)
    numeric = _NUMERIC_KEYS | {SOURCE_EXCEL_ROW}
    wb = xlsxwriter.Workbook(str(out_path), {"constant_memory": False, "remove_timezone": True})
    ws = wb.add_worksheet("УПД свод")
    header_fmt = wb.add_format(
        {
            "bold": True,
            "font_color": "#000000",
            "bg_color": _SUMMARY_HEADER_FILL,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "text_wrap": True,
        }
    )
    cell_fmt = wb.add_format({"align": "left", "valign": "vcenter", "border": 1})
    num_fmt = wb.add_format(
        {"align": "right", "valign": "vcenter", "border": 1, "num_format": "0.####"}
    )
    ws.set_row(0, 36)
    for ci, (_key, label, width) in enumerate(columns):
        ws.write(0, ci, label, header_fmt)
        ws.set_column(ci, ci, width)
    for row_idx, row in enumerate(rows, start=1):
        for ci, (key, _label, _width) in enumerate(columns):
            value = row.values.get(key, "")
            if value in (None, ""):
                ws.write_blank(row_idx, ci, None, cell_fmt)
                continue
            if key in numeric and isinstance(value, (int, float)) and not isinstance(
                value, bool
            ):
                ws.write_number(row_idx, ci, float(value), num_fmt)
            else:
                ws.write(row_idx, ci, value, cell_fmt)
    last = len(rows)
    if last > 0:
        ws.autofilter(0, 0, last, len(columns) - 1)
    ws.freeze_panes(1, 0)
    try:
        wb.close()
    except xlsxwriter.exceptions.FileCreateError:
        retry = out_path.with_name(
            f"{UPD_SUMMARY_STEM}_"
            f"{datetime.now().strftime('%Y.%m.%d_%H.%M.%S.%f')}"
            f"{out_path.suffix}"
        )
        print(f"Сводный файл УПД не записан, повтор: {retry.name}")
        return save_upd_summary_xlsx(rows, extra_captions, retry)
    return out_path


def _format_stats_report(
    stats: UpdLoadStats,
    root: str,
    extra_captions: dict[str, str],
) -> str:
    lines = [
        f"root: {root}",
        stats.format_short(),
        f"sheets_loaded: {stats.sheets_loaded}",
        f"sheets_skipped: {stats.sheets_skipped}",
        f"empty_rows: {stats.empty_rows}",
        f"other_rows: {stats.other_rows}",
        f"files_without_position: {stats.files_without_position}",
        "",
        "папки (файлов / position_row):",
    ]
    folders = sorted(
        set(stats.folder_files) | set(stats.folder_positions),
        key=lambda name: name.casefold(),
    )
    for folder in folders:
        lines.append(
            f"  {folder}: files={stats.folder_files[folder]} "
            f"positions={stats.folder_positions[folder]}"
        )
    if extra_captions:
        lines.append("")
        lines.append("доп. колонки:")
        for key, caption in sorted(extra_captions.items()):
            lines.append(f"  {key} = {caption}")
    if stats.failures:
        lines.append("")
        lines.append("ошибки:")
        lines.extend(f"  {item}" for item in stats.failures)
    if stats.remarks:
        lines.append("")
        lines.append("замечания:")
        lines.extend(f"  {item}" for item in stats.remarks)
    return "\n".join(lines).rstrip() + "\n"


def load_and_merge_upd(
    root: str | Path,
    *,
    output_dir: str | Path | None = None,
    workers: int | None = None,
) -> UpdLoadResult:
    """Walk UPD upload xlsx, classify ``position_row``, write merged summary.

    Args:
        root: Folder with 1C upload workbooks (recursive).
        output_dir: Summary/report directory; default network «УПД сводный файл».
        workers: Parallel file readers; ``1`` is sequential.

    Returns:
        Paths to summary xlsx and txt report, plus counters.

    Raises:
        FileNotFoundError: Root folder is missing.
        RuntimeError: No xlsx files, or none produced a position row.
    """
    global _last_upd_result

    root_path = Path(str(root or "").strip())
    if not root_path.is_dir():
        raise FileNotFoundError(f"Папка УПД не найдена: {root_path}")

    xlsx_files, unsupported = collect_upd_files(root_path)
    stats = UpdLoadStats(
        files_total=len(xlsx_files) + len(unsupported),
        files_xlsx=len(xlsx_files),
        files_unsupported=len(unsupported),
    )
    for path in unsupported:
        rel = _rel_under_root(root_path, path)
        stats.failures.append(f"{rel}: неподдерживаемое расширение {path.suffix}")
        stats.remarks.append(f"{rel}: пропуск {path.suffix} (нужен xlsx)")

    print(f"УПД: папка {root_path}")
    print(
        f"УПД: xlsx={len(xlsx_files)}, "
        f"пропуск не-xlsx={len(unsupported)}"
    )
    if not xlsx_files:
        raise RuntimeError("В папке нет файлов .xlsx / .xlsm для чтения УПД.")

    n_workers = _resolve_workers(len(xlsx_files), workers)
    outcomes: list[_UpdFileOutcome] = []
    if n_workers == 1:
        for path in xlsx_files:
            outcomes.append(_load_one_file(path, root_path))
    else:
        with ThreadPoolExecutor(max_workers=n_workers) as pool:
            futs = {
                pool.submit(_load_one_file, path, root_path): path
                for path in xlsx_files
            }
            done = 0
            for fut in as_completed(futs):
                outcomes.append(fut.result())
                done += 1
                if done % 25 == 0 or done == len(xlsx_files):
                    print(f"УПД: прочитано {done}/{len(xlsx_files)}")

    outcomes.sort(key=lambda item: item.source_rel.casefold())
    positions: list[UpdRow] = []
    extra_captions: dict[str, str] = {}
    for outcome in outcomes:
        extra_captions.update(outcome.extra_captions)
        stats.sheet_hits.extend(outcome.hits)
        stats.sheets_loaded += len(outcome.sheets_loaded)
        stats.sheets_skipped += len(outcome.sheets_skipped)
        stats.empty_rows += outcome.empty_rows
        stats.other_rows += outcome.other_rows
        stats.remarks.extend(outcome.remarks)
        stats.folder_files[outcome.source_folder] += 1
        if not outcome.ok:
            stats.files_failed += 1
            if outcome.format_mismatch:
                stats.files_format_mismatch += 1
            msg = outcome.error or "ошибка чтения"
            stats.failures.append(f"{outcome.source_rel}: {msg}")
            print(f"УПД: FAIL {outcome.source_rel} — {msg}")
            continue
        stats.files_ok += 1
        stats.folder_positions[outcome.source_folder] += len(outcome.positions)
        positions.extend(outcome.positions)
        layouts = {hit.layout for hit in outcome.hits if hit.status == "loaded"}
        if LAYOUT_WIDE in layouts:
            stats.layout_wide += 1
        elif LAYOUT_COMPACT in layouts:
            stats.layout_compact += 1
        if not outcome.positions:
            stats.files_without_position += 1
        npos = len(outcome.positions)
        print(
            f"УПД: OK {outcome.source_rel}  "
            f"sheets={len(outcome.sheets_loaded)} position_row={npos}"
        )

    stats.position_rows = len(positions)
    if not positions:
        raise RuntimeError(
            "Не найдено ни одной position_row УПД. "
            "Проверьте шапку (Наименование + Код 1С) и отчёт загрузки."
        )

    out_dir = upd_output_dir(output_dir)
    summary_path = out_dir / upd_summary_xlsx_name()
    try:
        saved = save_upd_summary_xlsx(positions, extra_captions, summary_path)
    except Exception as exc:
        raise RuntimeError(f"Не удалось записать свод УПД: {exc}") from exc
    report_path = out_dir / UPD_LOAD_REPORT_NAME
    report_path.write_text(
        _format_stats_report(stats, str(root_path), extra_captions),
        encoding="utf-8",
    )
    success = stats.files_ok > 0 and stats.position_rows > 0
    result = UpdLoadResult(
        summary_path=str(saved),
        report_path=str(report_path),
        stats=stats,
        success=success,
        extra_captions=extra_captions,
    )
    _last_upd_result = result
    print(f"УПД: {stats.format_short()}")
    print(f"УПД: свод {saved}")
    print(f"УПД: отчёт {report_path}")
    return result
