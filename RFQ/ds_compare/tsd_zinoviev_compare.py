"""Compare robot TSD packing cache vs Zinoviev manual summary workbook.

Manual source (default)::

    \\\\bcc\\eng\\...\\УЛ сводный файл\\Сводный от Зиновьева\\ТСД по всем ДС_общий.xlsx

Sheet ``общая `` concatenates packing-list forms. Significant columns (0-based),
after UL name was inserted in column B (07.2026):

| Col | Field |
|-----|-------|
| A 0 | № ДС (ДС15 / 4905 / ГФ 5) |
| B 1 | packing-list name (explicit) |
| C 2 | CODE |
| D 3 | SPECIFICATION_NAME |
| E 4 | TAGS |
| I 8 | NAME |
| J 9 | VALUES |
| K 10 | UNITS |
| N 13 | VENDOR |

Garbage rows: banners (Shipper/Total/…), repeated bilingual headers,
column-index rows (15,16,…), empty CODE, non-numeric Quantity.

Match key: ``(DS_TITLE, DS_SYSTEM, CODE, packing_list_name)``.
Packing-list name: Zinoviev col B; robot ``ANNOTATION`` source file
(basename / ``PL_…`` / bare id). Sheet/tab is **not** in the key yet.

Tags are normalized to individual strings (lists, ``;``-joins, ``['tag']``
reprs). Excel has separate robot/Zinoviev tag columns plus a diff column,
a Zinoviev coverage verdict (qty + tags), and a second sheet matching
unique UL names across both sides.
"""

from __future__ import annotations

import ast
import math
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import xlsxwriter
from openpyxl import load_workbook

from RFQ.ds_compare.ds_units_normalize import normalize_units_text
from RFQ.ds_compare.tsd_packing_load import parse_title_system_from_spec_name
from RFQ.packing_list_provider import (
    PackingQualityLevel,
    load_packing_dataset,
    normalize_packing_key_part,
    packing_match_key,
)
from base.base_classes import RowStd, RowType
from base.tables_columns import (
    ANNOTATION,
    CODE,
    DS_SYSTEM,
    DS_TITLE,
    NAME,
    SPECIFICATION_NAME,
    TAGS,
    TITLE,
    UNITS,
    VALUES,
    VENDOR,
)
from utils.colors import Color

DEFAULT_ZINOVIEV_DIR = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP"
    r"\УЛ сводный файл\Сводный от Зиновьева"
)
DEFAULT_ZINOVIEV_XLSX = DEFAULT_ZINOVIEV_DIR / "ТСД по всем ДС_общий.xlsx"

_RESULT_FOLDER_PREFIX = "_результат_сравнения_"
_RESULT_FOLDER_FMT = "%Y_%m_%d_%H_%M"
# Two reports per run: without UL in key (legacy) and with UL name in key.
_RESULT_XLSX_BY_CODE = "сравнение_титул_марка_код.xlsx"
_RESULT_XLSX_BY_UL = "сравнение_титул_марка_код_УЛ.xlsx"
_RESULT_REPORT_BY_CODE = "отчет_титул_марка_код.txt"
_RESULT_REPORT_BY_UL = "отчет_титул_марка_код_УЛ.txt"
# Backward-compatible aliases (UL-in-key variant).
_RESULT_XLSX_NAME = _RESULT_XLSX_BY_UL
_RESULT_REPORT_NAME = _RESULT_REPORT_BY_UL

_ZIN_SHEET_CANDIDATES = ("общая ", "общая", "Общая ", "Общая")

# 0-based column indexes on the Zinoviev packing form body (UL name in B).
_COL_DS_NUM = 0
_COL_UL_NAME = 1
_COL_CODE = 2
_COL_SPEC = 3
_COL_TAGS = 4
_COL_NAME = 8
_COL_QTY = 9
_COL_UNITS = 10
_COL_VENDOR = 13

_BANNER_MARKERS = (
    "грузоотправитель",
    "грузополучатель",
    "основание поставки",
    "упаковочный лист",
    "контактные",
    "итого",
    "shipper",
    "consignee",
    "total",
    "№ грузового",
    "номер грузового",
    "код рд",
    "code rd",
    "тип упаковки",
    "package types",
    "нетто,",
    "net weight",
)

_TAG_SPLIT_RE = re.compile(r"[;\n,]+")
_TAG_EDGE_JUNK_RE = re.compile(r"""^[\s\[\]\{\}'\"«»]+|[\s\[\]\{\}'\"«»]+$""")
# Packing-list ids: PL_2076…, PL-АГХК-…, PLAN004466FR / AN004466FR.
_PL_TOKEN_RE = re.compile(r"PL[_-][\w.-]+", re.IGNORECASE)
_PLAN_TOKEN_RE = re.compile(r"(?:PLAN|AN)(\d{4,}[A-Za-z0-9]*)", re.IGNORECASE)
_PACKING_LIST_NUM_RE = re.compile(
    r"(?:packing\s*list|упаковочный\s*лист)\s*(?:no\.?|№|#)?\s*([^\s\\/]+)",
    re.IGNORECASE,
)
# Robot filenames often append revision noise after the real UL id.
_UL_REV_TAIL_RE = re.compile(
    r"(?:"
    r"[_\s./-]+(?:ред(?:акция)?|версия|вер\.?|согл\w*|новая|new|rev)[\w.\s/-]*"
    r"|(?<=\d)ред(?:акция)?[\w.\s/-]*"
    r")$",
    re.IGNORECASE | re.DOTALL,
)
_UL_TRAIL_COPY_RE = re.compile(r"[_\-][1-9]$")
_UL_TRAIL_JUNK_RE = re.compile(r"[._\-\s]+$")
_XLSX_SUFFIXES = (".xlsx", ".xls", ".xlsm")

_STATUS_MATCH = "Совпадение"
_STATUS_QTY_DIFF = "Расхождение кол-ва"
_STATUS_ROBOT_ONLY = "Только в роботе"
_STATUS_ZIN_ONLY = "Только у Зиновьева"
_STATUS_NO_TITLE = "Нет титула/марки"
_STATUS_UNITS_DIFF = "Совпадение (ед. изм. различаются)"
_STATUS_ROBOT_UNAVAILABLE = "Кэш робота недоступен"

_UL_STATUS_MATCH = "Совпадение"
_UL_STATUS_SPELLING = "Совпадение (разное написание)"
_UL_STATUS_ZIN_ONLY = "Только у Зиновьева"
_UL_STATUS_ROBOT_ONLY = "Только у робота"

_COV_FULL = "Покрыто"
_COV_MISSING = "Нет в роботе"
_COV_QTY = "Недобор кол-ва"
_COV_TAGS = "Не хватает тегов"
_COV_QTY_TAGS = "Недобор кол-ва и тегов"
_COV_NA = "—"

_STATUS_COLORS: dict[str, str] = {
    _STATUS_MATCH: Color.green,
    _STATUS_UNITS_DIFF: Color.soft_yellow,
    _STATUS_QTY_DIFF: Color.yellow,
    _STATUS_ROBOT_ONLY: Color.soft_cyan,
    _STATUS_ZIN_ONLY: Color.match_added_mto,
    _STATUS_NO_TITLE: Color.red,
    _STATUS_ROBOT_UNAVAILABLE: Color.red,
}

_UL_STATUS_COLORS: dict[str, str] = {
    _UL_STATUS_MATCH: Color.green,
    _UL_STATUS_SPELLING: Color.soft_yellow,
    _UL_STATUS_ZIN_ONLY: Color.match_added_mto,
    _UL_STATUS_ROBOT_ONLY: Color.soft_cyan,
}

_COVERAGE_COLORS: dict[str, str] = {
    _COV_FULL: Color.green,
    _COV_MISSING: Color.red,
    _COV_QTY: Color.yellow,
    _COV_TAGS: Color.yellow,
    _COV_QTY_TAGS: Color.yellow,
    _COV_NA: Color.no,
}

_HEADER_FILL = {
    "key": "#dbbcdb",
    "status": "#ffcc00",
    "robot": "#c9daf8",
    "zin": "#cd7f32",
    "meta": "#ead1dc",
}


@dataclass
class ZinAgg:
    """Aggregated Zinoviev positions for one match key."""

    qty: float = 0.0
    units: set[str] = field(default_factory=set)
    names: set[str] = field(default_factory=set)
    tags: set[str] = field(default_factory=set)
    specs: set[str] = field(default_factory=set)
    ds_nums: set[str] = field(default_factory=set)
    vendors: set[str] = field(default_factory=set)
    ul_names: set[str] = field(default_factory=set)
    source_rows: int = 0
    title: str = ""
    system: str = ""
    code: str = ""
    has_title_system: bool = False


@dataclass
class RobotAgg:
    """Aggregated robot packing positions for one match key."""

    qty: float = 0.0
    units: set[str] = field(default_factory=set)
    names: set[str] = field(default_factory=set)
    tags: set[str] = field(default_factory=set)
    specs: set[str] = field(default_factory=set)
    sources: set[str] = field(default_factory=set)
    ul_names: set[str] = field(default_factory=set)
    source_rows: int = 0
    title: str = ""
    system: str = ""
    code: str = ""
    has_title_system: bool = False


@dataclass
class _RawZinPosition:
    """One Zinoviev position before key aggregation."""

    title: str
    system: str
    code: str
    ul_name: str
    qty: float
    units: str
    name: str
    tags: set[str]
    spec: str
    ds_num: str
    vendor: str
    has_title_system: bool


@dataclass
class CompareOutRow:
    """One Excel result row."""

    title: str
    system: str
    title_mark: str
    code: str
    ul_name: str
    ul_robot: str
    ul_zin: str
    status: str
    color: str
    zin_coverage: str
    coverage_color: str
    robot_qty: float | None
    zin_qty: float | None
    diff: float | None
    units_robot: str
    units_zin: str
    name: str
    tags_robot: str
    tags_zin: str
    tags_diff: str
    tags_diff_color: str
    spec_zin: str
    ds_num_zin: str
    robot_sources: str
    robot_rows: int
    zin_rows: int


@dataclass
class UlNameOutRow:
    """One row of the UL-name inventory sheet."""

    ul_zin: str
    ul_robot: str
    status: str
    color: str
    norm_key: str
    ul_robot_file: str = ""
    zin_rows: int = 0
    robot_rows: int = 0


@dataclass
class ZinCompareStats:
    """Counters for console / report."""

    zin_raw_rows: int = 0
    zin_positions: int = 0
    zin_skipped: int = 0
    zin_no_title: int = 0
    robot_positions: int = 0
    robot_quality: str = ""
    keys_total: int = 0
    match: int = 0
    qty_diff: int = 0
    units_diff: int = 0
    robot_only: int = 0
    zin_only: int = 0
    no_title: int = 0
    zin_keys: int = 0
    zin_covered: int = 0
    zin_not_in_robot: int = 0
    zin_qty_short: int = 0
    zin_tags_missing: int = 0
    tags_equal: int = 0
    tags_differ: int = 0


@dataclass
class ZinCompareResult:
    """Outcome of one compare run."""

    success: bool
    message: str
    result_dir: str
    excel_path: str | None = None
    stats: ZinCompareStats | None = None


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def packing_list_basename(value: object) -> str:
    """Return file basename from a path-like packing-list label."""
    text = _cell_text(value).replace("/", "\\")
    if "\\" in text:
        text = text.rsplit("\\", 1)[-1]
    return text.strip()


def _strip_xlsx_suffix(text: str) -> str:
    lower = text.lower()
    for suffix in _XLSX_SUFFIXES:
        if lower.endswith(suffix):
            return text[: -len(suffix)]
    return text


def _strip_ul_revision_tail(token: str) -> str:
    """Drop revision / copy suffixes from a packing-list id token.

    Examples::

        PL_2076961.10_2891.04_ред1_271125 → PL_2076961.10_2891.04
        …15.19ред 0502 → …15.19
        …2894.08_1 → …2894.08
        …3221.01- → …3221.01
    """
    text = token.strip()
    if not text:
        return ""
    text = _UL_REV_TAIL_RE.sub("", text)
    text = _UL_TRAIL_JUNK_RE.sub("", text)
    # Short copy marker after a full id (…2894.08_1); keep real endings like …14.9.
    if _UL_TRAIL_COPY_RE.search(text):
        stem = _UL_TRAIL_COPY_RE.sub("", text)
        if re.search(r"\d", stem):
            text = stem
    return _UL_TRAIL_JUNK_RE.sub("", text)


def extract_packing_list_token(value: object) -> str:
    """Extract canonical UL id from Zinoviev col B or robot filename."""
    text = _strip_xlsx_suffix(packing_list_basename(value)).strip()
    if not text:
        return ""
    match = _PL_TOKEN_RE.search(text)
    if match:
        return _strip_ul_revision_tail(match.group(0))
    plan = _PLAN_TOKEN_RE.search(text)
    if plan:
        # Zinoviev uses PLAN…; robot AGCC files often embed AN….
        return "PLAN" + plan.group(1)
    pl_num = _PACKING_LIST_NUM_RE.search(text)
    if pl_num:
        num = _strip_ul_revision_tail(pl_num.group(1).strip())
        bare = re.match(r"^\d[\d._-]*", num)
        if bare:
            bare_id = _strip_ul_revision_tail(bare.group(0))
            # Long vendor ids (2076961…) vs short sheet numbers (1, 5).
            digits = re.sub(r"\D", "", bare_id)
            if len(digits) >= 6 or "_" in bare_id or bare_id.count(".") >= 2:
                return bare_id
        return f"Packing list {num}"
    bare = re.match(r"^\d[\d._-]*", text.strip())
    if bare:
        return _strip_ul_revision_tail(bare.group(0))
    return _strip_ul_revision_tail(text)


def canonical_packing_list_label(value: object) -> str:
    """Short UL label for Excel display (aligned with Zinoviev col B style)."""
    token = extract_packing_list_token(value)
    return token or packing_list_basename(value)


def _normalize_ul_id_separators(text: str) -> str:
    """Treat ``_`` and ``.`` as equivalent in numeric UL ids.

    ``2076961.2_2451.01`` ↔ ``2076961.2.2451_01`` → ``2076961.2.2451.01``.
    """
    return text.replace("_", ".")


def normalize_packing_list_name(value: object) -> str:
    """Normalize packing-list name for the compare key (no sheet/tab).

    Prefers a ``PL_…`` / ``PL-…`` / bare ``2076961.…`` / ``AN######`` id when
    present so Zinoviev col B aligns with robot filenames; strips common
    revision suffixes (``_ред1_…``, ``ред 0502``, ``новая``, ``_1``);
    underscores in the id are treated as dots.
    """
    text = _strip_xlsx_suffix(packing_list_basename(value)).strip()
    if not text:
        return ""
    match = _PL_TOKEN_RE.search(text)
    if match:
        token = _strip_ul_revision_tail(match.group(0))
        # Drop optional PL_ / PL- prefix: PL_2076… ↔ 2076…; keep Cyrillic body.
        body = token[3:] if len(token) > 3 and token[2] in "_-" else token
        return normalize_packing_key_part(_normalize_ul_id_separators(body))
    plan = _PLAN_TOKEN_RE.search(text)
    if plan:
        return normalize_packing_key_part("an" + plan.group(1))
    pl_num = _PACKING_LIST_NUM_RE.search(text)
    if pl_num:
        num = _strip_ul_revision_tail(pl_num.group(1).strip())
        # «Packing list 2076961.2_2451.01» → same key as bare id in col B.
        bare_in_num = re.match(r"^\d[\d._-]*", num)
        if bare_in_num:
            return normalize_packing_key_part(
                _normalize_ul_id_separators(bare_in_num.group(0))
            )
        return normalize_packing_key_part("packinglist" + num)
    # Bare ids without PL_ prefix (seen on some Zinoviev col B values).
    bare = re.match(r"^\d[\d._-]*", text.strip())
    if bare:
        return normalize_packing_key_part(
            _normalize_ul_id_separators(_strip_ul_revision_tail(bare.group(0)))
        )
    # №5 / №6 style — keep as-is after casefold/space strip.
    return normalize_packing_key_part(_strip_ul_revision_tail(text))


def ul_compare_key(
    title: object,
    system: object,
    code: object,
    ul_name: object,
) -> tuple[str, str, str, str]:
    """Match key: title + system + code + packing-list name (not sheet)."""
    t, s, c = packing_match_key(title, system, code)
    return (t, s, c, normalize_packing_list_name(ul_name))


def _clean_tag_token(token: str) -> str:
    """Strip quotes/brackets/whitespace from one tag fragment."""
    text = token.strip()
    if not text:
        return ""
    # Repeatedly peel edge junk: ``['tag']`` → ``tag``.
    for _ in range(4):
        cleaned = _TAG_EDGE_JUNK_RE.sub("", text).strip()
        if cleaned == text:
            break
        text = cleaned
    return text


def parse_tags(value: object) -> set[str]:
    """Normalize TAGS cell to a set of individual tag strings.

    Accepts ``list``/``tuple`` (robot cache), plain strings, multiline cells,
    and messy joins like ``tag1; ['tag2']; ['tag1']``.
    """
    result: set[str] = set()
    if value is None or value == "":
        return result
    if isinstance(value, (list, tuple, set)):
        for item in value:
            result |= parse_tags(item)
        return result
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text or text.casefold() in {"none", "nan"}:
        return result
    # Whole-cell Python list/tuple repr.
    if (text.startswith("[") and text.endswith("]")) or (
        text.startswith("(") and text.endswith(")")
    ):
        try:
            parsed = ast.literal_eval(text)
        except (ValueError, SyntaxError, TypeError):
            parsed = None
        if isinstance(parsed, (list, tuple, set)):
            return parse_tags(parsed)
        if isinstance(parsed, str):
            return parse_tags(parsed)
    for part in _TAG_SPLIT_RE.split(text):
        cleaned = _clean_tag_token(part)
        if cleaned:
            result.add(cleaned)
    return result


def _format_tags(tags: set[str], *, limit: int = 40) -> str:
    items = sorted(tags, key=str.casefold)
    if not items:
        return ""
    if len(items) > limit:
        return "; ".join(items[:limit]) + f"; …(+{len(items) - limit})"
    return "; ".join(items)


def _format_tags_diff(robot_tags: set[str], zin_tags: set[str]) -> str:
    """Human-readable tag delta; empty when sets are equal."""
    robot_cf = {t.casefold(): t for t in robot_tags}
    zin_cf = {t.casefold(): t for t in zin_tags}
    only_zin = sorted(
        (zin_cf[k] for k in zin_cf.keys() - robot_cf.keys()),
        key=str.casefold,
    )
    only_robot = sorted(
        (robot_cf[k] for k in robot_cf.keys() - zin_cf.keys()),
        key=str.casefold,
    )
    parts: list[str] = []
    if only_zin:
        parts.append("только Зиновьев: " + "; ".join(only_zin))
    if only_robot:
        parts.append("только робот: " + "; ".join(only_robot))
    return " | ".join(parts)


def _zin_coverage(
    *,
    has_robot: bool,
    has_zin: bool,
    robot_qty: float | None,
    zin_qty: float | None,
    robot_tags: set[str],
    zin_tags: set[str],
) -> str:
    """Whether robot fully covers this Zinoviev key (qty + tags)."""
    if not has_zin:
        return _COV_NA
    if not has_robot:
        return _COV_MISSING
    qty_short = (
        robot_qty is not None
        and zin_qty is not None
        and robot_qty + 1e-6 < zin_qty
    )
    robot_cf = {t.casefold() for t in robot_tags}
    missing_tags = {t for t in zin_tags if t.casefold() not in robot_cf}
    if qty_short and missing_tags:
        return _COV_QTY_TAGS
    if qty_short:
        return _COV_QTY
    if missing_tags:
        return _COV_TAGS
    return _COV_FULL


def _parse_qty(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        num = float(value)
        return num if math.isfinite(num) and num >= 0 else None
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        num = float(text)
    except ValueError:
        return None
    return num if math.isfinite(num) and num >= 0 else None


def _is_banner_or_header(code_cell: str) -> bool:
    low = code_cell.casefold()
    if any(marker in low for marker in _BANNER_MARKERS):
        return True
    # Column-index banner under the bilingual header (15, 16, 17, …).
    if code_cell.isdigit() and len(code_cell) <= 3:
        return True
    return False


def _join_unique(values: Iterable[str], *, limit: int = 8) -> str:
    items = sorted({v for v in values if v})
    if len(items) > limit:
        return "; ".join(items[:limit]) + f"; …(+{len(items) - limit})"
    return "; ".join(items)


def _pick_sheet_name(sheetnames: list[str]) -> str:
    stripped = {name.strip().casefold(): name for name in sheetnames}
    for candidate in _ZIN_SHEET_CANDIDATES:
        key = candidate.strip().casefold()
        if key in stripped:
            return stripped[key]
    for name in sheetnames:
        if "общ" in name.casefold():
            return name
    return sheetnames[0]


def load_zinoviev_raw_positions(
    xlsx_path: str | Path,
) -> tuple[list[_RawZinPosition], ZinCompareStats]:
    """Read Zinoviev workbook once into raw positions (before key aggregation).

    Args:
        xlsx_path: Path to ``ТСД по всем ДС_общий.xlsx``.

    Returns:
        Raw positions and load counters (``zin_*`` fields filled).

    Raises:
        FileNotFoundError: Workbook missing.
        ValueError: Workbook has no sheets.
    """
    path = Path(xlsx_path)
    if not path.is_file():
        raise FileNotFoundError(f"Файл Зиновьева не найден: {path}")

    stats = ZinCompareStats()
    raw: list[_RawZinPosition] = []
    current_ul_name = ""

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if not wb.sheetnames:
            raise ValueError(f"В файле нет вкладок: {path}")
        ws = wb[_pick_sheet_name(list(wb.sheetnames))]
        for row in ws.iter_rows(min_row=2, max_col=_COL_VENDOR + 1, values_only=True):
            stats.zin_raw_rows += 1
            # Prefer explicit UL name from column B on every row.
            ul_cell = _cell_text(
                row[_COL_UL_NAME] if len(row) > _COL_UL_NAME else None
            )
            if ul_cell:
                current_ul_name = ul_cell

            code = _cell_text(row[_COL_CODE] if len(row) > _COL_CODE else None)
            if not code:
                stats.zin_skipped += 1
                continue
            if _is_banner_or_header(code):
                stats.zin_skipped += 1
                continue
            qty = _parse_qty(row[_COL_QTY] if len(row) > _COL_QTY else None)
            if qty is None:
                stats.zin_skipped += 1
                continue
            name = _cell_text(row[_COL_NAME] if len(row) > _COL_NAME else None)
            units_raw = _cell_text(row[_COL_UNITS] if len(row) > _COL_UNITS else None)
            if not name and not units_raw:
                stats.zin_skipped += 1
                continue

            spec = _cell_text(row[_COL_SPEC] if len(row) > _COL_SPEC else None)
            title, system = parse_title_system_from_spec_name(spec)
            has_ts = bool(title and system)
            if not has_ts:
                stats.zin_no_title += 1

            tags_raw = row[_COL_TAGS] if len(row) > _COL_TAGS else None
            ds_num = _cell_text(row[_COL_DS_NUM] if len(row) > _COL_DS_NUM else None)
            vendor = _cell_text(row[_COL_VENDOR] if len(row) > _COL_VENDOR else None)
            # Position rows: column B; fallback to last seen B value.
            ul_name = ul_cell or current_ul_name
            raw.append(
                _RawZinPosition(
                    title=title,
                    system=system,
                    code=code,
                    ul_name=ul_name,
                    qty=qty,
                    units=normalize_units_text(units_raw),
                    name=name,
                    tags=parse_tags(tags_raw),
                    spec=spec,
                    ds_num=ds_num,
                    vendor=vendor,
                    has_title_system=has_ts,
                )
            )
            stats.zin_positions += 1
    finally:
        wb.close()

    return raw, stats


def aggregate_zin_positions(
    raw_rows: list[_RawZinPosition],
    *,
    include_ul_in_key: bool,
) -> dict[tuple[str, str, str, str], ZinAgg]:
    """Aggregate raw Zinoviev positions by compare key."""
    agg: dict[tuple[str, str, str, str], ZinAgg] = {}
    for item in raw_rows:
        ul_for_key = item.ul_name if include_ul_in_key else ""
        key = ul_compare_key(item.title, item.system, item.code, ul_for_key)
        bucket = agg.get(key)
        if bucket is None:
            bucket = ZinAgg(
                title=item.title,
                system=item.system,
                code=item.code,
                has_title_system=item.has_title_system,
            )
            agg[key] = bucket
        bucket.qty += item.qty
        bucket.source_rows += 1
        if item.units:
            bucket.units.add(item.units)
        if item.name:
            bucket.names.add(item.name)
        bucket.tags |= item.tags
        if item.spec:
            bucket.specs.add(item.spec)
        if item.ds_num:
            bucket.ds_nums.add(item.ds_num)
        if item.vendor:
            bucket.vendors.add(item.vendor)
        if item.ul_name:
            bucket.ul_names.add(item.ul_name)
    return agg


def load_zinoviev_positions(
    xlsx_path: str | Path,
    *,
    include_ul_in_key: bool = True,
) -> tuple[dict[tuple[str, str, str, str], ZinAgg], ZinCompareStats]:
    """Load and aggregate Zinoviev positions (convenience wrapper).

    Args:
        xlsx_path: Path to ``ТСД по всем ДС_общий.xlsx``.
        include_ul_in_key: If True, key includes packing-list name.

    Returns:
        Aggregated lookup and load counters.
    """
    raw, stats = load_zinoviev_raw_positions(xlsx_path)
    return aggregate_zin_positions(raw, include_ul_in_key=include_ul_in_key), stats


def _robot_row_qty(row: RowStd) -> float | None:
    return _parse_qty(row.get_value(VALUES))


def aggregate_robot_positions(
    rows: list[RowStd],
    *,
    include_ul_in_key: bool = True,
) -> dict[tuple[str, str, str, str], RobotAgg]:
    """Aggregate robot packing ``position_row`` list by match key."""
    agg: dict[tuple[str, str, str, str], RobotAgg] = {}
    for row in rows:
        if row.row_type != RowType.position_row:
            continue
        code = _cell_text(row.get_value(CODE))
        if not code:
            continue
        qty = _robot_row_qty(row)
        if qty is None:
            continue
        title = _cell_text(row.get_value(DS_TITLE))
        system = _cell_text(row.get_value(DS_SYSTEM))
        has_ts = bool(title and system)
        source_rel = _cell_text(row.get_value(ANNOTATION))
        ul_display = packing_list_basename(source_rel)
        ul_for_key = source_rel if include_ul_in_key else ""
        key = ul_compare_key(title, system, code, ul_for_key)
        bucket = agg.get(key)
        if bucket is None:
            bucket = RobotAgg(
                title=title,
                system=system,
                code=code,
                has_title_system=has_ts,
            )
            agg[key] = bucket
        bucket.qty += qty
        bucket.source_rows += 1
        units = normalize_units_text(row.get_value(UNITS))
        if units:
            bucket.units.add(units)
        name = _cell_text(row.get_value(NAME))
        if name:
            bucket.names.add(name)
        tags_el = row.el.get(TAGS)
        bucket.tags |= parse_tags(tags_el.value if tags_el is not None else None)
        if SPECIFICATION_NAME in row.el:
            spec = _cell_text(row.get_value(SPECIFICATION_NAME))
            if spec:
                bucket.specs.add(spec)
        sheet = _cell_text(row.get_value(TITLE))
        if source_rel or sheet:
            bucket.sources.add(f"{source_rel} · {sheet}".strip(" ·"))
        if ul_display:
            bucket.ul_names.add(ul_display)
    return agg


def _qty_close(a: float, b: float, *, eps: float = 1e-6) -> bool:
    return abs(a - b) <= eps


def _units_joined(units: set[str]) -> str:
    return "/".join(sorted(units)) if units else ""


def compare_aggregates(
    robot: dict[tuple[str, str, str, str], RobotAgg],
    zin: dict[tuple[str, str, str, str], ZinAgg],
    *,
    robot_available: bool,
    stats: ZinCompareStats,
) -> list[CompareOutRow]:
    """Build compare rows for the union of robot and Zinoviev keys."""
    out: list[CompareOutRow] = []
    all_keys = sorted(set(robot) | set(zin))
    stats.keys_total = len(all_keys)
    stats.zin_keys = len(zin)

    for key in all_keys:
        r = robot.get(key)
        z = zin.get(key)
        title = (r.title if r else "") or (z.title if z else "")
        system = (r.system if r else "") or (z.system if z else "")
        code = (r.code if r else "") or (z.code if z else "") or key[2]
        ul_robot = _join_unique(r.ul_names if r else set(), limit=6)
        ul_zin = _join_unique(z.ul_names if z else set(), limit=6)
        # Prefer Zinoviev label(s); else robot basename(s); else key part.
        ul_name = ul_zin or ul_robot or key[3]
        title_mark = f"{title}-{system}" if title or system else ""
        has_ts = (r.has_title_system if r else False) or (
            z.has_title_system if z else False
        )

        robot_qty = r.qty if r is not None else None
        zin_qty = z.qty if z is not None else None
        units_robot = _units_joined(r.units) if r else ""
        units_zin = _units_joined(z.units) if z else ""
        name = _join_unique(
            (r.names if r else set()) | (z.names if z else set()),
            limit=3,
        )
        robot_tags = set(r.tags) if r else set()
        zin_tags = set(z.tags) if z else set()
        tags_robot_s = _format_tags(robot_tags)
        tags_zin_s = _format_tags(zin_tags)
        tags_diff = _format_tags_diff(robot_tags, zin_tags)
        if tags_diff:
            stats.tags_differ += 1
            tags_diff_color = Color.yellow
        else:
            stats.tags_equal += 1
            tags_diff_color = Color.no

        zin_coverage = _zin_coverage(
            has_robot=r is not None and robot_available,
            has_zin=z is not None,
            robot_qty=robot_qty,
            zin_qty=zin_qty,
            robot_tags=robot_tags,
            zin_tags=zin_tags,
        )
        if z is not None:
            if zin_coverage == _COV_FULL:
                stats.zin_covered += 1
            elif zin_coverage == _COV_MISSING:
                stats.zin_not_in_robot += 1
            if zin_coverage in (_COV_QTY, _COV_QTY_TAGS):
                stats.zin_qty_short += 1
            if zin_coverage in (_COV_TAGS, _COV_QTY_TAGS):
                stats.zin_tags_missing += 1

        spec_zin = _join_unique(z.specs if z else set(), limit=3)
        ds_num = _join_unique(z.ds_nums if z else set(), limit=6)
        robot_sources = _join_unique(r.sources if r else set(), limit=4)
        robot_rows = r.source_rows if r else 0
        zin_rows = z.source_rows if z else 0

        if not robot_available and z is not None and r is None:
            status = _STATUS_ROBOT_UNAVAILABLE
            stats.zin_only += 1
            diff = None
        elif not has_ts and normalize_packing_key_part(title) == "":
            status = _STATUS_NO_TITLE
            stats.no_title += 1
            diff = (
                (robot_qty or 0.0) - (zin_qty or 0.0)
                if robot_qty is not None or zin_qty is not None
                else None
            )
        elif r is not None and z is None:
            status = _STATUS_ROBOT_ONLY
            stats.robot_only += 1
            diff = robot_qty
        elif z is not None and r is None:
            status = _STATUS_ZIN_ONLY
            stats.zin_only += 1
            diff = -(zin_qty or 0.0)
        else:
            assert r is not None and z is not None
            diff = r.qty - z.qty
            if _qty_close(r.qty, z.qty):
                if units_robot and units_zin and units_robot != units_zin:
                    status = _STATUS_UNITS_DIFF
                    stats.units_diff += 1
                else:
                    status = _STATUS_MATCH
                    stats.match += 1
            else:
                status = _STATUS_QTY_DIFF
                stats.qty_diff += 1

        out.append(
            CompareOutRow(
                title=title,
                system=system,
                title_mark=title_mark,
                code=code,
                ul_name=ul_name,
                ul_robot=ul_robot,
                ul_zin=ul_zin,
                status=status,
                color=_STATUS_COLORS.get(status, Color.no),
                zin_coverage=zin_coverage,
                coverage_color=_COVERAGE_COLORS.get(zin_coverage, Color.no),
                robot_qty=robot_qty,
                zin_qty=zin_qty,
                diff=diff,
                units_robot=units_robot,
                units_zin=units_zin,
                name=name,
                tags_robot=tags_robot_s,
                tags_zin=tags_zin_s,
                tags_diff=tags_diff,
                tags_diff_color=tags_diff_color,
                spec_zin=spec_zin,
                ds_num_zin=ds_num,
                robot_sources=robot_sources,
                robot_rows=robot_rows,
                zin_rows=zin_rows,
            )
        )

    # Prefer incomplete Zinoviev coverage first, then qty status.
    priority = {
        _STATUS_NO_TITLE: 0,
        _STATUS_ROBOT_UNAVAILABLE: 0,
        _STATUS_QTY_DIFF: 1,
        _STATUS_UNITS_DIFF: 2,
        _STATUS_ROBOT_ONLY: 3,
        _STATUS_ZIN_ONLY: 4,
        _STATUS_MATCH: 5,
    }
    cov_priority = {
        _COV_MISSING: 0,
        _COV_QTY_TAGS: 1,
        _COV_QTY: 2,
        _COV_TAGS: 3,
        _COV_FULL: 4,
        _COV_NA: 5,
    }
    out.sort(
        key=lambda row: (
            cov_priority.get(row.zin_coverage, 9),
            priority.get(row.status, 9),
            row.ul_name.casefold(),
            row.title_mark.casefold(),
            row.code.casefold(),
        )
    )
    return out


def _result_dir(parent: Path, *, when: datetime | None = None) -> Path:
    stamp = (when or datetime.now()).strftime(_RESULT_FOLDER_FMT)
    path = parent / f"{_RESULT_FOLDER_PREFIX}{stamp}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def build_ul_name_inventory(
    zin_raw: list[_RawZinPosition],
    robot_rows: list[RowStd],
) -> list[UlNameOutRow]:
    """Match unique packing-list names from Zinoviev (col B) vs robot files.

    Pairing uses ``normalize_packing_list_name`` (revision suffixes stripped).
    Robot display uses the extracted UL id (same style as Zinoviev); full
    filename is kept in ``ul_robot_file``.
    """
    zin_names: dict[str, set[str]] = {}
    zin_counts: dict[str, int] = {}
    for item in zin_raw:
        display = item.ul_name.strip() if item.ul_name else ""
        if not display:
            display = "(пусто)"
        key = normalize_packing_list_name(display) or normalize_packing_key_part(display)
        zin_names.setdefault(key, set()).add(display)
        zin_counts[key] = zin_counts.get(key, 0) + 1

    robot_canon: dict[str, set[str]] = {}
    robot_files: dict[str, set[str]] = {}
    robot_counts: dict[str, int] = {}
    for row in robot_rows:
        if row.row_type != RowType.position_row:
            continue
        source_rel = _cell_text(row.get_value(ANNOTATION))
        file_name = packing_list_basename(source_rel)
        if not file_name:
            continue
        key = normalize_packing_list_name(source_rel) or normalize_packing_key_part(
            file_name
        )
        robot_canon.setdefault(key, set()).add(canonical_packing_list_label(source_rel))
        robot_files.setdefault(key, set()).add(file_name)
        robot_counts[key] = robot_counts.get(key, 0) + 1

    out: list[UlNameOutRow] = []
    for key in sorted(set(zin_names) | set(robot_canon), key=str.casefold):
        z_set = zin_names.get(key, set())
        r_set = robot_canon.get(key, set())
        f_set = robot_files.get(key, set())
        z_text = _join_unique(z_set, limit=8)
        r_text = _join_unique(r_set, limit=8)
        f_text = _join_unique(f_set, limit=6)
        if z_set and r_set:
            z_cf = {x.casefold() for x in z_set}
            r_cf = {x.casefold() for x in r_set}
            same_display = bool(z_cf & r_cf)
            if not same_display:
                # Zinoviev id contained in robot file / canonical label.
                for z in z_set:
                    z_low = z.casefold()
                    if any(z_low in r.casefold() for r in r_set | f_set):
                        same_display = True
                        break
            status = _UL_STATUS_MATCH if same_display else _UL_STATUS_SPELLING
        elif z_set:
            status = _UL_STATUS_ZIN_ONLY
        else:
            status = _UL_STATUS_ROBOT_ONLY
        out.append(
            UlNameOutRow(
                ul_zin=z_text,
                ul_robot=r_text,
                status=status,
                color=_UL_STATUS_COLORS.get(status, Color.no),
                norm_key=key,
                ul_robot_file=f_text,
                zin_rows=zin_counts.get(key, 0),
                robot_rows=robot_counts.get(key, 0),
            )
        )

    priority = {
        _UL_STATUS_SPELLING: 0,
        _UL_STATUS_ZIN_ONLY: 1,
        _UL_STATUS_ROBOT_ONLY: 2,
        _UL_STATUS_MATCH: 3,
    }
    out.sort(
        key=lambda row: (
            priority.get(row.status, 9),
            row.ul_zin.casefold(),
            row.ul_robot.casefold(),
        )
    )
    return out


def _write_sheet_header(
    wb: Any,
    ws: Any,
    columns: list[tuple[str, str, str, int]],
) -> dict[str, Any]:
    """Write colored header row and column widths; return header formats."""
    header_fmts: dict[str, Any] = {}
    for section, fill in _HEADER_FILL.items():
        header_fmts[section] = wb.add_format(
            {
                "bold": True,
                "font_color": "#000000",
                "bg_color": fill,
                "align": "center",
                "valign": "vcenter",
                "border": 1,
                "text_wrap": True,
            }
        )
    ws.set_row(0, 36)
    for ci, (_key, label, section, width) in enumerate(columns):
        ws.write(0, ci, label, header_fmts[section])
        ws.set_column(ci, ci, width)
    return header_fmts


def save_compare_excel(
    rows: list[CompareOutRow],
    out_path: Path,
    *,
    ul_inventory: list[UlNameOutRow] | None = None,
) -> Path:
    """Write comparison workbook (headers, autofilter, freeze, status colors).

    Sheet ``Робот vs Зиновьев`` — position compare.
    Sheet ``УЛ имена`` — unique packing-list names from both sides (optional).
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)

    columns: list[tuple[str, str, str, int]] = [
        # key, label, section, width
        ("title", "Титул", "key", 10),
        ("system", "Марка", "key", 10),
        ("title_mark", "Титул-Марка", "key", 16),
        ("code", "Код РД", "key", 15),
        ("ul_name", "Упаковочный лист", "key", 28),
        ("ul_robot", "УЛ робот (файл)", "robot", 36),
        ("ul_zin", "УЛ Зиновьев (№)", "zin", 22),
        ("zin_coverage", "Покрытие Зиновьева", "status", 22),
        ("status", "Статус кол-ва", "status", 28),
        ("robot_qty", "Кол-во робот", "robot", 12),
        ("zin_qty", "Кол-во Зиновьев", "zin", 13),
        ("diff", "Робот − Зиновьев", "status", 14),
        ("units_robot", "Ед. изм. робот", "robot", 12),
        ("units_zin", "Ед. изм. Зиновьев", "zin", 13),
        ("name", "Наименование", "meta", 36),
        ("tags_robot", "Теги робот", "robot", 32),
        ("tags_zin", "Теги Зиновьев", "zin", 32),
        ("tags_diff", "Разница тегов", "status", 40),
        ("spec_zin", "Спецификация (Зиновьев)", "zin", 36),
        ("ds_num_zin", "№ ДС (Зиновьев)", "zin", 14),
        ("robot_rows", "Строк робот", "robot", 11),
        ("zin_rows", "Строк Зиновьев", "zin", 12),
        ("robot_sources", "Источник робот (файл · вкладка)", "robot", 42),
    ]

    wb = xlsxwriter.Workbook(str(out_path), {"constant_memory": False})
    ws = wb.add_worksheet("Робот vs Зиновьев")
    _write_sheet_header(wb, ws, columns)

    base_cell = {
        "align": "left",
        "valign": "vcenter",
        "border": 1,
        "text_wrap": True,
    }
    cell_fmts: dict[str, Any] = {Color.no: wb.add_format(base_cell)}
    num_fmts: dict[str, Any] = {
        Color.no: wb.add_format({**base_cell, "align": "right", "num_format": "0.####"})
    }

    def _fmt(color: str, *, numeric: bool = False) -> Any:
        store = num_fmts if numeric else cell_fmts
        if color not in store:
            opts = dict(base_cell)
            if color and color != Color.no:
                opts["bg_color"] = f"#{color}" if not color.startswith("#") else color
            if numeric:
                opts["align"] = "right"
                opts["num_format"] = "0.####"
                opts.pop("text_wrap", None)
            store[color] = wb.add_format(opts)
        return store[color]

    qty_keys = {"robot_qty", "zin_qty", "diff", "robot_rows", "zin_rows"}
    for row_idx, row in enumerate(rows):
        excel_row = row_idx + 1
        color = row.color if row.color != Color.no else Color.no
        values = {
            "title": row.title,
            "system": row.system,
            "title_mark": row.title_mark,
            "code": row.code,
            "ul_name": row.ul_name,
            "ul_robot": row.ul_robot,
            "ul_zin": row.ul_zin,
            "zin_coverage": row.zin_coverage,
            "status": row.status,
            "robot_qty": row.robot_qty,
            "zin_qty": row.zin_qty,
            "diff": row.diff,
            "units_robot": row.units_robot,
            "units_zin": row.units_zin,
            "name": row.name,
            "tags_robot": row.tags_robot,
            "tags_zin": row.tags_zin,
            "tags_diff": row.tags_diff,
            "spec_zin": row.spec_zin,
            "ds_num_zin": row.ds_num_zin,
            "robot_rows": row.robot_rows or None,
            "zin_rows": row.zin_rows or None,
            "robot_sources": row.robot_sources,
        }
        for ci, (key, _label, _section, _width) in enumerate(columns):
            value = values[key]
            if key in qty_keys:
                if value is None or value == "":
                    ws.write(excel_row, ci, "", _fmt(color))
                else:
                    ws.write_number(excel_row, ci, float(value), _fmt(color, numeric=True))
            elif key == "status":
                ws.write(excel_row, ci, value, _fmt(color))
            elif key == "zin_coverage":
                ws.write(
                    excel_row,
                    ci,
                    value,
                    _fmt(row.coverage_color if row.coverage_color != Color.no else Color.no),
                )
            elif key == "tags_diff":
                ws.write(
                    excel_row,
                    ci,
                    value if value else "",
                    _fmt(row.tags_diff_color if value else Color.no),
                )
            else:
                ws.write(excel_row, ci, value if value is not None else "", _fmt(color))

    last = len(rows)
    if last > 0:
        ws.autofilter(0, 0, last, len(columns) - 1)
    ws.freeze_panes(1, 0)

    if ul_inventory is not None:
        ul_columns: list[tuple[str, str, str, int]] = [
            ("ul_zin", "УЛ Зиновьев (столбец B)", "zin", 36),
            ("ul_robot", "УЛ робот (id)", "robot", 36),
            ("status", "Статус", "status", 28),
            ("ul_robot_file", "УЛ робот (файл)", "robot", 48),
            ("norm_key", "Ключ нормализации", "key", 28),
            ("zin_rows", "Строк Зиновьев", "zin", 14),
            ("robot_rows", "Строк робот", "robot", 12),
        ]
        ws_ul = wb.add_worksheet("УЛ имена")
        _write_sheet_header(wb, ws_ul, ul_columns)
        for row_idx, row in enumerate(ul_inventory):
            excel_row = row_idx + 1
            color = row.color if row.color != Color.no else Color.no
            values = {
                "ul_zin": row.ul_zin,
                "ul_robot": row.ul_robot,
                "status": row.status,
                "ul_robot_file": row.ul_robot_file,
                "norm_key": row.norm_key,
                "zin_rows": row.zin_rows or None,
                "robot_rows": row.robot_rows or None,
            }
            for ci, (key, _label, _section, _width) in enumerate(ul_columns):
                value = values[key]
                if key in {"zin_rows", "robot_rows"}:
                    if value is None:
                        ws_ul.write(excel_row, ci, "", _fmt(color))
                    else:
                        ws_ul.write_number(
                            excel_row, ci, float(value), _fmt(color, numeric=True)
                        )
                elif key == "status":
                    ws_ul.write(excel_row, ci, value, _fmt(color))
                else:
                    ws_ul.write(excel_row, ci, value or "", _fmt(color))
        ul_last = len(ul_inventory)
        if ul_last > 0:
            ws_ul.autofilter(0, 0, ul_last, len(ul_columns) - 1)
        ws_ul.freeze_panes(1, 0)

    wb.close()
    return out_path


def _quality_ru(quality: str) -> str:
    mapping = {
        "ok": "ок (замечаний нет)",
        "partial": "частичный (есть замечания по исходным УЛ)",
        "unavailable": "недоступен (нужно перечитать УЛ)",
    }
    return mapping.get(quality, quality or "неизвестно")


def _write_report(
    path: Path,
    *,
    zin_path: Path,
    robot_quality: str,
    stats: ZinCompareStats,
    excel_path: Path,
    include_ul_in_key: bool,
) -> None:
    """Write a Russian human-readable summary next to the Excel result."""
    covered_pct = (
        (100.0 * stats.zin_covered / stats.zin_keys) if stats.zin_keys else 0.0
    )
    if stats.zin_keys == 0:
        verdict = "Нет ключей Зиновьева для оценки покрытия."
    elif stats.zin_covered == stats.zin_keys:
        verdict = (
            "Робот полностью покрывает свод Зиновьева "
            "(по количеству и тегам на всех ключах)."
        )
    else:
        gaps = stats.zin_keys - stats.zin_covered
        verdict = (
            f"Робот покрывает {stats.zin_covered} из {stats.zin_keys} ключей "
            f"Зиновьева ({covered_pct:.1f}%). "
            f"Не полностью покрыто: {gaps}."
        )

    if include_ul_in_key:
        key_title = "Титул + Марка + Код РД + название УЛ"
        key_hint = (
            "(вкладка/лист пока не входит в ключ; имя УЛ у Зиновьева — "
            "столбец B, у робота — basename файла / PL_… / AN######)."
        )
        keys_label = "Уникальных ключей (титул+марка+код+УЛ):"
        mode_title = "режим с названием упаковочного листа"
    else:
        key_title = "Титул + Марка + Код РД"
        key_hint = (
            "(название УЛ и вкладка не входят в ключ — количества/теги "
            "суммируются по всем УЛ с одним титулом/маркой/кодом)."
        )
        keys_label = "Уникальных ключей (титул+марка+код):"
        mode_title = "режим без названия УЛ (сумма по всем листам)"

    lines = [
        "Сравнение свода УЛ: робот ↔ Зиновьев",
        f"Вариант: {mode_title}",
        "=" * 52,
        "",
        "Исходные данные",
        "-" * 52,
        f"Файл Зиновьева:  {zin_path}",
        f"Кэш робота:      {_quality_ru(robot_quality)}",
        f"Excel-результат: {excel_path}",
        "",
        "Загрузка Зиновьева",
        "-" * 52,
        f"Строк в файле (всего просмотрено):     {stats.zin_raw_rows}",
        f"Позиций (после отсева мусора):         {stats.zin_positions}",
        f"Пропущено (баннеры/заголовки/мусор):   {stats.zin_skipped}",
        f"Позиций без титула/марки в спеке:      {stats.zin_no_title}",
        f"{keys_label:<40} {stats.zin_keys}",
        "",
        "Загрузка робота",
        "-" * 52,
        f"Позиций в кэше (сумма по ключам):      {stats.robot_positions}",
        f"Ключей в сравнении (объединение):      {stats.keys_total}",
        "",
        f"Ключ сопоставления: {key_title}",
        key_hint,
        "",
        "Покрытие данных Зиновьева роботом",
        "-" * 52,
        f"Покрыто полностью:                     {stats.zin_covered}",
        f"Нет позиции в роботе:                  {stats.zin_not_in_robot}",
        f"Недобор количества:                    {stats.zin_qty_short}",
        f"Не хватает тегов у робота:             {stats.zin_tags_missing}",
        f"Доля полного покрытия:                 {covered_pct:.1f}%",
        "",
        "Сравнение количества (по ключам)",
        "-" * 52,
        f"Совпадение количества:                 {stats.match}",
        f"Расхождение количества:                {stats.qty_diff}",
        f"Совпадение, но разные ед. изм.:        {stats.units_diff}",
        f"Только в роботе:                       {stats.robot_only}",
        f"Только у Зиновьева:                    {stats.zin_only}",
        f"Без титула/марки:                      {stats.no_title}",
        "",
        "Сравнение тегов (по ключам)",
        "-" * 52,
        f"Наборы тегов совпадают:                {stats.tags_equal}",
        f"Наборы тегов различаются:              {stats.tags_differ}",
        "",
        "Вывод",
        "-" * 52,
        verdict,
        "",
        "Подсказка: в Excel смотрите столбец «Покрытие Зиновьева»",
        "и фильтр ≠ «Покрыто» — там все дыры относительно ручного свода.",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def _fresh_compare_stats(base: ZinCompareStats) -> ZinCompareStats:
    """Copy load-side counters; reset compare counters for a second key mode."""
    return ZinCompareStats(
        zin_raw_rows=base.zin_raw_rows,
        zin_positions=base.zin_positions,
        zin_skipped=base.zin_skipped,
        zin_no_title=base.zin_no_title,
        robot_positions=base.robot_positions,
        robot_quality=base.robot_quality,
    )


def _run_one_key_mode(
    *,
    robot_rows: list[RowStd],
    zin_raw: list[_RawZinPosition],
    load_stats: ZinCompareStats,
    robot_available: bool,
    include_ul_in_key: bool,
    result_dir: Path,
    excel_name: str,
    report_name: str,
    zin_path: Path,
    ul_inventory: list[UlNameOutRow] | None = None,
) -> tuple[ZinCompareStats, Path, Path]:
    """Aggregate, compare, write Excel+report for one key mode."""
    stats = _fresh_compare_stats(load_stats)
    zin_agg = aggregate_zin_positions(zin_raw, include_ul_in_key=include_ul_in_key)
    robot_agg: dict[tuple[str, str, str, str], RobotAgg] = {}
    if robot_available:
        robot_agg = aggregate_robot_positions(
            robot_rows, include_ul_in_key=include_ul_in_key
        )
        stats.robot_positions = sum(b.source_rows for b in robot_agg.values())

    rows = compare_aggregates(
        robot_agg,
        zin_agg,
        robot_available=robot_available,
        stats=stats,
    )
    excel_path = result_dir / excel_name
    save_compare_excel(rows, excel_path, ul_inventory=ul_inventory)
    report_path = result_dir / report_name
    _write_report(
        report_path,
        zin_path=zin_path,
        robot_quality=stats.robot_quality,
        stats=stats,
        excel_path=excel_path,
        include_ul_in_key=include_ul_in_key,
    )
    return stats, excel_path, report_path


def compare_robot_vs_zinoviev(
    zinoviev_path: str | Path | None = None,
    *,
    output_parent: str | Path | None = None,
) -> ZinCompareResult:
    """Compare robot packing cache with Zinoviev manual summary.

    Writes **two** Excel reports and two text reports into one timestamped
    folder:

    - without UL in key (титул+марка+код);
    - with packing-list name in key (титул+марка+код+УЛ).

    Each Excel also has sheet ``УЛ имена`` (Zinoviev col B vs robot files).

    Args:
        zinoviev_path: Manual workbook; default ``DEFAULT_ZINOVIEV_XLSX``.
        output_parent: Parent for ``_результат_сравнения_*`` folder;
            default ``DEFAULT_ZINOVIEV_DIR``.

    Returns:
        Paths and counters (``stats`` = variant **without** UL in key).
        ``excel_path`` points to the without-UL workbook.
    """
    zin_path = Path(zinoviev_path) if zinoviev_path else DEFAULT_ZINOVIEV_XLSX
    parent = Path(output_parent) if output_parent else DEFAULT_ZINOVIEV_DIR

    print(f"Зиновьев: чтение {zin_path}")
    zin_raw, load_stats = load_zinoviev_raw_positions(zin_path)
    print(
        f"Зиновьев: позиций={load_stats.zin_positions}, "
        f"пропущено={load_stats.zin_skipped}, "
        f"без титула/марки={load_stats.zin_no_title}"
    )

    dataset = load_packing_dataset()
    load_stats.robot_quality = dataset.quality.value
    robot_available = dataset.quality != PackingQualityLevel.UNAVAILABLE
    robot_rows = dataset.rows if robot_available else []
    if robot_available:
        print(
            f"Робот: quality={dataset.quality.value}, "
            f"position_rows={len(dataset.rows)}"
        )
    else:
        print(
            f"Робот: кэш недоступен ({dataset.format_short()}). "
            "Сравнение только со стороны Зиновьева."
        )

    ul_inventory = build_ul_name_inventory(
        zin_raw, robot_rows if robot_available else []
    )
    ul_match = sum(1 for r in ul_inventory if r.status == _UL_STATUS_MATCH)
    ul_spell = sum(1 for r in ul_inventory if r.status == _UL_STATUS_SPELLING)
    ul_zin_only = sum(1 for r in ul_inventory if r.status == _UL_STATUS_ZIN_ONLY)
    ul_robot_only = sum(1 for r in ul_inventory if r.status == _UL_STATUS_ROBOT_ONLY)
    print(
        f"УЛ имена: всего={len(ul_inventory)}, совпадение={ul_match}, "
        f"разное написание={ul_spell}, "
        f"только Зиновьев={ul_zin_only}, только робот={ul_robot_only}"
    )

    result_dir = _result_dir(parent)

    print("Сравнение 1/2: ключ титул+марка+код (без УЛ)…")
    stats_code, excel_code, _report_code = _run_one_key_mode(
        robot_rows=robot_rows,
        zin_raw=zin_raw,
        load_stats=load_stats,
        robot_available=robot_available,
        include_ul_in_key=False,
        result_dir=result_dir,
        excel_name=_RESULT_XLSX_BY_CODE,
        report_name=_RESULT_REPORT_BY_CODE,
        zin_path=zin_path,
        ul_inventory=ul_inventory,
    )
    pct_code = (
        (100.0 * stats_code.zin_covered / stats_code.zin_keys)
        if stats_code.zin_keys
        else 0.0
    )
    print(
        f"  без УЛ: ключей={stats_code.keys_total}, "
        f"покрытие={stats_code.zin_covered}/{stats_code.zin_keys} ({pct_code:.1f}%)"
    )

    print("Сравнение 2/2: ключ титул+марка+код+УЛ…")
    stats_ul, excel_ul, _report_ul = _run_one_key_mode(
        robot_rows=robot_rows,
        zin_raw=zin_raw,
        load_stats=load_stats,
        robot_available=robot_available,
        include_ul_in_key=True,
        result_dir=result_dir,
        excel_name=_RESULT_XLSX_BY_UL,
        report_name=_RESULT_REPORT_BY_UL,
        zin_path=zin_path,
        ul_inventory=ul_inventory,
    )
    pct_ul = (
        (100.0 * stats_ul.zin_covered / stats_ul.zin_keys) if stats_ul.zin_keys else 0.0
    )
    print(
        f"  с УЛ: ключей={stats_ul.keys_total}, "
        f"покрытие={stats_ul.zin_covered}/{stats_ul.zin_keys} ({pct_ul:.1f}%)"
    )

    msg = (
        f"Два отчёта в {result_dir}. "
        f"Без УЛ: покрытие {stats_code.zin_covered}/{stats_code.zin_keys} "
        f"({pct_code:.1f}%), совпадение кол-ва={stats_code.match}. "
        f"С УЛ: покрытие {stats_ul.zin_covered}/{stats_ul.zin_keys} "
        f"({pct_ul:.1f}%), совпадение кол-ва={stats_ul.match}. "
        f"УЛ имена: совп.={ul_match}, разн.напис.={ul_spell}, "
        f"только Зин.={ul_zin_only}, только робот={ul_robot_only}. "
        f"Файлы: {excel_code.name}, {excel_ul.name} (лист «УЛ имена»)."
    )
    print(msg)
    success = robot_available or load_stats.zin_positions > 0
    return ZinCompareResult(
        success=success,
        message=msg,
        result_dir=str(result_dir),
        excel_path=str(excel_code),
        stats=stats_code,
    )


__all__ = [
    "DEFAULT_ZINOVIEV_DIR",
    "DEFAULT_ZINOVIEV_XLSX",
    "ZinCompareResult",
    "ZinCompareStats",
    "build_ul_name_inventory",
    "compare_robot_vs_zinoviev",
    "load_zinoviev_positions",
    "normalize_packing_list_name",
    "parse_tags",
    "ul_compare_key",
]
