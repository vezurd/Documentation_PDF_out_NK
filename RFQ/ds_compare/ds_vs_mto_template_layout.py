"""
Read header labels (single-line), column widths, and header fill colors from
``templates/ds_vs_mto_template.xlsx`` so xlsxwriter export matches the template.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict

import openpyxl
from openpyxl.utils import get_column_letter

from base.tables_columns import ColNames


@dataclass(frozen=True)
class TemplateColLayout:
    """Per-column layout from row 1 of the DS vs MTO template."""

    header_one_line: str
    width: float
    header_bg_rgb6: str  # RRGGBB, no leading # (e.g. 92D050)


def _header_to_one_line(value: object) -> str:
    """Collapse newlines and repeated spaces (one column = one header line in Excel)."""
    if value is None:
        return ""
    t = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return " ".join(t.split()).strip()


def _fill_rgb6_from_cell(cell: object) -> str | None:
    """Extract solid fill as 6 hex digits; openpyxl often uses ARGB like FF92D050."""
    f = cell.fill
    if f is None:
        return None
    for attr in ("fgColor", "start_color"):
        col = getattr(f, attr, None)
        if col is None or not getattr(col, "rgb", None):
            continue
        r = str(col.rgb)
        if len(r) == 8 and (r.startswith("FF") or r.startswith("ff")):
            return r[2:].upper()
        if len(r) == 6:
            return r.upper()
    return None


def read_ds_vs_mto_template_layout(
    template_path: Path | str,
) -> Dict[str, TemplateColLayout] | None:
    """
    Reads first row (headers), column widths, and header background from the template.

    Keys are ``ColNames.DsVsMto`` field names in table order (0..33).

    Returns:
        Map col_name -> TemplateColLayout, or None if the file is missing/unreadable.
    """
    path = Path(template_path)
    if not path.is_file():
        return None

    order = [ColNames.DsVsMto.column_dict[k] for k in sorted(ColNames.DsVsMto.column_dict)]
    try:
        wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    except Exception:
        return None
    try:
        ws = wb.active
        out: dict[str, TemplateColLayout] = {}
        for i, col_name in enumerate(order, start=1):
            cell = ws.cell(1, i)
            letter = get_column_letter(i)
            wdim = ws.column_dimensions.get(letter)
            w = float(wdim.width) if wdim and wdim.width is not None else 14.0
            raw = cell.value
            header_1 = _header_to_one_line(raw) if raw is not None else ""
            if not header_1:
                header_1 = col_name
            bg = _fill_rgb6_from_cell(cell) or "D9D9D9"
            out[col_name] = TemplateColLayout(
                header_one_line=header_1,
                width=w,
                header_bg_rgb6=bg,
            )
        return out
    finally:
        wb.close()
