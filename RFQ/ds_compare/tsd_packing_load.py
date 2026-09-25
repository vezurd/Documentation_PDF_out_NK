"""Load TSD packing-list xlsx trees, validate sheets, cache RowStd + summary Excel."""

from __future__ import annotations

import hashlib
import os
import pickle
import re
import tempfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import xlsxwriter
from openpyxl.styles import Alignment, Font
from prettytable import PrettyTable

import base.t_comm_initial_classes as t_com_init_cls
import utils.path
from base.base_classes import CheckElement, RowStd, RowType, TableComments
from base.base_mto import trim_empty_rows_from_end
from base.base_xlsx_load import list_raw_to_std, load_rows_from_worksheet
from base.tables_columns import (
    ANNOTATION,
    CODE,
    DS_SYSTEM,
    DS_TITLE,
    NAME,
    SPECIFICATION_NAME,
    TAGS,
    TITLE,
    TYPE_MARK,
    UNITS,
    VALUES,
    VENDOR,
)
from RFQ.ds_compare.ds_quantity_parse import try_parse_quantity
from RFQ.packing_list_provider import (
    DEFAULT_PACKING_OUTPUT_DIR,
    PACKING_CACHE_NAME,
    PACKING_CACHE_VERSION,
    PackingIssue,
    PackingQualityLevel,
    load_packing_dataset,
    slim_packing_rows,
)
from utils.colors import Color
from utils.file_name_converts import ProjectFileName

# Deliverable folder (свод xlsx, отчёты, pickle) — сетевая «УЛ сводный файл».
DEFAULT_TSD_PACKING_OUTPUT_DIR = DEFAULT_PACKING_OUTPUT_DIR
_CACHE_DIR = DEFAULT_TSD_PACKING_OUTPUT_DIR
TSD_PACKING_CACHE_VERSION = PACKING_CACHE_VERSION
TSD_PACKING_ROWS_CACHE_NAME = PACKING_CACHE_NAME
# Canonical summary name: ``tsd_packing_summary_YYYY.MM.DD_HH.MM.SS.xlsx``
# (one helper below — do not hardcode the filename elsewhere).
TSD_PACKING_SUMMARY_STEM = "tsd_packing_summary"
TSD_PACKING_SUMMARY_TIMESTAMP_FMT = "%Y.%m.%d_%H.%M.%S"
TSD_PACKING_LOAD_REPORT_NAME = "tsd_packing_load_report.txt"
TSD_PACKING_SHEETS_REPORT_NAME = "tsd_packing_sheets_report.txt"
TSD_PACKING_CRITICAL_REPORT_NAME = "tsd_packing_critical.txt"

_last_tsd_critical_report: TsdCriticalReport | None = None

DEFAULT_TSD_PACKING_ROOT = (
    r"\\bcc\root\CurProjects\di_manegers\Пятницкий_П"
    r"\Амурский ГХК\Поставки\ТСД по всем ДС"
)
# Parallel file reads (I/O-bound UNC/xlsx). Threads keep RowStd in-process.
DEFAULT_TSD_LOAD_WORKERS = 8
_MIN_FILES_FOR_PARALLEL = 2

# Master 1 / Master 2 / … — skip; every other sheet is scanned for position_row.
_MASTER_SHEET_RE = re.compile(r"^Master\s*\d+\.?$", re.IGNORECASE)
# Wide SO - PL layout (госфин AGCC.323 / some PL-АГХК-*).
_SO_PL_SHEET_RE = re.compile(r"SO\s*-\s*PL", re.IGNORECASE)
_WHITESPACE_RE = re.compile(r"[\s\u00a0]+")
_LEADING_TITLE_RE = re.compile(r"^(\d{4})")

# Flat summary columns: (RowStd field or pseudo-key, header label, width).
# Widths aligned with step4_6_save_match_result_to_excel.py (CODE=15, TAGS=25,
# NAME=25..30, TYPE_MARK=20, VALUES=7, VENDOR=18); source/spec tuned to ТСД data.
_SUMMARY_COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("source_rel", "Файл", 45),
    ("sheet", "Вкладка", 12),
    ("source_excel_row", "Строка в исходном УЛ", 12),
    ("CODE", "Код РД", 15),
    ("SPECIFICATION_NAME", "Спецификация", 35),
    ("DS_TITLE", "Титул", 10),
    ("DS_SYSTEM", "Марка", 10),
    ("TAGS", "Теги", 25),
    ("NAME", "Наименование", 30),
    ("TYPE_MARK", "Тип / хар-ки", 20),
    ("VALUES", "Кол-во", 7),
    ("UNITS", "Ед. изм.", 8),
    ("VENDOR", "Поставщик", 18),
)
_SUMMARY_HEADER_FILL = "#dbbcdb"  # same as step4 RFP section
UL_FOLDER_STATS_SHEET = "Статистика по папкам"
_UL_FOLDER_STATS_HEADERS = (
    "Папка",
    "Всего строк",
    "Строк с кодом РД",
    "Строк с наименованием",
    "Строк с кол-вом",
)
_UL_FOLDER_STATS_NOTE = (
    "Число строк свода, в которых ячейка заполнена. "
    "Это не сумма колонки «Кол-во» и не число разных значений. "
    "Папка — каталог, в котором лежит файл УЛ."
)


@dataclass
class TsdSheetPlan:
    """Sheet classification for one workbook."""

    all_names: list[str]
    master_sheets: list[str]
    data_sheets: list[str]

    @property
    def can_load(self) -> bool:
        return bool(self.data_sheets)


@dataclass
class TsdSheetHit:
    """One workbook sheet and how many position rows it contributed."""

    source_rel: str
    sheet_name: str
    position_rows: int
    status: str  # loaded | skipped_master | load_error


@dataclass
class _TsdFileOutcome:
    """Result of loading one packing-list workbook (worker → main)."""

    source_rel: str
    folder: str
    ok: bool
    error: str | None = None
    positions: list[RowStd] = field(default_factory=list)
    all_rows: list[RowStd] = field(default_factory=list)
    data_sheets: list[str] = field(default_factory=list)
    hits: list[TsdSheetHit] = field(default_factory=list)
    row_issues: list[PackingIssue] = field(default_factory=list)


@dataclass
class TsdCriticalReport:
    """Human-readable critical remarks about input packing files.

    ``ok=True`` means no action needed on source files (detail reports optional).
    """

    ok: bool
    remarks: list[str] = field(default_factory=list)
    files_failed: int = 0
    files_without_position: int = 0
    format_issues: int = 0

    def format_short(self) -> str:
        """One-line verdict for status labels."""
        if self.ok:
            return "OK — критичных замечаний нет. Детальный отчёт можно не смотреть."
        parts: list[str] = [f"ЕСТЬ ЗАМЕЧАНИЯ ({len(self.remarks)})"]
        if self.files_failed:
            parts.append(f"ошибок файлов={self.files_failed}")
        if self.format_issues:
            parts.append(f"сдвиг/формат={self.format_issues}")
        if self.files_without_position:
            parts.append(f"без позиций={self.files_without_position}")
        parts.append("смотрите список ниже и отчёты в папке свода.")
        return " — ".join(parts)

    def format_detail(self) -> str:
        """Full text for GUI panel / critical report file."""
        lines = [self.format_short(), ""]
        if self.ok:
            lines.append("(замечаний нет)")
        else:
            for i, remark in enumerate(self.remarks, start=1):
                lines.append(f"{i}. {remark}")
                lines.append("")
        return "\n".join(lines).rstrip() + "\n"


@dataclass
class TsdLoadStats:
    """Aggregate counters for a TSD packing load run."""

    files_total: int = 0
    files_ok: int = 0
    files_multi_data_sheets: int = 0
    files_failed: int = 0
    position_rows: int = 0
    empty_rows: int = 0
    other_rows: int = 0
    no_title_system: int = 0
    files_without_position: int = 0
    folder_files: Counter[str] = field(default_factory=Counter)
    folder_positions: Counter[str] = field(default_factory=Counter)
    failures: list[str] = field(default_factory=list)
    sheet_hits: list[TsdSheetHit] = field(default_factory=list)
    row_issues: list[PackingIssue] = field(default_factory=list)


def tsd_packing_cache_dir() -> Path:
    """Return TSD packing output dir (свод/кэш/отчёты); create if missing.

    Default: ``\\\\bcc\\eng\\...\\_RFP\\УЛ сводный файл``.
    """
    _CACHE_DIR.mkdir(parents=True, exist_ok=True)
    return _CACHE_DIR


def tsd_packing_summary_xlsx_name(*, when: datetime | None = None) -> str:
    """Return canonical summary filename with local date/time suffix.

    Example: ``tsd_packing_summary_2026.07.29_16.54.37.xlsx``.
    """
    stamp = (when or datetime.now()).strftime(TSD_PACKING_SUMMARY_TIMESTAMP_FMT)
    return f"{TSD_PACKING_SUMMARY_STEM}_{stamp}.xlsx"


def tsd_packing_summary_xlsx_path(
    *,
    cache_dir: Path | None = None,
    when: datetime | None = None,
) -> Path:
    """Return full path for a new summary xlsx in the packing output dir."""
    return (cache_dir or tsd_packing_cache_dir()) / tsd_packing_summary_xlsx_name(
        when=when
    )


def find_latest_tsd_packing_summary(cache_dir: Path | None = None) -> Path | None:
    """Newest ``tsd_packing_summary*.xlsx`` in the output dir, or ``None``.

    Matches both timestamped names and the legacy fixed
    ``tsd_packing_summary.xlsx``.
    """
    directory = cache_dir or tsd_packing_cache_dir()
    if not directory.is_dir():
        return None
    files = [
        path
        for path in directory.glob(f"{TSD_PACKING_SUMMARY_STEM}*.xlsx")
        if path.is_file()
    ]
    if not files:
        return None
    return max(files, key=lambda path: path.stat().st_mtime)


def get_last_tsd_critical_report() -> TsdCriticalReport | None:
    """Return critical remarks from the last ``load_and_cache_tsd_packing`` run."""
    return _last_tsd_critical_report


def is_master_sheet(name: str) -> bool:
    """True when sheet name looks like Master 1 / Master 2 / …"""
    return bool(_MASTER_SHEET_RE.match(str(name).strip()))


def is_sopl_sheet(name: str) -> bool:
    """True when sheet uses the wide ``SO - PL`` column layout."""
    return bool(_SO_PL_SHEET_RE.search(str(name).strip()))


def collapse_all_whitespace(text: object) -> str:
    """Strip and remove all whitespace (for SPEC / title-system matching)."""
    if text is None:
        return ""
    return _WHITESPACE_RE.sub("", str(text).strip())


def normalize_spaces(text: object) -> str:
    """Strip edges and squeeze internal whitespace to a single space."""
    if text is None:
        return ""
    return _WHITESPACE_RE.sub(" ", str(text)).strip()


def classify_tsd_sheets(sheet_names: list[str]) -> TsdSheetPlan:
    """Split workbook sheets into Master* (skip) and all other data sheets."""
    masters: list[str] = []
    data: list[str] = []
    for name in sheet_names:
        if is_master_sheet(name):
            masters.append(name)
        else:
            data.append(name)
    return TsdSheetPlan(
        all_names=list(sheet_names),
        master_sheets=masters,
        data_sheets=data,
    )


def _find_std_code_rd_col(ws: Any) -> int | None:
    """Return 0-based column of STD «Код РД / Code RD» header, or None."""
    for row in ws.iter_rows(min_row=1, max_row=25, max_col=20, values_only=True):
        for j, value in enumerate(row):
            if value is None:
                continue
            text = str(value).upper().replace("\n", " ")
            if "CODE RD" in text or "КОД РД" in text:
                return j
    return None


def diagnose_std_column_offset(file_path: str, sheet_name: str) -> str | None:
    """Return action text when STD packing columns are shifted; else None."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return None
        code_col = _find_std_code_rd_col(wb[sheet_name])
    finally:
        wb.close()
    if code_col is None or code_col == 1:
        return None
    if code_col == 0:
        return (
            "Не правильный формат – надо вставить пустую колонку A "
            "(сейчас CODE стоит в A, пустая слева нет)"
        )
    if code_col == 2:
        return (
            "Не правильный формат – надо убрать колонку B "
            "(должна быть пустая только одна)"
        )
    return (
        "Не правильный формат – надо убрать лишние колонки слева от «Код РД» "
        f"(сейчас «Код РД» в столбце {code_col + 1}, должен быть в B; "
        "пустая только A)"
    )


def _absolute_source_path(root: str, source_rel: str) -> str:
    return str(Path(root) / source_rel)


def build_tsd_critical_report(
    stats: TsdLoadStats,
    *,
    root: str,
) -> TsdCriticalReport:
    """Build critical remarks from load stats (failures, zero files, offsets)."""
    remarks: list[str] = []
    format_issues = 0

    for fail in stats.failures:
        remarks.append(f"Ошибка чтения файла — {fail}")
    for issue in stats.row_issues:
        remarks.append(issue.format_line())

    # Files with zero position_row across all loaded sheets.
    positions_by_file: Counter[str] = Counter()
    sheets_by_file: dict[str, list[TsdSheetHit]] = {}
    for hit in stats.sheet_hits:
        if hit.status == "skipped_master":
            continue
        positions_by_file[hit.source_rel] += hit.position_rows
        sheets_by_file.setdefault(hit.source_rel, []).append(hit)

    for source_rel, n_pos in sorted(positions_by_file.items(), key=lambda x: x[0].lower()):
        if n_pos > 0:
            continue
        abs_path = _absolute_source_path(root, source_rel)
        sheet_notes: list[str] = []
        offset_found = False
        for hit in sheets_by_file.get(source_rel, []):
            if hit.status.startswith("load_error"):
                sheet_notes.append(f"{hit.sheet_name}: {hit.status}")
                continue
            if is_sopl_sheet(hit.sheet_name):
                sheet_notes.append(f"{hit.sheet_name}: 0 позиций (формат SO - PL)")
                continue
            # Prefer diagnosing Single* / main packing sheets.
            name_u = hit.sheet_name.strip().upper()
            if name_u.startswith("SINGLE") or name_u.startswith("SINGL"):
                try:
                    action = diagnose_std_column_offset(abs_path, hit.sheet_name)
                except Exception as exc:
                    sheet_notes.append(f"{hit.sheet_name}: не удалось проверить формат ({exc})")
                    continue
                if action:
                    remarks.append(f"{action}\n{abs_path}")
                    format_issues += 1
                    offset_found = True
                else:
                    sheet_notes.append(
                        f"{hit.sheet_name}: 0 позиций (шаблон столбцов похож на OK — нет строк данных?)"
                    )
        if not offset_found:
            detail = "; ".join(sheet_notes) if sheet_notes else "нет position_row"
            remarks.append(
                f"Нет позиций после чтения — {detail}\n{abs_path}"
            )

    report = TsdCriticalReport(
        ok=not remarks,
        remarks=remarks,
        files_failed=stats.files_failed,
        files_without_position=stats.files_without_position,
        format_issues=format_issues,
    )
    return report


def save_tsd_packing_critical_report(
    report: TsdCriticalReport,
    *,
    root: str,
    out_path: Path | None = None,
) -> Path:
    """Write critical remarks file under the TSD cache dir."""
    out = out_path or (tsd_packing_cache_dir() / TSD_PACKING_CRITICAL_REPORT_NAME)
    text = f"root: {root}\n\n{report.format_detail()}"
    out.write_text(text, encoding="utf-8")
    return out


def read_workbook_sheet_names(file_path: str) -> list[str]:
    """Return sheet names from an xlsx (read-only)."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _path_doc_factory(file_path: str) -> Any:
    name = os.path.basename(file_path)
    return type(
        "Doc",
        (),
        {
            "file_full_path": file_path,
            "file_name": name,
            "doc_Number_for_sort": name,
        },
    )()


def collect_tsd_files(root: str) -> list[Any]:
    """Recursive ``*.xlsx`` under *root* (skips ``~`` temps via get_files_single)."""
    return utils.path.get_files_single(
        root,
        endswith=(".xlsx", ".XLSX"),
        sub_folders=True,
        factory=_path_doc_factory,
    )


def _rel_under_root(root: str, file_path: str) -> str:
    try:
        return str(Path(file_path).resolve().relative_to(Path(root).resolve()))
    except Exception:
        return file_path


def _parent_folder_key(rel_path: str) -> str:
    parent = str(Path(rel_path).parent)
    return parent if parent not in (".", "") else "(root)"


def split_name_type_mark(raw: object) -> tuple[str, str]:
    """Split column H ``NAME;TYPE_MARK`` (first ``;`` only)."""
    text = "" if raw is None else str(raw).strip()
    if not text:
        return "", ""
    if ";" not in text:
        return text, ""
    name, type_mark = text.split(";", 1)
    return name.strip(), type_mark.strip()


def parse_title_system_from_spec_name(spec_name: object) -> tuple[str, str]:
    """Parse DS_TITLE / DS_SYSTEM from specification filename via AGCC patterns."""
    text = collapse_all_whitespace(spec_name)
    if not text:
        return "", ""
    token = ProjectFileName.scan_title_system(text)
    if not token:
        return "", ""
    if "-" not in token:
        return token, ""
    title, system = token.split("-", 1)
    return title.strip(), system.strip()


def parse_sopl_title_system(title_cell: object, mark_cell: object) -> tuple[str, str]:
    """Parse title/system from SO - PL columns BE/BF (Титул / Марка)."""
    title_raw = collapse_all_whitespace(title_cell)
    system = collapse_all_whitespace(mark_cell)
    match = _LEADING_TITLE_RE.match(title_raw)
    title = match.group(1) if match else title_raw
    return title, system


def _set_cell(row: RowStd, key: str, value: object) -> None:
    row.el[key] = CheckElement(value, Color.no)


def _normalize_row_text_fields(row: RowStd, *, sopl: bool) -> None:
    """Normalize whitespace on mapped text fields before title/system parse."""
    collapse_keys = {SPECIFICATION_NAME, DS_TITLE, DS_SYSTEM}
    squeeze_keys = {CODE, NAME, TAGS, UNITS, VENDOR, TYPE_MARK}
    for key in collapse_keys | squeeze_keys:
        el = row.el.get(key)
        if el is None:
            continue
        raw = el.value
        if raw is None or isinstance(raw, list):
            continue
        if key in collapse_keys:
            _set_cell(row, key, collapse_all_whitespace(raw))
        else:
            _set_cell(row, key, normalize_spaces(raw))
    if sopl:
        # SO - PL has no NAME;TYPE_MARK column — keep full NAME, empty TYPE_MARK.
        if TYPE_MARK not in row.el:
            _set_cell(row, TYPE_MARK, "")


def enrich_tsd_row(
    row: RowStd,
    *,
    source_rel: str,
    sheet_name: str,
) -> bool:
    """Post-process NAME/TYPE_MARK, title/system, source path. Returns title_ok."""
    sopl = row.t_com.tabel_type == t_com_init_cls.TsdPackingSoPl.tabel_type
    _normalize_row_text_fields(row, sopl=sopl)

    if sopl:
        title, system = parse_sopl_title_system(
            row.el.get(DS_TITLE).value if row.el.get(DS_TITLE) else "",
            row.el.get(DS_SYSTEM).value if row.el.get(DS_SYSTEM) else "",
        )
        if not (title and system):
            parsed_title, parsed_system = parse_title_system_from_spec_name(
                row.el[SPECIFICATION_NAME].value if SPECIFICATION_NAME in row.el else ""
            )
            title = title or parsed_title
            system = system or parsed_system
    else:
        name, type_mark = split_name_type_mark(row.el[NAME].value)
        _set_cell(row, NAME, name)
        _set_cell(row, TYPE_MARK, type_mark)
        title, system = parse_title_system_from_spec_name(row.el[SPECIFICATION_NAME].value)

    _set_cell(row, DS_TITLE, title)
    _set_cell(row, DS_SYSTEM, system)
    _set_cell(row, ANNOTATION, source_rel)
    _set_cell(row, TITLE, sheet_name)
    return bool(title and system)


def _code_looks_like_header_label(code: object) -> bool:
    """True when CODE cell is a column banner / header, not a product code."""
    text = str(code or "").strip().upper()
    if not text:
        return False
    return (
        text in ("CODE", "КОД", "КОД РД", "BCC", "PO ITEM")
        or text.startswith("КОД РД")
        or "CODE RD" in text
        or text.startswith("СПЕЦИФИКАЦИИ")
        or "NAME OF ZIP" in text
        or "TYPE OF FILE" in text
    )


# Compact product / PO-item codes (BCC…, ECBL…, short alnum with a digit).
_PRODUCT_CODE_RE = re.compile(
    r"^[A-Za-zА-Яа-я0-9][A-Za-zА-Яа-я0-9\-_.]{0,39}$"
)


def _code_looks_like_product_code(code: object) -> bool:
    """True when CODE looks like a packing-list product / PO item code.

    Used to gate ``source_required_fields`` outside the position table band so
    document letterhead / addresses / package legends are not treated as
    incomplete positions. Real mid-table incomplete rows still use the looser
    ``_malformed_source_row_issue`` heuristics.
    """
    text = str(code or "").strip()
    if not text or _code_looks_like_header_label(text):
        return False
    if len(text) > 40 or "\n" in text:
        return False
    # Bilingual form labels: ``Shipper / …``, ``001 = case / Ящик``.
    if "/" in text or "=" in text:
        return False
    if not _PRODUCT_CODE_RE.fullmatch(text):
        return False
    upper = text.upper()
    # BCC… / ECBL… are always product-like (even without digits in the suffix).
    if upper.startswith("BCC") and len(upper) > 3:
        return True
    if upper.startswith("ECBL") and len(upper) > 4:
        return True
    # Other short codes must contain a digit (PO item / numeric id).
    return any(ch.isdigit() for ch in text)


def _row_code_value(row: RowStd) -> object:
    el = row.el.get(CODE)
    return None if el is None else el.value


def _malformed_source_row_issue(
    row: RowStd,
    *,
    source_rel: str,
    sheet_name: str,
    excel_row: int,
) -> PackingIssue | None:
    """Return an issue for a near-position row missing required packing fields.

    Near-position signal (avoids document header / signature / footer noise):
    - ``CODE`` is filled with a non-header value, or both ``NAME`` and ``VALUES``
      are filled (column-shift case without code).
    Callers must also restrict checks to the sheet table band when positions
    exist (see ``load_tsd_file_rows``).
    """
    if row.row_type != RowType.other_row:
        return None
    required = (CODE, NAME, VALUES, UNITS)
    values = {
        key: row.el[key].value if key in row.el else None
        for key in required
    }
    nonempty = {
        key
        for key, value in values.items()
        if value is not None and str(value).strip()
    }
    missing = [key for key in required if key not in nonempty]
    if not missing:
        return None
    has_code = CODE in nonempty and not _code_looks_like_header_label(values[CODE])
    has_name_and_values = NAME in nonempty and VALUES in nonempty
    if not (has_code or has_name_and_values):
        return None
    if len(nonempty) < 2:
        return None
    return PackingIssue(
        code="source_required_fields",
        message=f"Похоже на позицию УЛ, но не заполнены обязательные поля: {', '.join(missing)}",
        source_file=source_rel,
        sheet=sheet_name,
        excel_row=excel_row,
        field=",".join(missing),
        action="Заполните CODE, NAME, VALUES и UNITS либо удалите лишнюю строку",
    )


def _load_tsd_sheet_light(ws: Any, t_com: TableComments) -> list[RowStd]:
    """Map worksheet cells → RowStd without MTO cabinet/section helpers."""
    # Packing templates can contain long visual gaps; scan the whole worksheet
    # so positions below 300 empty rows are never silently discarded.
    raw = load_rows_from_worksheet(
        ws,
        t_com.column_dict,
        max_empty_rows=None,
        skip_empty_rows=True,
        track_source_rows=True,
    )
    rows = list_raw_to_std(raw, t_com)
    return trim_empty_rows_from_end(rows)


def load_tsd_file_rows(
    file_path: str,
    *,
    root: str,
) -> tuple[
    list[RowStd],
    TsdSheetPlan,
    list[RowStd],
    list[TsdSheetHit],
    list[PackingIssue],
]:
    """Load all non-Master sheets from one file; look for position_row on each.

    Opens the workbook once (``read_only``), reads every data sheet from that
    handle, and skips MTO-only post-processors (cabinet / section_type).

    Returns:
        Position rows, sheet plan, all rows, per-sheet hits (including skipped
        Master), and actionable malformed-row issues.

    Raises:
        ValueError: workbook has no non-Master sheets to scan.
    """
    wb = openpyxl.load_workbook(
        file_path, read_only=True, data_only=True, rich_text=True
    )
    try:
        names = list(wb.sheetnames)
        plan = classify_tsd_sheets(names)
        source_rel = _rel_under_root(root, file_path)
        hits: list[TsdSheetHit] = []

        for master in plan.master_sheets:
            hits.append(
                TsdSheetHit(
                    source_rel=source_rel,
                    sheet_name=master,
                    position_rows=0,
                    status="skipped_master",
                )
            )

        if not plan.data_sheets:
            raise ValueError(
                f"no data sheets after skipping Master* (all={plan.all_names!r})"
            )

        all_rows: list[RowStd] = []
        positions: list[RowStd] = []
        row_issues: list[PackingIssue] = []
        for sheet in plan.data_sheets:
            tabel_class = (
                t_com_init_cls.TsdPackingSoPl
                if is_sopl_sheet(sheet)
                else t_com_init_cls.TsdPacking
            )
            t_com = TableComments(
                file_full_path=file_path,
                dir_path="-1",
                tabel_class=tabel_class,
            )
            t_com.sheet_name = sheet
            try:
                if sheet not in wb.sheetnames:
                    raise KeyError(sheet)
                rows = _load_tsd_sheet_light(wb[sheet], t_com)
            except Exception as exc:
                hits.append(
                    TsdSheetHit(
                        source_rel=source_rel,
                        sheet_name=sheet,
                        position_rows=0,
                        status=f"load_error:{exc}",
                    )
                )
                row_issues.append(
                    PackingIssue(
                        code="sheet_load_error",
                        message=f"Не удалось прочитать вкладку ({type(exc).__name__}: {exc})",
                        source_file=source_rel,
                        sheet=sheet,
                        action="Откройте исходный файл и проверьте вкладку; затем повторите чтение УЛ",
                    )
                )
                continue
            n_pos = 0
            # Defer source_required_fields until the position table band is known
            # so document header / signature rows outside it are not flagged.
            other_candidates: list[tuple[RowStd, int]] = []
            position_excel_rows: list[int] = []
            for fallback_row, row in enumerate(rows, start=1):
                excel_row = getattr(row, "_xlsx_source_row", fallback_row)
                row._packing_source_row = excel_row
                enrich_tsd_row(row, source_rel=source_rel, sheet_name=sheet)
                all_rows.append(row)
                if row.row_type == RowType.position_row:
                    positions.append(row)
                    n_pos += 1
                    position_excel_rows.append(excel_row)
                    missing_key_parts = [
                        key
                        for key in (DS_TITLE, DS_SYSTEM, CODE)
                        if not str(row.get_value(key) or "").strip()
                    ]
                    if missing_key_parts:
                        row_issues.append(
                            PackingIssue(
                                code="packing_key_missing",
                                message=(
                                    "У позиции не заполнены части ключа сопоставления: "
                                    + ", ".join(missing_key_parts)
                                ),
                                source_file=source_rel,
                                sheet=sheet,
                                excel_row=excel_row,
                                field=",".join(missing_key_parts),
                                action="Исправьте спецификацию/титул/марку/код в исходном УЛ",
                            )
                        )
                    raw_quantity = row.el[VALUES].value
                    quantity_ok, quantity = try_parse_quantity(raw_quantity)
                    tags_count = sum(
                        1
                        for tag in row.get_tags_list()
                        if str(tag).strip()
                    )
                    if not quantity_ok or quantity < 0:
                        row_issues.append(
                            PackingIssue(
                                code="packing_quantity",
                                message=(
                                    "Количество УЛ должно быть конечным "
                                    "неотрицательным числом"
                                ),
                                source_file=source_rel,
                                sheet=sheet,
                                excel_row=excel_row,
                                field=VALUES,
                                raw_value=raw_quantity,
                                action="Исправьте количество в исходном УЛ",
                            )
                        )
                    elif tags_count > quantity:
                        row_issues.append(
                            PackingIssue(
                                code="packing_tags_exceed_quantity",
                                message=(
                                    "Количество тегов УЛ не может превышать "
                                    "количество позиции"
                                ),
                                source_file=source_rel,
                                sheet=sheet,
                                excel_row=excel_row,
                                field=f"{TAGS},{VALUES}",
                                raw_value={
                                    "tags": tags_count,
                                    "quantity": raw_quantity,
                                },
                                action=(
                                    "Исправьте теги или количество в исходном УЛ"
                                ),
                            )
                        )
                else:
                    other_candidates.append((row, excel_row))
            if position_excel_rows:
                table_lo = min(position_excel_rows)
                table_hi = max(position_excel_rows)
                malformed_candidates: list[tuple[RowStd, int]] = []
                for row, excel_row in other_candidates:
                    if table_lo <= excel_row <= table_hi:
                        # Inside the position band: keep full incomplete-row checks.
                        malformed_candidates.append((row, excel_row))
                        continue
                    # Outside band: only product-like CODE (not letterhead/legend).
                    if _code_looks_like_product_code(_row_code_value(row)):
                        malformed_candidates.append((row, excel_row))
            else:
                # No complete positions: still catch product-like incomplete rows
                # without flooding critical with document banner/footer cells.
                malformed_candidates = [
                    (row, excel_row)
                    for row, excel_row in other_candidates
                    if _code_looks_like_product_code(_row_code_value(row))
                ]
            for row, excel_row in malformed_candidates:
                issue = _malformed_source_row_issue(
                    row,
                    source_rel=source_rel,
                    sheet_name=sheet,
                    excel_row=excel_row,
                )
                if issue is not None:
                    row_issues.append(issue)
            hits.append(
                TsdSheetHit(
                    source_rel=source_rel,
                    sheet_name=sheet,
                    position_rows=n_pos,
                    status="loaded",
                )
            )
    finally:
        wb.close()

    return positions, plan, all_rows, hits, row_issues


def resolve_tsd_load_workers(
    files_total: int,
    workers: int | None = None,
) -> int:
    """Clamp parallel worker count for TSD file reads."""
    if files_total < _MIN_FILES_FOR_PARALLEL:
        return 1
    requested = DEFAULT_TSD_LOAD_WORKERS if workers is None else int(workers)
    if requested <= 1:
        return 1
    return max(1, min(requested, files_total))


def _load_one_tsd_file(file_path: str, root: str) -> _TsdFileOutcome:
    """Worker: load one workbook; never raises (errors go into outcome)."""
    rel = _rel_under_root(root, file_path)
    folder = _parent_folder_key(rel)
    try:
        positions, plan, all_rows, hits, row_issues = load_tsd_file_rows(
            file_path, root=root
        )
    except Exception as exc:
        return _TsdFileOutcome(
            source_rel=rel,
            folder=folder,
            ok=False,
            error=f"{rel}: {exc}",
        )
    return _TsdFileOutcome(
        source_rel=rel,
        folder=folder,
        ok=True,
        positions=positions,
        all_rows=all_rows,
        data_sheets=list(plan.data_sheets),
        hits=hits,
        row_issues=row_issues,
    )


def _apply_tsd_file_outcome(
    stats: TsdLoadStats,
    all_positions: list[RowStd],
    outcome: _TsdFileOutcome,
    *,
    done: int,
    total: int,
) -> None:
    """Merge one file outcome into aggregate stats (main thread only)."""
    stats.folder_files[outcome.folder] += 1
    if not outcome.ok:
        stats.files_failed += 1
        msg = outcome.error or outcome.source_rel
        stats.failures.append(msg)
        stats.row_issues.append(
            PackingIssue(
                code="file_load_error",
                message=msg,
                source_file=outcome.source_rel,
                action="Исправьте/закройте исходный xlsx и повторите чтение УЛ",
            )
        )
        print(f"  [{done}/{total}] FAIL {msg}")
        return

    stats.sheet_hits.extend(outcome.hits)
    stats.row_issues.extend(outcome.row_issues)
    if len(outcome.data_sheets) > 1:
        stats.files_multi_data_sheets += 1
    stats.files_ok += 1
    n_pos = 0
    for row in outcome.all_rows:
        if row.row_type == RowType.position_row:
            stats.position_rows += 1
            n_pos += 1
            if not (row.get_value(DS_TITLE) and row.get_value(DS_SYSTEM)):
                stats.no_title_system += 1
        elif row.row_type == RowType.empty_row:
            stats.empty_rows += 1
        else:
            stats.other_rows += 1
    if n_pos == 0:
        stats.files_without_position += 1
        stats.row_issues.append(
            PackingIssue(
                code="file_without_positions",
                message="После чтения файла не найдено ни одной позиции УЛ",
                source_file=outcome.source_rel,
                action="Проверьте раскладку колонок и обязательные поля",
            )
        )
    stats.folder_positions[outcome.folder] += n_pos
    all_positions.extend(outcome.positions)
    print(
        f"  [{done}/{total}] OK {outcome.source_rel} "
        f"data_sheets={outcome.data_sheets!r} positions={n_pos}"
    )


def _files_fingerprint(root: str, file_paths: list[str]) -> str:
    parts: list[str] = [TSD_PACKING_CACHE_VERSION, os.path.normpath(root)]
    for fp in sorted(file_paths):
        mtime = os.path.getmtime(fp) if os.path.exists(fp) else -1.0
        parts.append(f"{fp}|{mtime}")
    return hashlib.md5("|".join(parts).encode(errors="replace")).hexdigest()


def save_tsd_packing_pickle(
    rows: list[RowStd],
    *,
    root: str,
    file_paths: list[str],
    stats: TsdLoadStats,
    critical: TsdCriticalReport,
    report_paths: dict[str, str],
    fingerprint: str | None = None,
) -> Path:
    """Write pickle payload with meta fingerprint under cache dir."""
    cache_dir = tsd_packing_cache_dir()
    path = cache_dir / TSD_PACKING_ROWS_CACHE_NAME
    # Same position_row set as summary xlsx, but slim columns only (cache v6+).
    slim_rows = slim_packing_rows(rows)
    payload = {
        "meta": {
            "version": TSD_PACKING_CACHE_VERSION,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "root": root,
            "fingerprint": fingerprint or _files_fingerprint(root, file_paths),
            "files_total": stats.files_total,
            "position_rows": stats.position_rows,
            "counters": {
                "files_total": stats.files_total,
                "files_ok": stats.files_ok,
                "files_multi_data_sheets": stats.files_multi_data_sheets,
                "files_failed": stats.files_failed,
                "position_rows": stats.position_rows,
                "empty_rows": stats.empty_rows,
                "other_rows": stats.other_rows,
                "no_title_system": stats.no_title_system,
                "files_without_position": stats.files_without_position,
                "malformed_rows": len(stats.row_issues),
            },
            "critical": {
                "ok": critical.ok,
                "remarks": list(critical.remarks),
            },
            "report_paths": dict(report_paths),
            "loader_issues": [
                {
                    **asdict(issue),
                    "severity": issue.severity.value,
                }
                for issue in stats.row_issues
            ],
        },
        "data": slim_rows,
    }
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            prefix=f".{path.name}.",
            suffix=".tmp",
            dir=path.parent,
            delete=False,
        ) as stream:
            temp_path = Path(stream.name)
            pickle.dump(payload, stream, protocol=pickle.HIGHEST_PROTOCOL)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp_path, path)
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()
    return path


def load_tsd_packing_pickle(
    root: str | None = None,
    *,
    file_paths: list[str] | None = None,
) -> list[RowStd] | None:
    """Compatibility wrapper over the structured shared packing provider."""
    path = tsd_packing_cache_dir() / TSD_PACKING_ROWS_CACHE_NAME
    expected_fingerprint = (
        _files_fingerprint(root, file_paths)
        if root is not None and file_paths is not None
        else None
    )
    dataset = load_packing_dataset(
        path,
        expected_version=TSD_PACKING_CACHE_VERSION,
        expected_root=root,
        expected_fingerprint=expected_fingerprint,
    )
    return dataset.rows if dataset.available else None


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return ";".join(str(x) for x in value)
    return str(value)


def _summary_row_values(row: RowStd) -> list[object]:
    """Build one summary data row (same order as ``_SUMMARY_COLUMNS``)."""
    source_excel_row = getattr(row, "_packing_source_row", None)
    if not isinstance(source_excel_row, int) or source_excel_row <= 0:
        source_excel_row = ""
    return [
        _cell_str(row.el[ANNOTATION].value),  # source_rel
        _cell_str(row.el[TITLE].value),  # sheet
        source_excel_row,
        _cell_str(row.el[CODE].value),
        _cell_str(row.el[SPECIFICATION_NAME].value),
        _cell_str(row.el[DS_TITLE].value),
        _cell_str(row.el[DS_SYSTEM].value),
        _cell_str(row.el[TAGS].value),
        _cell_str(row.el[NAME].value),
        _cell_str(row.el[TYPE_MARK].value),
        _cell_str(row.el[VALUES].value),
        _cell_str(row.el[UNITS].value),
        _cell_str(row.el[VENDOR].value),
    ]


def _ul_cell_filled(value: object) -> bool:
    """True when a summary cell is non-blank, including numeric zero."""
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def _ul_folder_sort_key(name: str) -> tuple:
    parts = re.split(r"(\d+)", name)
    key: list[tuple[int, int | str]] = []
    for part in parts:
        if not part:
            continue
        if part.isdigit():
            key.append((0, int(part)))
        else:
            key.append((1, part.casefold()))
    return tuple(key)


def ul_folder_stat_rows(
    samples: list[tuple[object, object, object, object]],
) -> list[tuple[str, int, int, int, int]]:
    """Count summary rows per source folder that contain code, name, and qty.

    Args:
        samples: ``(folder, code, name, qty)`` for each data row.

    Returns:
        One tuple per folder, then a total row labelled ``ИТОГО``.
    """
    buckets: dict[str, list[int]] = {}
    for folder, code, name, qty in samples:
        label = str(folder or "").strip() or "(корень)"
        counts = buckets.setdefault(label, [0, 0, 0, 0])
        counts[0] += 1
        if _ul_cell_filled(code):
            counts[1] += 1
        if _ul_cell_filled(name):
            counts[2] += 1
        if _ul_cell_filled(qty):
            counts[3] += 1
    rows = [
        (label, *buckets[label])
        for label in sorted(buckets, key=_ul_folder_sort_key)
    ]
    totals = [0, 0, 0, 0]
    for _label, *counts in rows:
        for index, value in enumerate(counts):
            totals[index] += value
    rows.append(("ИТОГО", *totals))
    return rows


def _folder_stat_samples(rows: list[RowStd]) -> list[tuple[str, str, str, str]]:
    samples: list[tuple[str, str, str, str]] = []
    for row in rows:
        if row.row_type != RowType.position_row:
            continue
        source_rel = _cell_str(row.el[ANNOTATION].value)
        samples.append(
            (
                _parent_folder_key(source_rel),
                _cell_str(row.el[CODE].value),
                _cell_str(row.el[NAME].value),
                _cell_str(row.el[VALUES].value),
            )
        )
    return samples


def _write_ul_folder_stats_sheet(
    wb: xlsxwriter.Workbook,
    samples: list[tuple[object, object, object, object]],
) -> None:
    """Second sheet: row counts of filled code, name, and quantity by folder."""
    ws = wb.add_worksheet(UL_FOLDER_STATS_SHEET)
    header_fmt = wb.add_format(
        {
            "bold": True,
            "align": "center",
            "valign": "vcenter",
            "text_wrap": True,
            "border": 1,
        }
    )
    total_fmt = wb.add_format({"bold": True, "border": 1})
    cell_fmt = wb.add_format({"border": 1})
    ws.set_row(0, 32)
    widths = (42, 14, 20, 24, 18)
    for col, (label, width) in enumerate(zip(_UL_FOLDER_STATS_HEADERS, widths)):
        ws.write(0, col, label, header_fmt)
        ws.set_column(col, col, width)
    written = ul_folder_stat_rows(samples)
    for row_idx, item in enumerate(written, start=1):
        fmt = total_fmt if item[0] == "ИТОГО" else cell_fmt
        for col, value in enumerate(item):
            ws.write(row_idx, col, value, fmt)
    note_row = len(written) + 2
    ws.merge_range(note_row, 0, note_row, 4, _UL_FOLDER_STATS_NOTE)
    if written:
        ws.autofilter(0, 0, len(written), 4)
    ws.freeze_panes(1, 1)


def add_ul_folder_stats_sheet(path: Path, dest: Path | None = None) -> int:
    """Add or replace «Статистика по папкам» on an existing UL summary.

    Args:
        path: ``tsd_packing_summary_*.xlsx`` to read.
        dest: Where to write. Defaults to ``path``. Use another path when
            the source workbook is open in Excel.

    Returns:
        Number of folder rows, excluding the total.

    Raises:
        ValueError: The summary sheet or a required header is missing.
    """
    wb = openpyxl.load_workbook(path)
    try:
        if "ТСД свод" not in wb.sheetnames:
            raise ValueError("в книге нет листа «ТСД свод»")
        ws = wb["ТСД свод"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        needed = ("Файл", "Код РД", "Наименование", "Кол-во")
        missing = [title for title in needed if title not in headers]
        if missing:
            raise ValueError(f"нет колонок: {', '.join(missing)}")
        index = {title: headers.index(title) for title in needed}

        def _at(row: tuple[object, ...], title: str) -> object:
            pos = index[title]
            return row[pos] if pos < len(row) else None

        samples: list[tuple[object, object, object, object]] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(cell is not None and cell != "" for cell in row):
                continue
            source = _at(row, "Файл")
            samples.append(
                (
                    _parent_folder_key(str(source or "")),
                    _at(row, "Код РД"),
                    _at(row, "Наименование"),
                    _at(row, "Кол-во"),
                )
            )
        if UL_FOLDER_STATS_SHEET in wb.sheetnames:
            del wb[UL_FOLDER_STATS_SHEET]
        stats = wb.create_sheet(UL_FOLDER_STATS_SHEET, 1)
        stats.append(list(_UL_FOLDER_STATS_HEADERS))
        for cell in stats[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                wrap_text=True, vertical="center", horizontal="center"
            )
        written = ul_folder_stat_rows(samples)
        for item in written:
            stats.append(list(item))
        for cell in stats[stats.max_row]:
            cell.font = Font(bold=True)
        note_row = stats.max_row + 2
        stats.cell(row=note_row, column=1, value=_UL_FOLDER_STATS_NOTE)
        stats.merge_cells(
            start_row=note_row, start_column=1, end_row=note_row, end_column=5
        )
        for col_letter, width in {
            "A": 42,
            "B": 14,
            "C": 20,
            "D": 24,
            "E": 18,
        }.items():
            stats.column_dimensions[col_letter].width = width
        stats.row_dimensions[1].height = 32
        stats.freeze_panes = "B2"
        stats.auto_filter.ref = f"A1:E{max(len(written) + 1, 1)}"
        wb.active = wb["ТСД свод"]
        target = dest or path
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
    finally:
        wb.close()
    return max(len(ul_folder_stat_rows(samples)) - 1, 0)


def save_tsd_packing_summary_xlsx(rows: list[RowStd], out_path: Path | None = None) -> Path:
    """Write flat summary workbook for manual review (xlsxwriter).

    Default name is always timestamped via ``tsd_packing_summary_xlsx_path``
    (single naming source for DS load and RFP/compare paths).

    Header row: bold fill, freeze panes, autofilter. Column widths follow
    step4-style defaults (see ``_SUMMARY_COLUMNS``).
    """
    out = out_path or tsd_packing_summary_xlsx_path()
    out.parent.mkdir(parents=True, exist_ok=True)

    data_rows = [r for r in rows if r.row_type == RowType.position_row]
    num_cols = len(_SUMMARY_COLUMNS)
    last_data_row = len(data_rows)  # 0-based header at row 0 → last data index

    wb = xlsxwriter.Workbook(str(out), {"constant_memory": False, "remove_timezone": True})
    ws = wb.add_worksheet("ТСД свод")

    header_fmt = wb.add_format(
        {
            "bold": True,
            "font_color": "#000000",
            "bg_color": _SUMMARY_HEADER_FILL,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "text_wrap": True,
        }
    )
    cell_fmt = wb.add_format(
        {
            "align": "left",
            "valign": "vcenter",
            "border": 1,
        }
    )

    ws.set_row(0, 36)
    for ci, (_key, label, width) in enumerate(_SUMMARY_COLUMNS):
        ws.write(0, ci, label, header_fmt)
        ws.set_column(ci, ci, width)

    for row_idx, row in enumerate(data_rows):
        excel_row = row_idx + 1
        for ci, value in enumerate(_summary_row_values(row)):
            ws.write(excel_row, ci, value, cell_fmt)

    if last_data_row > 0:
        ws.autofilter(0, 0, last_data_row, num_cols - 1)
    ws.freeze_panes(1, 0)
    _write_ul_folder_stats_sheet(wb, _folder_stat_samples(rows))

    try:
        wb.close()
    except xlsxwriter.exceptions.FileCreateError:
        if out_path is not None:
            raise
        # Rare: same-second collision or transient lock — retry with µs stamp.
        retry = out.with_name(
            f"{TSD_PACKING_SUMMARY_STEM}_"
            f"{datetime.now().strftime('%Y.%m.%d_%H.%M.%S.%f')}"
            f"{out.suffix}"
        )
        print(f"Сводный файл УЛ не записан, повтор: {retry.name}")
        return save_tsd_packing_summary_xlsx(rows, retry)
    return out


def _format_stats_report(stats: TsdLoadStats, root: str) -> str:
    lines: list[str] = [
        f"root: {root}",
        f"files_total: {stats.files_total}",
        f"files_ok: {stats.files_ok}",
        f"files_multi_data_sheets: {stats.files_multi_data_sheets}",
        f"files_failed: {stats.files_failed}",
        f"position_rows: {stats.position_rows}",
        f"empty_rows: {stats.empty_rows}",
        f"other_rows: {stats.other_rows}",
        f"no_title_system: {stats.no_title_system}",
        f"files_without_position: {stats.files_without_position}",
        f"malformed_rows: {len(stats.row_issues)}",
        "",
        "top folders by files:",
    ]
    for folder, cnt in stats.folder_files.most_common(30):
        lines.append(f"  {cnt:5d}  {folder}")
    lines.append("")
    lines.append("top folders by position rows:")
    for folder, cnt in stats.folder_positions.most_common(30):
        lines.append(f"  {cnt:5d}  {folder}")
    lines.append("")
    lines.append("FAILURES:")
    if not stats.failures:
        lines.append("  (none)")
    else:
        for item in stats.failures:
            lines.append(f"  {item}")
    lines.append("")
    lines.append("MALFORMED SOURCE ROWS:")
    if not stats.row_issues:
        lines.append("  (none)")
    else:
        for issue in stats.row_issues:
            lines.append(f"  {issue.format_line()}")
    return "\n".join(lines) + "\n"


def _aggregate_sheet_name_stats(
    hits: list[TsdSheetHit],
) -> list[tuple[str, int, int, int, int]]:
    """Aggregate by sheet name: files, positions, skipped_master, load_errors."""
    files: Counter[str] = Counter()
    positions: Counter[str] = Counter()
    skipped: Counter[str] = Counter()
    errors: Counter[str] = Counter()
    for hit in hits:
        files[hit.sheet_name] += 1
        positions[hit.sheet_name] += hit.position_rows
        if hit.status == "skipped_master":
            skipped[hit.sheet_name] += 1
        elif hit.status.startswith("load_error"):
            errors[hit.sheet_name] += 1
    rows: list[tuple[str, int, int, int, int]] = []
    for name in sorted(files.keys(), key=lambda s: (-positions[s], s.lower())):
        rows.append(
            (
                name,
                files[name],
                positions[name],
                skipped[name],
                errors[name],
            )
        )
    return rows


def save_tsd_packing_sheets_report(
    stats: TsdLoadStats,
    *,
    root: str,
    out_path: Path | None = None,
) -> Path:
    """Write sheet-name inventory: each sheet as a row + position_row counts.

    Also appends a per-file detail section (source_rel, sheet, positions, status).
    """
    out = out_path or (tsd_packing_cache_dir() / TSD_PACKING_SHEETS_REPORT_NAME)
    agg = _aggregate_sheet_name_stats(stats.sheet_hits)

    table = PrettyTable()
    table.field_names = [
        "sheet_name",
        "files",
        "position_rows",
        "skipped_master",
        "load_errors",
    ]
    table.align = "l"
    table.align["files"] = "r"
    table.align["position_rows"] = "r"
    table.align["skipped_master"] = "r"
    table.align["load_errors"] = "r"
    for name, n_files, n_pos, n_skip, n_err in agg:
        table.add_row([name, n_files, n_pos, n_skip, n_err])

    detail = PrettyTable()
    detail.field_names = ["source_rel", "sheet_name", "position_rows", "status"]
    detail.align = "l"
    detail.align["position_rows"] = "r"
    for hit in sorted(
        stats.sheet_hits,
        key=lambda h: (h.source_rel.lower(), h.sheet_name.lower()),
    ):
        detail.add_row(
            [hit.source_rel, hit.sheet_name, hit.position_rows, hit.status]
        )

    text = (
        f"root: {root}\n"
        f"sheet_name rows: {len(agg)}\n"
        f"file×sheet rows: {len(stats.sheet_hits)}\n\n"
        "=== BY SHEET NAME (строка = имя вкладки) ===\n"
        f"{table}\n\n"
        "=== DETAIL (каждый файл × вкладка) ===\n"
        f"{detail}\n"
    )
    out.write_text(text, encoding="utf-8")
    print(f"Отчёт по вкладкам: {out}")
    print(table)
    return out


@dataclass
class TsdLoadResult:
    """Outcome of ``load_and_cache_tsd_packing``."""

    summary_path: str
    stats: TsdLoadStats
    success: bool
    critical: TsdCriticalReport | None = None
    from_cache: bool = False


def load_and_cache_tsd_packing(
    root: str,
    *,
    force: bool = True,
    workers: int | None = None,
) -> TsdLoadResult:
    """Walk TSD root, load non-Master sheets, write pickle + summary + reports.

    ``Master*`` sheets are skipped. Every other sheet is scanned for
    ``position_row`` (including ``Single*``, ``SO - PL``, etc.).

    File reads run in a ``ThreadPoolExecutor`` (I/O-bound; default 8 workers).
    Each file is opened once; sheets are mapped without MTO cabinet/section helpers.

    Args:
        root: Source folder (recursive).
        force: If False and pickle fingerprint matches, reuse cache.
        workers: Parallel readers; ``None`` → ``DEFAULT_TSD_LOAD_WORKERS``;
            ``1`` forces sequential.

    Returns:
        ``TsdLoadResult`` with summary path, ``from_cache`` flag, and
        ``success=False`` if any file failed.

    Raises:
        FileNotFoundError: Root missing.
        RuntimeError: No successfully loaded files at all.
    """
    global _last_tsd_critical_report

    root = str(root or "").strip()
    if not root or not os.path.isdir(root):
        raise FileNotFoundError(f"Папка ТСД не найдена: {root}")

    docs = collect_tsd_files(root)
    file_paths = [d.file_full_path for d in docs]
    initial_fingerprint = _files_fingerprint(root, file_paths)
    stats = TsdLoadStats(files_total=len(docs))

    if not force:
        expected_fingerprint = initial_fingerprint
        dataset = load_packing_dataset(
            tsd_packing_cache_dir() / TSD_PACKING_ROWS_CACHE_NAME,
            expected_version=TSD_PACKING_CACHE_VERSION,
            expected_root=root,
            expected_fingerprint=expected_fingerprint,
        )
        if dataset.available:
            print(f"ТСД: загружено из кэша — {len(dataset.rows)} position rows")
            summary_path = ""
            if dataset.meta is not None:
                summary_path = str(dataset.meta.report_paths.get("summary", "")).strip()
            if not summary_path or not os.path.isfile(summary_path):
                print(
                    "ТСД: кэш актуален, но записанный свод отсутствует — "
                    "перечитываем исходные файлы."
                )
            else:
                counters = dataset.meta.counters if dataset.meta is not None else {}
                for field_name in (
                    "files_total",
                    "files_ok",
                    "files_multi_data_sheets",
                    "files_failed",
                    "position_rows",
                    "empty_rows",
                    "other_rows",
                    "no_title_system",
                    "files_without_position",
                ):
                    if field_name in counters:
                        setattr(stats, field_name, counters[field_name])
                critical = TsdCriticalReport(
                    ok=dataset.quality == PackingQualityLevel.OK,
                    remarks=[issue.format_line() for issue in dataset.issues],
                    files_failed=stats.files_failed,
                    files_without_position=stats.files_without_position,
                )
                _last_tsd_critical_report = critical
                return TsdLoadResult(
                    summary_path,
                    stats,
                    success=dataset.quality == PackingQualityLevel.OK,
                    critical=critical,
                    from_cache=True,
                )

    all_positions: list[RowStd] = []
    n_workers = resolve_tsd_load_workers(len(docs), workers)
    print(f"ТСД: найдено xlsx файлов: {len(docs)}")
    print(
        "Правило листов: Master* пропускаем; "
        "все остальные вкладки сканируем на position_row."
    )
    print(f"ТСД: параллельное чтение файлов, workers={n_workers}")

    total = len(docs)
    if n_workers <= 1:
        for i, doc in enumerate(docs, start=1):
            outcome = _load_one_tsd_file(doc.file_full_path, root)
            _apply_tsd_file_outcome(
                stats, all_positions, outcome, done=i, total=total
            )
    else:
        with ThreadPoolExecutor(max_workers=n_workers) as pool:
            futures = [
                pool.submit(_load_one_tsd_file, doc.file_full_path, root)
                for doc in docs
            ]
            done = 0
            for fut in as_completed(futures):
                done += 1
                outcome = fut.result()
                _apply_tsd_file_outcome(
                    stats, all_positions, outcome, done=done, total=total
                )

    final_fingerprint = _files_fingerprint(root, file_paths)
    if final_fingerprint != initial_fingerprint:
        stats.row_issues.append(
            PackingIssue(
                code="source_changed_during_load",
                message=(
                    "Состав или время изменения исходных УЛ изменились во время чтения; "
                    "свод может быть частичным"
                ),
                action="Дождитесь окончания правок файлов и повторите чтение УЛ",
            )
        )

    report_text = _format_stats_report(stats, root)
    report_path = tsd_packing_cache_dir() / TSD_PACKING_LOAD_REPORT_NAME
    report_path.write_text(report_text, encoding="utf-8")
    sheets_report_path = save_tsd_packing_sheets_report(stats, root=root)

    critical = build_tsd_critical_report(stats, root=root)
    critical_path = save_tsd_packing_critical_report(critical, root=root)
    _last_tsd_critical_report = critical

    table = PrettyTable()
    table.field_names = ["metric", "value"]
    table.align = "l"
    for key in (
        "files_total",
        "files_ok",
        "files_multi_data_sheets",
        "files_failed",
        "position_rows",
        "empty_rows",
        "other_rows",
        "no_title_system",
        "files_without_position",
    ):
        table.add_row([key, getattr(stats, key)])
    print(table)
    print(f"Отчёт: {report_path}")
    print(f"Отчёт вкладок: {sheets_report_path}")
    print(f"Критичные замечания: {critical_path}")
    print()
    print("=== КРИТИЧНЫЕ ЗАМЕЧАНИЯ ===")
    print(critical.format_detail())

    if not all_positions and stats.files_ok == 0:
        raise RuntimeError(
            f"ТСД: нет успешно загруженных файлов "
            f"(ошибок={stats.files_failed}, см. {report_path})."
        )

    summary_path = save_tsd_packing_summary_xlsx(all_positions)
    pickle_path = save_tsd_packing_pickle(
        all_positions,
        root=root,
        file_paths=file_paths,
        stats=stats,
        critical=critical,
        fingerprint=final_fingerprint,
        report_paths={
            "summary": str(summary_path),
            "load": str(report_path),
            "sheets": str(sheets_report_path),
            "critical": str(critical_path),
        },
    )
    print(f"Pickle: {pickle_path}")
    print(f"Свод: {summary_path}")

    ok = critical.ok
    if not ok:
        print(
            f"ТСД WARNING: ошибок файлов={stats.files_failed}, "
            f"критичных замечаний={len(critical.remarks)}; "
            f"свод частичный — {summary_path}; детали — {report_path}"
        )
    return TsdLoadResult(
        str(summary_path),
        stats,
        success=ok,
        critical=critical,
        from_cache=False,
    )


def ensure_tsd_packing_cache_current(
    root: str | os.PathLike[str] | None = None,
) -> TsdLoadResult:
    """Rebuild the packing cache when the source-tree fingerprint changed.

    Same check as grouped DS preflight: ``load_and_cache_tsd_packing(..., force=False)``.

    Args:
        root: TSD source folder. Defaults to ``DEFAULT_TSD_PACKING_ROOT``.

    Returns:
        Load result. ``from_cache=True`` means the pickle matched the folder.

    Raises:
        FileNotFoundError: Source folder is missing.
        RuntimeError: No packing files could be loaded.
    """
    packing_root = str(root or DEFAULT_TSD_PACKING_ROOT).strip()
    return load_and_cache_tsd_packing(packing_root, force=False)


def format_packing_cache_freshness_detail(result: TsdLoadResult) -> str:
    """Short Russian detail for the RFP ``ul_preflight`` milestone.

    Args:
        result: Outcome of ``ensure_tsd_packing_cache_current``.

    Returns:
        Compact ``unchanged|rebuilt; …`` text for the progress table.
    """
    action = "unchanged" if result.from_cache else "rebuilt"
    if result.from_cache:
        reason = "кэш совпадает с исходной папкой"
    else:
        reason = "исходные УЛ изменились или кэш отсутствовал"
    rows = int(getattr(result.stats, "position_rows", 0) or 0)
    detail = f"{action}; {reason}; position_rows={rows}"
    if not result.success:
        detail += "; есть критичные замечания"
    return detail


@dataclass
class TsdOneFileResult:
    """Outcome of checking one packing-list workbook."""

    success: bool
    message: str
    result_path: str | None = None
    critical: TsdCriticalReport | None = None


def inspect_one_tsd_file(file_path: str, out_dir: str | Path) -> TsdOneFileResult:
    """Run the full TSD reader on one workbook and write reports beside it.

    Uses ``_load_one_tsd_file``, the same sheet rules and critical report as
    ``load_and_cache_tsd_packing``. Does not write the production pickle,
    summary or ``tsd_packing_critical.txt`` under the UL cache folder, and
    does not replace the in-memory critical report of a full run.

    Args:
        file_path: One ``.xlsx`` packing list.
        out_dir: Folder for this check's reports. Created when missing.

    Returns:
        Success follows the critical report (a readable file with positions
        and no critical remarks). ``result_path`` is ``out_dir``.
    """

    path = Path(file_path)
    out = Path(out_dir)
    if not path.is_file():
        return TsdOneFileResult(False, f"файл не найден: {path}", None)
    if path.name.startswith("~$"):
        return TsdOneFileResult(False, "это временный файл Excel (~$)", None)
    if path.suffix.lower() != ".xlsx":
        return TsdOneFileResult(
            False,
            f"упаковочный лист читается как xlsx: {path.name}",
            None,
        )

    out.mkdir(parents=True, exist_ok=True)
    root = str(path.parent)
    print(f"УЛ один файл: {path}")
    print(
        "Правило листов: Master* пропускаем; "
        "все остальные вкладки сканируем на position_row."
    )
    outcome = _load_one_tsd_file(str(path), root)
    stats = TsdLoadStats(files_total=1)
    positions: list[RowStd] = []
    _apply_tsd_file_outcome(stats, positions, outcome, done=1, total=1)

    report_path = out / TSD_PACKING_LOAD_REPORT_NAME
    report_path.write_text(_format_stats_report(stats, root), encoding="utf-8")
    sheets_path = save_tsd_packing_sheets_report(
        stats,
        root=root,
        out_path=out / TSD_PACKING_SHEETS_REPORT_NAME,
    )
    critical = build_tsd_critical_report(stats, root=root)
    critical_path = save_tsd_packing_critical_report(
        critical,
        root=root,
        out_path=out / TSD_PACKING_CRITICAL_REPORT_NAME,
    )
    summary_path = save_tsd_packing_summary_xlsx(
        positions,
        out_path=out / tsd_packing_summary_xlsx_name(),
    )
    print(f"Отчёт: {report_path}")
    print(f"Отчёт вкладок: {sheets_path}")
    print(f"Критичные замечания: {critical_path}")
    print(f"Свод одного файла: {summary_path}")
    print("=== КРИТИЧНЫЕ ЗАМЕЧАНИЯ ===")
    print(critical.format_detail())
    message = (
        f"Один файл {path.name}: позиций {stats.position_rows}, "
        f"замечаний {len(critical.remarks)}. Отчёт: {out}"
    )
    return TsdOneFileResult(
        success=bool(critical.ok and outcome.ok),
        message=message,
        result_path=str(out),
        critical=critical,
    )
