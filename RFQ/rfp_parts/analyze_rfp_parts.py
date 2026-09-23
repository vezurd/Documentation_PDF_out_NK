"""Analyze split RFP part workbooks and sum them into a summary.

The script detects the relevant RFP sheet/header in each workbook, extracts the
columns currently used by the aggregated RFP loader, and writes a text report
with file-level stats, warnings, and summed totals.

Comparison with the legacy aggregated ``Сводная RFP.xlsx`` is optional
(``--summary-file``); the default run is parts-only.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import dataclass, field, replace
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from prettytable import PrettyTable

# Package lives at RFQ/rfp_parts/; repo root is two levels up.
ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from RFQ.rfp_parts.ds_checklist import (
    DEFAULT_CHECKLIST_FILE,
    DEFAULT_PARTS_DIR,
    ChecklistCompareResult,
    format_checklist_report_section,
    parse_ds_name_from_file_name,
    run_checklist_compare,
)
from RFQ.rfp_parts.file_status import (
    DUPLICATE_TAGS_XLSX_NAME,
    EMPTY_CODE_XLSX_NAME,
    build_file_status_payload,
    format_field_list,
    humanize_parts_message,
    write_file_status_json,
)
from base.base_google import load_base
from base.tables_columns import (
    CODE,
    ColNames,
    DS_CODE_1C,
    DS_NAME,
    DS_NUMBER,
    DS_SPECIFICATION,
    DS_TITLE,
    EQUIPMENT_CODE,
    NAME,
    NAME_2,
    RFP_SUPPLY_STATUS,
    TAGS,
    TYPE_MARK,
    UNITS,
    UNITS_CHECK_STATUS,
    UNITS_CONVERSION_TRACE,
    VALUES,
    VALUES_2,
)
from RFQ.ds_compare.ds_units_normalize import join_unique_units_text, normalize_units_text
from RFQ.tags_rfp_compare.rfp_supply_status import (
    CANONICAL_EXCLUDED_FROM_SUPPLY,
    KIND_EXCLUDED,
    KIND_TAG_REPLACED,
    KIND_UNKNOWN,
    RfpSupplyStatusIssue,
    apply_tag_replacement,
    is_canonical_excluded_from_supply,
    parse_rfp_supply_status,
    raise_if_rfp_supply_status_issues,
)
from RFQ.tags_rfp_compare.rfp_tags_utils import load_config
from RFQ.units_convert import (
    ALGORITHM_VERSION,
    ConversionInvariant,
    ConversionPlan,
    ConversionRequest,
    GoogleUnitsIndex,
    UnitsConversionError,
    build_conversion_plan,
    build_google_units_index,
    converted_quantity,
    converted_unit,
    write_fractional_conversion_logs,
)
from tags.tag_classes import TagClass
from tags.tag_parser import get_tag
from utils.colors import Color


DEFAULT_BASE = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP"
)
DEFAULT_SUMMARY_FILE = DEFAULT_BASE / "Сводная RFP.xlsx"  # legacy; not used by default
PARTS_KIND = "parts"
_PARTS_KINDS = frozenset({"parts", "increase", "decrease"})
# Per-run report folders (GUI + CLI default).
DEFAULT_REPORTS_BASE_DIR = DEFAULT_BASE / "RFP сводный файл"
# Stable local manifest only (not next to the package / not on UNC).
TMP_DIR = ROOT / "tmp"
DEFAULT_MANIFEST = TMP_DIR / "rfp_parts_manifest.json"
# Legacy prefix kept for docs only; new runs use RUN_DIR_STAMP_RE.
REPORT_DIR_PREFIX = "rfp_сбор_отчеты_"
NET_XLSX_NAME = "rfp_parts_net.xlsx"
NET_NO_TAGS_XLSX_NAME = "rfp_parts_net_no_tags.xlsx"
COLLISIONS_XLSX_NAME = "rfp_parts_collisions.xlsx"
BUILD_DEPS_JSON_NAME = "rfp_parts_build_deps.json"
RUN_DIR_STAMP_RE = re.compile(r"^\d{4}\.\d{2}\.\d{2}_\d{2}\.\d{2}(?:_\d+)?$")
PARTS_TIMING_PREFIX = "[parts timing]"
PARTS_TIMING_ARTIFACT_NAME = "rfp_parts_timing.txt"
PARTS_TIMING_SLOW_THRESHOLD_SEC = 2.0


def parts_timing_format_start(label: str) -> str:
    """Format a phase START line (stable prefix for tests and Job monitor)."""
    return f"{PARTS_TIMING_PREFIX} START: {label}"


def parts_timing_format_done(label: str, elapsed_sec: float, total_sec: float) -> str:
    """Format a phase DONE line with phase and cumulative elapsed seconds."""
    return (
        f"{PARTS_TIMING_PREFIX} DONE: {label} — "
        f"{elapsed_sec:.2f} с; всего {total_sec:.2f} с"
    )


def parts_timing_format_error(
    label: str,
    elapsed_sec: float,
    total_sec: float,
    exc_type_name: str,
) -> str:
    """Format a phase ERROR line; caller re-raises the original exception."""
    return (
        f"{PARTS_TIMING_PREFIX} ERROR: {label} — "
        f"{elapsed_sec:.2f} с; всего {total_sec:.2f} с; {exc_type_name}"
    )


def parts_timing_format_total(total_sec: float) -> str:
    """Format the final TOTAL line for a successful run."""
    return f"{PARTS_TIMING_PREFIX} TOTAL: {total_sec:.2f} с"


def parts_timing_format_progress(
    current: int,
    total: int,
    cumulative_elapsed_sec: float,
) -> str:
    """Format extraction progress (every 10 files and at the last file)."""
    return (
        f"{PARTS_TIMING_PREFIX} PROGRESS: файлы {current}/{total}; "
        f"извлечение {cumulative_elapsed_sec:.2f} с"
    )


def parts_timing_format_slow_file(
    index: int,
    total: int,
    file_name: str,
    elapsed_sec: float,
) -> str:
    """Format a slow per-workbook extraction line (>= threshold seconds)."""
    return (
        f"{PARTS_TIMING_PREFIX} SLOW: файл {index}/{total} "
        f"{file_name} — {elapsed_sec:.2f} с"
    )


class _PartsTimingSession:
    """Live stdout timing for ``run_rfp_parts_analyze``; buffers lines for artifact."""

    def __init__(self) -> None:
        self._baseline = time.perf_counter()
        self._lines: list[str] = []

    def _emit(self, line: str) -> None:
        print(line, flush=True)
        self._lines.append(line)

    def elapsed_total(self) -> float:
        return time.perf_counter() - self._baseline

    @contextmanager
    def phase(self, label: str):
        """Time one labeled phase; ERROR + re-raise on failure."""
        self._emit(parts_timing_format_start(label))
        phase_start = time.perf_counter()
        try:
            yield
        except Exception as exc:
            elapsed = time.perf_counter() - phase_start
            self._emit(
                parts_timing_format_error(
                    label,
                    elapsed,
                    self.elapsed_total(),
                    type(exc).__name__,
                )
            )
            raise
        else:
            elapsed = time.perf_counter() - phase_start
            self._emit(
                parts_timing_format_done(label, elapsed, self.elapsed_total())
            )

    def emit_progress(self, current: int, total: int, cumulative_elapsed_sec: float) -> None:
        self._emit(
            parts_timing_format_progress(current, total, cumulative_elapsed_sec)
        )

    def emit_slow_file(
        self,
        index: int,
        total: int,
        file_name: str,
        elapsed_sec: float,
    ) -> None:
        self._emit(
            parts_timing_format_slow_file(index, total, file_name, elapsed_sec)
        )

    def emit_total(self) -> None:
        self._emit(parts_timing_format_total(self.elapsed_total()))

    def write_artifact(self, reports_dir: Path) -> None:
        """Best-effort UTF-8 timing log; WARN on failure, never raises."""
        if not self._lines:
            return
        try:
            path = Path(reports_dir) / PARTS_TIMING_ARTIFACT_NAME
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("\n".join(self._lines) + "\n", encoding="utf-8")
        except OSError as exc:
            print(
                f"WARN: cannot write {PARTS_TIMING_ARTIFACT_NAME}: "
                f"{type(exc).__name__}: {exc}",
                flush=True,
            )


def make_reports_out_dir(*, base: Path | None = None, when: datetime | None = None) -> Path:
    """Build ``…/RFP сводный файл/YYYY.MM.DD_HH.MM`` (unique within the minute).

    Args:
        base: Parent directory for the stamped folder. Defaults to
            ``DEFAULT_REPORTS_BASE_DIR`` (UNC).
        when: Timestamp for the folder name; defaults to now.

    Returns:
        Path to the per-run reports directory (not created yet). If the stamp
        already exists, appends ``_2``, ``_3``, … so an open Excel file from a
        previous run cannot block writing.
    """
    parent = base or DEFAULT_REPORTS_BASE_DIR
    stamp = (when or datetime.now()).strftime("%Y.%m.%d_%H.%M")
    candidate = parent / stamp
    if not candidate.exists():
        return candidate
    suffix = 2
    while True:
        candidate = parent / f"{stamp}_{suffix}"
        if not candidate.exists():
            return candidate
        suffix += 1


def find_latest_rfp_parts_run_dir(*, base: Path | None = None) -> Path | None:
    """Newest stamp folder ``YYYY.MM.DD_HH.MM`` under the reports base.

    Args:
        base: Parent to scan; defaults to ``DEFAULT_REPORTS_BASE_DIR``.

    Returns:
        Path to the latest matching directory, or ``None``.
    """
    parent = base or DEFAULT_REPORTS_BASE_DIR
    try:
        if not parent.is_dir():
            return None
        dirs = [
            path
            for path in parent.iterdir()
            if path.is_dir() and RUN_DIR_STAMP_RE.fullmatch(path.name)
        ]
    except OSError:
        return None
    if not dirs:
        return None
    dirs.sort(key=lambda path: path.name, reverse=True)
    return dirs[0]


def _resolve_latest_stamp_file(name: str, *, base: Path | None = None) -> Path | None:
    """Return ``<latest stamp>/<name>`` when that file exists."""
    run_dir = find_latest_rfp_parts_run_dir(base=base)
    if run_dir is None:
        return None
    path = run_dir / name
    return path if path.is_file() else None


def resolve_latest_rfp_parts_net_xlsx(*, base: Path | None = None) -> Path | None:
    """``rfp_parts_net.xlsx`` inside the latest stamp folder, if present."""
    return _resolve_latest_stamp_file(NET_XLSX_NAME, base=base)


def resolve_latest_rfp_parts_net_no_tags_xlsx(
    *, base: Path | None = None
) -> Path | None:
    """``rfp_parts_net_no_tags.xlsx`` inside the latest stamp folder, if present."""
    return _resolve_latest_stamp_file(NET_NO_TAGS_XLSX_NAME, base=base)


def resolve_latest_rfp_parts_collisions_xlsx(*, base: Path | None = None) -> Path | None:
    """``rfp_parts_collisions.xlsx`` inside the latest stamp folder, if present."""
    run_dir = find_latest_rfp_parts_run_dir(base=base)
    if run_dir is None:
        return None
    path = run_dir / COLLISIONS_XLSX_NAME
    return path if path.is_file() else None


def resolve_latest_rfp_parts_duplicate_tags_xlsx(
    *, base: Path | None = None
) -> Path | None:
    """``Отчет по тегам - сбор частей.xlsx`` inside the latest stamp folder, if present."""
    run_dir = find_latest_rfp_parts_run_dir(base=base)
    if run_dir is None:
        return None
    path = run_dir / DUPLICATE_TAGS_XLSX_NAME
    return path if path.is_file() else None


EXPECTED_SHEET_NAME = "Перечень материалов"
# Cover / legend / contents tabs are never scanned for positions.
NEVER_READ_SHEET_PATTERNS: tuple[re.Pattern[str], ...] = tuple(
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"титульн",
        r"^титул$",
        r"^title(\s*page|\s*sheet)?$",
        r"cover",
        r"лист\s*регистрац",
        r"^содержание$",
        r"^contents?$",
        r"общие\s*(данные|сведения)",
        r"^легенда",
        r"инструкц",
        r"instruction",
        r"^примечан",
        r"revision\s*history",
        r"record\s*of\s*revision",
    )
)

PRIMARY_FIELDS = (
    "DS_NUMBER",
    "DS_TITLE",
    "DS_SPECIFICATION",
    "TAGS",
    "DS_CODE_1C",
    "CODE",
    "NAME",
    "TYPE_MARK",
    "VALUES",
    "UNITS",
)
OPTIONAL_FIELDS = ("NAME_2", "VENDOR", "VALUES_2", "RFP_SUPPLY_STATUS")
REQUIRED_FIELDS = ("DS_TITLE", "CODE", "NAME", "VALUES", "UNITS")
NO_TAG_KEY = "<NO_TAG>"
# Sum key for the net workbook. Stored in rfp_parts_build_deps.json so Launch
# rebuilds when this contract changes (old nets glued rows across DS names).
AGGREGATION_KEY = "ds_name+title+code+tag"
UnitKey = tuple[str, str, str, str]  # ds_name, title_system, code, tag
# Parts files have two qty blocks: RD (Excel 9-10) and Lot (Excel 17-18).
# Parts-mode always reads the Lot block. Summary still uses the RD block (+ DS_NAME).
EXPECTED_PART_HEADER = {
    "DS_NUMBER": 0,
    "DS_TITLE": 1,
    "DS_SPECIFICATION": 2,
    "TAGS": 3,
    "DS_CODE_1C": 4,
    "CODE": 5,
    "NAME": 6,
    "TYPE_MARK": 7,
    "VALUES": 16,  # Excel №17 — «Закупка по Лоту» / Кол-во
    "UNITS": 17,  # Excel №18 — «Закупка по Лоту» / Ед. изм
}
EXPECTED_SUMMARY_HEADER = {
    "DS_NUMBER": 1,
    "DS_TITLE": 2,
    "DS_SPECIFICATION": 3,
    "TAGS": 4,
    "DS_CODE_1C": 5,
    "CODE": 6,
    "NAME": 7,
    "TYPE_MARK": 8,
    "VALUES": 9,
    "UNITS": 10,
}
# When header row has two «Кол-во»/«Ед. изм», keep the rightmost (Lot) for parts.
_HEADER_RIGHTMOST_FIELDS = frozenset({"VALUES", "UNITS"})

HEADER_PATTERNS: dict[str, tuple[str, ...]] = {
    "DS_NUMBER": (r"№\s*п/?п", r"порядк"),
    "DS_TITLE": (r"титул",),
    "DS_SPECIFICATION": (r"спецификац",),
    "TAGS": (r"tag", r"тег", r"линия"),
    "DS_CODE_1C": (r"код\s*1\s*с", r"код\s*1c"),
    "CODE": (r"код\s*рд", r"\bbcc\b"),
    "NAME": (r"наименование\s*мтр",),
    "TYPE_MARK": (r"техническ.*характерист", r"тип.*марк"),
    "VALUES": (r"кол-?во", r"количество"),
    "UNITS": (r"ед\.?,?\s*изм", r"единиц.*измер"),
    "NAME_2": (r"наименование.*(мто|ркд|продукц)",),
    "VENDOR": (r"поставщик", r"изготовитель", r"производитель", r"страна\s*производства"),
    "VALUES_2": (r"кол-?во.*(мто|ркд|vo)",),
    "RFP_SUPPLY_STATUS": (r"статус\s*позиц",),
}

DIAGNOSTICS_LEGEND = (
    "Папка: источник RFP — единая папка частей RFP_Зиновьев.",
    "Файл: имя исходной книги RFP, где найдена строка с предупреждением.",
    "Строка: номер строки в выбранном рабочем листе (обычно 'Перечень материалов').",
    "Теги: исходная колонка TAGS / 'Линия, Tag-номер' (Excel-колонка D, индекс 3 в шаблоне частей).",
    "Код оборудования: исходная колонка CODE / 'Код РД / BCC' (Excel-колонка F, индекс 5 в шаблоне частей).",
    "Количество/ед.: блок «Закупка по Лоту» — Excel №17–18 (индексы 16–17 в шаблоне частей). "
    "Пустые или нечисловые значения в строке с данными — ERROR (не пропускаются тихо).",
    "Сообщение: краткое описание проверки; ERROR — шаблон/пустой файл/пустой Лот/"
    "нечисловое количество/в поле Tag есть подстрока RFQ/некорректный статус позиции. "
    "Дубли тегов и TAG≠лот — "
    f"в {DUPLICATE_TAGS_XLSX_NAME} (не в таблице GUI).",
)


@dataclass(frozen=True)
class HeaderCandidate:
    sheet: str
    row_number: int
    score: int
    columns: dict[str, int]
    labels: list[tuple[int, str]]


@dataclass(frozen=True)
class RfpRecord:
    kind: str
    file_name: str
    sheet: str
    excel_row: int
    ds_name: str
    ds_number: str
    ds_title: str
    ds_specification: str
    tags: str
    ds_code_1c: str
    code: str
    name: str
    type_mark: str
    values: Decimal
    units: str
    units_check_status: str = ""
    units_conversion_trace: str = ""
    rfp_supply_status: str = ""


@dataclass(frozen=True)
class SkippedEmptyCodeRow:
    """One parts row skipped because CODE (and maybe NAME/TITLE) is empty."""

    file_name: str
    sheet: str
    excel_row: int
    ds_name: str
    ds_title: str
    tags: str
    name: str
    type_mark: str
    values_text: str
    units: str
    missing_fields: str
    level: str


@dataclass
class FileStats:
    kind: str
    file_name: str
    file_id: str = ""
    ds_name: str = ""
    workbook_sheets: tuple[str, ...] = ()
    sheet: str = ""
    header_row: int = 0
    rows: int = 0
    qty_sum: Decimal = Decimal("0")
    warnings: int = 0
    header_signature: str = ""
    primary_signature: str = ""
    values_with_inner_spaces: int = 0


@dataclass
class PartsAggregation:
    """Tag-aware unit/decimal totals from the parts folder (summed into the summary)."""

    units: Counter[UnitKey]
    decimals: dict[UnitKey, Decimal]
    examples: dict[UnitKey, RfpRecord]
    units_status_by_key: dict[UnitKey, str] = field(default_factory=dict)
    units_trace_by_key: dict[UnitKey, str] = field(default_factory=dict)
    supply_status_by_key: dict[UnitKey, str] = field(default_factory=dict)


@dataclass
class CoarseCollisionResult:
    """Coarse title_system+code checks (units mismatch across summed parts)."""

    qty_by_key: dict[tuple[str, str], Decimal]
    examples: dict[tuple[str, str], RfpRecord]
    ambiguous_meta: list[tuple[tuple[str, str], RfpRecord, tuple[str, ...]]]
    records_by_key: dict[tuple[str, str], list[RfpRecord]] = field(
        default_factory=dict
    )


@dataclass(frozen=True)
class UnitsFixRow:
    """One source RFP line to align UNITS to the Google code base."""

    title_system: str
    code: str
    google_units: str
    file_units: str
    units_variants: tuple[str, ...]
    record: RfpRecord


def _norm_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").replace("\r", " ").strip()


def _tags_text_contains_rfq(value: str) -> bool:
    """True when the TAG cell contains the substring ``RFQ`` (any case)."""
    return "RFQ" in value.upper()


def _norm_key_text(value: Any) -> str:
    return "".join(_norm_text(value).split()).upper()


def _normalize_code(value: Any) -> str:
    return _norm_key_text(value)


def _normalize_numeric_text(value: Any) -> str:
    """Strip edges and all inner whitespace (spaces, NBSP, newlines) from qty text.

    Examples: ``'1 048,000'``, ``'1\\n891,000'``, ``'1\\xa0891,000'`` -> digits-only groups.
    """
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    raw = str(value).replace("\xa0", " ").strip()
    if not raw:
        return ""
    # Remove ALL whitespace between digit groups, including line breaks.
    return "".join(ch for ch in raw if not ch.isspace())


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    raw = _normalize_numeric_text(value)
    if not raw:
        return None
    raw = re.sub(r"[^0-9,.\-+]", "", raw)
    if not raw:
        return None
    raw = raw.replace(",", ".")
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def _has_inner_numeric_whitespace(value: Any) -> bool:
    """Return True for numeric-looking strings such as '1 048,000' or '1\\n891,000'."""
    if not isinstance(value, str):
        return False
    return re.search(r"\d\s+\d", value) is not None


def _parse_ds_name_from_file_name(file_name: str) -> str:
    """Extract the DS name prefix from a split RFP workbook name."""
    return parse_ds_name_from_file_name(file_name)


def _ds_name_from_row(kind: str, file_name: str, row: list[Any], columns: dict[str, int]) -> str:
    if kind in _PARTS_KINDS:
        return _parse_ds_name_from_file_name(file_name)
    ds_number_idx = columns.get("DS_NUMBER")
    if ds_number_idx is not None and ds_number_idx > 0 and ds_number_idx - 1 < len(row):
        return _norm_text(row[ds_number_idx - 1])
    return ""


def _expected_header_for_kind(kind: str) -> dict[str, int] | None:
    if kind in _PARTS_KINDS:
        return EXPECTED_PART_HEADER
    if kind == "summary":
        return EXPECTED_SUMMARY_HEADER
    return None


def _signature_from_columns(columns: dict[str, int], fields: Iterable[str] = PRIMARY_FIELDS) -> str:
    return ", ".join(f"{field}:{columns.get(field)}" for field in fields)


def _signature_from_expected(expected: dict[str, int]) -> str:
    return ", ".join(f"{field}:{expected.get(field)}" for field in PRIMARY_FIELDS)


def _is_total_row(row_values: list[Any]) -> bool:
    first_values = [_norm_text(v).upper() for v in row_values[:3] if _norm_text(v)]
    return any("ИТОГО" in v or "ВСЕГО" in v for v in first_values)


def _match_header(
    row_values: Iterable[Any],
    *,
    prefer_lot_qty: bool = False,
) -> dict[str, int]:
    """Map header labels to column indexes.

    Identity columns (title/code/name/…) take the leftmost match (RD block).
    For parts-mode (``prefer_lot_qty=True``) VALUES/UNITS take the rightmost
    match so they bind to «Закупка по Лоту». Summary keeps the leftmost (RD) pair.
    """
    found: dict[str, int] = {}
    rightmost = _HEADER_RIGHTMOST_FIELDS if prefer_lot_qty else frozenset()
    for idx, raw_value in enumerate(row_values):
        text = _norm_text(raw_value).lower()
        if not text:
            continue
        for field, patterns in HEADER_PATTERNS.items():
            if field in found and field not in rightmost:
                continue
            if any(re.search(pattern, text, re.IGNORECASE) for pattern in patterns):
                found[field] = idx
    return found


def _safe_iter_rows(ws: Any, warnings: list[tuple[str, str, str]], file_name: str, *, min_row: int = 1, max_row: int | None = None):
    """Yield worksheet rows and keep rows already read if openpyxl fails at the end."""
    try:
        yield from ws.iter_rows(min_row=min_row, max_row=max_row, values_only=True)
    except ValueError as exc:
        warnings.append(
            (
                "ERROR",
                file_name,
                f"{ws.title}: openpyxl остановился при разборе XML листа: {exc}",
            )
        )


def _build_file_id(kind: str, file_name: str) -> str:
    """Build a stable source ID from folder kind and RFP document number."""
    stem = Path(file_name).stem
    doc_match = re.search(
        r"AGCC\.287-[\w.\-]+?RFP-\d+(?:_[\wА-Яа-я]+)*",
        stem,
        re.IGNORECASE,
    )
    if doc_match:
        doc_id = doc_match.group(0)
    else:
        doc_id = stem
    doc_id = re.sub(r"\s+", "", doc_id).upper()
    return f"{kind}:{doc_id}"


def is_never_read_sheet(sheet_name: str) -> bool:
    """Return True for cover/legend/contents tabs that must not be scanned."""
    text = _norm_text(sheet_name)
    if not text:
        return True
    folded = text.casefold()
    return any(pattern.search(folded) for pattern in NEVER_READ_SHEET_PATTERNS)


def is_expected_sheet_name(sheet_name: str) -> bool:
    """Return True when the tab name matches ``Перечень материалов`` (any case)."""
    return _norm_text(sheet_name).casefold() == EXPECTED_SHEET_NAME.casefold()


def split_part_sheets(sheetnames: Iterable[str]) -> tuple[tuple[str, ...], tuple[str, ...]]:
    """Split workbook tabs into never-read vs candidate sheets.

    Args:
        sheetnames: Workbook sheet titles in order.

    Returns:
        ``(skipped, candidates)`` tuples.
    """
    skipped: list[str] = []
    candidates: list[str] = []
    for name in sheetnames:
        if is_never_read_sheet(name):
            skipped.append(name)
        else:
            candidates.append(name)
    return tuple(skipped), tuple(candidates)


def _file_has_error(
    warnings: list[tuple[str, str, str]],
    file_name: str,
) -> bool:
    return any(level == "ERROR" and fname == file_name for level, fname, _ in warnings)


def _note_empty_file(
    warnings: list[tuple[str, str, str]],
    stats: FileStats,
    *,
    detail: str,
) -> None:
    """ERROR when a workbook is present but no positions were extracted."""
    if _file_has_error(warnings, stats.file_name):
        return
    warnings.append(
        (
            "ERROR",
            stats.file_name,
            f"файл есть, но позиции не прочитаны: {detail}",
        )
    )
    stats.warnings += 1


def _header_has_required_fields(header: HeaderCandidate | None) -> bool:
    if header is None:
        return False
    return all(field in header.columns for field in REQUIRED_FIELDS)


def _find_header(
    wb: Any,
    warnings: list[tuple[str, str, str]],
    file_name: str,
    *,
    prefer_lot_qty: bool = False,
    sheet_name: str | None = None,
) -> HeaderCandidate | None:
    """Find the best RFP header on one sheet (first 100 rows)."""
    target = sheet_name or EXPECTED_SHEET_NAME
    if target not in wb.sheetnames:
        return None
    return _find_header_on_sheet(
        wb[target],
        warnings,
        file_name,
        prefer_lot_qty=prefer_lot_qty,
    )


def _find_header_on_sheet(
    ws: Any,
    warnings: list[tuple[str, str, str]],
    file_name: str,
    *,
    prefer_lot_qty: bool = False,
) -> HeaderCandidate | None:
    candidates: list[HeaderCandidate] = []
    for row_number, row in enumerate(
        _safe_iter_rows(ws, warnings, file_name, min_row=1, max_row=100),
        start=1,
    ):
        columns = _match_header(row, prefer_lot_qty=prefer_lot_qty)
        score = len([field for field in PRIMARY_FIELDS if field in columns])
        if score < 5:
            continue
        labels = [
            (idx, _norm_text(value)[:90])
            for idx, value in enumerate(row)
            if _norm_text(value)
        ]
        candidates.append(
            HeaderCandidate(
                sheet=ws.title,
                row_number=row_number,
                score=score,
                columns=columns,
                labels=labels,
            )
        )
    candidates.sort(key=lambda item: item.score, reverse=True)
    return candidates[0] if candidates else None


def _choose_part_sheet(
    candidates: tuple[str, ...],
    probes: list[tuple[str, HeaderCandidate | None]],
) -> tuple[str, HeaderCandidate | None]:
    """Pick one working sheet: the only candidate, else expected name, else best header."""
    by_name = {name: header for name, header in probes}
    if len(candidates) == 1:
        name = candidates[0]
        return name, by_name.get(name)
    usable = [
        (name, header)
        for name, header in probes
        if _header_has_required_fields(header)
    ]
    for name, header in usable:
        if is_expected_sheet_name(name):
            return name, header
    if usable:
        name, header = max(usable, key=lambda item: item[1].score if item[1] else 0)
        return name, header
    scored = [(name, header) for name, header in probes if header is not None]
    if scored:
        name, header = max(scored, key=lambda item: item[1].score if item[1] else 0)
        return name, header
    for name in candidates:
        if is_expected_sheet_name(name):
            return name, by_name.get(name)
    name = candidates[0]
    return name, by_name.get(name)


def _cell(row: list[Any], columns: dict[str, int], field: str) -> Any:
    idx = columns.get(field)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _bind_header_to_stats(stats: FileStats, header: HeaderCandidate) -> None:
    stats.sheet = header.sheet
    stats.header_row = header.row_number
    stats.header_signature = ", ".join(
        f"{field}:{idx}" for field, idx in sorted(header.columns.items())
    )
    stats.primary_signature = _signature_from_columns(header.columns)


def _warn_shifted_header(
    kind: str,
    path: Path,
    header: HeaderCandidate,
    stats: FileStats,
    warnings: list[tuple[str, str, str]],
) -> None:
    expected_header = _expected_header_for_kind(kind)
    if expected_header is None:
        return
    shifted = {
        field: {"expected": expected_idx, "actual": header.columns.get(field)}
        for field, expected_idx in expected_header.items()
        if header.columns.get(field) != expected_idx
    }
    if not shifted:
        return
    warnings.append(
        (
            "ERROR",
            path.name,
            "порядок основных столбцов RFP не совпадает с шаблоном: "
            f"{shifted}; ожидалось {_signature_from_expected(expected_header)}; "
            f"факт {stats.primary_signature}",
        )
    )


def _emit_supply_status_issue(
    *,
    warnings: list[tuple[str, str, str]],
    stats: FileStats,
    status_issues: list[RfpSupplyStatusIssue] | None,
    file_name: str,
    excel_row: int,
    code: str,
    raw: str,
    detail: str = "",
) -> None:
    issue = RfpSupplyStatusIssue(
        file_name=file_name,
        excel_row=excel_row,
        code=code,
        raw=raw,
        detail=detail,
    )
    if status_issues is not None:
        status_issues.append(issue)
    warnings.append(("ERROR", file_name, issue.as_warning_message()))
    stats.warnings += 1


def _append_records_from_sheet(
    *,
    ws: Any,
    header: HeaderCandidate,
    path: Path,
    kind: str,
    warnings: list[tuple[str, str, str]],
    stats: FileStats,
    status_issues: list[RfpSupplyStatusIssue] | None = None,
    skipped_empty_code: list[SkippedEmptyCodeRow] | None = None,
) -> list[RfpRecord]:
    records: list[RfpRecord] = []
    for offset, row_tuple in enumerate(
        _safe_iter_rows(ws, warnings, path.name, min_row=header.row_number + 1),
        start=header.row_number + 1,
    ):
        row = list(row_tuple)
        if _is_total_row(row):
            continue
        watched_values = [_cell(row, header.columns, field) for field in REQUIRED_FIELDS]
        if not any(_norm_text(value) for value in watched_values):
            continue

        missing_in_row = [
            field
            for field in REQUIRED_FIELDS
            if not _norm_text(_cell(row, header.columns, field))
        ]
        if missing_in_row:
            # Empty Lot qty/units on an otherwise filled row is a shift/template problem —
            # do not skip quietly. Fully blank rows are filtered above.
            level = (
                "ERROR"
                if any(field in missing_in_row for field in ("VALUES", "UNITS"))
                else "WARN"
            )
            warnings.append(
                (
                    level,
                    path.name,
                    f"строка {offset}: пропущена, нет полей: "
                    f"{format_field_list(missing_in_row)}",
                )
            )
            if level == "ERROR":
                stats.warnings += 1
            if skipped_empty_code is not None and "CODE" in missing_in_row:
                skipped_empty_code.append(
                    SkippedEmptyCodeRow(
                        file_name=path.name,
                        sheet=header.sheet,
                        excel_row=offset,
                        ds_name=_ds_name_from_row(
                            kind, path.name, row, header.columns
                        ),
                        ds_title=_norm_text(_cell(row, header.columns, "DS_TITLE")),
                        tags=_norm_text(_cell(row, header.columns, "TAGS")),
                        name=_norm_text(_cell(row, header.columns, "NAME")),
                        type_mark=_norm_text(_cell(row, header.columns, "TYPE_MARK")),
                        values_text=_norm_text(_cell(row, header.columns, "VALUES")),
                        units=_norm_text(_cell(row, header.columns, "UNITS")),
                        missing_fields=format_field_list(missing_in_row),
                        level=level,
                    )
                )
            continue

        raw_qty = _cell(row, header.columns, "VALUES")
        if _has_inner_numeric_whitespace(raw_qty):
            stats.values_with_inner_spaces += 1

        qty = _parse_decimal(raw_qty)
        if qty is None:
            warnings.append(
                (
                    "ERROR",
                    path.name,
                    f"строка {offset}: не разобрать количество лота "
                    f"(Excel №17)={raw_qty!r}",
                )
            )
            stats.warnings += 1
            continue

        units = _norm_text(_cell(row, header.columns, "UNITS"))
        if not units:
            warnings.append(
                (
                    "ERROR",
                    path.name,
                    f"строка {offset}: пустая единица измерения лота (Excel №18)",
                )
            )
            stats.warnings += 1
            continue

        tags = _norm_text(_cell(row, header.columns, "TAGS"))
        if _tags_text_contains_rfq(tags):
            warnings.append(
                (
                    "ERROR",
                    path.name,
                    f"строка {offset}: в поле Tag недопустима подстрока RFQ ({tags!r})",
                )
            )
            stats.warnings += 1
            continue

        code = _norm_text(_cell(row, header.columns, "CODE"))
        raw_status = _cell(row, header.columns, "RFP_SUPPLY_STATUS")
        parsed_status = parse_rfp_supply_status(raw_status)
        if parsed_status.kind == KIND_UNKNOWN:
            _emit_supply_status_issue(
                warnings=warnings,
                stats=stats,
                status_issues=status_issues,
                file_name=path.name,
                excel_row=offset,
                code=code,
                raw=parsed_status.text,
            )
            continue
        if parsed_status.kind == KIND_TAG_REPLACED:
            new_tags, replace_error = apply_tag_replacement(tags, parsed_status)
            if replace_error:
                _emit_supply_status_issue(
                    warnings=warnings,
                    stats=stats,
                    status_issues=status_issues,
                    file_name=path.name,
                    excel_row=offset,
                    code=code,
                    raw=parsed_status.text,
                    detail=replace_error,
                )
                continue
            tags = new_tags
            supply_status = parsed_status.text
        elif parsed_status.kind == KIND_EXCLUDED:
            supply_status = CANONICAL_EXCLUDED_FROM_SUPPLY
        else:
            supply_status = ""

        record = RfpRecord(
            kind=kind,
            file_name=path.name,
            sheet=header.sheet,
            excel_row=offset,
            ds_name=_ds_name_from_row(kind, path.name, row, header.columns),
            ds_number=_norm_text(_cell(row, header.columns, "DS_NUMBER")),
            ds_title=_norm_text(_cell(row, header.columns, "DS_TITLE")),
            ds_specification=_norm_text(_cell(row, header.columns, "DS_SPECIFICATION")),
            tags=tags,
            ds_code_1c=_norm_text(_cell(row, header.columns, "DS_CODE_1C")),
            code=code,
            name=_norm_text(_cell(row, header.columns, "NAME")),
            type_mark=_norm_text(_cell(row, header.columns, "TYPE_MARK")),
            values=qty,
            units=units,
            rfp_supply_status=supply_status,
        )
        records.append(record)
        stats.rows += 1
        stats.qty_sum += qty
    return records


def _extract_summary_records(
    wb: Any,
    path: Path,
    kind: str,
    warnings: list[tuple[str, str, str]],
    stats: FileStats,
    status_issues: list[RfpSupplyStatusIssue] | None = None,
    skipped_empty_code: list[SkippedEmptyCodeRow] | None = None,
) -> list[RfpRecord]:
    """Legacy сводная RFP: keep the strict ``Перечень материалов`` tab."""
    if EXPECTED_SHEET_NAME not in wb.sheetnames:
        warnings.append(
            (
                "ERROR",
                path.name,
                f"в книге нет листа «{EXPECTED_SHEET_NAME}»",
            )
        )
        stats.warnings += 1
        return []
    header = _find_header(
        wb, warnings, path.name, prefer_lot_qty=False, sheet_name=EXPECTED_SHEET_NAME
    )
    if header is None:
        warnings.append(
            (
                "ERROR",
                path.name,
                f"шапка RFP не найдена в первых 100 строках листа «{EXPECTED_SHEET_NAME}»",
            )
        )
        stats.warnings += 1
        return []
    _bind_header_to_stats(stats, header)
    missing_fields = [field for field in REQUIRED_FIELDS if field not in header.columns]
    if missing_fields:
        warnings.append(
            (
                "ERROR",
                path.name,
                f"не найдены обязательные столбцы: {format_field_list(missing_fields)}",
            )
        )
        stats.warnings += 1
        return []
    _warn_shifted_header(kind, path, header, stats, warnings)
    return _append_records_from_sheet(
        ws=wb[header.sheet],
        header=header,
        path=path,
        kind=kind,
        warnings=warnings,
        stats=stats,
        status_issues=status_issues,
        skipped_empty_code=skipped_empty_code,
    )


def _extract_part_records(
    wb: Any,
    path: Path,
    kind: str,
    warnings: list[tuple[str, str, str]],
    stats: FileStats,
    status_issues: list[RfpSupplyStatusIssue] | None = None,
    skipped_empty_code: list[SkippedEmptyCodeRow] | None = None,
) -> list[RfpRecord]:
    """Soft sheet pick: one tab any name; several tabs → skip cover, read one, WARN rest."""
    skipped, candidates = split_part_sheets(wb.sheetnames)
    if not candidates:
        skipped_txt = ", ".join(skipped) if skipped else "(нет вкладок)"
        warnings.append(
            (
                "ERROR",
                path.name,
                "нет рабочих листов для чтения (все вкладки в списке пропуска: "
                f"{skipped_txt})",
            )
        )
        stats.warnings += 1
        return []

    prefer_lot = kind in _PARTS_KINDS
    probes: list[tuple[str, HeaderCandidate | None]] = [
        (
            name,
            _find_header_on_sheet(
                wb[name], warnings, path.name, prefer_lot_qty=prefer_lot
            ),
        )
        for name in candidates
    ]
    chosen_name, header = _choose_part_sheet(candidates, probes)
    stats.sheet = chosen_name
    if header is None:
        warnings.append(
            (
                "ERROR",
                path.name,
                "шапка RFP не найдена в первых 100 строках листа "
                f"«{chosen_name}» (рабочие листы: {', '.join(candidates)})",
            )
        )
        stats.warnings += 1
        return []

    _bind_header_to_stats(stats, header)
    unread = [name for name in candidates if name != header.sheet]
    if len(candidates) > 1 and unread:
        warnings.append(
            (
                "WARN",
                path.name,
                f"в книге {len(candidates)} рабочих листов, позиции прочитаны только с "
                f"{header.sheet!r}; не читались: {', '.join(unread)}",
            )
        )
    missing_fields = [field for field in REQUIRED_FIELDS if field not in header.columns]
    if missing_fields:
        warnings.append(
            (
                "ERROR",
                path.name,
                f"не найдены обязательные столбцы: {format_field_list(missing_fields)}",
            )
        )
        stats.warnings += 1
        return []

    _warn_shifted_header(kind, path, header, stats, warnings)
    records = _append_records_from_sheet(
        ws=wb[header.sheet],
        header=header,
        path=path,
        kind=kind,
        warnings=warnings,
        stats=stats,
        status_issues=status_issues,
        skipped_empty_code=skipped_empty_code,
    )
    if not records:
        _note_empty_file(
            warnings,
            stats,
            detail=(
                f"лист {header.sheet!r}, шапка строка {header.row_number}, "
                "позиций с заполненными обязательными полями нет "
                "(сдвиг столбцов или пустая таблица)"
            ),
        )
    return records


def _extract_records(
    path: Path,
    kind: str,
    warnings: list[tuple[str, str, str]],
    status_issues: list[RfpSupplyStatusIssue] | None = None,
    skipped_empty_code: list[SkippedEmptyCodeRow] | None = None,
) -> tuple[list[RfpRecord], FileStats]:
    stats = FileStats(kind=kind, file_name=path.name)
    stats.file_id = _build_file_id(kind, path.name)
    records: list[RfpRecord] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True, rich_text=False)
    except Exception as exc:
        warnings.append(
            (
                "ERROR",
                path.name,
                f"не удалось открыть книгу: {type(exc).__name__}: {exc}",
            )
        )
        stats.warnings += 1
        return records, stats

    try:
        stats.workbook_sheets = tuple(wb.sheetnames)
        stats.ds_name = _parse_ds_name_from_file_name(path.name)
        if kind == "summary":
            records = _extract_summary_records(
                wb,
                path,
                kind,
                warnings,
                stats,
                status_issues=status_issues,
                skipped_empty_code=skipped_empty_code,
            )
        else:
            records = _extract_part_records(
                wb,
                path,
                kind,
                warnings,
                stats,
                status_issues=status_issues,
                skipped_empty_code=skipped_empty_code,
            )
        if not records:
            _note_empty_file(
                warnings,
                stats,
                detail="лист не выбран, шапка не найдена или колонки не совпали с шаблоном",
            )
    finally:
        wb.close()

    return records, stats


def _xlsx_files(directory: Path) -> list[Path]:
    if not directory.exists():
        raise FileNotFoundError(f"Directory does not exist: {directory}")
    if not directory.is_dir():
        raise NotADirectoryError(f"Path is not a directory: {directory}")
    return sorted(
        [
            path
            for path in directory.iterdir()
            if path.suffix.lower() in {".xlsx", ".xlsm"} and not path.name.startswith("~$")
        ],
        key=lambda path: path.name.lower(),
    )


def resolve_parts_workbooks(
    parts_dir: Path,
    only_files: Iterable[Path] | None = None,
) -> list[Path]:
    """Return part workbooks for a collection run.

    Args:
        parts_dir: Folder ``RFP_Зиновьев``. Used when ``only_files`` is omitted.
        only_files: When set, check only these workbooks. The same extractor
            runs on each file. An empty sequence is an error.

    Returns:
        Workbook paths, locks skipped.

    Raises:
        FileNotFoundError: Directory or a named file is missing, or the
            limited list has no workbook.
        NotADirectoryError: ``parts_dir`` is not a directory and ``only_files``
            is omitted.
        ValueError: A limited path is not an Excel workbook.
    """

    if only_files is None:
        return _xlsx_files(parts_dir)
    files: list[Path] = []
    for raw in only_files:
        path = Path(raw)
        if path.name.startswith("~$"):
            continue
        if not path.is_file():
            raise FileNotFoundError(f"RFP file not found: {path}")
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError(f"Not an Excel workbook: {path}")
        files.append(path)
    if not files:
        raise FileNotFoundError("No RFP workbooks in only_files")
    return files


def _record_key(record: RfpRecord) -> tuple[str, str]:
    return (_norm_key_text(record.ds_title), _norm_key_text(record.code))


def _unit_key(record: RfpRecord, tag: str) -> UnitKey:
    """Return the net sum key: ds_name + title + code + tag (or ``<NO_TAG>``)."""
    tag_key = _norm_key_text(tag) if tag else NO_TAG_KEY
    return (
        _norm_key_text(record.ds_name),
        _norm_key_text(record.ds_title),
        _norm_key_text(record.code),
        tag_key,
    )


def _normalize_tags_text(value: str) -> str:
    """Normalize common Excel spacing artifacts inside tag strings."""
    value = value.replace("\xa0", " ")
    value = re.sub(r"\s*-\s*", "-", value)
    return value.strip()


def _raw_tag_tokens(value: str) -> list[str]:
    return [token.strip() for token in re.split(r"[;,\s]+", value) if token.strip()]


def _record_tag_keys(record: RfpRecord, warnings: list[tuple[str, str, str]]) -> list[str]:
    raw_tags = _norm_text(record.tags)
    normalized_tags = _normalize_tags_text(raw_tags) if raw_tags else ""
    if not normalized_tags:
        return []

    parsed_tags = get_tag(normalized_tags)
    if parsed_tags:
        return parsed_tags

    fallback_tags = _raw_tag_tokens(normalized_tags)
    return fallback_tags


def _is_int_decimal(value: Decimal) -> bool:
    return value == value.to_integral_value()


def _sum_by_key(records: list[RfpRecord]) -> dict[tuple[str, str], Decimal]:
    result: dict[tuple[str, str], Decimal] = defaultdict(lambda: Decimal("0"))
    for record in records:
        result[_record_key(record)] += record.values
    return result


def _resolve_units_matrix_path(units_matrix_path: Path | None) -> Path:
    """Resolve conversion matrix path from explicit arg or main RFP profile."""
    if units_matrix_path is not None:
        return Path(units_matrix_path)
    config = load_config()
    paths = config.get("paths") if isinstance(config, dict) else None
    raw = paths.get("units_convert_matrix") if isinstance(paths, dict) else None
    if not raw:
        raise UnitsConversionError(
            "В конфигурации RFP отсутствует paths.units_convert_matrix"
        )
    return Path(str(raw))


def _load_google_rows_strict() -> list[Any]:
    """Load Google base rows; conversion treats failure as fatal."""
    try:
        rows = load_base()
    except Exception as exc:
        raise UnitsConversionError(
            f"Не удалось загрузить базу кодов Google: {type(exc).__name__}: {exc}"
        ) from exc
    if rows is None:
        raise UnitsConversionError("Загрузка базы кодов Google (load_base) вернула None")
    return rows


def _build_parts_conversion_requests(records: list[RfpRecord]) -> list[ConversionRequest]:
    """Build immutable conversion requests for parts records."""
    requests: list[ConversionRequest] = []
    silent_warnings: list[tuple[str, str, str]] = []
    for index, record in enumerate(records):
        tags = _record_tag_keys(record, silent_warnings)
        requests.append(
            ConversionRequest(
                request_id=f"parts-{index}",
                contour="parts",
                code=record.code,
                source_unit=record.units,
                quantity=record.values,
                tags_count=len(tags),
                invariant=(
                    ConversionInvariant.TAGS_EQUAL
                    if tags
                    else ConversionInvariant.NONE
                ),
                location=(
                    f"{record.file_name}·{record.sheet}·{record.excel_row}"
                ),
                item_name=str(record.name or "").strip(),
            )
        )
    return requests


def _apply_parts_conversion_plan(
    records: list[RfpRecord],
    plan: ConversionPlan,
) -> list[RfpRecord]:
    """Return converted records; never mutate the source list items."""
    actions = {action.request_id: action for action in plan.actions}
    converted: list[RfpRecord] = []
    for index, record in enumerate(records):
        request_id = f"parts-{index}"
        action = actions.get(request_id)
        if action is None:
            raise UnitsConversionError(
                f"В плане конвертации отсутствует действие для {request_id!r}"
            )
        converted.append(
            replace(
                record,
                values=converted_quantity(action),
                units=converted_unit(action),
                units_check_status=action.status,
                units_conversion_trace=action.trace,
            )
        )
    return converted


def _assert_unified_target_units(records: list[RfpRecord]) -> None:
    """Fail fast before sums if one title+code still has mixed target units."""
    units_by_key: dict[tuple[str, str], set[str]] = defaultdict(set)
    for record in records:
        if not record.units:
            continue
        units_by_key[_record_key(record)].add(normalize_units_text(record.units))
    mixed = [
        (key, sorted(values))
        for key, values in sorted(units_by_key.items())
        if len(values) > 1
    ]
    if mixed:
        sample = ", ".join(
            f"{title}+{code}: {' | '.join(units)}"
            for (title, code), units in mixed[:5]
        )
        extra = f" (+{len(mixed) - 5} ещё)" if len(mixed) > 5 else ""
        raise UnitsConversionError(
            "После конвертации остались смешанные целевые ЕИ: "
            f"{sample}{extra}"
        )


def _merge_units_status_trace(
    bucket: dict[UnitKey, str],
    key: UnitKey,
    value: str,
) -> None:
    if not value:
        return
    bucket[key] = join_unique_units_text([bucket.get(key, ""), value])


def _net_supply_status_flag(raw: str) -> str:
    """Map a source status to the net cell (exclusion, tag replacement, or empty)."""
    parsed = parse_rfp_supply_status(raw)
    if parsed.kind == KIND_EXCLUDED:
        return CANONICAL_EXCLUDED_FROM_SUPPLY
    if parsed.kind == KIND_TAG_REPLACED:
        return parsed.text
    return ""


def _merge_supply_status(
    bucket: dict[UnitKey, str],
    key: UnitKey,
    record: RfpRecord,
    warnings: list[tuple[str, str, str]],
    *,
    emit_warnings: bool = True,
) -> None:
    """Merge exclusion / tag-replacement flags on one aggregation key.

    Mixed values on one key emit WARN. Canonical excluded wins over a
    replacement text; otherwise the non-empty value is kept. Unknown
    statuses never reach this merge (extract is fatal first).
    """
    incoming = _net_supply_status_flag(record.rfp_supply_status)
    previous = bucket.get(key)
    if previous is None:
        bucket[key] = incoming
        return
    if previous == incoming:
        return
    if emit_warnings:
        ds_name, title, code, tag = key
        warnings.append(
            (
                "WARN",
                record.file_name,
                (
                    f"строка {record.excel_row}: смешение статуса поставки "
                    f"на ключе {ds_name}/{title}/{code}/{tag}"
                ),
            )
        )
    if (
        is_canonical_excluded_from_supply(previous)
        or is_canonical_excluded_from_supply(incoming)
    ):
        bucket[key] = CANONICAL_EXCLUDED_FROM_SUPPLY
    else:
        bucket[key] = previous or incoming


def extract_parts_records(parts_dir: Path) -> list[RfpRecord]:
    """Extract parts records from a folder without running the full analyze."""
    warnings: list[tuple[str, str, str]] = []
    status_issues: list[RfpSupplyStatusIssue] = []
    records: list[RfpRecord] = []
    for path in _xlsx_files(parts_dir):
        part_records, _stats = _extract_records(
            path, PARTS_KIND, warnings, status_issues=status_issues
        )
        records.extend(part_records)
    raise_if_rfp_supply_status_issues(status_issues)
    return records


def build_parts_build_deps_snapshot(
    plan: ConversionPlan,
    *,
    google_index: GoogleUnitsIndex,
    matrix_path: Path,
) -> dict[str, Any]:
    """Build deterministic dependency snapshot from an existing conversion plan."""
    relevant_codes = {dep.code for dep in plan.dependencies}
    google_units_by_code = {
        code: google_index.google_unit(code)
        for code in sorted(relevant_codes)
    }
    return {
        "ALGORITHM_VERSION": ALGORITHM_VERSION,
        "AGGREGATION_KEY": AGGREGATION_KEY,
        "matrix_path": str(Path(matrix_path).resolve()),
        "dependencies": [
            {
                "code": item.code,
                "source_unit": item.source_unit,
                "target_unit": item.target_unit,
                "coefficient": str(item.coefficient),
            }
            for item in plan.dependencies
        ],
        "google_units_by_code": google_units_by_code,
    }


def compute_parts_build_deps_snapshot(
    records: list[RfpRecord],
    *,
    google_index: GoogleUnitsIndex,
    matrix_path: Path,
) -> dict[str, Any]:
    """Build dependency snapshot for parts preflight via a fresh plan."""
    requests = _build_parts_conversion_requests(records)
    plan = build_conversion_plan(requests, google_index, matrix_path)
    return build_parts_build_deps_snapshot(
        plan,
        google_index=google_index,
        matrix_path=matrix_path,
    )


def write_parts_build_deps_snapshot(
    out_dir: Path,
    snapshot: dict[str, Any],
) -> Path:
    """Write build dependency snapshot next to ``rfp_parts_net.xlsx``."""
    out_path = Path(out_dir) / BUILD_DEPS_JSON_NAME
    out_path.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return out_path


def _build_unit_counters(
    records: list[RfpRecord],
    warnings: list[tuple[str, str, str]],
    *,
    emit_warnings: bool = True,
) -> PartsAggregation:
    """Build tag-aware unit counters plus decimal buckets for non-unit quantities."""
    unit_counter: Counter[UnitKey] = Counter()
    decimal_quantities: dict[UnitKey, Decimal] = defaultdict(lambda: Decimal("0"))
    examples: dict[UnitKey, RfpRecord] = {}
    units_status_by_key: dict[UnitKey, str] = {}
    units_trace_by_key: dict[UnitKey, str] = {}
    supply_status_by_key: dict[UnitKey, str] = {}

    for record in records:
        tags = _record_tag_keys(record, warnings)

        if tags:
            if emit_warnings and _is_int_decimal(record.values) and int(record.values) != len(tags):
                warnings.append(
                    (
                        "WARN",
                        record.file_name,
                        f"строка {record.excel_row}: теги={tags!r}; количество тегов "
                        f"({len(tags)}) не совпадает с количеством лота ({record.values})",
                    )
                )
            for tag in tags:
                key = _unit_key(record, tag)
                unit_counter[key] += 1
                examples.setdefault(key, record)
                _merge_units_status_trace(
                    units_status_by_key, key, record.units_check_status
                )
                _merge_units_status_trace(
                    units_trace_by_key, key, record.units_conversion_trace
                )
                _merge_supply_status(
                    supply_status_by_key,
                    key,
                    record,
                    warnings,
                    emit_warnings=emit_warnings,
                )
            # Integer leftover (lot > tags) stays on <NO_TAG>; no extra if tags >= lot
            # or if values is not an integer.
            if _is_int_decimal(record.values):
                extra = int(record.values) - len(tags)
                if extra > 0:
                    remainder_key = _unit_key(record, "")
                    unit_counter[remainder_key] += extra
                    examples.setdefault(remainder_key, record)
                    _merge_units_status_trace(
                        units_status_by_key,
                        remainder_key,
                        record.units_check_status,
                    )
                    _merge_units_status_trace(
                        units_trace_by_key,
                        remainder_key,
                        record.units_conversion_trace,
                    )
                    _merge_supply_status(
                        supply_status_by_key,
                        remainder_key,
                        record,
                        warnings,
                        emit_warnings=emit_warnings,
                    )
            continue

        key = _unit_key(record, "")
        examples.setdefault(key, record)
        _merge_units_status_trace(
            units_status_by_key, key, record.units_check_status
        )
        _merge_units_status_trace(
            units_trace_by_key, key, record.units_conversion_trace
        )
        _merge_supply_status(
            supply_status_by_key,
            key,
            record,
            warnings,
            emit_warnings=emit_warnings,
        )
        if record.values < 0:
            warnings.append(
                (
                    "ERROR",
                    record.file_name,
                    f"строка {record.excel_row}: отрицательное количество={record.values}",
                )
            )
            continue
        if _is_int_decimal(record.values):
            unit_counter[key] += int(record.values)
        else:
            decimal_quantities[key] += record.values

    return PartsAggregation(
        units=unit_counter,
        decimals=dict(decimal_quantities),
        examples=examples,
        units_status_by_key=units_status_by_key,
        units_trace_by_key=units_trace_by_key,
        supply_status_by_key=supply_status_by_key,
    )


_TAG_REMARK_KIND_LABEL = {
    "duplicate": "Дубль тега",
    "lot_gt_tags": "Лот > тегов",
    "tags_gt_lot": "Тегов > лота",
}
_TAG_REMARK_KIND_ORDER = {
    "duplicate": 0,
    "lot_gt_tags": 1,
    "tags_gt_lot": 2,
}


def _collect_tag_remarks(records: list[RfpRecord]) -> list[dict[str, Any]]:
    """Collect duplicate-tag and lot-vs-tag-count rows for the dedicated xlsx.

    A tag that appears on more than one source row (any file) is a duplicate.
    Integer lot vs parsed tag count uses ``lot_gt_tags`` / ``tags_gt_lot``
    (same rule as ``_build_unit_counters``). Rows are not written to the GUI
    file table.
    """
    silent: list[tuple[str, str, str]] = []
    parsed: list[tuple[RfpRecord, list[str]]] = []
    tag_to_records: dict[str, list[RfpRecord]] = defaultdict(list)
    for record in records:
        tags = list(dict.fromkeys(_record_tag_keys(record, silent)))
        parsed.append((record, tags))
        for tag in tags:
            tag_to_records[tag].append(record)

    duplicate_tags = {
        tag for tag, hits in tag_to_records.items() if len(hits) > 1
    }
    remarks: list[dict[str, Any]] = []
    for record, tags in parsed:
        if tags and _is_int_decimal(record.values):
            lot_int = int(record.values)
            if lot_int > len(tags):
                remarks.append(
                    _tag_remark_row(
                        kind="lot_gt_tags",
                        record=record,
                        tag="; ".join(tags),
                        tags_count=len(tags),
                        other_locations="",
                    )
                )
            elif len(tags) > lot_int:
                remarks.append(
                    _tag_remark_row(
                        kind="tags_gt_lot",
                        record=record,
                        tag="; ".join(tags),
                        tags_count=len(tags),
                        other_locations="",
                    )
                )
        for tag in tags:
            if tag not in duplicate_tags:
                continue
            others = [
                f"{item.file_name}·стр.{item.excel_row}"
                for item in tag_to_records[tag]
                if item is not record
            ]
            shown = ", ".join(others[:8])
            extra = f" и ещё {len(others) - 8}" if len(others) > 8 else ""
            remarks.append(
                _tag_remark_row(
                    kind="duplicate",
                    record=record,
                    tag=tag,
                    tags_count=len(tags),
                    other_locations=f"{shown}{extra}",
                )
            )
    remarks.sort(
        key=lambda item: (
            _TAG_REMARK_KIND_ORDER.get(item["kind"], 9),
            str(item["tag"]).casefold(),
            str(item["file_name"]).casefold(),
            int(item["excel_row"] or 0),
        )
    )
    return remarks


def _tag_remark_row(
    *,
    kind: str,
    record: RfpRecord,
    tag: str,
    tags_count: int,
    other_locations: str,
) -> dict[str, Any]:
    return {
        "kind": kind,
        "kind_label": _TAG_REMARK_KIND_LABEL.get(kind, kind),
        "file_name": record.file_name,
        "sheet": record.sheet,
        "excel_row": record.excel_row,
        "ds_name": record.ds_name,
        "ds_title": record.ds_title,
        "code": record.code,
        "name": record.name,
        "tag": tag,
        "tags_text": record.tags,
        "tags_count": tags_count,
        "lot_qty": record.values,
        "other_locations": other_locations,
    }


def _write_duplicate_tags_xlsx(
    path: Path,
    remarks: list[dict[str, Any]],
) -> int:
    """Write duplicate-tag / tag-count remarks next to the net workbook.

    Args:
        path: Stamp-folder path ``Отчет по тегам - сбор частей.xlsx``.
        remarks: Output of ``_collect_tag_remarks``.

    Returns:
        Number of data rows written.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Дубли тегов"
    ws.append(
        [
            "Тип",
            "Файл",
            "Лист",
            "Строка",
            "Тег",
            "Кол-во лота",
            "Число тегов",
            "Код",
            "Наименование",
            "Имя ДС",
            "Титул",
            "Теги в ячейке",
            "Где ещё встречается",
        ]
    )
    fill_yellow = PatternFill(
        start_color=Color.soft_yellow,
        end_color=Color.soft_yellow,
        fill_type="solid",
    )
    fill_cyan = PatternFill(
        start_color=Color.soft_cyan,
        end_color=Color.soft_cyan,
        fill_type="solid",
    )
    prev_tag = object()
    fill_toggle = 0
    for item in remarks:
        tag_key = item["tag"] if item["kind"] == "duplicate" else None
        if tag_key is not None and tag_key != prev_tag:
            fill_toggle ^= 1
            prev_tag = tag_key
        ws.append(
            [
                item["kind_label"],
                item["file_name"],
                item["sheet"],
                item["excel_row"],
                item["tag"],
                _qty_excel_value(item["lot_qty"]),
                item["tags_count"],
                item["code"],
                str(item["name"] or "")[:80],
                item["ds_name"],
                item["ds_title"],
                item["tags_text"],
                item["other_locations"],
            ]
        )
        if item["kind"] == "duplicate":
            fill = fill_cyan if fill_toggle else fill_yellow
            for cell in ws[ws.max_row]:
                cell.fill = fill
    widths = {
        "A": 22.0,
        "B": 42.0,
        "C": 22.0,
        "D": 10.0,
        "E": 28.0,
        "F": 12.0,
        "G": 14.0,
        "H": 16.0,
        "I": 40.0,
        "J": 16.0,
        "K": 14.0,
        "L": 36.0,
        "M": 48.0,
    }
    _apply_sheet_widths(ws, widths)
    _finish_data_sheet(ws, has_rows=bool(remarks))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(remarks)


def _write_empty_code_xlsx(
    path: Path,
    rows: list[SkippedEmptyCodeRow],
) -> int:
    """Write skipped empty-CODE positions next to the net workbook.

    Args:
        path: Stamp-folder path ``Отчет по позициям без кода - сбор частей.xlsx``.
        rows: Snapshots collected during extract.

    Returns:
        Number of data rows written. Writes nothing when ``rows`` is empty.
    """
    if not rows:
        return 0
    wb = Workbook()
    ws = wb.active
    ws.title = "Без кода"
    ws.append(
        [
            "Файл",
            "Лист",
            "Строка",
            "Имя ДС",
            "Титул",
            "Теги",
            "Наименование",
            "Тип/марка",
            "Кол-во лота",
            "Ед. изм.",
            "Нет полей",
            "Уровень",
        ]
    )
    for item in rows:
        qty = _parse_decimal(item.values_text)
        qty_cell: float | int | str
        if qty is not None:
            qty_cell = _qty_excel_value(qty)
        else:
            qty_cell = item.values_text
        ws.append(
            [
                item.file_name,
                item.sheet,
                item.excel_row,
                item.ds_name,
                item.ds_title,
                item.tags,
                item.name,
                item.type_mark,
                qty_cell,
                item.units,
                item.missing_fields,
                item.level,
            ]
        )
    widths = {
        "A": 42.0,
        "B": 22.0,
        "C": 10.0,
        "D": 16.0,
        "E": 16.0,
        "F": 28.0,
        "G": 40.0,
        "H": 28.0,
        "I": 14.0,
        "J": 12.0,
        "K": 28.0,
        "L": 12.0,
    }
    _apply_sheet_widths(ws, widths)
    _finish_data_sheet(ws, has_rows=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(rows)


def _compute_coarse_collisions(records: list[RfpRecord]) -> CoarseCollisionResult:
    """Aggregate title_system+code checks (units mismatch across summed parts)."""
    zero = Decimal("0")
    qty_by_key: dict[tuple[str, str], Decimal] = defaultdict(lambda: zero)
    examples: dict[tuple[str, str], RfpRecord] = {}
    units_by_key: dict[tuple[str, str], set[str]] = defaultdict(set)
    records_by_key: dict[tuple[str, str], list[RfpRecord]] = defaultdict(list)
    for record in records:
        key = _record_key(record)
        qty_by_key[key] += record.values
        examples.setdefault(key, record)
        records_by_key[key].append(record)
        if record.units:
            units_by_key[key].add(normalize_units_text(record.units))

    ambiguous_meta: list[tuple[tuple[str, str], RfpRecord, tuple[str, ...]]] = []
    for key in sorted(qty_by_key):
        units_variants = tuple(sorted(units_by_key.get(key, ())))
        if len(units_variants) > 1:
            ambiguous_meta.append((key, examples[key], units_variants))

    return CoarseCollisionResult(
        qty_by_key=dict(qty_by_key),
        examples=examples,
        ambiguous_meta=ambiguous_meta,
        records_by_key={
            key: records_by_key[key] for key, _, _ in ambiguous_meta
        },
    )


def _add_aggregation_report(
    lines: list[str],
    aggregation: PartsAggregation,
) -> None:
    """Append summed unit/decimal totals to the text report."""
    summary_table = PrettyTable()
    summary_table.field_names = ["Показатель", "Значение"]
    summary_table.add_row(["Единичных строк (сумма частей)", sum(aggregation.units.values())])
    summary_table.add_row(["Ключей единичных строк", len(aggregation.units)])
    summary_table.add_row(
        ["Дробное количество (сумма частей)", sum(aggregation.decimals.values(), Decimal("0"))]
    )
    summary_table.add_row(["Ключей дробного количества", len(aggregation.decimals)])
    _add_table(
        lines,
        "Сводка суммирования частей",
        summary_table,
        "Одинаковые ключи ds_name + title_system + code + tag/<NO_TAG> складываются; "
        "разные ДС не склеиваются.",
    )


def _add_table(lines: list[str], title: str, table: PrettyTable, note: str = "") -> None:
    lines.append("")
    lines.append(title)
    if note:
        lines.append(note)
    lines.append(str(table))


def _split_warning_row(message: str) -> tuple[str, str]:
    match = re.match(r"(?:row|строка)\s+(\d+)\s*:\s*(.*)", message, re.IGNORECASE)
    if not match:
        return "", message
    return match.group(1), match.group(2)


def _kind_label(kind: str) -> str:
    return {
        "parts": "части",
        "increase": "увеличение",
        "decrease": "уменьшение",
        "summary": "сводная",
    }.get(kind, kind or "")


def _level_label(level: str) -> str:
    return {
        "ERROR": "Ошибка",
        "WARN": "Предупреждение",
        "INFO": "Информация",
    }.get(level, level)


def _diagnostic_rows(
    warnings: list[tuple[str, str, str]],
    file_stats: list[FileStats],
    limit: int | None = None,
) -> list[list[str]]:
    kind_by_file = {stats.file_name: stats.kind for stats in file_stats}
    rows: list[list[str]] = []
    source_warnings = warnings if limit is None else warnings[:limit]
    for level, file_name, message in source_warnings:
        row_number, clean_message = _split_warning_row(humanize_parts_message(message))
        rows.append(
            [
                _level_label(level),
                _kind_label(kind_by_file.get(file_name, "")),
                file_name,
                row_number,
                clean_message,
            ]
        )
    return rows


def _write_summary_increase_diff_xlsx(
    path: Path,
    summary_records: list[RfpRecord],
    increase_records: list[RfpRecord],
) -> int:
    """Write summary-vs-increase differences to XLSX and return row count."""
    summary_by_key = _sum_by_key(summary_records)
    increase_by_key = _sum_by_key(increase_records)
    diff_keys = [
        key
        for key in sorted(set(summary_by_key) | set(increase_by_key))
        if summary_by_key.get(key, Decimal("0")) != increase_by_key.get(key, Decimal("0"))
    ]
    diff_rows = [
        (
            increase_by_key.get(key, Decimal("0")) - summary_by_key.get(key, Decimal("0")),
            key[0],
            key[1],
            summary_by_key.get(key, Decimal("0")),
            increase_by_key.get(key, Decimal("0")),
        )
        for key in diff_keys
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "summary_vs_increase"
    ws.append(["delta", "title_system", "code", "summary", "increase"])
    for delta, title_system, code, summary_value, increase_value in sorted(
        diff_rows,
        key=lambda item: abs(item[0]),
        reverse=True,
    ):
        ws.append([delta, title_system, code, summary_value, increase_value])

    widths = {
        "A": 18,
        "B": 32,
        "C": 18,
        "D": 18,
        "E": 18,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(diff_rows)


def _record_to_xlsx_row(record: RfpRecord) -> list[Any]:
    return [
        record.kind,
        record.file_name,
        record.excel_row,
        record.ds_name,
        record.ds_number,
        record.ds_title,
        record.ds_specification,
        record.tags,
        record.code,
        record.name,
        record.type_mark,
        record.values,
        record.units,
    ]


def _write_records_sheet(wb: Workbook, sheet_name: str, records: list[RfpRecord]) -> None:
    ws = wb.active
    ws.title = sheet_name
    headers = [
        "Вид",
        "Файл",
        "Строка",
        "Имя ДС",
        "№ п/п",
        "Титул",
        "Спецификация",
        "Теги",
        "Код РД",
        "Наименование",
        "Техн. характеристики",
        "Кол-во",
        "Ед. изм.",
    ]
    ws.append(headers)
    for record in records:
        kind_ru = {
            "parts": "части",
            "increase": "увеличение",
            "decrease": "уменьшение",
        }.get(
            record.kind, record.kind
        )
        row = _record_to_xlsx_row(record)
        row[0] = kind_ru
        ws.append(row)

    widths = {
        "A": 12,
        "B": 55,
        "C": 10,
        "D": 10,
        "E": 8,
        "F": 16,
        "G": 28,
        "H": 25,
        "I": 15,
        "J": 36,
        "K": 28,
        "L": 10,
        "M": 10,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _write_records_xlsx(
    path: Path,
    sheet_name: str,
    records: list[RfpRecord],
) -> int:
    """Write a full source records table to a single-sheet workbook."""
    wb = Workbook()
    _write_records_sheet(wb, sheet_name, records)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(records)


def _write_diagnostics_xlsx(
    path: Path,
    warnings: list[tuple[str, str, str]],
    file_stats: list[FileStats],
) -> int:
    """Write warnings/diagnostics table to XLSX with a legend sheet."""
    wb = Workbook()
    legend_ws = wb.active
    legend_ws.title = "legend"
    legend_ws.append(["Описание"])
    for item in DIAGNOSTICS_LEGEND:
        legend_ws.append([item])
    legend_ws.column_dimensions["A"].width = 110
    legend_ws.auto_filter.ref = f"A1:A{legend_ws.max_row}"
    legend_ws.freeze_panes = "A2"

    ws = wb.create_sheet("diagnostics")
    headers = ["Уровень", "Папка RFP", "Файл", "Строка", "Сообщение"]
    ws.append(headers)
    rows = _diagnostic_rows(warnings, file_stats)
    for row in rows:
        ws.append(row)
    widths = {
        "A": 16,
        "B": 14,
        "C": 55,
        "D": 10,
        "E": 105,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(rows)


def _clone_net_record(
    example: RfpRecord,
    *,
    values: Decimal,
    tags: str,
    ds_number: str,
    kind: str = "net",
    units_check_status: str = "",
    units_conversion_trace: str = "",
    rfp_supply_status: str | None = None,
) -> RfpRecord:
    """Build a summary-row RfpRecord from a parts-folder example."""
    supply_status = (
        example.rfp_supply_status
        if rfp_supply_status is None
        else rfp_supply_status
    )
    return RfpRecord(
        kind=kind,
        file_name=example.file_name,
        sheet=example.sheet,
        excel_row=example.excel_row,
        ds_name=example.ds_name,
        ds_number=ds_number,
        ds_title=example.ds_title,
        ds_specification=example.ds_specification,
        tags=tags,
        ds_code_1c=example.ds_code_1c,
        code=example.code,
        name=example.name,
        type_mark=example.type_mark,
        values=values,
        units=example.units,
        units_check_status=units_check_status,
        units_conversion_trace=units_conversion_trace,
        rfp_supply_status=supply_status,
    )


@dataclass
class NetSummaryRow:
    """One row of the combined summary workbook (data + collision filter flags)."""

    record: RfpRecord
    source: str
    flag_unmatched_decrease: str = ""
    flag_negative_net: str = ""
    flag_units_mismatch: str = ""
    flag_via_replacement: str = ""

    @property
    def flag_any_collision(self) -> str:
        if any(
            (
                self.flag_unmatched_decrease,
                self.flag_negative_net,
                self.flag_units_mismatch,
                self.flag_via_replacement,
            )
        ):
            return "Да"
        return ""


def _append_summary_unit_rows(
    rows: list[NetSummaryRow],
    *,
    example: RfpRecord,
    remaining: int | Decimal,
    tag: str,
    tag_is_unit: bool,
    seq_start: int,
    units_mismatch: set[tuple[str, str]],
    unit_key: UnitKey,
    units_check_status: str = "",
    units_conversion_trace: str = "",
    rfp_supply_status: str = "",
) -> int:
    """Append one or more summary rows; return next sequence number."""
    _ds_name, title_system, code, tag_key = unit_key
    title_code = (title_system, code)
    flag_units = "Да" if title_code in units_mismatch else ""
    tag_text = "" if tag_key == NO_TAG_KEY else tag
    seq = seq_start

    def _one(values: Decimal) -> None:
        nonlocal seq
        seq += 1
        rows.append(
            NetSummaryRow(
                record=_clone_net_record(
                    example,
                    values=values,
                    tags=tag_text,
                    ds_number=str(seq),
                    kind="net",
                    units_check_status=units_check_status,
                    units_conversion_trace=units_conversion_trace,
                    rfp_supply_status=rfp_supply_status,
                ),
                source="Сумма частей",
                flag_units_mismatch=flag_units,
            )
        )

    if tag_is_unit and tag_key != NO_TAG_KEY:
        for _ in range(int(remaining)):
            _one(Decimal(1))
    else:
        qty = Decimal(remaining) if not isinstance(remaining, Decimal) else remaining
        _one(qty)
    return seq


def _build_summary_rows(
    aggregation: PartsAggregation,
    coarse: CoarseCollisionResult,
) -> list[NetSummaryRow]:
    """Build the summary by summing parts (tag-expanded units + decimal leftovers)."""
    units_mismatch = {key for key, _, _ in coarse.ambiguous_meta}
    rows: list[NetSummaryRow] = []
    seq = 0

    for key in sorted(aggregation.units):
        remaining = aggregation.units.get(key, 0)
        if remaining <= 0:
            continue
        example = aggregation.examples.get(key)
        if example is None:
            continue
        status = aggregation.units_status_by_key.get(key, "")
        trace = aggregation.units_trace_by_key.get(key, "")
        supply = aggregation.supply_status_by_key.get(key, "")
        seq = _append_summary_unit_rows(
            rows,
            example=example,
            remaining=remaining,
            tag=key[3],
            tag_is_unit=True,
            seq_start=seq,
            units_mismatch=units_mismatch,
            unit_key=key,
            units_check_status=status,
            units_conversion_trace=trace,
            rfp_supply_status=supply,
        )

    for key in sorted(aggregation.decimals):
        remaining = aggregation.decimals.get(key, Decimal("0"))
        if remaining <= Decimal("0"):
            continue
        example = aggregation.examples.get(key)
        if example is None:
            continue
        status = aggregation.units_status_by_key.get(key, "")
        trace = aggregation.units_trace_by_key.get(key, "")
        supply = aggregation.supply_status_by_key.get(key, "")
        seq = _append_summary_unit_rows(
            rows,
            example=example,
            remaining=remaining,
            tag=key[3],
            tag_is_unit=False,
            seq_start=seq,
            units_mismatch=units_mismatch,
            unit_key=key,
            units_check_status=status,
            units_conversion_trace=trace,
            rfp_supply_status=supply,
        )

    return rows


def _build_no_tags_summary_rows(
    records: list[RfpRecord],
    coarse: CoarseCollisionResult,
) -> list[NetSummaryRow]:
    """One net row per source part: lot qty as in the DS, TAGS blank.

    Does not expand tags into ``VALUES=1`` rows and does not sum several
    source lines of the same code into one qty. N original positions stay
    N rows (``load_tags=false`` / ``rfp_parts_net_no_tags.xlsx``).
    """
    units_mismatch = {key for key, _, _ in coarse.ambiguous_meta}
    rows: list[NetSummaryRow] = []
    seq = 0
    for record in records:
        if record.values <= 0:
            continue
        seq += 1
        flag_units = "Да" if _record_key(record) in units_mismatch else ""
        rows.append(
            NetSummaryRow(
                record=_clone_net_record(
                    record,
                    values=record.values,
                    tags="",
                    ds_number=str(seq),
                    kind="net",
                    units_check_status=record.units_check_status,
                    units_conversion_trace=record.units_conversion_trace,
                    rfp_supply_status=_net_supply_status_flag(
                        record.rfp_supply_status
                    ),
                ),
                source="Сумма частей (без тегов)",
                flag_units_mismatch=flag_units,
            )
        )
    return rows


# Contiguous RFP_AGGREAGATED integer slots that Step1 actually reads (0–10).
# Extra slots 14/15/17/18/19 are padded; filters start after 19 (U–Z).
_STEP1_NET_CORE_MAX_INDEX = 10

_STEP1_NET_HEADER_BY_FIELD = {
    DS_NAME: "Имя ДС",
    DS_NUMBER: "№ п/п",
    DS_TITLE: "Титул",
    DS_SPECIFICATION: "Спецификация",
    TAGS: "Теги / Tag",
    DS_CODE_1C: "Тип оборудования",  # occupies the 1C slot; Step1 treats it optional
    CODE: "Код РД / BCC",
    NAME: "Наименование МТР",
    TYPE_MARK: "Технические характеристики",
    VALUES: "Кол-во",
    UNITS: "Ед. изм.",
    NAME_2: "Наименование 2",
    RFP_SUPPLY_STATUS: "Статус поставки",
    VALUES_2: "Закупка по Лоту, кол-во",
    UNITS_CHECK_STATUS: "Статус ед. изм.",
    UNITS_CONVERSION_TRACE: "Trace ед. изм.",
}

_NET_FILTER_HEADERS = (
    "Источник строки",
    "Есть коллизия",
    "Уменьшение без пары",
    "Отрицательный итог",
    "Разные ед. изм.",
    "Через замену кода",
)

_STEP1_NET_WIDTH_BY_FIELD = {
    DS_NAME: 10,
    DS_NUMBER: 8,
    DS_TITLE: 16,
    DS_SPECIFICATION: 34,
    TAGS: 28,
    DS_CODE_1C: 10,
    CODE: 14,
    NAME: 40,
    TYPE_MARK: 28,
    VALUES: 10,
    UNITS: 10,
    NAME_2: 16,
    RFP_SUPPLY_STATUS: 22,
    VALUES_2: 12,
    UNITS_CHECK_STATUS: 18,
    UNITS_CONVERSION_TRACE: 40,
}
_NET_FILTER_WIDTHS = (22, 14, 20, 18, 18, 18)


def step1_net_core_columns() -> dict[int, str]:
    """Return Excel indices Step1 reads from ``rfp_parts_net.xlsx``.

    Source of truth: integer keys 0..10 in ``ColNames.RFP_AGGREAGATED.column_dict``.
    Fractional keys (VENDOR 8.5) are skipped — ``load_rows_from_worksheet`` only
    sees integer ``cell_index``. Do not insert extra columns in this range:
    ``VALUES`` must stay at 9 and ``UNITS`` at 10.

    Returns:
        Mapping of 0-based column index to ``ColNames`` field id.
    """
    return {
        int(idx): name
        for idx, name in ColNames.RFP_AGGREAGATED.column_dict.items()
        if isinstance(idx, int) and 0 <= idx <= _STEP1_NET_CORE_MAX_INDEX
    }


def step1_net_extra_columns() -> dict[int, str]:
    """Integer ``RFP_AGGREAGATED`` slots after the core 0–10 block.

    ``NAME_2`` is 14, ``RFP_SUPPLY_STATUS`` is 15, ``VALUES_2`` is 17,
    units status/trace are 18/19. Filter flags must start after ``max(extra)``
    so they cannot occupy these indices.

    Returns:
        Mapping of 0-based column index to ``ColNames`` field id.
    """
    return {
        int(idx): name
        for idx, name in ColNames.RFP_AGGREAGATED.column_dict.items()
        if isinstance(idx, int) and idx > _STEP1_NET_CORE_MAX_INDEX
    }


def _ensure_net_layout_matches_step1() -> tuple[dict[int, str], dict[int, str]]:
    """Fail fast if the net writer drifted from Step1 ``RFP_AGGREAGATED`` slots."""
    core = step1_net_core_columns()
    extra = step1_net_extra_columns()
    core_header_fields = {
        field
        for field in _STEP1_NET_HEADER_BY_FIELD
        if field not in extra.values()
    }
    missing = sorted(core_header_fields - set(core.values()))
    unexpected = sorted(set(core.values()) - core_header_fields)
    if missing or unexpected:
        raise RuntimeError(
            "rfp_parts_net.xlsx must follow Step1 RFP_AGGREAGATED indices 0–10; "
            f"missing={missing}; unexpected={unexpected}"
        )
    if core.get(6) != CODE or core.get(9) != VALUES or core.get(10) != UNITS:
        raise RuntimeError(
            "Step1 RFP_AGGREAGATED moved CODE/VALUES/UNITS; update "
            "_write_net_xlsx before writing rfp_parts_net.xlsx: "
            f"{core!r}"
        )
    if extra.get(17) != VALUES_2:
        raise RuntimeError(
            "Step1 RFP_AGGREAGATED index 17 must stay VALUES_2 for the net "
            f"writer; got extra={extra!r}"
        )
    if extra.get(18) != UNITS_CHECK_STATUS or extra.get(19) != UNITS_CONVERSION_TRACE:
        raise RuntimeError(
            "Step1 RFP_AGGREAGATED indices 18/19 must stay "
            "UNITS_CHECK_STATUS / UNITS_CONVERSION_TRACE for the net writer; "
            f"got extra={extra!r}"
        )
    if extra.get(15) != RFP_SUPPLY_STATUS:
        raise RuntimeError(
            "Step1 RFP_AGGREAGATED index 15 must stay RFP_SUPPLY_STATUS "
            f"for the net writer; got extra={extra!r}"
        )
    return core, extra


def _net_step1_width() -> int:
    core, extra = _ensure_net_layout_matches_step1()
    return max(max(core), max(extra)) + 1


def _net_header_row() -> list[str]:
    core, extra = _ensure_net_layout_matches_step1()
    header = [""] * _net_step1_width()
    for idx, field in {**core, **extra}.items():
        header[idx] = _STEP1_NET_HEADER_BY_FIELD.get(field, field)
    header.extend(_NET_FILTER_HEADERS)
    return header


def _net_data_cells(
    item: NetSummaryRow,
    *,
    equipment_type: str,
    qty: float | int,
) -> list[object]:
    core, extra = _ensure_net_layout_matches_step1()
    record = item.record
    by_field: dict[str, object] = {
        DS_NAME: record.ds_name,
        DS_NUMBER: record.ds_number,
        DS_TITLE: record.ds_title,
        DS_SPECIFICATION: record.ds_specification,
        TAGS: record.tags,
        DS_CODE_1C: equipment_type,
        CODE: record.code,
        NAME: record.name,
        TYPE_MARK: record.type_mark,
        VALUES: qty,
        UNITS: record.units,
        NAME_2: "",
        RFP_SUPPLY_STATUS: record.rfp_supply_status,
        VALUES_2: qty,
        UNITS_CHECK_STATUS: record.units_check_status,
        UNITS_CONVERSION_TRACE: record.units_conversion_trace,
    }
    row: list[object] = [""] * _net_step1_width()
    for idx, field in {**core, **extra}.items():
        row[idx] = by_field.get(field, "")
    row.extend(
        [
            item.source,
            item.flag_any_collision,
            item.flag_unmatched_decrease,
            item.flag_negative_net,
            item.flag_units_mismatch,
            item.flag_via_replacement,
        ]
    )
    return row


def _net_column_widths(core: dict[int, str], extra: dict[int, str]) -> dict[str, float]:
    widths: dict[str, float] = {}
    for idx, field in {**core, **extra}.items():
        widths[get_column_letter(idx + 1)] = float(
            _STEP1_NET_WIDTH_BY_FIELD.get(field, 12)
        )
    filter_start = max(max(core), max(extra)) + 2
    for offset, width in enumerate(_NET_FILTER_WIDTHS):
        widths[get_column_letter(filter_start + offset)] = float(width)
    return widths


def _equipment_type_from_tags(tags: str) -> str:
    """Derive equipment codes from tag text (Step4 ``EQUIPMENT_TYPE_STATUS``).

    Parses each tag via ``TagClass`` and joins unique ``equipment`` parts with ``/``.
    """
    if not tags:
        return ""
    if isinstance(tags, list):
        tag_list = [str(t).strip() for t in tags if str(t).strip()]
    else:
        raw = str(tags).strip()
        if not raw:
            return ""
        tag_list = [t.strip() for t in raw.replace(";", ",").split(",") if t.strip()]
        if len(tag_list) == 1 and " " in tag_list[0] and "-" in tag_list[0]:
            # Single cell may hold one tag without commas.
            tag_list = [tag_list[0]]
    equipment_values: list[str] = []
    for tag in tag_list:
        try:
            equipment = TagClass(str(tag)).equipment or ""
        except Exception:
            equipment = ""
        if equipment and equipment not in equipment_values:
            equipment_values.append(equipment)
    return "/".join(equipment_values)


@dataclass(frozen=True)
class GoogleCodeLookup:
    """Google code base indexed by normalized BCC."""

    equipment_by_code: dict[str, str]
    units_by_code: dict[str, str]


def _row_field_text(row: Any, field: str) -> str:
    el = getattr(row, "el", {}).get(field) if hasattr(row, "el") else None
    if el is None:
        return ""
    value = getattr(el, "value", None)
    if value is None:
        return ""
    return str(value).strip()


def _build_google_code_lookup(rows: list[Any]) -> GoogleCodeLookup:
    """Map normalized BCC/CODE → Google ``EQUIPMENT_CODE`` and ``UNITS``.

    First non-empty value for each field wins.
    """
    equipment: dict[str, str] = {}
    units: dict[str, str] = {}
    for row in rows or []:
        code_el = getattr(row, "el", {}).get(CODE) if hasattr(row, "el") else None
        if code_el is None:
            continue
        code_key = _normalize_code(code_el.value)
        if not code_key:
            continue
        if code_key not in equipment:
            eq_text = _row_field_text(row, EQUIPMENT_CODE)
            if eq_text:
                equipment[code_key] = eq_text
        if code_key not in units:
            units_text = _row_field_text(row, UNITS)
            if units_text:
                units[code_key] = units_text
    return GoogleCodeLookup(equipment_by_code=equipment, units_by_code=units)


def _load_google_code_lookup() -> GoogleCodeLookup:
    """Load Google code base (cache/export) and index EQUIPMENT_CODE + UNITS by CODE."""
    empty = GoogleCodeLookup(equipment_by_code={}, units_by_code={})
    try:
        print("loading google code base (EQUIPMENT_CODE, UNITS)…")
        rows = load_base() or []
        mapping = _build_google_code_lookup(rows)
        print(
            "google indexed: "
            f"EQUIPMENT_CODE={len(mapping.equipment_by_code)} "
            f"UNITS={len(mapping.units_by_code)} codes"
        )
        return mapping
    except Exception as exc:
        print(
            f"WARN: google code base unavailable for equipment/units columns: "
            f"{type(exc).__name__}: {exc}"
        )
        return empty


def _google_units_for_code(
    code: str,
    units_by_code: dict[str, str] | None,
) -> str:
    if not units_by_code:
        return ""
    code_key = _normalize_code(code)
    if not code_key:
        return ""
    return units_by_code.get(code_key, "")


def _normalize_units_compare(value: str) -> str:
    """Strip whitespace and trailing dots so ``шт.`` matches ``шт``."""
    return normalize_units_text(value)


def _google_units_matches_variants(
    google_units: str, variants: tuple[str, ...]
) -> bool:
    google_norm = _normalize_units_compare(google_units)
    if not google_norm:
        return False
    return any(_normalize_units_compare(item) == google_norm for item in variants)


def _units_text_equal(left: str, right: str) -> bool:
    """Exact UNITS match after strip (``шт.`` ≠ ``шт``)."""
    return str(left or "").strip() == str(right or "").strip()


def _qty_excel_value(value: Decimal) -> float | int:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _collect_units_fix_rows(
    coarse: CoarseCollisionResult,
    units_by_code: dict[str, str] | None,
) -> list[UnitsFixRow]:
    """Source RFP lines whose UNITS should be aligned to Google.

    Scope is colliding keys (more than one unit spelling). If Google UNITS is
    known, skip lines that already match it exactly; otherwise keep every
    source line of the key for inspection.
    """
    out: list[UnitsFixRow] = []
    for key, example, variants in coarse.ambiguous_meta:
        google_units = _google_units_for_code(key[1], units_by_code)
        records = coarse.records_by_key.get(key) or [example]
        for record in records:
            if google_units and _units_text_equal(record.units, google_units):
                continue
            out.append(
                UnitsFixRow(
                    title_system=key[0],
                    code=key[1],
                    google_units=google_units,
                    file_units=record.units,
                    units_variants=variants,
                    record=record,
                )
            )
    out.sort(
        key=lambda item: (
            item.record.file_name.lower(),
            item.record.sheet.lower(),
            item.record.excel_row,
            item.code,
        )
    )
    return out


def _equipment_type_for_record(
    record: RfpRecord,
    *,
    equipment_by_code: dict[str, str] | None = None,
) -> str:
    """Equipment type from tags; if no tags — Google EQUIPMENT_CODE by BCC."""
    from_tags = _equipment_type_from_tags(record.tags)
    if from_tags:
        return from_tags
    if not equipment_by_code:
        return ""
    code_key = _normalize_code(record.code)
    if not code_key:
        return ""
    return equipment_by_code.get(code_key, "")


# Collision row fills (shared with legend sheet).
_COLLISION_FILL_UNITS = PatternFill(
    fill_type="solid", fgColor=Color.match_not_found
)
_COLLISION_FILL_OTHER = PatternFill(
    fill_type="solid", fgColor=Color.match_not_found
)
_COLLISION_FILL_OK = PatternFill(
    fill_type="solid", fgColor=Color.match_matched
)
_COLLISION_FILL_GOOGLE_UNITS_DIFF = PatternFill(
    fill_type="solid", fgColor=Color.yellow
)
_HEADER_FONT = Font(bold=True)
_TOTAL_FONT = Font(bold=True)

_COLLISION_LEGEND_ROWS: tuple[tuple[str, str, PatternFill, str], ...] = (
    (
        "—",
        "Без коллизий (сумма частей в своде)",
        _COLLISION_FILL_OK,
        Color.match_matched,
    ),
    (
        "1",
        "Разные ед. изм. (по титулу + коду)",
        _COLLISION_FILL_UNITS,
        Color.match_not_found,
    ),
)


def _collision_row_fill(item: NetSummaryRow) -> PatternFill | None:
    """Row fill for свод collision flags (units mismatch)."""
    if not item.flag_any_collision:
        return None
    if item.flag_units_mismatch:
        return _COLLISION_FILL_UNITS
    return _COLLISION_FILL_OTHER


def _write_collision_legend_sheet(wb: Workbook) -> None:
    """Add a legend sheet with collision color samples."""
    ws = wb.create_sheet("Легенда коллизий")
    ws.append(["Приоритет", "Пример цвета", "Смысл", "HEX"])
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    for priority, meaning, fill, hex_code in _COLLISION_LEGEND_ROWS:
        ws.append([priority, "", meaning, hex_code.upper()])
        sample = ws.cell(row=ws.max_row, column=2)
        sample.fill = fill
        sample.value = "■■■"
    ws.append([])
    ws.append(
        [
            "",
            "",
            "При нескольких флагах на строке применяется заливка с наивысшим приоритетом.",
            "",
        ]
    )
    ws.append(
        [
            "",
            "",
            "Фильтры коллизий — столбцы U–Z на листе свода (актуален «Разные ед. изм.»). "
            "Матрица по титулам — лист «Статистика».",
            "",
        ]
    )
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 12
    ws.freeze_panes = "A2"


def _write_title_stats_sheet(wb: Workbook, rows: list[NetSummaryRow]) -> None:
    """Matrix: title/mark × OK / collision kinds (counts from свод rows)."""
    by_title: dict[str, Counter[str]] = defaultdict(Counter)
    for item in rows:
        title = (item.record.ds_title or "").strip() or "(без титула)"
        c = by_title[title]
        c["total"] += 1
        if item.flag_any_collision:
            c["any"] += 1
        else:
            c["ok"] += 1
        if item.flag_units_mismatch:
            c["units"] += 1

    ws = wb.create_sheet("Статистика", 1)
    headers = [
        "Титул / Марка",
        "Всего строк",
        "Без коллизий",
        "С коллизией",
        "Разные ед. изм.",
    ]
    header_fills = [
        None,
        None,
        _COLLISION_FILL_OK,
        _COLLISION_FILL_UNITS,
        _COLLISION_FILL_UNITS,
    ]
    ws.append(headers)
    for col_idx, cell in enumerate(ws[1], start=1):
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            wrap_text=True, vertical="center", horizontal="center"
        )
        fill = header_fills[col_idx - 1]
        if fill is not None:
            cell.fill = fill

    totals: Counter[str] = Counter()
    col_fills: dict[int, PatternFill] = {
        3: _COLLISION_FILL_OK,
        4: _COLLISION_FILL_UNITS,
        5: _COLLISION_FILL_UNITS,
    }

    def _paint_metric_cells(row_idx: int, counts: Counter[str]) -> None:
        values_by_col = {
            3: counts["ok"],
            4: counts["any"],
            5: counts["units"],
        }
        for col_idx, fill in col_fills.items():
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            if values_by_col.get(col_idx, 0) > 0:
                cell.font = Font(bold=True)

    for title in sorted(by_title):
        c = by_title[title]
        ws.append(
            [
                title,
                c["total"],
                c["ok"],
                c["any"],
                c["units"],
            ]
        )
        _paint_metric_cells(ws.max_row, c)
        for key in ("total", "ok", "any", "units"):
            totals[key] += c[key]

    ws.append(
        [
            "ИТОГО",
            totals["total"],
            totals["ok"],
            totals["any"],
            totals["units"],
        ]
    )
    totals_row = ws.max_row
    for cell in ws[totals_row]:
        cell.font = _TOTAL_FONT
    _paint_metric_cells(totals_row, totals)
    ws.cell(row=totals_row, column=1).font = _TOTAL_FONT
    ws.cell(row=totals_row, column=2).font = _TOTAL_FONT

    ws.append([])
    note_row = ws.max_row + 1
    ws.append(
        [
            "Примечание: считается по строкам свода (сумма частей). "
            "Цвета столбцов — как на листе «Легенда коллизий»."
        ]
    )
    ws.merge_cells(
        start_row=note_row, start_column=1, end_row=note_row, end_column=5
    )

    for col_letter, width in {
        "A": 16,
        "B": 11,
        "C": 13,
        "D": 12,
        "E": 18,
    }.items():
        ws.column_dimensions[col_letter].width = width
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(5)}{totals_row}"


def _write_net_xlsx(
    path: Path,
    rows: list[NetSummaryRow],
    *,
    equipment_by_code: dict[str, str] | None = None,
) -> int:
    """Write the summed parts workbook that Step1 later loads.

    Core columns A–K follow integer keys 0–10 of ``RFP_AGGREAGATED``
    (``step1_net_core_columns``): ``VALUES`` at J, ``UNITS`` at K. Column F
    holds equipment type in the ``DS_CODE_1C`` slot. Do not insert VENDOR
    between TYPE_MARK and VALUES — the xlsx loader never reads key 8.5.
    Index 17 is ``VALUES_2`` (same lot qty as ``VALUES``) so Step1
    ``get_row_type`` sees a non-empty lot column. Index 15 is
    ``RFP_SUPPLY_STATUS``. Indices 18/19 hold units conversion status/trace.
    Filter flags start after index 19 (U–Z).
    Collision rows are filled (see ``_collision_row_fill``);
    sheets ``Статистика`` and ``Легенда коллизий`` summarize by title and colors.
    """
    core, extra = _ensure_net_layout_matches_step1()
    wb = Workbook()
    ws = wb.active
    ws.title = EXPECTED_SHEET_NAME
    ws.append(_net_header_row())
    for item in rows:
        record = item.record
        qty: float | int
        if record.values == record.values.to_integral_value():
            qty = int(record.values)
        else:
            qty = float(record.values)
        ws.append(
            _net_data_cells(
                item,
                equipment_type=_equipment_type_for_record(
                    record, equipment_by_code=equipment_by_code
                ),
                qty=qty,
            )
        )
        fill = _collision_row_fill(item)
        if fill is not None:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    for col_letter, width in _net_column_widths(core, extra).items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    _write_title_stats_sheet(wb, rows)
    _write_collision_legend_sheet(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(rows)


def _apply_sheet_widths(ws, widths: dict[str, float]) -> None:
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


def _finish_data_sheet(ws, *, has_rows: bool) -> None:
    ws.freeze_panes = "A2"
    if has_rows and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


_COLLISION_WIDTHS = {
    "Сводка": {"A": 52.0, "B": 14.0},
    "К исправлению": {
        "A": 36.0,
        "B": 24.0,
        "C": 10.0,
        "D": 12.0,
        "E": 18.0,
        "F": 16.0,
        "G": 16.0,
        "H": 16.0,
        "I": 44.0,
        "J": 10.0,
        "K": 22.0,
    },
    "Ключи": {
        "A": 20.71,
        "B": 17.86,
        "C": 17.57,
        "D": 16.0,
        "E": 18.0,
        "F": 12.0,
        "G": 44.0,
    },
}


def _write_collisions_xlsx(
    path: Path,
    coarse: CoarseCollisionResult,
    *,
    units_by_code: dict[str, str] | None = None,
) -> int:
    """Write collision workbook for aligning source UNITS to Google.

    Sheet ``К исправлению`` lists every source RFP line (file / sheet / Excel
    row) whose UNITS does not exactly match Google. Lines that already have
    the Google spelling are omitted. Sheet ``Ключи`` keeps one row per
    title+code with all unit variants.
    """
    fix_rows = _collect_units_fix_rows(coarse, units_by_code)
    google_hit = 0
    google_miss = 0
    google_diff = 0
    for key, _, units_variants in coarse.ambiguous_meta:
        google_units = _google_units_for_code(key[1], units_by_code)
        if not google_units:
            google_miss += 1
        elif _google_units_matches_variants(google_units, units_variants):
            google_hit += 1
        else:
            google_diff += 1
    files_to_fix = {item.record.file_name for item in fix_rows}

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Сводка"
    summary_rows = [
        ["Показатель", "Значение"],
        ["Ключей в своде", len(coarse.qty_by_key)],
        ["Ключей с разными ед. изм.", len(coarse.ambiguous_meta)],
        ["Строк к исправлению", len(fix_rows)],
        ["Файлов с ошибкой", len(files_to_fix)],
        ["Ед. изм. Google среди вариантов", google_hit],
        ["Ед. изм. Google не среди вариантов", google_diff],
        ["Нет ед. изм. в Google", google_miss],
    ]
    for row in summary_rows:
        ws_summary.append(row)
    ws_summary.append([])
    ws_summary.append(
        [
            "Лист «К исправлению»: только строки исходных RFP, где ед. изм. "
            "не совпадает с Google. Колонки Файл / Лист / Строка — координаты "
            "ячейки в исходном файле. Строки с уже верной ед. изм. скрыты."
        ]
    )
    _finish_data_sheet(ws_summary, has_rows=True)
    _apply_sheet_widths(ws_summary, _COLLISION_WIDTHS["Сводка"])

    ws_fix = wb.create_sheet("К исправлению")
    ws_fix.append(
        [
            "Файл",
            "Лист",
            "Строка",
            "Имя ДС",
            "Титул-система",
            "Код",
            "Ед. изм. (в файле)",
            "Ед. изм. (Google)",
            "Наименование",
            "Кол-во",
            "Ед. изм. (варианты)",
        ]
    )
    for item in fix_rows:
        record = item.record
        ws_fix.append(
            [
                record.file_name,
                record.sheet,
                record.excel_row,
                record.ds_name,
                item.title_system,
                item.code,
                item.file_units,
                item.google_units,
                record.name[:80],
                _qty_excel_value(record.values),
                " | ".join(item.units_variants),
            ]
        )
        file_units_cell = ws_fix.cell(row=ws_fix.max_row, column=7)
        file_units_cell.fill = _COLLISION_FILL_GOOGLE_UNITS_DIFF
        if item.google_units:
            google_cell = ws_fix.cell(row=ws_fix.max_row, column=8)
            google_cell.fill = _COLLISION_FILL_OK
    _finish_data_sheet(ws_fix, has_rows=bool(fix_rows))
    _apply_sheet_widths(ws_fix, _COLLISION_WIDTHS["К исправлению"])

    ws_keys = wb.create_sheet("Ключи")
    ws_keys.append(
        [
            "Титул-система",
            "Код",
            "Ед. изм. (варианты)",
            "Ед. изм. (Google)",
            "Строк к исправлению",
            "Файлов",
            "Наименование",
        ]
    )
    fix_count_by_key: Counter[tuple[str, str]] = Counter()
    files_by_key: dict[tuple[str, str], set[str]] = defaultdict(set)
    for item in fix_rows:
        key = (item.title_system, item.code)
        fix_count_by_key[key] += 1
        files_by_key[key].add(item.record.file_name)
    for key, example, units_variants in coarse.ambiguous_meta:
        google_units = _google_units_for_code(key[1], units_by_code)
        ws_keys.append(
            [
                key[0],
                key[1],
                " | ".join(units_variants),
                google_units,
                fix_count_by_key.get(key, 0),
                len(files_by_key.get(key, ())),
                example.name[:80],
            ]
        )
        if google_units:
            cell = ws_keys.cell(row=ws_keys.max_row, column=4)
            if _google_units_matches_variants(google_units, units_variants):
                cell.fill = _COLLISION_FILL_OK
            else:
                cell.fill = _COLLISION_FILL_GOOGLE_UNITS_DIFF
    _finish_data_sheet(ws_keys, has_rows=bool(coarse.ambiguous_meta))
    _apply_sheet_widths(ws_keys, _COLLISION_WIDTHS["Ключи"])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return len(fix_rows)


def _manifest_entries(file_stats: list[FileStats]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for stats in file_stats:
        if stats.kind == "summary":
            continue
        entries.append(
            {
                "kind": stats.kind,
                "file_id": stats.file_id,
                "file_name": stats.file_name,
                "ds_name": stats.ds_name,
                "selected_sheet": stats.sheet,
                "header_row": stats.header_row,
                "primary_signature": stats.primary_signature,
                "rows": stats.rows,
                "qty_sum": str(stats.qty_sum),
            }
        )
    return sorted(entries, key=lambda item: (item["kind"], item["file_id"], item["file_name"]))


def _write_manifest(path: Path, file_stats: list[FileStats]) -> None:
    payload = {
        "description": "Expected split RFP source files. Regenerate only after manual source review.",
        "entries": _manifest_entries(file_stats),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _validate_manifest(path: Path, file_stats: list[FileStats]) -> list[str]:
    if not path.exists():
        return [f"manifest file does not exist: {path}"]
    payload = json.loads(path.read_text(encoding="utf-8"))
    expected_entries = payload.get("entries")
    if not isinstance(expected_entries, list):
        return [f"manifest has no entries list: {path}"]

    expected_by_id = {entry.get("file_id"): entry for entry in expected_entries}
    actual_entries = _manifest_entries(file_stats)
    actual_by_id = {entry.get("file_id"): entry for entry in actual_entries}
    errors: list[str] = []

    duplicate_actual = [
        file_id
        for file_id, count in Counter(entry.get("file_id") for entry in actual_entries).items()
        if count > 1
    ]
    duplicate_expected = [
        file_id
        for file_id, count in Counter(entry.get("file_id") for entry in expected_entries).items()
        if count > 1
    ]
    for file_id in duplicate_expected:
        errors.append(f"duplicate file_id in manifest: {file_id}")
    for file_id in duplicate_actual:
        errors.append(f"duplicate file_id in current files: {file_id}")

    expected_ids = set(expected_by_id)
    actual_ids = set(actual_by_id)
    for file_id in sorted(expected_ids - actual_ids):
        errors.append(f"missing file_id: {file_id}")
    for file_id in sorted(actual_ids - expected_ids):
        errors.append(f"unexpected new file_id: {file_id}")

    checked_fields = (
        "kind",
        "selected_sheet",
        "header_row",
        "primary_signature",
    )
    for file_id in sorted(expected_ids & actual_ids):
        expected = expected_by_id[file_id]
        actual = actual_by_id[file_id]
        for field in checked_fields:
            if expected.get(field) != actual.get(field):
                errors.append(
                    f"{file_id}: {field} changed: expected {expected.get(field)!r}, actual {actual.get(field)!r}"
                )

    return errors


def _build_report(
    summary_records: list[RfpRecord],
    parts_records: list[RfpRecord],
    file_stats: list[FileStats],
    warnings: list[tuple[str, str, str]],
    manifest_errors: list[str] | None = None,
    diff_xlsx_path: Path | None = None,
    diff_xlsx_rows: int | None = None,
    parts_xlsx_path: Path | None = None,
    parts_xlsx_rows: int | None = None,
    diagnostics_xlsx_path: Path | None = None,
    diagnostics_xlsx_rows: int | None = None,
    checklist_compare: ChecklistCompareResult | None = None,
    aggregation: PartsAggregation | None = None,
    coarse: CoarseCollisionResult | None = None,
    net_xlsx_path: Path | None = None,
    net_xlsx_rows: int | None = None,
    collisions_xlsx_path: Path | None = None,
    collisions_xlsx_rows: int | None = None,
    duplicate_tags_xlsx_path: Path | None = None,
    duplicate_tags_xlsx_rows: int | None = None,
    empty_code_xlsx_path: Path | None = None,
    empty_code_xlsx_rows: int | None = None,
) -> str:
    lines: list[str] = []
    lines.append("Диагностический отчет по RFP из частей")
    lines.append("=" * 80)

    total_parts_qty = sum((record.values for record in parts_records), Decimal("0"))
    if summary_records:
        total_summary_qty = sum((record.values for record in summary_records), Decimal("0"))
        lines.append(
            f"summary records: {len(summary_records)}, qty sum: {total_summary_qty}"
        )
    lines.append(f"parts records: {len(parts_records)}, qty sum: {total_parts_qty}")

    if checklist_compare is not None:
        lines.append("")
        lines.append(format_checklist_report_section(checklist_compare))

    file_table = PrettyTable()
    file_table.field_names = [
        "kind",
        "file_id",
        "ds_name",
        "file",
        "selected_sheet",
        "header",
        "rows",
        "qty_sum",
        "warnings",
    ]
    for stats in file_stats:
        warn_count = sum(1 for _, file_name, _ in warnings if file_name == stats.file_name)
        file_table.add_row(
            [
                stats.kind,
                stats.file_id,
                stats.ds_name,
                stats.file_name,
                stats.sheet,
                stats.header_row or "",
                stats.rows,
                stats.qty_sum,
                warn_count,
            ]
        )
    _add_table(
        lines,
        "Файлы-источники и выбранные листы",
        file_table,
        "По каждому файлу показаны ID, имя ДС из имени файла, выбранный лист, строка шапки, "
        "количество строк и сумма VALUES. Warnings > 0 требует проверки.",
    )

    selected_sheet_counts = Counter((stats.kind, stats.sheet) for stats in file_stats if stats.sheet)
    selected_sheet_table = PrettyTable()
    selected_sheet_table.field_names = ["kind", "selected_sheet", "count"]
    for (kind, sheet), count in sorted(selected_sheet_counts.items()):
        selected_sheet_table.add_row([kind, sheet, count])
    _add_table(
        lines,
        "Сводка выбранных листов",
        selected_sheet_table,
        "Один рабочий лист читается с любым именем. Несколько рабочих листов: "
        "предпочитается 'Перечень материалов', остальные непрочитанные — WARN. "
        "Титульные/легенда/содержание не читаются.",
    )

    primary_signature_counts = Counter(
        (stats.kind, stats.primary_signature)
        for stats in file_stats
        if stats.kind in _PARTS_KINDS and stats.primary_signature
    )
    primary_signature_table = PrettyTable()
    primary_signature_table.field_names = ["kind", "count", "primary_signature"]
    for (kind, signature), count in sorted(primary_signature_counts.items()):
        primary_signature_table.add_row([kind, count, signature])
    _add_table(
        lines,
        "Порядок рабочих колонок",
        primary_signature_table,
        "Главная проверка шаблона: для файлов частей колонки 0..7 — блок РД, "
        "VALUES/UNITS — блок «Закупка по Лоту» (Excel №17–18, индексы 16–17). "
        "Отличие в любой книге попадает в ERROR ниже.",
    )

    if diff_xlsx_path is not None and diff_xlsx_rows is not None:
        try:
            display_diff_path = diff_xlsx_path.relative_to(ROOT)
        except ValueError:
            display_diff_path = diff_xlsx_path
        lines.append("")
        lines.append("Отличия сводной RFP от частей")
        lines.append(
            "Подробная таблица вынесена в отдельный XLSX: "
            f"{display_diff_path.as_posix()} (строк отличий: {diff_xlsx_rows})."
        )

    if parts_xlsx_path is not None:
        try:
            display_parts_path = parts_xlsx_path.relative_to(ROOT)
        except ValueError:
            display_parts_path = parts_xlsx_path
        lines.append("")
        lines.append("Общая таблица частей")
        lines.append(
            f"- parts: {display_parts_path.as_posix()} ({parts_xlsx_rows or 0} строк)"
        )

    if net_xlsx_path is not None:
        try:
            display_net = net_xlsx_path.relative_to(ROOT)
        except ValueError:
            display_net = net_xlsx_path
        lines.append("")
        lines.append("Свод (сумма частей)")
        lines.append(
            f"- свод: {display_net.as_posix()} ({net_xlsx_rows or 0} строк; "
            "фильтры коллизий в столбцах U–Z)"
        )
    if collisions_xlsx_path is not None:
        try:
            display_col = collisions_xlsx_path.relative_to(ROOT)
        except ValueError:
            display_col = collisions_xlsx_path
        lines.append(
            f"- collisions: {display_col.as_posix()} "
            f"({collisions_xlsx_rows or 0} строк к исправлению)"
        )

    if aggregation is not None:
        _add_aggregation_report(lines, aggregation)

    if coarse is None:
        coarse = _compute_coarse_collisions(parts_records)

    summary_table = PrettyTable()
    summary_table.field_names = ["Показатель", "Значение"]
    summary_table.add_row(["Всего агрегированных ключей", len(coarse.qty_by_key)])
    summary_table.add_row(["Ключей с разными ед. изм.", len(coarse.ambiguous_meta)])
    _add_table(
        lines,
        "Агрегированная сводка суммирования частей",
        summary_table,
        "Контроль по title_system + code. Коллизия — разные ед. изм. на одном ключе.",
    )

    if manifest_errors:
        manifest_table = PrettyTable()
        manifest_table.field_names = ["error"]
        for message in manifest_errors:
            manifest_table.add_row([message])
        _add_table(
            lines,
            "Ошибки проверки manifest",
            manifest_table,
            "Если таблица не пуста, список файлов или шаблон изменился относительно зафиксированных ID.",
        )

    if diagnostics_xlsx_path is not None and diagnostics_xlsx_rows is not None:
        try:
            display_diagnostics_path = diagnostics_xlsx_path.relative_to(ROOT)
        except ValueError:
            display_diagnostics_path = diagnostics_xlsx_path
        lines.append("")
        lines.append("XLSX диагностики")
        lines.append(
            "Таблица предупреждений дополнительно сохранена в XLSX: "
            f"{display_diagnostics_path.as_posix()} ({diagnostics_xlsx_rows} строк)."
        )

    if duplicate_tags_xlsx_path is not None and duplicate_tags_xlsx_rows:
        try:
            display_tags_path = duplicate_tags_xlsx_path.relative_to(ROOT)
        except ValueError:
            display_tags_path = duplicate_tags_xlsx_path
        lines.append("")
        lines.append("XLSX дублей тегов")
        lines.append(
            "Дубли тегов и несовпадение числа тегов с количеством лота: "
            f"{display_tags_path.as_posix()} ({duplicate_tags_xlsx_rows} строк). "
            "В таблице GUI эти замечания не показываются."
        )

    if empty_code_xlsx_path is not None and empty_code_xlsx_rows:
        try:
            display_empty = empty_code_xlsx_path.relative_to(ROOT)
        except ValueError:
            display_empty = empty_code_xlsx_path
        lines.append("")
        lines.append("XLSX позиций без кода")
        lines.append(
            "Строки без Кода РД (не входят в свод): "
            f"{display_empty.as_posix()} ({empty_code_xlsx_rows} строк)."
        )

    warn_table = PrettyTable()
    warn_table.field_names = ["Уровень", "Папка RFP", "Файл", "Строка", "Сообщение"]
    for row in _diagnostic_rows(warnings, file_stats, limit=300):
        warn_table.add_row(row)
    diagnostic_note = (
        "ERROR означает, что исходные файлы нельзя принимать в загрузку без исправления. "
        "WARN означает, что отчёт собран, но указанную строку или правило нужно проверить вручную.\n"
        "Легенда полей диагностики:\n"
        + "\n".join(f"- {item}" for item in DIAGNOSTICS_LEGEND)
    )
    _add_table(
        lines,
        f"Диагностика и предупреждения (показано {min(len(warnings), 300)} из {len(warnings)})",
        warn_table,
        diagnostic_note,
    )

    return "\n".join(lines) + "\n"


def run_rfp_parts_analyze(
    *,
    parts_dir: Path = DEFAULT_PARTS_DIR,
    out_dir: Path | None = None,
    summary_file: Path | None = None,
    checklist: Path | None = DEFAULT_CHECKLIST_FILE,
    no_checklist: bool = True,
    diff_xlsx: Path | None = None,
    parts_xlsx: Path | None = None,
    diagnostics_xlsx: Path | None = None,
    out: Path | None = None,
    write_manifest: Path | None = None,
    manifest: Path | None = None,
    strict: bool = False,
    units_matrix_path: Path | None = None,
    only_files: Iterable[Path] | None = None,
) -> Path:
    """Sum workbooks from the parts folder into ``rfp_parts_net.xlsx``.

    This is the same collection as ``python -m RFQ.rfp_parts`` / the GUI tab
    «RFP · Сбор частей». Also writes sibling ``rfp_parts_net_no_tags.xlsx``
    (one row per source line, TAGS blank) for ``load_tags=false``.

    Args:
        parts_dir: Folder ``RFP_Зиновьев``.
        out_dir: Stamp folder for artifacts. Created when omitted.
        summary_file: Optional legacy aggregated RFP for a diff workbook.
        checklist: DS checklist path; unused unless ``no_checklist`` is False.
        no_checklist: Skip checklist compare (default True — no DS registry yet).
        diff_xlsx: Optional summary-vs-parts diff path.
        parts_xlsx: Optional records workbook path.
        diagnostics_xlsx: Optional diagnostics workbook path.
        out: Text report path.
        write_manifest: Optional manifest output path.
        manifest: Optional manifest to validate against.
        strict: Exit 2 when the run has errors.
        units_matrix_path: Optional conversion matrix override. When omitted,
            ``load_config()['paths']['units_convert_matrix']`` is used.
        only_files: When set, run the same collection on these workbooks
            only. ``None`` reads every workbook directly in ``parts_dir``.

    Returns:
        Path to the written tagged ``rfp_parts_net.xlsx``.
    """
    from RFQ.rfp_parts.parts_net_preflight import write_parts_sources_snapshot

    timing = _PartsTimingSession()
    reports_dir = out_dir or make_reports_out_dir()
    reports_dir.mkdir(parents=True, exist_ok=True)
    report_path = out or (reports_dir / "rfp_parts_report.txt")
    diff_path = None
    if summary_file:
        diff_path = diff_xlsx or (reports_dir / "rfp_summary_increase_diff.xlsx")
    parts_xlsx_path = parts_xlsx or (reports_dir / "rfp_parts_records.xlsx")
    diagnostics_xlsx_path = diagnostics_xlsx or (
        reports_dir / "rfp_parts_diagnostics.xlsx"
    )
    net_xlsx = reports_dir / NET_XLSX_NAME
    net_no_tags_xlsx = reports_dir / NET_NO_TAGS_XLSX_NAME
    collisions_xlsx = reports_dir / COLLISIONS_XLSX_NAME
    print(f"reports dir: {reports_dir}")
    print(f"parts dir: {parts_dir}")

    try:
        warnings: list[tuple[str, str, str]] = []
        status_issues: list[RfpSupplyStatusIssue] = []
        file_stats: list[FileStats] = []
        summary_records: list[RfpRecord] = []
        all_parts: list[RfpRecord] = []
        skipped_empty_code: list[SkippedEmptyCodeRow] = []

        checklist_compare: ChecklistCompareResult | None = None
        if not no_checklist:
            with timing.phase("сверка checklist ДС"):
                checklist_compare = run_checklist_compare(
                    checklist_path=checklist,
                    parts_dir=parts_dir,
                    progress=True,
                )

        if summary_file:
            with timing.phase("извлечение сводной RFP"):
                records, stats = _extract_records(
                    summary_file,
                    "summary",
                    warnings,
                    status_issues=status_issues,
                    skipped_empty_code=skipped_empty_code,
                )
                summary_records.extend(records)
                file_stats.append(stats)

        with timing.phase("список файлов частей"):
            part_files = resolve_parts_workbooks(parts_dir, only_files)

        with timing.phase("извлечение всех частей"):
            extraction_started = time.perf_counter()
            part_total = len(part_files)
            for part_index, path in enumerate(part_files, start=1):
                file_started = time.perf_counter()
                records, stats = _extract_records(
                    path,
                    PARTS_KIND,
                    warnings,
                    status_issues=status_issues,
                    skipped_empty_code=skipped_empty_code,
                )
                file_elapsed = time.perf_counter() - file_started
                all_parts.extend(records)
                file_stats.append(stats)
                if file_elapsed >= PARTS_TIMING_SLOW_THRESHOLD_SEC:
                    timing.emit_slow_file(
                        part_index, part_total, path.name, file_elapsed
                    )
                if part_index % 10 == 0 or part_index == part_total:
                    timing.emit_progress(
                        part_index,
                        part_total,
                        time.perf_counter() - extraction_started,
                    )

        raise_if_rfp_supply_status_issues(status_issues)

        with timing.phase("загрузка Google base"):
            google_rows = _load_google_rows_strict()

        with timing.phase("построение Google index и lookups"):
            google_index = build_google_units_index(google_rows)
            google_lookup = _build_google_code_lookup(google_rows)

        matrix_path = _resolve_units_matrix_path(units_matrix_path)

        with timing.phase("построение запросов конвертации"):
            conversion_requests = _build_parts_conversion_requests(all_parts)

        try:
            with timing.phase("чтение матрицы и построение плана ЕИ"):
                conversion_plan = build_conversion_plan(
                    conversion_requests,
                    google_index,
                    matrix_path,
                )
            for path in write_fractional_conversion_logs(conversion_plan, reports_dir):
                print(
                    f"written fractional conversion xlsx: {path} "
                    f"({len(conversion_plan.fractional_issues)} rows)"
                )
            with timing.phase("применение плана конвертации и проверка целевых ЕИ"):
                all_parts = _apply_parts_conversion_plan(all_parts, conversion_plan)
                _assert_unified_target_units(all_parts)
        except UnitsConversionError as exc:
            warnings.append(("ERROR", "", f"фатальная ошибка конвертации ед. изм.: {exc}"))
            if strict:
                raise SystemExit(2) from exc
            raise

        with timing.phase("снимок зависимостей сборки"):
            build_deps_snapshot = build_parts_build_deps_snapshot(
                conversion_plan,
                google_index=google_index,
                matrix_path=matrix_path,
            )

        with timing.phase("агрегация единиц и количеств"):
            aggregation = _build_unit_counters(all_parts, warnings)

        with timing.phase("грубые коллизии ед. изм."):
            coarse = _compute_coarse_collisions(all_parts)

        with timing.phase("построение строк свода"):
            net_rows = _build_summary_rows(aggregation, coarse)
            no_tags_rows = _build_no_tags_summary_rows(all_parts, coarse)

        manifest_errors: list[str] = []
        if write_manifest:
            _write_manifest(write_manifest, file_stats)
        if manifest:
            manifest_errors = _validate_manifest(manifest, file_stats)

        diff_xlsx_rows: int | None = None
        if diff_path and summary_records:
            with timing.phase("запись diff сводной vs части"):
                diff_xlsx_rows = _write_summary_increase_diff_xlsx(
                    diff_path,
                    summary_records,
                    all_parts,
                )

        parts_xlsx_rows: int | None = None
        if parts_xlsx_path:
            with timing.phase("запись rfp_parts_records.xlsx"):
                parts_xlsx_rows = _write_records_xlsx(
                    parts_xlsx_path,
                    "Части",
                    all_parts,
                )

        diagnostics_xlsx_rows: int | None = None
        if diagnostics_xlsx_path:
            with timing.phase("запись rfp_parts_diagnostics.xlsx"):
                diagnostics_xlsx_rows = _write_diagnostics_xlsx(
                    diagnostics_xlsx_path,
                    warnings,
                    file_stats,
                )

        tag_remarks = _collect_tag_remarks(all_parts)
        tag_xlsx_path = reports_dir / DUPLICATE_TAGS_XLSX_NAME
        tag_xlsx_rows = 0
        if tag_remarks:
            with timing.phase(f"запись {DUPLICATE_TAGS_XLSX_NAME}"):
                tag_xlsx_rows = _write_duplicate_tags_xlsx(tag_xlsx_path, tag_remarks)

        empty_code_xlsx_path = reports_dir / EMPTY_CODE_XLSX_NAME
        empty_code_xlsx_rows = 0
        if skipped_empty_code:
            with timing.phase(f"запись {EMPTY_CODE_XLSX_NAME}"):
                empty_code_xlsx_rows = _write_empty_code_xlsx(
                    empty_code_xlsx_path, skipped_empty_code
                )

        with timing.phase("запись rfp_parts_net.xlsx"):
            net_xlsx_rows = _write_net_xlsx(
                net_xlsx, net_rows, equipment_by_code=google_lookup.equipment_by_code
            )

        with timing.phase("запись rfp_parts_net_no_tags.xlsx"):
            net_no_tags_xlsx_rows = _write_net_xlsx(
                net_no_tags_xlsx,
                no_tags_rows,
                equipment_by_code=google_lookup.equipment_by_code,
            )

        with timing.phase("запись rfp_parts_collisions.xlsx"):
            collisions_xlsx_rows = _write_collisions_xlsx(
                collisions_xlsx, coarse, units_by_code=google_lookup.units_by_code
            )

        file_status_payload: dict[str, Any] = {}
        with timing.phase("снимки sources и build deps"):
            write_parts_sources_snapshot(reports_dir, Path(parts_dir), part_files)
            write_parts_build_deps_snapshot(reports_dir, build_deps_snapshot)
            file_status_payload = build_file_status_payload(
                file_stats, warnings, tag_remarks=tag_xlsx_rows
            )
            write_file_status_json(reports_dir, file_status_payload)

        with timing.phase("формирование текстового отчёта"):
            report = _build_report(
                summary_records,
                all_parts,
                file_stats,
                warnings,
                manifest_errors,
                diff_path if diff_path and summary_records else None,
                diff_xlsx_rows,
                parts_xlsx_path if parts_xlsx_path else None,
                parts_xlsx_rows,
                diagnostics_xlsx_path if diagnostics_xlsx_path else None,
                diagnostics_xlsx_rows,
                checklist_compare,
                aggregation,
                coarse,
                net_xlsx,
                net_xlsx_rows,
                collisions_xlsx,
                collisions_xlsx_rows,
                tag_xlsx_path if tag_xlsx_rows else None,
                tag_xlsx_rows,
                empty_code_xlsx_path if empty_code_xlsx_rows else None,
                empty_code_xlsx_rows,
            )
            report_path.parent.mkdir(parents=True, exist_ok=True)
            report_path.write_text(report, encoding="utf-8")

        print(f"written report: {report_path}")
        if diff_path and diff_xlsx_rows is not None:
            print(f"written summary/parts diff xlsx: {diff_path} ({diff_xlsx_rows} rows)")
        if parts_xlsx_path and parts_xlsx_rows is not None:
            print(f"written parts xlsx: {parts_xlsx_path} ({parts_xlsx_rows} rows)")
        if diagnostics_xlsx_path and diagnostics_xlsx_rows is not None:
            print(
                f"written diagnostics xlsx: {diagnostics_xlsx_path} "
                f"({diagnostics_xlsx_rows} rows)"
            )
        if tag_xlsx_rows:
            print(
                f"written duplicate tags xlsx: {tag_xlsx_path} ({tag_xlsx_rows} rows)"
            )
        if empty_code_xlsx_rows:
            print(
                f"written empty-code xlsx: {empty_code_xlsx_path} "
                f"({empty_code_xlsx_rows} rows)"
            )
        print(f"written net xlsx: {net_xlsx} ({net_xlsx_rows} rows)")
        print(
            f"written net xlsx (no tags): {net_no_tags_xlsx} "
            f"({net_no_tags_xlsx_rows} rows)"
        )
        if file_status_payload:
            print(
                "written file status json: "
                f"ok={file_status_payload.get('ok', 0)}, "
                f"error={file_status_payload.get('error', 0)}, "
                f"tag_remarks={file_status_payload.get('tag_remarks', 0)}"
            )
        print(
            f"written collisions xlsx: {collisions_xlsx} "
            f"({collisions_xlsx_rows} rows to fix)"
        )
        if write_manifest:
            print(f"written manifest: {write_manifest}")
        print(f"parts records: {len(all_parts)}")
        if summary_records:
            print(f"summary records: {len(summary_records)}")
        print(f"warnings: {len(warnings)}")
        print(f"manifest_errors: {len(manifest_errors)}")
        print(
            f"collisions: units_mismatch_keys={len(coarse.ambiguous_meta)} "
            f"rows_to_fix={collisions_xlsx_rows}"
        )
        if checklist_compare is not None:
            print(
                "checklist_compare: "
                f"errors={checklist_compare.error_count}, "
                f"warns={checklist_compare.warn_count}, "
                f"ok={checklist_compare.ok_count}"
            )

        checklist_has_errors = bool(
            checklist_compare is not None
            and (checklist_compare.load_error or checklist_compare.error_count)
        )
        has_errors = (
            any(level == "ERROR" for level, _, _ in warnings)
            or bool(manifest_errors)
            or checklist_has_errors
        )
        timing.emit_total()
        if strict and has_errors:
            raise SystemExit(2)
        return net_xlsx
    finally:
        timing.write_artifact(reports_dir)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--summary-file",
        type=Path,
        default=None,
        help=(
            "Optional legacy Сводная RFP.xlsx for diff vs parts. "
            "Default: off (parts-only). Pass path to enable."
        ),
    )
    parser.add_argument("--parts-dir", type=Path, default=DEFAULT_PARTS_DIR)
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=None,
        help=(
            "Directory for report artifacts. "
            "Default: <RFP сводный файл>/YYYY.MM.DD_HH.MM"
        ),
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="Text report path (default: <out-dir>/rfp_parts_report.txt)",
    )
    parser.add_argument(
        "--diff-xlsx",
        type=Path,
        default=None,
        help=(
            "Summary-vs-parts diff xlsx (only with --summary-file). "
            "Default: <out-dir>/rfp_summary_increase_diff.xlsx"
        ),
    )
    parser.add_argument(
        "--parts-xlsx",
        type=Path,
        default=None,
        help="Default: <out-dir>/rfp_parts_records.xlsx",
    )
    parser.add_argument(
        "--diagnostics-xlsx",
        type=Path,
        default=None,
        help="Default: <out-dir>/rfp_parts_diagnostics.xlsx",
    )
    parser.add_argument("--write-manifest", type=Path, nargs="?", const=DEFAULT_MANIFEST, default=None)
    parser.add_argument("--manifest", type=Path, nargs="?", const=DEFAULT_MANIFEST, default=None)
    parser.add_argument("--strict", action="store_true")
    parser.add_argument(
        "--checklist",
        type=Path,
        default=DEFAULT_CHECKLIST_FILE,
        help="Maintained DS checklist xlsx (used only with --with-checklist)",
    )
    parser.add_argument(
        "--with-checklist",
        action="store_true",
        help="Opt-in DS registry compare (off by default while there is no registry)",
    )
    parser.add_argument(
        "--no-checklist",
        action="store_true",
        help="Deprecated: checklist compare is already off by default",
    )
    args = parser.parse_args()
    run_rfp_parts_analyze(
        parts_dir=args.parts_dir,
        out_dir=args.out_dir,
        summary_file=args.summary_file,
        checklist=args.checklist,
        no_checklist=not args.with_checklist,
        diff_xlsx=args.diff_xlsx,
        parts_xlsx=args.parts_xlsx,
        diagnostics_xlsx=args.diagnostics_xlsx,
        out=args.out,
        write_manifest=args.write_manifest,
        manifest=args.manifest,
        strict=args.strict,
    )



if __name__ == "__main__":
    main()
