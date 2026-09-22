"""Strict DS audit, quality gate and Step1-compatible baseline workbook.

Reads recursive DS xlsx/xlsm against an already loaded ``DsRegistryDocument``.
Reports are always written. ``Свод ДС для запуска.xlsx`` is written only when
there are zero blocking DS/registry issues. Output directory is required;
this module never defaults to a UNC path.
"""

from __future__ import annotations

import hashlib
import os
import re
import time
import warnings
from collections import defaultdict
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field, replace
from datetime import datetime
from decimal import Decimal
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Literal, Protocol, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import WINDOWS_EPOCH
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.worksheet.worksheet import Worksheet

warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
    category=UserWarning,
    module=r"openpyxl\..*",
)

from RFQ.ds_compare.ds_quantity_parse import try_parse_quantity
from RFQ.ds_compare.ds_units_normalize import normalize_units_text
from RFQ.rfp_parts.analyze_rfp_parts import NetSummaryRow, RfpRecord, _write_net_xlsx
from RFQ.rfp_parts.ds_progress import DsHeartbeat
from RFQ.rfp_parts.ds_registry import (
    MODE_FILTER,
    MODE_NEEDS_SPLIT,
    MODE_NO_UL,
    MODE_WHOLE,
    DsRegistryDocument,
    DsRegistryRelation,
    DsRegistryRow,
    canonical_supply_group_id,
    cell_text,
    detect_registry_format,
)
from RFQ.units_convert.fractional_log import write_fractional_conversion_logs
from RFQ.units_convert.gate import build_conversion_plan
from RFQ.units_convert.models import (
    STATUS_IDENTITY,
    ConversionInvariant,
    ConversionPlan,
    ConversionRequest,
    GoogleUnitsIndex,
    UnitsConversionError,
    converted_quantity,
    converted_unit,
    normalize_code,
    parse_decimal_quantity,
)

IssueLevel = Literal["ERROR", "WARN", "OVERLAY"]

ALGORITHM_VERSION = "ds_baseline_v2"
BASELINE_XLSX_NAME = "Свод ДС для запуска.xlsx"
STRUCTURE_REPORT_PREFIX = "Отчет по структуре файлов - ДС"
QUALITY_REPORT_PREFIX = "Отчет по качеству данных - ДС"
EMPTY_CODE_REPORT_PREFIX = "Отчет по позициям без кода - ДС"
DUPLICATE_TAGS_REPORT_PREFIX = "Отчет по дублям тегов - ДС"
TAG_MISMATCH_REPORT_PREFIX = "Отчет по несоответствию тегов - ДС"
STAMP_FORMAT = "%Y%m%d_%H%M%S"
CORE_WIDTH = 12
HEADER_SCAN_ROWS = 80
SCAN_MAX_COL = 64
MIN_CANDIDATE_ROLES = 5
REQUIRED_CANDIDATE_ROLES = frozenset({"title", "name", "code", "units", "qty"})

CORE_ROLES: tuple[str, ...] = (
    "npp",
    "title",
    "system",
    "specification",
    "rfq",
    "name",
    "code_1c",
    "code",
    "supplier",
    "type",
    "units",
    "qty",
)
ROLE_INDEX = {role: index for index, role in enumerate(CORE_ROLES)}
_ROLE_SHORT_RU: dict[str, str] = {
    "npp": "№ п/п",
    "title": "Титул",
    "system": "Раздел",
    "specification": "Спецификация",
    "rfq": "RFQ",
    "name": "Наименование",
    "code_1c": "Код 1С",
    "code": "Код РД",
    "supplier": "Поставщик",
    "type": "Техтребования",
    "units": "Ед. изм.",
    "qty": "Кол-во",
}
_LAYOUT_HINT_MAX_CHARS = 500
_LAYOUT_HINT_MAX_PARTS = 6
_QTY_SLOT_SAMPLE_ROWS = 40
_QTY_SLOT_UNIT_WORDS: tuple[str, ...] = (
    "шт",
    "м2",
    "м3",
    "кг",
    "компл",
    "упак",
    "м",
)

ISSUE_UNRESOLVED_ID = "unresolved_source_id"
ISSUE_AMBIGUOUS_ID = "ambiguous_source_id"
ISSUE_MULTI_DATA_SHEET = "multiple_data_sheets"
ISSUE_NO_HEADER = "no_ds_header"
ISSUE_INTERNAL_SHIFT = "internal_column_shift"
ISSUE_MISSING_ROLE = "missing_core_role"
ISSUE_NOT_ENOUGH_COLUMNS = "not_enough_columns"
ISSUE_QTY_FORMULA = "qty_formula"
ISSUE_QTY_EMPTY = "qty_empty"
ISSUE_QTY_NON_NUMERIC = "qty_non_numeric"
ISSUE_QTY_NON_FINITE = "qty_non_finite"
ISSUE_QTY_NEGATIVE = "qty_negative"
ISSUE_QTY_ZERO = "qty_zero"
ISSUE_EMPTY_CODE = "empty_procurement_code"
ISSUE_UNKNOWN_GOOGLE = "unknown_google_code"
ISSUE_EXACT_DUPLICATE = "exact_duplicate_row"
ISSUE_TAG_DUPLICATE = "tag_duplicate"
ISSUE_TAG_MISMATCH = "tag_qty_mismatch"
ISSUE_GROUP_FALLBACK = "group_fallback"
ISSUE_WORKBOOK = "workbook_read_error"
ISSUE_UNITS_CONVERT = "units_conversion_error"
ISSUE_REGISTRY = "registry_error"

_EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm"})
_WS_RE = re.compile(r"\s+")
_TAG_SPLIT_RE = re.compile(r"[;,\n\r]+")
_ID_STRIP_PREFIX_RE = re.compile(r"^(?:ДС|DS)[_ ]?", re.IGNORECASE)
_TAG_HEADER_RE = re.compile(r"(?:\btags?\b)|(?:^тег)|(?:тег[иа]?\b)|(?:линия.*tag)", re.IGNORECASE)

_HEADER_FONT = Font(bold=True)
_LINK_FONT = Font(color="0563C1", underline="single")
_FILL_HEADER = PatternFill(fill_type="solid", fgColor="D9E2F3")
_FILL_ERROR = PatternFill(fill_type="solid", fgColor="FFC7CE")
_FILL_WARN = PatternFill(fill_type="solid", fgColor="FFF2CC")
_ALIGN_WRAP = Alignment(wrap_text=True, vertical="top")

_BOILERPLATE_NEEDLES = (
    "общая сумма",
    "условия поставки",
    "грузоотправитель",
    "грузополучатель",
    "изготовитель",
    "стороны согласовали",
    "поставщик подтверждает",
    "во всем остальном",
    "в соответствии с",
    "со стороны поставщика",
    "со стороны покупателя",
)

_HEAD_REPEAT_NEEDLES = (
    "наименование позиций товара по рд",
    "код 1с соу",
    "технические требования",
    "ед. изм.",
)
_FOOTER_NEEDLES = (
    "итого",
    "всего",
    "на сумму",
    "общая сумма",
    "передан через диадок",
    "банковские реквизиты",
    "инн ",
    "кпп ",
    "р/счет",
    "р/счёт",
    "бик:",
)

# More specific roles first so «Наименование … Поставщика» is not NAME.
_ROLE_PATTERNS: tuple[tuple[str, tuple[re.Pattern[str], ...]], ...] = (
    (
        "code_1c",
        (
            re.compile(r"код\s*1\s*с", re.IGNORECASE),
            re.compile(r"код\s*1c", re.IGNORECASE),
        ),
    ),
    (
        "supplier",
        (
            re.compile(r"наименование.*поставщик", re.IGNORECASE),
            re.compile(r"^поставщик$", re.IGNORECASE),
        ),
    ),
    (
        "name",
        (
            re.compile(r"наименование\s*позиций\s*товара\s*по\s*рд", re.IGNORECASE),
            re.compile(r"наименование\s*мтр", re.IGNORECASE),
            re.compile(r"^наименование$", re.IGNORECASE),
        ),
    ),
    (
        "specification",
        (re.compile(r"спецификац", re.IGNORECASE),),
    ),
    (
        "code",
        (
            re.compile(r"код\s*рд", re.IGNORECASE),
            re.compile(r"\bbcc\b", re.IGNORECASE),
        ),
    ),
    (
        "type",
        (
            re.compile(r"техническ.*требован", re.IGNORECASE),
            re.compile(r"техническ.*характерист", re.IGNORECASE),
            re.compile(r"гост\s*/\s*ту", re.IGNORECASE),
            re.compile(r"тип.*марк", re.IGNORECASE),
        ),
    ),
    (
        "units",
        (
            re.compile(r"ед\.?\s*изм", re.IGNORECASE),
            re.compile(r"единиц.*измер", re.IGNORECASE),
        ),
    ),
    (
        "qty",
        (
            re.compile(r"^кол-?во$", re.IGNORECASE),
            re.compile(r"^количество$", re.IGNORECASE),
        ),
    ),
    (
        "title",
        (re.compile(r"^титул\b", re.IGNORECASE),),
    ),
    (
        "system",
        (
            re.compile(r"^раздел\b", re.IGNORECASE),
            re.compile(r"^марка\b", re.IGNORECASE),
            re.compile(r"^система\b", re.IGNORECASE),
        ),
    ),
    (
        "rfq",
        (
            re.compile(r"^rfq$", re.IGNORECASE),
            re.compile(r"номер\s*rfq", re.IGNORECASE),
            re.compile(r"^рфк$", re.IGNORECASE),
        ),
    ),
    (
        "npp",
        (
            re.compile(r"№\s*п/?п", re.IGNORECASE),
            re.compile(r"^npp$", re.IGNORECASE),
            re.compile(r"порядк", re.IGNORECASE),
            re.compile(r"^№$", re.IGNORECASE),
        ),
    ),
)

_SKIP_NAME_PREFIXES = (
    "реестр_дс",
    "реестр дс",
    "свод_дс",
    "свод дс",
    "отчет по структуре файлов - дс_",
    "отчет по качеству данных - дс_",
    "отчет по позициям без кода - дс_",
    "отчет по дублям тегов - дс_",
    "отчет по несоответствию тегов - дс_",
)
_SKIP_EXACT_NAMES = frozenset(
    {
        BASELINE_XLSX_NAME.casefold(),
        "дробные значения после конвертации дс.xlsx",
    }
)


# ---------------------------------------------------------------------------
# Public dataclasses
# ---------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class DsBaselineIssue:
    """One audit / quality / grouping finding."""

    code: str
    level: IssueLevel
    message: str
    blocking: bool = False
    path: Path | None = None
    relpath: str = ""
    sheet: str = ""
    excel_row: int | None = None
    source_id: str = ""
    group_id: str = ""


@dataclass(frozen=True, slots=True)
class DsSourceFile:
    """One collected DS workbook with a content fingerprint."""

    path: Path
    relpath: str
    size: int
    mtime_ns: int
    fingerprint: str


@dataclass(frozen=True, slots=True)
class DsSkippedFile:
    """Workbook under the source root that is not a DS specification."""

    path: Path
    relpath: str
    reason: str


@dataclass(frozen=True, slots=True)
class DsSourceIdResolution:
    """Result of matching a file path to an active registry source ID."""

    source_id: str | None
    method: str
    candidates: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class DsSheetScan:
    """Per-sheet DS-header scan used for the structure report."""

    name: str
    is_active: bool
    is_candidate: bool
    header_row: int | None
    score: int
    roles: tuple[str, ...]
    leading_empty: int


@dataclass(frozen=True, slots=True)
class DsBaselinePosition:
    """One DS position with provenance, grouping and conversion fields."""

    source_id: str
    group_id: str
    group_label: str
    grouping_mode: str
    grouping_fallback: bool
    overlay_blocked: bool
    relpath: str
    file_name: str
    path: Path
    sheet: str
    excel_row: int
    ds_number: str
    title: str
    system: str
    specification: str
    rfq: str
    name: str
    code_1c: str
    code: str
    code_normalized: str
    supplier_name: str
    type_mark: str
    units: str
    qty: Decimal | None
    qty_raw: str
    qty_source: Decimal | None
    unit_source: str
    qty_target: Decimal | None = None
    unit_target: str = ""
    conversion_status: str = ""
    conversion_trace: str = ""
    conversion_coefficient: Decimal | None = None
    diagnostic_tags: tuple[str, ...] = ()
    leading_empty: int = 0


@dataclass(frozen=True, slots=True)
class DsSupplyGroupInfo:
    """Grouping metadata consumed by the later RFP overlay stage."""

    group_id: str
    group_label: str
    source_ids: tuple[str, ...]
    modes: tuple[str, ...]
    fallback: bool
    overlay_blocked: bool
    position_count: int


@dataclass(frozen=True, slots=True)
class DsConversionBatch:
    """Converter output: positions with qty/unit/status filled, optional plan."""

    converted: tuple[DsBaselinePosition, ...]
    plan: ConversionPlan | None = None
    warnings: tuple[str, ...] = ()


@dataclass
class DsBaselineResult:
    """Public result of ``build_ds_baseline``."""

    source_root: Path
    output_dir: Path
    registry_path: Path
    stamp: str
    fingerprint: str
    files: list[DsSourceFile] = field(default_factory=list)
    skipped: list[DsSkippedFile] = field(default_factory=list)
    positions: list[DsBaselinePosition] = field(default_factory=list)
    groups: list[DsSupplyGroupInfo] = field(default_factory=list)
    issues: list[DsBaselineIssue] = field(default_factory=list)
    file_count: int = 0
    position_count: int = 0
    group_count: int = 0
    blocking_issue_count: int = 0
    warn_count: int = 0
    overlay_count: int = 0
    blocking: bool = False
    baseline_path: Path | None = None
    structure_report_path: Path | None = None
    quality_report_path: Path | None = None
    empty_code_report_path: Path | None = None
    duplicate_tags_report_path: Path | None = None
    tag_mismatch_report_path: Path | None = None
    fractional_log_paths: tuple[Path, ...] = ()

    def summary_line(self) -> str:
        if not self.blocking:
            return (
                f"OK: ДС baseline готов, файлов={self.file_count}, "
                f"позиций={self.position_count}, WARN={self.warn_count}"
            )
        return (
            f"BLOCKED: ERROR={self.blocking_issue_count}, "
            f"WARN={self.warn_count}, OVERLAY={self.overlay_count}, "
            f"файлов={self.file_count}, позиций={self.position_count}"
        )


class DsUnitsConverter(Protocol):
    """Injectable units conversion. Default talks to RFQ.units_convert APIs."""

    def convert_positions(
        self, positions: Sequence[DsBaselinePosition]
    ) -> DsConversionBatch:
        """Return converted copies; must not fake a matrix/Google result."""


# ---------------------------------------------------------------------------
# Path / text helpers
# ---------------------------------------------------------------------------


def _cell_text(value: object) -> str:
    return cell_text(value)


def _is_empty(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _norm_header(text: object) -> str:
    raw = _cell_text(text).replace("\n", " ").replace("\r", " ")
    return _WS_RE.sub(" ", raw).strip()


def _norm_filter(value: object) -> str:
    return _WS_RE.sub(" ", _cell_text(value)).casefold()


def _leading_empty_count(values: Sequence[object]) -> int:
    count = 0
    for value in values:
        if not _is_empty(value):
            break
        count += 1
    return count


def _row_used(values: Sequence[object]) -> bool:
    return any(not _is_empty(value) for value in values)


def _posix_relpath(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        return path.name


def _file_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@lru_cache(maxsize=4096)
def _path_uri(path_key: str) -> str | None:
    """Build a ``file:`` URI without ``Path.resolve()``.

    ``resolve()`` hits the network for every UNC path. Quality reports used
    to call it twice per position row, which looks like a hang after the
    workbooks are already parsed. Excel also caps a sheet at ~65 530
    hyperlinks — same trap as Step4 RFP/MTO path columns.
    """
    try:
        as_path = Path(path_key)
        if as_path.is_absolute():
            return as_path.as_uri()
        return as_path.resolve(strict=False).as_uri()
    except (OSError, ValueError):
        return None


def _set_path_cell(cell, path: Path | str | None) -> None:
    """Put a clickable file link on a *file-list* cell, not on every position."""
    if path is None or path == "":
        return
    if cell.value is None or cell.value == "":
        cell.value = str(path)
    cell.alignment = _ALIGN_WRAP
    uri = _path_uri(str(path))
    if uri:
        cell.hyperlink = uri
        cell.font = _LINK_FONT


def _atomic_write_bytes(target: Path, data: bytes) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = target.with_name(f".{target.stem}.{os.getpid()}.tmp{target.suffix}")
    try:
        tmp_path.write_bytes(data)
        os.replace(tmp_path, target)
    except OSError:
        tmp_path.unlink(missing_ok=True)
        raise
    return target


def _save_workbook_atomic(
    target: Path,
    workbook: Workbook,
    progress: Callable[[str], None] | None = None,
) -> Path:
    if progress is not None:
        progress("сериализация xlsx")
    buffer = BytesIO()
    try:
        workbook.save(buffer)
    finally:
        workbook.close()
    data = buffer.getvalue()
    if progress is not None:
        progress(f"запись на диск ({max(len(data) // 1024, 1)} КБ)")
    return _atomic_write_bytes(target, data)


def _issue(
    code: str,
    level: IssueLevel,
    message: str,
    *,
    blocking: bool | None = None,
    path: Path | None = None,
    relpath: str = "",
    sheet: str = "",
    excel_row: int | None = None,
    source_id: str = "",
    group_id: str = "",
) -> DsBaselineIssue:
    is_blocking = (level == "ERROR") if blocking is None else blocking
    return DsBaselineIssue(
        code=code,
        level=level,
        message=message,
        blocking=is_blocking,
        path=path,
        relpath=relpath,
        sheet=sheet,
        excel_row=excel_row,
        source_id=source_id,
        group_id=group_id,
    )


# ---------------------------------------------------------------------------
# Collect
# ---------------------------------------------------------------------------


def _skip_reason(
    path: Path,
    *,
    registry_path: Path | None,
) -> str | None:
    name = path.name
    if name.startswith("~$"):
        return "excel_lock"
    folded = name.casefold()
    stem = path.stem.casefold()
    if folded in _SKIP_EXACT_NAMES:
        return "own_output"
    for prefix in _SKIP_NAME_PREFIXES:
        if stem.startswith(prefix) or folded.startswith(prefix):
            return "own_or_manual_or_registry_name"
    if registry_path is not None:
        try:
            if path.resolve() == Path(registry_path).resolve():
                return "canonical_registry"
        except OSError:
            pass
    return None


def collect_ds_workbooks(
    source_root: str | Path,
    *,
    registry_path: str | Path | None = None,
) -> tuple[list[DsSourceFile], list[DsSkippedFile]]:
    """Recursively collect DS xlsx/xlsm, skipping locks, registry and own outputs.

    Args:
        source_root: Folder to walk.
        registry_path: Canonical registry path; that file is skipped if inside
            the tree.

    Returns:
        ``(files, skipped)`` where ``files`` is sorted by posix relpath.
    """

    root = Path(source_root)
    registry = Path(registry_path) if registry_path else None
    collected: list[DsSourceFile] = []
    skipped: list[DsSkippedFile] = []
    if not root.is_dir():
        return collected, skipped

    candidates: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in _EXCEL_SUFFIXES:
            continue
        candidates.append(path)

    for path in candidates:
        relpath = _posix_relpath(path, root)
        reason = _skip_reason(path, registry_path=registry)
        if reason:
            skipped.append(DsSkippedFile(path=path, relpath=relpath, reason=reason))
            continue
        try:
            fmt = detect_registry_format(path)
        except Exception:
            fmt = "unknown"
        if fmt in {"new", "legacy"}:
            skipped.append(
                DsSkippedFile(path=path, relpath=relpath, reason=f"registry_{fmt}")
            )
            continue
        try:
            data = path.read_bytes()
            stat = path.stat()
        except OSError:
            skipped.append(DsSkippedFile(path=path, relpath=relpath, reason="unreadable"))
            continue
        collected.append(
            DsSourceFile(
                path=path,
                relpath=relpath,
                size=stat.st_size,
                mtime_ns=stat.st_mtime_ns,
                fingerprint=_file_sha256(data),
            )
        )
    collected.sort(key=lambda item: item.relpath)
    return collected, skipped


# ---------------------------------------------------------------------------
# Source ID
# ---------------------------------------------------------------------------


def _active_ids(registry: DsRegistryDocument) -> tuple[str, ...]:
    ids = [row.source_id for row in registry.active_rows if row.source_id]
    return tuple(sorted(set(ids), key=lambda item: (-len(item), item)))


def _folder_exact_id(folder_name: str, active_ids: Sequence[str]) -> str | None:
    folded = folder_name.strip().casefold()
    for source_id in active_ids:
        variants = (
            f"дс_{source_id}".casefold(),
            f"дс{source_id}".casefold(),
        )
        if folded in variants:
            return source_id
    return None


def _id_boundary_ok(remainder: str) -> bool:
    """True when ``remainder`` does not continue the matched source ID.

    ``4905_1`` must not collapse to ``4905`` (``_`` + digits only). An
    actual-like sequential/revision tail ``_24Б`` is a valid boundary.
    """

    if not remainder:
        return True
    if remainder[0].isalnum():
        return False
    if remainder.startswith("_") and len(remainder) > 1 and remainder[1].isdigit():
        index = 1
        while index < len(remainder) and remainder[index].isdigit():
            index += 1
        if index < len(remainder) and remainder[index].isalpha():
            return True
        return False
    return True



def _filename_prefix_id(file_name: str, active_ids: Sequence[str]) -> str | None:
    stem = Path(file_name).stem.strip()
    stripped = _ID_STRIP_PREFIX_RE.sub("", stem, count=1).strip()
    variants = (stem, stripped)
    ordered = sorted(set(active_ids), key=lambda item: (-len(item), item))
    for source_id in ordered:
        for text in variants:
            if text == source_id:
                return source_id
            if text.startswith(source_id) and _id_boundary_ok(text[len(source_id) :]):
                return source_id
    return None


def _token_hits(text: str, source_id: str) -> bool:
    escaped = re.escape(source_id)
    pattern = re.compile(
        rf"(?:^|[/\\]|[^0-9A-Za-zА-Яа-я])(?:ДС[_\s]?)?{escaped}(?![0-9A-Za-zА-Яа-я])(?!_\d)",
        re.IGNORECASE,
    )
    return pattern.search(text) is not None


def resolve_ds_source_id(
    relpath: str,
    active_ids: Sequence[str],
) -> DsSourceIdResolution:
    """Resolve a DS file to one active registry source ID.

    Order: nearest ancestor exact ``ДС_<id>`` / ``ДС{id}``, then filename
    prefix longest exact ID (authoritative even if another ``ДС`` token
    appears later in the name), then a unique token in the relative path.
    Ambiguous only when ancestor and prefix both miss and several tokens
    remain. ``4905_1`` never collapses to ``4905``.
    """

    if not active_ids:
        return DsSourceIdResolution(None, "unresolved", ())

    posix = relpath.replace("\\", "/")
    parts = posix.split("/")
    file_name = parts[-1] if parts else posix
    ancestors = list(reversed(parts[:-1]))
    for folder in ancestors:
        hit = _folder_exact_id(folder, active_ids)
        if hit is not None:
            return DsSourceIdResolution(hit, "ancestor", (hit,))

    prefix = _filename_prefix_id(file_name, active_ids)
    if prefix is not None:
        return DsSourceIdResolution(prefix, "filename_prefix", (prefix,))

    hits: list[str] = []
    for source_id in active_ids:
        if _token_hits(posix, source_id) and source_id not in hits:
            hits.append(source_id)
    if len(hits) == 1:
        return DsSourceIdResolution(hits[0], "unique_token", tuple(hits))
    if len(hits) > 1:
        return DsSourceIdResolution(None, "ambiguous", tuple(hits))
    return DsSourceIdResolution(None, "unresolved", ())



# ---------------------------------------------------------------------------
# Header / schema
# ---------------------------------------------------------------------------


def classify_header_role(text: object) -> str | None:
    """Return the core DS role for a header cell, or None."""

    label = _norm_header(text)
    if not label:
        return None
    for role, patterns in _ROLE_PATTERNS:
        if any(pattern.search(label) for pattern in patterns):
            return role
    return None


def _is_tag_header(text: object) -> bool:
    """True for a dedicated Tag column, not a core role that merely mentions TAG."""

    label = _norm_header(text)
    if not label:
        return False
    if classify_header_role(label) is not None:
        return False
    return bool(_TAG_HEADER_RE.search(label))


def _role_map(values: Sequence[object]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, value in enumerate(values):
        role = classify_header_role(value)
        if role and role not in found:
            found[role] = index
    return found


def _header_score(values: Sequence[object]) -> tuple[int, tuple[str, ...]]:
    roles = tuple(sorted(_role_map(values)))
    return len(roles), roles


def _is_candidate_header(values: Sequence[object]) -> bool:
    mapping = _role_map(values)
    roles = set(mapping)
    if not REQUIRED_CANDIDATE_ROLES.issubset(roles):
        return False
    return len(roles) >= MIN_CANDIDATE_ROLES


def _is_boilerplate_shifted(shifted: Sequence[object]) -> bool:
    blob = " ".join(_cell_text(value).casefold() for value in shifted[:CORE_WIDTH])
    return any(needle in blob for needle in _BOILERPLATE_NEEDLES)


def _is_repeated_header(shifted: Sequence[object]) -> bool:
    blob = " ".join(_cell_text(value).casefold() for value in shifted[:CORE_WIDTH])
    if classify_header_role(shifted[ROLE_INDEX["title"]] if len(shifted) > 1 else "") == "title":
        code_role = classify_header_role(
            shifted[ROLE_INDEX["code"]] if len(shifted) > ROLE_INDEX["code"] else ""
        )
        if code_role == "code":
            return True
    return any(needle in blob for needle in _HEAD_REPEAT_NEEDLES)


def _units_party_label(shifted: Sequence[object]) -> bool:
    units = _cell_text(shifted[ROLE_INDEX["units"]] if len(shifted) > ROLE_INDEX["units"] else "")
    return units.strip() in {"ПОКУПАТЕЛЬ", "ПОСТАВЩИК", "Покупатель", "Поставщик"}


def _core_text_blob(shifted: Sequence[object]) -> str:
    parts = [
        _cell_text(shifted[index] if index < len(shifted) else "")
        for index in range(CORE_WIDTH)
    ]
    return " ".join(parts).casefold()


def _has_footer_needle(text: object) -> bool:
    folded = _cell_text(text).casefold()
    if not folded:
        return False
    return any(needle in folded for needle in _FOOTER_NEEDLES)


def _qty_formula_is_aggregate(text: object) -> bool:
    folded = _cell_text(text).strip().casefold().replace(" ", "")
    return folded.startswith("=sum") or folded.startswith("=subtotal")


def _all_core_ref(shifted: Sequence[object]) -> bool:
    if CORE_WIDTH <= 0:
        return False
    for index in range(CORE_WIDTH):
        value = shifted[index] if index < len(shifted) else None
        if _cell_text(value).strip().casefold() != "#ref!":
            return False
    return True


def _unit_words_in_text(text: object) -> tuple[str, ...]:
    folded = normalize_units_text(_cell_text(text)).casefold()
    if not folded:
        return ()
    found: list[str] = []
    for word in _QTY_SLOT_UNIT_WORDS:
        if word == "м":
            if re.search(r"(?<![0-9a-zа-яё])м(?![0-9a-zа-яё23])", folded):
                found.append(word)
        elif word in folded:
            found.append(word)
    return tuple(found)


# ---------------------------------------------------------------------------
# Workbook load
# ---------------------------------------------------------------------------


def _open_pair(
    data: bytes,
) -> tuple[OpenpyxlWorkbook, OpenpyxlWorkbook, BytesIO, BytesIO]:
    formula_buf = BytesIO(data)
    try:
        formula = load_workbook(
            formula_buf, read_only=True, data_only=False, rich_text=False
        )
    except Exception:
        formula_buf.close()
        raise
    values_buf = BytesIO(data)
    try:
        values = load_workbook(
            values_buf, read_only=True, data_only=True, rich_text=False
        )
    except Exception:
        formula.close()
        formula_buf.close()
        values_buf.close()
        raise
    return formula, values, formula_buf, values_buf


def _close_pair(
    wb_f: object | None,
    wb_v: object | None,
    buf_f: BytesIO | None = None,
    buf_v: BytesIO | None = None,
) -> None:
    for item in (wb_f, wb_v, buf_f, buf_v):
        if item is None:
            continue
        closer = getattr(item, "close", None)
        if closer is None:
            continue
        try:
            closer()
        except Exception:  # noqa: BLE001 - best-effort resource release
            pass


def _is_tabular_sheet(ws: object) -> bool:
    if isinstance(ws, Worksheet):
        return True
    return hasattr(ws, "iter_rows") and hasattr(ws, "_get_source")


def _prepare_read_only_sheet(ws: object) -> None:
    """Drop stored dimensions so inflated max_row/max_column cannot pad gaps."""

    reset = getattr(ws, "reset_dimensions", None)
    if callable(reset):
        reset()


def _iter_regular_raw_rows(
    ws: object,
) -> Iterator[tuple[int, list[dict[str, object]]]]:
    for excel_row, cells in enumerate(ws.iter_rows(), start=1):
        payload: list[dict[str, object]] = []
        for index, cell in enumerate(cells, start=1):
            payload.append(
                {
                    "column": index,
                    "value": getattr(cell, "value", None),
                    "data_type": getattr(cell, "data_type", None),
                }
            )
        yield excel_row, payload


def _iter_xml_raw_rows(
    ws: object,
) -> Iterator[tuple[int, list[dict[str, object]]]]:
    get_source = getattr(ws, "_get_source", None)
    if not callable(get_source):
        yield from _iter_regular_raw_rows(ws)
        return
    src = get_source()
    try:
        parent = getattr(ws, "parent", None)
        parser = WorkSheetParser(
            src,
            getattr(ws, "_shared_strings", []),
            data_only=bool(getattr(parent, "data_only", False)),
            epoch=getattr(parent, "epoch", WINDOWS_EPOCH),
            date_formats=getattr(parent, "_date_formats", set()),
        )
        for excel_row, cells in parser.parse():
            yield int(excel_row), list(cells)
    finally:
        src.close()


def _materialize_xml_row(
    cells: Sequence[dict[str, object]],
    max_col: int | None,
) -> tuple[list[object], set[int]]:
    selected: list[tuple[int, object, bool]] = []
    last = 0
    for cell in cells:
        column = int(cell.get("column") or 0)
        if column < 1:
            continue
        if max_col is not None and column > max_col:
            continue
        value = cell.get("value")
        data_type = cell.get("data_type")
        is_formula = data_type == "f" or (
            isinstance(value, str) and value.startswith("=")
        )
        selected.append((column - 1, value, is_formula))
        if not _is_empty(value) or is_formula:
            last = max(last, column)
    if last <= 0:
        return [], set()
    values: list[object] = [None] * last
    formula_cols: set[int] = set()
    for col0, value, is_formula in selected:
        if col0 >= last:
            continue
        values[col0] = value
        if is_formula:
            formula_cols.add(col0)
    return values, formula_cols


def _iter_xml_sheet_rows(
    ws: object,
    *,
    max_col: int | None = None,
    max_row: int | None = None,
) -> Iterator[tuple[int, list[object], set[int]]]:
    """Yield XML-present rows only; do not fill dimension gaps with blank lists."""

    for excel_row, cells in _iter_xml_raw_rows(ws):
        if max_row is not None and excel_row > max_row:
            break
        values, formula_cols = _materialize_xml_row(cells, max_col)
        yield excel_row, values, formula_cols


def _read_prefix_rows(
    ws: object, *, max_row: int, max_col: int
) -> list[tuple[int, list[object]]]:
    rows: list[tuple[int, list[object]]] = []
    for excel_row, values, _cols in _iter_xml_sheet_rows(
        ws, max_col=max_col, max_row=max_row
    ):
        rows.append((excel_row, values))
    return rows


def _rstrip_empty(values: Sequence[object]) -> list[object]:
    end = len(values)
    while end > 0 and _is_empty(values[end - 1]):
        end -= 1
    return list(values[:end])


def _read_max_col(offset: int, tag_columns: Sequence[int]) -> int:
    needed = offset + CORE_WIDTH
    for index in tag_columns:
        needed = max(needed, index + 1)
    return max(needed, SCAN_MAX_COL)


def _iter_paired_sheet_rows(
    ws_f: object,
    ws_v: object | None,
    *,
    max_col: int,
) -> Iterator[tuple[int, list[object], list[object], set[int]]]:
    value_iter = (
        _iter_xml_sheet_rows(ws_v, max_col=max_col)
        if ws_v is not None
        else iter(())
    )
    pending = next(value_iter, None)
    for excel_row, formula_values, formula_cols in _iter_xml_sheet_rows(
        ws_f, max_col=max_col
    ):
        while pending is not None and pending[0] < excel_row:
            pending = next(value_iter, None)
        if pending is not None and pending[0] == excel_row:
            value_values = pending[1]
            pending = next(value_iter, None)
        else:
            value_values = formula_values
        yield excel_row, formula_values, value_values, formula_cols


def _used_offset(rows: Sequence[tuple[int, list[object]]]) -> int:
    counts = [
        _leading_empty_count(values)
        for _excel_row, values in rows
        if _row_used(values)
    ]
    if not counts:
        return 0
    return min(counts)


def _scan_sheet(
    name: str,
    *,
    is_active: bool,
    rows: Sequence[tuple[int, list[object]]],
) -> DsSheetScan:
    best_row: int | None = None
    best_score = 0
    best_roles: tuple[str, ...] = ()
    candidate = False
    header_values: list[object] = []
    for excel_row, values in rows[:HEADER_SCAN_ROWS]:
        if not _row_used(values):
            continue
        score, roles = _header_score(values)
        if _is_candidate_header(values) and score >= best_score:
            candidate = True
            best_score = score
            best_row = excel_row
            best_roles = roles
            header_values = list(values)
        elif score > best_score:
            best_score = score
            best_row = excel_row
            best_roles = roles
    if candidate and header_values:
        offset = _leading_empty_count(header_values)
    else:
        offset = _used_offset(rows)
    return DsSheetScan(
        name=name,
        is_active=is_active,
        is_candidate=candidate,
        header_row=best_row if candidate else None,
        score=best_score,
        roles=best_roles,
        leading_empty=offset,
    )


def _pad(values: Sequence[object], width: int) -> list[object]:
    padded = list(values)
    if len(padded) < width:
        padded.extend([None] * (width - len(padded)))
    return padded


def _composite_title(title: str, system: str) -> str:
    title = title.strip()
    system = system.strip()
    if title and system:
        return f"{title}-{system}"
    return title or system


def _parse_qty(
    raw: object,
) -> tuple[Decimal | None, str | None]:
    """Return (qty, error_code). error_code is None on success including zero."""

    if _is_empty(raw):
        return None, ISSUE_QTY_EMPTY
    ok, parsed = try_parse_quantity(raw)
    if not ok:
        text = _cell_text(raw)
        if text.startswith("="):
            return None, ISSUE_QTY_FORMULA
        return None, ISSUE_QTY_NON_NUMERIC
    try:
        qty = parse_decimal_quantity(raw)
    except UnitsConversionError:
        qty = Decimal(str(parsed))
    if qty.is_nan() or qty.is_infinite():
        return None, ISSUE_QTY_NON_FINITE
    if qty < 0:
        return qty, ISSUE_QTY_NEGATIVE
    return qty, None


def _fallback_group_id(source_id: str) -> str:
    return f"ДС{source_id}" if source_id else ""


def _assign_group(
    *,
    source_id: str,
    title: str,
    system: str,
    registry_row: DsRegistryRow | None,
) -> tuple[str, str, str, bool, bool, str | None]:
    """Return group_id, label, mode, fallback, overlay_blocked, issue message.

    The bag is the actual DS number. Extra UL folders and RFP files of that
    number stay in the same bag. Title and mark do not split it.
    """

    del title, system
    if registry_row is None or not source_id:
        label = _fallback_group_id(source_id)
        return label, label, "", True, True, "нет строки реестра для ID"
    group_id = canonical_supply_group_id(source_id)
    return group_id, group_id, "", False, False, None


def _row_matches_filter(
    title: str, system: str, rel: DsRegistryRelation
) -> bool:
    return (
        _norm_filter(title) == _norm_filter(rel.title_filter)
        and _norm_filter(system) == _norm_filter(rel.mark_filter)
    )


def _diagnostic_tags(values: Sequence[object], tag_columns: Sequence[int]) -> tuple[str, ...]:
    tags: list[str] = []
    for index in tag_columns:
        if index >= len(values):
            continue
        raw = _cell_text(values[index])
        if not raw:
            continue
        for part in _TAG_SPLIT_RE.split(raw):
            token = part.strip()
            if token:
                tags.append(token)
    return tuple(tags)


def _core_duplicate_key(shifted: Sequence[object]) -> tuple[str, ...]:
    width = min(CORE_WIDTH, len(shifted))
    return tuple(_cell_text(shifted[i]) for i in range(width))


def _header_is_canon(shifted_header: Sequence[object]) -> bool:
    """True when shifted A–L classify to ``CORE_ROLES`` in order."""

    for expected, role in enumerate(CORE_ROLES):
        cell = shifted_header[expected] if expected < len(shifted_header) else None
        if classify_header_role(cell) != role:
            return False
    return True


def _sample_qty_slot_unit_words(
    ws_f: object,
    *,
    header_row: int | None,
    qty_col_index: int,
    max_used: int = _QTY_SLOT_SAMPLE_ROWS,
) -> tuple[str, ...]:
    """Return unit-word hits in the fixed canon qty slot on the first data rows."""

    if qty_col_index < 0:
        return ()
    seen: set[str] = set()
    used = 0
    for excel_row, values, _cols in _iter_xml_sheet_rows(
        ws_f, max_col=qty_col_index + 1
    ):
        if header_row is not None and excel_row <= header_row:
            continue
        if not _row_used(values):
            continue
        used += 1
        cell = values[qty_col_index] if qty_col_index < len(values) else None
        seen.update(_unit_words_in_text(cell))
        if used >= max_used:
            break
    return tuple(word for word in _QTY_SLOT_UNIT_WORDS if word in seen)


def _build_layout_hint(
    header_values: Sequence[object],
    *,
    offset: int,
    qty_slot_unit_words: Sequence[str] = (),
) -> str:
    """One Russian sentence describing why the header is not the A–L canon."""

    parts: list[str] = []
    raw_found = _role_map(header_values)
    shifted = list(header_values[offset:]) if offset else list(header_values)
    for expected, role in enumerate(CORE_ROLES):
        short = _ROLE_SHORT_RU[role]
        canon_col0 = offset + expected
        canon_letter = get_column_letter(canon_col0 + 1)
        cell = shifted[expected] if expected < len(shifted) else None
        if classify_header_role(cell) == role:
            continue
        found_col0 = raw_found.get(role)
        if _is_empty(cell):
            parts.append(f"пустой столбец {canon_letter} между ролями")
        if found_col0 is None:
            parts.append(f"нет «{short}»")
        elif found_col0 != canon_col0:
            found_letter = get_column_letter(found_col0 + 1)
            parts.append(f"{short} в {found_letter}, в каноне {canon_letter}")
    role_parts = parts[:_LAYOUT_HINT_MAX_PARTS]
    if qty_slot_unit_words:
        shown = ", ".join(qty_slot_unit_words)
        role_parts.append(f"в слоте количества значения ед. изм. ({shown})")
    message = "Раскладка не канон A–L."
    if role_parts:
        message = f"{message} {'. '.join(role_parts)}."
    if len(message) > _LAYOUT_HINT_MAX_CHARS:
        return message[: _LAYOUT_HINT_MAX_CHARS - 1] + "…"
    return message


def _parse_workbook(
    source: DsSourceFile,
    *,
    active_ids: Sequence[str],
    registry_by_id: dict[str, DsRegistryRow],
    issues: list[DsBaselineIssue],
    sheet_scans: list[tuple[str, DsSheetScan]],
    column_notes: list[dict[str, object]],
) -> list[DsBaselinePosition]:
    positions: list[DsBaselinePosition] = []
    wb_f: OpenpyxlWorkbook | None = None
    wb_v: OpenpyxlWorkbook | None = None
    buf_f: BytesIO | None = None
    buf_v: BytesIO | None = None
    try:
        data = source.path.read_bytes()
        wb_f, wb_v, buf_f, buf_v = _open_pair(data)
    except Exception as exc:  # noqa: BLE001 - workbook may be corrupt
        issues.append(
            _issue(
                ISSUE_WORKBOOK,
                "ERROR",
                f"не прочитан workbook: {type(exc).__name__}: {exc}",
                path=source.path,
                relpath=source.relpath,
            )
        )
        return positions

    try:
        return _parse_open_workbook(
            source,
            wb_f=wb_f,
            wb_v=wb_v,
            active_ids=active_ids,
            registry_by_id=registry_by_id,
            issues=issues,
            sheet_scans=sheet_scans,
            column_notes=column_notes,
        )
    finally:
        _close_pair(wb_f, wb_v, buf_f, buf_v)


def _parse_open_workbook(
    source: DsSourceFile,
    *,
    wb_f: OpenpyxlWorkbook,
    wb_v: OpenpyxlWorkbook,
    active_ids: Sequence[str],
    registry_by_id: dict[str, DsRegistryRow],
    issues: list[DsBaselineIssue],
    sheet_scans: list[tuple[str, DsSheetScan]],
    column_notes: list[dict[str, object]],
) -> list[DsBaselinePosition]:
    positions: list[DsBaselinePosition] = []
    active_name = wb_f.active.title if wb_f.active is not None else ""
    scans: list[DsSheetScan] = []
    for name in wb_f.sheetnames:
        ws_f = wb_f[name]
        if not _is_tabular_sheet(ws_f):
            continue
        _prepare_read_only_sheet(ws_f)
        prefix = _read_prefix_rows(
            ws_f, max_row=HEADER_SCAN_ROWS, max_col=SCAN_MAX_COL
        )
        scan = _scan_sheet(name, is_active=name == active_name, rows=prefix)
        scans.append(scan)
        sheet_scans.append((source.relpath, scan))

    candidates = [item for item in scans if item.is_candidate]
    candidate_names = [item.name for item in candidates]
    resolution = resolve_ds_source_id(source.relpath, active_ids)
    source_id = resolution.source_id or ""
    if resolution.method == "unresolved":
        issues.append(
            _issue(
                ISSUE_UNRESOLVED_ID,
                "ERROR",
                "не сопоставлен ID ДС источника с активным реестром",
                path=source.path,
                relpath=source.relpath,
            )
        )
    elif resolution.method == "ambiguous":
        issues.append(
            _issue(
                ISSUE_AMBIGUOUS_ID,
                "ERROR",
                "неоднозначный ID ДС источника: "
                + ", ".join(resolution.candidates),
                path=source.path,
                relpath=source.relpath,
            )
        )

    if len(candidates) > 1:
        issues.append(
            _issue(
                ISSUE_MULTI_DATA_SHEET,
                "ERROR",
                "больше одного листа-кандидата с шапкой ДС: "
                + ", ".join(candidate_names),
                path=source.path,
                relpath=source.relpath,
                sheet=active_name,
                source_id=source_id,
            )
        )
    active_scan = next((item for item in scans if item.is_active), None)
    if active_scan is None or not active_scan.is_candidate:
        extra_bits: list[str] = []
        if candidate_names and active_name not in candidate_names:
            extra_bits.append(f"шапка на {', '.join(candidate_names)}")
        found_roles = active_scan.roles if active_scan is not None else ()
        if found_roles:
            labels = [_ROLE_SHORT_RU.get(role, role) for role in found_roles]
            extra_bits.append("найдены роли: " + ", ".join(labels))
        extra = ("; " + "; ".join(extra_bits)) if extra_bits else ""
        issues.append(
            _issue(
                ISSUE_NO_HEADER,
                "ERROR",
                f"активный лист {active_name!r} не является единственным data-sheet ДС"
                + extra,
                path=source.path,
                relpath=source.relpath,
                sheet=active_name,
                source_id=source_id,
            )
        )
        return positions

    sheet_name = active_scan.name
    ws_f = wb_f[sheet_name]
    ws_v = wb_v[sheet_name] if sheet_name in wb_v.sheetnames else None
    _prepare_read_only_sheet(ws_f)
    if ws_v is not None:
        _prepare_read_only_sheet(ws_v)

    header_row_num = active_scan.header_row
    offset = active_scan.leading_empty
    header_values: list[object] = []
    if header_row_num is not None:
        for excel_row, values, _cols in _iter_xml_sheet_rows(
            ws_f, max_row=header_row_num
        ):
            if excel_row == header_row_num:
                header_values = _rstrip_empty(values)
                break
        if header_values:
            offset = _leading_empty_count(header_values)

    width = max(len(header_values), offset + CORE_WIDTH)
    shifted_header = _pad(header_values, width)[offset:]
    tag_columns = [
        index
        for index, value in enumerate(header_values)
        if _is_tag_header(value)
    ]
    read_max_col = _read_max_col(offset, tag_columns)
    column_notes.append(
        {
            "relpath": source.relpath,
            "path": source.path,
            "sheet": sheet_name,
            "header_row": header_row_num,
            "leading_empty": offset,
            "roles": ", ".join(
                f"{role}={get_column_letter(ROLE_INDEX[role] + 1)}"
                for role in CORE_ROLES
            ),
            "tag_columns": ", ".join(
                get_column_letter(index + 1) for index in tag_columns
            ),
            "tail_after_L": max(len(header_values) - offset - CORE_WIDTH, 0),
        }
    )
    if not _header_is_canon(shifted_header):
        qty_slot_index = offset + ROLE_INDEX["qty"]
        issues.append(
            _issue(
                ISSUE_INTERNAL_SHIFT,
                "ERROR",
                _build_layout_hint(
                    header_values,
                    offset=offset,
                    qty_slot_unit_words=_sample_qty_slot_unit_words(
                        ws_f,
                        header_row=header_row_num,
                        qty_col_index=qty_slot_index,
                    ),
                ),
                path=source.path,
                relpath=source.relpath,
                sheet=sheet_name,
                excel_row=header_row_num,
                source_id=source_id,
            )
        )
        return positions

    registry_row = registry_by_id.get(source_id)
    seen_core: dict[tuple[str, ...], int] = {}
    qty_col_index = offset + ROLE_INDEX["qty"]
    emitted_group_fallback = False
    for excel_row, formula_values, value_values, formula_cols in _iter_paired_sheet_rows(
        ws_f, ws_v, max_col=read_max_col
    ):
        if header_row_num is not None and excel_row <= header_row_num:
            continue
        padded = _pad(formula_values, read_max_col)
        if not _row_used(padded):
            continue
        shifted = padded[offset:]
        if not _row_used(shifted[:CORE_WIDTH]):
            continue
        if _is_repeated_header(shifted) or _units_party_label(shifted):
            continue
        if _is_boilerplate_shifted(shifted):
            continue

        value_row = _pad(value_values, read_max_col)
        shifted_values = value_row[offset:]

        def at(role: str, from_values: Sequence[object] = shifted_values) -> object:
            index = ROLE_INDEX[role]
            return from_values[index] if index < len(from_values) else None

        npp = _cell_text(at("npp"))
        if _has_footer_needle(npp) or _has_footer_needle(
            _core_text_blob(shifted_values)
        ):
            continue
        qty_formula_obj = at("qty", from_values=shifted)
        if _qty_formula_is_aggregate(qty_formula_obj):
            continue
        title = _cell_text(at("title"))
        system = _cell_text(at("system"))
        specification = _cell_text(at("specification"))
        rfq = _cell_text(at("rfq"))
        name = _cell_text(at("name"))
        code_1c = _cell_text(at("code_1c"))
        code = _cell_text(at("code"))
        supplier = _cell_text(at("supplier"))
        type_mark = _cell_text(at("type"))
        units_raw = _cell_text(at("units"))
        units = normalize_units_text(units_raw)
        qty_raw_obj = at("qty")
        qty_raw_text = _cell_text(qty_raw_obj)
        qty_probe, _qty_probe_error = _parse_qty(qty_raw_obj)
        if _all_core_ref(shifted) or _all_core_ref(shifted_values):
            continue
        has_identity = bool(title or system or name or supplier or code)
        if not has_identity and qty_probe is None:
            continue
        has_formula = qty_col_index in formula_cols
        if has_formula:
            issues.append(
                _issue(
                    ISSUE_QTY_FORMULA,
                    "ERROR",
                    f"в количестве формула {qty_raw_text or '(data_only пусто)'}",
                    path=source.path,
                    relpath=source.relpath,
                    sheet=sheet_name,
                    excel_row=excel_row,
                    source_id=source_id,
                )
            )
            qty: Decimal | None = None
            qty_error = ISSUE_QTY_FORMULA
        else:
            qty, qty_error = _parse_qty(qty_raw_obj)
            if qty_error == ISSUE_QTY_NEGATIVE:
                issues.append(
                    _issue(
                        ISSUE_QTY_NEGATIVE,
                        "ERROR",
                        f"отрицательное количество {qty_raw_text}",
                        path=source.path,
                        relpath=source.relpath,
                        sheet=sheet_name,
                        excel_row=excel_row,
                        source_id=source_id,
                    )
                )
            elif qty_error == ISSUE_QTY_EMPTY:
                issues.append(
                    _issue(
                        ISSUE_QTY_EMPTY,
                        "ERROR",
                        "пустое количество",
                        path=source.path,
                        relpath=source.relpath,
                        sheet=sheet_name,
                        excel_row=excel_row,
                        source_id=source_id,
                    )
                )
            elif qty_error == ISSUE_QTY_NON_FINITE:
                issues.append(
                    _issue(
                        ISSUE_QTY_NON_FINITE,
                        "ERROR",
                        f"количество не конечное {qty_raw_text}",
                        path=source.path,
                        relpath=source.relpath,
                        sheet=sheet_name,
                        excel_row=excel_row,
                        source_id=source_id,
                    )
                )
            elif qty_error == ISSUE_QTY_NON_NUMERIC:
                issues.append(
                    _issue(
                        ISSUE_QTY_NON_NUMERIC,
                        "ERROR",
                        f"количество не число {qty_raw_text!r}",
                        path=source.path,
                        relpath=source.relpath,
                        sheet=sheet_name,
                        excel_row=excel_row,
                        source_id=source_id,
                    )
                )
            elif qty is not None and qty == 0:
                issues.append(
                    _issue(
                        ISSUE_QTY_ZERO,
                        "WARN",
                        "нулевое количество",
                        path=source.path,
                        relpath=source.relpath,
                        sheet=sheet_name,
                        excel_row=excel_row,
                        source_id=source_id,
                    )
                )

        code_normalized = normalize_code(code)
        if not code_normalized:
            issues.append(
                _issue(
                    ISSUE_EMPTY_CODE,
                    "ERROR",
                    "пустой закупочный код",
                    path=source.path,
                    relpath=source.relpath,
                    sheet=sheet_name,
                    excel_row=excel_row,
                    source_id=source_id,
                )
            )

        group_id, group_label, mode, fallback, overlay_blocked, group_msg = (
            _assign_group(
                source_id=source_id,
                title=title,
                system=system,
                registry_row=registry_row,
            )
        )
        if fallback and group_msg and not emitted_group_fallback:
            issues.append(
                _issue(
                    ISSUE_GROUP_FALLBACK,
                    "OVERLAY",
                    group_msg,
                    blocking=False,
                    path=source.path,
                    relpath=source.relpath,
                    sheet=sheet_name,
                    excel_row=excel_row,
                    source_id=source_id,
                    group_id=group_id,
                )
            )
            emitted_group_fallback = True

        dup_key = _core_duplicate_key(shifted_values)
        previous = seen_core.get(dup_key)
        if previous is not None:
            issues.append(
                _issue(
                    ISSUE_EXACT_DUPLICATE,
                    "WARN",
                    f"точный дубль исходной строки (как строка {previous})",
                    path=source.path,
                    relpath=source.relpath,
                    sheet=sheet_name,
                    excel_row=excel_row,
                    source_id=source_id,
                    group_id=group_id,
                )
            )
        else:
            seen_core[dup_key] = excel_row

        tags = _diagnostic_tags(padded, tag_columns)
        folded_tags = [item.casefold() for item in tags]
        if len(folded_tags) != len(set(folded_tags)):
            issues.append(
                _issue(
                    ISSUE_TAG_DUPLICATE,
                    "WARN",
                    "дубли тегов внутри строки",
                    path=source.path,
                    relpath=source.relpath,
                    sheet=sheet_name,
                    excel_row=excel_row,
                    source_id=source_id,
                    group_id=group_id,
                )
            )
        if (
            tags
            and qty is not None
            and qty == qty.to_integral_value()
            and len(tags) != int(qty)
        ):
            issues.append(
                _issue(
                    ISSUE_TAG_MISMATCH,
                    "WARN",
                    f"число тегов {len(tags)} ≠ количество {qty}",
                    path=source.path,
                    relpath=source.relpath,
                    sheet=sheet_name,
                    excel_row=excel_row,
                    source_id=source_id,
                    group_id=group_id,
                )
            )

        positions.append(
            DsBaselinePosition(
                source_id=source_id,
                group_id=group_id,
                group_label=group_label,
                grouping_mode=mode,
                grouping_fallback=fallback,
                overlay_blocked=overlay_blocked,
                relpath=source.relpath,
                file_name=source.path.name,
                path=source.path,
                sheet=sheet_name,
                excel_row=excel_row,
                ds_number=npp,
                title=title,
                system=system,
                specification=specification,
                rfq=rfq,
                name=name,
                code_1c=code_1c,
                code=code,
                code_normalized=code_normalized,
                supplier_name=supplier,
                type_mark=type_mark,
                units=units,
                qty=qty,
                qty_raw=qty_raw_text,
                qty_source=qty,
                unit_source=units,
                diagnostic_tags=tags,
                leading_empty=offset,
            )
        )
    return positions



# ---------------------------------------------------------------------------
# Units conversion
# ---------------------------------------------------------------------------


def _identity_trace(
    *,
    code: str,
    unit: str,
    qty: Decimal,
) -> str:
    return (
        f"code={code}; src={unit}; tgt={unit}; "
        f"qty={qty}; coef=1; result={qty}; status={STATUS_IDENTITY}"
    )


class IdentityDsUnitsConverter:
    """Offline converter: identity qty/unit, status/trace preserved, no I/O."""

    def convert_positions(
        self, positions: Sequence[DsBaselinePosition]
    ) -> DsConversionBatch:
        converted: list[DsBaselinePosition] = []
        for item in positions:
            if item.qty is None:
                converted.append(item)
                continue
            converted.append(
                replace(
                    item,
                    qty_target=item.qty,
                    unit_target=item.units,
                    conversion_status=STATUS_IDENTITY,
                    conversion_trace=_identity_trace(
                        code=item.code or item.code_normalized,
                        unit=item.units,
                        qty=item.qty,
                    ),
                    conversion_coefficient=Decimal("1"),
                )
            )
        return DsConversionBatch(converted=tuple(converted), plan=None)


class RfQDsUnitsConverter:
    """Default adapter: ``build_conversion_plan`` / apply without RowStd."""

    def __init__(
        self,
        *,
        google_index: GoogleUnitsIndex | None = None,
        matrix_path: str | Path | None = None,
    ) -> None:
        self.google_index = google_index
        self.matrix_path = Path(matrix_path) if matrix_path else None

    def convert_positions(
        self, positions: Sequence[DsBaselinePosition]
    ) -> DsConversionBatch:
        if self.google_index is None or self.matrix_path is None:
            raise UnitsConversionError(
                "конвертация ДС требует google_index и matrix_path; "
                "для offline-тестов передайте IdentityDsUnitsConverter"
            )
        requests: list[ConversionRequest] = []
        for index, item in enumerate(positions):
            if item.qty is None or not item.code_normalized:
                continue
            requests.append(
                ConversionRequest(
                    request_id=f"ds-{index}",
                    contour="ds",
                    code=item.code,
                    source_unit=item.unit_source or item.units,
                    quantity=item.qty,
                    tags_count=0,
                    invariant=ConversionInvariant.NONE,
                    location=f"{item.relpath}·{item.sheet}·{item.excel_row}",
                    item_name=item.name,
                )
            )
        if not requests:
            return IdentityDsUnitsConverter().convert_positions(positions)

        plan = build_conversion_plan(
            requests, self.google_index, self.matrix_path
        )
        actions = {action.request_id: action for action in plan.actions}
        converted: list[DsBaselinePosition] = []
        for index, item in enumerate(positions):
            action = actions.get(f"ds-{index}")
            if action is None:
                converted.append(item)
                continue
            converted.append(
                replace(
                    item,
                    qty=converted_quantity(action),
                    units=converted_unit(action),
                    qty_target=converted_quantity(action),
                    unit_target=converted_unit(action),
                    conversion_status=action.status,
                    conversion_trace=action.trace,
                    conversion_coefficient=action.coefficient,
                )
            )
        return DsConversionBatch(
            converted=tuple(converted),
            plan=plan,
            warnings=plan.warnings,
        )


# ---------------------------------------------------------------------------
# Reports + baseline writer
# ---------------------------------------------------------------------------


def _style_header(ws: Worksheet, width: int) -> None:
    for col in range(1, width + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(width)}{max(ws.max_row, 1)}"


def _write_summary_sheet(
    ws: Worksheet,
    rows: Sequence[tuple[str, object]],
) -> None:
    ws.title = "Сводка"
    ws.append(["Показатель", "Значение"])
    for label, value in rows:
        ws.append([label, value if not isinstance(value, Path) else str(value)])
        if isinstance(value, Path):
            _set_path_cell(ws.cell(row=ws.max_row, column=2), value)
    _style_header(ws, 2)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 80


def _mark_level(cell, level: str) -> None:
    # WARN/OVERLAY fills on tens of thousands of rows are an openpyxl bottleneck.
    if level == "ERROR":
        cell.fill = _FILL_ERROR


def _emit_phase(
    phase_callback: Callable[[str], None] | None,
    text: str,
) -> None:
    if phase_callback is not None:
        phase_callback(text)


def _phase_timer(
    phase_callback: Callable[[str], None] | None,
    label: str,
) -> Callable[..., None]:
    started = time.perf_counter()
    _emit_phase(phase_callback, label)

    def done(suffix: str = "") -> None:
        elapsed = time.perf_counter() - started
        extra = f" ({suffix})" if suffix else ""
        _emit_phase(phase_callback, f"{label}{extra} — {elapsed:.1f} с")

    return done


def _display_tag(owners: Sequence[DsBaselinePosition], tag_key: str) -> str:
    for item in owners:
        for tag in item.diagnostic_tags:
            if tag.casefold() == tag_key:
                return tag
    return tag_key


def _write_structure_report(
    path: Path,
    *,
    result_head: Sequence[tuple[str, object]],
    files: Sequence[DsSourceFile],
    skipped: Sequence[DsSkippedFile],
    sheet_scans: Sequence[tuple[str, DsSheetScan]],
    column_notes: Sequence[dict[str, object]],
    issues: Sequence[DsBaselineIssue],
    progress: Callable[[str], None] | None = None,
) -> Path:
    hb = DsHeartbeat(progress, "отчёт по структуре")
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    _write_summary_sheet(summary, result_head)

    files_ws = wb.create_sheet("Файлы ДС")
    files_ws.append(
        [
            "Относительный путь",
            "Полный путь",
            "Размер",
            "mtime_ns",
            "fingerprint",
        ]
    )
    for item in files:
        files_ws.append(
            [item.relpath, str(item.path), item.size, item.mtime_ns, item.fingerprint]
        )
        _set_path_cell(files_ws.cell(row=files_ws.max_row, column=1), item.path)
        _set_path_cell(files_ws.cell(row=files_ws.max_row, column=2), item.path)
    if skipped:
        files_ws.append([])
        files_ws.append(["Пропущенные файлы", "Причина", "Путь"])
        for item in skipped:
            files_ws.append([item.relpath, item.reason, str(item.path)])
            _set_path_cell(files_ws.cell(row=files_ws.max_row, column=3), item.path)
    _style_header(files_ws, 5)
    files_ws.column_dimensions["A"].width = 40
    files_ws.column_dimensions["B"].width = 70

    cols_ws = wb.create_sheet("Колонки")
    cols_ws.append(
        [
            "Файл",
            "Путь",
            "Лист",
            "Строка шапки",
            "Сдвиг пустых",
            "Роли A-L",
            "Колонки тегов",
            "Хвост после L",
        ]
    )
    for note in column_notes:
        path_obj = note["path"]
        cols_ws.append(
            [
                note["relpath"],
                str(path_obj),
                note["sheet"],
                note["header_row"],
                note["leading_empty"],
                note["roles"],
                note["tag_columns"],
                note["tail_after_L"],
            ]
        )
        _set_path_cell(cols_ws.cell(row=cols_ws.max_row, column=1), path_obj)
        _set_path_cell(cols_ws.cell(row=cols_ws.max_row, column=2), path_obj)
    _style_header(cols_ws, 8)
    cols_ws.column_dimensions["A"].width = 36
    cols_ws.column_dimensions["B"].width = 70
    cols_ws.column_dimensions["F"].width = 40

    sheets_ws = wb.create_sheet("Листы")
    sheets_ws.append(
        [
            "Файл",
            "Лист",
            "Активный",
            "Кандидат ДС",
            "Строка шапки",
            "Оценка",
            "Роли",
            "Сдвиг",
        ]
    )
    file_by_rel = {item.relpath: item.path for item in files}
    for relpath, scan in sheet_scans:
        sheets_ws.append(
            [
                relpath,
                scan.name,
                "да" if scan.is_active else "",
                "да" if scan.is_candidate else "",
                scan.header_row,
                scan.score,
                ", ".join(scan.roles),
                scan.leading_empty,
            ]
        )
        _set_path_cell(sheets_ws.cell(row=sheets_ws.max_row, column=1), file_by_rel.get(relpath))
    _style_header(sheets_ws, 8)
    sheets_ws.column_dimensions["A"].width = 40

    problems = wb.create_sheet("Проблемы строк")
    problems.append(
        [
            "Уровень",
            "Код",
            "Файл",
            "Путь",
            "Лист",
            "Строка",
            "ID ДС",
            "Группа",
            "Сообщение",
        ]
    )
    issue_count = len(issues)
    for index, item in enumerate(issues, start=1):
        problems.append(
            [
                item.level,
                item.code,
                item.relpath,
                str(item.path) if item.path else "",
                item.sheet,
                item.excel_row,
                item.source_id,
                item.group_id,
                item.message,
            ]
        )
        _mark_level(problems.cell(row=problems.max_row, column=1), item.level)
        if index == 1 or index % 500 == 0 or index == issue_count:
            hb.tick(f"замечания {index}/{issue_count}")
    _style_header(problems, 9)
    problems.column_dimensions["C"].width = 36
    problems.column_dimensions["D"].width = 70
    problems.column_dimensions["I"].width = 70
    hb.tick("сохранение xlsx", force=True)
    saved = _save_workbook_atomic(path, wb, progress=progress)
    hb.finish("сохранён")
    return saved


def _conversion_needs_report(item: DsBaselinePosition) -> bool:
    """True when a position is not a pure identity conversion."""

    status = (item.conversion_status or "").strip()
    if not status or status == STATUS_IDENTITY:
        return False
    return True


def _write_quality_report(
    path: Path,
    *,
    result_head: Sequence[tuple[str, object]],
    positions: Sequence[DsBaselinePosition],
    issues: Sequence[DsBaselineIssue],
    files: Sequence[DsSourceFile],
    progress: Callable[[str], None] | None = None,
) -> Path:
    hb = DsHeartbeat(progress, "отчёт по качеству")
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    _write_summary_sheet(summary, result_head)

    def _pos_row(item: DsBaselinePosition, extra: Sequence[object] = ()) -> list[object]:
        return [
            item.relpath,
            str(item.path),
            item.sheet,
            item.excel_row,
            item.source_id,
            item.group_label,
            item.code,
            item.name,
            str(item.qty) if item.qty is not None else item.qty_raw,
            item.units,
            *extra,
        ]

    def _write_pos_sheet(
        title: str,
        rows: Sequence[DsBaselinePosition],
        extra_headers: Sequence[str] = (),
        extra_fn=None,
    ) -> None:
        ws = wb.create_sheet(title)
        headers = [
            "Файл",
            "Путь",
            "Лист",
            "Строка",
            "ID ДС",
            "Группа",
            "Код",
            "Наименование",
            "Кол-во",
            "Ед. изм.",
            *extra_headers,
        ]
        ws.append(list(headers))
        for item in rows:
            extra = extra_fn(item) if extra_fn else ()
            ws.append(_pos_row(item, extra))
        _style_header(ws, len(headers))
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 70

    empty_code = [item for item in positions if not item.code_normalized]
    _write_pos_sheet("Позиции без кода", empty_code)

    qty_issues = [
        item
        for item in issues
        if item.code
        in {
            ISSUE_QTY_FORMULA,
            ISSUE_QTY_EMPTY,
            ISSUE_QTY_NON_NUMERIC,
            ISSUE_QTY_NON_FINITE,
            ISSUE_QTY_NEGATIVE,
            ISSUE_QTY_ZERO,
        }
    ]
    qty_ws = wb.create_sheet("Количества")
    qty_ws.append(
        ["Уровень", "Код", "Файл", "Путь", "Лист", "Строка", "Сообщение"]
    )
    for item in qty_issues:
        qty_ws.append(
            [
                item.level,
                item.code,
                item.relpath,
                str(item.path) if item.path else "",
                item.sheet,
                item.excel_row,
                item.message,
            ]
        )
        _mark_level(qty_ws.cell(row=qty_ws.max_row, column=1), item.level)
    _style_header(qty_ws, 7)
    qty_ws.column_dimensions["C"].width = 36
    qty_ws.column_dimensions["D"].width = 70
    qty_ws.column_dimensions["G"].width = 60

    units_rows = [
        item for item in positions if _conversion_needs_report(item)
    ]
    if units_rows:
        units_ws = wb.create_sheet("Единицы и конвертация")
        units_ws.append(
            [
                "Файл",
                "Путь",
                "Лист",
                "Строка",
                "Код",
                "Исходное кол-во",
                "Исходная ЕИ",
                "Коэффициент",
                "Целевое кол-во",
                "Целевая ЕИ",
                "Статус",
                "Trace",
            ]
        )
        units_count = len(units_rows)
        for index, item in enumerate(units_rows, start=1):
            units_ws.append(
                [
                    item.relpath,
                    str(item.path),
                    item.sheet,
                    item.excel_row,
                    item.code,
                    str(item.qty_source) if item.qty_source is not None else "",
                    item.unit_source,
                    str(item.conversion_coefficient)
                    if item.conversion_coefficient is not None
                    else "",
                    str(item.qty_target) if item.qty_target is not None else "",
                    item.unit_target,
                    item.conversion_status,
                    item.conversion_trace,
                ]
            )
            if index == 1 or index % 500 == 0 or index == units_count:
                hb.tick(f"единицы {index}/{units_count}")
        _style_header(units_ws, 12)
        units_ws.column_dimensions["A"].width = 36
        units_ws.column_dimensions["B"].width = 70
        units_ws.column_dimensions["L"].width = 70

    tag_notes: dict[tuple[str, int | None], list[str]] = defaultdict(list)
    for iss in issues:
        if iss.code in {ISSUE_TAG_DUPLICATE, ISSUE_TAG_MISMATCH}:
            tag_notes[(iss.relpath, iss.excel_row)].append(iss.message)
    tags_ws = wb.create_sheet("Теги")
    tags_ws.append(
        [
            "Файл",
            "Путь",
            "Лист",
            "Строка",
            "Код",
            "Кол-во",
            "Число тегов",
            "Теги",
            "Замечание",
        ]
    )
    for item in positions:
        note = "; ".join(tag_notes.get((item.relpath, item.excel_row), ()))
        if not note:
            continue
        tags_ws.append(
            [
                item.relpath,
                str(item.path),
                item.sheet,
                item.excel_row,
                item.code,
                str(item.qty) if item.qty is not None else "",
                len(item.diagnostic_tags),
                "; ".join(item.diagnostic_tags),
                note,
            ]
        )
    _style_header(tags_ws, 9)
    tags_ws.column_dimensions["A"].width = 36
    tags_ws.column_dimensions["B"].width = 70

    dup_ws = wb.create_sheet("Дубли")
    dup_ws.append(
        ["Уровень", "Код", "Файл", "Путь", "Лист", "Строка", "Сообщение"]
    )
    for item in issues:
        if item.code not in {ISSUE_EXACT_DUPLICATE, ISSUE_TAG_DUPLICATE}:
            continue
        dup_ws.append(
            [
                item.level,
                item.code,
                item.relpath,
                str(item.path) if item.path else "",
                item.sheet,
                item.excel_row,
                item.message,
            ]
        )
        _mark_level(dup_ws.cell(row=dup_ws.max_row, column=1), item.level)
    _style_header(dup_ws, 7)
    dup_ws.column_dimensions["C"].width = 36
    dup_ws.column_dimensions["D"].width = 70

    google_ws = wb.create_sheet("Коды вне Google")
    google_ws.append(
        ["Файл", "Путь", "Лист", "Строка", "Код", "Наименование", "Сообщение"]
    )
    for item in issues:
        if item.code != ISSUE_UNKNOWN_GOOGLE:
            continue
        google_ws.append(
            [
                item.relpath,
                str(item.path) if item.path else "",
                item.sheet,
                item.excel_row,
                "",
                "",
                item.message,
            ]
        )
    pos_by_loc = {(p.relpath, p.excel_row): p for p in positions}
    for row in range(2, google_ws.max_row + 1):
        rel = str(google_ws.cell(row=row, column=1).value or "")
        excel_row = google_ws.cell(row=row, column=4).value
        pos = pos_by_loc.get((rel, excel_row))
        if pos is not None:
            google_ws.cell(row=row, column=5, value=pos.code)
            google_ws.cell(row=row, column=6, value=pos.name)
    _style_header(google_ws, 7)
    google_ws.column_dimensions["A"].width = 36
    google_ws.column_dimensions["B"].width = 70

    src_ws = wb.create_sheet("Источники")
    src_ws.append(
        ["Относительный путь", "Полный путь", "Размер", "mtime_ns", "fingerprint"]
    )
    for item in files:
        src_ws.append(
            [item.relpath, str(item.path), item.size, item.mtime_ns, item.fingerprint]
        )
        _set_path_cell(src_ws.cell(row=src_ws.max_row, column=1), item.path)
        _set_path_cell(src_ws.cell(row=src_ws.max_row, column=2), item.path)
    _style_header(src_ws, 5)
    src_ws.column_dimensions["A"].width = 40
    src_ws.column_dimensions["B"].width = 70
    hb.tick("сохранение xlsx", force=True)
    saved = _save_workbook_atomic(path, wb, progress=progress)
    hb.finish("сохранён")
    return saved


def _write_sidecar_positions(
    path: Path,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    path_cols: Sequence[int],
    progress: Callable[[str], None] | None = None,
) -> Path:
    del path_cols
    hb = DsHeartbeat(progress, f"sidecar {title}")
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title
    ws.append(list(headers))
    row_count = len(rows)
    for index, row in enumerate(rows, start=1):
        ws.append(list(row))
        if index == 1 or index % 500 == 0 or index == row_count:
            hb.tick(f"{index}/{row_count}")
    _style_header(ws, len(headers))
    hb.tick("сохранение xlsx", force=True)
    saved = _save_workbook_atomic(path, wb, progress=progress)
    hb.finish("сохранён")
    return saved


def _positions_to_net_rows(
    positions: Sequence[DsBaselinePosition],
) -> list[NetSummaryRow]:
    rows: list[NetSummaryRow] = []
    seq = 0
    for item in positions:
        if item.qty is None:
            continue
        seq += 1
        record = RfpRecord(
            kind="ds",
            file_name=item.file_name,
            sheet=item.sheet,
            excel_row=item.excel_row,
            ds_name=item.group_label,
            ds_number=str(seq),
            ds_title=_composite_title(item.title, item.system),
            ds_specification=item.specification,
            tags="",
            ds_code_1c=item.code_1c,
            code=item.code,
            name=item.name,
            type_mark=item.type_mark,
            values=item.qty,
            units=item.units,
            units_check_status=item.conversion_status,
            units_conversion_trace=item.conversion_trace,
        )
        rows.append(NetSummaryRow(record=record, source="ДС baseline"))
    return rows


def _build_groups(
    positions: Sequence[DsBaselinePosition],
) -> list[DsSupplyGroupInfo]:
    buckets: dict[str, list[DsBaselinePosition]] = defaultdict(list)
    for item in positions:
        key = item.group_id or item.group_label or item.source_id
        buckets[key].append(item)
    groups: list[DsSupplyGroupInfo] = []
    for group_id in sorted(buckets):
        items = buckets[group_id]
        source_ids = tuple(dict.fromkeys(item.source_id for item in items if item.source_id))
        modes = tuple(dict.fromkeys(item.grouping_mode for item in items if item.grouping_mode))
        groups.append(
            DsSupplyGroupInfo(
                group_id=group_id,
                group_label=items[0].group_label or group_id,
                source_ids=source_ids,
                modes=modes,
                fallback=any(item.grouping_fallback for item in items),
                overlay_blocked=any(item.overlay_blocked for item in items),
                position_count=len(items),
            )
        )
    return groups


def _tree_fingerprint(
    files: Sequence[DsSourceFile],
    *,
    registry: DsRegistryDocument,
) -> str:
    digest = hashlib.sha256()
    digest.update(ALGORITHM_VERSION.encode("utf-8"))
    digest.update(str(registry.path).encode("utf-8", errors="replace"))
    if registry.path.is_file():
        try:
            digest.update(registry.path.read_bytes())
        except OSError:
            pass
    for item in files:
        digest.update(
            f"{item.relpath}|{item.size}|{item.mtime_ns}|{item.fingerprint}".encode(
                "utf-8"
            )
        )
    return digest.hexdigest()


def _copy_registry_issues(
    registry: DsRegistryDocument,
) -> list[DsBaselineIssue]:
    copied: list[DsBaselineIssue] = []
    for item in registry.validation.issues:
        blocking = item.level == "ERROR"
        copied.append(
            DsBaselineIssue(
                code=item.code or ISSUE_REGISTRY,
                level=item.level,
                message=item.message,
                blocking=blocking,
                path=registry.path,
                relpath=registry.path.name,
                excel_row=item.excel_row,
                source_id=item.source_id,
                group_id=item.group_id,
            )
        )
    return copied


# ---------------------------------------------------------------------------
# Public entry
# ---------------------------------------------------------------------------


def build_ds_baseline(
    source_root: str | Path,
    registry: DsRegistryDocument,
    output_dir: str | Path,
    *,
    write_baseline: bool = True,
    converter: DsUnitsConverter | None = None,
    google_index: GoogleUnitsIndex | None = None,
    matrix_path: str | Path | None = None,
    stamp: str | None = None,
    equipment_by_code: dict[str, str] | None = None,
    progress_callback: Callable[[int, int, str, float | None], object] | None = None,
    phase_callback: Callable[[str], object] | None = None,
    files_discovered_callback: Callable[[int, int], object] | None = None,
) -> DsBaselineResult:
    """Audit DS workbooks, write reports, and optionally write the Step1 baseline.

    Args:
        source_root: Recursive folder of DS xlsx/xlsm.
        registry: Already loaded canonical registry document.
        output_dir: Destination for reports and the baseline. Required; never
            defaults to UNC.
        write_baseline: When True, write ``Свод ДС для запуска.xlsx`` iff there
            are zero blocking issues.
        converter: Units conversion. Default uses RFQ plan/matrix/Google.
            Pass ``IdentityDsUnitsConverter`` for offline tests.
        google_index: Optional Google units index for unknown-code warnings
            and the default converter.
        matrix_path: Units matrix for the default converter. Not used by
            the identity converter.
        stamp: Report timestamp ``YYYYMMDD_HHMMSS``. Default is now.
        equipment_by_code: Optional Google equipment map for net column F.
        progress_callback: Optional ``(index, total, relpath, file_elapsed)`` hook.
            ``file_elapsed`` is ``None`` before parse and seconds after.
        phase_callback: Optional short Russian phase label for live Job monitor.
        files_discovered_callback: ``(file_count, skipped_count)`` after scan.

    Returns:
        ``DsBaselineResult`` with counts, issues, paths, rows, groups and
        fingerprint. Reports are always written.
    """

    root = Path(source_root)
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp_value = stamp or datetime.now().strftime(STAMP_FORMAT)
    issues: list[DsBaselineIssue] = _copy_registry_issues(registry)
    if phase_callback is not None:
        phase_callback("обход папки ДС")
    files, skipped = collect_ds_workbooks(root, registry_path=registry.path)
    if files_discovered_callback is not None:
        files_discovered_callback(len(files), len(skipped))
    if phase_callback is not None:
        phase_callback(
            f"список ДС: {len(files)} файлов, пропущено {len(skipped)}"
        )
    active_ids = _active_ids(registry)
    registry_by_id = {row.source_id: row for row in registry.active_rows if row.source_id}

    sheet_scans: list[tuple[str, DsSheetScan]] = []
    column_notes: list[dict[str, object]] = []
    positions: list[DsBaselinePosition] = []
    total = len(files)
    for index, source in enumerate(files, start=1):
        relpath = str(source.relpath)
        if progress_callback is not None:
            progress_callback(index, total, relpath, None)
        file_start = time.perf_counter()
        positions.extend(
            _parse_workbook(
                source,
                active_ids=active_ids,
                registry_by_id=registry_by_id,
                issues=issues,
                sheet_scans=sheet_scans,
                column_notes=column_notes,
            )
        )
        if progress_callback is not None:
            progress_callback(
                index, total, relpath, time.perf_counter() - file_start
            )

    done_convert = _phase_timer(phase_callback, "конвертация единиц измерения")
    if google_index is not None:
        known = set(google_index.units_by_code)
        for item in positions:
            if item.code_normalized and item.code_normalized not in known:
                issues.append(
                    _issue(
                        ISSUE_UNKNOWN_GOOGLE,
                        "WARN",
                        f"кода {item.code!r} нет в Google",
                        path=item.path,
                        relpath=item.relpath,
                        sheet=item.sheet,
                        excel_row=item.excel_row,
                        source_id=item.source_id,
                        group_id=item.group_id,
                    )
                )

    if converter is None:
        converter = RfQDsUnitsConverter(
            google_index=google_index, matrix_path=matrix_path
        )
    conversion_plan: ConversionPlan | None = None
    try:
        batch = converter.convert_positions(positions)
        positions = list(batch.converted)
        conversion_plan = batch.plan
        for warning in batch.warnings:
            issues.append(
                _issue(ISSUE_UNITS_CONVERT, "WARN", warning, blocking=False)
            )
    except UnitsConversionError as exc:
        issues.append(
            _issue(
                ISSUE_UNITS_CONVERT,
                "ERROR",
                str(exc),
            )
        )
    done_convert(f"{len(positions)} позиций")

    done_tags = _phase_timer(phase_callback, "дубли тегов между строками")
    cross_tags: dict[str, list[DsBaselinePosition]] = defaultdict(list)
    for item in positions:
        seen: set[str] = set()
        for tag in item.diagnostic_tags:
            key = tag.casefold()
            if key in seen:
                continue
            seen.add(key)
            cross_tags[key].append(item)
    shared_tag_count = 0
    for tag_key, owners in cross_tags.items():
        if len(owners) < 2:
            continue
        shared_tag_count += 1
        sample = owners[0]
        tag_text = _display_tag(owners, tag_key)
        issues.append(
            _issue(
                ISSUE_TAG_DUPLICATE,
                "WARN",
                f"дубль тега {tag_text!r} в {len(owners)} строках",
                path=sample.path,
                relpath=sample.relpath,
                sheet=sample.sheet,
                excel_row=sample.excel_row,
                source_id=sample.source_id,
                group_id=sample.group_id,
            )
        )
    done_tags(f"{shared_tag_count} уникальных")

    done_groups = _phase_timer(phase_callback, "группы поставки")
    groups = _build_groups(positions)
    fingerprint = _tree_fingerprint(files, registry=registry)
    blocking_count = sum(1 for item in issues if item.blocking)
    warn_count = sum(1 for item in issues if item.level == "WARN")
    overlay_count = sum(1 for item in issues if item.level == "OVERLAY")
    blocking = blocking_count > 0
    done_groups(f"{len(groups)} групп")

    structure_path = out_dir / f"{STRUCTURE_REPORT_PREFIX}_{stamp_value}.xlsx"
    quality_path = out_dir / f"{QUALITY_REPORT_PREFIX}_{stamp_value}.xlsx"
    summary_rows: list[tuple[str, object]] = [
        ("Корень ДС", root),
        ("Реестр", registry.path),
        ("Папка отчётов", out_dir),
        ("Алгоритм", ALGORITHM_VERSION),
        ("Штамп", stamp_value),
        ("Файлов ДС", len(files)),
        ("Пропущено", len(skipped)),
        ("Позиций", len(positions)),
        ("Групп поставки", len(groups)),
        ("ERROR (блокеры)", blocking_count),
        ("WARN", warn_count),
        ("OVERLAY", overlay_count),
        ("Baseline записан", "нет" if blocking or not write_baseline else "да"),
        ("Fingerprint", fingerprint),
        ("Сводка реестра", registry.validation.summary_line()),
    ]
    _emit_phase(
        phase_callback,
        f"отчёт по структуре ({len(files)} файлов, {len(issues)} замечаний, "
        f"{len(positions)} позиций)",
    )
    _write_structure_report(
        structure_path,
        result_head=summary_rows,
        files=files,
        skipped=skipped,
        sheet_scans=sheet_scans,
        column_notes=column_notes,
        issues=issues,
        progress=phase_callback,
    )
    _emit_phase(
        phase_callback,
        f"отчёт по качеству ({len(positions)} позиций)",
    )
    _write_quality_report(
        quality_path,
        result_head=summary_rows,
        positions=positions,
        issues=issues,
        files=files,
        progress=phase_callback,
    )

    empty_code_path = None
    empty_code_rows = [item for item in positions if not item.code_normalized]
    if empty_code_rows:
        empty_code_path = out_dir / f"{EMPTY_CODE_REPORT_PREFIX}_{stamp_value}.xlsx"
        _write_sidecar_positions(
            empty_code_path,
            "Без кода",
            [
                "Файл",
                "Путь",
                "Лист",
                "Строка",
                "ID ДС",
                "Группа",
                "Наименование",
                "Кол-во",
                "Ед. изм.",
            ],
            [
                [
                    item.relpath,
                    str(item.path),
                    item.sheet,
                    item.excel_row,
                    item.source_id,
                    item.group_label,
                    item.name,
                    item.qty_raw,
                    item.units,
                ]
                for item in empty_code_rows
            ],
            path_cols=(1, 2),
            progress=phase_callback,
        )

    duplicate_tags_path = None
    dup_tag_issues = [item for item in issues if item.code == ISSUE_TAG_DUPLICATE]
    if dup_tag_issues:
        duplicate_tags_path = (
            out_dir / f"{DUPLICATE_TAGS_REPORT_PREFIX}_{stamp_value}.xlsx"
        )
        _write_sidecar_positions(
            duplicate_tags_path,
            "Дубли тегов",
            ["Файл", "Путь", "Лист", "Строка", "ID ДС", "Сообщение"],
            [
                [
                    item.relpath,
                    str(item.path) if item.path else "",
                    item.sheet,
                    item.excel_row,
                    item.source_id,
                    item.message,
                ]
                for item in dup_tag_issues
            ],
            path_cols=(1, 2),
            progress=phase_callback,
        )

    tag_mismatch_path = None
    mismatch_issues = [item for item in issues if item.code == ISSUE_TAG_MISMATCH]
    if mismatch_issues:
        tag_mismatch_path = (
            out_dir / f"{TAG_MISMATCH_REPORT_PREFIX}_{stamp_value}.xlsx"
        )
        _write_sidecar_positions(
            tag_mismatch_path,
            "Несоответствие тегов",
            ["Файл", "Путь", "Лист", "Строка", "ID ДС", "Сообщение"],
            [
                [
                    item.relpath,
                    str(item.path) if item.path else "",
                    item.sheet,
                    item.excel_row,
                    item.source_id,
                    item.message,
                ]
                for item in mismatch_issues
            ],
            path_cols=(1, 2),
            progress=phase_callback,
        )

    fractional_paths: tuple[Path, ...] = ()
    if conversion_plan is not None and conversion_plan.fractional_issues:
        fractional_paths = tuple(
            write_fractional_conversion_logs(conversion_plan, out_dir)
        )

    baseline_path = None
    if write_baseline and not blocking:
        done_baseline = _phase_timer(phase_callback, f"запись {BASELINE_XLSX_NAME}")
        target = out_dir / BASELINE_XLSX_NAME
        tmp_path = target.with_name(
            f".{target.stem}.{os.getpid()}.tmp{target.suffix}"
        )
        net_rows = _positions_to_net_rows(positions)
        try:
            _write_net_xlsx(
                tmp_path, net_rows, equipment_by_code=equipment_by_code
            )
            os.replace(tmp_path, target)
            baseline_path = target
        except Exception:
            tmp_path.unlink(missing_ok=True)
            raise
        done_baseline(f"{len(net_rows)} строк")

    return DsBaselineResult(
        source_root=root,
        output_dir=out_dir,
        registry_path=registry.path,
        stamp=stamp_value,
        fingerprint=fingerprint,
        files=list(files),
        skipped=list(skipped),
        positions=positions,
        groups=groups,
        issues=issues,
        file_count=len(files),
        position_count=len(positions),
        group_count=len(groups),
        blocking_issue_count=blocking_count,
        warn_count=warn_count,
        overlay_count=overlay_count,
        blocking=blocking,
        baseline_path=baseline_path,
        structure_report_path=structure_path,
        quality_report_path=quality_path,
        empty_code_report_path=empty_code_path,
        duplicate_tags_report_path=duplicate_tags_path,
        tag_mismatch_report_path=tag_mismatch_path,
        fractional_log_paths=fractional_paths,
    )
