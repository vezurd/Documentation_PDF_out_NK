"""Count equipment tags and positions in each root RFP workbook.

A DS workbook pasted into the RFP folder by rearranging columns has positions
and no equipment tags. The census is that check: same header as the parts
reader, tags only when ``get_tag`` recognizes an equipment tag.
"""

from __future__ import annotations

import html
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from RFQ.rfp_parts.analyze_rfp_parts import (
    DEFAULT_REPORTS_BASE_DIR,
    make_reports_out_dir,
    _cell,
    _choose_part_sheet,
    _find_header_on_sheet,
    _is_total_row,
    _norm_text,
    _normalize_tags_text,
    _safe_iter_rows,
    resolve_parts_workbooks,
    split_part_sheets,
)
from tags.tag_parser import get_tag

_COPY_HEADERS = ("Имя файла", "Кол-во тегов", "Кол-во позиций", "Примечание")
CENSUS_DIR_NAME = "_теги_позиций_RFP"
CENSUS_XLSX_NAME = "Теги и позиции RFP.xlsx"
_SUSPECT_FILL = PatternFill(fill_type="solid", fgColor="FFF8C5")
_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(vertical="center", wrap_text=True)


@dataclass(frozen=True, slots=True)
class RfpTagCensusRow:
    """One workbook in the RFP root."""

    file_name: str
    tag_count: int
    position_count: int
    note: str = ""

    @property
    def suspect(self) -> bool:
        """Positions exist and no equipment tag was recognized."""
        return self.position_count > 0 and self.tag_count == 0


@dataclass(frozen=True, slots=True)
class RfpTagCensus:
    """Scan of the RFP root, suspects first."""

    rfp_root: Path
    rows: tuple[RfpTagCensusRow, ...]

    @property
    def suspect_count(self) -> int:
        return sum(1 for row in self.rows if row.suspect)

    def summary(self) -> str:
        """One line for the job monitor and the tab caption."""
        return (
            f"файлов {len(self.rows)}, без тегов при живых позициях "
            f"{self.suspect_count}"
        )

    def copy_tsv(self) -> str:
        """Tab-separated table for paste into a mail or Excel."""
        lines = ["\t".join(_COPY_HEADERS)]
        for row in self.rows:
            lines.append(
                "\t".join(
                    (
                        row.file_name,
                        str(row.tag_count),
                        str(row.position_count),
                        row.note,
                    )
                )
            )
        return "\n".join(lines)

    def copy_html(self) -> str:
        """HTML table so Outlook pastes columns, not one line of text."""
        head = "".join(f"<th>{html.escape(name)}</th>" for name in _COPY_HEADERS)
        body: list[str] = []
        for row in self.rows:
            cells = (
                row.file_name,
                str(row.tag_count),
                str(row.position_count),
                row.note,
            )
            body.append(
                "<tr>"
                + "".join(f"<td>{html.escape(value)}</td>" for value in cells)
                + "</tr>"
            )
        return (
            "<table border='1' cellspacing='0' cellpadding='4'>"
            f"<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"
        )


@dataclass(frozen=True, slots=True)
class RfpTagCensusJobResult:
    """Duck-typed ``FunctionJobRunner`` result."""

    success: bool
    message: str
    result_path: str | None = None


_last_census: RfpTagCensus | None = None


def get_last_rfp_tag_census() -> RfpTagCensus | None:
    """Return the census from the last successful scan, if any."""
    return _last_census


def _emit(message: str) -> None:
    print(message, flush=True)


def _count_positions_and_tags(ws: object, header: object) -> tuple[int, int, str]:
    """Return ``(positions, equipment tags, note)`` under one header."""
    columns = header.columns
    note = ""
    if "TAGS" not in columns:
        note = "нет колонки тегов"
    if "CODE" not in columns and "NAME" not in columns:
        return 0, 0, "нет колонок код/наименование"
    positions = 0
    tags = 0
    warnings: list[tuple[str, str, str]] = []
    for row_tuple in _safe_iter_rows(
        ws, warnings, header.sheet, min_row=header.row_number + 1
    ):
        row = list(row_tuple)
        if _is_total_row(row):
            continue
        code = _norm_text(_cell(row, columns, "CODE"))
        name = _norm_text(_cell(row, columns, "NAME"))
        if not code and not name:
            continue
        positions += 1
        raw_tags = _norm_text(_cell(row, columns, "TAGS"))
        if not raw_tags:
            continue
        tags += len(get_tag(_normalize_tags_text(raw_tags)))
    return positions, tags, note


def _scan_workbook(path: Path) -> RfpTagCensusRow:
    try:
        wb = load_workbook(path, read_only=True, data_only=True, rich_text=False)
    except Exception as exc:
        return RfpTagCensusRow(
            path.name, 0, 0, f"не открыть: {type(exc).__name__}: {exc}"
        )
    try:
        _skipped, candidates = split_part_sheets(wb.sheetnames)
        if not candidates:
            return RfpTagCensusRow(path.name, 0, 0, "нет рабочего листа")
        warnings: list[tuple[str, str, str]] = []
        probes = [
            (
                name,
                _find_header_on_sheet(
                    wb[name], warnings, path.name, prefer_lot_qty=True
                ),
            )
            for name in candidates
        ]
        _chosen, header = _choose_part_sheet(candidates, probes)
        if header is None:
            return RfpTagCensusRow(path.name, 0, 0, "шапка RFP не найдена")
        positions, tags, note = _count_positions_and_tags(wb[header.sheet], header)
        if positions == 0 and not note:
            note = "позиций нет"
        return RfpTagCensusRow(path.name, tags, positions, note)
    finally:
        wb.close()


def _sort_key(row: RfpTagCensusRow) -> tuple[int, int, int, str]:
    if row.suspect:
        band = 0
    elif row.position_count == 0 and row.note:
        band = 1
    else:
        band = 2
    return (band, row.tag_count, -row.position_count, row.file_name.lower())


def scan_rfp_tag_census(rfp_root: str | Path) -> RfpTagCensus:
    """Scan direct ``.xlsx`` / ``.xlsm`` children of the RFP root.

    Args:
        rfp_root: Folder ``RFP_Зиновьев``. Subfolders are not read.

    Returns:
        Rows ordered so files with positions and zero equipment tags come first.

    Raises:
        FileNotFoundError: The folder is missing or has no workbooks.
        NotADirectoryError: ``rfp_root`` is not a directory.
    """
    root = Path(rfp_root)
    files = resolve_parts_workbooks(root)
    total = len(files)
    rows: list[RfpTagCensusRow] = []
    for index, path in enumerate(files, start=1):
        if index == 1 or index == total or index % 10 == 0:
            _emit(f"[rfp tags] {index}/{total} {path.name}")
        rows.append(_scan_workbook(path))
    rows.sort(key=_sort_key)
    return RfpTagCensus(rfp_root=root, rows=tuple(rows))


def census_reports_dir() -> Path:
    """Folder ``RFP сводный файл/_теги_позиций_RFP`` (not a parts stamp)."""
    return DEFAULT_REPORTS_BASE_DIR / CENSUS_DIR_NAME


def write_rfp_tag_census_xlsx(census: RfpTagCensus, dest: Path) -> Path:
    """Write the census table. Suspect rows (positions, no tags) are yellow.

    Args:
        census: Scan result.
        dest: Target ``.xlsx`` path. Parent folders are created.

    Returns:
        ``dest``.
    """
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Теги и позиции"
    ws.append(list(_COPY_HEADERS))
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
    for row in census.rows:
        ws.append([row.file_name, row.tag_count, row.position_count, row.note])
        if row.suspect:
            for cell in ws[ws.max_row]:
                cell.fill = _SUSPECT_FILL
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 36
    wb.save(dest)
    wb.close()
    return dest


def run_rfp_tag_census_job(rfp_root: str | Path) -> RfpTagCensusJobResult:
    """Scan the RFP root, print a copyable table, and save an xlsx.

    Args:
        rfp_root: Folder ``RFP_Зиновьев``.

    Returns:
        Success when the folder was read and the workbook was saved.
        The table stays in :func:`get_last_rfp_tag_census`. ``result_path``
        is the xlsx under ``_теги_позиций_RFP``.
    """
    global _last_census
    root = Path(rfp_root)
    if not root.is_dir():
        return RfpTagCensusJobResult(
            False, f"Папка RFP не найдена: {root}", None
        )
    try:
        census = scan_rfp_tag_census(root)
    except (FileNotFoundError, NotADirectoryError, OSError) as exc:
        return RfpTagCensusJobResult(False, str(exc), None)
    _last_census = census
    _emit(census.summary())
    _emit(census.copy_tsv())
    dest = make_reports_out_dir(base=census_reports_dir()) / CENSUS_XLSX_NAME
    try:
        write_rfp_tag_census_xlsx(census, dest)
    except OSError as exc:
        return RfpTagCensusJobResult(
            False,
            f"{census.summary()}. xlsx не записан: {exc}",
            None,
        )
    _emit(str(dest))
    return RfpTagCensusJobResult(True, f"{census.summary()}. {dest.name}", str(dest))
