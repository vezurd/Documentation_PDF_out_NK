"""Group-level DS↔RFP reconciliation and atomic hybrid overlay.

RFP is read from the folder root only (no recursion). The hybrid workbook
``Свод ДС-RFP для запуска.xlsx`` is written only when the DS baseline has
zero blockers and this stage has zero global blockers. The Russian audit
report is always written. Overlay is atomic per supply group: MATCH takes
every RFP row of the group; any other status takes every DS baseline row.
RFP-only groups are reported and excluded. Never mix rows inside a group.
"""

from __future__ import annotations

import hashlib
import os
import time
from collections import defaultdict
from collections.abc import Callable
from dataclasses import dataclass, field, replace
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Literal, Protocol, Sequence

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from RFQ.ds_compare.ds_units_normalize import normalize_units_text
from RFQ.rfp_parts.analyze_rfp_parts import (
    PARTS_KIND,
    FileStats,
    NetSummaryRow,
    RfpRecord,
    _apply_parts_conversion_plan,
    _build_parts_conversion_requests,
    _build_summary_rows,
    _build_unit_counters,
    _compute_coarse_collisions,
    _extract_records,
    _write_net_xlsx,
)
from RFQ.rfp_parts.ds_baseline import (
    BASELINE_XLSX_NAME,
    _AUDIT_DIR_RE,
    DsBaselineIssue,
    DsBaselinePosition,
    DsBaselineResult,
    DsSkippedFile,
    DsSourceFile,
    _FILL_ERROR,
    _FILL_WARN,
    _file_sha256,
    _identity_trace,
    _mark_level,
    _positions_to_net_rows,
    _save_workbook_atomic,
    _set_path_cell,
    _style_header,
    _write_summary_sheet,
    make_audit_stamp_dir,
)
from RFQ.rfp_parts.ds_identity import DsIdentity, parse_rfp_ds_identity
from RFQ.rfp_parts.ds_registry import DsRegistryDocument
from RFQ.tags_rfp_compare.rfp_supply_status import (
    RfpSupplyStatusIssue,
    is_canonical_excluded_from_supply,
)
from RFQ.units_convert.fractional_log import write_fractional_conversion_logs
from RFQ.units_convert.gate import build_conversion_plan
from RFQ.units_convert.models import (
    STATUS_CONVERTED,
    STATUS_IDENTITY,
    ConversionPlan,
    GoogleUnitsIndex,
    UnitsConversionError,
    normalize_code,
)

IssueLevel = Literal["ERROR", "WARN", "OVERLAY"]
GroupStatus = Literal["MATCH", "MISMATCH", "DS_ONLY", "RFP_ONLY", "BLOCKED"]
SelectedSource = Literal["RFP", "ДС", ""]

ALGORITHM_VERSION = "ds_rfp_hybrid_v1"
HYBRID_XLSX_NAME = "Свод ДС-RFP для запуска.xlsx"
HYBRID_REPORT_PREFIX = "Отчет по сверке ДС-RFP-УЛ"
REPORT_XLSX_NAME = HYBRID_REPORT_PREFIX
STAMP_FORMAT = "%Y%m%d_%H%M%S"
SOURCE_RFP = "RFP overlay"
SOURCE_DS = "ДС baseline"

STATUS_MATCH: GroupStatus = "MATCH"
STATUS_MISMATCH: GroupStatus = "MISMATCH"
STATUS_DS_ONLY: GroupStatus = "DS_ONLY"
STATUS_RFP_ONLY: GroupStatus = "RFP_ONLY"
STATUS_BLOCKED: GroupStatus = "BLOCKED"

ISSUE_UNPARSED_RFP = "unparsed_rfp"
ISSUE_DUPLICATE_RFP_KEY = "duplicate_rfp_key"
ISSUE_AMBIGUOUS_MAPPING = "ambiguous_rfp_mapping"
ISSUE_RFP_EXTRACT = "rfp_extract_error"
ISSUE_UNITS_CONVERT = "units_conversion_error"
ISSUE_OVERLAY_BLOCKED = "overlay_blocked"
ISSUE_GROUP_FALLBACK = "group_fallback"
ISSUE_BASELINE_BLOCKED = "baseline_blocked"
ISSUE_RFP_ROOT = "rfp_root_missing"

_EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm"})
_OK_CONVERSION = frozenset({STATUS_IDENTITY, STATUS_CONVERTED, ""})

_SKIP_EXACT_NAMES = frozenset(
    {
        HYBRID_XLSX_NAME.casefold(),
        f"{HYBRID_REPORT_PREFIX}.xlsx".casefold(),
        "отчет по сверке дс-rfp.xlsx",
        BASELINE_XLSX_NAME.casefold(),
        "rfp_parts_net.xlsx",
        "rfp_parts_net_no_tags.xlsx",
        "rfp_parts_collisions.xlsx",
        "отчет по тегам - сбор частей.xlsx",
        "отчет по позициям без кода - сбор частей.xlsx",
        "дробные значения после конвертации rfp.xlsx",
        "дробные значения после конвертации дс.xlsx",
    }
)
_SKIP_PREFIXES = (
    "отчет по сверке дс-rfp",
    "отчет по структуре файлов - дс",
    "отчет по качеству данных - дс",
    "отчет по позициям без кода - дс",
    "отчет по дублям тегов - дс",
    "отчет по несоответствию тегов - дс",
    "rfp_parts_net",
    "свод дс",
    "свод дс-rfp",
)

_FILL_MATCH = PatternFill(fill_type="solid", fgColor="C6EFCE")
_FILL_DS_ONLY = PatternFill(fill_type="solid", fgColor="D9E2F3")
_FILL_RFP_ONLY = PatternFill(fill_type="solid", fgColor="D9D9D9")
_FILL_MISMATCH = _FILL_WARN
_FILL_BLOCKED = _FILL_ERROR
_ZERO = Decimal("0")


def hybrid_report_filename(stamp: str) -> str:
    """Return ``Отчет по сверке ДС-RFP-УЛ_<штамп>.xlsx``."""

    return f"{HYBRID_REPORT_PREFIX}_{stamp}.xlsx"


def find_latest_hybrid_report(directory: str | Path | None) -> Path | None:
    """Newest hybrid audit workbook in ``directory``, or None."""

    if directory is None:
        return None
    root = Path(directory)
    prefix = HYBRID_REPORT_PREFIX.casefold()
    matches: list[Path] = []
    scan_dirs = [root]
    try:
        for child in root.iterdir():
            if child.is_dir() and _AUDIT_DIR_RE.match(child.name):
                scan_dirs.append(child)
    except OSError:
        return None
    try:
        for folder in scan_dirs:
            for path in folder.iterdir():
                if not path.is_file():
                    continue
                name = path.name
                if name.startswith("~$"):
                    continue
                folded = name.casefold()
                if folded.startswith(prefix) and folded.endswith(".xlsx"):
                    matches.append(path)
    except OSError:
        return None
    if not matches:
        return None
    return max(matches, key=lambda item: item.stat().st_mtime)


# ---------------------------------------------------------------------------
# Public dataclasses
# ---------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class DsRfpHybridIssue:
    """One mapping / extract / conversion / overlay finding."""

    code: str
    level: IssueLevel
    message: str
    blocking: bool = False
    global_blocker: bool = False
    path: Path | None = None
    relpath: str = ""
    excel_row: int | None = None
    rfp_key: str = ""
    group_id: str = ""


@dataclass(frozen=True, slots=True)
class DsRfpCodeDelta:
    """One (code, unit) bag comparison inside a supply group."""

    group_id: str
    code: str
    unit: str
    ds_qty: Decimal
    rfp_qty: Decimal
    delta: Decimal
    kind: str


@dataclass(frozen=True, slots=True)
class DsRfpGroupResult:
    """Reconciliation verdict for one supply group."""

    group_id: str
    group_label: str
    status: GroupStatus
    selected_source: SelectedSource
    reason: str
    overlay_blocked: bool
    fallback: bool
    ds_position_count: int
    rfp_record_count: int
    rfp_files: tuple[str, ...]
    ds_source_ids: tuple[str, ...]
    deltas: tuple[DsRfpCodeDelta, ...] = ()


@dataclass
class RfpFileExtract:
    """One root RFP workbook after identity, extract and conversion."""

    source: DsSourceFile
    identity: DsIdentity
    rfp_key: str
    group_id: str
    records: list[RfpRecord] = field(default_factory=list)
    warnings: list[tuple[str, str, str]] = field(default_factory=list)
    status_issues: list[RfpSupplyStatusIssue] = field(default_factory=list)
    stats: FileStats | None = None
    extract_error: bool = False
    mapping_issue: str = ""


@dataclass
class RfpConversionBatch:
    """Converter output for RFP records."""

    converted: tuple[RfpRecord, ...]
    plan: ConversionPlan | None = None
    warnings: tuple[str, ...] = ()


@dataclass
class DsRfpHybridResult:
    """Public result of ``build_ds_rfp_hybrid``."""

    rfp_root: Path
    output_dir: Path
    registry_path: Path
    stamp: str
    fingerprint: str
    algorithm_version: str = ALGORITHM_VERSION
    files: list[DsSourceFile] = field(default_factory=list)
    skipped: list[DsSkippedFile] = field(default_factory=list)
    extracts: list[RfpFileExtract] = field(default_factory=list)
    groups: list[DsRfpGroupResult] = field(default_factory=list)
    issues: list[DsRfpHybridIssue] = field(default_factory=list)
    hybrid_rows: list[NetSummaryRow] = field(default_factory=list)
    file_count: int = 0
    group_count: int = 0
    match_count: int = 0
    mismatch_count: int = 0
    ds_only_count: int = 0
    rfp_only_count: int = 0
    blocked_count: int = 0
    blocking_issue_count: int = 0
    global_blocker_count: int = 0
    blocking: bool = False
    hybrid_path: Path | None = None
    report_path: Path | None = None
    fractional_log_paths: tuple[Path, ...] = ()

    def summary_line(self) -> str:
        if self.hybrid_path is not None:
            return (
                f"OK: гибрид ДС-RFP готов, MATCH={self.match_count}, "
                f"MISMATCH={self.mismatch_count}, DS_ONLY={self.ds_only_count}, "
                f"RFP_ONLY={self.rfp_only_count}, BLOCKED={self.blocked_count}"
            )
        return (
            f"BLOCKED: hybrid не записан, GLOBAL={self.global_blocker_count}, "
            f"ERROR={self.blocking_issue_count}, MATCH={self.match_count}, "
            f"BLOCKED_групп={self.blocked_count}"
        )


class RfpWorkbookLoader(Protocol):
    """Injectable RFP workbook reader. Default uses production extract."""

    def extract(self, path: Path) -> tuple[list[RfpRecord], FileStats, list[tuple[str, str, str]], list[RfpSupplyStatusIssue]]:
        """Return records, stats, warnings and supply-status issues for one file."""


class RfpUnitsConverter(Protocol):
    """Injectable units conversion. Default talks to RFQ.units_convert APIs."""

    def convert_records(self, records: Sequence[RfpRecord]) -> RfpConversionBatch:
        """Return converted copies; must not fake a matrix/Google result."""


# ---------------------------------------------------------------------------
# Default loader / converter
# ---------------------------------------------------------------------------


class ProductionRfpLoader:
    """Production parts extractor: ``_extract_records(path, "parts", ...)``."""

    def extract(
        self, path: Path
    ) -> tuple[
        list[RfpRecord],
        FileStats,
        list[tuple[str, str, str]],
        list[RfpSupplyStatusIssue],
    ]:
        warnings: list[tuple[str, str, str]] = []
        status_issues: list[RfpSupplyStatusIssue] = []
        records, stats = _extract_records(
            path,
            PARTS_KIND,
            warnings,
            status_issues=status_issues,
        )
        return records, stats, warnings, status_issues


class IdentityRfpUnitsConverter:
    """Offline converter: identity qty/unit, status/trace preserved, no I/O."""

    def convert_records(self, records: Sequence[RfpRecord]) -> RfpConversionBatch:
        converted: list[RfpRecord] = []
        for item in records:
            converted.append(
                replace(
                    item,
                    units_check_status=STATUS_IDENTITY,
                    units_conversion_trace=_identity_trace(
                        code=item.code,
                        unit=item.units,
                        qty=item.values,
                    ),
                )
            )
        return RfpConversionBatch(converted=tuple(converted), plan=None)


class PartsRfpUnitsConverter:
    """Default adapter: same plan/matrix/Google contour as parts analyze."""

    def __init__(
        self,
        *,
        google_index: GoogleUnitsIndex | None = None,
        matrix_path: str | Path | None = None,
    ) -> None:
        self.google_index = google_index
        self.matrix_path = Path(matrix_path) if matrix_path else None

    def convert_records(self, records: Sequence[RfpRecord]) -> RfpConversionBatch:
        if not records:
            return IdentityRfpUnitsConverter().convert_records(records)
        if self.google_index is None or self.matrix_path is None:
            raise UnitsConversionError(
                "конвертация RFP требует google_index и matrix_path; "
                "для offline-тестов передайте IdentityRfpUnitsConverter"
            )
        requests = _build_parts_conversion_requests(list(records))
        plan = build_conversion_plan(
            requests, self.google_index, self.matrix_path
        )
        converted = _apply_parts_conversion_plan(list(records), plan)
        return RfpConversionBatch(
            converted=tuple(converted),
            plan=plan,
            warnings=plan.warnings,
        )


# ---------------------------------------------------------------------------
# Collect
# ---------------------------------------------------------------------------


def _skip_reason(path: Path) -> str | None:
    name = path.name
    if name.startswith("~$"):
        return "excel_lock"
    folded = name.casefold()
    stem = path.stem.casefold()
    if folded in _SKIP_EXACT_NAMES:
        return "own_output"
    for prefix in _SKIP_PREFIXES:
        if stem.startswith(prefix) or folded.startswith(prefix):
            return "own_or_legacy_or_report"
    return None


def collect_rfp_workbooks(
    rfp_root: str | Path,
) -> tuple[list[DsSourceFile], list[DsSkippedFile]]:
    """Collect root-only RFP xlsx/xlsm, skipping locks, own outputs and nets.

    Args:
        rfp_root: Folder whose *direct* children are scanned. Subfolders are
            ignored (no recursion).

    Returns:
        ``(files, skipped)`` sorted by casefold name.
    """

    root = Path(rfp_root)
    collected: list[DsSourceFile] = []
    skipped: list[DsSkippedFile] = []
    if not root.is_dir():
        return collected, skipped

    candidates = [
        path
        for path in root.iterdir()
        if path.is_file() and path.suffix.lower() in _EXCEL_SUFFIXES
    ]
    candidates.sort(key=lambda item: item.name.casefold())
    for path in candidates:
        relpath = path.name
        reason = _skip_reason(path)
        if reason:
            skipped.append(DsSkippedFile(path=path, relpath=relpath, reason=reason))
            continue
        try:
            data = path.read_bytes()
            stat = path.stat()
        except OSError:
            skipped.append(
                DsSkippedFile(path=path, relpath=relpath, reason="unreadable")
            )
            continue
        collected.append(
            DsSourceFile(
                path=path,
                relpath=relpath,
                size=stat.st_size,
                mtime_ns=stat.st_mtime_ns,
                fingerprint=_file_sha256(data),
            )
        )
    return collected, skipped


# ---------------------------------------------------------------------------
# Mapping / issues
# ---------------------------------------------------------------------------


def _issue(
    code: str,
    level: IssueLevel,
    message: str,
    *,
    blocking: bool | None = None,
    global_blocker: bool = False,
    path: Path | None = None,
    relpath: str = "",
    excel_row: int | None = None,
    rfp_key: str = "",
    group_id: str = "",
) -> DsRfpHybridIssue:
    is_blocking = (level == "ERROR") if blocking is None else blocking
    return DsRfpHybridIssue(
        code=code,
        level=level,
        message=message,
        blocking=is_blocking,
        global_blocker=global_blocker,
        path=path,
        relpath=relpath,
        excel_row=excel_row,
        rfp_key=rfp_key,
        group_id=group_id,
    )


def rfp_identity_key(identity: DsIdentity) -> str | None:
    """Return the registry ``rfp_key`` (bare actual) for a parsed RFP name."""

    if identity.kind == "unparsed" or identity.actual is None:
        return None
    return str(identity.actual)


def _rfp_key_groups(registry: DsRegistryDocument) -> dict[str, tuple[str, ...]]:
    """Map an RFP number to the actual-DS group that named it.

    The group is ``ДС{актуальный номер}``, not the folder name and not the
    number parsed from the file when the registry names a different number.
    """

    from RFQ.rfp_parts.ds_registry import canonical_supply_group_id

    mapping: dict[str, list[str]] = {}
    for row in registry.active_rows:
        if not row.source_id:
            continue
        group_id = canonical_supply_group_id(row.source_id)
        for rel in row.relations:
            if not rel.rfp_key:
                continue
            groups = mapping.setdefault(rel.rfp_key, [])
            if group_id not in groups:
                groups.append(group_id)
    return {key: tuple(groups) for key, groups in mapping.items()}


def _rfp_file_groups(registry: DsRegistryDocument) -> dict[str, tuple[str, ...]]:
    """Map a registry RFP file name to the actual-DS group that named it."""

    from RFQ.rfp_parts.ds_registry import _ds_file_names, canonical_supply_group_id

    mapping: dict[str, list[str]] = {}
    for row in registry.active_rows:
        if not row.source_id:
            continue
        group_id = canonical_supply_group_id(row.source_id)
        for rel in row.relations:
            for name in _ds_file_names(rel.rfp_file):
                groups = mapping.setdefault(name.casefold(), [])
                if group_id not in groups:
                    groups.append(group_id)
    return {key: tuple(groups) for key, groups in mapping.items()}


def _qty_excel(value: Decimal | None) -> object:
    if value is None:
        return ""
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _compare_key(code: object, unit: object) -> tuple[str, str]:
    return (normalize_code(code), normalize_units_text(unit))


def _add_qty(
    bag: dict[tuple[str, str], Decimal],
    code: object,
    unit: object,
    qty: Decimal | None,
) -> None:
    if qty is None:
        return
    key = _compare_key(code, unit)
    if not key[0]:
        return
    bag[key] = bag.get(key, _ZERO) + qty


def _nonzero(bag: dict[tuple[str, str], Decimal]) -> dict[tuple[str, str], Decimal]:
    return {key: value for key, value in bag.items() if value != 0}


def _ds_bag(positions: Sequence[DsBaselinePosition]) -> dict[tuple[str, str], Decimal]:
    bag: dict[tuple[str, str], Decimal] = {}
    for item in positions:
        _add_qty(bag, item.code or item.code_normalized, item.units, item.qty)
    return bag


def _rfp_compare_bag(records: Sequence[RfpRecord]) -> dict[tuple[str, str], Decimal]:
    bag: dict[tuple[str, str], Decimal] = {}
    for item in records:
        if is_canonical_excluded_from_supply(item.rfp_supply_status):
            continue
        _add_qty(bag, item.code, item.units, item.values)
    return bag


def _build_deltas(
    group_id: str,
    ds_bag: dict[tuple[str, str], Decimal],
    rfp_bag: dict[tuple[str, str], Decimal],
) -> tuple[DsRfpCodeDelta, ...]:
    keys = sorted(set(ds_bag) | set(rfp_bag))
    rows: list[DsRfpCodeDelta] = []
    for code, unit in keys:
        ds_qty = ds_bag.get((code, unit), _ZERO)
        rfp_qty = rfp_bag.get((code, unit), _ZERO)
        if ds_qty == rfp_qty:
            kind = "equal"
        elif rfp_qty == 0 and ds_qty != 0:
            kind = "missing"
        elif ds_qty == 0 and rfp_qty != 0:
            kind = "excess"
        else:
            kind = "delta"
        rows.append(
            DsRfpCodeDelta(
                group_id=group_id,
                code=code,
                unit=unit,
                ds_qty=ds_qty,
                rfp_qty=rfp_qty,
                delta=rfp_qty - ds_qty,
                kind=kind,
            )
        )
    return tuple(rows)


def _conversion_ok(status: str) -> bool:
    return status in _OK_CONVERSION


def _file_extract_error(
    stats: FileStats,
    warnings: Sequence[tuple[str, str, str]],
    status_issues: Sequence[RfpSupplyStatusIssue],
) -> bool:
    if status_issues:
        return True
    if stats.warnings > 0:
        return True
    return any(level == "ERROR" for level, _file, _msg in warnings)


def _hybrid_fingerprint(
    *,
    baseline: DsBaselineResult,
    registry: DsRegistryDocument,
    files: Sequence[DsSourceFile],
) -> str:
    digest = hashlib.sha256()
    digest.update(ALGORITHM_VERSION.encode("utf-8"))
    digest.update(baseline.fingerprint.encode("utf-8"))
    digest.update(str(registry.path).encode("utf-8", errors="replace"))
    if registry.path.is_file():
        try:
            digest.update(registry.path.read_bytes())
        except OSError:
            pass
    for item in files:
        digest.update(
            f"{item.relpath}|{item.size}|{item.mtime_ns}|{item.fingerprint}".encode(
                "utf-8"
            )
        )
    return digest.hexdigest()


def _relabel_rfp(records: Sequence[RfpRecord], group_label: str) -> list[RfpRecord]:
    return [replace(item, ds_name=group_label) for item in records]


def _rfp_net_rows(records: Sequence[RfpRecord], group_label: str) -> list[NetSummaryRow]:
    relabeled = _relabel_rfp(records, group_label)
    if not relabeled:
        return []
    silent: list[tuple[str, str, str]] = []
    aggregation = _build_unit_counters(relabeled, silent, emit_warnings=False)
    coarse = _compute_coarse_collisions(relabeled)
    rows = _build_summary_rows(aggregation, coarse)
    for item in rows:
        item.source = SOURCE_RFP
        item.record = replace(item.record, ds_name=group_label)
    return rows


def _ds_net_rows(positions: Sequence[DsBaselinePosition]) -> list[NetSummaryRow]:
    rows = _positions_to_net_rows(positions)
    for item in rows:
        item.source = SOURCE_DS
    return rows


def _renumber(rows: Sequence[NetSummaryRow]) -> list[NetSummaryRow]:
    out: list[NetSummaryRow] = []
    for index, item in enumerate(rows, start=1):
        out.append(
            NetSummaryRow(
                record=replace(item.record, ds_number=str(index), kind="net"),
                source=item.source,
                flag_unmatched_decrease=item.flag_unmatched_decrease,
                flag_negative_net=item.flag_negative_net,
                flag_units_mismatch=item.flag_units_mismatch,
                flag_via_replacement=item.flag_via_replacement,
            )
        )
    return out


def _status_fill(status: str) -> PatternFill | None:
    if status == STATUS_MATCH:
        return _FILL_MATCH
    if status == STATUS_MISMATCH:
        return _FILL_MISMATCH
    if status == STATUS_BLOCKED:
        return _FILL_BLOCKED
    if status == STATUS_DS_ONLY:
        return _FILL_DS_ONLY
    if status == STATUS_RFP_ONLY:
        return _FILL_RFP_ONLY
    return None


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------


def _paint_status(cell, status: str) -> None:
    fill = _status_fill(status)
    if fill is not None:
        cell.fill = fill


def _write_hybrid_report(
    path: Path,
    *,
    result_head: Sequence[tuple[str, object]],
    groups: Sequence[DsRfpGroupResult],
    extracts: Sequence[RfpFileExtract],
    files: Sequence[DsSourceFile],
    skipped: Sequence[DsSkippedFile],
    issues: Sequence[DsRfpHybridIssue],
    baseline_issues: Sequence[DsBaselineIssue],
    hybrid_rows: Sequence[NetSummaryRow],
    baseline_files: Sequence[DsSourceFile],
) -> Path:
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    _write_summary_sheet(summary, result_head)

    groups_ws = wb.create_sheet("Группы поставки")
    groups_ws.append(
        [
            "Группа",
            "Ярлык",
            "Статус",
            "Выбранный источник",
            "Причина",
            "Overlay blocked",
            "Fallback",
            "Позиций ДС",
            "Записей RFP",
            "Файлы RFP",
            "ID ДС источника",
            "Кодов с расхождением",
        ]
    )
    for item in groups:
        mismatch_n = sum(1 for delta in item.deltas if delta.kind != "equal")
        groups_ws.append(
            [
                item.group_id,
                item.group_label,
                item.status,
                item.selected_source,
                item.reason,
                "да" if item.overlay_blocked else "",
                "да" if item.fallback else "",
                item.ds_position_count,
                item.rfp_record_count,
                "; ".join(item.rfp_files),
                "; ".join(item.ds_source_ids),
                mismatch_n,
            ]
        )
        _paint_status(groups_ws.cell(row=groups_ws.max_row, column=3), item.status)
    _style_header(groups_ws, 12)
    groups_ws.column_dimensions["A"].width = 18
    groups_ws.column_dimensions["E"].width = 70
    groups_ws.column_dimensions["J"].width = 36

    delta_ws = wb.create_sheet("Разница по кодам")
    delta_ws.append(
        [
            "Группа",
            "Код",
            "Ед. изм.",
            "Кол-во ДС",
            "Кол-во RFP",
            "Дельта (RFP−ДС)",
            "Вид",
            "Статус группы",
        ]
    )
    status_by_group = {item.group_id: item.status for item in groups}
    for item in groups:
        for delta in item.deltas:
            delta_ws.append(
                [
                    delta.group_id,
                    delta.code,
                    delta.unit,
                    _qty_excel(delta.ds_qty),
                    _qty_excel(delta.rfp_qty),
                    _qty_excel(delta.delta),
                    delta.kind,
                    status_by_group.get(delta.group_id, ""),
                ]
            )
            if delta.kind != "equal":
                _mark_level(delta_ws.cell(row=delta_ws.max_row, column=7), "WARN")
    _style_header(delta_ws, 8)
    delta_ws.column_dimensions["B"].width = 18
    delta_ws.column_dimensions["G"].width = 14

    cover = wb.create_sheet("Покрытие файлов")
    cover.append(
        [
            "Контур",
            "Файл",
            "Путь",
            "Ключ RFP",
            "Группа",
            "Kind identity",
            "Записей",
            "Ошибка извлечения",
            "Причина пропуска / mapping",
            "fingerprint",
        ]
    )
    extract_by_name = {item.source.relpath: item for item in extracts}
    for item in files:
        extracted = extract_by_name.get(item.relpath)
        cover.append(
            [
                "RFP",
                item.relpath,
                str(item.path),
                extracted.rfp_key if extracted else "",
                extracted.group_id if extracted else "",
                extracted.identity.kind if extracted else "",
                len(extracted.records) if extracted else 0,
                "да" if extracted is not None and extracted.extract_error else "",
                extracted.mapping_issue if extracted else "",
                item.fingerprint,
            ]
        )
        _set_path_cell(cover.cell(row=cover.max_row, column=2), item.path)
        _set_path_cell(cover.cell(row=cover.max_row, column=3), item.path)
    for item in skipped:
        cover.append(
            [
                "RFP",
                item.relpath,
                str(item.path),
                "",
                "",
                "",
                0,
                "",
                item.reason,
                "",
            ]
        )
        _set_path_cell(cover.cell(row=cover.max_row, column=2), item.path)
        _set_path_cell(cover.cell(row=cover.max_row, column=3), item.path)
    for item in baseline_files:
        cover.append(
            [
                "ДС",
                item.relpath,
                str(item.path),
                "",
                "",
                "",
                "",
                "",
                "",
                item.fingerprint,
            ]
        )
        _set_path_cell(cover.cell(row=cover.max_row, column=2), item.path)
        _set_path_cell(cover.cell(row=cover.max_row, column=3), item.path)
    _style_header(cover, 10)
    cover.column_dimensions["B"].width = 40
    cover.column_dimensions["C"].width = 70
    cover.column_dimensions["I"].width = 40

    problems = wb.create_sheet("Проблемы")
    problems.append(
        [
            "Уровень",
            "Код",
            "Контур",
            "Файл",
            "Путь",
            "Строка",
            "Ключ RFP",
            "Группа",
            "Global",
            "Сообщение",
        ]
    )
    for item in issues:
        problems.append(
            [
                item.level,
                item.code,
                "RFP",
                item.relpath,
                str(item.path) if item.path else "",
                item.excel_row,
                item.rfp_key,
                item.group_id,
                "да" if item.global_blocker else "",
                item.message,
            ]
        )
        _mark_level(problems.cell(row=problems.max_row, column=1), item.level)
        if item.path is not None:
            _set_path_cell(problems.cell(row=problems.max_row, column=4), item.path)
            _set_path_cell(problems.cell(row=problems.max_row, column=5), item.path)
    for item in baseline_issues:
        problems.append(
            [
                item.level,
                item.code,
                "ДС",
                item.relpath,
                str(item.path) if item.path else "",
                item.excel_row,
                "",
                item.group_id,
                "да" if item.blocking else "",
                item.message,
            ]
        )
        _mark_level(problems.cell(row=problems.max_row, column=1), item.level)
        if item.path is not None:
            _set_path_cell(problems.cell(row=problems.max_row, column=4), item.path)
            _set_path_cell(problems.cell(row=problems.max_row, column=5), item.path)
    _style_header(problems, 10)
    problems.column_dimensions["D"].width = 36
    problems.column_dimensions["E"].width = 70
    problems.column_dimensions["J"].width = 70

    origin = wb.create_sheet("Происхождение")
    origin.append(
        [
            "№",
            "Группа / Имя ДС",
            "Источник строки",
            "Код",
            "Наименование",
            "Теги",
            "Кол-во",
            "Ед. изм.",
            "Статус поставки",
            "Файл",
            "Лист",
            "Строка",
            "Статус ед. изм.",
        ]
    )
    for item in hybrid_rows:
        record = item.record
        origin.append(
            [
                record.ds_number,
                record.ds_name,
                item.source,
                record.code,
                record.name,
                record.tags,
                _qty_excel(record.values),
                record.units,
                record.rfp_supply_status,
                record.file_name,
                record.sheet,
                record.excel_row,
                record.units_check_status,
            ]
        )
    _style_header(origin, 13)
    origin.column_dimensions["B"].width = 18
    origin.column_dimensions["C"].width = 16
    origin.column_dimensions["E"].width = 36
    origin.column_dimensions["F"].width = 28
    origin.column_dimensions["J"].width = 36
    return _save_workbook_atomic(path, wb)


# ---------------------------------------------------------------------------
# Public entry
# ---------------------------------------------------------------------------


def build_ds_rfp_hybrid(
    baseline: DsBaselineResult,
    registry: DsRegistryDocument,
    rfp_root: str | Path,
    output_dir: str | Path,
    *,
    write_hybrid: bool = True,
    converter: RfpUnitsConverter | None = None,
    loader: RfpWorkbookLoader | None = None,
    google_index: GoogleUnitsIndex | None = None,
    matrix_path: str | Path | None = None,
    stamp: str | None = None,
    equipment_by_code: dict[str, str] | None = None,
    progress_callback: Callable[[int, int, str, float | None], object] | None = None,
    phase_callback: Callable[[str], object] | None = None,
) -> DsRfpHybridResult:
    """Reconcile root RFP against a DS baseline and write the hybrid overlay.

    Args:
        baseline: Result of ``build_ds_baseline`` (positions already converted).
        registry: Loaded canonical registry (rfp_key ↔ group_id).
        rfp_root: Root-only folder of RFP workbooks.
        output_dir: Stable folder for ``Свод ДС-RFP для запуска.xlsx``. The audit
            report goes into a ``YYYY.MM.DD_HH.MM`` child (the same child as
            the baseline reports when they were just written here). Required;
            never defaults to UNC.
        write_hybrid: When True, write ``Свод ДС-RFP для запуска.xlsx`` iff
            the baseline has zero blockers and this stage has zero global
            blockers.
        converter: Units conversion. Default uses the parts plan/matrix/Google
            contour. Pass ``IdentityRfpUnitsConverter`` for offline tests.
        loader: Workbook extractor. Default is production ``_extract_records``.
        google_index: Google units index for the default converter.
        matrix_path: Units matrix for the default converter.
        stamp: Report timestamp stored in the summary sheet.
        equipment_by_code: Optional Google equipment map for net column F.

    Returns:
        ``DsRfpHybridResult`` with groups, issues, fingerprint and paths.
        The audit report is always written.
    """

    root = Path(rfp_root)
    stable_dir = Path(output_dir)
    stable_dir.mkdir(parents=True, exist_ok=True)
    baseline_reports = baseline.output_dir
    if (
        baseline_reports.parent == stable_dir
        and _AUDIT_DIR_RE.match(baseline_reports.name)
    ):
        report_dir = baseline_reports
    else:
        report_dir = make_audit_stamp_dir(stable_dir)
    stamp_value = stamp or datetime.now().strftime(STAMP_FORMAT)
    issues: list[DsRfpHybridIssue] = []
    if not root.is_dir():
        issues.append(
            _issue(
                ISSUE_RFP_ROOT,
                "WARN",
                f"корень RFP не найден: {root}",
                blocking=False,
                path=root,
                relpath=str(root),
            )
        )

    if baseline.blocking:
        issues.append(
            _issue(
                ISSUE_BASELINE_BLOCKED,
                "ERROR",
                "baseline ДС содержит блокеры; гибрид не записывается",
                global_blocker=True,
                path=baseline.output_dir,
                relpath=str(baseline.output_dir),
            )
        )

    files, skipped = collect_rfp_workbooks(root)
    key_groups = _rfp_key_groups(registry)
    file_groups = _rfp_file_groups(registry)
    loader_impl: RfpWorkbookLoader = loader or ProductionRfpLoader()
    extracts: list[RfpFileExtract] = []
    if phase_callback is not None:
        phase_callback(
            f"разбор RFP: {len(files)} файлов, пропущено {len(skipped)}"
        )
    total_rfp = len(files)

    for index, source in enumerate(files, start=1):
        relpath = str(source.relpath)
        if progress_callback is not None:
            progress_callback(index, total_rfp, relpath, None)
        file_start = time.perf_counter()
        identity = parse_rfp_ds_identity(source.path.name)
        rfp_key = rfp_identity_key(identity) or ""
        group_id = ""
        mapping_issue = ""
        named = file_groups.get(source.path.name.casefold(), ())
        if len(named) == 1:
            group_id = named[0]
        elif len(named) > 1:
            mapping_issue = "неоднозначный файл RFP"
            issues.append(
                _issue(
                    ISSUE_AMBIGUOUS_MAPPING,
                    "ERROR",
                    (
                        f"файл RFP {source.path.name!r} относится к нескольким группам: "
                        + ", ".join(named)
                    ),
                    global_blocker=True,
                    path=source.path,
                    relpath=source.relpath,
                    rfp_key=rfp_key,
                )
            )
        elif not rfp_key:
            mapping_issue = "имя файла не разобрано"
            issues.append(
                _issue(
                    ISSUE_UNPARSED_RFP,
                    "ERROR",
                    f"не разобрано имя RFP {source.path.name!r}",
                    global_blocker=True,
                    path=source.path,
                    relpath=source.relpath,
                )
            )
        else:
            mapped = key_groups.get(rfp_key, ())
            if len(mapped) == 1:
                group_id = mapped[0]
            elif len(mapped) > 1:
                mapping_issue = "неоднозначный ключ RFP"
                issues.append(
                    _issue(
                        ISSUE_AMBIGUOUS_MAPPING,
                        "ERROR",
                        (
                            f"ключ RFP {rfp_key!r} относится к нескольким группам: "
                            + ", ".join(mapped)
                        ),
                        global_blocker=True,
                        path=source.path,
                        relpath=source.relpath,
                        rfp_key=rfp_key,
                    )
                )
            else:
                mapping_issue = "ключ RFP нет в реестре"
        records, stats, warnings, status_issues = loader_impl.extract(source.path)
        extract_error = _file_extract_error(stats, warnings, status_issues)
        if extract_error:
            detail = "; ".join(
                message
                for level, _name, message in warnings
                if level == "ERROR"
            ) or "ошибка разбора файла RFP"
            issues.append(
                _issue(
                    ISSUE_RFP_EXTRACT,
                    "ERROR",
                    detail,
                    path=source.path,
                    relpath=source.relpath,
                    rfp_key=rfp_key,
                    group_id=group_id,
                )
            )
        extracts.append(
            RfpFileExtract(
                source=source,
                identity=identity,
                rfp_key=rfp_key,
                group_id=group_id,
                records=list(records),
                warnings=list(warnings),
                status_issues=list(status_issues),
                stats=stats,
                extract_error=extract_error,
                mapping_issue=mapping_issue,
            )
        )
        if progress_callback is not None:
            progress_callback(
                index,
                total_rfp,
                relpath,
                time.perf_counter() - file_start,
            )

    hashed: dict[tuple[str, str], list[RfpFileExtract]] = defaultdict(list)
    for item in extracts:
        if not item.group_id:
            continue
        try:
            digest = hashlib.sha256(item.source.path.read_bytes()).hexdigest()
        except OSError:
            continue
        hashed[(item.group_id, digest)].append(item)
    duplicate_groups: set[str] = set()
    for items in hashed.values():
        if len(items) < 2:
            continue
        group_id = items[0].group_id
        duplicate_groups.add(group_id)
        names = ", ".join(item.source.relpath for item in items)
        issues.append(
            _issue(
                ISSUE_DUPLICATE_RFP_KEY,
                "ERROR",
                f"точный дубль файла RFP в группе {group_id}: {names}",
                path=items[0].source.path,
                relpath=items[0].source.relpath,
                rfp_key=items[0].rfp_key,
                group_id=group_id,
            )
        )
        for item in items:
            if not item.mapping_issue:
                item.mapping_issue = "точный дубль файла RFP"

    if converter is None:
        converter = PartsRfpUnitsConverter(
            google_index=google_index, matrix_path=matrix_path
        )
    conversion_plan: ConversionPlan | None = None
    all_records: list[RfpRecord] = []
    spans: list[tuple[RfpFileExtract, int]] = []
    for item in extracts:
        spans.append((item, len(item.records)))
        all_records.extend(item.records)
    try:
        batch = converter.convert_records(all_records)
        converted = list(batch.converted)
        conversion_plan = batch.plan
        for warning in batch.warnings:
            issues.append(
                _issue(ISSUE_UNITS_CONVERT, "WARN", warning, blocking=False)
            )
        offset = 0
        for item, count in spans:
            item.records = converted[offset : offset + count]
            offset += count
            bad = [
                rec
                for rec in item.records
                if not _conversion_ok(rec.units_check_status)
            ]
            if bad:
                sample = bad[0]
                issues.append(
                    _issue(
                        ISSUE_UNITS_CONVERT,
                        "ERROR",
                        (
                            f"конвертация RFP неуспешна для {sample.code!r} "
                            f"({sample.units_check_status})"
                        ),
                        path=item.source.path,
                        relpath=item.source.relpath,
                        excel_row=sample.excel_row,
                        rfp_key=item.rfp_key,
                        group_id=item.group_id,
                    )
                )
                item.extract_error = True
    except UnitsConversionError as exc:
        issues.append(
            _issue(
                ISSUE_UNITS_CONVERT,
                "ERROR",
                str(exc),
                global_blocker=True,
            )
        )

    baseline_groups = {item.group_id: item for item in baseline.groups}
    positions_by_group: dict[str, list[DsBaselinePosition]] = defaultdict(list)
    for pos in baseline.positions:
        key = pos.group_id or pos.group_label or pos.source_id
        positions_by_group[key].append(pos)

    overlay_blocked_ids = set(registry.validation.overlay_blocked_group_ids)
    for item in baseline.groups:
        if item.overlay_blocked or item.fallback:
            overlay_blocked_ids.add(item.group_id)
    for pos in baseline.positions:
        if pos.overlay_blocked or pos.grouping_fallback:
            overlay_blocked_ids.add(pos.group_id or pos.group_label)

    group_ids: list[str] = []
    seen_groups: set[str] = set()

    def _remember(group_id: str) -> None:
        if group_id and group_id not in seen_groups:
            seen_groups.add(group_id)
            group_ids.append(group_id)

    for item in baseline.groups:
        _remember(item.group_id)
    for item in extracts:
        if item.group_id:
            _remember(item.group_id)
        elif item.rfp_key:
            _remember(f"RFP_ONLY:{item.rfp_key}")
        else:
            _remember(f"RFP_ONLY:{item.source.relpath}")
    group_ids.sort()

    extracts_by_group: dict[str, list[RfpFileExtract]] = defaultdict(list)
    for item in extracts:
        if item.group_id:
            extracts_by_group[item.group_id].append(item)
        elif item.rfp_key:
            extracts_by_group[f"RFP_ONLY:{item.rfp_key}"].append(item)
        else:
            extracts_by_group[f"RFP_ONLY:{item.source.relpath}"].append(item)

    convert_failed_groups = {
        item.group_id
        for item in issues
        if item.code == ISSUE_UNITS_CONVERT and item.blocking and item.group_id
    }
    extract_failed_groups = {
        item.group_id
        for item in extracts
        if item.extract_error and item.group_id
    }
    ambiguous_groups = {
        group_id
        for key, groups in key_groups.items()
        if len(groups) > 1
        for group_id in groups
    }

    group_results: list[DsRfpGroupResult] = []
    chosen_rows: list[NetSummaryRow] = []
    for group_id in group_ids:
        info = baseline_groups.get(group_id)
        positions = list(positions_by_group.get(group_id, ()))
        file_extracts = list(extracts_by_group.get(group_id, ()))
        rfp_names = tuple(item.source.relpath for item in file_extracts)
        group_label = (
            info.group_label
            if info is not None and info.group_label
            else group_id
        )
        fallback = bool(info.fallback) if info is not None else False
        overlay_blocked = group_id in overlay_blocked_ids
        if info is not None:
            overlay_blocked = overlay_blocked or info.overlay_blocked
            fallback = fallback or info.fallback

        usable = [
            item
            for item in file_extracts
            if item.group_id
            and not item.extract_error
            and not item.mapping_issue
        ]
        rfp_records = [rec for item in usable for rec in item.records]
        listed_rfp_records = [
            rec for item in file_extracts for rec in item.records
        ]
        ds_bag = _ds_bag(positions)
        rfp_bag = _rfp_compare_bag(rfp_records)
        deltas = _build_deltas(group_id, ds_bag, rfp_bag)
        ds_nz = _nonzero(ds_bag)
        rfp_nz = _nonzero(rfp_bag)
        has_ds = bool(positions)
        has_rfp_side = bool(file_extracts)
        block_reasons: list[str] = []
        if overlay_blocked:
            block_reasons.append("группа overlay_blocked в реестре/baseline")
            issues.append(
                _issue(
                    ISSUE_OVERLAY_BLOCKED,
                    "OVERLAY",
                    "overlay группы заблокирован реестром или fallback",
                    blocking=False,
                    group_id=group_id,
                )
            )
        if fallback:
            block_reasons.append("grouping fallback")
            issues.append(
                _issue(
                    ISSUE_GROUP_FALLBACK,
                    "OVERLAY",
                    "группа в режиме fallback / неоднозначного распределения",
                    blocking=False,
                    group_id=group_id,
                )
            )
        if group_id in duplicate_groups:
            block_reasons.append("точный дубль файла RFP")
        if group_id in extract_failed_groups or group_id in convert_failed_groups:
            block_reasons.append("ошибка разбора или конвертации RFP")
        if group_id in ambiguous_groups:
            block_reasons.append("неоднозначный ключ RFP")

        if block_reasons:
            status: GroupStatus = STATUS_BLOCKED
            selected: SelectedSource = "ДС" if has_ds else ""
            reason = "; ".join(dict.fromkeys(block_reasons))
        elif has_ds and not has_rfp_side:
            status = STATUS_DS_ONLY
            selected = "ДС"
            reason = "в корне нет RFP этой группы"
        elif has_rfp_side and not has_ds:
            status = STATUS_RFP_ONLY
            selected = ""
            reason = "только RFP, в свод не входит (ДС авторитетен)"
        elif ds_nz == rfp_nz:
            status = STATUS_MATCH
            selected = "RFP"
            reason = "полное совпадение code+unit+qty"
        else:
            status = STATUS_MISMATCH
            selected = "ДС"
            reason = "расхождение количеств, кодов или единиц"

        if status == STATUS_MATCH:
            chosen_rows.extend(_rfp_net_rows(rfp_records, group_label))
        elif selected == "ДС":
            chosen_rows.extend(_ds_net_rows(positions))

        group_results.append(
            DsRfpGroupResult(
                group_id=group_id,
                group_label=group_label,
                status=status,
                selected_source=selected,
                reason=reason,
                overlay_blocked=overlay_blocked or fallback,
                fallback=fallback,
                ds_position_count=len(positions),
                rfp_record_count=len(listed_rfp_records),
                rfp_files=rfp_names,
                ds_source_ids=info.source_ids if info is not None else (),
                deltas=deltas,
            )
        )

    hybrid_rows = _renumber(chosen_rows)
    fingerprint = _hybrid_fingerprint(
        baseline=baseline, registry=registry, files=files
    )
    global_blocker_count = sum(1 for item in issues if item.global_blocker)
    blocking_count = sum(1 for item in issues if item.blocking)
    blocking = baseline.blocking or global_blocker_count > 0
    match_count = sum(1 for item in group_results if item.status == STATUS_MATCH)
    mismatch_count = sum(
        1 for item in group_results if item.status == STATUS_MISMATCH
    )
    ds_only_count = sum(1 for item in group_results if item.status == STATUS_DS_ONLY)
    rfp_only_count = sum(
        1 for item in group_results if item.status == STATUS_RFP_ONLY
    )
    blocked_count = sum(1 for item in group_results if item.status == STATUS_BLOCKED)

    fractional_paths: tuple[Path, ...] = ()
    if conversion_plan is not None and conversion_plan.fractional_issues:
        fractional_paths = tuple(
            write_fractional_conversion_logs(conversion_plan, report_dir)
        )

    hybrid_path: Path | None = None
    if write_hybrid and not blocking:
        target = stable_dir / HYBRID_XLSX_NAME
        tmp_path = target.with_name(
            f".{target.stem}.{os.getpid()}.tmp{target.suffix}"
        )
        try:
            _write_net_xlsx(
                tmp_path, hybrid_rows, equipment_by_code=equipment_by_code
            )
            os.replace(tmp_path, target)
            hybrid_path = target
        except Exception:
            tmp_path.unlink(missing_ok=True)
            raise

    report_path = report_dir / hybrid_report_filename(stamp_value)
    summary_rows: list[tuple[str, object]] = [
        ("Корень RFP", root),
        ("Реестр", registry.path),
        ("Папка отчётов", report_dir),
        ("Алгоритм", ALGORITHM_VERSION),
        ("Штамп", stamp_value),
        ("Fingerprint", fingerprint),
        ("Fingerprint baseline", baseline.fingerprint),
        ("Файлов RFP", len(files)),
        ("Пропущено", len(skipped)),
        ("Групп", len(group_results)),
        ("MATCH", match_count),
        ("MISMATCH", mismatch_count),
        ("DS_ONLY", ds_only_count),
        ("RFP_ONLY", rfp_only_count),
        ("BLOCKED", blocked_count),
        ("ERROR", blocking_count),
        ("GLOBAL blockers", global_blocker_count),
        ("Baseline blocking", "да" if baseline.blocking else "нет"),
        ("Гибрид записан", "да" if hybrid_path is not None else "нет"),
        ("Строк гибрида", len(hybrid_rows)),
        ("Сводка реестра", registry.validation.summary_line()),
        ("Сводка baseline", baseline.summary_line()),
    ]
    _write_hybrid_report(
        report_path,
        result_head=summary_rows,
        groups=group_results,
        extracts=extracts,
        files=files,
        skipped=skipped,
        issues=issues,
        baseline_issues=baseline.issues,
        hybrid_rows=hybrid_rows,
        baseline_files=baseline.files,
    )

    return DsRfpHybridResult(
        rfp_root=root,
        output_dir=report_dir,
        registry_path=registry.path,
        stamp=stamp_value,
        fingerprint=fingerprint,
        algorithm_version=ALGORITHM_VERSION,
        files=list(files),
        skipped=list(skipped),
        extracts=extracts,
        groups=group_results,
        issues=issues,
        hybrid_rows=hybrid_rows,
        file_count=len(files),
        group_count=len(group_results),
        match_count=match_count,
        mismatch_count=mismatch_count,
        ds_only_count=ds_only_count,
        rfp_only_count=rfp_only_count,
        blocked_count=blocked_count,
        blocking_issue_count=blocking_count,
        global_blocker_count=global_blocker_count,
        blocking=blocking,
        hybrid_path=hybrid_path,
        report_path=report_path,
        fractional_log_paths=fractional_paths,
    )
