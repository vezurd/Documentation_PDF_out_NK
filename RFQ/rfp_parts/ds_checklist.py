"""Load maintained DS checklist xlsx and compare it to RFP part folders.

The checklist (``ДС_увеличение_уменьшение.xlsx``) lists expected DS numbers.
Tab/CLI collection (``run_checklist_compare``) checks the increase column
against the single parts folder ``RFP_Зиновьев``. Decrease data is no longer
collected; that column is not compared.

The two-folder helper ``compare_checklist_to_folders`` remains for the main
RFP pipeline preflight (increase/decrease dirs from JSON).

Status cells may contain ``ЕСТЬ`` / ``НЕТ`` / ``Не требуется``;
only ``Не требуется`` is trusted from the file — presence is always checked
against the actual folder contents at compare time.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook

DEFAULT_BASE = Path(
    r"\\bcc\eng\PrDoc\377_НИПИГАЗ\АГХК\КСБ\RFP_MTO_VO\_RFP"
)
DEFAULT_CHECKLIST_FILE = DEFAULT_BASE / "ДС_увеличение_уменьшение.xlsx"
DEFAULT_INCREASE_DIR = DEFAULT_BASE / "RFP на увеличение"
DEFAULT_DECREASE_DIR = DEFAULT_BASE / "RFP на уменьшение"
DEFAULT_PARTS_DIR = DEFAULT_BASE / "RFP_Зиновьев"

Kind = Literal["increase", "decrease", "parts"]
IssueLevel = Literal["OK", "WARN", "ERROR"]

_NOT_REQUIRED_RE = re.compile(r"^не\s*требуется$", re.IGNORECASE)
_KIND_LABEL = {
    "increase": "увеличение",
    "decrease": "уменьшение",
    "parts": "части",
}


def parse_ds_name_from_file_name(file_name: str) -> str:
    """Extract the DS name prefix from a split RFP workbook name.

    Accepts ``ДС4...``, ``ДС 4. ...``, ``ДС47_13А_...``; spaces after ``ДС``
    are stripped from the result.
    """
    match = re.match(
        r"\s*(ДС\s*[\wА-Яа-я-]*)",
        Path(file_name).stem,
        re.IGNORECASE,
    )
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(1)).upper()


# Backward-compatible alias used by folder helpers in this module.
_parse_ds_name_from_file_name = parse_ds_name_from_file_name


@dataclass(frozen=True)
class ChecklistEntry:
    """One DS row from the increase or decrease column of the checklist."""

    ds: str
    kind: Kind
    comment: str
    not_required: bool


@dataclass
class DsChecklist:
    """Parsed checklist workbook contents."""

    path: Path
    increase: list[ChecklistEntry] = field(default_factory=list)
    decrease: list[ChecklistEntry] = field(default_factory=list)
    other: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ChecklistIssue:
    """One compare finding for a DS or folder file."""

    level: IssueLevel
    kind: Kind
    ds: str
    code: str
    message: str


@dataclass
class ChecklistCompareResult:
    """Result of comparing the checklist to folder contents."""

    checklist_path: Path
    issues: list[ChecklistIssue] = field(default_factory=list)
    load_error: str | None = None
    compare_mode: Literal["folders", "parts"] | None = None

    @property
    def ok_count(self) -> int:
        return sum(1 for item in self.issues if item.level == "OK")

    @property
    def warn_count(self) -> int:
        return sum(1 for item in self.issues if item.level == "WARN")

    @property
    def error_count(self) -> int:
        return sum(1 for item in self.issues if item.level == "ERROR")

    @property
    def is_ok(self) -> bool:
        return self.load_error is None and self.error_count == 0 and self.warn_count == 0

    def summary_line(self) -> str:
        """Short one-line status for console / GUI."""
        if self.load_error:
            return f"ERROR: список ДС — {self.load_error}"
        if self.is_ok:
            if self.compare_mode == "folders":
                return (
                    f"OK: список ДС совпадает с папками увеличения/уменьшения "
                    f"(проверено {self.ok_count} позиций)"
                )
            return (
                f"OK: список ДС совпадает с папкой частей "
                f"(проверено {self.ok_count} позиций)"
            )
        parts = [
            f"ERROR={self.error_count}" if self.error_count else None,
            f"WARN={self.warn_count}" if self.warn_count else None,
            f"OK={self.ok_count}",
        ]
        return "Сверка списка ДС: " + ", ".join(p for p in parts if p)


# Last compare result (GUI Job monitor / finish handler without a second UNC scan).
_LAST_COMPARE_RESULT: ChecklistCompareResult | None = None


def get_last_checklist_compare() -> ChecklistCompareResult | None:
    """Return the most recent ``run_checklist_compare`` result, if any."""
    return _LAST_COMPARE_RESULT


def _progress(message: str, *, enabled: bool) -> None:
    if enabled:
        print(f"[checklist] {message}", flush=True)


def is_not_required_comment(comment: str | None) -> bool:
    """Return True when the checklist comment means the DS is intentionally skipped."""
    if comment is None:
        return False
    return bool(_NOT_REQUIRED_RE.match(str(comment).strip()))


def _norm_ds_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _norm_comment(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def collect_folder_ds(
    folder: Path,
    *,
    progress: bool = False,
    label: str = "",
) -> dict[str, str]:
    """Map upper-case DS keys from xlsx/xlsm/xls names in ``folder``.

    Args:
        folder: Parts directory (``RFP_Зиновьев`` or a pipeline increase/decrease folder).
        progress: When True, print scan steps (Job monitor / CLI).
        label: Human folder label for progress lines (e.g. ``увеличение``).
    """
    tag = label or str(folder)
    found: dict[str, str] = {}
    _progress(f"Проверка папки «{tag}»: {folder}", enabled=progress)
    if not folder.exists():
        _progress(f"Папка «{tag}» не найдена (exists=False)", enabled=progress)
        return found
    _progress(f"Чтение имён файлов в «{tag}»…", enabled=progress)
    file_count = 0
    for path in folder.iterdir():
        if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
            continue
        file_count += 1
        name = _parse_ds_name_from_file_name(path.name)
        if name:
            found[name.upper()] = name
    _progress(
        f"«{tag}»: готово — xlsx={file_count}, ДС в имени={len(found)}",
        enabled=progress,
    )
    return found


def table_ds_to_file_keys(ds: str) -> list[str]:
    """Convert a checklist DS label (``14``, ``47/13А``, ``ДС100/45А``) to folder keys.

    Slash forms from the checklist become underscore forms used in file names
    (``47/13А`` / ``ДС100/45А`` → ``ДС47_13А`` / ``ДС100_45А``).
    """
    s = str(ds).strip().upper().replace(" ", "").replace("/", "_")
    if not s:
        return []
    if s.startswith("ДС"):
        base = s
    elif re.match(r"^\d", s):
        base = "ДС" + s
    else:
        base = s
    return [base]


def present_in(folder_ds: dict[str, str], ds: str) -> bool:
    """Return True if checklist DS matches any file stem in ``folder_ds``."""
    keys = table_ds_to_file_keys(ds)
    folder_keys = list(folder_ds.keys())
    for key in keys:
        if key in folder_ds:
            return True
        for folder_key in folder_keys:
            if folder_key == key or folder_key.startswith(key + "_"):
                return True
            if key.startswith(folder_key) or folder_key.startswith(key):

                def _stem(value: str) -> str:
                    match = re.match(r"(ДС\d+(?:_\d+)?)", value)
                    return match.group(1) if match else value

                if _stem(key) == _stem(folder_key):
                    return True
    return False


def matched_folder_keys(folder_ds: dict[str, str], ds: str) -> set[str]:
    """Return folder DS keys that match the given checklist DS label."""
    matched: set[str] = set()
    keys = table_ds_to_file_keys(ds)
    for folder_key, display in folder_ds.items():
        for key in keys:
            if folder_key == key or folder_key.startswith(key + "_"):
                matched.add(folder_key)
                break
            if key.startswith(folder_key) or folder_key.startswith(key):
                stem_re = re.compile(r"(ДС\d+(?:_\d+)?)")
                stem_key = stem_re.match(key)
                stem_folder = stem_re.match(folder_key)
                if stem_key and stem_folder and stem_key.group(1) == stem_folder.group(1):
                    matched.add(folder_key)
                    break
                if key == folder_key:
                    matched.add(folder_key)
                    break
    return matched


def load_ds_checklist(path: Path | None = None) -> DsChecklist:
    """Read increase/decrease DS lists and comments from the checklist xlsx.

    Args:
        path: Checklist workbook. Defaults to UNC ``ДС_увеличение_уменьшение.xlsx``.

    Returns:
        Parsed ``DsChecklist``.

    Raises:
        FileNotFoundError: If the workbook does not exist.
        ValueError: If the workbook has no usable sheet/headers.
    """
    checklist_path = Path(path or DEFAULT_CHECKLIST_FILE)
    if not checklist_path.exists():
        raise FileNotFoundError(f"checklist not found: {checklist_path}")

    workbook = load_workbook(checklist_path, read_only=True, data_only=True)
    try:
        sheet = workbook["ДС"] if "ДС" in workbook.sheetnames else workbook.active
        increase: list[ChecklistEntry] = []
        decrease: list[ChecklistEntry] = []
        other: list[str] = []
        seen_inc: set[str] = set()
        seen_dec: set[str] = set()
        seen_other: set[str] = set()

        for row in sheet.iter_rows(min_row=2, max_col=8, values_only=True):
            inc_ds = _norm_ds_cell(row[0] if len(row) > 0 else None)
            inc_comment = _norm_comment(row[1] if len(row) > 1 else None)
            dec_ds = _norm_ds_cell(row[3] if len(row) > 3 else None)
            dec_comment = _norm_comment(row[4] if len(row) > 4 else None)
            other_ds = _norm_ds_cell(row[5] if len(row) > 5 else None)

            if inc_ds and inc_ds.upper() not in seen_inc:
                seen_inc.add(inc_ds.upper())
                increase.append(
                    ChecklistEntry(
                        ds=inc_ds,
                        kind="increase",
                        comment=inc_comment,
                        not_required=is_not_required_comment(inc_comment),
                    )
                )
            if dec_ds and dec_ds.upper() not in seen_dec:
                seen_dec.add(dec_ds.upper())
                decrease.append(
                    ChecklistEntry(
                        ds=dec_ds,
                        kind="decrease",
                        comment=dec_comment,
                        not_required=is_not_required_comment(dec_comment),
                    )
                )
            if other_ds and other_ds.upper() not in seen_other:
                seen_other.add(other_ds.upper())
                other.append(other_ds)
    finally:
        workbook.close()

    if not increase and not decrease:
        raise ValueError(f"checklist has no DS rows: {checklist_path}")

    return DsChecklist(
        path=checklist_path,
        increase=increase,
        decrease=decrease,
        other=other,
    )


def _compare_entries_to_folder(
    result: ChecklistCompareResult,
    *,
    kind: Kind,
    entries: list[ChecklistEntry],
    folder: Path,
    progress: bool,
) -> None:
    """Append OK/WARN/ERROR issues for one checklist column vs one folder."""
    label = _KIND_LABEL[kind]
    _progress(
        f"Сверка «{label}»: в списке {len(entries)} позиций",
        enabled=progress,
    )
    folder_ds = collect_folder_ds(folder, progress=progress, label=label)
    covered: set[str] = set()

    for entry in entries:
        matched = matched_folder_keys(folder_ds, entry.ds)
        covered |= matched
        present = bool(matched) or present_in(folder_ds, entry.ds)
        if entry.not_required:
            if present:
                result.issues.append(
                    ChecklistIssue(
                        level="WARN",
                        kind=kind,
                        ds=entry.ds,
                        code="not_required_present",
                        message=(
                            f"{label}: ДС {entry.ds} помечен «Не требуется», "
                            "но файл есть в папке"
                        ),
                    )
                )
            else:
                result.issues.append(
                    ChecklistIssue(
                        level="OK",
                        kind=kind,
                        ds=entry.ds,
                        code="not_required_absent",
                        message=(
                            f"{label}: ДС {entry.ds} не требуется, файла нет"
                        ),
                    )
                )
        elif present:
            result.issues.append(
                ChecklistIssue(
                    level="OK",
                    kind=kind,
                    ds=entry.ds,
                    code="present",
                    message=f"{label}: ДС {entry.ds} ожидается и найден",
                )
            )
        else:
            result.issues.append(
                ChecklistIssue(
                    level="ERROR",
                    kind=kind,
                    ds=entry.ds,
                    code="missing",
                    message=(
                        f"{label}: ДС {entry.ds} есть в списке, "
                        "но файла нет в папке"
                    ),
                )
            )

    for folder_key, display in sorted(folder_ds.items()):
        if folder_key in covered:
            continue
        result.issues.append(
            ChecklistIssue(
                level="WARN",
                kind=kind,
                ds=display,
                code="extra",
                message=(
                    f"{label}: файл {display} есть в папке, "
                    "но отсутствует в списке ДС"
                ),
            )
        )
    _progress(
        f"«{label}»: сверка позиций списка завершена",
        enabled=progress,
    )


def compare_checklist_to_folders(
    checklist: DsChecklist,
    *,
    increase_dir: Path | None = None,
    decrease_dir: Path | None = None,
    progress: bool = False,
) -> ChecklistCompareResult:
    """Compare checklist to increase/decrease folders (RFP Step1 preflight).

    Args:
        checklist: Parsed checklist.
        increase_dir: Folder ``RFP на увеличение``.
        decrease_dir: Folder ``RFP на уменьшение``.
        progress: When True, print scan/compare steps.

    Returns:
        ``ChecklistCompareResult`` with OK/WARN/ERROR issues.
    """
    inc_dir = Path(increase_dir or DEFAULT_INCREASE_DIR)
    dec_dir = Path(decrease_dir or DEFAULT_DECREASE_DIR)
    result = ChecklistCompareResult(
        checklist_path=checklist.path, compare_mode="folders"
    )
    _compare_entries_to_folder(
        result,
        kind="increase",
        entries=checklist.increase,
        folder=inc_dir,
        progress=progress,
    )
    _compare_entries_to_folder(
        result,
        kind="decrease",
        entries=checklist.decrease,
        folder=dec_dir,
        progress=progress,
    )
    return result


def compare_checklist_to_parts_folder(
    checklist: DsChecklist,
    *,
    parts_dir: Path | None = None,
    progress: bool = False,
) -> ChecklistCompareResult:
    """Compare the increase column to the single parts folder ``RFP_Зиновьев``.

    The decrease column is ignored: those files are no longer collected.

    Args:
        checklist: Parsed checklist.
        parts_dir: Folder with RFP part workbooks. Defaults to ``DEFAULT_PARTS_DIR``.
        progress: When True, print scan/compare steps.

    Returns:
        ``ChecklistCompareResult`` with OK/WARN/ERROR issues.
    """
    folder = Path(parts_dir or DEFAULT_PARTS_DIR)
    result = ChecklistCompareResult(
        checklist_path=checklist.path, compare_mode="parts"
    )
    _compare_entries_to_folder(
        result,
        kind="parts",
        entries=checklist.increase,
        folder=folder,
        progress=progress,
    )
    return result


def run_checklist_compare(
    *,
    checklist_path: Path | None = None,
    parts_dir: Path | None = None,
    increase_dir: Path | None = None,
    decrease_dir: Path | None = None,
    progress: bool = False,
) -> ChecklistCompareResult:
    """Load checklist and compare to the parts folder (tab / CLI collection).

    Args:
        checklist_path: Optional override for the checklist xlsx.
        parts_dir: Optional parts folder (``RFP_Зиновьев``).
        increase_dir: Unused; kept so old callers do not break.
        decrease_dir: Unused; kept so old callers do not break.
        progress: When True, print steps and the result table (Job monitor / CLI).

    Returns:
        Compare result; ``load_error`` set if the checklist could not be read.
    """
    del increase_dir, decrease_dir
    global _LAST_COMPARE_RESULT
    path = Path(checklist_path or DEFAULT_CHECKLIST_FILE)
    _progress("=" * 60, enabled=progress)
    _progress("Старт сверки списка ДС с папкой частей", enabled=progress)
    _progress(f"Файл списка: {path}", enabled=progress)
    try:
        _progress("Чтение xlsx списка ДС…", enabled=progress)
        checklist = load_ds_checklist(path)
    except (OSError, FileNotFoundError, ValueError) as exc:
        result = ChecklistCompareResult(checklist_path=path, load_error=str(exc))
        _LAST_COMPARE_RESULT = result
        _progress(f"Ошибка чтения списка: {exc}", enabled=progress)
        if progress:
            print(format_checklist_report_section(result), flush=True)
        return result

    _progress(
        f"Список загружен: ожидаемых ДС={len(checklist.increase)} "
        f"(колонка увеличения; уменьшение не сверяется), "
        f"прочее={len(checklist.other)}",
        enabled=progress,
    )
    result = compare_checklist_to_parts_folder(
        checklist,
        parts_dir=parts_dir,
        progress=progress,
    )
    _LAST_COMPARE_RESULT = result
    _progress(result.summary_line(), enabled=progress)
    if progress:
        print(flush=True)
        print(format_checklist_report_section(result), flush=True)
        print(flush=True)
        print(format_checklist_full_table(result), flush=True)
        _progress("Сверка завершена", enabled=True)
    return result


def format_checklist_report_section(result: ChecklistCompareResult) -> str:
    """Build a text-report section for checklist vs folders compare."""
    lines: list[str] = []
    lines.append("Сверка со списком ДС (ДС_увеличение_уменьшение.xlsx)")
    lines.append("-" * 80)
    lines.append(f"Файл списка: {result.checklist_path}")
    lines.append(result.summary_line())
    if result.load_error:
        lines.append(f"Ошибка чтения: {result.load_error}")
        return "\n".join(lines)

    if result.compare_mode == "folders":
        lines.append(
            "Из файла списка учитываются колонки увеличения и уменьшения "
            "и пометка «Не требуется»; наличие файлов проверяется по папкам "
            "«RFP на увеличение» и «RFP на уменьшение»."
        )
    else:
        lines.append(
            "Из файла списка учитывается колонка увеличения и пометка «Не требуется»; "
            "наличие файлов проверяется по папке частей (RFP_Зиновьев). "
            "Колонка уменьшения не сверяется."
        )

    problems = [i for i in result.issues if i.level != "OK"]
    if not problems:
        lines.append("Расхождений нет.")
        return "\n".join(lines)

    lines.append("")
    lines.append(f"{'Уровень':<6} {'Папка':<12} {'ДС':<16} Сообщение")
    for issue in problems:
        lines.append(
            f"{issue.level:<6} {_KIND_LABEL[issue.kind]:<12} "
            f"{issue.ds:<16} {issue.message}"
        )
    return "\n".join(lines)


def format_checklist_full_table(result: ChecklistCompareResult) -> str:
    """Build a full OK/WARN/ERROR issues table for console / Job monitor."""
    lines: list[str] = []
    lines.append("Полная таблица сверки (все позиции)")
    lines.append("-" * 80)
    if result.load_error:
        lines.append(f"Ошибка чтения: {result.load_error}")
        return "\n".join(lines)
    if not result.issues:
        lines.append("(нет позиций)")
        return "\n".join(lines)
    lines.append(
        f"{'Уровень':<6} {'Папка':<12} {'ДС':<16} {'Код':<22} Сообщение"
    )
    for issue in result.issues:
        lines.append(
            f"{issue.level:<6} {_KIND_LABEL[issue.kind]:<12} "
            f"{issue.ds:<16} {issue.code:<22} {issue.message}"
        )
    lines.append("-" * 80)
    lines.append(
        f"Итого: OK={result.ok_count}, WARN={result.warn_count}, "
        f"ERROR={result.error_count}"
    )
    return "\n".join(lines)
