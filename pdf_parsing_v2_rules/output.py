from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
import hashlib
import os
from pathlib import Path
from typing import Any, Sequence

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import utils.excel as utils_excel
import utils.path


NORMCONTROL_SHEET_TITLE = "Замечания НК"
NORMCONTROL_HEADER_ROW = 1
NORMCONTROL_NO_ERRORS_ROW_KEY = "__NK_NO_ERRORS_SUMMARY__"
NORMCONTROL_SUMMARY_PLACEHOLDER_FIELD = "nk_summary_placeholder"
NORMCONTROL_NO_ERRORS_MESSAGE_RU = "Анализ прошел успешно, ошибок не выявлено."
_EXCEL_SUMMARY_GRAY = "FF808080"
_HEADER_FILL_MAIN = "D9EAF7"
_HEADER_FILL_REVIEW = "E2F0D9"
_RESOLVED_FILL = "E2F0D9"
_BORDER_THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_ALIGN_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_ALIGN_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)


@dataclass(frozen=True)
class NormcontrolColumnDef:
    """Excel column definition for the normcontrol report."""

    key: str
    header_label: str
    width: float
    alignment: Alignment
    hidden: bool = False


NORMCONTROL_COLUMNS: list[NormcontrolColumnDef] = [
    NormcontrolColumnDef("result", "Ошибка", 10, _ALIGN_CENTER_WRAP),
    NormcontrolColumnDef("c_code", "Код ошибки", 12, _ALIGN_CENTER_WRAP),
    NormcontrolColumnDef("c_description", "Описание проверки", 36, _ALIGN_LEFT_WRAP),
    NormcontrolColumnDef("doc_name", "Имя документа", 34, _ALIGN_LEFT_WRAP),
    NormcontrolColumnDef("page_num", "Страница", 10, _ALIGN_CENTER_WRAP),
    NormcontrolColumnDef("text", "Замечание", 80, _ALIGN_LEFT_WRAP),
    NormcontrolColumnDef("resolved", "Исправлено", 12, _ALIGN_CENTER_WRAP),
    NormcontrolColumnDef("resolved_at", "Дата отметки", 20, _ALIGN_CENTER_WRAP),
    NormcontrolColumnDef("comment", "Комментарий", 28, _ALIGN_LEFT_WRAP),
    NormcontrolColumnDef("row_key", "row_key", 18, _ALIGN_LEFT_WRAP, hidden=True),
    NormcontrolColumnDef("pdf_path", "pdf_path", 40, _ALIGN_LEFT_WRAP, hidden=True),
]
_REVIEW_COLUMN_KEYS = {"resolved", "resolved_at", "comment", "row_key", "pdf_path"}


@dataclass
class CheckRow:
    """One normcontrol check result."""

    result: bool
    c_code: int
    c_description: str
    doc_name: str
    page_num: int | str
    text: str
    pdf_path: str = ""
    """Optional PDF path when ``doc_name`` is not in ``curr_proj`` (e.g. файл ОД)."""


def make_check_row(
    result: bool,
    c_code_pair: Sequence[Any],
    doc: str,
    page: Any,
    text: str,
    *,
    pdf_path: str = "",
) -> CheckRow:
    """Build CheckRow; page=\"нет\" → page_num=\"нет\", else page.page_num."""
    if page == "нет":
        page_num: int | str = "нет"
    else:
        page_num = page.page_num
    return CheckRow(
        result=result,
        c_code=c_code_pair[0],
        c_description=c_code_pair[1],
        doc_name=doc,
        page_num=page_num,
        text=text,
        pdf_path=str(pdf_path or "").strip(),
    )


def print_check_list(
    check_list: list[CheckRow],
    tag: str = "",
    print_check_mode: int = 0,
    c_code_filter: int = 0,
) -> list[CheckRow]:
    # print_check_mode == 0 - только ошибки, все проверки
    # print_check_mode == 1 - и ошибки и не ошибки, все проверки
    # print_check_mode == 2 - только ошибки, проверки по c_code
    # print_check_mode == 3 - и ошибки и не ошибки, проверки по c_code
    check_list_out: list[CheckRow] = []
    from prettytable import PrettyTable

    table = PrettyTable()
    table.field_names = [
        "Флаг",
        "Код пров.",
        "Описание проверки",
        "Имя документа",
        "№ стр.",
        "Вывод проверки",
    ]
    table.border = 1
    table.align = "l"
    print(f"Выводим список ошибок... <{tag}>")
    for row in check_list:
        row_tuple = [
            row.result,
            row.c_code,
            row.c_description,
            row.doc_name,
            row.page_num,
            row.text,
        ]
        if print_check_mode == 0:
            if not row.result:
                table.add_row(row_tuple)
                check_list_out.append(row)
        elif print_check_mode == 1:
            table.add_row(row_tuple)
            check_list_out.append(row)
        elif print_check_mode == 2:
            if c_code_filter == row.c_code:
                if not row.result:
                    table.add_row(row_tuple)
                    check_list_out.append(row)
        elif print_check_mode == 3:
            if c_code_filter == row.c_code:
                table.add_row(row_tuple)
                check_list_out.append(row)
    print(f"   Количество ошибок - {len(check_list_out)}")
    print(table)
    print("Конец вывода ошибок.\n")
    return check_list_out


def _page_sort_value(page_num: int | str) -> tuple[int, int | str]:
    """Return a stable sort key for page numbers."""
    if isinstance(page_num, int):
        return (0, page_num)
    text = str(page_num).strip()
    if text.isdigit():
        return (0, int(text))
    return (1, text.lower())


def _row_key_seed(row: CheckRow) -> str:
    """Build a stable seed string for row-key generation."""
    return "|".join(
        [
            "1" if row.result else "0",
            str(row.c_code),
            row.c_description.strip(),
            row.doc_name.strip(),
            str(row.page_num).strip(),
            row.text.strip(),
        ]
    )


def _row_key_from_seed(seed: str, occurrence_index: int) -> str:
    digest = hashlib.sha1(seed.encode("utf-8")).hexdigest()[:12]
    return f"{digest}:{occurrence_index}"


def _doc_pdf_paths(curr_proj: Sequence[Any]) -> dict[str, str]:
    """Build doc-name -> pdf-path map for UI navigation."""
    mapping: dict[str, str] = {}
    for doc in curr_proj:
        doc_name = str(getattr(doc, "doc_OD_style_file_name", "") or "").strip()
        file_path = str(getattr(doc, "file_full_path", "") or "").strip()
        if doc_name and file_path and doc_name not in mapping:
            mapping[doc_name] = file_path
    return mapping


def _serialize_rows(check_list: list[CheckRow], curr_proj: Sequence[Any]) -> list[dict[str, Any]]:
    """Convert CheckRow objects into UI- and Excel-friendly dicts."""
    doc_paths = _doc_pdf_paths(curr_proj)
    occurrence_by_seed: dict[str, int] = defaultdict(int)
    rows: list[dict[str, Any]] = []
    for row in check_list:
        seed = _row_key_seed(row)
        occurrence_by_seed[seed] += 1
        row_key = _row_key_from_seed(seed, occurrence_by_seed[seed])
        page_num = row.page_num
        pdf_path = str(getattr(row, "pdf_path", "") or "").strip()
        if not pdf_path:
            pdf_path = doc_paths.get(str(row.doc_name).strip(), "")
        rows.append(
            {
                "row_key": row_key,
                "result": bool(row.result),
                "result_label": "OK" if row.result else "Ошибка",
                "c_code": int(row.c_code),
                "c_description": str(row.c_description),
                "doc_name": str(row.doc_name),
                "page_num": page_num,
                "page_sort": _page_sort_value(page_num),
                "text": str(row.text),
                "pdf_path": pdf_path,
                "resolved": False,
                "resolved_at": "",
                "comment": "",
            }
        )
    return rows


def normcontrol_no_errors_summary_row() -> dict[str, Any]:
    """One synthetic row when checks ran and every check passed (no flagged rows)."""
    return {
        "row_key": NORMCONTROL_NO_ERRORS_ROW_KEY,
        "result": True,
        "result_label": "OK",
        "c_code": "",
        "c_description": "",
        "doc_name": "",
        "page_num": "",
        "page_sort": (-2, ""),
        "text": NORMCONTROL_NO_ERRORS_MESSAGE_RU,
        "pdf_path": "",
        "resolved": False,
        "resolved_at": "",
        "comment": "",
        NORMCONTROL_SUMMARY_PLACEHOLDER_FIELD: True,
    }


def _row_to_excel_values(row: dict[str, Any]) -> dict[str, Any]:
    """Normalize one serialized row for writing to Excel."""
    return {
        "result": "Нет" if row.get("result") else "Да",
        "c_code": row.get("c_code", ""),
        "c_description": row.get("c_description", ""),
        "doc_name": row.get("doc_name", ""),
        "page_num": row.get("page_num", ""),
        "text": row.get("text", ""),
        "resolved": "Да" if row.get("resolved") else "",
        "resolved_at": row.get("resolved_at", ""),
        "comment": row.get("comment", ""),
        "row_key": row.get("row_key", ""),
        "pdf_path": row.get("pdf_path", ""),
    }


def _apply_excel_row_style(
    ws: Any,
    row_idx: int,
    resolved: bool,
    *,
    summary_placeholder: bool = False,
) -> None:
    """Apply borders/alignment/fill (and optional gray text) for one worksheet row."""
    fill = PatternFill(fill_type="solid", fgColor=_RESOLVED_FILL) if resolved else None
    gray_font = Font(color=_EXCEL_SUMMARY_GRAY)
    for col_idx, col_def in enumerate(NORMCONTROL_COLUMNS, start=1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.alignment = col_def.alignment
        cell.border = _BORDER_THIN
        if summary_placeholder:
            cell.font = gray_font
            cell.fill = PatternFill(fill_type=None)
        else:
            cell.font = Font()
            if fill is not None:
                cell.fill = fill
            else:
                cell.fill = PatternFill(fill_type=None)


def _prepare_normcontrol_worksheet(ws: Any) -> None:
    """Create header row, widths, freeze panes and autofilter."""
    ws.title = NORMCONTROL_SHEET_TITLE
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = True
    ws.row_dimensions[1].height = 30
    for col_idx, col_def in enumerate(NORMCONTROL_COLUMNS, start=1):
        cell = ws.cell(row=NORMCONTROL_HEADER_ROW, column=col_idx)
        cell.value = col_def.header_label
        cell.font = Font(bold=True)
        cell.alignment = _ALIGN_CENTER_WRAP
        cell.border = _BORDER_THIN
        header_fill = _HEADER_FILL_REVIEW if col_def.key in _REVIEW_COLUMN_KEYS else _HEADER_FILL_MAIN
        cell.fill = PatternFill(fill_type="solid", fgColor=header_fill)
        dim = ws.column_dimensions[get_column_letter(col_idx)]
        dim.width = col_def.width
        dim.hidden = col_def.hidden


def build_normcontrol_report(
    check_list: list[CheckRow],
    flagged_check_list: list[CheckRow],
    curr_proj: Sequence[Any],
) -> dict[str, Any]:
    """Build a structured normcontrol payload for GUI and Excel export."""
    all_rows = _serialize_rows(check_list, curr_proj)
    flagged_rows = [row for row in all_rows if not row["result"]]
    all_rows.sort(key=lambda item: (item["doc_name"].lower(), item["page_sort"], item["c_code"], item["row_key"]))
    flagged_rows.sort(
        key=lambda item: (item["doc_name"].lower(), item["page_sort"], item["c_code"], item["row_key"])
    )
    summary_placeholder_row: dict[str, Any] | None = None
    if all_rows and not flagged_rows:
        summary_placeholder_row = normcontrol_no_errors_summary_row()
    return {
        "all_rows": all_rows,
        "flagged_rows": flagged_rows,
        "stats": {
            "total_checks": len(all_rows),
            "total_errors": len(flagged_rows),
            "passed_checks": max(len(all_rows) - len(flagged_rows), 0),
            "resolved_errors": 0,
            "active_errors": len(flagged_rows),
        },
        "summary_placeholder_row": summary_placeholder_row,
        "xlsx_path": "",
    }


def formats_convolution(input_list: list[str]) -> str:
    dict_list: dict[str, int] = {}
    for i in input_list:
        if i in dict_list:
            dict_list[i] = dict_list[i] + 1
        else:
            dict_list[i] = 1
    output_str = ""
    for k in dict_list:
        if dict_list[k] > 1:
            output_str = output_str + str(dict_list[k]) + k + ", "
        else:
            output_str = output_str + k + ", "
    output_str = output_str.strip().strip(",")
    return output_str


def write_normcontrol_workbook(
    flagged_rows: list[dict[str, Any]],
    xlsx_path: str,
    *,
    sanitize_illegal_chars: bool = True,
) -> str:
    """Write the current flagged normcontrol rows into an xlsx workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    _prepare_normcontrol_worksheet(ws)
    for row_idx, row in enumerate(flagged_rows, start=2):
        excel_row = _row_to_excel_values(row)
        for col_idx, col_def in enumerate(NORMCONTROL_COLUMNS, start=1):
            raw = excel_row.get(col_def.key, "")
            ws.cell(row=row_idx, column=col_idx).value = utils_excel.sanitize_cell_value_for_openpyxl(
                raw,
                enabled=sanitize_illegal_chars,
            )
        is_summary = bool(row.get(NORMCONTROL_SUMMARY_PLACEHOLDER_FIELD))
        _apply_excel_row_style(
            ws,
            row_idx,
            bool(row.get("resolved")) and not is_summary,
            summary_placeholder=is_summary,
        )
    last_col = get_column_letter(len(NORMCONTROL_COLUMNS))
    last_row = max(len(flagged_rows) + 1, 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    utils.path.make_dir(str(Path(xlsx_path).parent))
    wb.save(str(xlsx_path))
    wb.close()
    return xlsx_path


def excel_check_list_out(
    flagged_rows: list[CheckRow],
    full_check_list: list[CheckRow],
    curr_proj: Any,
    path_out_dir: str,
    *,
    sanitize_illegal_chars: bool = True,
) -> str:
    """Write flagged check rows to the normcontrol xlsx report.

    When there are no flagged rows but ``full_check_list`` is non-empty, writes one
    gray summary row (analysis completed, no issues).
    """
    title_mark = curr_proj[0].doc_Short_Title
    serialized = _serialize_rows(flagged_rows, curr_proj)
    if not serialized and full_check_list:
        serialized = [normcontrol_no_errors_summary_row()]
    serialized.sort(key=lambda item: (item["doc_name"].lower(), item["page_sort"], item["c_code"], item["row_key"]))
    out_path = Path(path_out_dir) / f"{title_mark}_результат_проверки.xlsx"
    try:
        return write_normcontrol_workbook(
            serialized,
            str(out_path),
            sanitize_illegal_chars=sanitize_illegal_chars,
        )
    except PermissionError as err:
        print(err)
        print("\tFAIL save Excel (excel_check_list_out)\n")
    return str(out_path)


def load_review_state_from_workbook(xlsx_path: str) -> dict[str, dict[str, Any]]:
    """Load review-state columns from an existing normcontrol workbook."""
    if not xlsx_path or not os.path.exists(xlsx_path):
        return {}
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        ws = wb[NORMCONTROL_SHEET_TITLE] if NORMCONTROL_SHEET_TITLE in wb.sheetnames else wb.active
        header_to_col: dict[str, int] = {}
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row=NORMCONTROL_HEADER_ROW, column=col_idx).value
            if value:
                header_to_col[str(value).strip()] = col_idx
        required = {"row_key", "Исправлено", "Дата отметки", "Комментарий"}
        if not required.issubset(header_to_col):
            return {}
        rows: dict[str, dict[str, Any]] = {}
        for row_idx in range(2, ws.max_row + 1):
            row_key = str(ws.cell(row=row_idx, column=header_to_col["row_key"]).value or "").strip()
            if not row_key:
                continue
            resolved_text = str(
                ws.cell(row=row_idx, column=header_to_col["Исправлено"]).value or ""
            ).strip().lower()
            rows[row_key] = {
                "resolved": resolved_text in {"да", "true", "1", "yes", "x"},
                "resolved_at": str(
                    ws.cell(row=row_idx, column=header_to_col["Дата отметки"]).value or ""
                ).strip(),
                "comment": str(ws.cell(row=row_idx, column=header_to_col["Комментарий"]).value or "").strip(),
            }
        return rows
    finally:
        wb.close()


def _header_to_col_map(ws: Any) -> dict[str, int]:
    """Return worksheet header label -> column index map."""
    header_to_col: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = ws.cell(row=NORMCONTROL_HEADER_ROW, column=col_idx).value
        if value:
            header_to_col[str(value).strip()] = col_idx
    return header_to_col


def _review_flag_from_cell(value: Any) -> bool:
    """Parse a checkbox-like cell value into bool."""
    text = str(value or "").strip().lower()
    return text in {"да", "true", "1", "yes", "x"}


def _optional_cell_text(ws: Any, row_idx: int, header_to_col: dict[str, int], header: str) -> str:
    """Return optional cell text by header name or empty string."""
    col_idx = header_to_col.get(header)
    if not col_idx:
        return ""
    return str(ws.cell(row=row_idx, column=col_idx).value or "").strip()


def load_normcontrol_report_from_workbook(xlsx_path: str) -> dict[str, Any]:
    """Load full normcontrol rows from an existing xlsx workbook."""
    if not xlsx_path or not os.path.exists(xlsx_path):
        return {
            "all_rows": [],
            "flagged_rows": [],
            "stats": {},
            "summary_placeholder_row": None,
            "xlsx_path": xlsx_path,
        }
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        ws = wb[NORMCONTROL_SHEET_TITLE] if NORMCONTROL_SHEET_TITLE in wb.sheetnames else wb.active
        header_to_col = _header_to_col_map(ws)
        required = {"Ошибка", "Код ошибки", "Описание проверки", "Имя документа", "Страница", "Замечание"}
        if not required.issubset(header_to_col):
            return {
                "all_rows": [],
                "flagged_rows": [],
                "stats": {},
                "summary_placeholder_row": None,
                "xlsx_path": xlsx_path,
            }
        rows: list[dict[str, Any]] = []
        for row_idx in range(2, ws.max_row + 1):
            doc_name = str(ws.cell(row=row_idx, column=header_to_col["Имя документа"]).value or "").strip()
            c_description = str(
                ws.cell(row=row_idx, column=header_to_col["Описание проверки"]).value or ""
            ).strip()
            text = str(ws.cell(row=row_idx, column=header_to_col["Замечание"]).value or "").strip()
            if not doc_name and not c_description and not text:
                continue
            page_raw = ws.cell(row=row_idx, column=header_to_col["Страница"]).value
            if isinstance(page_raw, int):
                page_num: int | str = page_raw
            else:
                page_text = str(page_raw or "").strip()
                page_num = int(page_text) if page_text.isdigit() else page_text
            row_key = _optional_cell_text(ws, row_idx, header_to_col, "row_key")
            if not row_key:
                seed_row = CheckRow(
                    result=str(ws.cell(row=row_idx, column=header_to_col["Ошибка"]).value or "").strip().lower()
                    not in {"да", "error"},
                    c_code=int(ws.cell(row=row_idx, column=header_to_col["Код ошибки"]).value or 0),
                    c_description=c_description,
                    doc_name=doc_name,
                    page_num=page_num,
                    text=text,
                )
                row_key = _row_key_from_seed(_row_key_seed(seed_row), row_idx - 1)
            result_is_ok = (
                str(ws.cell(row=row_idx, column=header_to_col["Ошибка"]).value or "").strip().lower()
                in {"нет", "ok", "false", ""}
            )
            row = {
                "row_key": row_key,
                "result": result_is_ok,
                "result_label": "OK" if result_is_ok else "Ошибка",
                "c_code": int(ws.cell(row=row_idx, column=header_to_col["Код ошибки"]).value or 0),
                "c_description": c_description,
                "doc_name": doc_name,
                "page_num": page_num,
                "page_sort": _page_sort_value(page_num),
                "text": text,
                "pdf_path": _optional_cell_text(ws, row_idx, header_to_col, "pdf_path"),
                "resolved": _review_flag_from_cell(
                    _optional_cell_text(ws, row_idx, header_to_col, "Исправлено")
                ),
                "resolved_at": _optional_cell_text(ws, row_idx, header_to_col, "Дата отметки"),
                "comment": _optional_cell_text(ws, row_idx, header_to_col, "Комментарий"),
            }
            rows.append(row)
        summary_placeholder_row: dict[str, Any] | None = None
        data_rows: list[dict[str, Any]] = []
        for row in rows:
            if str(row.get("row_key") or "").strip() == NORMCONTROL_NO_ERRORS_ROW_KEY:
                row[NORMCONTROL_SUMMARY_PLACEHOLDER_FIELD] = True
                row["c_code"] = ""
                summary_placeholder_row = row
            else:
                data_rows.append(row)
        data_rows.sort(
            key=lambda item: (item["doc_name"].lower(), item["page_sort"], item["c_code"], item["row_key"])
        )
        flagged = [row for row in data_rows if not row.get("result", False)]
        resolved_errors = sum(1 for row in flagged if row.get("resolved", False))
        return {
            "all_rows": data_rows,
            "flagged_rows": flagged,
            "stats": {
                "total_checks": len(data_rows),
                "total_errors": len(flagged),
                "passed_checks": max(len(data_rows) - len(flagged), 0),
                "resolved_errors": resolved_errors,
                "active_errors": len(flagged) - resolved_errors,
            },
            "summary_placeholder_row": summary_placeholder_row,
            "xlsx_path": xlsx_path,
        }
    finally:
        wb.close()


def merge_review_state(
    rows: list[dict[str, Any]],
    review_state_by_key: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    """Apply imported/saved review-state to structured normcontrol rows."""
    merged_rows: list[dict[str, Any]] = []
    for row in rows:
        state = review_state_by_key.get(str(row.get("row_key") or ""), {})
        updated = dict(row)
        updated["resolved"] = bool(state.get("resolved", row.get("resolved", False)))
        updated["resolved_at"] = str(state.get("resolved_at", row.get("resolved_at", "")) or "")
        updated["comment"] = str(state.get("comment", row.get("comment", "")) or "")
        merged_rows.append(updated)
    return merged_rows


def _ensure_review_columns(ws: Any) -> dict[str, int]:
    """Make sure workbook has all review/state columns and return key->column map."""
    header_to_col: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = ws.cell(row=NORMCONTROL_HEADER_ROW, column=col_idx).value
        if value:
            header_to_col[str(value).strip()] = col_idx
    next_col = ws.max_column + 1
    for col_def in NORMCONTROL_COLUMNS:
        if col_def.header_label in header_to_col:
            continue
        cell = ws.cell(row=NORMCONTROL_HEADER_ROW, column=next_col)
        cell.value = col_def.header_label
        cell.font = Font(bold=True)
        cell.alignment = _ALIGN_CENTER_WRAP
        cell.border = _BORDER_THIN
        header_fill = _HEADER_FILL_REVIEW if col_def.key in _REVIEW_COLUMN_KEYS else _HEADER_FILL_MAIN
        cell.fill = PatternFill(fill_type="solid", fgColor=header_fill)
        dim = ws.column_dimensions[get_column_letter(next_col)]
        dim.width = col_def.width
        dim.hidden = col_def.hidden
        header_to_col[col_def.header_label] = next_col
        next_col += 1
    return {col_def.key: header_to_col[col_def.header_label] for col_def in NORMCONTROL_COLUMNS}


def save_review_state_to_workbook(
    rows: list[dict[str, Any]],
    xlsx_path: str,
    *,
    sanitize_illegal_chars: bool = True,
) -> str:
    """Persist current review-state and service columns into the xlsx workbook."""
    if not xlsx_path:
        raise ValueError("xlsx_path is required")
    wb = openpyxl.load_workbook(xlsx_path)
    try:
        ws = wb[NORMCONTROL_SHEET_TITLE] if NORMCONTROL_SHEET_TITLE in wb.sheetnames else wb.active
        key_to_row_idx: dict[str, int] = {}
        key_col_idx: int | None = None
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row=NORMCONTROL_HEADER_ROW, column=col_idx).value
            if str(value or "").strip() == "row_key":
                key_col_idx = col_idx
                break
        if key_col_idx is not None:
            for row_idx in range(2, ws.max_row + 1):
                row_key = str(ws.cell(row=row_idx, column=key_col_idx).value or "").strip()
                if row_key:
                    key_to_row_idx[row_key] = row_idx
        col_idx_by_key = _ensure_review_columns(ws)
        for row in rows:
            row_key = str(row.get("row_key") or "").strip()
            row_idx = key_to_row_idx.get(row_key)
            if row_idx is None:
                continue
            excel_row = _row_to_excel_values(row)
            for key, col_idx in col_idx_by_key.items():
                raw = excel_row.get(key, "")
                ws.cell(row=row_idx, column=col_idx).value = (
                    utils_excel.sanitize_cell_value_for_openpyxl(
                        raw,
                        enabled=sanitize_illegal_chars,
                    )
                )
            _apply_excel_row_style(ws, row_idx, bool(row.get("resolved")))
        last_col = get_column_letter(len(NORMCONTROL_COLUMNS))
        last_row = max(ws.max_row, 1)
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"
        wb.save(xlsx_path)
        return xlsx_path
    finally:
        wb.close()
